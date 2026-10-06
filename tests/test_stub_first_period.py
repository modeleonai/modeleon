"""A window whose first period is short — a valuation on 1 May.

A quarter or year window may start on the 1st of any month: its first
period then runs from that day to the end of its quarter or year (a stub)
and keeps its label; every later period is whole. The engine never scales
an amount — an annual rate times ``mo.time.months / 12`` is the period's
amount. A day inside a month is refused: a period's months are whole.
"""

from __future__ import annotations

import shutil
import subprocess
import warnings
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pytest

import modeleon as mo


def _model(start=date(2026, 5, 1), grain='year', periods=3):
    m = mo.Model('m', default_grain=grain, default_periods=periods)
    with m:
        m.inputs = mo.MultiVariable('Inputs')
        m.inputs.start = mo.Variable(start, display_name='Start')
        m.default_start = m.inputs.start
        m.t = mo.MultiVariable('Time')
        m.t.start = mo.time.start
        m.t.end = mo.time.end
        m.t.days = mo.time.days
        m.t.months = mo.time.months
        m.t.index = mo.time.index
    return m


def _fields(m):
    return {f: getattr(m.t, f).value for f in ('start', 'end', 'days', 'months', 'index')}


# ══════════════════════════════════════════════════════════════════════
# What the window holds
# ══════════════════════════════════════════════════════════════════════


def test_a_year_window_from_1_may_has_an_eight_month_first_period():
    got = _fields(_model())
    assert got['start'] == [date(2026, 5, 1), date(2027, 1, 1), date(2028, 1, 1)]
    assert got['end'] == [date(2026, 12, 31), date(2027, 12, 31), date(2028, 12, 31)]
    assert got['days'] == [245, 365, 366]
    assert got['months'] == [8, 12, 12]
    assert got['index'] == [1, 2, 3]


def test_a_quarter_window_from_1_may_has_a_two_month_first_quarter():
    got = _fields(_model(grain='quarter'))
    assert got['start'] == [date(2026, 5, 1), date(2026, 7, 1), date(2026, 10, 1)]
    assert got['days'] == [61, 92, 92]
    assert got['months'] == [2, 3, 3]


def test_the_short_period_keeps_its_calendar_label():
    m = _model()
    assert m.default_start == '2026'
    assert m.t.days.time.start == '2026' and m.t.days.time.first_day == date(2026, 5, 1)


def test_a_window_starting_on_a_boundary_is_whole():
    m = _model(start=date(2026, 1, 1))
    assert _fields(m)['days'] == [365, 365, 366]
    assert m.t.days.time.first_day is None


def test_an_annual_amount_scaled_by_the_months_is_the_periods_amount():
    m = _model()
    with m:
        m.rate = mo.Variable([120.0, 120.0, 120.0])
        m.flow = m.rate * mo.time.months / 12
    assert m.flow.value == [80.0, 120.0, 120.0]


def test_a_day_inside_a_month_is_refused():
    with pytest.raises(ValueError) as refused:
        _model(start=date(2026, 5, 15))
    message = str(refused.value)
    assert 'starts on the 1st of a month' in message
    assert 'keep the exact day as an input' in message


def test_a_month_window_has_no_short_first_period():
    got = _fields(_model(grain='month', start=date(2026, 5, 1)))
    assert got['days'] == [31, 30, 31]
    with pytest.raises(ValueError, match='starts on the 1st of a month'):
        _model(grain='month', start=date(2026, 5, 15))


def test_a_typed_start_date_keeps_its_day():
    m = mo.Model('m', default_grain='year', default_periods=2, default_start=date(2026, 5, 1))
    with m:
        m.days = mo.time.days
    assert m.days.value == [245, 365]


def test_a_new_grain_rereads_the_start():
    m = mo.Model('m', default_grain='year', default_periods=4)
    with m:
        m.inputs = mo.MultiVariable('Inputs')
        m.inputs.start = mo.Variable(date(2026, 5, 1))
        m.default_start = m.inputs.start
    m.default_grain = 'quarter'
    assert m.default_start == '2026-Q2'
    with m:
        m.days = mo.time.days
    assert m.days.value == [61, 92, 92, 90]


# ══════════════════════════════════════════════════════════════════════
# Lines that meet across windows
# ══════════════════════════════════════════════════════════════════════


def test_a_whole_year_never_meets_the_short_one_of_the_same_label():
    stub = _model()
    whole = mo.Model('w', default_grain='year', default_start='2026', default_periods=3)
    with whole:
        whole.rate = mo.Variable([1.0, 2.0, 3.0])
    with pytest.raises(ValueError, match='do not line up'):
        stub.t.days + whole.rate


def test_a_short_period_may_open_a_combined_window():
    stub = _model()
    later = mo.Model('l', default_grain='year', default_start='2027', default_periods=2)
    with later:
        later.rate = mo.Variable([10.0, 20.0], extend=mo.zero())
    total = stub.t.days + later.rate
    assert total.value == [245, 375, 386]
    assert total.time.first_day == date(2026, 5, 1)


def test_a_mean_over_a_short_quarter_weighs_its_days():
    m = _model(grain='quarter', periods=3)
    with m:
        m.rate = mo.Variable([0.10, 0.20, 0.30], regrain=mo.up('mean'))
    year = m.rate.at('year')
    assert year.value == pytest.approx([(0.10 * 61 + 0.20 * 92 + 0.30 * 92) / 245])


def test_time_read_for_a_whole_window_does_not_join_the_short_one():
    whole = mo.Model('w', default_grain='year', default_start='2026', default_periods=3)
    with whole:
        days = mo.time.days * 1
    stub = _model()
    with pytest.raises(ValueError, match='from 2026-05-01'):
        with stub:
            stub.extra = mo.MultiVariable('Extra')
            stub.extra.d = days


def test_a_tracked_whole_year_never_meets_the_short_one():
    m = mo.Model('m', default_grain='year', default_periods=3,
                 default_start=date(2026, 5, 1), tracks=mo.Tracks('plan', 'actual'))
    with pytest.raises(ValueError, match='do not line up'):
        with m:
            m.budget = mo.MultiVariable('Budget', default_grain='year',
                                        default_start='2026', default_periods=3)
            m.budget.rev = mo.Variable(plan=[1200., 1300., 1400.], actual=[1250., None, None])
            m.c = mo.MultiVariable('C')
            m.c.other = mo.Variable([100., 100., 100.])
            m.c.total = m.budget.rev + m.c.other


def test_the_refusal_names_both_lines_and_their_first_days():
    stub = _model()
    whole = mo.Model('w', default_grain='year', default_start='2026', default_periods=3)
    with whole:
        whole.rate = mo.Variable([1.0, 2.0, 3.0])
    with pytest.raises(ValueError) as refused:
        stub.t.days + whole.rate
    message = str(refused.value)
    assert "starts its first period, 2026, on 2026-05-01" in message
    assert "starts with the whole period 2026" in message


# ══════════════════════════════════════════════════════════════════════
# A coarser view of a short first period
# ══════════════════════════════════════════════════════════════════════


def _quarterly(periods):
    m = mo.Model('m', default_grain='quarter', default_periods=periods,
                 default_start=date(2026, 5, 1))
    with m:
        m.c = mo.MultiVariable('C')
        m.c.rev = mo.Variable([100., 110, 120, 130, 140, 150, 160][:periods],
                              regrain=mo.up('sum'))
        m.c.profit = mo.Variable([10., 12, 14, 16, 18, 20, 22][:periods], regrain=mo.up('sum'))
        m.c.margin = mo.Variable(m.c.profit / m.c.rev, regrain=mo.ratio('profit', 'rev'))
        m.c.target = mo.Variable([0.1] * periods, regrain=mo.up('mean'))
        m.c.gap = m.c.margin - m.c.target
        m.c.debt = mo.Variable([1000., 900, 800, 700, 600, 500, 400][:periods],
                               regrain=mo.up('last'))
        m.c.daily = m.c.debt / mo.time.days
    return m


def test_a_year_view_of_a_short_first_quarter_keeps_every_line_on_one_calendar():
    y = _quarterly(7).at('year')
    assert y.c.gap.value == pytest.approx([36 / 330 - 0.1, 76 / 580 - 0.1])
    assert y.c.daily.value == pytest.approx([800 / 245, 400 / 365])


def test_a_year_view_of_a_window_ending_inside_a_year():
    y = _quarterly(6).at('year')
    assert y.c.daily.value == pytest.approx([800 / 245, 500 / 273])


def test_a_monthly_section_rolled_up_meets_the_short_year_it_starts_with():
    m = _model(periods=3)
    with m:
        m.ops = mo.MultiVariable('Ops', default_grain='month', default_start=date(2026, 5, 1),
                                 default_periods=32)
        m.ops.rev = mo.Variable([10.0] * 32, regrain=mo.up('sum'))
        m.ops_y = m.ops.rev.at('year')
        m.fixed = mo.Variable([5.0, 5.0, 5.0])
        m.total = m.fixed + m.ops_y
    assert m.total.value == [85.0, 125.0, 125.0]


# ══════════════════════════════════════════════════════════════════════
# The workbook
# ══════════════════════════════════════════════════════════════════════


def _book(model, tmp_path):
    out = tmp_path / 'book.xlsx'
    with warnings.catch_warnings():
        warnings.simplefilter('ignore')
        model.to_excel(str(out))
    return openpyxl.load_workbook(str(out))


def _row(ws, label):
    for row in ws.iter_rows():
        if row[0].value == label:
            return [c.value for c in row[1:] if c.value is not None]
    raise KeyError(label)


def test_the_header_ends_each_period_with_its_year_whatever_month_it_starts(tmp_path):
    m = _model()
    m.default_header = m.t
    ws = _book(m, tmp_path)['Time']
    ends = _row(ws, 'End')
    assert all('MOD(MONTH(' in f for f in ends), ends
    months = _row(ws, 'Months')
    assert all(f.startswith('=(YEAR(') for f in months), months


def test_a_start_cell_on_a_boundary_is_spelled_so_it_may_move(tmp_path):
    m = _model(start=date(2026, 1, 1))
    m.default_header = m.t
    ends = _row(_book(m, tmp_path)['Time'], 'End')
    assert all('MOD(MONTH(' in f for f in ends), ends


def test_a_window_started_at_a_label_keeps_its_formulas(tmp_path):
    m = mo.Model('m', default_grain='year', default_start='2026', default_periods=3)
    with m:
        m.t = mo.MultiVariable('Time')
        m.t.end = mo.time.end
    m.default_header = m.t
    ends = _row(_book(m, tmp_path)['Time'], 'End')
    assert not any('MOD(' in f for f in ends), ends


def test_the_months_row_is_counted_from_its_dates_whatever_the_start(tmp_path):
    # A window started at a label has no date cell, but its header has the
    # period's first and last day: months come from them, as days do.
    m = mo.Model('m', default_grain='year', default_start='2026', default_periods=3)
    with m:
        m.t = mo.MultiVariable('Time')
        m.t.start = mo.time.start
        m.t.end = mo.time.end
        m.t.months = mo.time.months
    m.default_header = m.t
    months = _row(_book(m, tmp_path)['Time'], 'Months')
    assert months[0] == '=(YEAR(B3 + 1) - YEAR(B2)) * 12 + MONTH(B3 + 1) - MONTH(B2)', months
    assert not any(f in ('=12', 12) for f in months), months


def test_without_a_start_row_the_months_run_from_the_previous_end(tmp_path):
    m = mo.Model('m', default_grain='year', default_start='2026', default_periods=3)
    with m:
        m.t = mo.MultiVariable('Time')
        m.t.end = mo.time.end
        m.t.months = mo.time.months
    m.default_header = m.t
    months = _row(_book(m, tmp_path)['Time'], 'Months')
    assert months[1] == '=(YEAR(C2 + 1) - YEAR(B2 + 1)) * 12 + MONTH(C2 + 1) - MONTH(B2 + 1)', months


def test_a_whole_year_line_is_never_written_under_the_short_column(tmp_path):
    m = _model(periods=3)
    with m:
        m.ops = mo.MultiVariable('Ops')
        m.ops.capex = mo.Variable([120.0, 120.0, 120.0], start='2026', grain='year')
    with pytest.raises(ValueError, match=r'starts 2026-05-01 \(inside 2026\)'):
        _book(m, tmp_path)


def test_a_count_of_days_shows_as_a_number(tmp_path):
    m = _model()
    m.default_header = m.t
    ws = _book(m, tmp_path)['Time']
    for row in ws.iter_rows():
        if row[0].value in ('Days', 'Months', 'Index'):
            assert {c.number_format for c in row[1:] if c.value is not None} == {'0'}, row[0].value


def test_a_typed_start_spells_its_offset_as_a_number(tmp_path):
    m = mo.Model('m', default_grain='year', default_periods=2, default_start=date(2026, 5, 1))
    with m:
        m.c = mo.MultiVariable('C')
        m.c.days = mo.time.days * 1
    days = _row(_book(m, tmp_path)['C'], 'Days')
    assert days[0] == '=(EOMONTH(DATE(2026,5,1),7)-DATE(2026,5,1)+1) * 1', days


def _soffice():
    found = shutil.which('soffice') or shutil.which('libreoffice')
    if found:
        return found
    mac = Path('/Applications/LibreOffice.app/Contents/MacOS/soffice')
    return str(mac) if mac.exists() else None


def _recalculate(path: Path, tmp_path: Path):
    soffice = _soffice()
    if soffice is None:
        pytest.skip('LibreOffice is not installed')
    outdir = tmp_path / 'calc'
    profile = (tmp_path / 'lo-profile').as_uri()
    subprocess.run(
        [soffice, f'-env:UserInstallation={profile}', '--headless',
         '--convert-to', 'xlsx', '--outdir', str(outdir), str(path)],
        check=True, capture_output=True, timeout=120,
    )
    return openpyxl.load_workbook(str(outdir / path.name), data_only=True)


def _serial(v):
    if isinstance(v, datetime):
        v = v.date()
    if isinstance(v, date):
        return (v - date(1899, 12, 30)).days
    return v


def _flows(m):
    """A line built on time, read off the header — what a model computes."""
    with m:
        m.c = mo.MultiVariable('Flows')
        m.c.rate = mo.Variable([1200.0] * len(m.t.days.value), display_name='rate')
        m.c.flow = (m.c.rate * mo.time.months / 12).set_display_name('flow')
        m.c.accrual = (m.c.rate * 0.05 * mo.time.days / 365).set_display_name('accrual')
    return m


@pytest.mark.slow
@pytest.mark.parametrize('grain', ['quarter', 'year'])
@pytest.mark.parametrize('header', [True, False], ids=['header', 'inline'])
@pytest.mark.parametrize('typed', [False, True], ids=['start-cell', 'typed-date'])
def test_the_workbook_recalculates_to_the_engine(tmp_path, grain, header, typed):
    if typed:
        m = mo.Model('m', default_grain=grain, default_periods=4,
                     default_start=date(2026, 5, 1))
        with m:
            m.t = mo.MultiVariable('Time')
            for f in ('start', 'end', 'days', 'months', 'index'):
                setattr(m.t, f, getattr(mo.time, f))
    else:
        m = _model(grain=grain, periods=4)
    if header:
        m.default_header = m.t
    m = _flows(m)
    out = tmp_path / 'book.xlsx'
    with warnings.catch_warnings():
        warnings.simplefilter('ignore')
        m.to_excel(str(out))
    wb = _recalculate(out, tmp_path)
    for sheet, mv in (('Time', m.t), ('Flows', m.c)):
        for name, var in mv._components.items():
            got = [_serial(v) for v in _row(wb[sheet], var.display_name or name)]
            want = [_serial(v) for v in var.value]
            assert got[-len(want):] == pytest.approx(want), (grain, header, typed, name)


@pytest.mark.slow
@pytest.mark.parametrize('grain, start', [('year', '2026'), ('quarter', '2026-Q1'),
                                          ('month', '2026-01')])
def test_a_label_window_recalculates_its_months(tmp_path, grain, start):
    m = mo.Model('m', default_grain=grain, default_start=start, default_periods=4)
    with m:
        m.t = mo.MultiVariable('Time')
        m.t.end = mo.time.end
        m.t.months = mo.time.months
    m.default_header = m.t
    out = tmp_path / 'book.xlsx'
    with warnings.catch_warnings():
        warnings.simplefilter('ignore')
        m.to_excel(str(out))
    wb = _recalculate(out, tmp_path)
    assert _row(wb['Time'], 'Months')[-4:] == m.t.months.value


@pytest.mark.slow
def test_year_totals_of_a_short_first_quarter_recalculate_to_the_engine(tmp_path):
    m = _quarterly(7)
    with m:
        m.inputs = mo.MultiVariable('Inputs')
    m.default_excel_view = mo.ExcelView(timeline={'totals': ['year']})
    out = tmp_path / 'book.xlsx'
    with warnings.catch_warnings():
        warnings.simplefilter('ignore')
        m.to_excel(str(out))
    wb = _recalculate(out, tmp_path)
    year = m.at('year')
    for name in ('rev', 'gap', 'target', 'daily'):
        label = getattr(m.c, name).display_name
        row = next(r for r in wb['C'].iter_rows() if r[0].value == label)
        header = next(r for r in wb['C'].iter_rows() if any(
            isinstance(c.value, str) and c.value.startswith('Total') for c in r))
        totals = [row[i].value for i, c in enumerate(header) if
                  isinstance(c.value, str) and c.value.startswith('Total')]
        assert totals == pytest.approx(getattr(year.c, name).value), name


@pytest.mark.slow
@pytest.mark.parametrize('moved_to', [date(2026, 5, 1), date(2026, 10, 1), date(2026, 1, 1)])
def test_a_start_moved_in_the_workbook_moves_the_periods(tmp_path, moved_to):
    built = _flows(_model(start=date(2026, 1, 1) if moved_to != date(2026, 1, 1)
                          else date(2026, 5, 1), periods=3))
    built.default_header = built.t
    out = tmp_path / 'book.xlsx'
    with warnings.catch_warnings():
        warnings.simplefilter('ignore')
        built.to_excel(str(out))
    wb = openpyxl.load_workbook(str(out))
    ws = wb['Inputs']
    # The input row, not the header mirror above it.
    cell = next(r[1] for r in ws.iter_rows()
                if r[0].value == 'Start' and isinstance(r[1].value, datetime))
    cell.value = datetime(moved_to.year, moved_to.month, moved_to.day)
    wb.save(str(out))
    calc = _recalculate(out, tmp_path)
    want = _flows(_model(start=moved_to, periods=3))
    for name in ('days', 'months'):
        got = [_serial(v) for v in _row(calc['Time'], name.capitalize())][-3:]
        assert got == getattr(want.t, name).value, name
    assert _row(calc['Flows'], 'flow')[-3:] == pytest.approx(want.c.flow.value)
