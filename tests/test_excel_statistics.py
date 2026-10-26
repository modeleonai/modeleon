# SPDX-License-Identifier: Apache-2.0
"""Statistics in the workbook.

A row is its cells' range (``Data!B3:F3``); a list of scalar lines whose
cells stand one under another is a vertical range — one field down the
records of a table, ``Comps!C2:C6`` — and ``RANK`` of each record among
its peers is one call per record against that range. ``MEDIAN`` and the
spreads take any number of arguments, two rows and more one call per
period over each period's cells; the others one range — a list of yearly
lines one range per period, the run of their cells in its column — and
cells that form no run carry the computed value (``check_compat`` lists
them).
The names Excel 2010 added are stored with the ``_xlfn.`` prefix
(``=_xlfn.PERCENTILE.INC(...)``). Every function recalculates in
LibreOffice to the engine's numbers, except where LibreOffice departs
from Excel, pinned by cell as an expected failure.
"""

import datetime as dt
import shutil
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook

import modeleon as mo
from modeleon.compile import check_compat
from modeleon.functions.statistics import (
    LARGE,
    MEDIAN,
    PERCENTILE,
    PERCENTILE_EXC,
    PERCENTILE_INC,
    QUARTILE,
    QUARTILE_EXC,
    QUARTILE_INC,
    RANK,
    RANK_AVG,
    RANK_EQ,
    SMALL,
    STDEV,
    STDEV_P,
    STDEV_S,
    STDEVP,
    VAR,
    VAR_P,
    VAR_S,
    VARP,
)


def _rows(ws):
    out = {}
    for row in ws.iter_rows():
        label = row[0].value
        if label:
            out[str(label).split(" (")[0]] = [c.value for c in row[1:]]
    return out


def _book(model, tmp_path):
    path = tmp_path / "book.xlsx"
    model.to_excel(path)
    return load_workbook(path)


class Comp(mo.MultiVariableClass):
    def compute(self, multiple):
        self.multiple = mo.Variable(multiple, display_name="Multiple", unit="x")


def _model():
    m = mo.Model("m", display_name="M", default_grain="month", default_start="2026-01",
                 default_periods=5)
    with m:
        m.inputs = mo.MultiVariable("Inputs", excel_props={"tab": True})
        m.inputs.k = mo.Variable(0.3, display_name="K")
        m.inputs.quart = mo.Variable(3, display_name="Quart")
        m.data = mo.MultiVariable("Data", excel_props={"tab": True})
        m.data.v = mo.Variable([10.0, 20.0, 30.0, 40.0, 50.0], display_name="V", unit="USD")
        m.data.u = mo.Variable([40.0, 10.0, 30.0, 20.0, 50.0], display_name="U", unit="USD")
        m.data.ties = mo.Variable([5, 3, 5, 1, 3], display_name="Ties")
        m.comps = mo.MultiVariable("Comps", excel_props={"tab": True, "orient": "records"})
        with m.comps as c:
            for key, x in (("a", 8.0), ("b", 12.0), ("c", 10.0), ("d", 9.0), ("e", 10.0)):
                setattr(c, key, Comp(multiple=x, display_name=key.upper()))
        m.s = mo.MultiVariable("Summary", excel_props={"tab": True})
    return m


def _peers(m):
    c = m.comps
    return [c.a.multiple, c.b.multiple, c.c.multiple, c.d.multiple, c.e.multiple]


# ══════════════════════════════════════════════════════════════════════
# Rows are ranges, the names as the file stores them
# ══════════════════════════════════════════════════════════════════════


class TestRowsAreRanges:
    def test_every_function_over_rows(self, tmp_path):
        m = _model()
        d, i = m.data, m.inputs
        with m.s as s:
            s.median = MEDIAN(d.v).set_display_name("median")
            s.stdev_s = STDEV_S(d.v).set_display_name("stdev_s")
            s.stdev_p = STDEV_P(d.v).set_display_name("stdev_p")
            s.var_s = VAR_S(d.v).set_display_name("var_s")
            s.var_p = VAR_P(d.v).set_display_name("var_p")
            s.stdev = STDEV(d.v).set_display_name("stdev")
            s.stdevp = STDEVP(d.v).set_display_name("stdevp")
            s.var = VAR(d.v).set_display_name("var")
            s.varp = VARP(d.v).set_display_name("varp")
            s.p_inc = PERCENTILE_INC(d.v, i.k).set_display_name("p_inc")
            s.p_exc = PERCENTILE_EXC(d.v, 0.4).set_display_name("p_exc")
            s.q_inc = QUARTILE_INC(d.v, i.quart).set_display_name("q_inc")
            s.q_exc = QUARTILE_EXC(d.v, 1).set_display_name("q_exc")
            s.perc = PERCENTILE(d.v, 0.3).set_display_name("perc")
            s.quart = QUARTILE(d.v, 1).set_display_name("quart")
            s.large = LARGE(d.v, 2).set_display_name("large")
            s.small = SMALL(d.v, 2).set_display_name("small")
            s.rank_eq = RANK_EQ(3, d.ties).set_display_name("rank_eq")
            s.rank_avg = RANK_AVG(3, d.ties, 1).set_display_name("rank_avg")
            s.rank = RANK(i.quart, d.ties).set_display_name("rank")
        rows = _rows(_book(m, tmp_path)["Summary"])
        assert rows["median"][0] == "=MEDIAN(Data!B2:F2)"
        assert rows["stdev_s"][0] == "=_xlfn.STDEV.S(Data!B2:F2)"
        assert rows["stdev_p"][0] == "=_xlfn.STDEV.P(Data!B2:F2)"
        assert rows["var_s"][0] == "=_xlfn.VAR.S(Data!B2:F2)"
        assert rows["var_p"][0] == "=_xlfn.VAR.P(Data!B2:F2)"
        assert rows["stdev"][0] == "=STDEV(Data!B2:F2)"
        assert rows["stdevp"][0] == "=STDEVP(Data!B2:F2)"
        assert rows["var"][0] == "=VAR(Data!B2:F2)"
        assert rows["varp"][0] == "=VARP(Data!B2:F2)"
        assert rows["p_inc"][0] == "=_xlfn.PERCENTILE.INC(Data!B2:F2, Inputs!B1)"
        assert rows["p_exc"][0] == "=_xlfn.PERCENTILE.EXC(Data!B2:F2, 0.4)"
        assert rows["q_inc"][0] == "=_xlfn.QUARTILE.INC(Data!B2:F2, Inputs!B2)"
        assert rows["q_exc"][0] == "=_xlfn.QUARTILE.EXC(Data!B2:F2, 1)"
        assert rows["perc"][0] == "=PERCENTILE(Data!B2:F2, 0.3)"
        assert rows["quart"][0] == "=QUARTILE(Data!B2:F2, 1)"
        assert rows["large"][0] == "=LARGE(Data!B2:F2, 2)"
        assert rows["small"][0] == "=SMALL(Data!B2:F2, 2)"
        assert rows["rank_eq"][0] == "=_xlfn.RANK.EQ(3, Data!B4:F4)"
        assert rows["rank_avg"][0] == "=_xlfn.RANK.AVG(3, Data!B4:F4, 1)"
        assert rows["rank"][0] == "=RANK(Inputs!B2, Data!B4:F4)"
        # The values the cells cache are the engine's.
        assert (m.s.median.value, m.s.var_s.value, m.s.p_inc.value) == (30.0, 250.0, 22.0)
        assert (m.s.q_inc.value, m.s.large.value, m.s.rank_avg.value) == (40.0, 40.0, 2.5)

    def test_rows_are_one_call_per_period_cells_and_typed_values_arguments(self, tmp_path):
        m = _model()
        d, c = m.data, m.comps
        with m.s as s:
            s.two = MEDIAN(d.v, d.u).set_display_name("two")
            s.typed = STDEV_S(d.v, 5, True, "3").set_display_name("typed")
            s.cells = MEDIAN(c.a.multiple, c.c.multiple, 7.5).set_display_name("cells")
            s.half = (MEDIAN(d.v) / 2).set_display_name("half")
        rows = _rows(_book(m, tmp_path)["Summary"])
        # The scalar lines take column B; the periods move to C.
        assert rows["two"][1:] == [f"=MEDIAN(Data!{col}2, Data!{col}3)" for col in "BCDEF"]
        assert rows["typed"][1:] == [
            f'=_xlfn.STDEV.S(Data!{col}2, 5, TRUE, "3")' for col in "BCDEF"]
        assert rows["cells"][0] == "=MEDIAN(Comps!B2, Comps!B4, 7.5)"
        assert rows["half"][0] == "=MEDIAN(Data!B2:F2) / 2"
        assert m.s.two.value == [25.0, 15.0, 30.0, 30.0, 50.0]

    def test_on_the_rows_own_sheet(self, tmp_path):
        m = _model()
        with m.data as d:
            d.mid = MEDIAN(d.v).set_display_name("mid")
            d.top = LARGE(d.v, 1).set_display_name("top")
        rows = _rows(_book(m, tmp_path)["Data"])
        # The scalar lines take column B; the periods move to C.
        assert rows["mid"][0] == "=MEDIAN(C2:G2)"
        assert rows["top"][0] == "=LARGE(C2:G2, 1)"


# ══════════════════════════════════════════════════════════════════════
# One field down the records is a vertical range
# ══════════════════════════════════════════════════════════════════════


class TestRecordsAreVerticalRanges:
    def test_a_field_down_the_records(self, tmp_path):
        m = _model()
        peers = _peers(m)
        with m.s as s:
            s.median = MEDIAN(peers).set_display_name("median")
            s.q1 = PERCENTILE_INC(peers, m.inputs.k).set_display_name("q1")
            s.q3 = QUARTILE_EXC(peers, 3).set_display_name("q3")
            s.sd = STDEV_S(peers).set_display_name("sd")
            s.top = LARGE(peers, 1).set_display_name("top")
            s.rank_a = RANK_AVG(m.comps.c.multiple, peers).set_display_name("rank_a")
        rows = _rows(_book(m, tmp_path)["Summary"])
        assert rows["median"][0] == "=MEDIAN(Comps!B2:B6)"
        assert rows["q1"][0] == "=_xlfn.PERCENTILE.INC(Comps!B2:B6, Inputs!B1)"
        assert rows["q3"][0] == "=_xlfn.QUARTILE.EXC(Comps!B2:B6, 3)"
        assert rows["sd"][0] == "=_xlfn.STDEV.S(Comps!B2:B6)"
        assert rows["top"][0] == "=LARGE(Comps!B2:B6, 1)"
        assert rows["rank_a"][0] == "=_xlfn.RANK.AVG(Comps!B4, Comps!B2:B6)"
        assert (m.s.median.value, m.s.top.value, m.s.rank_a.value) == (10.0, 12.0, 2.5)
        assert str(m.s.median._unit) == "x" and m.s.rank_a._unit is None
        assert check_compat(m, "excel") == []

    def test_each_record_ranked_among_its_peers(self, tmp_path):
        m = _model()
        peers = _peers(m)
        with m.comps as c:
            for rec in (c.a, c.b, c.c, c.d, c.e):
                rec.rank = RANK_EQ(rec.multiple, peers).set_display_name("Rank")
        ws = _book(m, tmp_path)["Comps"]
        header = [cell.value for cell in ws[1]]
        at = header.index("Rank")
        ranks = {row[0].value: row[at].value for row in ws.iter_rows(min_row=2)}
        assert ranks == {
            "A": "=_xlfn.RANK.EQ(B2, B2:B6)", "B": "=_xlfn.RANK.EQ(B3, B2:B6)",
            "C": "=_xlfn.RANK.EQ(B4, B2:B6)", "D": "=_xlfn.RANK.EQ(B5, B2:B6)",
            "E": "=_xlfn.RANK.EQ(B6, B2:B6)",
        }
        assert [r.rank.value for r in (c.a, c.b, c.c, c.d, c.e)] == [5, 1, 2, 4, 2]
        assert check_compat(m, "excel") == []


# ══════════════════════════════════════════════════════════════════════
# A series parameter is one call per period over the same range
# ══════════════════════════════════════════════════════════════════════


class TestOneCallPerPeriod:
    def test_a_series_k_and_each_period_ranked_among_its_row(self, tmp_path):
        m = _model()
        with m.data as d:
            d.ks = mo.Variable([0.0, 0.25, 0.5, 0.75, 1.0], display_name="Ks")
            d.pct = PERCENTILE_INC(d.v, d.ks).set_display_name("Pct")
            d.rank = RANK_EQ(d.u, d.u).set_display_name("Rank")
        rows = _rows(_book(m, tmp_path)["Data"])
        assert rows["Pct"][:5] == [f"=_xlfn.PERCENTILE.INC(B2:F2, {col}5)" for col in "BCDEF"]
        assert rows["Rank"][:5] == [f"=_xlfn.RANK.EQ({col}3, B3:F3)" for col in "BCDEF"]
        assert d.pct.value == [10.0, 20.0, 30.0, 40.0, 50.0]
        assert d.rank.value == [2, 5, 3, 4, 1]


# ══════════════════════════════════════════════════════════════════════
# Cells that form no run carry the value
# ══════════════════════════════════════════════════════════════════════


class TestNoRunShipsTheValue:
    def test_one_range_over_a_gap_two_sheets_or_a_literal(self, tmp_path):
        m = _model()
        c = m.comps
        with m.s as s:
            s.gap = PERCENTILE_INC([c.a.multiple, c.c.multiple], 0.5).set_display_name("gap")
            s.sheets = LARGE([c.a.multiple, m.inputs.k], 1).set_display_name("sheets")
            s.literal = RANK_EQ(8.0, [c.a.multiple, 7.0]).set_display_name("literal")
            s.backwards = SMALL([c.c.multiple, c.b.multiple, c.a.multiple],
                                1).set_display_name("backwards")
        rows = _rows(_book(m, tmp_path)["Summary"])
        assert rows["gap"][0] == 9 and rows["sheets"][0] == 8
        assert rows["literal"][0] == 1 and rows["backwards"][0] == 8
        issues = check_compat(m, "excel")
        assert [(i.kind, i.qpath) for i in issues] == [
            ("value", m.s.gap.id), ("value", m.s.sheets.id),
            ("value", m.s.literal.id), ("value", m.s.backwards.id)]

    def test_a_reduction_over_a_gap_carries_its_value_over_cells_it_lists_them(
            self, tmp_path):
        # A list of lines is ONE range: cells that form no run carry the
        # value; the lines as separate arguments are a live comma list.
        m = _model()
        c = m.comps
        with m.s as s:
            s.gap = MEDIAN([c.a.multiple, c.c.multiple]).set_display_name("gap")
            s.args = MEDIAN(c.a.multiple, c.c.multiple).set_display_name("args")
        rows = _rows(_book(m, tmp_path)["Summary"])
        assert rows["gap"][0] == 9
        assert rows["args"][0] == "=MEDIAN(Comps!B2, Comps!B4)"
        assert [(i.kind, i.qpath) for i in check_compat(m, "excel")] == [
            ("value", m.s.gap.id)]

    def test_a_row_split_by_a_totals_column(self, tmp_path):
        m = mo.Model("m", display_name="M", default_grain="month", default_start="2026-01",
                     default_periods=6,
                     default_excel_view=mo.ExcelView(timeline={"totals": ["quarter"]}))
        with m:
            m.data = mo.MultiVariable("Data", excel_props={"tab": True})
            m.data.v = mo.Variable([10.0, 20.0, 30.0, 40.0, 50.0, 60.0], display_name="V",
                                   regrain=mo.up("sum"))
            m.s = mo.MultiVariable("Summary", excel_props={"tab": True})
            m.s.mid = MEDIAN(m.data.v).set_display_name("mid")
            m.s.p = PERCENTILE_INC(m.data.v, 0.5).set_display_name("p")
        rows = _rows(_book(m, tmp_path)["Summary"])
        # Never =MEDIAN(Data!B3:H3), which would read the Q1 total too.
        assert rows["mid"][0] == 35 and rows["p"][0] == 35
        assert [(i.kind, i.qpath) for i in check_compat(m, "excel")] == [
            ("value", m.s.mid.id), ("value", m.s.p.id)]

    def test_a_row_the_book_has_no_cells_for(self, tmp_path):
        m = _model()
        with m.s as s:
            s.scaled = MEDIAN(m.data.v * 2.0).set_display_name("scaled")
        with pytest.warns(mo.CrossScopeReferenceWarning, match="MEDIAN"):
            rows = _rows(_book(m, tmp_path)["Summary"])
        assert rows["scaled"][0] == 60


# ══════════════════════════════════════════════════════════════════════
# LibreOffice recalculates these functions to the engine's numbers
# ══════════════════════════════════════════════════════════════════════


def _soffice():
    for candidate in ("soffice", "libreoffice"):
        found = shutil.which(candidate)
        if found:
            return found
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


def _recalculated(model, tmp_path):
    """Every sheet after LibreOffice recalculates the written workbook
    (conftest seeds the profile to recalculate on load)."""
    soffice = _soffice()
    if soffice is None:
        pytest.skip("LibreOffice not installed")
    path = tmp_path / "statistics.xlsx"
    model.to_excel(path)
    subprocess.run(
        [soffice, f"-env:UserInstallation={(tmp_path / 'prof').as_uri()}",
         "--headless", "--convert-to", "xlsx", "--outdir",
         str(tmp_path / "calc"), str(path)],
        check=True, capture_output=True,
    )
    return load_workbook(tmp_path / "calc" / "statistics.xlsx", data_only=True)


def _serial(value):
    """A date as the cell holds it, so a recalculated date (read back as
    a datetime) compares with the engine's date."""
    if isinstance(value, dt.datetime):
        return (value - dt.datetime(1899, 12, 30)) / dt.timedelta(days=1)
    if isinstance(value, dt.date):
        return (value - dt.date(1899, 12, 30)).days
    return value


class Peer(mo.MultiVariableClass):
    def compute(self, rev):
        self.rev = mo.Variable(rev, display_name="Revenue", unit="USD")


#: The per-period statistics written on Data, read back period by period.
_DATA_PER_PERIOD = (
    "pct", "rank_row", "rank_ties", "kth",
    "med_pp", "med3_pp", "med_list_pp", "med_typed_pp", "med_mixed_pp", "med_err_pp",
    "med_dates_pp", "sd_pp", "stdevp_pp", "var_pp", "var_list_pp",
    "p_pp", "p_k_pp", "p_exc_pp", "q_pp", "large_pp", "small_pp", "rank_pp",
    "rank_avg_pp", "large_err_pp", "p_none_pp",
)


def _parity_model():
    """Every function and its edges: rows on Data (with the per-period
    calls — over the same range, and over sibling rows one period at a
    time), the peers on Comps (with a rank per record), a records table of
    yearly lines on Peers (with each year's rank per record), one cell per
    call on Summary — a row of them where the call is one per period."""
    m = _model()
    i, d, c = m.inputs, m.data, m.comps
    peers = _peers(m)
    m.peers = mo.MultiVariable("Peers", excel_props={"tab": True, "orient": "records"})
    with m.peers as pp:
        pp.a = Peer(rev=[10.0, 30.0, 25.0, None, 5.0], display_name="PA")
        pp.b = Peer(rev=[20.0, 10.0, 45.0, 7.0, 5.0], display_name="PB")
        pp.c = Peer(rev=[35.0, 40.0, 10.0, 8.0, 5.0], display_name="PC")
        pp.d = Peer(rev=[15.0, 40.0, 30.0, 9.0, 5.0], display_name="PD")
    revs = [r.rev for r in (pp.a, pp.b, pp.c, pp.d)]
    three = peers[:3]                               # 8, 12, 10
    i.blank = mo.Variable(None, display_name="Blank")
    i.text_k = mo.Variable("0.5", display_name="TextK")
    i.text_two = mo.Variable("2", display_name="TextTwo")
    i.err = mo.Variable("#DIV/0!", display_name="Err")
    i.na = mo.Variable("#N/A", display_name="Na")
    i.flag = mo.Variable(True, display_name="Flag")
    d.zero = mo.Variable([1.0, 0.0, 1.0, 1.0, 1.0], display_name="Zero")
    d.err = (d.v / d.zero).set_display_name("Err")           # 10, #DIV/0!, 30, 40, 50
    d.none = mo.Variable(["a", None, "b", None, "c"], display_name="None")
    d.mixed = mo.Variable([3.0, None, "x", 1.0, 4.0], display_name="Mixed")
    d.numtext = mo.Variable(["10", "20", 5.0, None, "x"], display_name="Numtext")
    d.bools = mo.Variable([3.0, True, 1.0, 4.0, False], display_name="Bools")
    d.dates = mo.Variable([dt.date(2026, 1, 31), dt.date(2026, 3, 31), dt.date(2026, 2, 28),
                           dt.date(2026, 2, 27), dt.date(2025, 12, 31)], display_name="Dates")
    d.ks = mo.Variable([0.0, 0.25, 0.5, 0.75, 1.0], display_name="Ks")
    # One call per period over the same range.
    d.pct = PERCENTILE_INC(d.v, d.ks).set_display_name("Pct")
    d.rank_row = RANK_EQ(d.u, d.u).set_display_name("RankRow")
    d.rank_ties = RANK_AVG(d.ties, d.ties, 1).set_display_name("RankTies")
    d.kth = LARGE(d.u, d.ties).set_display_name("Kth")
    # Several rows, one call per period over each period's cells.
    rows = [d.v, d.u, d.ties]                       # rows 2-4: one run per column
    d.med_pp = MEDIAN(d.v, d.u).set_display_name("MedPP")
    d.med3_pp = MEDIAN(d.v, d.u, d.ties).set_display_name("Med3PP")
    d.med_list_pp = MEDIAN(rows).set_display_name("MedListPP")
    d.med_typed_pp = MEDIAN(d.v, [1.0, 50.0, 25.0, 0.0, 100.0]).set_display_name("MedTypedPP")
    d.med_mixed_pp = MEDIAN(d.mixed, d.v).set_display_name("MedMixedPP")
    d.med_err_pp = MEDIAN(d.err, d.u).set_display_name("MedErrPP")
    d.med_dates_pp = MEDIAN(d.dates, dt.date(2026, 2, 1)).set_display_name("MedDatesPP")
    d.sd_pp = STDEV_S(d.v, d.u, 5.0).set_display_name("SdPP")
    d.stdevp_pp = STDEVP(d.v, d.u).set_display_name("StdevpPP")
    d.var_pp = VAR(d.v, d.ties).set_display_name("VarPP")
    d.var_list_pp = VAR_P(rows).set_display_name("VarListPP")
    d.p_pp = PERCENTILE_INC(rows, 0.25).set_display_name("PPP")
    d.p_k_pp = PERCENTILE_INC(rows, d.ks).set_display_name("PKPP")
    d.p_exc_pp = PERCENTILE_EXC(rows, 0.5).set_display_name("PExcPP")
    d.q_pp = QUARTILE_INC(rows, 3).set_display_name("QPP")
    d.large_pp = LARGE(rows, 1).set_display_name("LargePP")
    d.small_pp = SMALL(rows, 2).set_display_name("SmallPP")
    d.rank_pp = RANK_EQ(d.u, rows).set_display_name("RankPP")
    d.rank_avg_pp = RANK_AVG(d.ties, rows, 1).set_display_name("RankAvgPP")
    d.large_err_pp = LARGE([d.zero, d.err], 1).set_display_name("LargeErrPP")     # rows 5-6
    d.p_none_pp = PERCENTILE_INC([d.none, d.mixed], 0.5).set_display_name("PNonePP")  # 7-8
    # Each record among its peers.
    with m.comps:
        for rec in (c.a, c.b, c.c, c.d, c.e):
            rec.rank = RANK_EQ(rec.multiple, peers).set_display_name("Rank")
            rec.rank_avg = RANK_AVG(rec.multiple, peers, 1).set_display_name("RankAvg")
    # Each record's yearly line among its peers', year by year.
    with m.peers:
        for rec in (pp.a, pp.b, pp.c, pp.d):
            rec.rank = RANK_EQ(rec.rev, revs).set_display_name("Rank")
            rec.rank_avg = RANK_AVG(rec.rev, revs, 1).set_display_name("RankAvg")
    with m.s as s:
        # MEDIAN
        s.med = MEDIAN(d.v)
        s.med_even = MEDIAN(d.u, 60.0)
        s.med_two = MEDIAN(d.v, d.ties)
        s.med_mixed = MEDIAN(d.mixed)
        s.med_numtext = MEDIAN(d.numtext)
        s.med_typed = MEDIAN(1, True, 5)
        s.med_none = MEDIAN(1, None, 5)
        s.med_text = MEDIAN(1, "abc")
        s.med_err = MEDIAN(d.err)
        s.med_err_text = MEDIAN(i.err, "x")
        s.med_text_err = MEDIAN("x", i.err)
        s.med_two_errors = MEDIAN(i.err, i.na)
        s.med_dates = MEDIAN(d.dates)
        s.med_peers = MEDIAN(peers)
        s.med_cells = MEDIAN(c.a.multiple, c.c.multiple)
        s.med_empty = MEDIAN(d.none)
        s.med_numtext_typed = MEDIAN(1, "3")
        s.med_bools = MEDIAN(d.bools)
        s.med_flag = MEDIAN(i.flag, 4)
        # The spreads
        s.sd_s = STDEV_S(d.v)
        s.sd_p = STDEV_P(d.v)
        s.var_s = VAR_S(d.v)
        s.var_p = VAR_P(d.v)
        s.stdev = STDEV(d.u)
        s.stdevp = STDEVP(d.u)
        s.var = VAR(d.u)
        s.varp = VARP(d.u)
        s.var_two = VAR_S(d.v, d.ties)
        s.sd_mixed = STDEV_S(d.mixed)
        s.sd_typed = STDEV_S(1, True, 5)
        s.sd_big = STDEV_S(1e15 + 1, 1e15 + 2, 1e15 + 3)
        s.var_small = VAR_S(0.1, 0.2, 0.3)
        s.sd_s_one = STDEV_S(c.a.multiple)
        s.sd_p_one = STDEV_P(c.a.multiple)
        s.var_s_empty = VAR_S(d.none)
        s.var_p_empty = VAR_P(d.none)
        s.sd_err = STDEV_S(d.err)
        s.sd_dates = STDEV_S(d.dates)
        s.var_peers = VAR_P(peers)
        s.sd_numtext_typed = STDEV_S(1, "3", 5)
        s.sd_text_typed = STDEV_S(1, "abc", 5)
        s.sd_bools = STDEV_S(d.bools)
        s.sd_two_errors = STDEV_S(i.err, i.na)
        # Percentiles
        s.p_inc = PERCENTILE_INC(d.v, i.k)
        s.p_inc_0 = PERCENTILE_INC(d.v, 0)
        s.p_inc_1 = PERCENTILE_INC(d.v, 1)
        s.p_inc_unsorted = PERCENTILE_INC(d.u, 0.3)
        s.p_inc_text_k = PERCENTILE_INC(d.v, "0.5")
        s.p_inc_cell_text_k = PERCENTILE_INC(d.v, i.text_k)
        s.p_inc_true = PERCENTILE_INC(d.v, True)
        s.p_inc_blank = PERCENTILE_INC(d.v, i.blank)
        s.p_inc_err = PERCENTILE_INC(d.err, 0.5)
        s.p_inc_kerr = PERCENTILE_INC(d.v, i.err)
        s.p_inc_kerr_err = PERCENTILE_INC(d.err, i.na)
        s.p_inc_dates = PERCENTILE_INC(d.dates, 0.25)
        s.p_inc_peers = PERCENTILE_INC(peers, i.k)
        s.p_inc_one = PERCENTILE_INC(c.a.multiple, 0.3)
        s.p_inc_mixed = PERCENTILE_INC(d.mixed, 0.25)
        s.p_exc = PERCENTILE_EXC(d.v, 0.4)
        s.p_exc_lo = PERCENTILE_EXC(three, 0.25)
        s.p_exc_hi = PERCENTILE_EXC(three, 0.75)
        s.p_exc_near = PERCENTILE_EXC(d.v, 0.167)
        s.perc = PERCENTILE(d.u, 0.3)
        s.p_inc_bad = PERCENTILE_INC(d.v, 1.1)
        s.p_inc_empty = PERCENTILE_INC(d.none, 0.5)
        s.p_inc_err_badk = PERCENTILE_INC(d.err, 2)
        s.p_exc_below = PERCENTILE_EXC(d.v, 0.16)
        s.p_exc_0 = PERCENTILE_EXC(d.v, 0)
        s.p_exc_empty = PERCENTILE_EXC(d.none, 0.5)
        # Quartiles
        s.q_inc_1 = QUARTILE_INC(d.v, 1)
        s.q_inc_2 = QUARTILE_INC(d.u, 2)
        s.q_inc_cell = QUARTILE_INC(d.v, i.quart)
        s.q_inc_15 = QUARTILE_INC(d.v, 1.5)
        s.q_inc_45 = QUARTILE_INC(d.v, 4.5)
        s.q_exc_1 = QUARTILE_EXC(d.v, 1)
        s.q_exc_3 = QUARTILE_EXC(d.u, 3)
        s.q_exc_35 = QUARTILE_EXC(d.v, 3.5)
        s.q_exc_two = QUARTILE_EXC([c.a.multiple, c.b.multiple], 2)
        s.q_inc_err = QUARTILE_INC(d.err, 1)
        s.q_inc_kerr = QUARTILE_INC(d.err, i.na)
        s.quart = QUARTILE(d.u, 1)
        s.q_inc_neg = QUARTILE_INC(d.v, -0.5)
        s.q_inc_5 = QUARTILE_INC(d.v, 5)
        s.q_exc_0 = QUARTILE_EXC(d.v, 0)
        s.q_exc_one = QUARTILE_EXC(c.a.multiple, 1)
        s.q_inc_empty = QUARTILE_INC(d.none, 1)
        # The k-th largest and smallest
        s.large_1 = LARGE(d.u, 1)
        s.large_25 = LARGE(d.u, 2.5)
        s.large_05 = LARGE(d.u, 0.5)
        s.small_25 = SMALL(d.u, 2.5)
        s.small_55 = SMALL(d.u, 5.5)
        s.large_ties = LARGE(d.ties, 2)
        s.large_true = LARGE(d.v, True)
        s.large_err = LARGE(d.err, 1)
        s.large_kerr = LARGE(d.err, i.na)
        s.large_err_badk = LARGE(d.err, 0)
        s.large_dates = LARGE(d.dates, 2)
        s.small_dates = SMALL(d.dates, 1)
        s.large_peers = LARGE(peers, 2)
        s.small_peers = SMALL(peers, 1)
        s.small_mixed = SMALL(d.mixed, 2)
        s.large_0 = LARGE(d.v, 0)
        s.large_55 = LARGE(d.v, 5.5)
        s.small_6 = SMALL(d.v, 6)
        s.small_05 = SMALL(d.v, 0.5)
        s.large_empty = LARGE(d.none, 1)
        s.large_text_k = LARGE(d.v, "2")
        s.large_cell_text_k = LARGE(d.v, i.text_two)
        s.small_cell_text_k = SMALL(d.v, i.text_two)
        s.large_bools = LARGE(d.bools, 4)
        # Ranks
        s.rank_desc = RANK_EQ(3, d.ties)
        s.rank_asc = RANK_EQ(3, d.ties, 1)
        s.rank_avg = RANK_AVG(3, d.ties)
        s.rank_avg_asc = RANK_AVG(5, d.ties, 1)
        s.rank_old = RANK(5, d.ties)
        s.rank_absent = RANK_EQ(4, d.ties)
        s.rank_err = RANK_EQ(1, d.err)
        s.rank_text_num = RANK_EQ("3", d.ties)
        s.rank_true = RANK_EQ(True, d.ties)
        s.rank_blank = RANK_EQ(i.blank, d.ties)
        s.rank_order_true = RANK_EQ(3, d.ties, True)
        s.rank_order_text = RANK_EQ(3, d.ties, "1")
        s.rank_order_cell = RANK_EQ(3, d.ties, i.flag)
        s.rank_num_err = RANK_EQ(i.err, d.v)
        s.rank_order_err = RANK_EQ(10, d.v, i.err)
        s.rank_num_err_ref_err = RANK_EQ(i.na, d.err)
        s.rank_dates = RANK_EQ(dt.date(2026, 2, 28), d.dates)
        s.rank_mixed = RANK_EQ(1, d.mixed)
        s.rank_peer = RANK_AVG(c.c.multiple, peers)
        s.rank_none = RANK_EQ(1, d.none)
        s.rank_bools = RANK_EQ(0, d.bools)
        # A records table of yearly lines: each year's statistic.
        s.peer_med = MEDIAN(revs)
        s.peer_q1 = PERCENTILE_INC(revs, i.k)
        s.peer_sd = STDEV_S(revs)
        s.peer_top = LARGE(revs, 2)
    for name, line in m.s.walk():
        if isinstance(line, mo.Variable):
            line.set_display_name(name)
    return m


#: Where LibreOffice's answer is not the engine's: what LibreOffice 25.2
#: writes instead, pinned by line. Each is where it departs from Excel,
#: whose rules the engine computes: a ``k`` or a quart out of its range
#: and a range without numbers are LibreOffice's "invalid argument" /
#: "no value", written as #VALUE! where Excel says #NUM! (a rank among no
#: numbers: #N/A); numeric text typed into MEDIAN or a spread is #VALUE!
#: there, where Excel counts it, and so is numeric text as LARGE's or
#: SMALL's k, typed or in a cell — LibreOffice reads that k from a number
#: only, the engine as Excel reads a value argument (Microsoft documents
#: no rule for this k; unmeasured); a boolean in a range is a number
#: there (TRUE 1, FALSE 0), where Excel skips it; a negative fractional
#: quart is rounded down there (-0.5 is -1, out of range) where Excel
#: truncates it (to 0, the smallest); and LibreOffice's spreads read their
#: arguments from the last, so of two errors the later one shows.
_LIBREOFFICE_INSTEAD = {
    "med_empty": "#VALUE!", "p_inc_bad": "#VALUE!", "p_inc_empty": "#VALUE!",
    "p_inc_err_badk": "#VALUE!", "p_exc_below": "#VALUE!", "p_exc_0": "#VALUE!",
    "p_exc_empty": "#VALUE!", "q_inc_5": "#VALUE!", "q_exc_0": "#VALUE!",
    "q_exc_one": "#VALUE!", "q_inc_empty": "#VALUE!", "large_0": "#VALUE!",
    "large_55": "#VALUE!", "small_6": "#VALUE!", "small_05": "#VALUE!",
    "large_empty": "#VALUE!", "rank_none": "#VALUE!",
    "med_numtext_typed": "#VALUE!", "sd_numtext_typed": "#VALUE!",
    "large_text_k": "#VALUE!", "large_cell_text_k": "#VALUE!",
    "small_cell_text_k": "#VALUE!",
    "med_bools": 1, "med_flag": 2.5, "sd_bools": 1.6431676725155, "large_bools": 1,
    "rank_bools": 5,
    "q_inc_neg": "#VALUE!",
    "sd_two_errors": "#N/A",
    "p_none_pp[1]": "#VALUE!", "p_none_pp[2]": "#VALUE!",
}


def _blocks(ws):
    """A records sheet of yearly lines: ``{(field, record): its cells}`` —
    each field a block of rows under its own label, a row per record."""
    out, field = {}, None
    for row in ws.iter_rows(min_row=2):
        label = row[0].value
        if label is None:
            continue
        cells = [c.value for c in row[1:]]
        if all(v is None for v in cells):
            field = str(label).split(" (")[0]
        else:
            out[(field, str(label))] = cells
    return out


def _cells(m, book):
    """Every function cell as ``(where, engine value, LibreOffice value)``."""
    # The scalar lines take Summary's column B, the periods C onwards.
    summary = _rows(book["Summary"])
    for name, line in m.s.walk():
        if isinstance(line, mo.Variable):
            if isinstance(line.value, list):
                for period, (ours, theirs) in enumerate(zip(line.value, summary[name][1:6])):
                    yield f"{name}[{period}]", ours, theirs
            else:
                yield name, line.value, summary[name][0]
    data = _rows(book["Data"])
    for name in _DATA_PER_PERIOD:
        line = getattr(m.data, name)
        for period, (ours, theirs) in enumerate(zip(line.value, data[line.display_name][:5])):
            yield f"{name}[{period}]", ours, theirs
    ws = book["Comps"]
    header = [cell.value for cell in ws[1]]
    for row in ws.iter_rows(min_row=2):
        rec = getattr(m.comps, str(row[0].value).lower())
        for field, label in (("rank", "Rank"), ("rank_avg", "RankAvg")):
            yield (f"{row[0].value}.{field}", getattr(rec, field).value,
                   row[header.index(label)].value)
    blocks = _blocks(book["Peers"])
    for key in ("a", "b", "c", "d"):
        rec = getattr(m.peers, key)
        for field, label in (("rank", "Rank"), ("rank_avg", "RankAvg")):
            cells = blocks[(label, rec.display_name)]
            for period, (ours, theirs) in enumerate(zip(getattr(rec, field).value, cells[:5])):
                yield f"{rec.display_name}.{field}[{period}]", ours, theirs


def _same(ours, theirs):
    if isinstance(ours, str) or isinstance(theirs, str):
        return ours == theirs
    return _serial(theirs) == pytest.approx(_serial(ours), rel=1e-9, abs=1e-12)


@pytest.mark.slow
def test_the_workbook_recalculates_to_the_python_values(tmp_path):
    m = _parity_model()
    book = _recalculated(m, tmp_path)
    cells = list(_cells(m, book))
    assert len(cells) == 338
    names = {name for name, _ours, _theirs in cells}
    assert set(_LIBREOFFICE_INSTEAD) <= names
    for name, ours, theirs in cells:
        expected = _LIBREOFFICE_INSTEAD.get(name, ours)
        assert _same(expected, theirs), (name, ours, theirs)
    # A few the comparison would pass for the wrong reason.
    s = m.s
    assert (s.med_dates.value, s.large_dates.value) == (dt.date(2026, 2, 27), dt.date(2026, 2, 28))
    assert (s.large_25.value, s.small_25.value, s.large_05.value) == (30.0, 20.0, 50.0)
    assert (s.rank_num_err_ref_err.value, s.p_inc_kerr_err.value) == ("#DIV/0!", "#N/A")
    assert (s.large_kerr.value, s.large_err_badk.value) == ("#N/A", "#DIV/0!")
    assert (s.med_two_errors.value, s.sd_two_errors.value) == ("#DIV/0!", "#DIV/0!")
    assert (s.large_cell_text_k.value, s.small_cell_text_k.value) == (40.0, 20.0)
    # One call per period, never one number over every cell.
    d = m.data
    assert d.med_pp.value == [25.0, 15.0, 30.0, 30.0, 50.0]
    assert d.p_pp.value == [7.5, 6.5, 17.5, 10.5, 26.5]
    assert d.rank_pp.value == [1, 2, 1, 2, 1]
    assert s.peer_med.value == [17.5, 35.0, 27.5, 8.0, 5.0]
    assert m.peers.a.rank.value == [4, 3, 3, "#N/A", 1]


@pytest.mark.slow
@pytest.mark.xfail(strict=True, reason=(
    "LibreOffice writes a k or a quart out of range and a range without numbers as #VALUE! "
    "(Excel: #NUM!, a rank #N/A), refuses numeric text typed into MEDIAN or a spread "
    "and as LARGE's or SMALL's k (typed or in a cell), counts a boolean in a range, "
    "rounds a negative quart down and reads a spread's arguments from the last; the "
    "engine keeps Excel's rules — see _LIBREOFFICE_INSTEAD."
))
def test_where_libreoffice_departs_from_excel(tmp_path):
    m = _parity_model()
    book = _recalculated(m, tmp_path)
    departing = [(name, ours, theirs) for name, ours, theirs in _cells(m, book)
                 if name in _LIBREOFFICE_INSTEAD and not _same(ours, theirs)]
    assert departing == [], departing
