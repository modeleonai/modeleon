# SPDX-License-Identifier: Apache-2.0
"""``excel_props={'hidden': True}`` — helper rows and sheets.

A hidden row is not a skipped one: it is laid out and written in full
(values, formulas, labels), so everything that references it stays
live. Excel's NATIVE mechanism hides it (``row_dimensions[r].hidden`` /
``sheet_state='hidden'``) — the workbook stays auditable and the eye
is not cluttered. The flag inherits down the MV tree, so one key on a
helper section is enough.
"""

import openpyxl
import pytest

import modeleon as mo


@pytest.fixture()
def book(tmp_path):
    def _build(model):
        path = tmp_path / "out.xlsx"
        model.to_excel(str(path))
        return openpyxl.load_workbook(str(path))

    return _build


class TestHiddenRows:
    def test_flagged_row_is_written_but_natively_hidden(self, book):
        m = mo.Model("m")
        with m:
            m.лист = mo.MultiVariable("Лист", excel_props={"tab": True})
            with m.лист as л:
                л.выручка = mo.Variable([10.0, 20.0], display_name="Выручка")
                л.служебная = mo.Variable(
                    [1.0, 2.0],
                    display_name="Служебная",
                    excel_props={"hidden": True},
                )
        wb = book(m)
        ws = wb["Лист"]
        rows = {
            c.value: c.row
            for col in ws.iter_cols(max_col=2)
            for c in col
            if isinstance(c.value, str)
        }
        # Written — and the values are in place…
        served = rows["Служебная"]
        assert ws.cell(row=served, column=2).value == 1.0
        # …but the row is natively hidden; the ordinary one is not.
        assert ws.row_dimensions[served].hidden is True
        assert not ws.row_dimensions[rows["Выручка"]].hidden

    def test_flag_inherits_down_the_mv_tree(self, book):
        m = mo.Model("m")
        with m:
            m.лист = mo.MultiVariable("Лист", excel_props={"tab": True})
            with m.лист as л:
                л.видимая = mo.Variable([5.0], display_name="Видимая")
                л.расчёты = mo.MultiVariable(
                    "Расчёты", excel_props={"hidden": True}
                )
                л.расчёты.промежуточная = mo.Variable(
                    [7.0], display_name="Промежуточная"
                )
        wb = book(m)
        ws = wb["Лист"]
        by_label = {
            c.value: c.row
            for col in ws.iter_cols(max_col=2)
            for c in col
            if isinstance(c.value, str)
        }
        # The section header and its child row are hidden by ONE key.
        assert ws.row_dimensions[by_label["Расчёты"]].hidden is True
        assert ws.row_dimensions[by_label["Промежуточная"]].hidden is True
        assert not ws.row_dimensions[by_label["Видимая"]].hidden

    def test_references_to_hidden_rows_stay_live_formulas(self, book):
        # A hidden row keeps its addresses — a formula on top of it stays
        # a live reference, not a baked-in literal.
        m = mo.Model("m")
        with m:
            m.лист = mo.MultiVariable("Лист", excel_props={"tab": True})
            with m.лист as л:
                л.база = mo.Variable(
                    [3.0], display_name="База", excel_props={"hidden": True}
                )
                л.двойная = mo.Variable(
                    л.база * 2, display_name="Двойная"
                )
        wb = book(m)
        ws = wb["Лист"]
        rows = {
            c.value: c.row
            for col in ws.iter_cols(max_col=2)
            for c in col
            if isinstance(c.value, str)
        }
        formula = ws.cell(row=rows["Двойная"], column=2).value
        assert isinstance(formula, str) and formula.startswith("=")
        # A live reference to the hidden row's cell — not a baked-in literal.
        assert f"B{rows['База']}" in formula


class TestHiddenSheet:
    def test_flagged_tab_hides_the_whole_sheet(self, book):
        m = mo.Model("m")
        with m:
            m.данные = mo.MultiVariable("Данные", excel_props={"tab": True})
            with m.данные as д:
                д.x = mo.Variable([1.0], display_name="X")
            m.служебный = mo.MultiVariable(
                "Служебный",
                excel_props={"tab": True, "hidden": True},
            )
            with m.служебный as с:
                с.y = mo.Variable([2.0], display_name="Y")
        wb = book(m)
        assert wb["Данные"].sheet_state == "visible"
        assert wb["Служебный"].sheet_state == "hidden"
        # The sheet is written in full — it can still be audited.
        assert wb["Служебный"].max_row >= 1


class TestHiddenSurvivesTabWrappers:
    def test_hidden_mv_wraps_into_a_hidden_virtual_sheet(self, book):
        # A top-level MV without the tab flag is wrapped by the engine
        # into a VIRTUAL tab — the wrapper must not LOSE the hidden
        # flag: hidden source → hidden sheet.
        m = mo.Model("m")
        with m:
            m.данные = mo.MultiVariable("Данные", excel_props={"tab": True})
            with m.данные as д:
                д.x = mo.Variable([1.0], display_name="X")
            m.служебный = mo.MultiVariable(
                "Служебный", excel_props={"hidden": True}
            )
            with m.служебный as с:
                с.y = mo.Variable([9.0], display_name="Y")
        wb = book(m)
        служ = next(n for n in wb.sheetnames if "Служебный" in n)
        assert wb[служ].sheet_state == "hidden"
