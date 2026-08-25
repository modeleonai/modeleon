"""``('row_sum', caption)`` projected column — the row's LIVE total.

A view that projects the ``row_sum`` field gives every opted-in row
(``excel_props={'row_sum': True}``) a ``=SUM(...)`` over its own value
cells in that lead column; rows without the flag leave it empty (a
ratio row's sum is nonsense). Non-contiguous value runs (subtotal
columns interleave with months) compress to comma-joined ranges.
"""

import os
import tempfile

import openpyxl
import pytest

import modeleon as mo
from modeleon.compile.excel.writer import row_sum_formula


def _sheet(model, name):
    path = os.path.join(tempfile.mkdtemp(), "t.xlsx")
    model.to_excel(path)
    return openpyxl.load_workbook(path)[name]


def _windowed():
    m = mo.Model("m", display_name="M")
    m.default_start, m.default_grain, m.default_periods = "2026-01", "month", 3
    m.default_excel_view = mo.ExcelView(
        meta={"fields": [("row_sum", "Итого")]}
    )
    return m


class TestRowSumFormula:
    def test_contiguous_run_is_one_range(self):
        assert row_sum_formula(["B5", "C5", "D5"]) == "=SUM(B5:D5)"

    def test_interleaved_runs_join_with_commas(self):
        assert row_sum_formula(["B5", "C5", "F5", "G5"]) == "=SUM(B5:C5,F5:G5)"

    def test_single_cell_and_empty(self):
        assert row_sum_formula(["B5"]) == "=SUM(B5)"
        assert row_sum_formula([]) is None


class TestRowSumColumn:
    def test_flagged_row_gets_a_live_sum(self):
        m = _windowed()
        with m:
            m.план = mo.MultiVariable("План", excel_props={"tab": True})
            with m.план as p:
                p.доля = mo.Variable(
                    [1.0, 0.5, 0.0],
                    display_name="Доля",
                    excel_props={"row_sum": True},
                )
        ws = _sheet(m, "План")
        # Caption heads the projected column; the row carries =SUM.
        grid = [
            [ws.cell(r, c).value for c in range(1, 7)]
            for r in range(1, ws.max_row + 1)
        ]
        flat = [x for row in grid for x in row if x is not None]
        assert "Итого" in flat
        sums = [x for x in flat if isinstance(x, str) and x.startswith("=SUM(")]
        assert len(sums) == 1
        assert ":" in sums[0]

    def test_unflagged_row_leaves_the_column_empty(self):
        m = _windowed()
        with m:
            m.план = mo.MultiVariable("План", excel_props={"tab": True})
            with m.план as p:
                p.вес = mo.Variable([0.5, 0.5, 0.5], display_name="Вес")
        ws = _sheet(m, "План")
        flat = [
            ws.cell(r, c).value
            for r in range(1, ws.max_row + 1)
            for c in range(1, 7)
        ]
        assert not any(
            isinstance(x, str) and x.startswith("=SUM(") for x in flat
        )

    def test_scalar_and_all_blank_rows_refuse_the_sum(self):
        # One law with the grid: a scalar is not a run of periods
        # (=SUM(D5) beside the number is noise), and an all-blank row
        # must not total a fabricated 0.
        m = _windowed()
        with m:
            m.план = mo.MultiVariable("План", excel_props={"tab": True})
            with m.план as p:
                p.ставка = mo.Variable(
                    500_000, display_name="Ставка",
                    excel_props={"row_sum": True},
                )
                p.пусто = mo.Variable(
                    [None, None, None], display_name="Пусто",
                    excel_props={"row_sum": True},
                )
                p.живая = mo.Variable(
                    [1.0, 0.5, None], display_name="Живая",
                    excel_props={"row_sum": True},
                )
        ws = _sheet(m, "План")
        sums = [
            ws.cell(r, c).value
            for r in range(1, ws.max_row + 1)
            for c in range(1, 8)
            if isinstance(ws.cell(r, c).value, str)
            and ws.cell(r, c).value.startswith("=SUM(")
        ]
        assert len(sums) == 1  # только «Живая»
