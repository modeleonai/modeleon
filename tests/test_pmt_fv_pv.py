# SPDX-License-Identifier: Apache-2.0
"""``PMT``, ``FV`` and ``PV``: one call per period, inside a loop container
too — over the loop's own rows, as the schedule functions (``IPMT``,
``PPMT``, ``NPER``, ``RATE``) already run there — and, over scalars as
over a series, what the recalculated book shows at the edges of their
domain."""

from __future__ import annotations

import math
import shutil
import subprocess
from decimal import Decimal
from pathlib import Path

import numpy as np
import pytest
from openpyxl import load_workbook

import modeleon as mo

N = 12
RATE = 0.01
LOAN = 1200.0
CASH = [0.0, 0.0, 300.0, 0.0, 0.0, 0.0, 200.0, 0.0, 0.0, 0.0, 0.0, 0.0]


def _pmt(rate, nper, pv, fv=0.0):
    """Excel's PMT at type 0, as the payment the borrower makes."""
    growth = (1 + rate) ** nper
    return rate * (pv * growth + fv) / (growth - 1)


def _schedule(cash):
    """The debt worked by hand: each month's annuity on the balance as it
    stands over the months left, the interest, then the cash swept."""
    rows = {"payment": [], "close": []}
    balance = LOAN
    for t in range(N):
        payment = _pmt(RATE, N - t, balance)
        interest = balance * RATE
        sweep = min(cash[t], balance + interest - payment)
        balance = balance + interest - payment - sweep
        rows["payment"].append(payment)
        rows["close"].append(balance)
    return rows


def _debt(cash=None, balloon_unit="USD"):
    m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=N)
    with m:
        m.i = mo.MultiVariable("Inputs")
        with m.i as i:
            i.rate = mo.Variable(RATE)
            i.loan = mo.Variable(LOAN, unit="USD")
            i.balloon = mo.Variable(0.0, unit=balloon_unit)
            i.cash = mo.Variable(list(cash or [0.0] * N), unit="USD")
        m.d = mo.MultiVariable("Debt", loop=True)
        with m.d as d:
            d.open = mo.lag(d.close, fill=m.i.loan)
            d.left = N - mo.time.index + 1
            d.payment = -mo.PMT(m.i.rate, d.left, d.open, m.i.balloon)
            d.interest = d.open * m.i.rate
            d.sweep = mo.MIN(m.i.cash, d.open + d.interest - d.payment)
            d.close = d.open + d.interest - d.payment - d.sweep
            d.grown = mo.FV(m.i.rate, 1, 0, -d.open)
            d.worth = mo.PV(m.i.rate, 1, 0, -d.close)
    return m


def _lines(ws, periods):
    """A sheet's lines by label: each one's first ``periods`` cells."""
    return {row[0].value: [c.value for c in row[1:periods + 1]]
            for row in ws.iter_rows() if row[0].value}


def test_a_payment_on_the_balance_as_it_stands_is_one_call_per_period():
    # Re-worked each month on what is left, an annuity pays the same every
    # month and closes the loan at zero.
    m = _debt()
    assert m.d.payment.value == pytest.approx([_pmt(RATE, N, LOAN)] * N)
    assert m.d.close.value[-1] == pytest.approx(0.0, abs=1e-9)
    assert {str(m.d.payment.unit), str(m.d.grown.unit), str(m.d.worth.unit)} == {"USD"}


def test_a_cash_sweep_matches_the_schedule_worked_by_hand():
    m = _debt(cash=CASH)
    hand = _schedule(CASH)
    assert m.d.payment.value == pytest.approx(hand["payment"])
    assert m.d.close.value == pytest.approx(hand["close"])
    assert m.d.grown.value == pytest.approx([o * (1 + RATE) for o in m.d.open.value])
    assert m.d.worth.value == pytest.approx([c / (1 + RATE) for c in m.d.close.value])


def test_an_amount_in_another_unit_is_refused_when_the_loop_closes():
    with pytest.raises(ValueError, match="mo.PMT takes its values in one unit; "
                                         "'open' and 'balloon' are in USD and EUR"):
        _debt(balloon_unit="EUR")


def test_the_book_holds_a_live_call_per_period(tmp_path):
    _debt(cash=CASH).to_excel(tmp_path / "debt.xlsx")
    rows = _lines(load_workbook(tmp_path / "debt.xlsx")["Debt"], N)
    assert rows["Payment (USD)"][:2] == ["=-PMT(Inputs!B2, B3, B2, Inputs!B4, 0)",
                                         "=-PMT(Inputs!B2, C3, C2, Inputs!B4, 0)"]
    assert rows["Grown (USD)"][:2] == ["=FV(Inputs!B2, 1, 0, -B2, 0)",
                                       "=FV(Inputs!B2, 1, 0, -C2, 0)"]
    assert rows["Worth (USD)"][:2] == ["=PV(Inputs!B2, 1, 0, -B7, 0)",
                                       "=PV(Inputs!B2, 1, 0, -C7, 0)"]


def _past_the_term(payment):
    """A five-month loan in an eight-month window: from month six there is
    no term left. ``payment(rate, d)`` writes the payment row."""
    m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=8)
    with m:
        m.i = mo.MultiVariable("Inputs")
        with m.i as i:
            i.rate = mo.Variable(RATE)
            i.loan = mo.Variable(1000.0, unit="USD")
        m.d = mo.MultiVariable("Debt", loop=True)
        with m.d as d:
            d.open = mo.lag(d.close, fill=m.i.loan)
            d.left = mo.IF(mo.time.index <= 5, 6 - mo.time.index, 0)
            d.payment = payment(m.i.rate, d)
            d.close = d.open * (1 + m.i.rate) - d.payment
    return m


def test_a_payment_past_the_term_is_an_error_a_guard_reads():
    # PMT over no term left is #NUM!, as in the book: the period carries
    # it, and an IF on the term or an IFERROR around the call reads it.
    due = _pmt(RATE, 5, 1000.0)
    bare = _past_the_term(lambda r, d: -mo.PMT(r, d.left, d.open))
    assert bare.d.payment.value[:5] == pytest.approx([due] * 5)
    assert bare.d.payment.value[5] == '#NUM!'
    for guard in (lambda r, d: mo.IF(d.left > 0, -mo.PMT(r, d.left, d.open), 0),
                  lambda r, d: mo.IFERROR(-mo.PMT(r, d.left, d.open), 0)):
        m = _past_the_term(guard)
        assert m.d.payment.value[:5] == pytest.approx([due] * 5)
        assert m.d.payment.value[5:] == [0, 0, 0]
        assert m.d.close.value[4:] == pytest.approx([0.0] * 4, abs=1e-9)


def test_an_error_cell_reaching_the_loop_is_carried_from_its_period():
    # A rate that is #DIV/0! in March: the payment carries it from March
    # on, as IPMT's does, and the months before keep their numbers.
    m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=6)
    with m:
        m.i = mo.MultiVariable("Inputs")
        m.i.rate = RATE / mo.Variable([1, 1, 0, 1, 1, 1])
        m.d = mo.MultiVariable("Debt", loop=True)
        with m.d as d:
            d.open = mo.lag(d.close, fill=1000.0)
            d.left = 6 - mo.time.index + 1
            d.payment = -mo.PMT(m.i.rate, d.left, d.open)
            d.grown = mo.FV(m.i.rate, 1, 0, -d.open)
            d.close = d.open * (1 + RATE) - d.payment
    due = _pmt(RATE, 6, 1000.0)
    assert m.d.payment.value[:2] == pytest.approx([due, due])
    assert m.d.payment.value[2:] == ['#DIV/0!'] * 4
    assert m.d.grown.value[2:] == ['#DIV/0!'] * 4


class TestWhatTheBookShows:
    """Over scalars as over a series, a call gives what the recalculated
    book shows (the numbers below are LibreOffice's): ``#NUM!`` where its
    result is ``#NUM!``, the error a cell holds, and its numbers at the
    edges of the domain."""

    @pytest.mark.parametrize("call, expected", [
        (lambda: mo.PMT(0.01, 0, 1000), "#NUM!"),           # no term
        (lambda: mo.PMT(0, 0, 1000), "#NUM!"),
        (lambda: mo.PMT(-1, 0, 1000, 100), "#NUM!"),        # at -100%, no term or a negative one
        (lambda: mo.PMT(-1, -1, 1000, 100), "#NUM!"),
        (lambda: mo.PMT(-1, 12, 1000, 0, 1), "#NUM!"),      # paid at the start, at -100%
        (lambda: mo.PMT(-1.5, 2, 1000), "#NUM!"),           # below -100%: its power takes log1p
        (lambda: mo.PV(-1, 5, -10), "#NUM!"),               # discounted at -100%
        (lambda: mo.FV(-1, -1, -10, -1000), "#NUM!"),
        (lambda: mo.FV(-1.5, 0.5, -10), "#NUM!"),           # a fractional power of a negative
        (lambda: mo.PMT(1.0, 5000, 1000), "#NUM!"),         # (1 + rate) ** nper too large for a
                                                            # cell, though -1000 would fit
        (lambda: mo.FV(1.0, 5000, -10), "#NUM!"),           # a result too large for a cell
        (lambda: mo.FV(1, 5000, 0, 0), "#NUM!"),            # 1 as 1.0: a cell holds a double
        (lambda: mo.FV(1, 1000, 0, -1e10), "#NUM!"),        # a product too large for a cell
        (lambda: mo.PMT(1, 1000, 1e10), "#NUM!"),
        (lambda: mo.PV(-0.9, 10, 0, 1e300), "#NUM!"),
        (lambda: mo.PMT(mo.Variable("#DIV/0!"), 4, 1000), "#DIV/0!"),  # a cell's error
        (lambda: mo.PV(0.01, "x", -10), "#VALUE!"),         # text that is no number
    ])
    def test_where_the_book_shows_an_error(self, call, expected):
        assert call().value == expected

    @pytest.mark.parametrize("call, expected", [
        (lambda: mo.FV(-1.5, 2, 0, -100), 25.0),            # a whole power of a negative
        (lambda: mo.PV(1.0, 5000, -10), 10.0),              # discounted by (1 + rate) ** -nper
        (lambda: mo.PV(-1, -1, -10), -10.0),
        (lambda: mo.FV(0.01, 4, -10, 0, 1), 41.0100501),    # paid at the start of a period
        (lambda: mo.FV(0, 12, -10, -1000), 1120.0),         # at a rate of 0
        (lambda: mo.PV(0, 2, -10, 1000), -980.0),
        (lambda: mo.PV(0, 12, -10, 1000, 1), -880.0),
        (lambda: mo.PV(1e-12, 2, -10, 1000), -979.998221986353),
        (lambda: mo.PMT(0.01, "12", 1000), -88.8487886783417),  # numeric text is its number
    ])
    def test_where_the_book_shows_a_number(self, call, expected):
        assert call().value == pytest.approx(expected, rel=1e-12)

    def test_over_a_series(self):
        assert mo.PMT(0.01, mo.Variable([4, 0, 4]), 1000).value[1] == "#NUM!"
        assert mo.FV(mo.Variable([-1.5, 0.01]), 0.5, -10).value[0] == "#NUM!"

    def test_a_zero_is_never_negative(self):
        for call in (mo.PMT(0.01, 12, 0), mo.FV(0.01, 12, 0, 0), mo.PV(0.01, 12, 0, 0)):
            assert call.value == 0 and math.copysign(1.0, call.value) == 1.0

    def test_a_tiny_rate_keeps_its_digits(self):
        # (1 + rate) ** nper - 1 through expm1, as the spreadsheet computes
        # it and IPMT and PPMT split it: the subtraction would keep four.
        r = Decimal(1e-12)
        t = (1 + r) ** 12
        exact = float(-1000 * r * t / (t - 1))
        assert mo.PMT(1e-12, 12, 1000).value == pytest.approx(exact, rel=1e-12)

    def test_a_type_other_than_zero_reads_as_one(self):
        assert mo.PMT(0.01, 4, 1000, 0, 2).value == mo.PMT(0.01, 4, 1000, 0, 1).value

    def test_a_decimal_or_a_numpy_value_is_its_number(self):
        # As the book reads the cell; IPMT reads it so too.
        due, early = mo.PMT(0.05, 12, 1000).value, mo.PMT(0.05, 12, 1000, 0, 1).value
        assert mo.PMT(Decimal("0.05"), 12, 1000).value == pytest.approx(due)
        assert mo.PMT(np.array(0.05), 12, 1000).value == pytest.approx(due)
        assert mo.PMT(0.05, 12, 1000, 0, np.True_).value == early
        assert mo.PMT(0.05, 12, 1000, 0, mo.Variable(np.float64(0.05)) > 0).value == early
        assert mo.IPMT(Decimal("0.05"), 2, 12, 1000).value == pytest.approx(
            mo.IPMT(0.05, 2, 12, 1000).value)

    def test_the_text_keeps_the_numbers_given(self):
        rate, loan, pay = mo.Variable(0.05), mo.Variable(1000), mo.Variable(-10)
        assert mo.PMT(rate, 12, loan).source_code == "PMT(0.05, 12, 1000, 0, 0)"
        assert mo.FV(rate, 12, pay).source_code == "FV(0.05, 12, -10, 0, 0)"
        assert mo.PV(rate, 12, pay).source_code == "PV(0.05, 12, -10, 0, 0)"


def _soffice():
    for candidate in ("soffice", "libreoffice"):
        found = shutil.which(candidate)
        if found:
            return found
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


def _recalculated(soffice, book, tmp_path):
    # The profile recalculates on load: conftest seeds every profile a
    # test starts LibreOffice with.
    subprocess.run(
        [soffice, f"-env:UserInstallation={(tmp_path / 'prof').as_uri()}",
         "--headless", "--convert-to", "xlsx", "--outdir", str(tmp_path / "calc"), str(book)],
        check=True, capture_output=True, timeout=180,
    )
    return load_workbook(tmp_path / "calc" / book.name, data_only=True)


@pytest.mark.slow
def test_the_book_recalculates_to_the_python_values(tmp_path):
    soffice = _soffice()
    if soffice is None:
        pytest.skip("LibreOffice is not installed")
    m = _debt(cash=CASH)
    m.to_excel(tmp_path / "debt.xlsx")
    book = _lines(_recalculated(soffice, tmp_path / "debt.xlsx", tmp_path)["Debt"], N)
    for label, line in (("Payment (USD)", m.d.payment), ("Close (USD)", m.d.close),
                        ("Grown (USD)", m.d.grown), ("Worth (USD)", m.d.worth)):
        assert book[label] == pytest.approx(line.value, rel=1e-9, abs=1e-9), label


#: Rates, terms and types at the edges of the three functions' domain, one
#: period each: below and at -100%, 0, a tiny rate, 100%; no term, a
#: fraction of one, a negative one, one whose power with 100% leaves a
#: product too large for a cell, one whose power is.
_EDGES = [(r, n, w) for r in (-1.5, -1, -0.5, 0, 1e-12, 0.01, 1.0)
          for n in (0, 0.5, 2, 12, -1, 1020, 5000) for w in (0, 1)]
_EDGES += [(0.01, 12, 2), (0.01, 12, 0.5)]


def _edges():
    m = mo.Model("m", default_grain="month", default_start="2026-01",
                 default_periods=len(_EDGES))
    with m:
        m.i = mo.MultiVariable("Inputs")
        with m.i as i:
            i.rate = mo.Variable([r for r, _, _ in _EDGES])
            i.term = mo.Variable([n for _, n, _ in _EDGES])
            i.kind = mo.Variable([w for _, _, w in _EDGES])
        m.e = mo.MultiVariable("Edges")
        with m.e as e:
            e.payment = mo.PMT(m.i.rate, m.i.term, 1000, 100, m.i.kind)
            e.future = mo.FV(m.i.rate, m.i.term, -10, -1000, m.i.kind)
            e.present = mo.PV(m.i.rate, m.i.term, -10, 1000, m.i.kind)
    return m


@pytest.mark.slow
def test_at_the_edges_the_book_recalculates_to_the_python_values(tmp_path):
    soffice = _soffice()
    if soffice is None:
        pytest.skip("LibreOffice is not installed")
    m = _edges()
    m.to_excel(tmp_path / "edges.xlsx")
    book = _lines(_recalculated(soffice, tmp_path / "edges.xlsx", tmp_path)["Edges"],
                  len(_EDGES))
    for label, line in (("Payment", m.e.payment), ("Future", m.e.future),
                        ("Present", m.e.present)):
        assert len(book[label]) == len(line.value) == len(_EDGES)
        for case, python, calc in zip(_EDGES, line.value, book[label]):
            if isinstance(python, str):
                assert calc == python, (label, case)
            else:
                assert calc == pytest.approx(python, rel=1e-12, abs=1e-12), (label, case)
