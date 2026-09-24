"""A section that follows plain variable rows breathes before its header.

The walk always left a gap AFTER a section, but a section following a
run of Variables glued its header straight to the last row. The pre-gap
follows the view's ``nesting.gap_rows``; dense books (``gap_rows: 0``)
are untouched.
"""

import os
import tempfile

import openpyxl

import modeleon as mo


def _sheet(model, name):
    path = os.path.join(tempfile.mkdtemp(), "t.xlsx")
    model.to_excel(path)
    return openpyxl.load_workbook(path)[name]


def _labels(ws):
    return [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]


def _build(gap_rows=None):
    m = mo.Model("m", display_name="M")
    with m:
        if gap_rows is not None:
            m.default_excel_view = mo.ExcelView(
                nesting={"gap_rows": gap_rows, "indent": 2}
            )
        m.sheet = mo.MultiVariable("Sheet", excel_props={"tab": True})
        with m.sheet as s:
            s.a = mo.Variable(1, display_name="A")
            s.b = mo.Variable(2, display_name="B")
            s.params = mo.MultiVariable(display_name="Params")
            with s.params as p:
                p.x = mo.Variable(3, display_name="X")
    return m


def test_default_gap_separates_rows_from_a_following_section():
    labels = _labels(_sheet(_build(), "Sheet"))
    b_row = labels.index("B")
    header_row = labels.index("Params")
    # one blank row between the last variable and the section header
    assert header_row == b_row + 2
    assert labels[b_row + 1] is None


def test_dense_view_keeps_sections_glued():
    labels = _labels(_sheet(_build(gap_rows=0), "Sheet"))
    b_row = labels.index("B")
    assert labels.index("Params") == b_row + 1
