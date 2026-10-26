# SPDX-License-Identifier: Apache-2.0
"""A series where a function used to take one value is one call per period.

``ROUND``'s digits, ``EDATE`` / ``EOMONTH``'s months, ``DAYS360``'s
method and every argument of ``PMT``, ``FV`` and ``PV`` may be a line:
the call is one per period, as the spreadsheet writes it in every column
(``=ROUND(C5, C9)``, ``=EDATE(C5, C9)``, ``=PMT(C9 / 100, 12, 1000, 0,
0)``) — a scalar broadcast, lines on other windows read by date and
lines along keys key by key, as ``a + b`` reads them. A rate ``NPV``
discounts a whole row by, and the guess ``IRR`` / ``XIRR`` start from,
stay one value: a series there is refused by name.
"""

import shutil
import subprocess
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

import modeleon as mo

pytestmark = pytest.mark.filterwarnings("ignore::UserWarning")


def _rows(ws):
    out = {}
    for row in ws.iter_rows():
        label = row[0].value
        if label:
            out[str(label)] = [c.value for c in row[1:]]
    return out


def _row_of(ws):
    return {str(row[0].value): row[0].row for row in ws.iter_rows() if row[0].value}


def _lines():
    """Three years of inputs on X, a line from 2027 on T (held), keyed
    lines on K."""
    m = mo.Model("m", default_grain="year", default_start="2026", default_periods=3)
    with m:
        m.x = mo.MultiVariable("X")
        m.x.v = mo.Variable([1.2345, 2.3456, -3.4567], display_name="v")
        m.x.d = mo.Variable([0.0, 1.0, 2.0], display_name="d")
        m.x.dt = mo.Variable([date(2026, 1, 31), date(2026, 2, 28), date(2026, 3, 31)],
                             display_name="dt")
        m.x.k = mo.Variable([1.0, 2.0, -1.0], display_name="k")
        m.x.pct = mo.Variable([5.0, 6.0, 7.0], display_name="pct")
        m.x.rt = mo.Variable([0.05, 0.06, 0.07], display_name="rt")
        m.x.n = mo.Variable([12.0, 24.0, 36.0], display_name="n")
        m.x.amt = mo.Variable([1000.0, 2000.0, 3000.0], display_name="amt")
        m.x.eu = mo.Variable([True, False, True], display_name="eu")
        m.x.end = mo.Variable([date(2026, 3, 31), date(2026, 3, 31), date(2026, 5, 31)],
                              display_name="end")
        m.t = mo.MultiVariable("T", default_grain="year", default_start="2027",
                               default_periods=2)
        m.t.dg = mo.Variable([3.0, 1.0], display_name="dg", extend=mo.hold())
        m.t.nofill = mo.Variable([3.0, 1.0], display_name="nofill")
    return m


def _one_by_one(make, *series):
    """The scalar call made in each period — what a period's cell computes."""
    return [make(*(s[i] if isinstance(s, list) else s for s in series)).value
            for i in range(max(len(s) for s in series if isinstance(s, list)))]


#: Each call over X's lines, with the values the scalar call gives period by
#: period.
_CALLS = {
    "rnd": (lambda m: mo.ROUND(m.x.v, m.x.d),
            lambda m: _one_by_one(mo.ROUND, m.x.v.value, m.x.d.value)),
    "rnd2": (lambda m: mo.ROUND(2.71828, m.x.d),
             lambda m: _one_by_one(mo.ROUND, 2.71828, m.x.d.value)),
    "ed": (lambda m: mo.EDATE(m.x.dt, m.x.k),
           lambda m: _one_by_one(mo.EDATE, m.x.dt.value, m.x.k.value)),
    "ed2": (lambda m: mo.EDATE(date(2026, 1, 31), m.x.k),
            lambda m: _one_by_one(mo.EDATE, date(2026, 1, 31), m.x.k.value)),
    "eo": (lambda m: mo.EOMONTH(m.x.dt, m.x.k),
           lambda m: _one_by_one(mo.EOMONTH, m.x.dt.value, m.x.k.value)),
    "pm": (lambda m: mo.PMT(m.x.pct / 100, 12, 1000),
           lambda m: _one_by_one(mo.PMT, [p / 100 for p in m.x.pct.value], 12, 1000)),
    "pm2": (lambda m: mo.PMT(0.05, m.x.n, m.x.amt, 0, 1),
            lambda m: _one_by_one(mo.PMT, 0.05, m.x.n.value, m.x.amt.value, 0, 1)),
    "fv": (lambda m: mo.FV(m.x.pct / 100, 12, -100),
           lambda m: _one_by_one(mo.FV, [p / 100 for p in m.x.pct.value], 12, -100)),
    "pv": (lambda m: mo.PV(0.05, m.x.n, -100, m.x.amt),
           lambda m: _one_by_one(mo.PV, 0.05, m.x.n.value, -100, m.x.amt.value)),
    "d360": (lambda m: mo.DAYS360(m.x.dt, m.x.end, m.x.eu),
             lambda m: _one_by_one(mo.DAYS360, m.x.dt.value, m.x.end.value, m.x.eu.value)),
}


def _model():
    m = _lines()
    with m:
        for name, (make, _) in _CALLS.items():
            setattr(m.x, name, make(m).set_display_name(name))
    return m


class TestEachPeriodIsTheScalarCall:
    @pytest.mark.parametrize("name", list(_CALLS))
    def test_the_value_is_the_scalar_call_in_each_period(self, name):
        m = _lines()
        make, expected = _CALLS[name]
        got = make(m)
        want = expected(m)
        assert got.var_type == "list" and len(got.value) == 3
        if isinstance(want[0], float):
            assert got.value == pytest.approx(want)
        else:
            assert got.value == want

    def test_the_digits_of_each_period(self):
        m = _lines()
        assert mo.ROUND(m.x.v, m.x.d).value == [1.0, 2.3, -3.46]
        assert mo.ROUND(2.71828, m.x.d).value == [3.0, 2.7, 2.72]

    def test_the_months_of_each_period(self):
        m = _lines()
        assert mo.EDATE(m.x.dt, m.x.k).value == [
            date(2026, 2, 28), date(2026, 4, 28), date(2026, 2, 28)]
        assert mo.EOMONTH(m.x.dt, m.x.k).value == [
            date(2026, 2, 28), date(2026, 4, 30), date(2026, 2, 28)]

    def test_the_method_of_each_period(self):
        # 28 Feb to 31 Mar: the U.S. method counts 30, the European 32.
        m = _lines()
        assert mo.DAYS360(m.x.dt, m.x.end, m.x.eu).value == [60, 30, 60]
        assert mo.DAYS360(m.x.dt, m.x.end, True).value == [60, 32, 60]

    def test_one_value_keeps_the_call_as_it_was(self):
        m = _lines()
        assert mo.ROUND(m.x.v, 1).value == [1.2, 2.3, -3.5]
        assert mo.ROUND(m.x.v, 1).formula == "ROUND(v, 1)"
        assert mo.EDATE(m.x.dt, 1).value[0] == date(2026, 2, 28)
        assert mo.PMT(0.05, 12, 1000).value == pytest.approx(-112.82541002081534)
        assert mo.PMT(0.05, 12, 1000).var_type == "scalar"
        assert mo.DAYS360(m.x.dt, m.x.end, True).value == [60, 32, 60]

    def test_a_blank_month_is_refused_by_name(self):
        m = _lines()
        with m:
            m.x.gap = mo.Variable([1.0, None, 2.0], display_name="gap")
        with pytest.raises(TypeError, match="'gap' is blank in"):
            mo.EDATE(m.x.dt, m.x.gap)


class TestAsTheOperatorsReadTheirOperands:
    def test_a_line_on_another_window_is_read_by_date(self):
        # dg runs from 2027 and is held back to 2026: digits [3, 3, 1].
        m = _lines()
        got = mo.ROUND(m.x.v, m.t.dg)
        assert got.value == [1.235, 2.346, -3.5]
        assert got.time is not None and got.time.start == "2026"
        assert mo.EDATE(m.x.dt, m.t.dg).value == [
            date(2026, 4, 30), date(2026, 5, 28), date(2026, 4, 30)]
        assert mo.PMT(m.x.rt, m.t.dg, 1000).value == pytest.approx(
            [mo.PMT(r, n, 1000).value for r, n in zip([0.05, 0.06, 0.07], [3, 3, 1])])
        assert mo.DAYS360(m.x.dt, m.x.end, m.t.dg).value == [60, 32, 60]

    def test_a_line_without_a_rule_is_refused(self):
        m = _lines()
        for call in (lambda: mo.ROUND(m.x.v, m.t.nofill),
                     lambda: mo.EDATE(m.x.dt, m.t.nofill),
                     lambda: mo.EOMONTH(m.x.dt, m.t.nofill),
                     lambda: mo.PMT(0.05, m.t.nofill, m.x.amt),
                     lambda: mo.FV(m.x.rt, m.t.nofill, -100),
                     lambda: mo.PV(0.05, m.t.nofill, -100, m.x.amt),
                     lambda: mo.DAYS360(m.x.dt, m.x.end, m.t.nofill)):
            with pytest.raises(ValueError, match="'nofill' does not cover the combined window"):
                call()

    def test_lines_along_keys_combine_key_by_key(self):
        m = mo.Model("m", default_grain="year", default_start="2026", default_periods=3)
        with m:
            m.r = mo.Variable({"Bear": 0.04, "Base": 0.05, "Bull": 0.06}, display_name="r")
            m.n = mo.Variable({"Bull": 10.0, "Base": 20.0, "Bear": 30.0}, display_name="n")
            m.s = mo.Variable([1.0, 2.0, 3.0], display_name="s")
        got = mo.PMT(m.r, m.n, 1000)
        assert got._keys == ["Bear", "Base", "Bull"]
        assert got.value == pytest.approx([mo.PMT(0.04, 30, 1000).value,
                                           mo.PMT(0.05, 20, 1000).value,
                                           mo.PMT(0.06, 10, 1000).value])
        assert mo.ROUND(1.23456, m.n)._keys == ["Bull", "Base", "Bear"]
        with pytest.raises(ValueError, match="runs along its keys"):
            mo.PMT(m.r, m.s, 1000)


class TestOneValueStaysOneValue:
    def test_a_series_rate_for_npv_is_refused_by_name(self):
        m = _lines()
        with pytest.raises(TypeError, match="NPV takes a single discount rate"):
            mo.NPV(m.x.pct / 100, m.x.amt)
        with pytest.raises(TypeError, match="NPV takes a single discount rate"):
            mo.NPV([0.05, 0.06, 0.07], [100.0, 200.0, 300.0])

    @pytest.mark.parametrize("call", ["IRR", "XIRR"])
    def test_a_series_guess_is_refused_by_name(self, call):
        m = _lines()
        with m:
            m.x.cf = mo.Variable([-100.0, 60.0, 70.0], display_name="cf")
            m.x.on = mo.Variable([date(2026, 1, 1), date(2027, 1, 1), date(2028, 1, 1)],
                                 display_name="on")
        args = (m.x.cf,) if call == "IRR" else (m.x.cf, m.x.on)
        with pytest.raises(TypeError, match=f"{call}'s guess is one number"):
            getattr(mo, call)(*args, m.x.pct / 100)


class TestTheWorkbook:
    def test_each_column_writes_its_period_s_cells(self, tmp_path):
        m = _model()
        path = tmp_path / "book.xlsx"
        m.to_excel(path)
        ws = load_workbook(path)["X"]
        x, at = _rows(ws), _row_of(ws)

        def each(text):
            return [text.format(c=c, **{k: f"{c}{r}" for k, r in at.items()})
                    for c in "BCD"]

        assert x["rnd"] == each("=ROUND({v}, {d})")
        assert x["rnd2"] == each("=ROUND(2.71828, {d})")
        assert x["ed"] == each("=EDATE({dt}, {k})")
        assert x["ed2"][1] == f"=EDATE(DATE(2026, 1, 31), C{at['k']})"
        assert x["eo"] == each("=EOMONTH({dt}, {k})")
        assert x["pm"] == each("=PMT({pct} / 100, 12, 1000, 0, 0)")
        assert x["pm2"] == each("=PMT(0.05, {n}, {amt}, 0, 1)")
        assert x["fv"] == each("=FV({pct} / 100, 12, -100, 0, 0)")
        assert x["pv"] == each("=PV(0.05, {n}, -100, {amt}, 0)")
        assert x["d360"] == each("=DAYS360({dt}, {end}, {eu})")


def _soffice():
    for candidate in ("soffice", "libreoffice"):
        found = shutil.which(candidate)
        if found:
            return found
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


@pytest.mark.slow
def test_the_workbook_recalculates_each_period_s_call(tmp_path):
    soffice = _soffice()
    if soffice is None:
        pytest.skip("LibreOffice not installed")
    m = _model()
    with m:
        m.x.rdg = mo.ROUND(m.x.v, m.t.dg).set_display_name("rdg")
    path = tmp_path / "series.xlsx"
    m.to_excel(path)
    subprocess.run(
        [soffice, f"-env:UserInstallation={(tmp_path / 'prof').as_uri()}",
         "--headless", "--convert-to", "xlsx", "--outdir", str(tmp_path / "calc"), str(path)],
        check=True, capture_output=True,
    )
    x = _rows(load_workbook(tmp_path / "calc" / "series.xlsx", data_only=True)["X"])
    for name in [*_CALLS, "rdg"]:
        want = getattr(m.x, name).value
        got = x[name][:3]
        for g, w in zip(got, want):
            if isinstance(w, date):
                assert isinstance(g, datetime) and g.date() == w, (name, got, want)
            else:
                assert g == pytest.approx(w), (name, got, want)
