# SPDX-License-Identifier: Apache-2.0
"""XIRR computes the rate the spreadsheet's XIRR computes.

The spreadsheet's XIRR finds the rate at which

    sum( CF_i / (1 + rate) ** ((d_i - d_0) / 365) ) = 0

where ``d_0`` is the FIRST date in the list: a year is 365 days, leap
years included, every date counts as its whole day (a time of day is
dropped), and every other date must fall on or after the first one
(they may come in any order). A 365.25-day year looks harmless but
moves the rate in the fourth decimal on an ordinary schedule, so the
number Python shows and the number in the written workbook disagree.
These tests pin the day count and recalculate the written workbook to
show the file computes what Python computed.
"""

from __future__ import annotations

import subprocess
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pytest

import modeleon as mo

# (cash flows, dates) — irregular gaps, longer than a year, a leap day,
# dates after the first one out of order, a losing project, times of
# day, and cash flows in the trillions and in thousandths.
SCHEDULES = {
    "Irregular": (
        [-100_000, 20_000, 30_000, 60_000],
        [date(2025, 1, 1), date(2025, 3, 15), date(2025, 7, 10), date(2025, 12, 31)],
    ),
    "EighteenMonths": (
        [-100_000, 20_000, 30_000, 45_000, 60_000],
        [date(2024, 1, 15), date(2024, 6, 30), date(2024, 12, 31),
         date(2025, 3, 1), date(2025, 7, 15)],
    ),
    "Unordered": (
        [-250_000, 40_000, 90_000, 70_000, 120_000],
        [date(2023, 6, 30), date(2025, 12, 31), date(2024, 2, 29),
         date(2024, 9, 30), date(2026, 6, 30)],
    ),
    "Loss": (
        [-100_000, 30_000, 25_000, 20_000],
        [date(2024, 3, 1), date(2025, 2, 28), date(2026, 3, 2), date(2027, 3, 1)],
    ),
    "TimeOfDay": (
        [-100_000, 60_000, 60_000],
        [datetime(2024, 1, 1, 18), datetime(2024, 7, 1, 6), datetime(2025, 1, 1, 6)],
    ),
    "Trillions": (
        [-3.7e13, 1.1e13, 1.3e13, 1.7e13],
        [date(2025, 1, 1), date(2025, 3, 15), date(2025, 7, 10), date(2025, 12, 31)],
    ),
    "Thousandths": (
        [-0.001, 0.0002, 0.0003, 0.0006],
        [date(2025, 1, 1), date(2025, 3, 15), date(2025, 7, 10), date(2025, 12, 31)],
    ),
}


def _excel_residual(rate, cash_flows, dates):
    """The spreadsheet's XIRR equation at ``rate`` — zero at the answer.

    Returned as a share of the cash flows' total size, so a schedule in
    trillions and one in thousandths are held to the same precision."""
    first = dates[0].toordinal()  # whole days: a time of day is dropped
    residual = sum(
        cf / (1 + rate) ** ((d.toordinal() - first) / 365)
        for cf, d in zip(cash_flows, dates)
    )
    return residual / sum(abs(cf) for cf in cash_flows)


@pytest.mark.parametrize("name", sorted(SCHEDULES))
def test_xirr_counts_a_year_as_365_days(name):
    cash_flows, dates = SCHEDULES[name]
    rate = mo.XIRR(cash_flows, dates).value
    assert _excel_residual(rate, cash_flows, dates) == pytest.approx(0, abs=1e-12)


def test_xirr_rejects_a_date_before_the_first():
    # The first date starts the schedule; a later entry dated before it
    # is an error in the spreadsheet, so Python refuses it too rather
    # than show a rate the workbook will not.
    with pytest.raises(ValueError, match="first date"):
        mo.XIRR(
            [-100_000, 30_000, 50_000, 60_000],
            [date(2024, 6, 1), date(2024, 1, 1), date(2025, 1, 1), date(2025, 6, 1)],
        )


def test_xirr_allows_flows_on_the_first_date():
    cash_flows = [-60_000, -40_000, 130_000]
    dates = [date(2024, 1, 1), date(2024, 1, 1), date(2025, 7, 1)]
    rate = mo.XIRR(cash_flows, dates).value
    assert _excel_residual(rate, cash_flows, dates) == pytest.approx(0, abs=1e-6)


# ─── The written workbook, recalculated ───────────────────────────


def _soffice() -> str | None:
    """Path to the LibreOffice command line, or None when not installed."""
    import shutil

    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


@pytest.fixture(scope="module")
def recalculated(tmp_path_factory):
    """Write one tab per schedule, let LibreOffice calculate the book, and
    return ``{tab: (python rate, workbook rate)}``."""
    soffice = _soffice()
    if soffice is None:
        pytest.skip("LibreOffice is not installed")
    tmp_path = tmp_path_factory.mktemp("xirr")

    model = mo.MultiVariable("Returns")
    python = {}
    for name, (cash_flows, dates) in SCHEDULES.items():
        tab = mo.MultiVariable(name, excel_props={'tab': True})
        setattr(model, name.lower(), tab)
        tab.cash_flow = mo.Variable(cash_flows)
        tab.paid_on = mo.Variable(dates)
        try:
            tab.xirr = mo.XIRR(tab.cash_flow, tab.paid_on)
        except (ValueError, TypeError, RuntimeError) as exc:
            python[name] = exc  # fails that schedule's test, not the others
        else:
            python[name] = tab.xirr.value

    src = tmp_path / "book.xlsx"
    model.to_excel(str(src))
    outdir = tmp_path / "calc"
    profile = (tmp_path / "lo-profile").as_uri()  # private: parallel runs don't clash
    subprocess.run(
        [soffice, f"-env:UserInstallation={profile}", "--headless",
         "--convert-to", "xlsx", "--outdir", str(outdir), str(src)],
        check=True, capture_output=True, timeout=120,
    )

    written = openpyxl.load_workbook(str(src))
    calculated = openpyxl.load_workbook(str(outdir / "book.xlsx"), data_only=True)
    result = {}
    for name in SCHEDULES:
        if isinstance(python[name], Exception):
            result[name] = (python[name], None)
            continue
        cells = [
            c.coordinate for row in written[name].iter_rows() for c in row
            if isinstance(c.value, str) and c.value.startswith("=XIRR(")
        ]
        assert len(cells) == 1, cells
        result[name] = (python[name], calculated[name][cells[0]].value)
    return result


@pytest.mark.parametrize("name", sorted(SCHEDULES))
def test_xirr_matches_the_recalculated_workbook(recalculated, name):
    python_rate, workbook_rate = recalculated[name]
    if isinstance(python_rate, Exception):
        raise python_rate
    assert isinstance(workbook_rate, float), workbook_rate
    assert python_rate == pytest.approx(workbook_rate, abs=1e-9)
