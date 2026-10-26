# SPDX-License-Identifier: Apache-2.0
"""The grain vocabulary is one list, and every entry point says no by name.

A word the engine does not know — a grain, a totals kind, a day label
that is really a week — must be refused where it is given, never
accepted and left to die (or lie) in a writer later.
"""
from datetime import date

import pytest

import modeleon as mo
from modeleon.compile.excel.header import spell_anchor, spell_inline
from modeleon.compile.excel.layout import LayoutEngine
from modeleon.compile.excel.view import resolve_excel_view
from modeleon.core.time import GRAINS, _parse


class TestDayLabels:
    def test_a_week_label_is_not_a_day_start(self):
        # date.fromisoformat reads '2026-W01' as the Monday it names; a
        # day window would then start on 2025-12-29 under a label that
        # never spells that day.
        with pytest.raises(ValueError, match="Time day start must be 'YYYY-MM-DD'"):
            mo.Model('m', default_grain='day', default_start='2026-W01', default_periods=3)

    @pytest.mark.parametrize("label", ["20260105", "2026-1-5", "2026-01-05T10:00", 20260105])
    def test_only_the_dashed_form_is_a_day_label(self, label):
        with pytest.raises(ValueError, match="Time day start must be 'YYYY-MM-DD'"):
            _parse(label, 'day')

    def test_a_dashed_day_label_still_parses(self):
        assert _parse('2026-01-05', 'day') == date(2026, 1, 5)
        m = mo.Model('m', default_grain='day', default_start='2026-01-05', default_periods=2)
        with m:
            m.x = mo.Variable([1.0, 2.0])
        assert m.x.time.start == '2026-01-05'


class TestGrainOnAssignment:
    def test_an_unknown_grain_is_refused_when_assigned(self):
        # The constructor refused it; plain assignment used to accept it
        # and the workbook writer died on it later, naming no attribute.
        m = mo.Model('m')
        with pytest.raises(ValueError, match="Unknown default_grain 'minute'; supported: hour, "):
            m.default_grain = 'minute'
        with pytest.raises(ValueError, match="Unknown default_grain 'fortnight'"):
            m.default_grain = 'fortnight'
        m.default_grain = 'hour'
        assert m.default_grain == 'hour'

    def test_a_known_grain_assigns(self):
        m = mo.Model('m')
        m.default_grain = 'quarter'
        assert m.default_grain == 'quarter'
        m.default_grain = None
        assert m.default_grain is None


class TestTotalsKinds:
    def test_a_view_refuses_a_totals_kind_it_cannot_draw(self):
        # An unknown kind used to be accepted and dropped: the model
        # declared the columns and the book showed none, without a word.
        with pytest.raises(ValueError, match="ExcelView timeline totals may name day, week, month, quarter, half, year, iso_year; got 'fortnight'"):
            mo.ExcelView(timeline={'totals': ['fortnight']})
        # Nothing is finer than an hour: it is never a total.
        with pytest.raises(ValueError, match="got 'hour'"):
            mo.ExcelView(timeline={'totals': ['quarter', 'hour']})

    def test_the_known_kinds_resolve_in_their_order(self):
        m = mo.Model('m', default_grain='month', default_start='2026-01', default_periods=6,
                     default_excel_view=mo.ExcelView(timeline={'totals': ['year', 'quarter']}))
        with m:
            m.x = mo.Variable([1.0] * 6, regrain=mo.up('sum'))
        assert resolve_excel_view(m).time_totals == ('year', 'quarter')

    def test_a_layout_override_refuses_an_unknown_kind(self):
        m = mo.Model('m', default_grain='month', default_start='2026-01', default_periods=3)
        with m:
            m.x = mo.Variable([1.0] * 3, regrain=mo.up('sum'))
        with pytest.raises(ValueError, match="A totals column is one of hour, day, week, month, quarter, half, year, iso_year; got 'fortnight'"):
            LayoutEngine([m], totals_override=('fortnight',))
        # The override's own vocabulary is the engine's: a caller may ask
        # for a kind the declaration cannot (a month column over days).
        LayoutEngine([m], totals_override=('month',))
        LayoutEngine([m], totals_override=())


class TestHeaderSpelling:
    def test_the_header_refuses_a_grain_it_cannot_step(self):
        # An unknown word used to be spelled as one month (STEP.get(grain,
        # 1)) — a silent 'C2-C2+1' in the days row.
        with pytest.raises(ValueError, match="got grain 'minute'"):
            spell_anchor('days', 1, 'minute', 'B2', lambda f, i: None)
        # An hour spells: its fraction of a day.
        assert spell_anchor('days', 1, 'hour', 'B2', lambda f, i: None) == '1/24'
        # Every ruler of the calendar is spelled; a word it does not know
        # is refused by name, the spelled rulers listed.
        with pytest.raises(ValueError, match="spells periods of hour, day, week, month, quarter, half, year or iso_year; got grain 'minute'"):
            spell_inline('end', 1, 'minute', 'B2')
        assert spell_inline('end', 0, 'iso_year', 'B2').startswith('((DATE(YEAR(B2+3)+1,1,4)')

    def test_the_known_grains_still_spell(self):
        assert spell_inline('end', 1, 'quarter', 'B2') == 'EOMONTH(B2,5)'
        assert spell_inline('end', 0, 'day', 'B2') == 'B2'

    def test_a_day_windows_anchor_spells_its_index_row(self, tmp_path):
        # The index chain needs no month stride; refusing a day grain
        # there would break every day model with a header that
        # counts its periods.
        cell = lambda f, i: f"I{i + 3}" if i >= 0 else None  # noqa: E731
        assert spell_anchor('index', 1, 'day', 'B2', cell) == 'I3+1'
        assert spell_anchor('index', 0, 'day', 'B2', lambda f, i: None) == '1'
        m = mo.Model('m', default_grain='day', default_start='2026-01-05', default_periods=3)
        with m:
            m.header = mo.MultiVariable('Time')
            with m.header as h:
                h.start = mo.time.start
                h.index = mo.time.index
            m.default_header = m.header
            m.x = mo.Variable([1.0, 2.0, 3.0]) * mo.time.index
        path = tmp_path / 'day.xlsx'
        m.to_excel(path)
        from openpyxl import load_workbook
        ws = load_workbook(path)['Time']
        row = next(r for r in ws.iter_rows() if r[0].value == 'Index')
        assert [c.value for c in row[1:4]] == ['=1', '=B3+1', '=C3+1']
        assert m.x.value == [1.0, 4.0, 9.0]


def test_the_vocabulary_is_one_tuple():
    assert GRAINS == ('hour', 'day', 'week', 'month', 'quarter', 'half', 'year', 'iso_year')
