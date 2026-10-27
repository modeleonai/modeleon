# SPDX-License-Identifier: Apache-2.0
"""Hours: the wall clock of one fixed-offset zone.

An hour is ``[13:00, 14:00)``, keyed ``2026-10-21T13``. Its instants are
the clock of the window's zone (``default_zone='+01:00'``), where every
day has 24 hours — a zone that moves its clocks is refused where it is
declared. ``mo.time.end`` is the hour's stop, ``mo.time.days`` its 1/24
of a day; an hour nests in its day, and so in everything a day nests in.
A sheet holds 16,384 columns: an hourly window that needs more is
refused with the count and the two ways out, never cut.
"""
import warnings
from datetime import date, datetime, timezone

import openpyxl
import pytest

import modeleon as mo
from modeleon.compile.excel import header as H
from modeleon.compile.excel.layout import LayoutEngine
from modeleon.compile.excel.renderer import _value_to_excel_literal
from modeleon.compile.excel.totals import build_totals_plan, header_rows
from modeleon.core.calendar import GREGORIAN, HOUR
from modeleon.core.time import (
    SheetWidthError, _period_labels, bucket_ranges, period_bounds, period_weights,
    time_field_series, window_start_label, zone_offset,
)
from modeleon.core.variable_ops import _to_serial

from test_excel_time_header import _recalculated


def _hourly(n=30, start='2026-10-21T22', totals=None, zone='+01:00'):
    kw = {'default_excel_view': mo.ExcelView(timeline={'totals': totals})} if totals else {}
    m = mo.Model('hr', default_grain='hour', default_start=start, default_periods=n,
                 default_zone=zone, **kw)
    with m:
        m.header = mo.MultiVariable('Time')
        with m.header as h:
            h.start = mo.time.start
            h.end = mo.time.end
            h.days = mo.time.days
            h.index = mo.time.index
        m.default_header = m.header
        m.ops = mo.MultiVariable('Ops')
        with m.ops as o:
            o.load = mo.Variable([float(i % 5) for i in range(n)], regrain=mo.up('sum'),
                                 display_name='Load')
            o.rate = mo.Variable(120.0, display_name='Rate per day')
            o.cost = (o.rate * mo.time.days).set_display_name('Cost')
    return m


def _spans(n=30):
    """Time read inside formulas on an hourly sheet with no Time header."""
    m = mo.Model('sp', default_grain='hour', default_start='2026-10-21T22', default_periods=n,
                 default_zone='+01:00', default_excel_view=mo.ExcelView(timeline={'totals': ['day']}))
    with m:
        m.ops = mo.MultiVariable('Ops')
        with m.ops as o:
            o.span = ((mo.time.end - mo.time.start) * 24).set_display_name('Span')
            o.twice = (mo.time.start * 2).set_display_name('Twice start')
    return m


def _book(m, path):
    out = path / 'hours.xlsx'
    with warnings.catch_warnings():
        warnings.simplefilter('ignore')
        m.to_excel(str(out))
    return openpyxl.load_workbook(out)


def _row(ws, label):
    return next(r for r in ws.iter_rows() if r[0].value == label)


class TestTheRuler:
    def test_an_hour_is_keyed_and_read_back_strictly(self):
        first = HOUR.parse('2026-10-21T13')
        assert first == datetime(2026, 10, 21, 13)
        assert HOUR.label(first) == '2026-10-21T13'
        assert HOUR.label(first, 'finance') == '21 Oct 2026 13:00'
        assert HOUR.label(first, 'compact') == '21-Oct-26 13:00'
        for bad in ('2026-10-21', '2026-10-21T24', '2026-10-21T13:30', '2026-10-21 13'):
            with pytest.raises(ValueError, match="'YYYY-MM-DDTHH'|00-23"):
                HOUR.parse(bad)

    def test_an_hour_stops_where_the_next_starts(self):
        assert HOUR.next_start(datetime(2026, 12, 31, 23)) == datetime(2027, 1, 1, 0)
        assert HOUR.start_of(datetime(2026, 10, 21, 13, 59, 59)) == datetime(2026, 10, 21, 13)
        assert HOUR.start_of(date(2026, 10, 21)) == datetime(2026, 10, 21, 0)
        assert HOUR.measure(datetime(2026, 10, 21, 13)) == 1 / 24
        assert HOUR.count(datetime(2026, 10, 21, 22), datetime(2026, 10, 23, 3)) == 30

    def test_a_fixed_offset_day_has_24_hours_even_when_europe_moves_its_clocks(self):
        # 29 March 2026: clocks move in Europe. The wall clock of a fixed
        # offset never skips an hour.
        assert len(bucket_ranges('hour', 'day', '2026-03-29T00', 24)) == 1
        assert _period_labels('2026-03-29T01', None, 3, 'hour') == [
            '2026-03-29T01', '2026-03-29T02', '2026-03-29T03']

    def test_an_hour_nests_in_its_day_and_everything_above(self):
        assert GREGORIAN.rolls_up_to('hour') == (
            'day', 'week', 'month', 'quarter', 'half', 'year', 'iso_year')
        GREGORIAN.verify()


class TestTheWindow:
    def test_a_window_starts_on_the_hour(self):
        assert window_start_label(datetime(2026, 10, 21, 13), 'hour') == '2026-10-21T13'
        assert window_start_label(date(2026, 10, 21), 'hour') == '2026-10-21T00'
        with pytest.raises(ValueError, match="starts on the hour; 2026-10-21T13:30:00"):
            window_start_label(datetime(2026, 10, 21, 13, 30), 'hour')
        with pytest.raises(ValueError, match="wall-clock time of its zone"):
            window_start_label(datetime(2026, 10, 21, 13, tzinfo=timezone.utc), 'hour')
        # The other rulers still refuse a clock.
        with pytest.raises(ValueError, match="starts at a day, not at a time of day"):
            window_start_label(datetime(2026, 10, 21, 13), 'day')

    def test_the_zone_is_a_fixed_offset(self):
        assert zone_offset('UTC') == zone_offset('Z') == '+00:00'
        assert zone_offset('+0100') == '+01:00' and zone_offset('-03:30') == '-03:30'
        with pytest.raises(ValueError, match="fixed offset from UTC.*Asia/Tokyo is \\+09:00"):
            zone_offset('Asia/Tokyo')
        with pytest.raises(ValueError, match="fixed offset"):
            zone_offset('+01:20')
        m = mo.Model('m', default_grain='hour', default_start='2026-10-21T00', default_periods=2,
                      default_zone='UTC')
        assert m.default_zone == '+00:00'
        m.default_zone = '+01:00'
        assert m.default_zone == '+01:00'
        with pytest.raises(ValueError, match="fixed offset"):
            m.default_zone = 'Europe/Berlin'

    def test_mo_time_of_an_hour(self):
        assert time_field_series('start', '2026-10-21T23', 'hour', 2) == [
            datetime(2026, 10, 21, 23), datetime(2026, 10, 22, 0)]
        assert time_field_series('end', '2026-10-21T23', 'hour', 2) == [
            datetime(2026, 10, 22, 0), datetime(2026, 10, 22, 1)]
        assert time_field_series('days', '2026-10-21T23', 'hour', 2) == [1 / 24, 1 / 24]
        assert period_bounds('2026-10-21T23', 'hour', 1) == (
            [datetime(2026, 10, 21, 23)], [datetime(2026, 10, 22, 0)])
        m = _hourly(n=3)
        assert m.header.days.value_type == 'float'
        assert m.ops.cost.value == pytest.approx([5.0, 5.0, 5.0])

    def test_an_hour_has_no_months(self):
        m = mo.Model('m', default_grain='hour', default_start='2026-10-21T00', default_periods=2)
        with m, pytest.raises(ValueError, match="An hour window has none - read mo.time.days"):
            m.months = mo.time.months


class TestRollUps:
    def test_hours_roll_up_to_days_by_sum_last_and_weighted_mean(self):
        m = mo.Model('m', default_grain='hour', default_start='2026-10-21T22', default_periods=5)
        with m:
            m.flow = mo.Variable([1.0, 2.0, 3.0, 4.0, 5.0], regrain=mo.up('sum'))
            m.level = mo.Variable([1.0, 2.0, 3.0, 4.0, 5.0], regrain=mo.up('last'))
            m.price = mo.Variable([1.0, 2.0, 3.0, 4.0, 5.0], regrain=mo.up('mean'))
        assert m.flow.at('day').value == [3.0, 12.0]
        assert m.level.at('day').value == [2.0, 5.0]
        assert m.price.at('day').value == pytest.approx([1.5, 4.0])
        assert m.flow.at('week').value == [15.0] and m.flow.at('month').value == [15.0]
        assert period_weights('hour', '2026-10-21T22', 2) == [1 / 24, 1 / 24]

    def test_a_day_line_is_never_viewed_by_hours(self):
        with pytest.raises(ValueError, match="coarsens only: day -> hour"):
            bucket_ranges('day', 'hour', '2026-10-21', 2)

    def test_an_hourly_line_and_a_daily_line_meet_only_through_at(self):
        h = mo.Model('h', default_grain='hour', default_start='2026-10-21T00', default_periods=2)
        with h:
            h.a = mo.Variable([1.0, 2.0], regrain=mo.up('sum'))
        d = mo.Model('d', default_grain='day', default_start='2026-10-21', default_periods=1)
        with d:
            d.b = mo.Variable([10.0], regrain=mo.up('sum'))
        with pytest.raises(ValueError, match=r"\(hours\) and 'b' \(days\).*a.at\('day'\)"):
            h.a + d.b


class TestTheBook:
    def test_an_hourly_header_is_a_closed_form_off_the_start_cell(self, tmp_path):
        ws = _book(_hourly(n=4), tmp_path)['Time']
        start, end, days, index = (_row(ws, k) for k in ('Start', 'End', 'Days', 'Index'))
        assert [c.value for c in start[1:4]] == [
            '=DATE(2026,10,21)+TIME(22,0,0)',
            '=DATE(2026,10,21)+TIME(22,0,0)+1/24',
            '=DATE(2026,10,21)+TIME(22,0,0)+2/24',
        ]
        assert [c.value for c in end[1:3]] == ['=B2+1/24', '=C2+1/24']
        assert [c.value for c in days[1:3]] == ['=B3-B2', '=C3-C2']
        assert days[1].number_format == '0.0000'
        assert [c.value for c in index[1:3]] == ['=1', '=B5+1']

    def test_captions_show_the_hour_itself(self, tmp_path):
        ws = _book(_hourly(n=3), tmp_path)['Ops']
        # A flat sheet's caption is the hour's start cell — its own
        # mirrored start row — shown as a clock.
        assert ws['A2'].value == 'Start' and ws['C1'].value == '=C2'
        assert ws['C1'].number_format == 'dd mmm yyyy hh:mm'

    def test_day_totals_are_live_sums_of_their_hours(self, tmp_path):
        wb = _book(_hourly(n=30, totals=['day']), tmp_path)
        ws = wb['Ops']
        load = _row(ws, 'Load')
        sums = [c.value for c in load if isinstance(c.value, str) and c.value.startswith('=SUM')]
        assert sums[0] == '=SUM(C8,D8)'                     # 22:00 and 23:00 of the 21st
        assert sums[1].count(',') == 23                      # all 24 hours of the 22nd
        assert len(sums) == 3
        # The day reads in full over hours alone; the hours keep their clock.
        assert ws['C1'].value == '21 Oct 2026' and ws['C2'].value == '22:00'
        assert ws['E1'].value == 'Total 21 Oct 2026'

    def test_a_days_time_in_its_total_is_read_off_its_hours(self, tmp_path):
        ws = _book(_hourly(n=30, totals=['day']), tmp_path)['Time']
        totals = [c for c in range(1, ws.max_column + 1)
                  if str(ws.cell(1, c).value or '').startswith('Total')]
        start, end, days = (_row(ws, k) for k in ('Start', 'End', 'Days'))
        # The 21st holds 22:00 and 23:00: its total starts at the first
        # hour's start, stops at the last hour's stop, lasts their days.
        # (Two header rows over the time rows: the day tier, the hours.)
        assert start[totals[0] - 1].value == '=B3'
        assert end[totals[0] - 1].value == '=C4'
        assert days[totals[0] - 1].value == '=SUM(B5,C5)'
        assert days[totals[1] - 1].value.count(',') == 23
        # Each keeps what the engine computed, for a reader that does not
        # recalculate.
        cached = openpyxl.load_workbook(tmp_path / 'hours.xlsx', data_only=True)['Time']
        col = openpyxl.utils.get_column_letter(totals[0])
        assert cached[f'{col}3'].value == datetime(2026, 10, 21, 22)
        assert cached[f'{col}4'].value == datetime(2026, 10, 22)
        assert cached[f'{col}5'].value == pytest.approx(2 / 24)

    def test_an_instant_inside_a_formula_stands_in_brackets(self, tmp_path):
        ws = _book(_spans(), tmp_path)['Ops']
        span, twice = _row(ws, 'Span'), _row(ws, 'Twice start')
        # 22:00 is DATE()+TIME(), a sum: wherever a formula reads it — the
        # first hour, a day's total — it stands in brackets, or Excel reads
        # x - start as (x - DATE()) + TIME().
        assert span[1].value == ('=((DATE(2026,10,21)+TIME(22,0,0)+1/24)'
                                 ' - (DATE(2026,10,21)+TIME(22,0,0))) * 24')
        assert twice[1].value == '=(DATE(2026,10,21)+TIME(22,0,0)) * 2'
        assert span[3].value == '=(DATE(2026, 10, 22) - (DATE(2026, 10, 21)+TIME(22, 0, 0))) * 24'
        # a formula of the period's date with no re-grain rule has no day
        # total — a duration (end minus start) does
        assert twice[3].value == '#VALUE!'

    def test_a_month_tier_over_days_over_hours(self):
        plan = build_totals_plan(('day', 'month'), 'hour', '2026-10-31T22', 4, 3)
        rows = header_rows(plan, _period_labels('2026-10-31T22', None, 4, 'hour', 'finance'))
        tiers = [[c['label'] for c in r['cells']] for r in rows]
        assert tiers[0][0] == 'Oct' and tiers[1][0] == '31'
        assert [c['label'] for c in rows[-1]['cells'] if c['kind'] == 'period'] == [
            '22:00', '23:00', '00:00', '01:00']

    def test_the_opening_column_is_where_the_first_hour_starts(self):
        assert H.spell_anchor('start', 0, 'hour', 'B2', lambda f, i: 'O9' if i < 0 else None) == 'B2'
        from modeleon.compile.excel.writer import opening_value
        m = _hourly(n=2)
        assert opening_value(m.header.end) == datetime(2026, 10, 21, 22)

    def test_an_instant_literal_keeps_its_clock(self):
        assert _value_to_excel_literal(datetime(2026, 1, 5, 13)) == '(DATE(2026, 1, 5)+TIME(13, 0, 0))'
        assert _value_to_excel_literal(datetime(2026, 1, 5)) == 'DATE(2026, 1, 5)'
        assert H.date_literal(datetime(2026, 1, 5, 13)) == 'DATE(2026,1,5)+TIME(13,0,0)'
        assert H.date_literal(date(2026, 1, 5)) == 'DATE(2026,1,5)'

    @pytest.mark.slow
    def test_the_hourly_book_recalculates_to_the_engine(self, tmp_path):
        m = _hourly(n=30, totals=['day'])
        wb = _recalculated(m, tmp_path)
        for sheet, mv, names in (('Time', m.header, ('start', 'end', 'days')),
                                 ('Ops', m.ops, ('cost',))):
            ws = wb[sheet]
            caption = next(r for r in ws.iter_rows(max_row=4)
                           if any(isinstance(c.value, str) and c.value.endswith(':00') for c in r))
            hours = [c.column for c in caption if isinstance(c.value, str) and c.value.endswith(':00')]
            assert len(hours) == 30
            rows = {r[0].value: r for r in ws.iter_rows() if r[0].value is not None}
            for name in names:
                var = getattr(mv, name)
                got = [round(_to_serial(rows[var.display_name][c - 1].value), 6) for c in hours]
                want = [round(_to_serial(v), 6) for v in var.value]
                assert got == pytest.approx(want), (sheet, name, got[:3], want[:3])
        # The Time sheet's day totals recalculate to the hours they hold.
        t = wb['Time']
        tot = [c for c in range(1, t.max_column + 1) if str(t.cell(1, c).value or '').startswith('Total')]
        days_row = next(r for r in t.iter_rows() if r[0].value == 'Days')
        assert [round(days_row[c - 1].value, 6) for c in tot] == [round(2 / 24, 6), 1.0, round(4 / 24, 6)]
        # A day's total of the cost row is its hours' cost: 2, 24 and 4 hours.
        ws = wb['Ops']
        totals = [c for c in _row(ws, 'Cost') if c.column in
                  [x.column for x in ws[1] if isinstance(x.value, str) and x.value.startswith('Total')]]
        assert [round(c.value, 6) for c in totals] == [10.0, 120.0, 20.0]

    @pytest.mark.slow
    def test_an_instant_inside_a_formula_recalculates_to_the_engine(self, tmp_path):
        m = _spans()
        ws = _recalculated(m, tmp_path)['Ops']
        hours = [c.column for c in ws[2] if isinstance(c.value, str) and c.value.endswith(':00')]
        totals = [c.column for c in ws[1] if isinstance(c.value, str) and c.value.startswith('Total')]
        assert len(hours) == 30 and len(totals) == 3
        span, twice = _row(ws, 'Span'), _row(ws, 'Twice start')
        assert [span[c - 1].value for c in hours] == pytest.approx([1.0] * 30)
        assert [span[c - 1].value for c in totals] == pytest.approx([2.0, 24.0, 4.0])
        want = [round(_to_serial(v), 6) for v in m.ops.twice.value]
        assert [round(twice[c - 1].value, 6) for c in hours] == pytest.approx(want)


class TestTheWidthLaw:
    def test_a_window_wider_than_a_sheet_is_refused_where_it_is_declared(self):
        with pytest.raises(SheetWidthError, match="A window of 20000 hours needs 20000 columns.*"
                                                  "a sheet holds 16384; hours fit about 22 months. "
                                                  "View it by day"):
            mo.Model('w', default_grain='hour', default_start='2026-01-01T00', default_periods=20000)
        m = mo.Model('w', default_grain='hour', default_start='2026-01-01T00', default_periods=2)
        with pytest.raises(SheetWidthError, match="20000 hours"):
            m.default_periods = 20000

    def test_totals_count_toward_the_width(self):
        m = mo.Model('big', default_grain='hour', default_start='2026-01-01T00', default_periods=16100,
                     default_excel_view=mo.ExcelView(timeline={'totals': ['day']}))
        with m:
            m.v = mo.Variable([1.0] * 16100, regrain=mo.up('sum'))
        with warnings.catch_warnings():
            warnings.simplefilter('ignore')
            with pytest.raises(SheetWidthError, match="16100 hours needs 16771 columns with its totals"):
                LayoutEngine([m]).compute_addresses()
        fits = mo.Model('fits', default_grain='hour', default_start='2026-01-01T00', default_periods=15000)
        with fits:
            fits.v = mo.Variable([1.0] * 15000, regrain=mo.up('sum'))
        LayoutEngine([fits]).compute_addresses()


def test_growth_by_the_hour_a_year_back():
    from modeleon.compile.excel.writer import _growth_by_calendar
    from modeleon.core.time import DefaultWindow
    n = 8761
    child = mo.Variable([float(i + 1) for i in range(n)])
    out = _growth_by_calendar(child, DefaultWindow('2025-01-01T00', 'hour', n), 'hour', n, child.value)
    # 2026-01-01T00 (index 8760) grows on the same hour a year back,
    # 2025-01-01T00 (index 0).
    assert out.value[8760] == 8761.0 / 1.0
    assert all(v is None for v in out.value[:8760])
