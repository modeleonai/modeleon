"""ExcelView tracks= dict form — the declared default track selection.

The string form stays the emission mode; the dict form
``tracks={'mode': ..., 'shown': [...]}`` additionally declares which
tracks a renderer shows by default (presentation metadata — emission ignores
it, other renderers read it from the resolved view).
"""

import pytest

import modeleon as mo
from modeleon.compile.excel.view import ExcelView, resolve_excel_view


def _model_with_view(view):
    m = mo.Model("m", tracks=mo.Tracks("plan", "fact"),
                 default_grain="month", default_start="2026-01",
                 default_periods=3)
    m.default_excel_view = view
    m.pnl = mo.MultiVariable(display_name="PnL")
    m.pnl.revenue = mo.Variable(plan=[100.0] * 3, fact=[110.0, None, None])
    return m


def test_dict_form_resolves_mode_and_shown():
    m = _model_with_view(ExcelView(
        tracks={"mode": "blend", "shown": ["fact", "plan"]}
    ))
    rv = resolve_excel_view(m.pnl.revenue)
    assert rv.tracks == "blend"
    assert rv.tracks_shown == ("fact", "plan")


def test_string_form_unchanged_and_shown_stays_none():
    m = _model_with_view(ExcelView(tracks="blend"))
    rv = resolve_excel_view(m.pnl.revenue)
    assert rv.tracks == "blend"
    assert rv.tracks_shown is None


def test_dict_form_validates_mode_and_keys():
    with pytest.raises(ValueError, match="mode must be one of"):
        ExcelView(tracks={"mode": "sideways"})
    with pytest.raises(ValueError, match="accepts 'mode' and 'shown'"):
        ExcelView(tracks={"shown": ["plan"], "extra": 1})
    with pytest.raises(ValueError, match="list of track names"):
        ExcelView(tracks={"shown": "plan"})


def test_dict_form_emits_like_its_mode(tmp_path):
    # 'shown' is presentation metadata: the workbook is byte-equivalent to the
    # plain mode string — blend prints the single display-default row.
    m = _model_with_view(ExcelView(
        tracks={"mode": "blend", "shown": ["fact", "plan"]}
    ))
    path = tmp_path / "t.xlsx"
    m.to_excel(str(path))
    from openpyxl import load_workbook
    ws = load_workbook(str(path)).active
    labels = [c.value for row in ws.iter_rows() for c in row
              if isinstance(c.value, str)]
    # blend mode: no per-track labeled rows
    assert not any("· plan" in s or "· fact" in s for s in labels)
