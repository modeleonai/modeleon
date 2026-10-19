# SPDX-License-Identifier: Apache-2.0
"""A caption that stands in a value column heads no row of the sheet,
so it takes no lead cells.

The lead columns — the article number, the unit, a projected field —
belong to a ROW: they are written beside its label. A scalar across an
``orient='columns'`` board and a field of an ``orient='records'`` table
have a caption in a value column instead, and the writer placed their
unit one column right of it, over the neighbour's caption: ``IRR``'s
``%`` vanished under ``MOIC``, and ``MOIC``'s ``x`` stood alone past the
board. Such a caption now writes nothing beside itself: on a sheet
without a unit column it carries its unit the way every caption does —
``IRR (%)`` — and on a sheet with one its unit stands in the unit row
under the captions (``test_excel_caption_unit_row``). A line whose label
stands in the label column keeps its lead cells.
"""

from openpyxl import load_workbook

import modeleon as mo


def _rows(model, sheet, tmp_path, n, cols="ABCDE"):
    path = tmp_path / "t.xlsx"
    model.to_excel(str(path))
    ws = load_workbook(str(path))[sheet]
    return [[ws[f"{c}{r}"].value for c in cols] for r in range(1, n + 1)]


def _kpis(unit_column):
    m = mo.Model("m", display_name="M")
    if unit_column:
        m.default_excel_view = mo.ExcelView(meta={"unit": "Unit"})
    m.kpis = mo.MultiVariable("KPIs", excel_props={"tab": True, "orient": "columns"})
    m.kpis.irr = mo.Variable(0.2, display_name="IRR", unit="%")
    m.kpis.moic = mo.Variable(2.5, display_name="MOIC", unit="x")
    return m


class Comp(mo.MultiVariableClass):
    def compute(self, ev, ebitda):
        self.ev = mo.Variable(ev, display_name="EV", unit="USD m",
                              excel_props={"article": "1.1"})
        self.ebitda = mo.Variable(ebitda, display_name="EBITDA", unit="USD m")


def _comps(meta):
    m = mo.Model("m", display_name="M")
    m.default_excel_view = mo.ExcelView(meta=meta)
    with m:
        m.comps = mo.MultiVariable("Comps", excel_props={"tab": True, "orient": "records"})
        with m.comps as c:
            c.a = Comp(ev=1200, ebitda=100, display_name="Alpha")
            c.b = Comp(ev=900, ebitda=90, display_name="Beta")
    return m


def test_a_scalar_across_a_board_carries_its_unit_in_the_unit_row(tmp_path):
    rows = _rows(_kpis(unit_column=True), "KPIs", tmp_path, 4)
    captions = next(r for r in rows if any(isinstance(v, str) and "IRR" in v for v in r))
    first = next(i for i, v in enumerate(captions) if v)
    assert captions[first:] == ["IRR", "MOIC", None]
    assert rows[rows.index(captions) + 1][first:] == ["%", "x", None]
    assert rows[rows.index(captions) + 2][first:] == [0.2, 2.5, None]


def test_without_a_unit_column_the_caption_reads_as_before(tmp_path):
    rows = _rows(_kpis(unit_column=False), "KPIs", tmp_path, 3)
    captions = next(r for r in rows if any(isinstance(v, str) and "IRR" in v for v in r))
    first = next(i for i, v in enumerate(captions) if v)
    assert captions[first:first + 2] == ["IRR (%)", "MOIC (x)"]


def test_a_records_field_carries_its_unit_in_the_unit_row(tmp_path):
    rows = _rows(_comps({"unit": "Unit"}), "Comps", tmp_path, 4)
    assert rows[0] == [None, None, "EV", "EBITDA", None]
    assert rows[1] == [None, None, "USD m", "USD m", None]
    assert rows[2] == ["Alpha", None, 1200, 100, None]
    assert rows[3] == ["Beta", None, 900, 90, None]


def test_a_records_field_writes_no_article_or_field_beside_its_caption(tmp_path):
    meta = {"article": "No.", "unit": "Unit", "fields": [("description", "Note")]}
    rows = _rows(_comps(meta), "Comps", tmp_path, 3, cols="ABCDEFG")
    header, units = rows[0], rows[1]
    # The article column, the unit column and the field column of the
    # header row and the unit row hold no field's article, unit or
    # description.
    assert header[:3] == [None, None, None] and units[:3] == [None, None, None]
    assert [v for v in header if v] == ["EV", "EBITDA"]
    assert [v for v in units if v] == ["USD m", "USD m"]


def test_a_board_field_in_the_label_column_keeps_its_lead_cells(tmp_path):
    m = mo.Model("m", display_name="M")
    m.default_excel_view = mo.ExcelView(meta={"unit": "Unit"})
    m.board = mo.MultiVariable("Board", excel_props={"tab": True, "orient": "columns"})
    m.board.p1 = mo.MultiVariable("P1")
    m.board.p1.cost = mo.Variable(50.0, display_name="Cost", unit="USD")
    m.board.p2 = mo.MultiVariable("P2")
    m.board.p2.cost = mo.Variable(80.0, display_name="Cost", unit="USD")
    rows = _rows(m, "Board", tmp_path, 2, cols="ABCD")
    assert rows[1] == ["Cost", "USD", 50.0, 80.0]
