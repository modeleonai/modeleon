# SPDX-License-Identifier: Apache-2.0
"""Conditional aggregates in the workbook.

A row is its cells' range (``Data!B2:D2``); a list of scalar lines whose
cells stand one under another is a vertical range (``Deals!B3:B5``) —
the field down the records of a table, or one period's column of their
blocks (``Deals!D8:D10``); cells that form no run cannot be one range,
so the cell carries its computed value and ``check_compat`` lists it.
The criteria is written as the spreadsheet writes it: ``">=10"``, a
cell, ``">=" & Inputs!B1``. ``MAXIFS`` / ``MINIFS`` are stored with the
``_xlfn.`` prefix. Every function recalculates in LibreOffice to the
engine's numbers, except where LibreOffice departs from Excel (a
boolean cell, numeric text, a wildcard over a number), pinned as an
expected failure.
"""

import datetime as dt
import shutil
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook

import modeleon as mo
from modeleon.compile import check_compat
from modeleon.compile.excel.renderer import _contiguous_run
from modeleon.functions.criteria import (
    AVERAGEIF,
    AVERAGEIFS,
    COUNT,
    COUNTA,
    COUNTIF,
    COUNTIFS,
    MAXIFS,
    MINIFS,
    SUMIF,
    SUMIFS,
)


def _rows(ws):
    out = {}
    for row in ws.iter_rows():
        label = row[0].value
        if label:
            out[str(label).split(" (")[0]] = [c.value for c in row[1:]]
    return out


def _book(model, tmp_path):
    path = tmp_path / "book.xlsx"
    model.to_excel(path)
    return load_workbook(path)


class Deal(mo.MultiVariableClass):
    def compute(self, region, amount):
        self.region = mo.Variable(region, display_name="Region")
        self.amount = mo.Variable(amount, display_name="Amount", unit="USD")
        self.flow = mo.Variable([amount * 0.1, amount * 0.2, amount * 0.3], display_name="Flow")


class Close(mo.MultiVariableClass):
    def compute(self, region, closed):
        self.region = mo.Variable(region, display_name="Region")
        self.closed = mo.Variable(closed, display_name="Closed")


def _model():
    m = mo.Model("m", display_name="M", default_grain="month", default_start="2026-01",
                 default_periods=3)
    with m:
        m.inputs = mo.MultiVariable("Inputs", excel_props={"tab": True})
        m.inputs.floor = mo.Variable(150.0, display_name="Floor")
        m.inputs.pick = mo.Variable("north", display_name="Pick")
        m.inputs.rule = mo.Variable("<>100", display_name="Rule")
        m.data = mo.MultiVariable("Data", excel_props={"tab": True})
        m.data.region = mo.Variable(["North", "South", "Northeast"], display_name="Region")
        m.data.amount = mo.Variable([100.0, 200.0, 300.0], display_name="Amount", unit="USD")
        m.data.year = mo.Variable([2026, 2026, 2027], display_name="Year")
        m.deals = mo.MultiVariable("Deals", excel_props={"tab": True, "orient": "records"})
        with m.deals as d:
            d.a = Deal(region="North", amount=100.0, display_name="Alpha")
            d.b = Deal(region="South", amount=200.0, display_name="Beta")
            d.c = Deal(region="North", amount=300.0, display_name="Gamma")
        m.closes = mo.MultiVariable("Closes", excel_props={"tab": True, "orient": "records"})
        with m.closes as c:
            c.a = Close(region="North", closed=dt.date(2026, 3, 31), display_name="A")
            c.b = Close(region="South", closed=dt.date(2026, 6, 30), display_name="B")
            c.c = Close(region="North", closed=dt.date(2024, 2, 29), display_name="C")
        m.s = mo.MultiVariable("Summary", excel_props={"tab": True})
    return m


# ══════════════════════════════════════════════════════════════════════
# Rows are ranges, the criteria written as the spreadsheet writes it
# ══════════════════════════════════════════════════════════════════════


class TestRowsAreRanges:
    def test_every_function_over_rows(self, tmp_path):
        m = _model()
        d = m.data
        with m.s as s:
            s.sumif = SUMIF(d.region, "North", d.amount).set_display_name("sumif")
            s.sumif_self = SUMIF(d.amount, ">=150").set_display_name("sumif_self")
            s.sumifs = SUMIFS(d.amount, d.region, "North", d.year, 2026).set_display_name("sumifs")
            s.countif = COUNTIF(d.region, "n*").set_display_name("countif")
            s.countifs = COUNTIFS(d.region, "<>South", d.amount, ">100").set_display_name("countifs")
            s.averageif = AVERAGEIF(d.region, "N*", d.amount).set_display_name("averageif")
            s.averageifs = AVERAGEIFS(d.amount, d.year, 2026).set_display_name("averageifs")
            s.maxifs = MAXIFS(d.amount, d.region, "N*").set_display_name("maxifs")
            s.minifs = MINIFS(d.amount, d.region, "N*").set_display_name("minifs")
            s.count = (COUNT(d.amount) + COUNT(d.region)).set_display_name("count")
            s.counta = COUNTA(d.region).set_display_name("counta")
            # Two rows are counted period by period: one call per
            # column, after the scalars' column.
            s.count_pp = COUNT(d.amount, d.region).set_display_name("count_pp")
        rows = _rows(_book(m, tmp_path)["Summary"])
        assert rows["sumif"][0] == '=SUMIF(Data!B2:D2, "North", Data!B3:D3)'
        assert rows["sumif_self"][0] == '=SUMIF(Data!B3:D3, ">=150")'
        assert rows["sumifs"][0] == '=SUMIFS(Data!B3:D3, Data!B2:D2, "North", Data!B4:D4, 2026)'
        assert rows["countif"][0] == '=COUNTIF(Data!B2:D2, "n*")'
        assert rows["countifs"][0] == '=COUNTIFS(Data!B2:D2, "<>South", Data!B3:D3, ">100")'
        assert rows["averageif"][0] == '=AVERAGEIF(Data!B2:D2, "N*", Data!B3:D3)'
        assert rows["averageifs"][0] == '=AVERAGEIFS(Data!B3:D3, Data!B4:D4, 2026)'
        assert rows["maxifs"][0] == '=_xlfn.MAXIFS(Data!B3:D3, Data!B2:D2, "N*")'
        assert rows["minifs"][0] == '=_xlfn.MINIFS(Data!B3:D3, Data!B2:D2, "N*")'
        assert rows["count"][0] == "=COUNT(Data!B3:D3) + COUNT(Data!B2:D2)"
        assert rows["counta"][0] == "=COUNTA(Data!B2:D2)"
        assert rows["count_pp"][1:4] == ["=COUNT(Data!B3, Data!B2)", "=COUNT(Data!C3, Data!C2)",
                                         "=COUNT(Data!D3, Data!D2)"]
        assert m.s.count_pp.value == [1, 1, 1]
        # The values the cells cache are the engine's.
        assert (m.s.sumif.value, m.s.sumif_self.value, m.s.sumifs.value) == (100.0, 500.0, 100.0)
        assert (m.s.countif.value, m.s.countifs.value) == (2, 1)
        assert (m.s.averageif.value, m.s.averageifs.value) == (200.0, 150.0)
        assert (m.s.maxifs.value, m.s.minifs.value, m.s.count.value, m.s.counta.value) == (
            300.0, 100.0, 3, 3)

    def test_every_criteria_kind(self, tmp_path):
        m = _model()
        d, i = m.data, m.inputs
        with m.s as s:
            s.text = SUMIF(d.region, "North", d.amount).set_display_name("text")
            s.number = SUMIF(d.year, 2026, d.amount).set_display_name("number")
            s.cell = SUMIF(d.region, i.pick, d.amount).set_display_name("cell")
            s.rule = SUMIF(d.amount, i.rule).set_display_name("rule")
            s.joined = SUMIF(d.amount, (">=", i.floor)).set_display_name("joined")
            s.concat = SUMIF(d.amount, mo.CONCAT(">=", i.floor)).set_display_name("concat")
            s.pair = SUMIF(d.amount, (">", 150)).set_display_name("pair")
            s.quoted = COUNTIF(d.region, 'say "hi"').set_display_name("quoted")
        rows = _rows(_book(m, tmp_path)["Summary"])
        assert rows["text"][0] == '=SUMIF(Data!B2:D2, "North", Data!B3:D3)'
        assert rows["number"][0] == "=SUMIF(Data!B4:D4, 2026, Data!B3:D3)"
        assert rows["cell"][0] == "=SUMIF(Data!B2:D2, Inputs!B2, Data!B3:D3)"
        assert rows["rule"][0] == "=SUMIF(Data!B3:D3, Inputs!B3)"
        assert rows["joined"][0] == '=SUMIF(Data!B3:D3, ">=" & Inputs!B1)'
        assert rows["concat"][0] == '=SUMIF(Data!B3:D3, ">=" & Inputs!B1)'
        assert rows["pair"][0] == '=SUMIF(Data!B3:D3, ">150")'
        assert rows["quoted"][0] == '=COUNTIF(Data!B2:D2, "say ""hi""")'
        assert (m.s.cell.value, m.s.rule.value, m.s.joined.value, m.s.pair.value) == (
            100.0, 500.0, 500.0, 500.0)

    def test_on_the_rows_own_sheet_and_inside_a_larger_formula(self, tmp_path):
        m = _model()
        with m.data as d:
            d.north = SUMIF(d.region, "North", d.amount).set_display_name("north")
            d.share = (SUMIF(d.region, "North", d.amount) / mo.SUM(d.amount)).set_display_name("share")
        rows = _rows(_book(m, tmp_path)["Data"])
        # The scalar lines take column B; the periods move to C.
        assert rows["north"][0] == '=SUMIF(C2:E2, "North", C3:E3)'
        assert rows["share"][0] == '=SUMIF(C2:E2, "North", C3:E3) / SUM(C3:E3)'


# ══════════════════════════════════════════════════════════════════════
# A list of scalar lines is a vertical range
# ══════════════════════════════════════════════════════════════════════


class TestRecordsAreVerticalRanges:
    def test_a_field_down_the_records(self, tmp_path):
        m = _model()
        d = m.deals
        regions = [d.a.region, d.b.region, d.c.region]
        amounts = [d.a.amount, d.b.amount, d.c.amount]
        with m.s as s:
            s.sumif = SUMIF(regions, "North", amounts).set_display_name("sumif")
            s.sumifs = SUMIFS(amounts, regions, "North", amounts, ">100").set_display_name("sumifs")
            s.countif = COUNTIF(regions, "North").set_display_name("countif")
            s.countifs = COUNTIFS(regions, "North", amounts, ">100").set_display_name("countifs")
            s.averageif = AVERAGEIF(regions, "North", amounts).set_display_name("averageif")
            s.averageifs = AVERAGEIFS(amounts, regions, "North").set_display_name("averageifs")
            s.maxifs = MAXIFS(amounts, regions, "North").set_display_name("maxifs")
            s.minifs = MINIFS(amounts, regions, "North").set_display_name("minifs")
            s.count = COUNT(amounts).set_display_name("count")
            s.counta = COUNTA(regions).set_display_name("counta")
            s.one = SUMIF([d.a.region], "North", [d.a.amount]).set_display_name("one")
        rows = _rows(_book(m, tmp_path)["Summary"])
        assert rows["sumif"][0] == '=SUMIF(Deals!B3:B5, "North", Deals!C3:C5)'
        assert rows["sumifs"][0] == '=SUMIFS(Deals!C3:C5, Deals!B3:B5, "North", Deals!C3:C5, ">100")'
        assert rows["countif"][0] == '=COUNTIF(Deals!B3:B5, "North")'
        assert rows["countifs"][0] == '=COUNTIFS(Deals!B3:B5, "North", Deals!C3:C5, ">100")'
        assert rows["averageif"][0] == '=AVERAGEIF(Deals!B3:B5, "North", Deals!C3:C5)'
        assert rows["averageifs"][0] == '=AVERAGEIFS(Deals!C3:C5, Deals!B3:B5, "North")'
        assert rows["maxifs"][0] == '=_xlfn.MAXIFS(Deals!C3:C5, Deals!B3:B5, "North")'
        assert rows["minifs"][0] == '=_xlfn.MINIFS(Deals!C3:C5, Deals!B3:B5, "North")'
        assert rows["count"][0] == "=COUNT(Deals!C3:C5)"
        assert rows["counta"][0] == "=COUNTA(Deals!B3:B5)"
        assert rows["one"][0] == '=SUMIF(Deals!B3, "North", Deals!C3)'
        assert (m.s.sumif.value, m.s.sumifs.value, m.s.countif.value) == (400.0, 300.0, 2)

    def test_one_period_s_column_of_the_records_blocks(self, tmp_path):
        m = _model()
        d = m.deals
        regions = [d.a.region, d.b.region, d.c.region]
        with m.s as s:
            s.jan = SUMIFS([d.a.flow[0], d.b.flow[0], d.c.flow[0]], regions, "North").set_display_name("jan")
            s.mar = SUMIFS([d.a.flow[-1], d.b.flow[-1], d.c.flow[-1]], regions, "North").set_display_name("mar")
        rows = _rows(_book(m, tmp_path)["Summary"])
        assert rows["jan"][0] == '=SUMIFS(Deals!D8:D10, Deals!B3:B5, "North")'
        assert rows["mar"][0] == '=SUMIFS(Deals!F8:F10, Deals!B3:B5, "North")'
        assert m.s.jan.value == pytest.approx(40.0) and m.s.mar.value == pytest.approx(120.0)

    def test_on_the_table_s_own_sheet(self, tmp_path):
        m = _model()
        with m.deals as d:
            d.north = SUMIF([d.a.region, d.b.region, d.c.region], "North",
                            [d.a.amount, d.b.amount, d.c.amount]).set_display_name("North total")
        rows = _rows(_book(m, tmp_path)["Deals"])
        assert rows["North total"][0] == '=SUMIF(B3:B5, "North", C3:C5)'

    def test_a_row_of_scalar_lines_across_a_board_is_a_horizontal_range(self, tmp_path):
        m = mo.Model("m", display_name="M")
        with m:
            m.board = mo.MultiVariable("Board", excel_props={"tab": True, "orient": "columns"})
            with m.board as b:
                for name, cost in (("p1", 50.0), ("p2", 70.0), ("p3", 90.0)):
                    col = mo.MultiVariable(name.upper())
                    setattr(b, name, col)
                    col.cost = mo.Variable(cost, display_name="Cost")
            m.s = mo.MultiVariable("Summary", excel_props={"tab": True})
            m.s.big = COUNTIF([m.board.p1.cost, m.board.p2.cost, m.board.p3.cost],
                              ">60").set_display_name("big")
        rows = _rows(_book(m, tmp_path)["Summary"])
        assert rows["big"][0] == '=COUNTIF(Board!B2:D2, ">60")' and m.s.big.value == 2


class TestNoRunShipsTheValue:
    def test_a_gap_two_sheets_or_a_literal_among_the_cells(self, tmp_path):
        m = _model()
        d = m.deals
        with m.s as s:
            s.gap = SUMIF([d.a.region, d.c.region], "North",
                          [d.a.amount, d.c.amount]).set_display_name("gap")
            s.sheets = SUMIF([d.a.region, m.inputs.pick], "North",
                             [d.a.amount, m.inputs.floor]).set_display_name("sheets")
            s.literal = SUMIF([d.a.region, "North"], "North",
                              [d.a.amount, 7.0]).set_display_name("literal")
            s.backwards = SUMIF([d.c.region, d.b.region, d.a.region], "North",
                                [d.c.amount, d.b.amount, d.a.amount]).set_display_name("backwards")
        rows = _rows(_book(m, tmp_path)["Summary"])
        # (``pick`` holds "north": the text criteria is case-insensitive.)
        assert rows["gap"][0] == 400 and rows["sheets"][0] == 250
        assert rows["literal"][0] == 107 and rows["backwards"][0] == 400
        issues = check_compat(m, "excel")
        assert [(i.kind, i.qpath) for i in issues] == [
            ("value", m.s.gap.id), ("value", m.s.sheets.id),
            ("value", m.s.literal.id), ("value", m.s.backwards.id)]
        assert all(i.literal is None and i.inlined.startswith("[") for i in issues), issues

    def test_a_run_is_listed_by_nobody(self, tmp_path):
        m = _model()
        d = m.deals
        with m.s as s:
            s.ok = SUMIF([d.a.region, d.b.region, d.c.region], "North",
                         [d.a.amount, d.b.amount, d.c.amount]).set_display_name("ok")
        _book(m, tmp_path)
        assert check_compat(m, "excel") == []

    def test_the_counts_list_the_cells(self, tmp_path):
        # COUNT and COUNTA take any number of arguments: a comma list of
        # the cells reads as their range would.
        m = _model()
        d = m.deals
        with m.s as s:
            s.count = COUNT([d.a.amount, d.c.amount]).set_display_name("count")
            s.counta = COUNTA([d.a.amount, d.c.amount]).set_display_name("counta")
        rows = _rows(_book(m, tmp_path)["Summary"])
        assert rows["count"][0] == "=COUNT(Deals!C3, Deals!C5)"
        assert rows["counta"][0] == "=COUNTA(Deals!C3, Deals!C5)"
        assert check_compat(m, "excel") == []

    def test_a_row_the_book_has_no_cells_for(self, tmp_path):
        m = _model()
        with m.s as s:
            s.scaled = SUMIF(m.data.region, "North", m.data.amount * 2.0).set_display_name("scaled")
        with pytest.warns(mo.CrossScopeReferenceWarning, match="SUMIF"):
            rows = _rows(_book(m, tmp_path)["Summary"])
        assert rows["scaled"][0] == 200


class TestContiguousRun:
    @pytest.mark.parametrize("refs, run", [
        (["B3", "B4", "B5"], True),
        (["Deals!B3", "Deals!B4"], True),
        (["$B$3", "$B$4"], True),
        (["B3"], True),
        (["B3", "C3", "D3"], True),
        (["B3", "B5"], False),
        (["B3", "C4"], False),
        (["B5", "B4", "B3"], False),
        (["D3", "C3"], False),
        (["B3", "B4", "C4"], False),
    ])
    def test_one_row_or_one_column_in_order(self, refs, run):
        assert _contiguous_run(refs) is run


# ══════════════════════════════════════════════════════════════════════
# LibreOffice recalculates these functions to the engine's numbers
# ══════════════════════════════════════════════════════════════════════


def _soffice():
    for candidate in ("soffice", "libreoffice"):
        found = shutil.which(candidate)
        if found:
            return found
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


def _recalculated(model, tmp_path):
    """The Summary sheet after LibreOffice recalculates the written
    workbook, by line label."""
    soffice = _soffice()
    if soffice is None:
        pytest.skip("LibreOffice not installed")
    path = tmp_path / "criteria.xlsx"
    model.to_excel(path)
    subprocess.run(
        [soffice, f"-env:UserInstallation={(tmp_path / 'prof').as_uri()}",
         "--headless", "--convert-to", "xlsx", "--outdir",
         str(tmp_path / "calc"), str(path)],
        check=True, capture_output=True,
    )
    return _rows(load_workbook(tmp_path / "calc" / "criteria.xlsx", data_only=True)["Summary"])


def _serial(value):
    """A date as the cell holds it, so a recalculated date (read back as
    a datetime) compares with the engine's date."""
    if isinstance(value, dt.datetime):
        return (value - dt.datetime(1899, 12, 30)) / dt.timedelta(days=1)
    if isinstance(value, dt.date):
        return (value - dt.date(1899, 12, 30)).days
    return value


@pytest.mark.slow
def test_the_workbook_recalculates_to_the_python_values(tmp_path):
    m = _model()
    d, i, t, c = m.data, m.inputs, m.deals, m.closes
    regions = [t.a.region, t.b.region, t.c.region]
    amounts = [t.a.amount, t.b.amount, t.c.amount]
    closes = [c.a.closed, c.b.closed, c.c.closed]
    close_regions = [c.a.region, c.b.region, c.c.region]
    i.blank = mo.Variable(None, display_name="Blank")
    i.empty = mo.Variable("", display_name="Empty")
    d.closed = mo.Variable([dt.date(2026, 1, 31), dt.date(2026, 2, 28), dt.date(2024, 2, 29)],
                           display_name="Closed")
    d.note = mo.Variable(["ab", "a~b", ""], display_name="Note")
    d.mix = mo.Variable([0.0, None, 5.0], display_name="Mix")
    d.flag = mo.Variable([True, False, True], display_name="Flag")
    d.gap = mo.IF(d.flag, "", "x").set_display_name("Gap")
    d.numtxt = mo.Variable(["10", 7.0, "-1"], display_name="Numtxt")
    d.p1 = mo.Variable([0.1, 0.2, 0.7], display_name="P1")
    d.p2 = mo.Variable([0.2, 0.1, -0.4], display_name="P2")
    d.share = (d.p1 + d.p2).set_display_name("Share")
    d.tf = mo.Variable(["TRUE", True, "x"], display_name="Tf")
    d.when = mo.Variable([dt.datetime(2026, 1, 31, 12), dt.datetime(2026, 1, 31, 6),
                          dt.datetime(2026, 1, 31, 18)], display_name="When")
    i.thr = mo.Variable(0.3, display_name="Thr")
    i.err = mo.Variable("#DIV/0!", display_name="Err")
    with m.s as s:
        s.sumif = SUMIF(d.region, "North", d.amount).set_display_name("sumif")
        s.sumif_ge = SUMIF(d.amount, ">=150").set_display_name("sumif_ge")
        s.sumif_cell = SUMIF(d.region, i.pick, d.amount).set_display_name("sumif_cell")
        s.sumif_rule = SUMIF(d.amount, i.rule).set_display_name("sumif_rule")
        s.sumif_joined = SUMIF(d.amount, (">=", i.floor)).set_display_name("sumif_joined")
        s.sumifs = SUMIFS(d.amount, d.region, "N*", d.year, 2026).set_display_name("sumifs")
        s.countif = COUNTIF(d.region, "n*").set_display_name("countif")
        s.countif_ne = COUNTIF(d.region, "<>South").set_display_name("countif_ne")
        s.countifs = COUNTIFS(d.region, "<>South", d.amount, ">100").set_display_name("countifs")
        s.averageif = AVERAGEIF(d.region, "N*", d.amount).set_display_name("averageif")
        s.averageifs = AVERAGEIFS(d.amount, d.year, 2026).set_display_name("averageifs")
        s.maxifs = MAXIFS(d.amount, d.region, "N*").set_display_name("maxifs")
        s.minifs = MINIFS(d.amount, d.region, "N*", d.year, "<>2026").set_display_name("minifs")
        s.count = (COUNT(d.amount) + COUNT(d.region)).set_display_name("count")
        s.counta = (COUNTA(d.region) + COUNTA(d.amount)).set_display_name("counta")
        # Two rows and a value typed in, counted period by period.
        s.count_pp = COUNT(d.amount, d.region, d.mix).set_display_name("count_pp")
        s.counta_pp = COUNTA(d.region, d.note, d.gap, "x").set_display_name("counta_pp")
        s.v_sumif = SUMIF(regions, "North", amounts).set_display_name("v_sumif")
        s.v_sumifs = SUMIFS(amounts, regions, "North", amounts, ">100").set_display_name("v_sumifs")
        s.v_countifs = COUNTIFS(regions, "North", amounts, "<300").set_display_name("v_countifs")
        s.v_averageif = AVERAGEIF(regions, "North", amounts).set_display_name("v_averageif")
        s.v_maxifs = MAXIFS(amounts, regions, "North").set_display_name("v_maxifs")
        s.v_minifs = MINIFS(amounts, regions, "North").set_display_name("v_minifs")
        s.v_jan = SUMIFS([t.a.flow[0], t.b.flow[0], t.c.flow[0]], regions, "North").set_display_name("v_jan")
        s.v_count = COUNT(amounts).set_display_name("v_count")
        # Empty text typed in is a blank cell; a formula's is a cell.
        s.typed_ne = COUNTIF(d.note, "<>").set_display_name("typed_ne")
        s.typed_blank = COUNTIF(d.note, "").set_display_name("typed_blank")
        s.typed_counta = COUNTA(d.note).set_display_name("typed_counta")
        s.formula_ne = COUNTIF(d.gap, "<>").set_display_name("formula_ne")
        s.formula_blank = COUNTIF(d.gap, "").set_display_name("formula_blank")
        s.formula_counta = COUNTA(d.gap).set_display_name("formula_counta")
        # ``~`` takes the next character literally.
        s.tilde = COUNTIF(d.note, "a~b").set_display_name("tilde")
        s.tilde_lead = COUNTIF(d.region, "~north").set_display_name("tilde_lead")
        # None, a blank cell and an empty-text cell as the criteria are 0.
        s.none = COUNTIF(d.mix, None).set_display_name("none")
        s.blank_cell = SUMIF(d.mix, i.blank, d.amount).set_display_name("blank_cell")
        s.empty_cell = COUNTIF(d.mix, i.empty).set_display_name("empty_cell")
        # Dates are numbers; the largest, the smallest and the mean come back as dates.
        s.dt_serial = COUNTIF(d.closed, ">45000").set_display_name("dt_serial")
        s.dt_ge = COUNTIF(d.closed, (">=", dt.date(2026, 1, 1))).set_display_name("dt_ge")
        s.dt_maxifs = MAXIFS(d.closed, d.region, "N*").set_display_name("dt_maxifs")
        s.dt_minifs = MINIFS(d.closed, d.region, "N*").set_display_name("dt_minifs")
        s.dt_averageif = AVERAGEIF(d.region, "<>South", d.closed).set_display_name("dt_averageif")
        s.dt_sumif = SUMIF(d.region, "N*", d.closed).set_display_name("dt_sumif")
        s.v_maxifs_date = MAXIFS(closes, close_regions, "North").set_display_name("v_maxifs_date")
        s.v_averageifs_date = AVERAGEIFS(closes, close_regions, "North").set_display_name("v_averageifs_date")
        # Numeric text meets an equality, never an ordering.
        s.nt_gt = COUNTIF(d.numtxt, ">5").set_display_name("nt_gt")
        s.nt_ge0 = COUNTIF(d.numtxt, ">=0").set_display_name("nt_ge0")
        s.nt_lt0 = COUNTIF(d.numtxt, "<0").set_display_name("nt_lt0")
        s.nt_sumif = SUMIF(d.numtxt, ">5", d.amount).set_display_name("nt_sumif")
        s.nt_text = COUNTIF(d.numtxt, "10").set_display_name("nt_text")
        # '=' is a blank cell only; an ordering alone compares text with ''.
        s.eq_formula = COUNTIF(d.gap, "=").set_display_name("eq_formula")
        s.eq_typed = COUNTIF(d.note, "=").set_display_name("eq_typed")
        s.ge_alone = COUNTIF(d.gap, ">=").set_display_name("ge_alone")
        s.gt_alone = COUNTIF(d.gap, ">").set_display_name("gt_alone")
        s.le_alone = COUNTIF(d.gap, "<=").set_display_name("le_alone")
        s.ge_typed = COUNTIF(d.note, ">=").set_display_name("ge_typed")
        s.eq_none = COUNTIF(d.mix, ("=", None)).set_display_name("eq_none")
        s.ne_none = COUNTIF(d.mix, ("<>", None)).set_display_name("ne_none")
        s.ge_blank = COUNTIF(d.region, (">=", i.blank)).set_display_name("ge_blank")
        # Numbers are equal as the spreadsheet compares them: 0.1 + 0.2 is 0.3.
        s.fl_eq = COUNTIF(d.share, 0.3).set_display_name("fl_eq")
        s.fl_cell = COUNTIF(d.share, i.thr).set_display_name("fl_cell")
        s.fl_ge = COUNTIF(d.share, ">=0.3").set_display_name("fl_ge")
        s.fl_gt = COUNTIF(d.share, ">0.3").set_display_name("fl_gt")
        s.fl_sumif = SUMIF(d.share, 0.3, d.amount).set_display_name("fl_sumif")
        # A criteria cell holding an error is the result.
        s.err_count = COUNTIF(d.region, i.err).set_display_name("err_count")
        s.err_sumif = SUMIF(d.region, i.err, d.amount).set_display_name("err_sumif")
        s.err_joined = COUNTIF(d.amount, (">", i.err)).set_display_name("err_joined")
        # A date written 2026-01-31 is that date.
        s.iso_ge = COUNTIF(d.closed, ">=2026-01-01").set_display_name("iso_ge")
        s.iso_eq = COUNTIF(d.closed, "2026-01-31").set_display_name("iso_eq")
        s.iso_sumif = SUMIF(d.closed, "<2026-02-01", d.amount).set_display_name("iso_sumif")
        # The text TRUE meets the criteria written "TRUE".
        s.true_text = COUNTIF(d.tf, "TRUE").set_display_name("true_text")
        s.not_true = COUNTIF(d.tf, "<>TRUE").set_display_name("not_true")
        # A time of day is written as the serial, never as DATE(...).
        s.noon = COUNTIF(d.when, dt.datetime(2026, 1, 31, 12)).set_display_name("noon")
        s.from_noon = COUNTIF(d.when, (">=", dt.datetime(2026, 1, 31, 12))).set_display_name("from_noon")
        # None typed into COUNT is the 0 the formula holds; COUNTA lists its cells.
        s.count_none = COUNT(1, None).set_display_name("count_none")
        s.counta_list = COUNTA([t.a.amount, t.c.amount]).set_display_name("counta_list")
    calc = _recalculated(m, tmp_path)
    lines = [(name, line) for name, line in m.s.walk() if isinstance(line, mo.Variable)]
    assert len(lines) == 75
    # Empty text typed into Note is a blank; Gap's, a formula's, is a cell.
    assert m.s.count_pp.value == [2, 1, 2] and m.s.counta_pp.value == [4, 4, 3]
    assert (m.s.nt_gt.value, m.s.nt_lt0.value, m.s.eq_formula.value, m.s.fl_eq.value) == (1, 0, 0, 3)
    assert (m.s.err_count.value, m.s.iso_ge.value, m.s.true_text.value) == ("#DIV/0!", 2, 2)
    assert (m.s.noon.value, m.s.from_noon.value) == (1, 2)
    assert m.s.dt_maxifs.value == dt.date(2026, 1, 31)
    assert m.s.dt_averageif.value == dt.date(2025, 2, 14)
    assert m.s.v_averageifs_date.value == dt.datetime(2025, 3, 15, 12)
    for name, line in lines:
        if isinstance(line.value, list):
            # One call per period, in the period columns after the scalars'.
            assert calc[name][1:4] == pytest.approx(line.value), (name, calc[name], line.value)
            continue
        assert _serial(calc[name][0]) == pytest.approx(_serial(line.value)), (
            name, calc[name][0], line.value)


@pytest.mark.slow
@pytest.mark.xfail(strict=True, reason=(
    "LibreOffice reads a boolean cell as a number, does not read numeric text in a "
    "criteria range as a number, and matches a wildcard against a number's digits; "
    "Excel does none of these, and the engine keeps Excel's rule"))
def test_where_libreoffice_departs_from_excel(tmp_path):
    m = mo.Model("m", display_name="M", default_grain="month", default_start="2026-01",
                 default_periods=6)
    with m:
        m.data = mo.MultiVariable("Data", excel_props={"tab": True})
        m.data.bools = mo.Variable([True, False, True, None, 1, 0], display_name="Bools")
        m.data.num = mo.Variable([100.0, -5.0, 0.0, None, "10", True], display_name="Num")
        m.s = mo.MultiVariable("Summary", excel_props={"tab": True})
    d = m.data
    with m.s as s:
        # A boolean cell: Excel 1, LibreOffice 3 (TRUE is 1 there); COUNT 2 against 5.
        s.boolean = COUNTIF(d.bools, ">0").set_display_name("boolean")
        s.boolean_count = COUNT(d.bools).set_display_name("boolean_count")
        # Numeric text in the criteria range: Excel 1, LibreOffice 0.
        s.numeric_text = COUNTIF(d.num, 10).set_display_name("numeric_text")
        # A wildcard over the number 100: Excel 1 (the text "10"), LibreOffice 2.
        s.wildcard = COUNTIF(d.num, "1*").set_display_name("wildcard")
        # "TRUE" over TRUE, FALSE, TRUE, 1, 0: Excel 2, LibreOffice 3 (1 is TRUE there).
        s.true_one = COUNTIF(d.bools, "TRUE").set_display_name("true_one")
    assert (m.s.boolean.value, m.s.boolean_count.value, m.s.true_one.value) == (1, 2, 2)
    assert (m.s.numeric_text.value, m.s.wildcard.value) == (1, 1)
    calc = _recalculated(m, tmp_path)
    departing = [(name, calc[name][0], line.value) for name, line in m.s.walk()
                 if isinstance(line, mo.Variable) and calc[name][0] != line.value]
    assert departing == [], departing
