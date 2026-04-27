# SPDX-License-Identifier: Apache-2.0
"""Tests for cohort_retention — active subscribers under a flat retention rate."""

import pytest
from openpyxl import load_workbook

import modeleon as mo


class TestValues:
    def test_first_period_equals_new(self):
        new = mo.Variable([100, 150, 200])
        r = mo.Variable(0.85)
        active = mo.cohort_retention(new, r)
        assert active._value[0] == pytest.approx(100.0)

    def test_recurrence_math(self):
        # active[t] = active[t-1] * r + new[t]
        new = mo.Variable([100, 150, 200, 250])
        r = mo.Variable(0.8)
        active = mo.cohort_retention(new, r)
        assert active._value[0] == pytest.approx(100.0)
        assert active._value[1] == pytest.approx(100 * 0.8 + 150)
        assert active._value[2] == pytest.approx(active._value[1] * 0.8 + 200)
        assert active._value[3] == pytest.approx(active._value[2] * 0.8 + 250)


class TestValidation:
    def test_scalar_new_customers_raises(self):
        with pytest.raises(ValueError, match="must be a list Variable"):
            mo.cohort_retention(mo.Variable(100), mo.Variable(0.85))

    def test_list_retention_raises(self):
        with pytest.raises(ValueError, match="must be a scalar"):
            mo.cohort_retention(mo.Variable([100, 150]), mo.Variable([0.85, 0.80]))

    def test_plain_value_new_customers_raises(self):
        with pytest.raises(TypeError, match="must be a Variable"):
            mo.cohort_retention([100, 150], mo.Variable(0.85))


class TestExcelEmission:
    """cohort_retention must produce live Excel formulas (one per period
    referencing the previous period), not the hand-calculated values.
    Prior version emitted ``=cohort_retention(...)`` which Excel can't evaluate.
    """

    def test_excel_cells_are_formulas(self, tmp_path):
        m = mo.MultiVariable("Cohort")
        m.c = mo.MultiVariable("C", excel_props={'tab': True})
        m.c.new_cust = mo.Variable([100, 150, 200, 250])
        m.c.retention = mo.Variable(0.85)
        m.c.active = mo.cohort_retention(m.c.new_cust, m.c.retention)

        m.to_excel(tmp_path / "c.xlsx")

        path = tmp_path / "c.xlsx"
        ws = load_workbook(path)["C"]
        active_row = None
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("=") and "+" in cell.value:
                    active_row = cell.row
                    break
            if active_row:
                break
        assert active_row is not None

        formulas = [
            cell.value for cell in ws[active_row]
            if isinstance(cell.value, str) and cell.value.startswith("=")
        ]
        assert len(formulas) >= 2
        # No broken ``=cohort_retention(...)`` formula
        assert not any("cohort_retention" in f for f in formulas)
