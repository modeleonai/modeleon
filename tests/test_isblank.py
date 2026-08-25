# SPDX-License-Identifier: Apache-2.0
"""``ISBLANK`` — the Excel-native "was this ever filled in?".

The case it exists for: an OPTIONAL date. A leaving date that hasn't
happened yet has to read as an empty cell — filling it with the end of
the window says «уволен 31 декабря», which an accountant will act on.
So the emptiness must survive into the workbook as a question Excel
itself can answer, not as a sentinel value the model invented.
"""

import logging

import pytest

import modeleon as mo


class TestValue:
    def test_none_and_empty_string_are_blank(self):
        assert mo.ISBLANK(mo.Variable(None)).value is True
        assert mo.ISBLANK(mo.Variable("")).value is True

    def test_a_real_value_is_not(self):
        assert mo.ISBLANK(mo.Variable("2026-03-01")).value is False
        assert mo.ISBLANK(mo.Variable(0)).value is False  # zero is a value

    def test_element_wise_over_a_list(self):
        v = mo.ISBLANK(mo.Variable(["", None, "x"]))
        assert v.value == [True, True, False]


class TestExcel:
    def test_renders_as_a_live_formula_over_the_cell(self, tmp_path, caplog):
        m = mo.Model("m", display_name="M")
        with m:
            m.дата = mo.Variable(None, display_name="Дата увольнения")
            m.работает = mo.ISBLANK(m.дата)
            m.работает._display_name = "Работает"
        path = tmp_path / "o.xlsx"
        with caplog.at_level(logging.WARNING, logger="modeleon"):
            m.to_excel(str(path))
        from openpyxl import load_workbook

        ws = load_workbook(str(path))["M"]
        row = next(
            r for r in range(1, 5) if ws.cell(r, 1).value == "Работает"
        )
        assert ws.cell(row, 2).value == "=ISBLANK(B1)"
        # The text alone proves nothing: an UNREGISTERED function emits
        # verbatim, so the cell reads correctly either way and the only
        # difference is a logged "Unknown function" the user never sees.
        # That is the omission this pins.
        assert "Unknown function" not in caplog.text


class TestUnderAGrainLens:
    def test_a_coarser_grain_refuses_rather_than_inventing_a_number(self):
        # ISBLANK is deliberately NOT in ``_RECOMPUTE_FUNCS``. The
        # boolean functions that are (``NOT``, ``AND``, ``OR``) come
        # back from a quarterly lens as [1, 2] — the truth series
        # SUMMED, a count of blanks wearing a boolean's name. Refusing
        # is the honest answer until that is fixed: the author is told
        # to declare what a quarter of "was it empty?" means.
        from modeleon.core.projection import project_model

        m = mo.Model("m", display_name="M", default_grain="month",
                     default_start="2026-01", default_periods=6)
        m.даты = mo.Variable(["a", "", "b", None, "c", ""],
                             display_name="Даты", regrain=mo.up("last"))
        m.пусто = mo.ISBLANK(m.даты)
        m.пусто._display_name = "Пусто"
        with pytest.raises(ValueError, match="re-grain"):
            project_model(m, grain="quarter")
