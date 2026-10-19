# SPDX-License-Identifier: Apache-2.0
"""A caption that heads no row says its unit the way the sheet says units.

The fields of a records table and the scalars across a board have a
caption in a value column, not a row with a unit column beside it. On
a sheet whose view declares a unit column the table writes a UNIT ROW
directly under its captions — one unit per field column, blank where
a field has none, in the unit column's own look — and the captions
stay bare, as the row labels do beside the unit column. A sheet that
declares no unit column suffixes the unit into the caption, as it
does into a row label; a view that hides units (``meta={'unit':
False}``) shows them nowhere — neither row nor suffix.
"""

from openpyxl import load_workbook

import modeleon as mo
from modeleon.compile.excel.layout import LayoutEngine

GREY = "#7F7F7F"


def _rows(model, sheet, tmp_path, n, cols="ABCDEF"):
    path = tmp_path / "t.xlsx"
    model.to_excel(str(path))
    ws = load_workbook(str(path))[sheet]
    return [[ws[f"{c}{r}"].value for c in cols] for r in range(1, n + 1)]


class Comp(mo.MultiVariableClass):
    def compute(self, ev, ebitda):
        self.ev = mo.Variable(ev, display_name="EV", unit="USD m")
        self.ebitda = mo.Variable(ebitda, display_name="EBITDA", unit="USD m")
        self.deals = mo.Variable(ebitda * 2, display_name="Deals")


def _comps(meta):
    m = mo.Model("m", display_name="M")
    m.default_excel_view = mo.ExcelView(meta=meta)
    with m:
        m.comps = mo.MultiVariable("Comps", excel_props={"tab": True, "orient": "records"})
        with m.comps as c:
            c.a = Comp(ev=1200, ebitda=100, display_name="Alpha")
            c.b = Comp(ev=900, ebitda=90, display_name="Beta")
    return m


def _kpis(meta, **view):
    m = mo.Model("m", display_name="M")
    m.default_excel_view = mo.ExcelView(meta=meta, **view)
    m.kpis = mo.MultiVariable("KPIs", excel_props={"tab": True, "orient": "columns"})
    m.kpis.irr = mo.Variable(0.2, display_name="IRR", unit="%")
    m.kpis.moic = mo.Variable(2.5, display_name="MOIC", unit="x")
    m.kpis.deals = mo.Variable(3, display_name="Deals")
    return m


class TestARecordsTable:
    def test_with_a_unit_column_the_units_stand_in_a_row_under_bare_captions(self, tmp_path):
        rows = _rows(_comps({"unit": "Unit"}), "Comps", tmp_path, 5)
        assert rows[0] == [None, None, "EV", "EBITDA", "Deals", None]
        assert rows[1] == [None, None, "USD m", "USD m", None, None]
        assert rows[2] == ["Alpha", None, 1200, 100, 200, None]
        assert rows[3] == ["Beta", None, 900, 90, 180, None]
        assert rows[4] == [None] * 6

    def test_without_a_unit_column_the_caption_carries_the_unit(self, tmp_path):
        rows = _rows(_comps({}), "Comps", tmp_path, 3)
        assert rows[0] == [None, "EV (USD m)", "EBITDA (USD m)", "Deals", None, None]
        assert rows[1] == ["Alpha", 1200, 100, 200, None, None]
        assert rows[2] == ["Beta", 900, 90, 180, None, None]

    def test_hidden_units_show_nowhere(self, tmp_path):
        # ``meta={'unit': False}`` hides units: no row under the captions,
        # no suffix in a caption or a row label.
        m = _comps({"unit": False})
        with m:
            m.comps.total = mo.Variable(2100, display_name="Total EV", unit="USD m")
        rows = _rows(m, "Comps", tmp_path, 4)
        assert rows[0] == [None, "EV", "EBITDA", "Deals", None, None]
        assert rows[3] == ["Total EV", 2100, None, None, None, None]

    def test_a_table_whose_fields_carry_no_unit_reserves_no_row(self, tmp_path):
        m = mo.Model("m", display_name="M")
        m.default_excel_view = mo.ExcelView(meta={"unit": "Unit"})
        with m:
            m.staff = mo.MultiVariable("Staff", excel_props={"tab": True, "orient": "records"})
            with m.staff as s:
                s.e1 = mo.MultiVariable("Employee 1")
                s.e1.salary = mo.Variable(100.0, display_name="Salary")
                s.e2 = mo.MultiVariable("Employee 2")
                s.e2.salary = mo.Variable(200.0, display_name="Salary")
        rows = _rows(m, "Staff", tmp_path, 3, cols="ABC")
        assert rows[0] == [None, None, "Salary"]
        assert rows[1] == ["Employee 1", None, 100.0]
        assert rows[2] == ["Employee 2", None, 200.0]

    def test_the_unit_row_wears_the_unit_columns_look_and_no_lead_cell(self, tmp_path):
        m = _comps({"article": "No.", "unit": "Unit"})
        m.default_excel_view = mo.ExcelView(
            meta={"article": "No.", "unit": "Unit"},
            sheet={"columns": [("unit", {"font_color": GREY, "italic": True})]},
        )
        path = tmp_path / "t.xlsx"
        m.to_excel(str(path))
        ws = load_workbook(str(path))["Comps"]
        rows = [[ws[f"{c}{r}"].value for c in "ABCDEFG"] for r in range(1, 4)]
        # The article, label and unit columns of the unit row stay empty:
        # the row is the table's, not a line's.
        assert rows[0] == [None, None, None, "EV", "EBITDA", "Deals", None]
        assert rows[1] == [None, None, None, "USD m", "USD m", None, None]
        assert rows[2] == [None, "Alpha", None, 1200, 100, 200, None]
        unit = ws["D2"]
        assert unit.font.i and unit.font.color.rgb == "FF7F7F7F"
        assert not ws["D1"].font.i

    def test_the_layout_publishes_the_unit_row_per_caption(self):
        m = _comps({"unit": "Unit"})
        engine = LayoutEngine([m.comps], flatten_nested_sheets=True, header_rows=True)
        addresses = engine.compute_addresses()
        assert addresses["m.comps.a.ev"].name == "C1"
        assert addresses["m.comps.a.ev"].values == ["C3"]
        assert engine.unit_rows == {"m.comps.a.ev": 2, "m.comps.a.ebitda": 2,
                                    "m.comps.a.deals": 2}
        plain = LayoutEngine([_comps({}).comps], flatten_nested_sheets=True, header_rows=True)
        plain.compute_addresses()
        assert plain.unit_rows == {}


class TestAScalarsAcrossBoard:
    def test_with_a_unit_column_the_units_stand_in_a_row_under_bare_captions(self, tmp_path):
        rows = _rows(_kpis({"unit": "Unit"}), "KPIs", tmp_path, 4)
        assert rows[1] == [None, None, "IRR", "MOIC", "Deals", None]
        assert rows[2] == [None, None, "%", "x", None, None]
        assert rows[3] == [None, None, 0.2, 2.5, 3, None]

    def test_without_a_unit_column_the_caption_carries_the_unit(self, tmp_path):
        rows = _rows(_kpis({}), "KPIs", tmp_path, 3)
        assert rows[1] == [None, "IRR (%)", "MOIC (x)", "Deals", None, None]
        assert rows[2] == [None, 0.2, 2.5, 3, None, None]

    def test_hidden_units_show_nowhere(self, tmp_path):
        rows = _rows(_kpis({"unit": False}), "KPIs", tmp_path, 3)
        assert rows[1] == [None, "IRR", "MOIC", "Deals", None, None]
        assert rows[2] == [None, 0.2, 2.5, 3, None, None]

    def test_the_unit_row_wears_the_unit_columns_look(self, tmp_path):
        m = _kpis({"unit": "Unit"},
                  sheet={"columns": [("unit", {"font_color": GREY, "italic": True})]})
        path = tmp_path / "t.xlsx"
        m.to_excel(str(path))
        ws = load_workbook(str(path))["KPIs"]
        assert ws["C3"].value == "%" and ws["C3"].font.i
        assert ws["C3"].font.color.rgb == "FF7F7F7F"
        assert ws["C2"].value == "IRR" and not ws["C2"].font.i

    def test_a_formula_reading_a_board_cell_follows_it(self, tmp_path):
        m = _kpis({"unit": "Unit"})
        with m:
            m.out = mo.MultiVariable("Out", excel_props={"tab": True})
            m.out.twice = (m.kpis.moic * 2).set_display_name("Twice")
        rows = _rows(m, "Out", tmp_path, 1, cols="ABC")
        assert rows[0] == ["Twice", None, "=KPIs!D4 * 2"]


class TestRecordsThatDisagreeOnAUnit:
    def test_the_record_that_names_the_column_names_its_unit(self, tmp_path):
        # The first record carrying a field owns the header cell; its unit
        # is the column's unit. A unit only a later record carries does not
        # stand under the caption, and reserves no row on its own.
        class Mixed(mo.MultiVariableClass):
            def compute(self, ev_unit, ebitda_unit):
                self.ev = mo.Variable(1.0, display_name="EV", unit=ev_unit)
                self.ebitda = mo.Variable(2.0, display_name="EBITDA", unit=ebitda_unit)

        m = mo.Model("m", display_name="M")
        m.default_excel_view = mo.ExcelView(meta={"unit": "Unit"})
        with m:
            m.comps = mo.MultiVariable("Comps", excel_props={"tab": True, "orient": "records"})
            with m.comps as c:
                c.a = Mixed(ev_unit=None, ebitda_unit="USD m", display_name="Alpha")
                c.b = Mixed(ev_unit="EUR m", ebitda_unit="%", display_name="Beta")
        rows = _rows(m, "Comps", tmp_path, 4, cols="ABCDE")
        assert rows[0] == [None, None, "EV", "EBITDA", None]
        assert rows[1] == [None, None, None, "USD m", None]
        assert rows[2] == ["Alpha", None, 1.0, 2.0, None]
        assert rows[3] == ["Beta", None, 1.0, 2.0, None]

        m2 = mo.Model("m", display_name="M")
        m2.default_excel_view = mo.ExcelView(meta={"unit": "Unit"})
        with m2:
            m2.comps = mo.MultiVariable("Comps", excel_props={"tab": True, "orient": "records"})
            with m2.comps as c:
                c.a = Mixed(ev_unit=None, ebitda_unit=None, display_name="Alpha")
                c.b = Mixed(ev_unit="EUR m", ebitda_unit="%", display_name="Beta")
        rows = _rows(m2, "Comps", tmp_path, 3, cols="ABCDE")
        assert rows[0] == [None, None, "EV", "EBITDA", None]
        assert rows[1] == ["Alpha", None, 1.0, 2.0, None]
        assert rows[2] == ["Beta", None, 1.0, 2.0, None]
