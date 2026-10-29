# SPDX-License-Identifier: Apache-2.0
"""A range is read once.

Every call over the same lines — the rank of each record among its peers,
a ``SUMIFS`` per record over the records' fields, a ``MATCH`` per record —
reads the range once and shares it: one AST node that every call's formula
holds and depends on as one dependency (the node depends on its lines), the
cells, the unit they share and what is derived from them (the numbers
sorted, the cells indexed by value); a call's source names the lines as
they were named when the range was read, in one text every call over it
shares. The workbook writes the range once per sheet. A range is read
again when one of its lines was written to since (its value replaced, its
unit, its path, its placement) — a write reads again only the ranges read
from that line —, when a window moved, or when a list a line holds was
changed in place.

The scale tests count the work, never the clock: a table twice as long
costs twice the work, not four times.
"""

import copy
import datetime as dt
import gc
import pickle
import warnings

import pytest
from openpyxl import load_workbook

import modeleon as mo
from modeleon.compile.excel.renderer import ExcelRenderer
from modeleon.core import variable
from modeleon.core.expr import ListExpr
from modeleon.functions import criteria, lookup, statistics


class Comp(mo.MultiVariableClass):
    def compute(self, multiple, rev=None):
        self.multiple = mo.Variable(multiple, display_name="Multiple", unit="x")
        if rev is not None:
            self.rev = mo.Variable(rev, display_name="Rev", unit="USD")


class Tx(mo.MultiVariableClass):
    def compute(self, region, amount, day):
        self.region = mo.Variable(region, display_name="Region")
        self.amount = mo.Variable(amount, display_name="Amount", unit="USD")
        self.day = mo.Variable(day, display_name="Day")


def _model():
    return mo.Model("m", default_grain="year", default_start="2026", default_periods=3)


def _comps(m, n):
    m.comps = mo.MultiVariable("Comps", excel_props={"tab": True, "orient": "records"})
    with m.comps as c:
        for k in range(n):
            setattr(c, f"r{k}", Comp(multiple=float(k * 7 % n), rev=[k, 2.0 * k, n - k],
                                     display_name=f"R{k}"))
    return [m.comps._component(f"r{k}") for k in range(n)]


def _txs(m, n, groups):
    m.tx = mo.MultiVariable("Tx", excel_props={"tab": True, "orient": "records"})
    with m.tx as t:
        for k in range(n):
            setattr(t, f"x{k}", Tx(region=f"G{k % groups}", amount=float(k % 13),
                                   day=dt.date(2026, 1, 1 + k % 28), display_name=f"x{k}"))
    return [m.tx._component(f"x{k}") for k in range(n)]


def _read_by_each_call(monkeypatch):
    """From here on every call reads its range itself, as before a range
    was read once: the answers to compare with."""
    monkeypatch.setattr(criteria, "_shared", lambda arg: None)
    monkeypatch.setattr(statistics, "_shared", lambda arg: None)


@pytest.fixture
def counted(monkeypatch):
    """The work, by kind: the cells read off a line, the cells a statistic
    counts, the cells a criteria or a lookup tests, the cells a renderer
    looks up for a range, the lists of lines walked for their lines."""
    work = {"cells": 0, "counted": 0, "tested": 0, "rendered": 0, "walked": 0}

    real_cells = criteria._cells

    def cells(var, raw):
        work["cells"] += 1
        return real_cells(var, raw)

    real_counted = statistics._counted

    def counted_(cells_):
        work["counted"] += len(cells_)
        return real_counted(cells_)

    def testing(real):
        def make(*args, **kwargs):
            test = real(*args, **kwargs)

            def tested(v):
                work["tested"] += 1
                return test(v)
            return tested
        return make

    real_cell_of = ExcelRenderer._cell_of

    def cell_of(self, *args, **kwargs):
        work["rendered"] += 1
        return real_cell_of(self, *args, **kwargs)

    monkeypatch.setattr(criteria, "_cells", cells)
    monkeypatch.setattr(statistics, "_counted", counted_)
    monkeypatch.setattr(criteria, "_matcher", testing(criteria._matcher))
    monkeypatch.setattr(lookup, "_exact_test", testing(lookup._exact_test))
    real_walk = ListExpr.iter_refs

    def walked(self):
        work["walked"] += 1
        return real_walk(self)

    monkeypatch.setattr(ExcelRenderer, "_cell_of", cell_of)
    monkeypatch.setattr(ListExpr, "iter_refs", walked)
    return work


def _rows(path, sheet):
    return [[c.value for c in row] for row in load_workbook(path)[sheet].iter_rows()]


# ══════════════════════════════════════════════════════════════════════
# The work grows with the table, not with its square
# ══════════════════════════════════════════════════════════════════════


class TestScale:
    def _ranked(self, counted, n, tmp_path):
        m = _model()
        with m:
            recs = _comps(m, n)
            peers = [r.multiple for r in recs]
            for k in counted:
                counted[k] = 0
            with m.comps:
                for r in recs:
                    r.rank = mo.RANK_EQ(r.multiple, peers).set_display_name("Rank")
            built = dict(counted)
            counted["rendered"] = counted["walked"] = 0
            m.to_excel(tmp_path / f"rank{n}.xlsx")
        return built, (counted["rendered"], counted["walked"])

    def test_a_rank_per_record_reads_and_writes_its_range_once(self, counted, tmp_path):
        (small, written_small), (large, written_large) = (
            self._ranked(counted, 60, tmp_path), self._ranked(counted, 120, tmp_path))
        # The range is read once: its cells once, counted and sorted once.
        assert (small["cells"], large["cells"]) == (60, 120)
        assert (small["counted"], large["counted"]) == (60, 120)
        # Written once per sheet: each peer's cell looked up once, and the
        # list walked once for what its lines read.
        assert (written_small, written_large) == ((60, 1), (120, 1))

    def test_a_rank_per_record_and_year_sorts_each_year_once(self, counted):
        def ranked(n):
            m = _model()
            with m:
                recs = _comps(m, n)
                revs = [r.rev for r in recs]
                counted["counted"] = 0
                with m.comps:
                    for r in recs:
                        r.rank = mo.RANK_EQ(r.rev, revs).set_display_name("Rank")
            return counted["counted"]

        assert (ranked(40), ranked(80)) == (3 * 40, 3 * 80)

    def test_a_sumifs_per_record_tests_only_the_cells_of_its_key(self, counted, tmp_path):
        def summed(n):
            m = _model()
            with m:
                recs = _txs(m, n, groups=n // 4)
                amounts, regions = [r.amount for r in recs], [r.region for r in recs]
                counted["tested"] = counted["cells"] = 0
                with m.tx:
                    for r in recs:
                        r.total = mo.SUMIFS(amounts, regions, r.region).set_display_name("Tot")
                built = dict(counted)
                counted["rendered"] = 0
                m.to_excel(tmp_path / f"sumifs{n}.xlsx")
            return built, counted["rendered"]

        (small, written_small), (large, written_large) = summed(60), summed(120)
        # The first call tests every cell, as an equality alone does; each
        # after tests the four cells of its key. Each line is read once (its
        # region once more, as the criteria of its own call).
        assert (small["tested"], large["tested"]) == (60 + 4 * 59, 120 + 4 * 119)
        assert (small["cells"], large["cells"]) == (3 * 60, 3 * 120)
        assert (written_small, written_large) == (2 * 60, 2 * 120)

    def test_a_match_per_record_tests_only_the_cells_of_its_key(self, counted):
        def matched(n):
            m = _model()
            with m:
                recs = _txs(m, n, groups=n)
                regions = [r.region for r in recs]
                counted["tested"] = 0
                with m.tx:
                    for r in recs:
                        r.at = mo.MATCH(r.region, regions, 0).set_display_name("At")
                        r.back = mo.XLOOKUP(r.region, regions, regions).set_display_name("Back")
                assert [r.at.value for r in recs] == list(range(1, n + 1))
            return counted["tested"]

        # The first call tests every cell; each after, the one of its key.
        assert (matched(60), matched(120)) == (60 + 2 * 60 - 1, 120 + 2 * 120 - 1)

    def test_an_equality_alone_reads_every_cell_and_a_second_indexes_them(self, counted):
        lines = [mo.Variable(f"t{k}") for k in range(40)]
        counted["tested"] = 0
        first = mo.COUNTIF(lines, "t7")
        assert counted["tested"] == 40
        shared = criteria._shared(lines)
        assert "index" not in shared.kept
        pattern = mo.COUNTIF(lines, "t1*")       # a pattern, which no index answers
        found = mo.MATCH("t1?", lines, 0)
        assert "index" not in shared.kept
        counted["tested"] = 0
        second = mo.MATCH("T9", lines, 0)
        assert "index" in shared.kept
        assert counted["tested"] == 1
        assert (first.value, pattern.value, found.value, second.value) == (1, 11, 11, 10)

    def test_each_call_checks_its_range_without_reading_the_lines_again(self, monkeypatch):
        reads = [0]
        real = criteria._VALUE

        def value(var):
            reads[0] += 1
            return real(var)

        monkeypatch.setattr(criteria, "_VALUE", value)

        def ranked(n, field):
            m = _model()
            with m:
                recs = _comps(m, n)
                peers = [getattr(r, field) for r in recs]
                reads[0] = 0
                with m.comps:
                    for r in recs:
                        r.rank = mo.RANK_EQ(getattr(r, field), peers).set_display_name("Rank")
            return reads[0]

        # Each line's value read once, by the first call: the check before
        # each reuse compares two counts, not every line — n reads for n
        # calls, where a check reading each line made n² of them. A line
        # holding a list keeps its cells compared at each reuse.
        assert (ranked(60, "multiple"), ranked(120, "multiple")) == (60, 120)
        assert (ranked(60, "rev"), ranked(120, "rev")) == (60, 120)

        # How a line is shown is no part of what a range reads: a peer named
        # or styled between the calls reads nothing again. Its unit is.
        m = _model()
        with m:
            recs = _comps(m, 10)
            peers = [r.multiple for r in recs]
            reads[0] = 0
            with m.comps:
                for r in recs:
                    r.rank = mo.RANK_EQ(r.multiple, peers)
                    r.multiple.set_display_name("Multiple (x)").set_style(bold=True)
            assert reads[0] == 10
            recs[0].multiple.unit = "x"
            assert mo.LARGE(peers, 1).value == 9.0 and reads[0] == 20

    def test_a_median_per_record_is_worked_out_once(self, counted):
        def mid(n):
            m = _model()
            with m:
                recs = _comps(m, n)
                peers = [r.multiple for r in recs]
                counted["counted"] = 0
                with m.comps:
                    for r in recs:
                        r.gap = (r.multiple - mo.MEDIAN(peers)).set_display_name("Gap")
            return counted["counted"]

        assert (mid(60), mid(120)) == (60, 120)


# ══════════════════════════════════════════════════════════════════════
# One node, shared by every call over the same lines
# ══════════════════════════════════════════════════════════════════════


class TestOneNode:
    def test_the_calls_hold_one_node_and_a_new_list_of_the_same_lines_finds_it(self):
        m = _model()
        with m:
            recs = _comps(m, 5)
            peers = [r.multiple for r in recs]
            with m.comps:
                for r in recs:
                    r.rank = mo.RANK_EQ(r.multiple, list(peers))
            nodes = {id(r.rank._expr.args[1]) for r in recs}
            assert len(nodes) == 1
            again = mo.PERCENTILE_INC(tuple(peers), 0.5)
            assert again._expr.args[0] is recs[0].rank._expr.args[1]
            other = mo.RANK_EQ(recs[0].multiple, peers[::-1])
            assert other._expr.args[1] is not again._expr.args[0]

    def test_a_call_over_the_same_list_finds_its_range_by_the_list(self, monkeypatch):
        looked_up = [0]
        real = criteria._find

        def find(*args):
            looked_up[0] += 1
            return real(*args)

        monkeypatch.setattr(criteria, "_find", find)
        m = _model()
        with m:
            recs = _comps(m, 6)
            peers = [r.multiple for r in recs]
            ranks = [mo.RANK_EQ(r.multiple, peers) for r in recs]
            # Looked up by its lines once; then found by the list itself.
            assert looked_up[0] == 1
            peers.reverse()                     # changed in place: another range
            top = mo.LARGE(peers, 1)
            assert looked_up[0] == 2
            assert top._expr.args[0] is not ranks[0]._expr.args[1]
            assert [item.var for item in top._expr.args[0].items] == peers
            assert mo.LARGE(peers, 6)._expr.args[0] is top._expr.args[0]
            assert looked_up[0] == 2

    def test_a_list_is_compared_with_its_range_a_tuple_is_its_range(self, monkeypatch):
        compared = [0]
        real = criteria.operator.is_

        def is_(a, b):
            compared[0] += 1
            return real(a, b)

        m = _model()
        with m:
            recs = _comps(m, 8)
            peers = [r.multiple for r in recs]
            fixed = tuple(peers)
            mo.LARGE(peers, 1)
            mo.LARGE(fixed, 1)
            monkeypatch.setattr(criteria.operator, "is_", is_)
            for r in recs:
                mo.RANK_EQ(r.multiple, peers)
            # A list may have changed since: each call compares its items.
            assert compared[0] == 8 * 8
            compared[0] = 0
            for r in recs:
                mo.RANK_EQ(r.multiple, fixed)
            # A tuple cannot: found once, it is its range at once after.
            assert compared[0] == 0
            assert mo.RANK_EQ(recs[0].multiple, fixed)._expr.args[1] \
                is mo.RANK_EQ(recs[0].multiple, peers)._expr.args[1]

    def test_a_copy_or_a_pickle_takes_the_formula_and_reads_the_lines_again(self):
        m = _model()
        with m:
            recs = _comps(m, 4)
            peers = [r.multiple for r in recs]
            with m.comps:
                for r in recs:
                    r.rank = mo.RANK_EQ(r.multiple, peers)
        for twin in (copy.deepcopy(recs[1].rank), pickle.loads(pickle.dumps(recs[1].rank))):
            node = twin._expr.args[1]
            assert node.__dict__.get("_lines") is None
            assert [item.var.value for item in node.items] == [0.0, 3.0, 2.0, 1.0]
            assert twin.value == recs[1].rank.value == 1

    def test_a_range_no_formula_holds_any_more_is_forgotten(self):
        m = _model()
        with m:
            recs = _comps(m, 4)
            key = tuple(id(r.multiple) for r in recs)
            rank = mo.RANK_EQ(recs[0].multiple, [r.multiple for r in recs])
            assert key in criteria._LISTS
            del rank
            gc.collect()
            assert key not in criteria._LISTS


class TestOneDependency:
    @staticmethod
    def _ranked(n):
        m = _model()
        with m:
            recs = _comps(m, n)
            peers = [r.multiple for r in recs]
            with m.comps:
                for r in recs:
                    r.rank = mo.RANK_EQ(r.multiple, peers).set_display_name("Rank")
        return recs, peers

    def test_a_call_per_record_holds_the_range_as_one_entry(self):
        for n in (60, 120):
            recs, peers = self._ranked(n)
            held = [r.rank._dependency_refs for r in recs]
            # Its number and the range: 2n entries over the table, not n².
            assert sum(map(len, held)) == 2 * n
            node = held[0][1]
            assert all(refs[0] is r.multiple and refs[1] is node for r, refs in zip(recs, held))
            assert [item.var for item in node.items] == peers
            assert recs[0].rank._dependency_lines() == peers
            assert recs[0].rank.dependencies == {line.id for line in peers}

    def test_a_sumifs_per_record_holds_its_ranges_and_its_key(self):
        def held(n):
            m = _model()
            with m:
                recs = _txs(m, n, groups=n // 4)
                amounts, regions = [r.amount for r in recs], [r.region for r in recs]
                with m.tx:
                    for r in recs:
                        r.total = mo.SUMIFS(amounts, regions, r.region).set_display_name("Tot")
            return sum(len(r.total._dependency_refs) for r in recs)

        assert (held(60), held(120)) == (3 * 60, 3 * 120)

    def test_every_line_reads_what_it_read_when_each_call_listed_its_lines(self, monkeypatch):
        def reads(var):
            ids = [line.id for line in var._dependency_lines()]
            return [i if not i.startswith("__floating__") else "<floating>" for i in ids]

        shared = _records_model()
        _read_by_each_call(monkeypatch)
        each = _records_model()
        for (name, a), (_, b) in zip(shared.walk(), each.walk()):
            if isinstance(a, mo.Variable):
                assert reads(a) == reads(b), name
                assert len(a.dependencies) == len(b.dependencies), name
                assert a.dependencies - {d for d in a.dependencies if "__floating__" in d} \
                    == b.dependencies - {d for d in b.dependencies if "__floating__" in d}, name
        assert shared.detect_cycles() is None and each.detect_cycles() is None

    def test_no_call_keeps_a_text_that_lists_its_lines(self):
        recs, peers = self._ranked(6)
        parts = [r.rank._source_call[1][1] for r in recs]
        assert all(part is parts[0] for part in parts)
        assert all("_source_code" not in r.rank.__dict__ for r in recs)
        listed = ", ".join(f"r{k}.multiple" for k in range(6))
        assert recs[2].rank.source_code == f"RANK_EQ(multiple, [{listed}])"

    def test_a_call_keeps_the_names_its_lines_had_when_it_was_made(self):
        # As the rest of a call's source does; a call made after the move
        # names the lines as they are named then.
        m = _model()
        with m:
            m.s = mo.MultiVariable("S")
            m.t = mo.MultiVariable("T")
            m.s.a, m.s.b = mo.Variable(1.0), mo.Variable(2.0)
            lines = [m.s.a, m.s.b]
            top = mo.LARGE(lines, 1)
            assert top.source_code == "LARGE([s.a, s.b], 1)"
            m.t.a = m.s.remove("a")
            assert top.source_code == "LARGE([s.a, s.b], 1)"
            assert mo.Variable(top).source_code == "LARGE([s.a, s.b], 1)"
            assert mo.LARGE(lines, 1).source_code == "LARGE([t.a, s.b], 1)"

    @pytest.mark.parametrize("peers_are", ["lines", "formulas", "per-period formulas",
                                           "elements", "ranked"])
    def test_rebuilding_calls_one_by_one_reads_the_range_once(self, monkeypatch, peers_are):
        # A call replaced by its own name detaches the one before it; the
        # peers it reads stay where they are, so their range is not read
        # again for each call rebuilt — named lines, unnamed formulas, which
        # each detach moving a window has asked where they run, elements of
        # rows; nor when a range reads the calls themselves (``ranked``): a
        # call detached marks that range alone. What a formula over the
        # range takes from its lines (``_range_reads``) is worked out once.
        reads = []
        worked_out = []
        original, original_reads = criteria._Lines._read, variable._range_reads

        def counted(self):
            reads.append(1)
            return original(self)

        def counted_reads(node):
            shared = node.__dict__.get("_lines")
            if shared is None or shared.stale or variable._RANGE_READS not in shared.kept:
                worked_out.append(1)
            return original_reads(node)

        def rebuild(n):
            m = _model()
            with m:
                m.p = mo.MultiVariable("P")
                for i in range(n):
                    value = {"elements": [float(i), 0.0, 0.0],
                             "per-period formulas": [float(i)] * 3}.get(peers_are, float(i))
                    setattr(m.p, f"x{i}", mo.Variable(value))
                lines = [getattr(m.p, f"x{i}") for i in range(n)]
                peers = {"formulas": [x * 1 for x in lines],
                         "per-period formulas": [x * 1 for x in lines],
                         "elements": [x[0] for x in lines]}.get(peers_are, lines)
                m.r = mo.MultiVariable("R")
                for i in range(n):
                    setattr(m.r, f"rank{i}", mo.RANK_EQ(peers[i], peers))
                ranks = [getattr(m.r, f"rank{i}") for i in range(n)]
                if peers_are == "ranked":
                    m.top = mo.COUNTIF(ranks, "<=10")
            monkeypatch.setattr(criteria._Lines, "_read", counted)
            monkeypatch.setattr(variable, "_range_reads", counted_reads)
            reads.clear()
            worked_out.clear()
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                with m:
                    for i in range(n):
                        setattr(m.r, f"rank{i}", mo.RANK_EQ(peers[i], peers))
            rebuilt = len(reads) + len(worked_out)
            if peers_are == "ranked":
                # The calls detached were written to: a range over them is
                # read again, once.
                before = len(reads)
                assert mo.COUNTIF(ranks, "<=10").value == 10 and len(reads) == before + 1
            monkeypatch.setattr(criteria._Lines, "_read", original)
            monkeypatch.setattr(variable, "_range_reads", original_reads)
            assert m.r.rank0.value == ([n] * 3 if peers_are == "per-period formulas" else n)
            return rebuilt

        assert rebuild(40) == 0 and rebuild(80) == 0

    def test_a_write_to_a_line_reads_again_only_the_ranges_read_from_it(self, monkeypatch):
        reads = []
        original = criteria._Lines._read

        def counted(self):
            reads.append(1)
            return original(self)

        m = _model()
        with m:
            recs = _comps(m, 6)
            multiples = [r.multiple for r in recs]
            firsts = [r.rev[0] for r in recs]       # elements of rows: 0, 1, ... 5
            assert (mo.LARGE(multiples, 1).value, mo.LARGE(firsts, 1).value) == (5.0, 5.0)
            monkeypatch.setattr(criteria._Lines, "_read", counted)
            recs[0].multiple._value = 9.0           # as a loop closing writes it
            assert mo.LARGE(firsts, 1).value == 5.0 and len(reads) == 0
            assert mo.LARGE(multiples, 1).value == 9.0 and len(reads) == 1
            recs[1].rev.unit = "EUR"                # the row an element of ``firsts`` names
            assert mo.LARGE(multiples, 2).value == 5.0 and len(reads) == 1
            assert mo.LARGE(firsts, 2).value == 4.0 and len(reads) == 2

    def test_a_source_text_pickles_with_every_protocol(self):
        m = _model()
        with m:
            m.s = mo.MultiVariable("S")
            m.s.a, m.s.b = mo.Variable(1.0), mo.Variable(2.0)
            m.s.top = mo.LARGE([m.s.a, m.s.b], 1)
        for protocol in range(pickle.HIGHEST_PROTOCOL + 1):
            twin = pickle.loads(pickle.dumps(m, protocol=protocol))
            assert twin.s.top.source_code == "LARGE([s.a, s.b], 1)", protocol
            assert twin.s.top.value == 2.0

    def test_a_range_over_a_line_reading_time_reads_time(self, monkeypatch):
        def calls():
            m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=3)
            with m:
                m.s = mo.MultiVariable("S")
                m.s.days = mo.time.days * 1.0
                m.s.a = mo.Variable([1.0, 2.0, 3.0])
                return [mo.MEDIAN([m.s.a, m.s.days]), mo.MEDIAN([m.s.a, m.s.a * 2])]

        shared = calls()
        _read_by_each_call(monkeypatch)
        each = calls()
        assert [c._reads_time for c in shared] == [c._reads_time for c in each] == [True, False]
        assert [c.value for c in shared] == [c.value for c in each]

    def test_a_model_of_calls_over_ranges_copies_and_pickles(self, tmp_path):
        m = _records_model()
        m.to_excel(tmp_path / "model.xlsx")
        for twin in (copy.deepcopy(m), pickle.loads(pickle.dumps(m))):
            for (name, a), (_, b) in zip(m.walk(), twin.walk()):
                if isinstance(a, mo.Variable):
                    assert (b.value, b.source_code) == (a.value, a.source_code), name
                    assert len(b.dependencies) == len(a.dependencies), name
            twin.to_excel(tmp_path / "twin.xlsx")
            for sheet in ("Comps", "Summary"):
                assert (_rows(tmp_path / "twin.xlsx", sheet)
                        == _rows(tmp_path / "model.xlsx", sheet))
            # A call over the twin's lines reads them, and a write to one is
            # seen; the model's own lines are not the twin's.
            peers = [twin.comps._component(f"r{k}").multiple for k in range(8)]
            assert mo.LARGE(peers, 2).value == m.s.top.value
            twin.comps._component("r0").multiple._value = 99.0
            assert mo.LARGE(peers, 1).value == 99.0
            assert mo.LARGE([m.comps._component(f"r{k}").multiple for k in range(8)],
                            1).value == 7.0


class TestTheWriterWalksAListOnce:
    def test_a_line_off_the_window_in_a_shared_list_is_named_for_each_formula(
            self, tmp_path):
        m = _model()
        # Read by date, as ``a + b`` reads it: a year later, no rule to
        # continue it — refused when the call is built.
        late = mo.Variable([5.0, 6.0, 7.0], start="2027", grain="year")
        # Two years from the sheet's first, continued by zero: read by date,
        # but a line the book does not write, whose values would stand in
        # the sheet's columns — refused by the export, for each formula.
        odd = mo.Variable([5.0, 6.0], start="2026", grain="year", extend=mo.zero())
        with m:
            m.s = mo.MultiVariable("S")
            m.s.a = mo.Variable([1.0, 2.0, 3.0])
            m.s.b = mo.Variable([3.0, 1.0, 2.0])
            with pytest.raises(ValueError, match="does not cover the combined window"):
                mo.RANK_EQ(m.s.a, [m.s.a, m.s.b, late])
            lines = [m.s.a, m.s.b, odd]
            m.s.r1 = mo.RANK_EQ(m.s.a, lines)
            m.s.r2 = mo.RANK_EQ(m.s.b, lines)
        with pytest.raises(ValueError, match="time window of its own") as refused:
            m.to_excel(tmp_path / "odd.xlsx")
        assert "(m.s.r1) reads" in str(refused.value)
        assert "(m.s.r2) reads" in str(refused.value)

    def test_a_list_met_again_on_its_own_way_down_is_not_taken_as_walked(self):
        from modeleon.compile.excel.writer import _reads_unwritten_off_window
        from modeleon.core.expr import FuncCall, VarRef

        a = mo.Variable()
        inner = ListExpr([VarRef(a)])
        a._set_expr(FuncCall("MEDIAN", [inner]))
        outer = ListExpr([VarRef(a)])
        clean = {}
        assert _reads_unwritten_off_window(FuncCall("F", [outer]), None, {}, set(),
                                           clean) is None
        assert clean == {}
        b = mo.Variable()
        b._set_expr(FuncCall("MEDIAN", [ListExpr([VarRef(mo.Variable(1.0))])]))
        shared = ListExpr([VarRef(b)])
        assert _reads_unwritten_off_window(FuncCall("F", [shared]), None, {}, set(),
                                           clean) is None
        assert shared in clean.values()


# ══════════════════════════════════════════════════════════════════════
# A line that changed is read again
# ══════════════════════════════════════════════════════════════════════


class TestReadAgain:
    def test_a_value_replaced_is_read_again(self):
        m = _model()
        with m:
            recs = _comps(m, 4)                     # 0, 3, 2, 1
            peers = [r.multiple for r in recs]
            assert mo.RANK_EQ(recs[1].multiple, peers).value == 1
            # As a loop closing, or a schedule laid out on adoption, writes it.
            recs[0].multiple._value = 9.0
            assert mo.RANK_EQ(recs[1].multiple, peers).value == 2
            assert mo.LARGE(peers, 1).value == 9.0
            assert mo.SUMIF(peers, 9.0).value == 9.0
            assert mo.MATCH(9.0, peers, 0).value == 1

    def test_a_row_changed_in_place_is_read_again(self):
        m = _model()
        with m:
            m.s = mo.MultiVariable("S")
            data = [1, 8.0, 9.0]
            m.s.u = mo.Variable(data)               # the line holds the caller's list
            m.s.k = mo.Variable(8.0)
            assert mo.MEDIAN(m.s.u).value == 8.0
            assert mo.COUNTIF(m.s.u, 8.0).value == 1
            assert mo.RANK_EQ(m.s.k, m.s.u).value == 2
            data[1] = 100.0
            assert mo.MEDIAN(m.s.u).value == 9.0
            assert mo.COUNTIF(m.s.u, 8.0).value == 0
            assert mo.RANK_EQ(m.s.k, m.s.u).value == "#N/A"
            assert mo.MATCH(100.0, m.s.u, 0).value == 2
            assert mo.COUNTIF(m.s.u, 1).value == 1
            m.s.u.value[0] = True                   # equal to 1, yet a boolean
            assert mo.COUNTIF(m.s.u, 1).value == 0
            assert mo.MEDIAN(m.s.u).value == 54.5
            assert mo.LARGE(m.s.u, 3).value == "#NUM!"

    def test_a_series_changed_in_place_is_read_again(self):
        m = _model()
        with m:
            m.s = mo.MultiVariable("S")
            m.s.a = mo.Variable([1.0, 2.0, 3.0])
            m.s.b = mo.Variable([4.0, 5.0, 6.0])
            m.s.c = mo.Variable([7.0, 0.0, 9.0])
            lines = [m.s.a, m.s.b, m.s.c]
            assert mo.RANK_EQ(m.s.a, lines).value == [3, 2, 3]
            assert mo.MEDIAN(lines).value == [4.0, 2.0, 6.0]
            m.s.a.value[0] = 100.0
            assert mo.RANK_EQ(m.s.a, lines).value == [1, 2, 3]
            assert mo.MEDIAN(lines).value == [7.0, 2.0, 6.0]
            # A cell moved from one line to the next: the same cells in the
            # same order, in lines of other lengths — refused as they now
            # are, as ``a + b`` refuses them.
            m.s.b.value.insert(0, m.s.a.value.pop())
            with pytest.raises(ValueError, match="does not cover the combined window"):
                m.s.a + m.s.b
            with pytest.raises(ValueError, match="does not cover the combined window"):
                mo.MEDIAN(lines)

    def test_a_unit_set_is_read_again(self):
        m = _model()
        with m:
            recs = _comps(m, 3)
            peers = [r.multiple for r in recs]
            assert str(mo.MEDIAN(peers)._unit) == "x"
            recs[2].multiple.unit = "USD"
            with pytest.raises(ValueError, match="one unit"):
                mo.MEDIAN(peers)
            recs[2].multiple.unit = "x"
            assert str(mo.LARGE(peers, 1)._unit) == "x"

    def test_a_line_adopted_or_moved_is_named_again(self):
        a, b = mo.Variable(1.0), mo.Variable(2.0)
        first = mo.RANK_EQ(a, [a, b])
        assert first.source_code == f"RANK_EQ({a.path.leaf}, [{a.path.leaf}, {b.path.leaf}])"
        m = _model()
        with m:
            m.s = mo.MultiVariable("S")
            m.s.a, m.s.b = a, b
            second = mo.RANK_EQ(m.s.a, [a, b])
        assert second.source_code == "RANK_EQ(a, [s.a, s.b])"

    def test_a_line_moved_to_another_container_is_named_again(self):
        m = _model()
        with m:
            m.s = mo.MultiVariable("S")
            m.t = mo.MultiVariable("T")
            m.s.a, m.s.b = mo.Variable(1.0), mo.Variable(2.0)
            lines = [m.s.a, m.s.b]
            assert mo.LARGE(lines, 1).source_code == "LARGE([s.a, s.b], 1)"
            line = m.s.remove("a")
            m.t.a = line
            assert mo.LARGE(lines, 1).source_code == "LARGE([t.a, s.b], 1)"

    def test_an_element_of_a_row_is_named_by_its_row_as_it_is_named_then(self):
        class Rec(mo.MultiVariableClass):
            def compute(self, rev):
                self.rev = mo.Variable(rev)

        m = _model()
        with m:
            m.comps = mo.MultiVariable("Comps")
        recs = [Rec(rev=[float(k), 2.0 * k, 3.0]) for k in range(4)]   # before they join
        firsts = [r.rev[0] for r in recs]
        before = mo.RANK_EQ(firsts[0], firsts)
        with m:
            for k, r in enumerate(recs):
                setattr(m.comps, f"r{k}", r)
        listed = ", ".join(f"r{k}.rev[0]" for k in range(4))
        assert mo.RANK_EQ(firsts[1], firsts).source_code.endswith(f", [{listed}])")
        # A call made before keeps the names its lines had then.
        assert before.source_code.endswith(", [rev[0], rev[0], rev[0], rev[0]])")

        with m:
            m.c = mo.MultiVariable("C")
            m.d = mo.MultiVariable("D")
            m.c.row = mo.Variable([1.0, 2.0, 3.0])
        row = m.c.row
        pair = [row[0], row[2]]
        assert mo.LARGE(pair, 1).source_code == "LARGE([c.row[0], c.row[2]], 1)"
        m.c.remove("row")
        assert "c.row" not in mo.LARGE(pair, 1).source_code
        with m:
            m.d.row2 = row
        assert mo.LARGE(pair, 1).source_code == "LARGE([d.row2[0], d.row2[2]], 1)"

    def test_lines_with_tracks_on_two_windows_are_read_by_their_dates(self, monkeypatch):
        def medians():
            m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=3,
                         tracks=mo.Tracks("plan", "actual"))
            with m:
                m.a = mo.MultiVariable("A")
                m.b = mo.MultiVariable("B", default_grain="month", default_start="2026-02",
                                       default_periods=3)
                m.a.x = mo.Variable(plan=[1.0, 2.0, 3.0], actual=[1.5, 2.5, 3.5])
                m.a.z = mo.Variable(plan=[5.0, 6.0, 7.0], actual=[5.0, 6.5, 7.0])
                m.b.y = mo.Variable(plan=[10.0, 20.0, 30.0], actual=[10.0, 20.0, 30.0])
                one = mo.MEDIAN([m.a.x, m.a.z]).value
                # Not paired by position: refused, as ``mo.MEDIAN(m.a.x, m.b.y)`` is.
                with pytest.raises(ValueError, match="holds tracks on 3 period"):
                    mo.MEDIAN([m.a.x, m.b.y])
            return {role: one[role] for role in one.roles}

        shared = medians()
        _read_by_each_call(monkeypatch)
        assert shared == medians() == {"plan": [3.0, 4.0, 5.0], "actual": [3.25, 4.5, 5.25]}

    def test_a_window_moved_is_read_again(self, monkeypatch):
        def medians():
            m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=4)
            with m:
                m.a = mo.MultiVariable("A")
                m.b = mo.MultiVariable("B", default_grain="month", default_start="2026-01",
                                       default_periods=4)
                m.a.x = mo.Variable([1.0, 2.0, 3.0, 4.0], extend=mo.zero())
                m.b.y = mo.Variable([10.0, 20.0, 30.0, 40.0], extend=mo.zero())
                lines = [m.a.x, m.b.y]
                first = mo.MEDIAN(lines).value
                m.b.default_start = "2026-02"       # B's line now runs a month later
                return first, mo.MEDIAN(lines).value

        shared = medians()
        _read_by_each_call(monkeypatch)
        assert shared == medians() == ([5.5, 11.0, 16.5, 22.0], [0.5, 6.0, 11.5, 17.0, 20.0])

    def test_a_line_moved_to_another_window_is_read_by_its_dates(self, monkeypatch):
        def medians():
            m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=4)
            with m:
                m.a = mo.MultiVariable("A")
                m.b = mo.MultiVariable("B", default_grain="month", default_start="2026-02",
                                       default_periods=4)
                m.a.x = mo.Variable([1.0, 2.0, 3.0, 4.0], extend=mo.zero())
                m.a.y = mo.Variable([10.0, 20.0, 30.0, 40.0], extend=mo.zero())
                lines = [m.a.x, m.a.y]
                first = mo.MEDIAN(lines).value
                m.b.y = m.a.remove("y")
                return first, mo.MEDIAN(lines).value

        shared = medians()
        _read_by_each_call(monkeypatch)
        assert shared == medians() == ([5.5, 11.0, 16.5, 22.0], [0.5, 6.0, 11.5, 17.0, 20.0])

    def test_a_formula_among_the_lines_placed_anew_is_read_again(self, monkeypatch):
        def medians():
            m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=4)
            with m:
                m.b = mo.MultiVariable("B", default_grain="month", default_start="2026-02",
                                       default_periods=4)
                m.b.y = mo.Variable([10.0, 20.0, 30.0, 40.0], extend=mo.zero())
                x = mo.Variable([1.0, 2.0, 3.0, 4.0], extend=mo.zero())
                lines = [x * 1, m.b.y * 1]
                # ``x`` has no window yet: its formula is read by position.
                first = mo.MEDIAN(lines).value
                m.a = mo.MultiVariable("A")
                m.a.x = x                       # from January now: no one window
                with pytest.raises(ValueError, match="does not cover the combined window"):
                    mo.MEDIAN(lines)
                return first

        shared = medians()
        _read_by_each_call(monkeypatch)
        assert shared == medians() == [5.5, 11.0, 16.5, 22.0]

    def test_a_formula_among_the_lines_is_read_again_once_it_runs_elsewhere(self, monkeypatch):
        # A line that a formula among the lines reads moves: the window the
        # formula runs on may have moved, so where it runs is asked again,
        # and the range is read again only when it runs elsewhere.
        reads = []
        original = criteria._Lines._read

        def counted(self):
            reads.append(1)
            return original(self)

        def medians(count):
            reads.clear()
            m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=4)
            with m:
                m.a = mo.MultiVariable("A")
                m.c = mo.MultiVariable("C")
                m.b = mo.MultiVariable("B", default_grain="month", default_start="2026-02",
                                       default_periods=4)
                m.a.x = mo.Variable([1.0, 2.0, 3.0, 4.0])
                m.a.y = mo.Variable([10.0, 20.0, 30.0, 40.0])
                lines = [m.a.x * 1, m.a.y * 1]
                first = mo.MEDIAN(lines).value
                if count:
                    monkeypatch.setattr(criteria._Lines, "_read", counted)
                m.c.y = m.a.remove("y")             # on the model's window still
                same = mo.MEDIAN(lines).value
                assert len(reads) == 0
                m.b.y = m.c.remove("y")             # from February: no one window
                with pytest.raises(ValueError, match="does not cover the combined window"):
                    mo.MEDIAN(lines)
                assert len(reads) == (1 if count else 0)
            monkeypatch.setattr(criteria._Lines, "_read", original)
            return first, same

        shared = medians(True)
        _read_by_each_call(monkeypatch)
        assert shared == medians(False) == ([5.5, 11.0, 16.5, 22.0], [5.5, 11.0, 16.5, 22.0])

    @pytest.mark.parametrize("window", [{}, {"default_grain": "month"}])
    def test_a_named_formula_on_no_window_of_its_own_is_read_again_once_it_runs_elsewhere(
            self, monkeypatch, window):
        # Named in containers that declare no dated window, a formula runs on
        # the window of the line it reads, as an unnamed one does.
        def medians():
            m = mo.Model("m", **window)
            with m:
                m.a = mo.MultiVariable("A", default_grain="month", default_start="2026-01",
                                       default_periods=3)
                m.c = mo.MultiVariable("C", default_grain="month", default_start="2026-02",
                                       default_periods=3)
                m.b = mo.MultiVariable("B", default_grain="month", default_start="2026-02",
                                       default_periods=2)
                m.s = mo.MultiVariable("S")
                m.a.x = mo.Variable([1.0, 2.0, 3.0])
                m.b.y = mo.Variable([10.0, 20.0], extend=mo.zero())
                m.s.g = m.a.x * 1
                lines = [m.s.g, m.b.y]
                first = mo.MEDIAN(lines)
                m.c.x = m.a.remove("x")             # x, so g, runs from February
                again = mo.MEDIAN(lines)
            return [(r.value, r.time.start) for r in (first, again)]

        shared = medians()
        _read_by_each_call(monkeypatch)
        assert shared == medians() == [([0.5, 6.0, 11.5], "2026-01"),
                                       ([5.5, 11.0, 1.5], "2026-02")]

    def test_a_formula_on_an_open_loop_is_read_again_once_its_line_is_adopted(self):
        # A formula on a loop not closed yet keeps no placement: adopted, the
        # line it reads moves it all the same.
        def close(read_first):
            m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=3)
            with m:
                m.e = mo.MultiVariable("E", default_grain="month", default_start="2026-02",
                                       default_periods=3)
                m.q = mo.Variable([10.0, 20.0, 30.0])
                x = mo.Variable([1.0, 2.0, 3.0])
                d = mo.MultiVariable("D", loop=True)
                with d:
                    d.a = mo.Variable([1.0, 2.0, 3.0])
                    p = x + mo.lag(d.b, fill=0.0)
                    held = mo.MEDIAN([p, m.q]) if read_first else None
                    m.e.x = x                       # p runs from February now
                    d.b = d.a + mo.MEDIAN([p, m.q])
                m.d = d
            return held

        for read_first in (False, True):
            with pytest.raises(ValueError, match="does not cover the combined window"):
                close(read_first)

    def test_a_line_lets_go_of_the_ranges_dropped(self):
        m = _model()
        with m:
            m.c = mo.MultiVariable("C")
            m.c.x = mo.Variable(1.0)
            longest = 0
            for k in range(100):
                mo.MEDIAN([m.c.x, mo.Variable(float(k))]).value     # read once, dropped
                gc.collect()
                longest = max(longest, len(m.c.x._watched))
            assert longest <= 8
            kept = [mo.MEDIAN([m.c.x, mo.Variable(float(k))]) for k in range(100)]
            assert len(m.c.x._watched) < 2 * len(kept) + 8
            del kept
            gc.collect()
            m.c.x.set_unit("x")                     # a write lets go of them too
            assert len(m.c.x._watched) == 0

    def test_a_list_changed_is_another_range(self):
        m = _model()
        with m:
            recs = _comps(m, 4)
            peers = [r.multiple for r in recs]
            assert mo.LARGE(peers, 1).value == 3.0
            peers.append(mo.Variable(7.0))
            assert mo.LARGE(peers, 1).value == 7.0
            peers[-1] = mo.Variable(-1.0)
            assert mo.SMALL(peers, 1).value == -1.0

    def test_empty_text_typed_in_then_given_a_formula(self):
        from modeleon.core.expr import Literal
        a, blank = mo.Variable(1.0), mo.Variable("")
        line = [a, blank]
        assert mo.COUNTIF(line, "=").value == 1        # typed in: a blank cell
        blank.expr = Literal("")
        assert mo.COUNTIF(line, "=").value == 0        # a formula's empty text

    def test_the_first_line_typed_again(self):
        a, b = mo.Variable(1), mo.Variable(2)
        assert mo.LARGE([a, b], 1).value_type == "int"
        a.value_type = "float"
        assert mo.LARGE([a, b], 1).value_type == "float"


class TestALoopClosing:
    """A call made while a loop is open reads provisional rows; when the
    loop closes the call runs again over the same lines, now holding their
    values — the range is read again, as a range read by each call is."""

    @staticmethod
    def _model():
        m = mo.Model("m", default_grain="year", default_start="2026", default_periods=4)
        with m:
            m.d = mo.MultiVariable("D", loop=True)
            with m.d as d:
                d.a = mo.Variable([1.0, 5.0, 2.0, 4.0])
                d.c = mo.lag(d.b, fill=0.0) + 1.0
                d.mid = mo.MEDIAN([d.a, d.c])
                d.rank = mo.RANK_EQ(d.a, [d.a, d.c])
                d.top = mo.LARGE([d.a, d.c], 1)
                d.b = d.a + d.c
        return m

    def test_a_row_read_ahead_among_the_lines_is_read_as_its_line(self, monkeypatch):
        def model():
            m = mo.Model("m", default_grain="year", default_start="2026", default_periods=4)
            with m:
                m.d = mo.MultiVariable("D", loop=True)
                with m.d as d:
                    d.a = mo.Variable([1.0, 5.0, 2.0, 4.0])
                    d.c = mo.lag(d.b, fill=0.0) + 1.0
                    d.mid = mo.MEDIAN([d.a, d.b])       # ``b`` read before its line
                    d.top = mo.LARGE([d.a, d.b], 1)
                    d.b = d.a + d.c
            return m

        from modeleon.core.loops import loop_rows
        shared = model()
        _read_by_each_call(monkeypatch)
        plain = model()
        assert shared.d.b.value == plain.d.b.value == [2.0, 8.0, 11.0, 16.0]
        assert shared.d.mid.value == plain.d.mid.value == [1.5, 6.5, 6.5, 10.0]
        assert shared.d.top.value == plain.d.top.value == [2.0, 8.0, 11.0, 16.0]
        assert ([v.path for v in loop_rows(shared.d.b)]
                == [v.path for v in loop_rows(plain.d.b)])

    def test_the_rows_are_read_again(self, monkeypatch):
        shared = self._model()
        _read_by_each_call(monkeypatch)
        plain = self._model()
        assert shared.d.c.value == plain.d.c.value == [1.0, 3.0, 9.0, 12.0]
        for name in ("mid", "rank", "top"):
            assert getattr(shared.d, name).value == getattr(plain.d, name).value, name
        assert shared.d.mid.value == [1.0, 4.0, 5.5, 8.0]
        assert shared.d.rank.value == [1, 1, 2, 2]
        assert shared.d.top.value == [1.0, 5.0, 9.0, 12.0]


# ══════════════════════════════════════════════════════════════════════
# An index answers as reading every cell does
# ══════════════════════════════════════════════════════════════════════

_CELLS = [
    10, 10.0, 10 + 1e-15, 10 + 1e-13, -0.0, 0, 0.1 + 0.2, 0.3, 1e300, 1e-300, 3, 2.5,
    dt.date(2026, 1, 31), dt.datetime(2026, 1, 31, 12, 0), dt.timedelta(days=3),
    True, False, "10", "1e1", "10%", " 10 ", "North", "north", "NORTH", "nörth",
    "ı", "I", "i", "İ", "ß", "SS", "ss", "ﬅ", "st", "\u212a", "k", "Σ", "σ", "ς",
    "\u0432", "\u1c80", "\u0412", "TRUE", "true", "x*", "x?y", "~", "#N/A", "#text", "", None,
]
_CRITERIA = [
    10, "10", "=10", 10.0, "1e1", True, "TRUE", "true", False, "north", "NORTH", "i",
    "I", "ı", "İ", "ss", "ß", "st", "k", "σ", "ς", "\u0432", "\u0412", dt.date(2026, 1, 31),
    ("=", dt.date(2026, 1, 31)), "#N/A", "", "=", "<>", ">5", "<>north", "n*", "x~*",
    0.3, 0.1 + 0.2, -0.0, 3, "3", None, 1e300,
]


def _answers(lines, amounts):
    out = []
    for crit in _CRITERIA:
        out.append(("SUMIF", crit, mo.SUMIF(lines, crit, amounts).value))
        out.append(("COUNTIF", crit, mo.COUNTIF(lines, crit).value))
        out.append(("AVERAGEIF", crit, mo.AVERAGEIF(lines, crit, amounts).value))
        out.append(("SUMIFS", crit, mo.SUMIFS(amounts, lines, crit, amounts, ">0").value))
        out.append(("COUNTIFS", crit, mo.COUNTIFS(amounts, ">0", lines, crit).value))
        out.append(("MAXIFS", crit, mo.MAXIFS(amounts, lines, crit).value))
    for key in _CELLS + ["x", "nörth", "NöRTH"]:
        if key is None:
            continue
        line = mo.Variable(key)
        out.append(("MATCH", key, mo.MATCH(line, lines, 0).value))
        for mode in (0, 2):
            for search in (1, -1):
                out.append(("XLOOKUP", (key, mode, search),
                            mo.XLOOKUP(line, lines, amounts, "none",
                                       match_mode=mode, search_mode=search).value))
    return out


def _same(a, b):
    return a == b or (isinstance(a, float) and isinstance(b, float) and a != a and b != b)


class TestTheIndexAnswersAsEveryCell:
    def test_criteria_and_lookups_over_awkward_cells(self, monkeypatch):
        lines = [mo.Variable(v) for v in _CELLS]
        amounts = [mo.Variable(float(k + 1)) for k in range(len(_CELLS))]
        indexed = _answers(lines, amounts)
        _read_by_each_call(monkeypatch)
        read = _answers(lines, amounts)
        assert len(indexed) == len(read)
        wrong = [(i, r) for i, r in zip(indexed, read) if not _same(i[2], r[2])]
        assert not wrong

    def test_a_letter_folds_as_the_text_pattern_compares_it(self):
        """Two letters a pattern of one matches fold alike — every letter
        the pattern engine compares without regard to case."""
        _sre = pytest.importorskip("_sre")
        casefix = pytest.importorskip("re._casefix")
        import sys
        lower = {}
        groups = {}
        for i in range(sys.maxunicode + 1):
            if 0xD800 <= i <= 0xDFFF:
                continue
            low = _sre.unicode_tolower(i)
            lower[i] = low
            groups.setdefault(low, []).append(i)
        for p, low in lower.items():
            if not _sre.unicode_iscased(p):
                continue
            key = criteria._fold(chr(p))
            if key is None:
                continue
            for same in (low, *casefix._EXTRA_CASES.get(low, ())):
                for c in groups.get(same, ()):
                    other = criteria._fold(chr(c))
                    assert other is None or other == key, (chr(p), chr(c))


# ══════════════════════════════════════════════════════════════════════
# The workbook is the one a range read by each call writes
# ══════════════════════════════════════════════════════════════════════


def _records_model():
    m = _model()
    with m:
        recs = _comps(m, 8)
        peers = [r.multiple for r in recs]
        revs = [r.rev for r in recs]
        with m.comps:
            for r in recs:
                r.rank = mo.RANK_EQ(r.multiple, peers).set_display_name("Rank")
                r.avg = mo.RANK_AVG(r.multiple, peers, 1).set_display_name("Avg")
                r.p = mo.PERCENTILE_INC(peers, 0.25).set_display_name("P25")
                r.gap = (r.multiple - mo.MEDIAN(peers)).set_display_name("Gap")
                r.year_rank = mo.RANK_EQ(r.rev, revs).set_display_name("YearRank")
                r.same = mo.COUNTIF(peers, r.multiple).set_display_name("Same")
                r.over = mo.SUMIFS(peers, peers, (">", r.multiple)).set_display_name("Over")
                r.at = mo.MATCH(r.multiple, peers, 0).set_display_name("At")
                r.back = mo.INDEX(peers, r.at).set_display_name("Back")
                r.look = mo.XLOOKUP(r.multiple, peers, peers).set_display_name("Look")
        m.s = mo.MultiVariable("Summary", excel_props={"tab": True})
        m.s.top = mo.LARGE(peers, 2)
        m.s.mid = mo.MEDIAN(revs)
    return m


class TestTheSameWorkbook:
    def test_every_cell(self, monkeypatch, tmp_path):
        _records_model().to_excel(tmp_path / "shared.xlsx")
        _read_by_each_call(monkeypatch)
        _records_model().to_excel(tmp_path / "each.xlsx")
        shared = load_workbook(tmp_path / "shared.xlsx")
        each = load_workbook(tmp_path / "each.xlsx")
        assert shared.sheetnames == each.sheetnames
        for name in shared.sheetnames:
            assert _rows(tmp_path / "shared.xlsx", name) == _rows(tmp_path / "each.xlsx", name)
        rows = _rows(tmp_path / "shared.xlsx", "Comps")
        header, first, last = rows[1], rows[2], rows[9]
        assert first[header.index("Rank")] == "=_xlfn.RANK.EQ(B3, B3:B10)"
        assert last[header.index("Rank")] == "=_xlfn.RANK.EQ(B10, B3:B10)"
        assert first[header.index("Over (x)")] == '=SUMIFS(B3:B10, B3:B10, ">" & B3)'
        assert first[header.index("At")] == "=MATCH(B3, B3:B10, 0)"

    def test_the_source_and_the_units(self, monkeypatch):
        shared = _records_model()
        _read_by_each_call(monkeypatch)
        each = _records_model()
        for a, b in zip(shared.walk(), each.walk()):
            va, vb = a[1], b[1]
            if isinstance(va, mo.Variable):
                assert va.source_code == vb.source_code, a[0]
                assert va.value == vb.value, a[0]
                assert str(va._unit) == str(vb._unit), a[0]
                assert va.value_type == vb.value_type, a[0]
