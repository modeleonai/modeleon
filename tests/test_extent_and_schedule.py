# SPDX-License-Identifier: Apache-2.0
"""The extent law (a bare list must cover the whole window) + ``mo.schedule``.

Both are adoption-time behaviors: the ambient window only exists once
the Variable joins a windowed tree, so the hook in
``_register_component`` / ``_run_adoption_hooks`` is where a schedule
materializes and where a wrong-length bare list dies loudly instead of
detonating later at an unrelated broadcast.
"""

from __future__ import annotations

import subprocess
import warnings
from pathlib import Path

import openpyxl
import pytest

import modeleon as mo


def _model(periods: int = 24) -> "mo.Model":
    return mo.Model(
        'm', default_grain='month', default_start='2025-01',
        default_periods=periods,
    )


def _written(model, tmp_path, sheet: str) -> dict:
    """Write ``model`` and return ``{label: [cells right of it]}`` of one
    sheet. Every line must resolve inside the workbook: a reference to a
    variable the book does not contain is a failure, not a warning."""
    out = tmp_path / "book.xlsx"
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        model.to_excel(str(out))
    stray = [w for w in caught
             if issubclass(w.category, mo.CrossScopeReferenceWarning)]
    assert not stray, str(stray[0].message)
    ws = openpyxl.load_workbook(str(out))[sheet]
    return {
        row[0].value: [c for c in row[1:] if c.value is not None]
        for row in ws.iter_rows()
        if row[0].value is not None
    }


def _soffice() -> "str | None":
    """Path to the LibreOffice command line, or None when not installed."""
    import shutil

    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


def _recalculated(model, tmp_path, sheet: str) -> dict:
    """Write ``model``, let LibreOffice calculate every formula, and
    return ``{label: [values right of it]}`` of one sheet — what the
    workbook itself computes."""
    soffice = _soffice()
    if soffice is None:
        pytest.skip("LibreOffice is not installed")
    src = tmp_path / "calc-src.xlsx"
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        model.to_excel(str(src))
    outdir = tmp_path / "calc"
    profile = (tmp_path / "lo-profile").as_uri()  # private: parallel runs don't clash
    subprocess.run(
        [soffice, f"-env:UserInstallation={profile}", "--headless",
         "--convert-to", "xlsx", "--outdir", str(outdir), str(src)],
        check=True, capture_output=True, timeout=120,
    )
    ws = openpyxl.load_workbook(
        str(outdir / "calc-src.xlsx"), data_only=True
    )[sheet]
    return {
        row[0].value: [c.value for c in row[1:] if c.value is not None]
        for row in ws.iter_rows()
        if row[0].value is not None
    }


class TestExtentLaw:
    def test_wrong_length_bare_list_dies_at_adoption(self):
        m = _model()
        with pytest.raises(ValueError, match="24-period window"):
            with m:
                m.короткий = mo.Variable([1.0, 2.0, 3.0])

    def test_full_length_scalar_and_len1_pass(self):
        m = _model(4)
        with m:
            m.полный = mo.Variable([1.0, 2.0, 3.0, 4.0])
            m.скаляр = mo.Variable(0.2)
            m.один = mo.Variable([5.0])
        assert m.полный._value == [1.0, 2.0, 3.0, 4.0]

    def test_own_location_exempts(self):
        # Declared partial: its own start= anchors it — the law's
        # explicit escape hatch for genuinely partial series.
        m = _model()
        with m:
            m.свой = mo.Variable([1.0, 2.0], start='2025-03', grain='month')
        assert m.свой.time.start == '2025-03'

    def test_keyed_list_exempts(self):
        # A labeled list (dict sugar) is not a time series.
        m = _model()
        with m:
            m.кейсы = mo.Variable({'Bear': 1.0, 'Base': 2.0, 'Bull': 3.0})
        assert m.кейсы._keys == ['Bear', 'Base', 'Bull']

    def test_axised_exempts(self):
        # Coordinates, not periods — the rank-1 guard owns that story.
        m = _model()
        сценарии = mo.Variable(['Bear', 'Bull'])
        with m:
            m.осевой = mo.Variable([1.0, 2.0], indexed_by=[сценарии])
        assert m.осевой.time is None

    def test_windowless_model_is_untouched(self):
        with mo.Model('w') as m:
            m.x = mo.Variable([1.0, 2.0, 3.0])
        assert m.x._value == [1.0, 2.0, 3.0]

    def test_operator_results_pass(self):
        m = _model(4)
        with m:
            m.a = mo.Variable([1.0] * 4)
            m.b = m.a * 2
        assert m.b._value == [2.0] * 4


class TestSchedule:
    def test_materializes_on_adoption(self):
        m = _model()
        with m:
            m.параметры = mo.MultiVariable()
            m.параметры.ндс = mo.schedule(
                {'2025-01': 0.12, '2026-01': 0.16}
            )
        v = m.параметры.ндс._value
        assert len(v) == 24
        assert v[0] == 0.12 and v[11] == 0.12
        assert v[12] == 0.16 and v[23] == 0.16

    def test_formula_over_schedule(self):
        m = _model()
        with m:
            m.ставка = mo.schedule({'2025-01': 0.10, '2026-01': 0.20})
            m.цена = mo.Variable([100.0] * 24)
            m.итог = m.цена * (1 + m.ставка)
        assert m.итог._value[0] == pytest.approx(110.0)
        assert m.итог._value[12] == pytest.approx(120.0)

    def test_default_regrain_is_mean(self):
        # A rate re-grains by mean, never by sum.
        m = _model(12)
        with m:
            m.ставка = mo.schedule({'2025-01': 0.12})
        q = m.ставка.at('quarter')
        assert q._value == pytest.approx([0.12, 0.12, 0.12, 0.12])

    def test_step_off_boundary_dies(self):
        m = _model()
        with pytest.raises(ValueError, match="boundary"):
            with m:
                m.ндс = mo.schedule({'2025-01': 0.12, '2026-02-15': 0.16})

    def test_window_start_uncovered_dies(self):
        m = _model()
        with pytest.raises(ValueError, match="window start"):
            with m:
                m.ндс = mo.schedule({'2025-06': 0.12})

    def test_step_beyond_window_is_legal(self):
        # The law only binds steps INSIDE the window span; a future
        # change beyond it simply doesn't materialize.
        m = _model(12)
        with m:
            m.ндс = mo.schedule({'2025-01': 0.12, '2027-01': 0.20})
        assert m.ндс._value == [0.12] * 12

    def test_empty_or_bad_keys_die(self):
        with pytest.raises(TypeError, match="non-empty"):
            mo.schedule({})
        with pytest.raises(TypeError, match="ISO period labels"):
            mo.schedule({2025: 0.12})


class TestExtend:
    """``extend=``: the declared continuation
    legalizes the prefix spelling the extent law otherwise refuses."""

    def test_zero_pads_a_flow(self):
        m = _model()
        with m:
            m.капекс = mo.Variable([500.0, 500.0, 500.0], extend=mo.zero())
        v = m.капекс._value
        assert len(v) == 24
        assert v[:3] == [500.0, 500.0, 500.0]
        assert v[3:] == [0.0] * 21

    def test_hold_carries_a_rate(self):
        m = _model()
        with m:
            m.ставка = mo.Variable([0.10, 0.10, 0.12], extend=mo.hold())
        v = m.ставка._value
        assert len(v) == 24
        assert v[2] == 0.12 and v[23] == 0.12

    def test_full_length_with_extend_is_untouched(self):
        m = _model(3)
        with m:
            m.x = mo.Variable([1.0, 2.0, 3.0], extend=mo.hold())
        assert m.x._value == [1.0, 2.0, 3.0]

    def test_extended_series_computes_downstream(self):
        m = _model(6)
        with m:
            m.поток = mo.Variable([100.0, 100.0], extend=mo.zero())
            m.итог = mo.cumsum(m.поток)
        assert m.итог._value == [100.0, 200.0, 200.0, 200.0, 200.0, 200.0]

    def test_bad_extend_value_dies(self):
        with pytest.raises(TypeError, match="mo.zero"):
            mo.Variable([1.0], extend="zero")

    def test_windowless_extend_stays_partial(self):
        # No window → nothing to materialize to; the declaration stays
        # on the Variable for date-aligned arithmetic to honor.
        with mo.Model('w') as m:
            m.x = mo.Variable([1.0, 2.0], extend=mo.zero())
        assert m.x._value == [1.0, 2.0]


class TestWrappedPendingSource:
    """A line wrapping (or cloned from) a source whose schedule or
    extension has not materialized yet — ``mo.Variable(mo.schedule(...))``.

    The line materializes the steps / extension itself when it joins
    the windowed model, so what it holds is its OWN input. The source
    it was built from never joins the model, so a reference to it has
    nothing true to point at: the workbook must carry the line's values,
    exactly as the bare ``mo.schedule(...)`` spelling does.
    """

    @staticmethod
    def _wrapped(periods: int = 3) -> "mo.Model":
        m = _model(periods)
        with m:
            m.p = mo.MultiVariable('P', excel_props={'tab': True})
            m.p.s = mo.Variable(
                mo.schedule({'2025-01': 0.1, '2025-03': 0.2}),
                display_name='S', unit='%',
                excel_props={'number_format': '0.0%'},
            )
            m.p.f = m.p.s * 100
        return m

    def test_wrapped_schedule_writes_its_steps(self, tmp_path):
        m = self._wrapped()
        assert m.p.s._value == [0.1, 0.1, 0.2]
        cells = _written(m, tmp_path, 'P')['S (%)']
        assert [c.value for c in cells] == [0.1, 0.1, 0.2]

    def test_wrapper_keeps_its_name_unit_and_format(self, tmp_path):
        m = self._wrapped()
        assert m.p.s._display_name == 'S'
        assert str(m.p.s._unit) == '%'
        cells = _written(m, tmp_path, 'P')['S (%)']
        assert {c.number_format for c in cells} == {'0.0%'}

    def test_formula_over_it_references_its_cells(self, tmp_path):
        m = self._wrapped()
        assert m.p.f._value == pytest.approx([10.0, 10.0, 20.0])
        row = _written(m, tmp_path, 'P')['F']
        assert [c.value for c in row] == [
            '=B2 * 100', '=C2 * 100', '=D2 * 100'
        ]
        calc = _recalculated(m, tmp_path, 'P')
        assert calc['S (%)'] == pytest.approx(m.p.s._value)
        assert calc['F'] == pytest.approx(m.p.f._value)

    def test_it_regrains_like_the_bare_schedule(self):
        # The wrapper IS the schedule line, so it keeps the schedule's
        # rule: a rate averages into a quarter, never sums.
        m = self._wrapped(periods=3)
        with m:
            m.p.bare = mo.schedule({'2025-01': 0.1, '2025-03': 0.2})
        quarter = m.p.s.at('quarter')._value
        assert quarter == pytest.approx(m.p.bare.at('quarter')._value)
        assert 0.1 < quarter[0] < 0.2

    def test_schedule_cloned_from_an_unwindowed_template(self, tmp_path):
        # Assigning an owned line into another container clones it; the
        # template has no window, so its own schedule never materialized.
        tmpl = mo.MultiVariable('tmpl')
        tmpl.bare = mo.schedule({'2025-01': 0.1, '2025-03': 0.2},
                                display_name='Bare')
        tmpl.wrap = mo.Variable(mo.schedule({'2025-01': 0.1}),
                                display_name='Wrap')
        m = _model(3)
        with m:
            m.p = mo.MultiVariable('P', excel_props={'tab': True})
            m.p.bare = tmpl.bare
            m.p.wrap = tmpl.wrap
        rows = _written(m, tmp_path, 'P')
        assert [c.value for c in rows['Bare']] == [0.1, 0.1, 0.2]
        assert [c.value for c in rows['Wrap']] == [0.1, 0.1, 0.1]

    def test_wrapped_partial_series_writes_its_extension(self, tmp_path):
        m = _model(4)
        with m:
            m.p = mo.MultiVariable('P', excel_props={'tab': True})
            m.p.e = mo.Variable(
                mo.Variable([1.0, 2.0], extend=mo.zero()), display_name='E'
            )
            m.p.h = mo.Variable(
                mo.Variable([5.0, 6.0], extend=mo.hold()), display_name='H'
            )
            m.p.t = m.p.e + m.p.h
        assert m.p.e._value == [1.0, 2.0, 0.0, 0.0]
        rows = _written(m, tmp_path, 'P')
        assert [c.value for c in rows['E']] == [1.0, 2.0, 0.0, 0.0]
        assert [c.value for c in rows['H']] == [5.0, 6.0, 6.0, 6.0]
        calc = _recalculated(m, tmp_path, 'P')
        assert calc['T'] == pytest.approx(m.p.t._value)

    def test_wrapped_schedule_as_a_shared_track(self, tmp_path):
        m = mo.Model(
            'm', default_grain='month', default_start='2025-01',
            default_periods=3,
            tracks=mo.Tracks(actual='fact', plan='plan'),
        )
        with m:
            m.p = mo.MultiVariable('P', excel_props={'tab': True})
            m.p.rate = mo.Variable(
                mo.schedule({'2025-01': 100.0, '2025-03': 120.0}),
                actual=[99.0, 98.0, 97.0], display_name='Rate',
            )
        rows = _written(m, tmp_path, 'P')
        assert [c.value for c in rows['Rate']] == [99.0, 98.0, 97.0]
        assert [c.value for c in rows['Rate · plan']] == [100.0, 100.0, 120.0]

    def test_wrapping_an_attached_schedule_stays_a_reference(self, tmp_path):
        # The source is a line of the model: the wrapper reads its
        # cells, and the link stays live in the workbook.
        m = _model(3)
        with m:
            m.p = mo.MultiVariable('P', excel_props={'tab': True})
            m.p.b = mo.schedule({'2025-01': 0.1, '2025-03': 0.2},
                                display_name='B')
            m.p.w = mo.Variable(m.p.b, display_name='W')
        rows = _written(m, tmp_path, 'P')
        assert [c.value for c in rows['W']] == ['=B2', '=C2', '=D2']

    def test_source_joining_first_keeps_the_live_link(self, tmp_path):
        # Built in a container before it joins the model: the schedule
        # and its wrapper both materialize at the mount, the schedule
        # first. The wrapper then re-reads a line of the book holding
        # the very same numbers, so its reference stays live.
        sub = mo.MultiVariable('P', excel_props={'tab': True})
        sub.b = mo.schedule({'2025-01': 0.1, '2025-03': 0.2},
                            display_name='B')
        sub.w = mo.Variable(sub.b, display_name='W')
        m = _model(3)
        m.p = sub
        assert m.p.w._value == [0.1, 0.1, 0.2]
        rows = _written(m, tmp_path, 'P')
        assert [c.value for c in rows['W']] == ['=B2', '=C2', '=D2']

    def test_extension_declared_over_a_formula_writes_its_values(
            self, tmp_path):
        # ``Z * 2`` covers only Z's two periods; the wrapper's extend=
        # supplies the rest. A formula cannot say "0 from here on", so
        # the line writes the numbers Python holds.
        m = _model(4)
        with m:
            m.p = mo.MultiVariable('P', excel_props={'tab': True})
            m.p.z = mo.Variable([3.0, 4.0], extend=mo.none(),
                                display_name='Z')
            m.p.w = mo.Variable(m.p.z * 2, extend=mo.zero(),
                                display_name='W')
            m.p.h = mo.Variable(m.p.z * 2, extend=mo.hold(),
                                display_name='H')
        assert m.p.w._value == [6.0, 8.0, 0.0, 0.0]
        assert m.p.h._value == [6.0, 8.0, 8.0, 8.0]
        calc = _recalculated(m, tmp_path, 'P')
        assert calc['W'] == pytest.approx(m.p.w._value)
        assert calc['H'] == pytest.approx(m.p.h._value)

    def test_wrapping_a_series_on_its_own_window_still_refuses(
            self, tmp_path):
        # The source starts in March, not with the sheet: its cells
        # cannot be placed under the sheet's dates, so the export must
        # keep refusing rather than write the wrapper as if the series
        # started with the window.
        m = _model(4)
        with m:
            m.p = mo.MultiVariable('P', excel_props={'tab': True})
            m.p.w = mo.Variable(
                mo.Variable([3.0, 4.0], start='2025-03', grain='month'),
                extend=mo.zero(), display_name='W',
            )
        with pytest.raises(ValueError, match='time window of its own'):
            m.to_excel(str(tmp_path / 'book.xlsx'))
