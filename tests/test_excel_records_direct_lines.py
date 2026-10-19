# SPDX-License-Identifier: Apache-2.0
"""A line declared directly on an ``orient='records'`` container — a
total under the table — is laid out beneath it, never dropped.

The records form lays out the container's RECORDS: a table of their
scalar fields, then a block per per-period field. A Variable set on
the container itself (``table.total = a.x + b.x``) belonged to neither
and vanished from the workbook without a word. It now follows the table
and the blocks as an ordinary row: a scalar in the sheet's scalar
column, a series in the period columns the blocks use, so the month
masthead stays where it was.
"""

import openpyxl

import modeleon as mo


def _grid(model, sheet, tmp_path, cols=6):
    path = tmp_path / "t.xlsx"
    model.to_excel(str(path))
    ws = openpyxl.load_workbook(str(path))[sheet]
    return [
        [ws.cell(r, c).value for c in range(1, cols + 1)]
        for r in range(1, ws.max_row + 1)
    ]


class Cust(mo.MultiVariableClass):
    def compute(self, raw):
        self.arr = mo.Variable(raw, display_name="ARR")


class Comp(mo.MultiVariableClass):
    def compute(self, ev, ebitda):
        self.ev = mo.Variable(ev, display_name="EV")
        self.ebitda = mo.Variable(ebitda, display_name="EBITDA")


def test_a_series_total_stands_under_the_blocks_in_the_period_columns(tmp_path):
    m = mo.Model("m", display_name="M", default_grain="month", default_start="2026-01",
                 default_periods=3)
    with m:
        m.active = mo.MultiVariable("Active", excel_props={"tab": True, "orient": "records"})
        with m.active as a:
            a.c1 = Cust(raw=[1.0, 2.0, 3.0], display_name="Customer 1")
            a.c2 = Cust(raw=[10.0, 20.0, 30.0], display_name="Customer 2")
            a.total = mo.Variable(a.c1.arr + a.c2.arr, display_name="Total ARR")
    g = _grid(m, "Active", tmp_path)
    assert g[0][1] == "Jan 2026"
    assert [row[0] for row in g] == [None, "ARR", "Customer 1", "Customer 2", "Total ARR"]
    assert g[2][1:4] == [1, 2, 3]
    assert g[4][1:4] == ["=B3 + B4", "=C3 + C4", "=D3 + D4"]


def test_a_scalar_total_stands_under_the_table(tmp_path):
    m = mo.Model("m", display_name="M")
    with m:
        m.comps = mo.MultiVariable("Comps", excel_props={"tab": True, "orient": "records"})
        with m.comps as c:
            c.a = Comp(ev=1200, ebitda=100, display_name="Alpha")
            c.b = Comp(ev=900, ebitda=90, display_name="Beta")
            c.total_ev = mo.Variable(c.a.ev + c.b.ev, display_name="Total EV")
    g = _grid(m, "Comps", tmp_path)
    assert g[0][:3] == [None, "EV", "EBITDA"]
    assert g[1][:3] == ["Alpha", 1200, 100]
    assert g[2][:3] == ["Beta", 900, 90]
    assert g[3][:3] == ["Total EV", "=B2 + B3", None]


def test_a_series_total_keeps_the_months_off_the_fields_table(tmp_path):
    class Project(mo.MultiVariableClass):
        def compute(self, price):
            self.price = mo.Variable(price, display_name="Price")
            self.profile = mo.Variable([0.2, 0.3, 0.5], display_name="Profile")
            self.revenue = self.profile * self.price

    m = mo.Model("m", display_name="M", default_grain="month", default_start="2026-01",
                 default_periods=3)
    with m:
        m.projects = mo.MultiVariable(
            "Projects", excel_props={"tab": True, "orient": "records"}
        )
        with m.projects as p:
            p.p1 = Project(price=100, display_name="Project 1")
            p.p2 = Project(price=200, display_name="Project 2")
            p.total = mo.Variable(p.p1.revenue + p.p2.revenue, display_name="Total revenue")
    g = _grid(m, "Projects", tmp_path)
    labels = [row[0] for row in g]
    # Column B is the table's "Price"; the months start past it, over the
    # blocks — and the total's cells stand in those same columns.
    assert g[0][1] is None and g[0][2] == "Jan 2026"
    i = labels.index("Revenue")
    assert labels[i + 1 : i + 4] == ["Project 1", "Project 2", "Total revenue"]
    assert g[i + 3][1] is None
    assert g[i + 3][2] == f"=C{i + 2} + C{i + 3}"


def test_a_container_of_only_direct_lines_lays_them_out(tmp_path):
    m = mo.Model("m", display_name="M")
    with m:
        m.empty = mo.MultiVariable("Empty", excel_props={"tab": True, "orient": "records"})
        m.empty.x = mo.Variable(1.0, display_name="X")
    assert _grid(m, "Empty", tmp_path)[0][:2] == ["X", 1]


def test_a_header_row_total_rides_the_section_row_and_no_other(tmp_path):
    # A ``header_row`` line is hoisted onto its section's row by the
    # parent's walk; the records form must not lay it out a second time.
    m = mo.Model("m", display_name="M")
    with m:
        m.book = mo.MultiVariable("Book", excel_props={"tab": True})
        with m.book as b:
            b.comps = mo.MultiVariable("Comps", excel_props={"orient": "records"})
            with b.comps as c:
                c.a = Comp(ev=1200, ebitda=100, display_name="Alpha")
                c.b = Comp(ev=900, ebitda=90, display_name="Beta")
                c.total_ev = mo.Variable(
                    c.a.ev + c.b.ev, display_name="Total EV", excel_props={"header_row": True}
                )
    g = _grid(m, "Book", tmp_path)
    assert g[0][:2] == ["Comps", "=B3 + B4"]
    cells = [v for row in g for v in row]
    assert "Total EV" not in cells
    assert cells.count("=B3 + B4") == 1
