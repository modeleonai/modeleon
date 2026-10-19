# SPDX-License-Identifier: Apache-2.0
"""A subscript lands on a cell, or the cell ships its number — never
the Python spelling.

``ufcf[-1]`` is the row's last period, as in Python; ``scen["Base"]``
is the keyed cell; an element of an unnamed line (``(x * 2)[1]``) is
its number inside a live formula. Anything the book has no cell for
takes the same route as every other fallback: the cell carries its
computed value. ``=inp.ufcf[-1] * 2`` used to be written into the cell,
a formula Excel reads as ``#NAME?``.
"""

import shutil
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook

import modeleon as mo
from modeleon.core.expr import Literal, MethodCall, Subscript, VarRef


def _rows(ws):
    out = {}
    for row in ws.iter_rows():
        label = row[0].value
        if label:
            out[str(label)] = [c.value for c in row[1:]]
    return out


def _book(model, tmp_path):
    path = tmp_path / "book.xlsx"
    model.to_excel(path)
    return load_workbook(path)


def _model():
    m = mo.Model("m", default_grain="year", default_start="2026", default_periods=3)
    with m:
        m.inp = mo.MultiVariable("Inp")
        m.inp.ufcf = mo.Variable([100.0, 110.0, 120.0], display_name="ufcf")
        m.inp.scen = mo.Variable({"Bear": 800.0, "Base": 1000.0, "Bull": 1200.0},
                                 display_name="scen")
        m.inp.byyear = mo.Variable([1.0, 2.0, 3.0], keys=[2026, 2027, 2028],
                                   display_name="byyear")
        m.out = mo.MultiVariable("Out")
    return m


def _cells(ws):
    return [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]


class TestANegativeIndexIsTheRealPeriod:
    def test_the_last_and_the_second_last_cell(self, tmp_path):
        m = _model()
        with m:
            m.out.last = (m.inp.ufcf[-1] * 2).set_display_name("last")
            m.out.second = (m.inp.ufcf[-2] + 1).set_display_name("second")
            m.out.first = (m.inp.ufcf[0] + 1).set_display_name("first")
        out = _rows(_book(m, tmp_path)["Out"])
        assert out["last"][0] == "=Inp!D2 * 2", out["last"]
        assert out["second"][0] == "=Inp!C2 + 1", out["second"]
        assert out["first"][0] == "=Inp!B2 + 1", out["first"]
        assert m.out.last.value == 240.0 and m.out.second.value == 111.0

    def test_inside_a_series_formula_every_period_reads_the_last_cell(self, tmp_path):
        m = _model()
        with m:
            m.out.share = (m.inp.ufcf / m.inp.ufcf[-1]).set_display_name("share")
        out = _rows(_book(m, tmp_path)["Out"])
        assert out["share"][0:3] == ["=Inp!B2 / Inp!D2", "=Inp!C2 / Inp!D2",
                                     "=Inp!D2 / Inp!D2"], out["share"]


class TestAKeyedSubscriptIsItsCell:
    """A keyed line stands under a caption row of its keys (row 3 names
    scen's, row 5 byyear's), so its cells are one row below its label's
    old place."""

    def test_by_label_and_by_year(self, tmp_path):
        m = _model()
        with m:
            m.out.base = (m.inp.scen["Base"] * 2).set_display_name("base")
            m.out.y27 = (m.inp.byyear[2027] * 2).set_display_name("y27")
        out = _rows(_book(m, tmp_path)["Out"])
        assert out["base"][0] == "=Inp!C4 * 2", out["base"]
        assert out["y27"][0] == "=Inp!C6 * 2", out["y27"]


class TestAnElementOfAnUnnamedLineIsItsNumber:
    def test_by_index_from_either_end(self, tmp_path):
        m = _model()
        with m:
            m.out.one = ((m.inp.ufcf * 2.0)[1] + 1).set_display_name("one")
            m.out.last = ((m.inp.ufcf * 2.0)[-1] + 1).set_display_name("last")
        out = _rows(_book(m, tmp_path)["Out"])
        assert out["one"][0] == "=220.0 + 1", out["one"]
        assert out["last"][0] == "=240.0 + 1", out["last"]

    def test_a_window_of_it_period_by_period(self, tmp_path):
        m = _model()
        with m:
            m.out.win = ((m.inp.ufcf * 2.0)[0:3] + 1).set_display_name("win")
        out = _rows(_book(m, tmp_path)["Out"])
        assert out["win"][0:3] == ["=200.0 + 1", "=220.0 + 1", "=240.0 + 1"], out["win"]


class TestWhatHasNoCellShipsItsValue:
    """Hand-built nodes the DSL never produces, so the renderer's own
    contract is what is pinned: a number in the cell, no Python text."""

    @staticmethod
    def _line(expr, value, name):
        v = mo.Variable()
        v._set_expr(expr)
        v._value = value
        v.var_type = "scalar"
        return v.set_display_name(name)

    def test_a_subscript_with_no_cell_and_a_method_with_no_translation(self, tmp_path):
        m = _model()
        with m:
            m.out.no_key = self._line(Subscript(VarRef(m.inp.ufcf), "zzz"), 7.0, "no_key")
            m.out.no_base = self._line(Subscript(Literal([1.0, 2.0]), 1), 2.0, "no_base")
            m.out.method = self._line(
                MethodCall(VarRef(m.inp.ufcf), "pct_change", [1], {}), 0.1, "method")
            m.out.shifted = self._line(
                MethodCall(Literal([1.0, 2.0]), "shift", [1], {}), 1.0, "shifted")
            m.out.in_formula = (self._line(
                Subscript(VarRef(m.inp.ufcf), "zzz"), 7.0, "k") * 2).set_display_name("in_formula")
        ws = _book(m, tmp_path)["Out"]
        out = _rows(ws)
        assert out["no_key"][0] == 7.0 and out["no_base"][0] == 2.0, out
        assert out["method"][0] == 0.1 and out["shifted"][0] == 1.0, out
        assert out["in_formula"][0] == 14.0, out["in_formula"]
        for text in _cells(ws):
            assert "[" not in text and ".pct_change(" not in text and ".shift(" not in text, text


def _soffice():
    for candidate in ("soffice", "libreoffice"):
        found = shutil.which(candidate)
        if found:
            return found
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


@pytest.mark.slow
def test_the_workbook_recalculates_to_the_python_values(tmp_path):
    soffice = _soffice()
    if soffice is None:
        pytest.skip("LibreOffice not installed")
    m = _model()
    with m:
        m.out.last = (m.inp.ufcf[-1] * 2).set_display_name("last")
        m.out.share = (m.inp.ufcf / m.inp.ufcf[-1]).set_display_name("share")
        m.out.base = (m.inp.scen["Base"] * 2).set_display_name("base")
    path = tmp_path / "sub.xlsx"
    m.to_excel(path)
    subprocess.run(
        [soffice, f"-env:UserInstallation={(tmp_path / 'prof').as_uri()}",
         "--headless", "--convert-to", "xlsx", "--outdir",
         str(tmp_path / "calc"), str(path)],
        check=True, capture_output=True,
    )
    out = _rows(load_workbook(tmp_path / "calc" / "sub.xlsx", data_only=True)["Out"])
    assert out["last"][0] == 240.0 and out["base"][0] == 2000.0, out
    assert [round(v, 6) for v in out["share"][1:4]] == [
        round(v, 6) for v in m.out.share.value], out["share"]
