# SPDX-License-Identifier: Apache-2.0
"""A keyed line gets a caption row of its own keys.

A keyed Variable (``mo.Variable({"FY2023": 400, "FY2024": 450})``) or a
line on one finite axis that is not time runs along KEYS, not periods.
Laid out on a period-headed sheet it used to sit under the year
captions with no caption of its own, so a reader took 400 under "2026"
for the 2026 figure. Such a line now stands under one reserved caption
row carrying its keys, each key over the cell that holds its value,
from the sheet's first period column; it is kept off the period
header's accounting (``keyed_var_ids``, not ``period_var_ids``), so the
header names the periods and nothing else. Consecutive keyed lines with
the same keys share one caption row.
"""

import shutil
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook

import modeleon as mo
from modeleon.compile.excel.layout import LayoutEngine


def _rows(model, sheet, tmp_path, n, cols="ABCDEFGH"):
    path = tmp_path / "t.xlsx"
    model.to_excel(str(path))
    ws = load_workbook(str(path))[sheet]
    return [[ws[f"{c}{r}"].value for c in cols] for r in range(1, n + 1)]


def _model(view=None):
    m = mo.Model("m", display_name="M", default_grain="year", default_start="2026",
                 default_periods=4)
    if view is not None:
        m.default_excel_view = view
    with m:
        m.inputs = mo.MultiVariable("Inputs", excel_props={"tab": True})
        with m.inputs as i:
            i.revenue = mo.Variable([100, 110, 121, 133], display_name="Revenue")
            i.hist = mo.Variable({"FY2023": 400, "FY2024": 450},
                                 display_name="Historical revenue")
            i.hist2 = mo.Variable({"FY2023": 40, "FY2024": 45},
                                  display_name="Historical EBITDA")
            i.growth = mo.Variable(0.1, display_name="Growth")
    return m


class TestTheBook:
    def test_the_keys_stand_over_the_values_and_the_header_names_the_periods_only(
        self, tmp_path,
    ):
        rows = _rows(_model(), "Inputs", tmp_path, 6)
        # The period header is the periods' — four years, not stretched or
        # shortened by the keyed lines beneath.
        assert rows[0] == [None, None, "2026", "2027", "2028", "2029", None, None]
        assert rows[1] == ["Revenue", None, 100, 110, 121, 133, None, None]
        # One caption row for the two lines that share the keys, each key
        # over its value, from the first period column.
        assert rows[2] == [None, None, "FY2023", "FY2024", None, None, None, None]
        assert rows[3] == ["Historical revenue", None, 400, 450, None, None, None, None]
        assert rows[4] == ["Historical EBITDA", None, 40, 45, None, None, None, None]
        assert rows[5] == ["Growth", 0.1, None, None, None, None, None, None]

    def test_lines_with_other_keys_get_a_caption_row_of_their_own(self, tmp_path):
        m = _model()
        with m.inputs as i:
            i.hist3 = mo.Variable({"FY2022": 1, "FY2023": 2, "FY2024": 3},
                                  display_name="Three years")
        rows = _rows(m, "Inputs", tmp_path, 8)
        assert rows[4] == ["Historical EBITDA", None, 40, 45, None, None, None, None]
        assert rows[5] == ["Growth", 0.1, None, None, None, None, None, None]
        assert rows[6] == [None, None, "FY2022", "FY2023", "FY2024", None, None, None]
        assert rows[7] == ["Three years", None, 1, 2, 3, None, None, None]

    def test_the_caption_row_wears_the_header_band(self, tmp_path):
        view = mo.ExcelView(bands={"header": {"bg": "#1F3864", "font_color": "#FFFFFF"}})
        path = tmp_path / "t.xlsx"
        _model(view).to_excel(str(path))
        ws = load_workbook(str(path))["Inputs"]
        assert ws["C3"].value == "FY2023"
        assert ws["C3"].font.b and ws["C3"].fill.fgColor.rgb == "FF1F3864"
        assert ws["C1"].fill.fgColor.rgb == "FF1F3864"

    def test_a_formula_reading_a_keyed_cell_follows_it(self, tmp_path):
        m = _model()
        with m:
            m.out = mo.MultiVariable("Out", excel_props={"tab": True})
            m.out.base = (m.inputs.hist["FY2024"] * 2).set_display_name("Base")
        rows = _rows(m, "Out", tmp_path, 2, cols="AB")
        assert rows[0] == ["Base", "=Inputs!D4 * 2"]

    def test_a_keyed_sheet_without_a_window_starts_with_its_caption_row(self, tmp_path):
        m = mo.Model("m", display_name="M")
        with m:
            m.hist = mo.MultiVariable("Hist", excel_props={"tab": True})
            m.hist.rev = mo.Variable({"FY2023": 400, "FY2024": 450}, display_name="Revenue")
        rows = _rows(m, "Hist", tmp_path, 3, cols="ABCD")
        assert rows[0] == [None, "FY2023", "FY2024", None]
        assert rows[1] == ["Revenue", 400, 450, None]
        assert rows[2] == [None, None, None, None]

    def test_a_line_on_a_finite_axis_is_captioned_by_the_axis_labels(self, tmp_path):
        m = mo.Model("m", display_name="M", default_grain="year", default_start="2026",
                     default_periods=3)
        scenarios = mo.Variable(["Bear", "Base", "Bull"])
        with m:
            m.s = mo.MultiVariable("S", excel_props={"tab": True})
            m.s.rev = mo.Variable([1, 2, 3], display_name="Revenue")
            m.s.share = mo.Variable([0.2, 0.5, 0.3], indexed_by=[scenarios],
                                    display_name="Share")
        rows = _rows(m, "S", tmp_path, 4, cols="ABCDE")
        assert rows[0] == [None, "2026", "2027", "2028", None]
        assert rows[1] == ["Revenue", 1, 2, 3, None]
        assert rows[2] == [None, "Bear", "Base", "Bull", None]
        assert rows[3] == ["Share", 0.2, 0.5, 0.3, None]

    def test_a_line_on_two_axes_is_still_refused(self, tmp_path):
        m = mo.Model("m", display_name="M")
        a = mo.Variable(["x", "y"])
        b = mo.Variable(["p", "q"])
        with m:
            m.s = mo.MultiVariable("S", excel_props={"tab": True})
            m.s.grid = mo.Variable([1, 2, 3, 4], indexed_by=[a, b], display_name="Grid")
        with pytest.raises(ValueError, match="2 axes"):
            m.to_excel(str(tmp_path / "t.xlsx"))


class TestTheLayout:
    def test_the_layout_publishes_the_caption_cells_and_keeps_the_line_off_the_periods(
        self,
    ):
        m = _model()
        engine = LayoutEngine([m.inputs], flatten_nested_sheets=True, header_rows=True)
        addresses = engine.compute_addresses()
        assert addresses["m.inputs.hist"].name == "A4"
        assert addresses["m.inputs.hist"].values == ["C4", "D4"]
        assert addresses["m.inputs.hist2"].values == ["C5", "D5"]
        assert engine.key_captions_by_sheet["Inputs"] == {"C3": "FY2023", "D3": "FY2024"}
        assert engine.key_captions_by_sheet[id(m.inputs)] is engine.key_captions_by_sheet[
            "Inputs"
        ]
        assert {"m.inputs.hist", "m.inputs.hist2"} <= engine.keyed_var_ids
        assert not {"m.inputs.hist", "m.inputs.hist2"} & engine.period_var_ids
        assert "m.inputs.revenue" in engine.period_var_ids
        assert engine.period_start_by_sheet["Inputs"] == 3

    def test_a_sheet_without_keyed_lines_publishes_nothing(self):
        m = mo.Model("m", display_name="M", default_grain="year", default_start="2026",
                     default_periods=2)
        with m:
            m.s = mo.MultiVariable("S", excel_props={"tab": True})
            m.s.rev = mo.Variable([1, 2], display_name="Revenue")
        engine = LayoutEngine([m.s], flatten_nested_sheets=True, header_rows=True)
        engine.compute_addresses()
        assert engine.key_captions_by_sheet == {}
        assert engine.keyed_var_ids == set()


class TestASheetWithSubtotalColumns:
    """A keyed line on a sheet with subtotal columns keeps every cell.

    Its cells run from the first period column across the bucket columns
    (it is off the period accounting, so the interleave does not move
    them), and it is no series: the totals pass writes no bucket on its
    row, where it used to put ``#VALUE!`` over the key that fell on a
    bucket column.
    """

    @staticmethod
    def _model():
        m = mo.Model("m", display_name="M", default_grain="month",
                     default_start="2026-01", default_periods=6)
        m.default_excel_view = mo.ExcelView(timeline={"totals": ["quarter", "year"]})
        with m:
            m.s = mo.MultiVariable("S", excel_props={"tab": True})
            with m.s as s:
                s.rev = mo.Variable([1, 2, 3, 4, 5, 6], display_name="Revenue",
                                    regrain=mo.up("sum"))
                s.hist = mo.Variable({"FY2022": 1, "FY2023": 2, "FY2024": 3, "FY2025": 4},
                                     display_name="Hist")
                s.chk = (s.hist["FY2025"] * 10).set_display_name("Check")
        return m

    def test_the_keyed_line_keeps_the_key_that_falls_on_a_bucket_column(self, tmp_path):
        rows = _rows(self._model(), "S", tmp_path, 7, cols="ABCDEFGHIJK")
        # Months C..E, Q1 at F, G..I, Q2 at J, the year at K.
        assert rows[2] == [None, None, "Jan", "Feb", "Mar", None, "Apr", "May", "Jun",
                           None, None]
        assert rows[3] == ["Revenue", None, 1, 2, 3, "=SUM(C4,D4,E4)", 4, 5, 6,
                           "=SUM(G4,H4,I4)", "=SUM(C4,D4,E4,G4,H4,I4)"]
        assert rows[4] == [None, None, "FY2022", "FY2023", "FY2024", "FY2025", None, None,
                           None, None, None]
        # Four keys, four cells — the fourth under the Q1 column, and no
        # bucket cell anywhere on the row.
        assert rows[5] == ["Hist", None, 1, 2, 3, 4, None, None, None, None, None]
        assert rows[6] == ["Check", "=F6 * 10", None, None, None, None, None, None, None,
                           None, None]

    @pytest.mark.slow
    def test_the_workbook_recalculates_to_the_python_values(self, tmp_path):
        soffice = _soffice()
        if soffice is None:
            pytest.skip("LibreOffice not installed")
        m = self._model()
        path = tmp_path / "keyed.xlsx"
        m.to_excel(str(path))
        subprocess.run(
            [soffice, f"-env:UserInstallation={(tmp_path / 'prof').as_uri()}",
             "--headless", "--convert-to", "xlsx", "--outdir",
             str(tmp_path / "calc"), str(path)],
            check=True, capture_output=True,
        )
        ws = load_workbook(tmp_path / "calc" / "keyed.xlsx", data_only=True)["S"]
        assert m.s.chk.value == 40
        assert [ws[f"{c}6"].value for c in "CDEFGHIJK"] == [1, 2, 3, 4, None, None, None,
                                                            None, None]
        assert ws["B7"].value == 40
        assert ws["K4"].value == 21


class TestAHiddenKeyedLine:
    def test_a_lone_hidden_line_hides_its_caption_row(self, tmp_path):
        m = mo.Model("m", display_name="M", default_grain="year", default_start="2026",
                     default_periods=3)
        with m:
            m.s = mo.MultiVariable("S", excel_props={"tab": True})
            with m.s as s:
                s.rev = mo.Variable([1, 2, 3], display_name="Revenue")
                s.hid = mo.Variable({"FY2023": 4, "FY2024": 5}, display_name="Hidden",
                                    excel_props={"hidden": True})
                s.after = mo.Variable([7, 8, 9], display_name="After")
        path = tmp_path / "t.xlsx"
        m.to_excel(str(path))
        ws = load_workbook(str(path))["S"]
        # Written in full, both rows natively hidden: no stray caption row
        # between Revenue and After.
        assert [ws[f"{c}3"].value for c in "ABC"] == [None, "FY2023", "FY2024"]
        assert [ws[f"{c}4"].value for c in "ABC"] == ["Hidden", 4, 5]
        assert sorted(r for r, d in ws.row_dimensions.items() if d.hidden) == [3, 4]

    def test_a_visible_sibling_keeps_the_shared_caption_row(self, tmp_path):
        m = mo.Model("m", display_name="M", default_grain="year", default_start="2026",
                     default_periods=3)
        with m:
            m.s = mo.MultiVariable("S", excel_props={"tab": True})
            with m.s as s:
                s.rev = mo.Variable([1, 2, 3], display_name="Revenue")
                s.hid = mo.Variable({"FY2023": 4, "FY2024": 5}, display_name="Hidden",
                                    excel_props={"hidden": True})
                s.vis = mo.Variable({"FY2023": 6, "FY2024": 7}, display_name="Visible")
        path = tmp_path / "t.xlsx"
        m.to_excel(str(path))
        ws = load_workbook(str(path))["S"]
        assert [ws[f"{c}5"].value for c in "ABC"] == ["Visible", 6, 7]
        assert sorted(r for r, d in ws.row_dimensions.items() if d.hidden) == [4]


class TestASheetOfKeyedLinesInASameColumnsBook:
    def test_the_sheet_shares_the_books_columns(self, tmp_path):
        m = mo.Model("m", display_name="M", default_grain="year", default_start="2026",
                     default_periods=3)
        m.default_excel_view = mo.ExcelView(sheet={"same_columns": True},
                                            meta={"constants": "Const"})
        with m:
            m.pnl = mo.MultiVariable("PnL", excel_props={"tab": True})
            m.pnl.rev = mo.Variable([1, 2, 3], display_name="Revenue")
            m.pnl.tax = mo.Variable(0.2, display_name="Tax")
            m.h = mo.MultiVariable("Hist", excel_props={"tab": True})
            m.h.hist = mo.Variable({"FY2023": 400, "FY2024": 450}, display_name="Hist")
            m.h.growth = mo.Variable(0.1, display_name="Growth")
        assert _rows(m, "PnL", tmp_path, 3, cols="ABCDE") == [
            [None, "Const", "2026", "2027", "2028"],
            ["Revenue", None, 1, 2, 3],
            ["Tax", 0.2, None, None, None],
        ]
        # The keyed sheet heads its columns like every other sheet of the
        # book: the constants column at B under its caption, the keys from
        # C where the periods stand elsewhere, in a caption row of their own.
        assert _rows(m, "Hist", tmp_path, 4, cols="ABCDE") == [
            [None, "Const", None, None, None],
            [None, None, "FY2023", "FY2024", None],
            ["Hist", None, 400, 450, None],
            ["Growth", 0.1, None, None, None],
        ]


def _soffice():
    for candidate in ("soffice", "libreoffice"):
        found = shutil.which(candidate)
        if found:
            return found
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None
