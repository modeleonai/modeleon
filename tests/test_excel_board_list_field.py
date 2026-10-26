# SPDX-License-Identifier: Apache-2.0
"""A field of several cells inside an instance of an ``orient='columns'``
board is refused at layout — never skipped.

The matrix gives every field ONE cell per instance. A per-period line,
a keyed line or a line on an axis declared inside an instance fitted
no cell, and was left out without a word: its row vanished from the
sheet, and every aggregate reading it — ``mo.SUM(self.revenue)`` — was
written as its computed number, a dead cell with a plausible look. The
layout now refuses such a field, naming the field, the instance and
the section, and says where the line does fit: outside the section, as
a row of the sheet, or — for a per-period field — under
``orient='records'``, where it gets a block of its own beneath the
table. The board's own direct lines, scalar and series alike, keep
their rows; a board of scalar fields lays out exactly as before.
"""

import warnings

import pytest
from openpyxl import load_workbook

import modeleon as mo


def _grid(model, sheet, tmp_path, cols=5):
    path = tmp_path / "t.xlsx"
    model.to_excel(str(path))
    ws = load_workbook(str(path))[sheet]
    return [
        [ws.cell(r, c).value for c in range(1, cols + 1)]
        for r in range(1, ws.max_row + 1)
    ]


def _labels(grid):
    return [row[0] for row in grid]


def _first_instance_column(grid):
    captions = next(row for row in grid if "P1" in row)
    first = captions.index("P1")
    assert captions[first:first + 2] == ["P1", "P2"]
    return first


def _model(orient="columns", *, view=False, field="period"):
    m = mo.Model("m", display_name="M", default_grain="month", default_start="2026-01",
                 default_periods=3)
    with m:
        if field == "axis":
            m.scen = mo.Variable(["Bear", "Base", "Bull"])
        m.sheet = mo.MultiVariable("Sheet", excel_props={"tab": True})
        with m.sheet as s:
            if view:
                s.board = mo.MultiVariable("Board")
                s.board.default_excel_view = mo.ExcelView(orient=orient)
            else:
                s.board = mo.MultiVariable("Board", excel_props={"orient": orient})
            with s.board as b:
                b.p1 = mo.MultiVariable("P1")
                b.p1.price = mo.Variable(10.0, display_name="Price")
                if field == "period":
                    b.p1.revenue = mo.Variable([1.0, 2.0, 3.0], display_name="Revenue")
                    b.p1.total = mo.Variable(mo.SUM(b.p1.revenue), display_name="Total")
                elif field == "keyed":
                    b.p1.growth = mo.Variable(
                        {"Bear": 0.9, "Base": 1.0, "Bull": 1.1}, display_name="Growth"
                    )
                elif field == "axis":
                    b.p1.growth = mo.Variable(
                        [0.9, 1.0, 1.1], indexed_by=[m.scen], display_name="Growth"
                    )
                b.p2 = mo.MultiVariable("P2")
                b.p2.price = mo.Variable(20.0, display_name="Price")
    return m


def _tracked_model(*, vol=True):
    m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=3,
                 tracks=mo.Tracks("plan", "actual"))
    with m:
        m.board = mo.MultiVariable("Board", excel_props={"tab": True, "orient": "columns"})
        with m.board as b:
            b.p1 = mo.MultiVariable("P1")
            b.p1.rate = mo.Variable(plan=5, display_name="Rate")
            if vol:
                b.p1.vol = mo.Variable(plan=[1, 2, 3], display_name="Vol")
            b.p2 = mo.MultiVariable("P2")
            b.p2.rate = mo.Variable(plan=6, display_name="Rate")
    return m


def test_a_per_period_field_in_an_instance_is_refused_with_both_fixes(tmp_path):
    path = tmp_path / "t.xlsx"
    with warnings.catch_warnings(record=True) as rec:
        warnings.simplefilter("always")
        with pytest.raises(ValueError) as exc:
            _model().to_excel(str(path))
    msg = str(exc.value)
    assert "'Revenue' in 'P1' is a line of 3 cell(s), one per period" in msg
    assert "'Board' is laid out with orient='columns'" in msg
    assert "outside 'Board'" in msg
    assert "orient='records'" in msg
    # Refused at layout, before the workbook exists: nothing is written.
    assert not path.exists()
    # ...and before any aggregate over the line is rendered as a dead
    # number — the warning that path fires never does.
    assert not [w for w in rec if w.category.__name__ == "CrossScopeReferenceWarning"]


@pytest.mark.parametrize("field", ["keyed", "axis"])
def test_a_keyed_field_in_an_instance_is_refused_toward_the_sheet(tmp_path, field):
    with pytest.raises(ValueError) as exc:
        _model(field=field).to_excel(str(tmp_path / "t.xlsx"))
    msg = str(exc.value)
    assert "'Growth' in 'P1' is a line of 3 cell(s), one per key" in msg
    assert "under a caption row of its keys" in msg
    # Under orient='records' a keyed field gets a block of bare values
    # with no caption row of its keys — a half layout, so not offered.
    assert "orient='records'" not in msg


def test_a_tracked_series_field_in_an_instance_is_refused_too(tmp_path):
    with pytest.raises(ValueError) as exc:
        _tracked_model().to_excel(str(tmp_path / "t.xlsx"))
    # After the writer expands the tracks, the head row keeps the name
    # 'Vol' and holds the list.
    assert "'Vol' in 'P1' is a line of 3 cell(s), one per period" in str(exc.value)


def test_the_view_spelling_of_the_block_is_refused_the_same_way(tmp_path):
    with pytest.raises(ValueError) as exc:
        _model(view=True).to_excel(str(tmp_path / "t.xlsx"))
    msg = str(exc.value)
    assert "'Revenue' in 'P1'" in msg
    assert "orient='columns'" in msg


def test_a_board_of_scalar_fields_lays_out_as_before(tmp_path):
    g = _grid(_model(field="none"), "Sheet", tmp_path)
    first = _first_instance_column(g)
    assert g[_labels(g).index("Price")][first:first + 2] == [10, 20]


def test_a_tracked_scalar_field_stays_a_cell(tmp_path):
    # A tracked SCALAR is a number once the writer expands the tracks,
    # not a list: it keeps its cell in the matrix.
    g = _grid(_tracked_model(vol=False), "Board", tmp_path)
    assert g[0][:3] == [None, "P1", "P2"]
    assert g[1][:3] == ["Rate", 5, 6]


def test_the_same_records_lay_out_with_orient_records(tmp_path):
    # The advice the refusal gives is true: under orient='records' the
    # per-period field gets a block of its own beneath the table — its
    # caption, then one row per instance, the values past the table's
    # columns — and the table's Total is a live SUM over that block,
    # not the dead number 6.
    g = _grid(_model(orient="records"), "Sheet", tmp_path, cols=7)
    labels = _labels(g)
    i = labels.index("Revenue")
    assert labels[i + 1] == "P1"
    assert g[i + 1][4:7] == [1, 2, 3]
    assert g[labels.index("P1")][2] == f"=SUM(E{i + 2}:G{i + 2})"


def test_the_layout_records_the_refusal_and_other_readers_keep_their_addresses(tmp_path):
    from modeleon.compile.excel.layout import LayoutEngine
    m = _model()
    engine = LayoutEngine([m.sheet], flatten_nested_sheets=True, header_rows=True)
    addresses = engine.compute_addresses()
    assert addresses, "the rest of the model keeps its addresses"
    assert len(engine.refusals) == 1 and "'Revenue' in 'P1'" in engine.refusals[0]
    # the notebook repr still renders (it never raises on a layout refusal)
    assert "P1" in m._repr_html_()
    # the compatibility check reports it instead of raising
    issues = mo.check_compat(m, "excel")
    layout_issues = [i for i in issues if i.kind == "layout"]
    assert len(layout_issues) == 1 and "'Revenue' in 'P1'" in str(layout_issues[0])


def test_a_one_value_bare_list_field_is_a_cell_like_everywhere_else(tmp_path):
    m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=3)
    m.board = mo.MultiVariable("Board", excel_props={"tab": True, "orient": "columns"})
    m.board.p1 = mo.MultiVariable("P1")
    m.board.p1.fee = mo.Variable([5.0], display_name="Fee")
    m.board.p1.fee2 = m.board.p1.fee + 1
    m.board.p2 = mo.MultiVariable("P2")
    m.board.p2.fee = mo.Variable([6.0], display_name="Fee")
    m.board.p2.fee2 = m.board.p2.fee + 1
    g = _grid(m, "Board", tmp_path, cols=3)
    assert g[1][:3] == ["Fee", 5, 6]
    assert g[2][:3] == ["Fee2", "=B2 + 1", "=C2 + 1"]
