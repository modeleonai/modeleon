# SPDX-License-Identifier: Apache-2.0
"""Subtotal columns — «янв фев мар (1 кв) … (год)» — and they are ALIVE.

``ExcelView(timeline={'totals': ['quarter', 'year']})`` interleaves a
quarter column after each quarter's last month and a year column after
the year. The parity promise extends to them: edit January in the
written file and the quarter, the year, and every formula standing
on them move. That rules out baked numbers everywhere a formula can
honestly stand:

* a rule line   → ``=SUM(месяцы)`` / ``=AVERAGE`` / last-month ref;
* a formula line → its OWN formula over the operands' bucket cells —
  exactly what a hand-built book writes at the quarter column;
* a ruleless literal → the projection's ``#VALUE!`` teaching token,
  never an invented sum.
"""

import os
import tempfile
from datetime import date, timedelta

import openpyxl
import pytest
from openpyxl.utils import get_column_letter

import modeleon as mo


def _book(model):
    path = os.path.join(tempfile.mkdtemp(), "t.xlsx")
    model.to_excel(path)
    return openpyxl.load_workbook(path)


def _tab(marked):
    """A first-level container is a tab either way: marked with
    ``excel_props={'tab': True}``, or unmarked — the default, where the
    writer makes the tab itself. The totals must not tell them apart."""
    return {"excel_props": {"tab": True}} if marked else {}


# Every content expectation below runs on both kinds of tab.
both_tabs = pytest.mark.parametrize(
    "marked", [True, False], ids=["marked-tab", "unmarked-tab"]
)


def _model(marked=True, **view_kw):
    m = mo.Model(
        "m",
        display_name="M",
        default_excel_view=mo.ExcelView(
            timeline={"totals": ["quarter", "year"], **view_kw}
        ),
    )
    m.default_start, m.default_grain, m.default_periods = "2026-01", "month", 6
    with m:
        m.план = mo.MultiVariable("План", **_tab(marked))
        with m.план as p:
            p.ставка = mo.Variable(0.1, display_name="Ставка")
            p.выручка = mo.Variable(
                [10, 20, 30, 40, 50, 60],
                display_name="Выручка",
                regrain=mo.up("sum"),
            )
            p.налог = p.выручка * p.ставка
    return m


class TestTheBookForm:
    def test_the_header_is_hierarchical(self):
        # «Год» over «кварталы» over «месяцы» — one header row per
        # grain. A bucket spans its months and the finer totals inside
        # it, and its OWN total stands beside that span, named, from
        # the bucket's own tier downwards: «Итого Q1» begins where «Q1»
        # begins, not a tier lower where it read as a thirteenth month.
        ws = _book(_model())["План"]
        assert ws.cell(1, 3).value == "2026"
        assert ws.cell(2, 3).value == "Q1"
        assert ws.cell(2, 7).value == "Q2"
        # The quarter total names itself ON the quarter row…
        assert ws.cell(2, 6).value == "Total Q1"
        assert ws.cell(2, 10).value == "Total Q2"
        # …so the month row carries months and nothing else.
        assert [ws.cell(3, c).value for c in range(3, 12)] == [
            "Jan", "Feb", "Mar", None,
            "Apr", "May", "Jun", None, None,
        ]
        # The year total starts on the YEAR row and runs the full stack.
        assert ws.cell(1, 11).value == "Total 2026"
        merges = sorted(str(r) for r in ws.merged_cells.ranges)
        assert "K1:K3" in merges
        assert "C1:J1" in merges          # the year: months + Q-totals
        assert "C2:E2" in merges          # Q1: its months only
        assert "F2:F3" in merges          # its total, two rows tall
        assert "G2:I2" in merges

    def test_data_starts_below_the_header_stack(self):
        ws = _book(_model())["План"]
        assert ws.cell(4, 1).value == "Ставка"
        assert ws.freeze_panes == "C4"

    @both_tabs
    def test_a_rule_line_sums_its_months_live(self, marked):
        ws = _book(_model(marked))["План"]
        assert ws["F5"].value == "=SUM(C5,D5,E5)"
        assert ws["J5"].value == "=SUM(G5,H5,I5)"

    @both_tabs
    def test_the_year_never_swallows_the_quarter_cells(self, marked):
        # The year's months are NOT contiguous once quarter columns
        # stand between them — a range SUM would double-count. Explicit
        # refs, months only.
        ws = _book(_model(marked))["План"]
        assert ws["K5"].value == "=SUM(C5,D5,E5,G5,H5,I5)"

    @both_tabs
    def test_a_formula_line_writes_its_own_formula_at_the_bucket(self, marked):
        # The hand-built book's rule: налог(Q1) = выручка(Q1) × ставка,
        # standing on the QUARTER cell — not a sum of monthly taxes,
        # not a baked number.
        ws = _book(_model(marked))["План"]
        assert ws["F6"].value == "=F5 * B4"
        assert ws["J6"].value == "=J5 * B4"
        assert ws["K6"].value == "=K5 * B4"

    @both_tabs
    def test_native_months_still_reference_their_own_cells(self, marked):
        # The interleave moves ADDRESSES, so native formulas follow.
        ws = _book(_model(marked))["План"]
        assert ws["C6"].value == "=C5 * B4"
        assert ws["G6"].value == "=G5 * B4"

    @both_tabs
    def test_a_ruleless_literal_teaches_instead_of_inventing(self, marked):
        m = mo.Model(
            "m",
            display_name="M",
            default_excel_view=mo.ExcelView(
                timeline={"totals": ["quarter"]}
            ),
        )
        m.default_start, m.default_grain, m.default_periods = (
            "2026-01", "month", 3,
        )
        with m:
            m.план = mo.MultiVariable("План", **_tab(marked))
            with m.план as p:
                p.курс = mo.Variable(
                    [480, 481, 482], display_name="Курс"
                )
        ws = _book(m)["План"]
        # No rule declared: summing an FX rate would be a lie; the
        # projection's token lands instead.
        assert ws["E3"].value == "#VALUE!"

    @both_tabs
    def test_partial_trailing_quarter_still_totals(self, marked):
        m = mo.Model(
            "m",
            display_name="M",
            default_excel_view=mo.ExcelView(
                timeline={"totals": ["quarter"]}
            ),
        )
        m.default_start, m.default_grain, m.default_periods = (
            "2026-01", "month", 4,
        )
        with m:
            m.план = mo.MultiVariable("План", **_tab(marked))
            with m.план as p:
                p.поток = mo.Variable(
                    [1, 2, 3, 4], display_name="Поток",
                    regrain=mo.up("sum"),
                )
        ws = _book(m)["План"]
        # Q1 after March; the lone April closes a partial Q2.
        assert ws["E3"].value == "=SUM(B3,C3,D3)"
        assert ws["G3"].value == "=SUM(F3)"

    def test_undeclared_totals_change_nothing(self):
        m = _model()
        m.default_excel_view = mo.ExcelView()
        ws = _book(m)["План"]
        row1 = [ws.cell(1, c).value for c in range(1, 10)]
        assert "Q1 2026" not in row1
        assert ws["C3"].value == 10  # contiguous months, untouched


class TestOtherRules:
    def _one_line(self, marked, **var_kw):
        m = mo.Model(
            "m",
            display_name="M",
            default_excel_view=mo.ExcelView(
                timeline={"totals": ["quarter"]}
            ),
        )
        m.default_start, m.default_grain, m.default_periods = (
            "2026-01", "month", 3,
        )
        with m:
            m.план = mo.MultiVariable("План", **_tab(marked))
            with m.план as p:
                p.линия = mo.Variable(
                    [10, 20, 30], display_name="Линия", **var_kw
                )
        return _book(m)["План"]

    @both_tabs
    def test_last(self, marked):
        # A balance's quarter is its closing month — a ref, not a sum.
        ws = self._one_line(marked, regrain=mo.up("last"))
        assert ws["E3"].value == "=D3"

    @both_tabs
    def test_mean(self, marked):
        ws = self._one_line(marked, regrain=mo.up("mean"))
        assert ws["E3"].value == "=AVERAGE(B3,C3,D3)"


def _grid(ws):
    return [
        [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        for r in range(1, ws.max_row + 1)
    ]


class TestEveryTabTotalsAlike:
    """A tab is a tab however it came to be: marked with
    ``excel_props={'tab': True}``, made by the writer from an unmarked
    first-level container (the default), the overview tab that collects
    Variables placed straight on the model, or the tab of a container
    written on its own with ``container.to_excel()``. Each one's total
    columns hold the same cells."""

    def _lines(self, p):
        p.ставка = mo.Variable(0.1, display_name="Ставка")
        p.выручка = mo.Variable(
            [10, 20, 30, 40, 50, 60], display_name="Выручка",
            regrain=mo.up("sum"),
        )
        p.налог = p.выручка * p.ставка
        p.курс = mo.Variable([480] * 6, display_name="Курс")
        p.остаток = mo.Variable(
            [1, 2, 3, 4, 5, 6], display_name="Остаток", regrain=mo.up("last"),
        )
        p.если = mo.IF(p.выручка > 25, p.выручка, 0)

    @staticmethod
    def _empty(display_name="M"):
        m = mo.Model(
            "m", display_name=display_name,
            default_excel_view=mo.ExcelView(
                timeline={"totals": ["quarter", "year"]}
            ),
        )
        m.default_start, m.default_grain, m.default_periods = (
            "2026-01", "month", 6,
        )
        return m

    def _model(self, marked, section=True):
        m = self._empty()
        with m:
            m.план = mo.MultiVariable("План", **_tab(marked))
            with m.план as p:
                self._lines(p)
                if section:
                    p.детали = mo.MultiVariable("Детали")
                    with p.детали as d:
                        d.двойной = p.налог * 2
        return m

    def test_an_unmarked_tab_writes_every_cell_a_marked_tab_writes(self):
        marked = _grid(_book(self._model(True))["План"])
        unmarked = _grid(_book(self._model(False))["План"])
        assert unmarked == marked
        # …and the cells are the live ones, not a shared emptiness.
        ws = _book(self._model(False))["План"]
        rows = {ws.cell(r, 1).value: r for r in range(1, ws.max_row + 1)}
        налог, курс = rows["Налог"], rows["Курс"]
        assert ws[f"F{налог}"].value == f"=F{rows['Выручка']} * B{rows['Ставка']}"
        assert ws[f"K{курс}"].value == "#VALUE!"
        assert ws[f"F{rows['Двойной']}"].value == f"=F{налог} * 2"

    def test_the_overview_tab_of_root_variables(self):
        # Lines placed straight on the model land on a tab named after
        # it — the same tab a marked container of those lines makes.
        m = self._empty("M")
        with m:
            self._lines(m)
        overview = _grid(_book(m)["M"])

        tabbed = self._empty("Книга")
        with tabbed:
            tabbed.m = mo.MultiVariable("M", excel_props={"tab": True})
            with tabbed.m as p:
                self._lines(p)
        assert overview == _grid(_book(tabbed)["M"])
        assert overview[5][5] == "=F5 * B4"      # Налог, Total Q1
        assert overview[6][10] == "#VALUE!"      # Курс, Total 2026

    @both_tabs
    def test_a_container_written_on_its_own(self, marked):
        # ``m.план.to_excel()`` writes the container's own lines onto one
        # tab — the same tab the whole model's book carries. (A section
        # inside it would become a tab of its own there, so none here.)
        whole = _grid(_book(self._model(True, section=False))["План"])
        m2 = self._model(marked, section=False)
        path = os.path.join(tempfile.mkdtemp(), "t.xlsx")
        m2.план.to_excel(path)
        alone = openpyxl.load_workbook(path)["План"]
        assert _grid(alone) == whole
        assert alone["F6"].value == "=F5 * B4"


class TestTheTotalsCalculate:
    """Recalculated by a spreadsheet program, a formula line's total
    column shows the number Python gets by re-graining the model:
    ``model.at('quarter')`` — the line's own formula over the quarter's
    operands, never the sum of its monthly results."""

    @staticmethod
    def _recalculated(model, tmp_path):
        import shutil
        import subprocess
        from pathlib import Path

        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
        if soffice is None and mac.exists():
            soffice = str(mac)
        if soffice is None:
            pytest.skip("LibreOffice is not installed")
        src = tmp_path / "book.xlsx"
        model.to_excel(str(src))
        outdir = tmp_path / "calc"
        profile = (tmp_path / "lo-profile").as_uri()  # private per run
        subprocess.run(
            [soffice, f"-env:UserInstallation={profile}", "--headless",
             "--convert-to", "xlsx", "--outdir", str(outdir), str(src)],
            check=True, capture_output=True, timeout=120,
        )
        return openpyxl.load_workbook(str(outdir / "book.xlsx"), data_only=True)

    @both_tabs
    def test_a_ratio_totals_to_the_ratio_of_totals(self, marked, tmp_path):
        m = _model(marked)
        with m.план as p:
            p.затраты = mo.Variable(
                [9.0, 12, 15, 38, 30, 33], display_name="Затраты",
                regrain=mo.up("sum"),
            )
            p.маржа = (p.выручка - p.затраты) / p.выручка
        ws = self._recalculated(m, tmp_path)["План"]
        row = next(r for r in range(1, ws.max_row + 1)
                   if ws.cell(r, 1).value == "Маржа")
        quarters = m.at("quarter").план.маржа.value
        year = m.at("year").план.маржа.value
        # Q1 = (60 − 36) / 60; the months' own margins would add up to 1.0.
        assert quarters[0] == pytest.approx(0.4)
        assert ws.cell(row, 6).value == pytest.approx(quarters[0])   # F
        assert ws.cell(row, 10).value == pytest.approx(quarters[1])  # J
        assert ws.cell(row, 11).value == pytest.approx(year[0])      # K


class TestTheTotalsWord:
    def test_the_model_speaks_its_own_language(self):
        m = mo.Model(
            "m",
            display_name="M",
            default_excel_view=mo.ExcelView(
                timeline={"totals": ["quarter"], "totals_word": "Итого"}
            ),
        )
        m.default_start, m.default_grain, m.default_periods = (
            "2026-01", "month", 3,
        )
        with m:
            m.план = mo.MultiVariable("План", excel_props={"tab": True})
            with m.план as p:
                p.поток = mo.Variable(
                    [1, 2, 3], display_name="Поток", regrain=mo.up("sum"),
                )
        ws = _book(m)["План"]
        # One tier above the months, so the quarter's own total names
        # itself on that tier and merges down through the month row.
        assert ws.cell(1, 5).value == "Итого Q1"
        assert "E1:E2" in {str(r) for r in ws.merged_cells.ranges}

    def test_tier_spans_read_from_the_left(self):
        # «2026» centered floats far from the row names — and the
        # month row stays centered under it.
        ws = _book(_model())["План"]
        assert ws.cell(1, 3).alignment.horizontal == "left"
        assert ws.cell(2, 3).alignment.horizontal == "left"
        assert ws.cell(3, 3).alignment.horizontal == "center"


class TestNativeColumnGroups:
    def test_the_file_carries_real_outline_levels(self):
        # Excel's own «свернуть 2027»: months sit two buckets deep,
        # a quarter's total one (inside the year), the year's total
        # outside every group — and the summary stands RIGHT of its
        # detail, where our totals are.
        ws = _book(_model())["План"]
        def level(c):
            return (ws.column_dimensions[c].outline_level
                    if c in ws.column_dimensions else 0)

        assert [level(c) for c in "CDEF"] == [2, 2, 2, 1]
        assert level("K") == 0  # «Итого 2026» survives every fold
        assert ws.sheet_properties.outlinePr.summaryRight is True


class TestTotalsOverride:
    def test_the_override_replaces_and_strips_the_declaration(self):
        # A caller-side override: LayoutEngine(totals_override=...) beats
        # the declared timeline.totals — ('quarter',) narrows, () strips,
        # None keeps. ``to_excel`` never passes it.
        from modeleon.compile.excel.layout import LayoutEngine

        m = _model()
        roots = [m.план]
        eng = LayoutEngine(roots, totals_override=("quarter",))
        eng.compute_addresses()
        plan = eng.totals_by_sheet.get("План")
        assert plan is not None and plan.kinds == ("quarter",)

        eng2 = LayoutEngine(roots, totals_override=())
        eng2.compute_addresses()
        assert eng2.totals_by_sheet.get("План") is None

        eng3 = LayoutEngine(roots)
        eng3.compute_addresses()
        assert eng3.totals_by_sheet.get("План").kinds == ("quarter", "year")


class TestRangesNeverSwallowATotal:
    """A row's cells stop being contiguous the moment totals declare a
    column between them — and ``SUM(C3:I3)`` then counts the quarter
    twice, in a file that looks perfectly valid.

    ``totals.rule_formula`` lists explicit refs for its own buckets;
    the same rule holds wherever a range is built.
    """

    def _book(self, totals, n=6, fn=None):
        import os
        import tempfile

        import openpyxl

        view = mo.ExcelView(timeline={"totals": totals}) if totals else None
        m = mo.Model("m", display_name="M",
                     **({"default_excel_view": view} if view else {}))
        m.default_start, m.default_grain, m.default_periods = (
            "2026-01", "month", n,
        )
        with m:
            m.p = mo.MultiVariable("P", excel_props={"tab": True})
            with m.p as p:
                p.поток = mo.Variable([10.0] * n, display_name="Поток",
                                      regrain=mo.up("sum"))
                p.всего = mo.Variable((fn or mo.SUM)(p.поток),
                                      display_name="Всего")
        path = os.path.join(tempfile.mkdtemp(), "t.xlsx")
        m.to_excel(path)
        ws = openpyxl.load_workbook(path)["P"]
        row = next(r for r in range(1, ws.max_row + 1)
                   if ws.cell(r, 1).value == "Всего")
        return ws.cell(row, 2).value

    def test_an_interleaved_row_lists_its_cells(self):
        # Months C,D,E then the Q1 total at F, then G,H,I. The range
        # form would add F twice over: 90 for a model that says 60.
        formula = self._book(["quarter"])
        assert ":" not in formula
        assert formula.replace(" ", "") == "=SUM(C3,D3,E3,G3,H3,I3)"

    def test_an_interleaved_row_on_another_sheet_qualifies_every_cell(self):
        # A list of cells is not a range: a sheet prefix binds to one
        # cell only. Unqualified, every cell after the first would be
        # read from the CURRENT sheet and the sum would silently be wrong.
        import os
        import tempfile

        import openpyxl

        m = mo.Model("m", display_name="M", default_excel_view=mo.ExcelView(
            timeline={"totals": ["quarter"]}))
        m.default_start, m.default_grain, m.default_periods = (
            "2026-01", "month", 6,
        )
        with m:
            m.p = mo.MultiVariable("P", excel_props={"tab": True})
            with m.p as p:
                p.поток = mo.Variable([10.0] * 6, display_name="Поток",
                                      regrain=mo.up("sum"))
            m.s = mo.MultiVariable("S", excel_props={"tab": True})
            with m.s as s:
                s.всего = mo.Variable(mo.SUM(m.p.поток), display_name="Всего")
        path = os.path.join(tempfile.mkdtemp(), "t.xlsx")
        m.to_excel(path)
        ws = openpyxl.load_workbook(path)["S"]
        row = next(r for r in range(1, ws.max_row + 1)
                   if ws.cell(r, 1).value == "Всего")
        formula = next(ws.cell(row, c).value for c in range(2, ws.max_column + 1)
                       if str(ws.cell(row, c).value or "").startswith("="))
        assert formula.startswith("=SUM(") and ":" not in formula
        cells = [a.strip() for a in formula[len("=SUM("):-1].split(",")]
        assert len(cells) == 6
        assert all(c.startswith("P!") for c in cells), formula

    def test_a_contiguous_row_keeps_its_range(self):
        # Nothing stands between the months — the compact form is both
        # correct and what a reader expects to see.
        assert self._book(None) == "=SUM(C2:H2)"

    def test_a_range_only_function_tells_the_truth_instead(self):
        # IRR's second argument is a GUESS, so a comma list would be
        # read as one. A formula that computes something else is worse
        # than the number itself: the cell falls back to its value.
        import os
        import tempfile

        import openpyxl

        m = mo.Model("m", display_name="M",
                     default_excel_view=mo.ExcelView(
                         timeline={"totals": ["quarter"]}))
        m.default_start, m.default_grain, m.default_periods = (
            "2026-01", "month", 6,
        )
        with m:
            m.p = mo.MultiVariable("P", excel_props={"tab": True})
            with m.p as p:
                p.поток = mo.Variable([-100.0, 30.0, 30.0, 30.0, 30.0, 30.0],
                                      display_name="Поток",
                                      regrain=mo.up("sum"))
                p.ставка = mo.Variable(mo.IRR(p.поток), display_name="Ставка")
        path = os.path.join(tempfile.mkdtemp(), "t.xlsx")
        m.to_excel(path)
        ws = openpyxl.load_workbook(path)["P"]
        row = next(r for r in range(1, ws.max_row + 1)
                   if ws.cell(r, 1).value == "Ставка")
        assert not str(ws.cell(row, 2).value).startswith("=")

    # A SLICE of the row (``SUM(x[0:4])``: January to April) is the same
    # question over fewer cells: its window still steps over the Q1
    # total, and the range form would count it in.

    def _sliced(self, fn, flows, *, other_sheet=False):
        """Write ``Итог = fn(поток)`` next to the flows, or on a tab of
        its own; return the book, what the «Итог» cell holds, and the
        value the model computed for it."""
        m = mo.Model("m", display_name="M", default_excel_view=mo.ExcelView(
            timeline={"totals": ["quarter"]}))
        m.default_start, m.default_grain, m.default_periods = (
            "2026-01", "month", len(flows),
        )
        with m:
            m.p = mo.MultiVariable("P", excel_props={"tab": True})
            with m.p as p:
                p.поток = mo.Variable(flows, display_name="Поток",
                                      regrain=mo.up("sum"))
                if not other_sheet:
                    p.итог = mo.Variable(fn(p.поток), display_name="Итог")
            if other_sheet:
                m.s = mo.MultiVariable("S", excel_props={"tab": True})
                with m.s as s:
                    s.итог = mo.Variable(fn(m.p.поток), display_name="Итог")
        home = m.s if other_sheet else m.p
        path = os.path.join(tempfile.mkdtemp(), "t.xlsx")
        m.to_excel(path)
        wb = openpyxl.load_workbook(path)
        ws = wb["S" if other_sheet else "P"]
        row = next(r for r in range(1, ws.max_row + 1)
                   if ws.cell(r, 1).value == "Итог")
        return wb, ws.cell(row, 2).value, home.итог.value

    @staticmethod
    def _sum_of_listed_cells(wb, formula, sheet):
        # Read ``=SUM(a, b, …)`` the way Excel does: each argument is
        # one cell, on its own sheet when it names one.
        assert formula.startswith("=SUM(") and ":" not in formula, formula
        total = 0.0
        for ref in formula[len("=SUM("):-1].split(","):
            where, _, addr = ref.strip().rpartition("!")
            total += wb[where or sheet][addr].value
        return total

    def test_a_sliced_window_lists_its_cells(self):
        # C,D,E are Jan–Mar, F the Q1 total, G April. ``C3:G3`` would
        # read 160 for a model that says 100.
        wb, formula, value = self._sliced(
            lambda x: mo.SUM(x[0:4]), [10.0, 20.0, 30.0, 40.0, 50.0, 60.0])
        assert value == 100.0
        assert formula.replace(" ", "") == "=SUM(C3,D3,E3,G3)"
        assert self._sum_of_listed_cells(wb, formula, "P") == value

    def test_a_sliced_window_on_another_sheet_qualifies_every_cell(self):
        wb, formula, value = self._sliced(
            lambda x: mo.SUM(x[0:4]), [10.0, 20.0, 30.0, 40.0, 50.0, 60.0],
            other_sheet=True)
        cells = [a.strip() for a in formula[len("=SUM("):-1].split(",")]
        assert len(cells) == 4
        assert all(c.startswith("P!") for c in cells), formula
        assert self._sum_of_listed_cells(wb, formula, "S") == value == 100.0

    def test_a_sliced_window_inside_one_quarter_keeps_its_range(self):
        wb, formula, value = self._sliced(
            lambda x: mo.SUM(x[3:6]), [10.0, 20.0, 30.0, 40.0, 50.0, 60.0])
        assert formula == "=SUM(G3:I3)"
        assert value == 150.0

    def test_a_sliced_window_inside_one_quarter_on_another_sheet(self):
        # A range names its sheet once, in front of the first cell.
        _wb, formula, value = self._sliced(
            lambda x: mo.SUM(x[3:6]), [10.0, 20.0, 30.0, 40.0, 50.0, 60.0],
            other_sheet=True)
        assert formula == "=SUM(P!F3:H3)"
        assert value == 150.0

    def test_a_stepped_window_lists_its_cells(self):
        # Every other month is never side by side, totals or not.
        wb, formula, value = self._sliced(
            lambda x: mo.SUM(x[0:6:2]), [10.0, 20.0, 30.0, 40.0, 50.0, 60.0])
        assert value == 90.0
        assert self._sum_of_listed_cells(wb, formula, "P") == value

    def test_a_sliced_window_for_a_range_only_function_is_its_value(self):
        # ``IRR(C3:H3, 0.1)`` would take the Q1 total as a cash flow.
        _wb, cell, value = self._sliced(
            lambda x: mo.IRR(x[0:5]), [-100.0, 30.0, 30.0, 30.0, 30.0, 30.0])
        assert not str(cell).startswith("="), cell
        assert cell == pytest.approx(value)


class TestADurationIsADayCount:
    """A spreadsheet stores a date as a number of days, so a duration
    added to a date is a plain number of days in the formula — never
    Python's ``5 days, 0:00:00``, which no spreadsheet can read."""

    def _formula(self, delta):
        m = mo.Model("m", display_name="M")
        with m:
            m.p = mo.MultiVariable("P", excel_props={"tab": True})
            with m.p as p:
                p.старт = mo.Variable(date(2026, 1, 1), display_name="Старт")
                p.срок = mo.Variable(p.старт + delta, display_name="Срок")
        ws = _book(m)["P"]
        assert ws["A1"].value == "Старт" and ws["A2"].value == "Срок"
        return ws["B2"].value

    def test_whole_days(self):
        assert self._formula(timedelta(days=90)) == "=B1 + 90"

    def test_part_of_a_day(self):
        assert self._formula(timedelta(hours=12)) == "=B1 + 0.5"

    def test_a_negative_duration(self):
        assert self._formula(timedelta(days=-5)) == "=B1 + -5"


class TestNativeOutlineGroups:
    """A collapsible outline, native to the file — both axes."""

    def test_rows_group_by_section_and_the_header_is_the_summary(self):
        m = mo.Model("m", display_name="M")
        m.default_start, m.default_grain, m.default_periods = (
            "2026-01", "month", 3,
        )
        with m:
            m.p = mo.MultiVariable("P", excel_props={"tab": True})
            with m.p as p:
                p.доходы = mo.MultiVariable("I. Доходы")
                with p.доходы as d:
                    d.выручка = mo.Variable([1, 2, 3], display_name="Выручка")
                    d.детали = mo.MultiVariable("Детали")
                    with d.детали as dd:
                        dd.курс = mo.Variable([480] * 3, display_name="Курс")
        ws = _book(m)["P"]
        rows = {ws.cell(r, 1).value: r for r in range(1, ws.max_row + 1)}
        # Section content is one level deep; nested content two.
        assert ws.row_dimensions[rows["Выручка"]].outline_level == 1
        assert ws.row_dimensions[rows["    Курс"] if "    Курс" in rows
                                 else rows[next(k for k in rows if k and "Курс" in k)]
                                 ].outline_level == 2
        # The header row is the SUMMARY, outside its own group…
        assert ws.row_dimensions[rows["I. Доходы"]].outline_level in (0, None)
        # …and Excel puts the ± on it because the summary sits above.
        assert ws.sheet_properties.outlinePr.summaryBelow is False

    def test_column_brackets_cover_the_whole_group(self):
        # Compare mode: a month is a GROUP of columns, and the outline
        # bracket must cover all of them — one column in four reads as
        # noise. Hidden operands sit one level deeper so expanding a
        # quarter does not unhide them.
        m = mo.Model(
            "m", display_name="M", default_grain="month",
            default_start="2026-01", default_periods=3,
            tracks=mo.Tracks(
                "план", "факт",
                blend=mo.blend(given="факт", follow="план",
                               until="2026-01", name="live")),
            default_excel_view=mo.ExcelView(
                tracks="compare",
                timeline={"totals": ["quarter"]}),
        )
        with m:
            m.p = mo.MultiVariable("P", excel_props={"tab": True})
            with m.p as p:
                p.поток = mo.Variable(план=[10.0] * 3,
                                      факт=[12.0, None, None],
                                      display_name="Поток",
                                      regrain=mo.up("sum"))
        ws = _book(m)["P"]
        # Group = факт | план | live | откл (stride 4); every month
        # column of January's group carries the month level.
        lv = {c: ws.column_dimensions[get_column_letter(c)].outline_level
              for c in range(2, 10)}
        assert lv[2] == lv[3] == lv[4] == lv[5] == 1
