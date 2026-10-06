# SPDX-License-Identifier: Apache-2.0
"""Names in any script: sheets, lines and tracks.

The non-Latin names below are deliberate — they are what these tests
check."""

from pathlib import Path

from openpyxl import load_workbook

import modeleon as mo


def test_a_non_latin_sheet_is_quoted_in_references(tmp_path: Path) -> None:
    root = mo.MultiVariable("root")
    root.a = mo.MultiVariable("Выручка")
    root.a.x = mo.Variable([1.0, 2.0], display_name="X")
    root.pl = mo.MultiVariable("PL")
    root.pl.y = mo.Variable(root.a.x * 2, display_name="Y")
    out = tmp_path / "unicode.xlsx"
    root.to_excel(out)
    ws = load_workbook(out)["PL"]
    y = next([c.value for c in r[1:] if c.value is not None]
             for r in ws.iter_rows() if r[0].value == "Y")
    assert y[0].startswith("='Выручка'!B") and y[0].endswith(" * 2")


def test_non_latin_lines_and_tracks_read_by_their_names() -> None:
    m = mo.Model("m", tracks=mo.Tracks("план", "факт"), default_grain="month",
                 default_start="2026-01", default_periods=2)
    with m:
        m.s = mo.MultiVariable("S")
        m.s.доход = mo.Variable(план=[1.0, 2.0], факт=[3.0, 4.0])
        m.s.двойной = m.s.доход * 2
    assert m.s.доход.факт.value == [3.0, 4.0]
    assert m.s.двойной.value["факт"] == [6.0, 8.0]
    assert m.s.двойной.formula == "доход * 2"
