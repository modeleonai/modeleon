# SPDX-License-Identifier: Apache-2.0
"""``SUM``, ``AVERAGE``, ``PRODUCT``, ``COUNT`` and ``COUNTA`` follow the one
rule every aggregate follows.

Several series are combined period by period, as ``a + b`` combines them:
``mo.SUM(a, b)`` is ``=SUM(C5, C9)`` in every column, scalars broadcast,
series of other lengths are refused, located series on other time windows
are aligned by date, a line with tracks is taken track by track, inside a
loop container too. One series alone is reduced over its
periods (``mo.SUM(row)`` is ``=SUM(C5:N5)``), and scalars only stay one
number (``mo.SUM(a, b)`` over two input cells is ``=SUM(B2, B3)``) — the
rule ``MIN`` / ``MAX`` and the statistics follow. The total of several
rows over all their periods is ``mo.SUM(a) + mo.SUM(b)``.
"""

import shutil
import subprocess
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

import modeleon as mo
from modeleon.functions.registry import admits_in_loop, per_period_arity

FIVE = ("SUM", "AVERAGE", "PRODUCT", "COUNT", "COUNTA")


def _rows(ws):
    out = {}
    for row in ws.iter_rows():
        label = row[0].value
        if label:
            out[str(label)] = [c.value for c in row[1:]]
    return out


def _book(model, tmp_path, sheet):
    path = tmp_path / "book.xlsx"
    model.to_excel(path)
    return _rows(load_workbook(path)[sheet])


def _model():
    m = mo.Model("m", default_grain="year", default_start="2026", default_periods=3)
    with m:
        m.i = mo.MultiVariable("I")
        m.i.floor = mo.Variable(4.0, display_name="floor")
        m.i.p = mo.Variable(3.0, display_name="p")
        m.i.q = mo.Variable(9.0, display_name="q")
        m.x = mo.MultiVariable("X")
        m.x.a = mo.Variable([3.0, 8.0, 5.0], display_name="a")
        m.x.b = mo.Variable([4.0, 4.0, 4.0], display_name="b")
    return m


def _a():
    return mo.Variable([3.0, 8.0, 5.0])


def _b():
    return mo.Variable([4.0, 4.0, 4.0])


def _windows_inside():
    """A model from January with sub-containers on windows inside its
    own: T from March (a flow continued by zero, a rate held, a scalar),
    U from January ending early, O of one period — one cell."""
    m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=4)
    with m:
        m.x = mo.MultiVariable("X")
        m.x.rev = mo.Variable([1.0, 2.0, 3.0, 4.0], display_name="rev")
        m.t = mo.MultiVariable("T", default_grain="month", default_start="2026-03",
                               default_periods=2)
        m.t.fee = mo.Variable([10.0, 20.0], display_name="fee", extend=mo.zero())
        m.t.rate = mo.Variable([0.5, 0.25], display_name="rate", extend=mo.hold())
        m.t.k = mo.Variable(7.0, display_name="k")
        m.u = mo.MultiVariable("U", default_grain="month", default_start="2026-01",
                               default_periods=2)
        m.u.early = mo.Variable([100.0, 200.0], display_name="early", extend=mo.zero())
        m.u.cap = mo.Variable([9.0, 8.0], display_name="cap", extend=mo.hold())
        m.o = mo.MultiVariable("O", default_grain="month", default_start="2026-03",
                               default_periods=1)
        m.o.once = mo.Variable([50.0], display_name="once", extend=mo.zero())
        m.o.level = mo.Variable([5.0], display_name="level", extend=mo.hold())
        m.x.s = mo.SUM(m.x.rev, m.t.fee).set_display_name("s")
        m.x.plus = (m.x.rev + m.t.fee).set_display_name("plus")
        m.x.av = mo.AVERAGE(m.x.rev, m.t.fee, m.t.rate).set_display_name("av")
        m.x.pr = mo.PRODUCT(m.x.rev, m.t.rate).set_display_name("pr")
        m.x.hp = (m.x.rev * m.t.rate).set_display_name("hp")
        m.x.mx = mo.MAX(m.x.rev, m.t.fee).set_display_name("mx")
        m.x.mn = mo.MIN(m.x.rev, m.u.cap).set_display_name("mn")
        m.x.n = mo.COUNT(m.x.rev, m.t.fee, m.u.early).set_display_name("n")
        m.x.na = mo.COUNTA(m.x.rev, m.o.once).set_display_name("na")
        m.x.sk = mo.SUM(m.x.rev, m.t.k).set_display_name("sk")
        m.x.se = mo.SUM(m.x.rev, m.u.early).set_display_name("se")
        m.x.so = mo.SUM(m.x.rev, m.o.once).set_display_name("so")
        m.x.sl = mo.SUM(m.x.rev, m.o.level).set_display_name("sl")
        m.x.pw = mo.POWER(m.x.rev, m.t.rate).set_display_name("pw")
    return m


_INSIDE = ("s", "plus", "av", "pr", "hp", "mx", "mn", "n", "na", "sk", "se", "so", "sl", "pw")


# ══════════════════════════════════════════════════════════════════════
# The values
# ══════════════════════════════════════════════════════════════════════


class TestSeveralSeriesArePerPeriod:
    def test_two_series(self):
        a, b = _a(), _b()
        assert mo.SUM(a, b).value == [7.0, 12.0, 9.0]
        assert mo.AVERAGE(a, b).value == [3.5, 6.0, 4.5]
        assert mo.PRODUCT(a, b).value == [12.0, 32.0, 20.0]
        assert mo.COUNT(a, b).value == [2, 2, 2]
        assert mo.COUNTA(a, b).value == [2, 2, 2]
        for name in FIVE:
            assert getattr(mo, name)(a, b).var_type == "list", name

    def test_a_scalar_broadcasts_on_either_side(self):
        a = _a()
        assert mo.SUM(a, 10.0).value == [13.0, 18.0, 15.0]
        assert mo.SUM(10.0, a).value == [13.0, 18.0, 15.0]
        assert mo.SUM(a, mo.Variable(1.0)).value == [4.0, 9.0, 6.0]
        assert mo.AVERAGE(a, 1.0).value == [2.0, 4.5, 3.0]
        assert mo.PRODUCT(2.0, a).value == [6.0, 16.0, 10.0]
        assert mo.COUNT(a, 5).value == [2, 2, 2]
        assert mo.COUNTA(mo.Variable("x"), a).value == [2, 2, 2]

    def test_a_list_typed_in_is_a_series(self):
        a = _a()
        assert mo.SUM(a, [1.0, 2.0, 3.0]).value == [4.0, 10.0, 8.0]
        assert mo.SUM([1.0, 2.0, 3.0], [10.0, 20.0, 30.0]).value == [11.0, 22.0, 33.0]
        assert mo.AVERAGE([1.0, 2.0, 3.0], a).value == [2.0, 5.0, 4.0]
        assert mo.PRODUCT(a, [1.0, 0.5, 2.0]).value == [3.0, 4.0, 10.0]

    def test_three_operands(self):
        a, b = _a(), _b()
        assert mo.SUM(a, b, 1.0).value == [8.0, 13.0, 10.0]
        assert mo.AVERAGE(a, b, mo.Variable(3.0)).value == pytest.approx(
            [10 / 3, 5.0, 4.0])
        assert mo.COUNT(a, b, 1.0).value == [3, 3, 3]

    def test_nested_in_arithmetic(self):
        a, b = _a(), _b()
        assert (mo.SUM(a, b) * 2).value == [14.0, 24.0, 18.0]
        assert mo.SUM(a * 2, b).value == [10.0, 20.0, 14.0]
        assert (mo.SUM(a, b) - mo.AVERAGE(a, b)).value == [3.5, 6.0, 4.5]

    def test_a_one_value_series_broadcasts_as_the_operators_do(self):
        a = _a()
        assert mo.SUM(mo.Variable([4.0]), a).value == [7.0, 12.0, 9.0]
        assert mo.COUNT(mo.Variable([4.0]), a).value == [2, 2, 2]

    @pytest.mark.parametrize("name", FIVE)
    def test_series_of_different_lengths_are_refused(self, name):
        with pytest.raises(ValueError, match="different lengths"):
            getattr(mo, name)(mo.Variable([1.0, 2.0]), _a())

    def test_the_formula_names_the_operands(self):
        m = _model()
        assert mo.SUM(m.x.a, m.x.b).formula == "SUM(a, b)"
        assert mo.AVERAGE(m.x.a, 0).formula == "AVERAGE(a, 0)"
        assert mo.PRODUCT(m.x.a, m.x.b).formula == "PRODUCT(a, b)"
        assert mo.COUNT(m.x.a, m.x.b).formula == "COUNT(a, b)"
        assert mo.COUNTA(m.x.a, "x").formula == "COUNTA(a, 'x')"

    def test_value_types(self):
        whole = mo.Variable([3, 8, 5])
        assert mo.SUM(whole, 4).value_type == "int"
        assert mo.SUM(whole, 4).value == [7, 12, 9]
        assert mo.AVERAGE(whole, 4).value_type == "float"
        assert mo.PRODUCT(whole, 2).value_type == "int"
        assert mo.PRODUCT(whole, 0.5).value_type == "float"
        assert mo.COUNT(whole, 4).value_type == "int"
        assert mo.COUNTA(whole, 4).value_type == "int"


class TestOneSeriesOrScalarsOnlyAreUnchanged:
    def test_one_series_is_reduced_over_its_periods(self):
        a = _a()
        assert mo.SUM(a).value == 16.0
        assert mo.AVERAGE(a).value == pytest.approx(16 / 3)
        assert mo.PRODUCT(a).value == 120.0
        assert mo.COUNT(a).value == 3
        assert mo.COUNTA(a).value == 3
        for name in FIVE:
            assert getattr(mo, name)(a).var_type == "scalar", name

    def test_a_list_typed_in_alone_is_a_range(self):
        assert mo.SUM([1.0, 2.0, 3.0]).value == 6.0
        assert mo.AVERAGE([1.0, 2.0, 3.0]).value == 2.0
        assert mo.PRODUCT([1.0, 2.0, 3.0]).value == 6.0
        assert mo.COUNT([1.0, "x", None]).value == 1
        assert mo.COUNTA([1.0, "x", None]).value == 2

    def test_scalars_only_are_one_number(self):
        total = mo.SUM(mo.Variable(3.0), mo.Variable(4.0))
        assert total.value == 7.0 and total.var_type == "scalar"
        assert mo.SUM(1, 2, 3).formula == "SUM(1, 2, 3)"
        assert mo.AVERAGE(mo.Variable(3.0), 5.0).value == 4.0
        assert mo.PRODUCT(mo.Variable(3.0), 5.0).value == 15.0
        assert mo.COUNT(1, "x", None).value == 2
        assert mo.COUNTA(1, "x", None).value == 3

    def test_a_reduction_as_an_operand_is_a_scalar(self):
        a = _a()
        out = mo.SUM(mo.SUM(a), 1.0)
        assert out.value == 17.0 and out.var_type == "scalar"

    def test_the_total_of_several_rows_is_the_sum_of_their_totals(self):
        a, b = _a(), _b()
        total = mo.SUM(a) + mo.SUM(b)
        assert total.value == 28.0 and total.var_type == "scalar"
        assert mo.SUM(mo.SUM(a, b)).value == 28.0


class TestErrorsBlanksAndText:
    def test_an_error_cell_is_the_period_s_result(self):
        e = mo.Variable([2.0, 2.0, 2.0]) / mo.Variable([1.0, 0.0, 2.0])
        assert mo.SUM(e, 1.0).value == [3.0, "#DIV/0!", 2.0]
        assert mo.AVERAGE(1.0, e).value == [1.5, "#DIV/0!", 1.0]
        assert mo.PRODUCT(e, 3.0).value == [6.0, "#DIV/0!", 3.0]

    def test_the_first_error_in_reading_order_wins(self):
        x = mo.Variable(["#N/A", 1.0, 2.0])
        y = mo.Variable(["#DIV/0!", "#VALUE!", 3.0])
        assert mo.SUM(x, y).value == ["#N/A", "#VALUE!", 5.0]
        assert mo.SUM(y, x).value == ["#DIV/0!", "#VALUE!", 5.0]

    def test_a_count_reads_an_error_as_a_cell(self):
        x = mo.Variable(["#N/A", 1.0, 2.0])
        assert mo.COUNT(x, _b()).value == [1, 2, 2]
        assert mo.COUNTA(x, _b()).value == [2, 2, 2]

    def test_a_blank_is_skipped(self):
        x = mo.Variable([None, 5.0, None])
        y = mo.Variable([None, None, 2.0])
        assert mo.SUM(x, y).value == [0, 5.0, 2.0]
        assert mo.AVERAGE(x, y).value == ["#DIV/0!", 5.0, 2.0]
        assert mo.PRODUCT(x, y).value == [0, 5.0, 2.0]
        assert mo.COUNT(x, y).value == [0, 1, 1]
        assert mo.COUNTA(x, y).value == [0, 1, 1]

    @pytest.mark.parametrize("name", ["SUM", "AVERAGE", "PRODUCT", "MIN", "MAX"])
    def test_text_in_a_line_is_refused_when_the_call_is_built(self, name):
        # The spreadsheet skips text in a referenced cell: no marker would
        # match the workbook, so the call is refused, as over one row —
        # with a message naming the function and the value.
        with pytest.raises(TypeError, match=rf"{name} \w+ numbers; 'x' is not one"):
            getattr(mo, name)(_a(), mo.Variable(["x", 1.0, 2.0]))
        with pytest.raises(TypeError, match=rf"{name} \w+ numbers; 'x' is not one"):
            getattr(mo, name)(_a(), "x")
        with pytest.raises(TypeError, match="give it the number"):
            getattr(mo, name)(mo.Variable([1.0, "x"]))

    def test_a_none_typed_into_the_call_is_the_0_the_formula_holds(self, tmp_path):
        # The renderer writes a typed None as 0 (=MIN(C2, 0)), so that is
        # what the call reads, as the counts read it; a line's blank cell
        # is skipped, as the sheet skips it.
        a = _a()
        assert mo.SUM(a, None).value == [3.0, 8.0, 5.0]
        assert mo.AVERAGE(a, None).value == [1.5, 4.0, 2.5]
        assert mo.MIN(a, None).value == [0, 0, 0]
        assert mo.MAX(a, None).value == [3.0, 8.0, 5.0]
        assert mo.PRODUCT(a, None).value == [0.0, 0.0, 0.0]
        assert mo.AVERAGE(a, [1.0, None, 3.0]).value == [2.0, 4.0, 4.0]
        assert mo.SUM(mo.Variable([3.0, None, 5.0]), a).value == [6.0, 8.0, 10.0]
        # Over one row or scalars only, the same reading.
        assert mo.AVERAGE(mo.Variable(4.0), None).value == 2.0
        assert mo.MIN(mo.Variable(4.0), None).value == 0
        assert mo.AVERAGE([1.0, None, 3.0]).value == pytest.approx(4 / 3)
        assert mo.AVERAGE(mo.Variable([1.0, None, 3.0])).value == 2.0
        m = _model()
        with m:
            m.x.mn = mo.MIN(m.x.a, None).set_display_name("mn")
        assert _book(m, tmp_path, "X")["mn"] == ["=MIN(B2, 0)", "=MIN(C2, 0)", "=MIN(D2, 0)"]

    def test_a_count_reads_a_line_s_cells_and_the_call_s_values_as_the_sheet_does(self):
        # In a line, a boolean and text are no numbers; typed into the call,
        # TRUE, numeric text and None (the 0 the formula holds) count.
        line = mo.Variable([True, 1.0, "7"])
        assert mo.COUNT(line, True).value == [1, 2, 1]
        assert mo.COUNT(line, "7").value == [1, 2, 1]
        assert mo.COUNT(line, None).value == [1, 2, 1]
        assert mo.COUNT(line, [None, "x", "2"]).value == [1, 1, 1]
        assert mo.COUNT(mo.Variable([date(2026, 1, 1), None, "x"]), _b()).value == [2, 1, 1]

    def test_counta_reads_empty_text_as_the_workbook_holds_it(self):
        # Typed into a line, empty text is written as nothing: a blank.
        # A formula's empty text is a cell, and counts.
        typed = mo.Variable(["", "x", 1.0])
        assert mo.COUNTA(typed, _b()).value == [1, 2, 2]
        a = _a()
        shown = mo.IF(a > 4.0, "big", "")
        assert shown.value == ["", "big", "big"]
        assert mo.COUNTA(shown, _b()).value == [2, 2, 2]
        assert mo.COUNTA(_b(), "").value == [2, 2, 2]


class TestTracks:
    def _tracked(self):
        m = mo.Model("m", tracks=mo.Tracks("plan", "actual"),
                     default_grain="month", default_start="2026-01", default_periods=3)
        with m:
            m.x = mo.Variable(plan=[1.0, 5.0, 3.0], actual=[2.0, None, 2.0])
            m.y = mo.Variable(plan=[1.0, 1.0, 1.0], actual=[3.0, 3.0, 3.0])
        return m

    def test_tracks_are_taken_track_by_track(self):
        m = self._tracked()
        assert dict(mo.SUM(m.x, m.y)._value.items()) == {
            "plan": [2.0, 6.0, 4.0], "actual": [5.0, 3.0, 5.0]}
        assert dict(mo.SUM(m.x, 1.0)._value.items()) == {
            "plan": [2.0, 6.0, 4.0], "actual": [3.0, 1.0, 3.0]}
        assert dict(mo.AVERAGE(m.x, m.y)._value.items()) == {
            "plan": [1.0, 3.0, 2.0], "actual": [2.5, 3.0, 2.5]}
        assert dict(mo.PRODUCT(m.x, m.y)._value.items()) == {
            "plan": [1.0, 5.0, 3.0], "actual": [6.0, 3.0, 6.0]}
        assert dict(mo.COUNT(m.x, m.y)._value.items()) == {
            "plan": [2, 2, 2], "actual": [2, 1, 2]}
        assert dict(mo.COUNTA(m.x, 0)._value.items()) == {
            "plan": [2, 2, 2], "actual": [2, 1, 2]}
        assert mo.SUM(m.x, m.y).var_type == "list"

    def test_one_line_with_tracks_alone_is_still_its_totals(self):
        m = self._tracked()
        assert dict(mo.SUM(m.x)._value.items()) == {"plan": 9.0, "actual": 4.0}


class TestUnits:
    def test_sum_and_average_keep_the_one_unit(self):
        a = mo.Variable([3.0, 8.0, 5.0], unit="$")
        b = mo.Variable([4.0, 4.0, 4.0], unit="$")
        c = mo.Variable([4.0, 4.0, 4.0], unit="hr")
        assert str(mo.SUM(a, b).unit) == "$"
        assert str(mo.AVERAGE(a, 0.0).unit) == "$"
        with pytest.raises(ValueError, match="mo.SUM takes its values in one unit"):
            mo.SUM(a, c)
        with pytest.raises(ValueError, match="mo.AVERAGE takes its values in one unit"):
            mo.AVERAGE(c, a)

    def test_a_product_per_period_multiplies_its_operands_units(self):
        price = mo.Variable([10.0, 11.0, 12.0], unit="USD/t")
        volume = mo.Variable([2.0, 3.0, 4.0], unit="t")
        revenue = mo.PRODUCT(price, volume)
        assert revenue.value == [20.0, 33.0, 48.0]
        assert str(revenue.unit) == "USD"
        assert str(mo.PRODUCT(volume, 2.0).unit) == "t"
        # One row of amounts alone is still refused: its cells' product is
        # in tonnes to the power of their count.
        with pytest.raises(ValueError, match="to the power of their count"):
            mo.PRODUCT(volume)

    def test_a_count_has_no_unit(self):
        a = mo.Variable([3.0, 8.0, 5.0], unit="$")
        assert mo.COUNT(a, a).unit is None
        assert mo.COUNTA(a, 1.0).unit is None


class TestLinesAlongKeys:
    """A line along keys is no period series: against one it is refused,
    as ``a + b`` refuses it; lines along the same keys combine key by key."""

    @staticmethod
    def _model():
        m = mo.Model("m", default_grain="year", default_start="2026", default_periods=3)
        with m:
            m.s = mo.Variable([1.0, 2.0, 3.0], display_name="s")
            m.k = mo.Variable({"Bear": 0.9, "Base": 1.0, "Bull": 1.1}, display_name="k")
            m.k2 = mo.Variable({"Bear": 1.0, "Base": 2.0, "Bull": 3.0}, display_name="k2")
            m.k3 = mo.Variable({"Bull": 1.0, "Base": 2.0, "Bear": 3.0}, display_name="k3")
        return m

    @pytest.mark.parametrize("name", FIVE + ("MIN", "MAX"))
    def test_against_a_period_series_it_is_refused_as_a_plus_b_refuses_it(self, name):
        m = self._model()
        with pytest.raises(ValueError, match="runs along its keys"):
            m.s + m.k
        with pytest.raises(ValueError, match="runs along its keys"):
            getattr(mo, name)(m.s, m.k)

    def test_lines_along_the_same_keys_combine_key_by_key(self, tmp_path):
        m = self._model()
        with m:
            m.t = mo.SUM(m.k, m.k2).set_display_name("t")
            m.n = mo.COUNT(m.k, m.k2, None).set_display_name("n")
        assert m.t.value == pytest.approx([1.9, 3.0, 4.1])
        assert m.t._keys == ["Bear", "Base", "Bull"] == (m.k + m.k2)._keys
        assert m.n.value == [3, 3, 3] and m.n._keys == ["Bear", "Base", "Bull"]
        assert mo.MAX(m.k, 1.0).value == [1.0, 1.0, 1.1]
        # One line along keys alone is still its total.
        assert mo.SUM(m.k).value == pytest.approx(3.0)

    def test_lines_along_keys_in_another_order_combine_key_by_key(self):
        m = self._model()
        # k3 is Bull 1, Base 2, Bear 3: read in k's order, as ``a + b`` reads it.
        assert mo.SUM(m.k, m.k3).value == pytest.approx((m.k + m.k3).value)
        assert mo.SUM(m.k, m.k3).value == pytest.approx([3.9, 3.0, 2.1])
        assert mo.COUNT(m.k, m.k3)._keys == ["Bear", "Base", "Bull"]
        m4 = mo.Variable({"Bear": 1.0, "Base": 2.0, "Mid": 3.0}, display_name="k4")
        with pytest.raises(ValueError, match="other keys"):
            mo.SUM(m.k, m4)


# ══════════════════════════════════════════════════════════════════════
# Inside a loop container
# ══════════════════════════════════════════════════════════════════════


class TestLinesOnOtherWindows:
    """Located series on other time windows are aligned by date, as ``a +
    b`` aligns them: each continued by its own ``extend=`` rule, one that
    does not cover the combined window refused, the result on that
    window. ``MIN`` / ``MAX`` paired them by position before."""

    @staticmethod
    def _windows():
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=4)
        with m:
            m.x = mo.MultiVariable("X")
            m.x.rev = mo.Variable([1.0, 2.0, 3.0, 4.0], display_name="rev", extend=mo.zero())
            m.t = mo.MultiVariable("T", default_grain="month", default_start="2026-03",
                                   default_periods=4)
            m.t.fee = mo.Variable([10.0, 20.0, 30.0, 40.0], display_name="fee",
                                  extend=mo.zero())
        return m

    def test_each_is_the_call_over_the_date_s_cells_as_its_operator_is(self):
        m = self._windows()
        rev, fee = m.x.rev, m.t.fee
        plus = rev + fee
        assert plus.value == [1.0, 2.0, 13.0, 24.0, 30.0, 40.0]
        assert mo.SUM(rev, fee).value == plus.value
        assert mo.AVERAGE(rev, fee).value == (plus / 2).value
        assert mo.PRODUCT(rev, fee).value == (rev * fee).value
        assert mo.MAX(rev, fee).value == [1.0, 2.0, 10.0, 20.0, 30.0, 40.0]
        assert mo.MIN(rev, fee).value == [0.0, 0.0, 3.0, 4.0, 0.0, 0.0]
        # A continued cell holds the 0 the sheet writes for it: counted.
        assert mo.COUNT(rev, fee).value == [2] * 6
        assert mo.COUNTA(rev, fee).value == [2] * 6
        for name in FIVE + ("MIN", "MAX"):
            assert getattr(mo, name)(rev, fee).time == plus.time, name

    @pytest.mark.parametrize("name", FIVE + ("MIN", "MAX"))
    def test_a_window_not_covered_is_refused_as_a_plus_b_refuses_it(self, name):
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=4)
        with m:
            m.x = mo.MultiVariable("X")
            m.x.rev = mo.Variable([1.0, 2.0, 3.0, 4.0], display_name="rev")
            m.t = mo.MultiVariable("T", default_grain="month", default_start="2026-03",
                                   default_periods=4)
            m.t.fee = mo.Variable([10.0, 20.0, 30.0, 40.0], display_name="fee")
        with pytest.raises(ValueError, match="'rev' does not cover the combined window"):
            m.x.rev + m.t.fee
        with pytest.raises(ValueError, match="'rev' does not cover the combined window"):
            getattr(mo, name)(m.x.rev, m.t.fee)

    def test_each_of_three_windows_continues_by_its_own_rule(self):
        m = self._windows()
        with m:
            m.u = mo.MultiVariable("U", default_grain="month", default_start="2026-05",
                                   default_periods=2)
            m.u.rate = mo.Variable([100.0, 200.0], display_name="rate", extend=mo.hold())
        s = mo.SUM(m.x.rev, m.t.fee, m.u.rate)
        assert s.value == [101.0, 102.0, 113.0, 124.0, 130.0, 240.0]
        assert s.time.start == "2026-01"

    def test_a_one_value_line_is_a_constant_as_the_operators_read_it(self):
        # One value on a window of three months is read as a
        # constant: every period reads it, as its one cell is written.
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=3)
        with m:
            m.x = mo.MultiVariable("X")
            m.x.a = mo.Variable([1.0, 2.0, 3.0], display_name="a")
            m.x.one = mo.Variable([5.0], display_name="one")
            m.x.held = mo.Variable([5.0], display_name="held", extend=mo.hold())
            m.o = mo.MultiVariable("O", default_grain="month", default_start="2026-02",
                                   default_periods=1)
            m.o.feb = mo.Variable([5.0], display_name="feb")
        assert mo.SUM(m.x.a, m.x.one).value == (m.x.a + m.x.one).value == [6.0, 7.0, 8.0]
        assert mo.SUM(m.x.a, m.x.held).value == (m.x.a + m.x.held).value == [6.0, 7.0, 8.0]
        # On a window of one period it is that period's value.
        with pytest.raises(ValueError, match="'feb' does not cover"):
            m.x.a + m.o.feb
        with pytest.raises(ValueError, match="'feb' does not cover"):
            mo.SUM(m.x.a, m.o.feb)

    def test_two_grains_are_refused(self):
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=3)
        with m:
            m.x = mo.MultiVariable("X")
            m.x.a = mo.Variable([1.0, 2.0, 3.0], display_name="a")
            m.q = mo.MultiVariable("Q", default_grain="quarter", default_start="2026-Q1",
                                   default_periods=3)
            m.q.b = mo.Variable([1.0, 2.0, 3.0], display_name="b")
        with pytest.raises(ValueError, match="different grains"):
            mo.SUM(m.x.a, m.q.b)
        with pytest.raises(ValueError, match="different grains"):
            mo.COUNT(m.x.a, m.q.b)

    def test_the_workbook_reads_the_cells_the_sheet_holds(self, tmp_path):
        # A line that starts with the window and ends early is written
        # over the window, continued as it declares: the call reads those
        # cells, live, and computes what Python computes.
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=4)
        with m:
            m.x = mo.MultiVariable("X")
            m.x.rev = mo.Variable([1.0, 2.0, 3.0, 4.0], display_name="rev")
            m.x.short = mo.Variable([5.0, 6.0], display_name="short", extend=mo.zero())
            m.x.s = mo.SUM(m.x.rev, m.x.short).set_display_name("s")
            m.x.n = mo.COUNT(m.x.rev, m.x.short).set_display_name("n")
        assert m.x.s.value == (m.x.rev + m.x.short).value == [6.0, 8.0, 3.0, 4.0]
        assert m.x.n.value == [2, 2, 2, 2]
        rows = _book(m, tmp_path, "X")
        assert rows["short"] == [5, 6, 0, 0]
        assert rows["s"] == ["=SUM(B2, B3)", "=SUM(C2, C3)", "=SUM(D2, D3)", "=SUM(E2, E3)"]
        assert rows["n"][3] == "=COUNT(E2, E3)"

    def test_a_line_on_a_window_inside_the_sheet_s_is_read_by_date(self, tmp_path):
        # A sub-container's sheet has the columns of its own window, so
        # a formula on the model's window reads its cell of the same
        # DATE — never the cell in the same position — and, for a
        # period it has no column for, the line's declared rule: 0 for
        # a flow, the first or last cell for a rate held. The operators'
        # cells, which paired by position before, follow the same rule.
        m = _windows_inside()
        assert m.x.s.value == m.x.plus.value == [1.0, 2.0, 13.0, 24.0]
        assert m.x.hp.value == m.x.pr.value == [0.5, 1.0, 1.5, 1.0]
        assert m.x.se.value == [101.0, 202.0, 3.0, 4.0]
        assert m.x.so.value == [1.0, 2.0, 53.0, 4.0]
        assert m.x.sl.value == [6.0, 7.0, 8.0, 9.0]
        path = tmp_path / "book.xlsx"
        m.to_excel(path)
        book = load_workbook(path)
        x, t, u, o = (_rows(book[name]) for name in ("X", "T", "U", "O"))
        assert t["fee"] == [None, 10, 20] and t["k"] == [7, None, None]
        assert u["early"] == [100, 200] and o["once"] == [50]
        assert x["s"] == ["=SUM(B2, 0)", "=SUM(C2, 0)", "=SUM(D2, T!C2)", "=SUM(E2, T!D2)"]
        assert x["plus"] == ["=B2 + 0", "=C2 + 0", "=D2 + T!C2", "=E2 + T!D2"]
        assert x["av"] == ["=AVERAGE(B2, 0, T!C3)", "=AVERAGE(C2, 0, T!C3)",
                           "=AVERAGE(D2, T!C2, T!C3)", "=AVERAGE(E2, T!D2, T!D3)"]
        assert x["pr"] == ["=PRODUCT(B2, T!C3)", "=PRODUCT(C2, T!C3)",
                           "=PRODUCT(D2, T!C3)", "=PRODUCT(E2, T!D3)"]
        assert x["hp"] == ["=B2 * T!C3", "=C2 * T!C3", "=D2 * T!C3", "=E2 * T!D3"]
        assert x["mx"][0] == "=MAX(B2, 0)" and x["mx"][3] == "=MAX(E2, T!D2)"
        assert x["pw"] == ["=POWER(B2, T!C3)", "=POWER(C2, T!C3)",
                           "=POWER(D2, T!C3)", "=POWER(E2, T!D3)"]
        # A scalar on that sheet is its one cell in every period, as before.
        assert x["sk"] == ["=SUM(B2, T!B4)"] + [f"=SUM({c}2, T!B4)" for c in "CDE"]
        # The same start, ending early: its last cell held, or 0.
        assert x["se"] == ["=SUM(B2, U!B2)", "=SUM(C2, U!C2)", "=SUM(D2, 0)", "=SUM(E2, 0)"]
        assert x["mn"] == ["=MIN(B2, U!B3)", "=MIN(C2, U!C3)", "=MIN(D2, U!C3)", "=MIN(E2, U!C3)"]
        assert x["n"] == ["=COUNT(B2, 0, U!B2)", "=COUNT(C2, 0, U!C2)",
                          "=COUNT(D2, T!C2, 0)", "=COUNT(E2, T!D2, 0)"]
        # One period, one cell: a located series still, read by date.
        assert x["so"] == ["=SUM(B2, 0)", "=SUM(C2, 0)", "=SUM(D2, O!B2)", "=SUM(E2, 0)"]
        assert x["sl"] == ["=SUM(B2, O!B3)"] + [f"=SUM({c}2, O!B3)" for c in "CDE"]
        assert x["na"] == ["=COUNTA(B2, 0)", "=COUNTA(C2, 0)", "=COUNTA(D2, O!B2)",
                           "=COUNTA(E2, 0)"]
        assert mo.check_compat(m, "excel") == []

    def test_a_combined_window_beyond_the_sheet_s_is_refused_by_the_export(self, tmp_path):
        # As the export refuses the operator's result on that window.
        m = self._windows()
        with m:
            m.x.s = mo.SUM(m.x.rev, m.t.fee).set_display_name("s")
            m.x.plus = (m.x.rev + m.t.fee).set_display_name("plus")
        with pytest.raises(ValueError, match="window of its own") as refused:
            m.to_excel(tmp_path / "w.xlsx")
        assert "'s'" in str(refused.value) and "'plus'" in str(refused.value)


def _loop(m):
    with m:
        m.d = mo.MultiVariable("D", loop=True)
        with m.d as d:
            prev = mo.lag(d.bal, fill=m.i.q)
            d.inflow = mo.SUM(m.x.a, prev, -10.0).set_display_name("inflow")
            d.mean = mo.AVERAGE(m.x.b, prev).set_display_name("mean")
            d.grown = mo.PRODUCT(prev, 0.5, m.x.b).set_display_name("grown")
            d.n = mo.COUNT(prev, m.x.a, "x").set_display_name("n")
            d.na = mo.COUNTA(prev, m.x.a, "x").set_display_name("na")
            d.bal = (prev + d.inflow - d.mean).set_display_name("bal")
    return m


class TestInsideALoop:
    def test_each_period_reads_the_previous_period(self):
        m = _loop(_model())
        # bal: 9 → inflow 3+9-10=2, mean (4+9)/2=6.5, bal 9+2-6.5=4.5;
        # 4.5 → inflow 8+4.5-10=2.5, mean 4.25, bal 2.75; 2.75 → inflow
        # -2.25, mean 3.375, bal -2.875.
        assert m.d.inflow.value == pytest.approx([2.0, 2.5, -2.25])
        assert m.d.mean.value == pytest.approx([6.5, 4.25, 3.375])
        assert m.d.grown.value == pytest.approx([18.0, 9.0, 5.5])
        assert m.d.n.value == [2, 2, 2]
        assert m.d.na.value == [3, 3, 3]
        assert m.d.bal.value == pytest.approx([4.5, 2.75, -2.875])

    def test_the_same_numbers_as_outside_the_loop(self):
        m = _loop(_model())
        prev = mo.Variable([9.0, 4.5, 2.75])
        assert mo.SUM(m.x.a, prev, -10.0).value == pytest.approx(m.d.inflow.value)
        assert mo.AVERAGE(m.x.b, prev).value == pytest.approx(m.d.mean.value)

    @pytest.mark.parametrize("name", FIVE)
    def test_one_series_alone_stays_refused_and_points_at_the_per_period_form(self, name):
        m = mo.Model("m", default_grain="year", default_start="2026", default_periods=3)
        with pytest.raises(ValueError, match=f"{name} reads a whole series") as err:
            with m:
                m.x = mo.MultiVariable("X", loop=True)
                with m.x as x:
                    x.a = mo.lag(x.b, fill=1.0) + 1.0
                    x.peak = getattr(mo, name)(x.a)
                    x.b = x.a - x.peak
        assert f"mo.{name}(a, b)" in str(err.value)


# ══════════════════════════════════════════════════════════════════════
# The workbook
# ══════════════════════════════════════════════════════════════════════


class TestTheWorkbook:
    def test_one_call_per_period(self, tmp_path):
        m = _model()
        with m:
            m.x.s = mo.SUM(m.x.a, m.x.b).set_display_name("s")
            m.x.av = mo.AVERAGE(m.x.a, m.i.floor).set_display_name("av")
            m.x.pr = mo.PRODUCT(m.x.a, m.x.b, 0.5).set_display_name("pr")
            m.x.n = mo.COUNT(m.x.a, [1.0, None, "x"]).set_display_name("n")
            m.x.na = mo.COUNTA(m.x.a, "x").set_display_name("na")
        x = _book(m, tmp_path, "X")
        assert x["s"] == ["=SUM(B2, B3)", "=SUM(C2, C3)", "=SUM(D2, D3)"], x["s"]
        assert x["av"] == ["=AVERAGE(B2, I!B1)", "=AVERAGE(C2, I!B1)",
                           "=AVERAGE(D2, I!B1)"], x["av"]
        assert x["pr"] == ["=PRODUCT(B2, B3, 0.5)", "=PRODUCT(C2, C3, 0.5)",
                           "=PRODUCT(D2, D3, 0.5)"], x["pr"]
        assert x["n"] == ["=COUNT(B2, 1.0)", "=COUNT(C2, 0)", '=COUNT(D2, "x")'], x["n"]
        assert x["na"] == ['=COUNTA(B2, "x")', '=COUNTA(C2, "x")', '=COUNTA(D2, "x")'], x["na"]

    def test_inside_a_larger_formula(self, tmp_path):
        m = _model()
        with m:
            m.x.twice = (mo.SUM(m.x.a, m.x.b) * 2).set_display_name("twice")
            m.x.inner = mo.SUM(m.x.a * 2, m.x.b).set_display_name("inner")
        x = _book(m, tmp_path, "X")
        assert x["twice"] == ["=SUM(B2, B3) * 2", "=SUM(C2, C3) * 2", "=SUM(D2, D3) * 2"]
        assert x["inner"] == ["=SUM(B2 * 2, B3)", "=SUM(C2 * 2, C3)", "=SUM(D2 * 2, D3)"]

    def test_one_series_and_scalars_only_are_written_as_before(self, tmp_path):
        m = _model()
        with m:
            m.x.red = mo.SUM(m.x.a).set_display_name("red")
            m.x.both = (mo.SUM(m.x.a) + mo.SUM(m.x.b)).set_display_name("both")
            m.x.pair = mo.SUM(m.i.p, m.i.q).set_display_name("pair")
            m.x.cnt = mo.COUNT(m.x.a).set_display_name("cnt")
        # A scalar line gives the sheet a constants column (B); the
        # periods move to C..E.
        x = _book(m, tmp_path, "X")
        assert x["red"][0] == "=SUM(C2:E2)", x["red"]
        assert x["both"][0] == "=SUM(C2:E2) + SUM(C3:E3)", x["both"]
        assert x["pair"][0] == "=SUM(I!B2, I!B3)", x["pair"]
        assert x["cnt"][0] == "=COUNT(C2:E2)", x["cnt"]

    def test_a_loop_reads_the_previous_column(self, tmp_path):
        m = _loop(_model())
        d = _book(m, tmp_path, "D")
        # bal is the sheet's last row, 7: each period reads its previous cell.
        assert d["inflow"] == ["=SUM(X!B2, I!B3, -10.0)", "=SUM(X!C2, B7, -10.0)",
                               "=SUM(X!D2, C7, -10.0)"], d["inflow"]
        assert d["n"][1] == '=COUNT(B7, X!C2, "x")', d["n"]

    def test_check_compat_lists_nothing(self):
        m = _loop(_model())
        with m:
            m.x.s = mo.SUM(m.x.a, m.x.b).set_display_name("s")
            m.x.n = mo.COUNTA(m.x.a, m.x.b).set_display_name("n")
        assert mo.check_compat(m, "excel") == []

    def test_json_and_the_notebook(self):
        m = _model()
        with m:
            m.x.s = mo.SUM(m.x.a, m.x.b).set_display_name("s")
        assert mo.to_json(m.x.s._expr) == {
            "kind": "funccall", "func": "SUM",
            "args": [{"kind": "varref", "path": "m.x.a"}, {"kind": "varref", "path": "m.x.b"}],
        }
        html = m._repr_html_()
        assert "SUM(B2, B3)" in html
        assert "B2:D2" not in html


# ══════════════════════════════════════════════════════════════════════
# Another grain: totals columns and the projection
# ══════════════════════════════════════════════════════════════════════


def _monthly():
    m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=6)
    with m:
        m.p = mo.MultiVariable("P")
        with m.p as p:
            p.a = mo.Variable([1.0, 2.0, 3.0, 4.0, 5.0, 6.0], display_name="a",
                              regrain=mo.up("sum"))
            p.b = mo.Variable([10.0, 10.0, 10.0, 20.0, 20.0, 20.0], display_name="b",
                              regrain=mo.up("sum"))
            p.cash = mo.Variable([5.0, 6.0, 7.0, 8.0, 9.0, 10.0],
                                 display_name="cash").set_regrain(mo.up("last"))
            p.price = mo.Variable([2.0, 2.0, 2.0, 3.0, 3.0, 3.0], display_name="price",
                                  regrain=mo.up("mean"))
            p.s = mo.SUM(p.a, p.b).set_display_name("s")
            p.av = mo.AVERAGE(p.a, p.b).set_display_name("av")
            p.pos = mo.SUM(p.cash, p.a).set_display_name("pos")
            p.rev = mo.PRODUCT(p.price, p.b).set_display_name("rev")
            p.twice = mo.PRODUCT(p.b, 2.0).set_display_name("twice")
    return m


class TestAnotherGrain:
    def test_a_quarter_is_the_call_over_the_operands_quarters(self):
        q = _monthly().at("quarter")
        assert q.p.s.value == [36.0, 75.0]
        assert q.p.av.value == [18.0, 37.5]
        # A balance takes its quarter's last month, a flow its sum, as
        # in cash + a.
        assert q.p.pos.value == [13.0, 25.0]

    def test_each_re_grains_as_its_operator_does(self):
        m = _monthly()
        p = m.p
        with m:
            p.s_op = p.a + p.b
            p.av_op = (p.a + p.b) / 2
            p.pos_op = p.cash + p.a
            p.rev_op = p.price * p.b
            p.twice_op = p.b * 2.0
        q = m.at("quarter").p
        for name in ("s", "av", "pos", "rev", "twice"):
            assert getattr(q, name).value == getattr(q, f"{name}_op").value, name
        # A product of two flows is a flow: the quarter sums its months.
        assert q.rev.value == [60.0, 180.0]

    def test_a_count_per_period_asks_for_its_rule(self):
        from modeleon.core.projection import project_model

        m = _monthly()
        with m:
            m.p.n = mo.COUNT(m.p.a, m.p.b).set_display_name("n")
        with pytest.raises(ValueError, match="explicit rule"):
            m.at("quarter")
        absorbed = project_model(m, "quarter", on_error="absorb")
        assert absorbed.p.n.value == "#VALUE!"
        assert absorbed.p.s.value == [36.0, 75.0]

    def test_a_totals_column_folds_it_live(self, tmp_path):
        from modeleon.compile.excel.view import ExcelView

        m = _monthly()
        m.default_excel_view = ExcelView(timeline={"totals": ["quarter"]})
        path = tmp_path / "t.xlsx"
        m.to_excel(path)
        rows = _rows(load_workbook(path)["P"])
        s = [v for v in rows["s"] if isinstance(v, str) and v.startswith("=SUM(")]
        assert len(s) == 8, rows["s"]
        assert all("," in f for f in s), s
        assert "#VALUE!" not in rows["s"] and "#VALUE!" not in rows["av"]


# ══════════════════════════════════════════════════════════════════════
# The registry
# ══════════════════════════════════════════════════════════════════════


class TestTheRegistry:
    @pytest.mark.parametrize("name", FIVE)
    def test_one_call_per_period_from_two_values(self, name):
        assert per_period_arity(name) == 2
        assert not admits_in_loop(name, 1)
        assert admits_in_loop(name, 2) and admits_in_loop(name, 3)


# ══════════════════════════════════════════════════════════════════════
# The spreadsheet computes the same numbers
# ══════════════════════════════════════════════════════════════════════


def _soffice():
    for candidate in ("soffice", "libreoffice"):
        found = shutil.which(candidate)
        if found:
            return found
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


@pytest.mark.slow
def test_the_workbook_recalculates_to_the_python_values(tmp_path):
    soffice = _soffice()
    if soffice is None:
        pytest.skip("LibreOffice not installed")
    m = _loop(_model())
    with m:
        m.x.s = mo.SUM(m.x.a, m.x.b).set_display_name("s")
        m.x.av = mo.AVERAGE(m.x.a, m.i.floor).set_display_name("av")
        m.x.pr = mo.PRODUCT(m.x.a, m.x.b, 0.5).set_display_name("pr")
        m.x.n = mo.COUNT(m.x.a, [1.0, None, "x"]).set_display_name("n")
        m.x.na = mo.COUNTA(m.x.a, "x").set_display_name("na")
        m.x.lit = mo.SUM(m.x.a, [1.0, 9.0, 1.0]).set_display_name("lit")
        m.x.twice = (mo.SUM(m.x.a, m.x.b) * 2).set_display_name("twice")
        m.x.inner = mo.AVERAGE(m.x.a * 2, m.x.b).set_display_name("inner")
        m.x.short = mo.Variable([5.0, 6.0], display_name="short", extend=mo.zero())
        m.x.s2 = mo.SUM(m.x.a, m.x.short).set_display_name("s2")
        m.x.n2 = mo.COUNT(m.x.a, m.x.short).set_display_name("n2")
        m.x.mn = mo.MIN(m.x.a, None).set_display_name("mn")
        m.x.hole = mo.AVERAGE(m.x.a, [1.0, None, 3.0]).set_display_name("hole")
        # On the scalars' sheet, so X keeps its period columns at B..D.
        m.i.red = mo.SUM(m.x.a).set_display_name("red")
        m.i.both = (mo.SUM(m.x.a) + mo.SUM(m.x.b)).set_display_name("both")
    path = tmp_path / "aggregates.xlsx"
    m.to_excel(path)
    subprocess.run(
        [soffice, f"-env:UserInstallation={(tmp_path / 'prof').as_uri()}",
         "--headless", "--convert-to", "xlsx", "--outdir",
         str(tmp_path / "calc"), str(path)],
        check=True, capture_output=True,
    )
    book = load_workbook(tmp_path / "calc" / "aggregates.xlsx", data_only=True)
    x = _rows(book["X"])
    d = _rows(book["D"])
    i = _rows(book["I"])
    for name in ("s", "av", "pr", "n", "na", "lit", "twice", "inner", "s2", "n2", "mn", "hole"):
        got = [round(v, 9) for v in x[name][:3]]
        assert got == pytest.approx(getattr(m.x, name).value), name
    assert i["red"][0] == m.i.red.value == 16.0
    assert i["both"][0] == m.i.both.value == 28.0
    for name in ("inflow", "mean", "grown", "n", "na", "bal"):
        got = [round(v, 9) for v in d[name][:3]]
        assert got == pytest.approx(getattr(m.d, name).value), name


@pytest.mark.slow
def test_the_workbook_recalculates_lines_on_windows_inside_the_sheet_s(tmp_path):
    soffice = _soffice()
    if soffice is None:
        pytest.skip("LibreOffice not installed")
    m = _windows_inside()
    path = tmp_path / "inside.xlsx"
    m.to_excel(path)
    subprocess.run(
        [soffice, f"-env:UserInstallation={(tmp_path / 'prof').as_uri()}",
         "--headless", "--convert-to", "xlsx", "--outdir",
         str(tmp_path / "calc"), str(path)],
        check=True, capture_output=True,
    )
    x = _rows(load_workbook(tmp_path / "calc" / "inside.xlsx", data_only=True)["X"])
    for name in _INSIDE:
        got = [round(v, 9) for v in x[name][:4]]
        assert got == pytest.approx(getattr(m.x, name).value), name


@pytest.mark.slow
def test_the_totals_recalculate_to_the_projection(tmp_path):
    soffice = _soffice()
    if soffice is None:
        pytest.skip("LibreOffice not installed")
    from modeleon.compile.excel.view import ExcelView

    m = _monthly()
    m.default_excel_view = ExcelView(timeline={"totals": ["quarter"]})
    path = tmp_path / "totals.xlsx"
    m.to_excel(path)
    subprocess.run(
        [soffice, f"-env:UserInstallation={(tmp_path / 'prof').as_uri()}",
         "--headless", "--convert-to", "xlsx", "--outdir",
         str(tmp_path / "calc"), str(path)],
        check=True, capture_output=True,
    )
    raw = _rows(load_workbook(path)["P"])
    calc = _rows(load_workbook(tmp_path / "calc" / "totals.xlsx", data_only=True)["P"])
    q = m.at("quarter")
    for name in ("s", "av", "pos", "rev", "twice"):
        native = getattr(m.p, name).value
        quarters = getattr(q.p, name).value
        cells = [v for v in calc[name] if v is not None]
        formulas = [f for f in raw[name] if f is not None]
        assert len(cells) == len(formulas) == 8, (name, raw[name])
        # Months then their quarter: Jan Feb Mar Q1 Apr May Jun Q2.
        assert cells == pytest.approx(native[:3] + quarters[:1] + native[3:] + quarters[1:]), name
