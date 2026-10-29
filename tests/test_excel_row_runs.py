# SPDX-License-Identifier: Apache-2.0
"""A row read as a range in a book with total columns between its periods.

``ExcelView(timeline={'totals': ['quarter', 'year']})`` puts a total
column after each quarter's last month and after the year, so a line's
period cells are no longer one run: ``Data!C4:N4`` would read the quarter
totals in. Wherever the spreadsheet takes several references the row is
written as its runs: ``=SUM(Data!C4:E4, Data!G4:I4, …)``; ``COUNTIF`` /
``SUMIF`` and their ``S`` forms are the sum of one call per run, and
``AVERAGEIF(S)`` the ratio of those sums. The readers of one ordered
range (``MATCH``, ``INDEX``, ``LARGE``, …) keep the computed value, and
``check_compat`` says why and names the way out. A book without total
columns is written as before, byte for byte.
"""

import datetime as dt
import os
import shutil
import subprocess
import tempfile
import warnings
from pathlib import Path

import pytest
from openpyxl import load_workbook

import modeleon as mo
from modeleon.compile import check_compat
from modeleon.compile.compat import CompatIssue
from modeleon.compile.excel.renderer import _contiguous_run, _run_spans, _united_spans

N = 24


def _customers():
    # Rises to one peak in month 17 (index 16), then falls.
    return [100.0 + 10 * i if i <= 16 else 260.0 - 15 * (i - 16) for i in range(N)]


def _revenue():
    return [1000.0 + 37 * ((i * 7) % 11) + i for i in range(N)]


def _regions():
    return [("North", "South", "East")[i % 3] for i in range(N)]


def _growth():
    return [1.0 + 0.001 * (i % 5) for i in range(N)]


def _flow():
    return [-1000.0] + [60.0 + i for i in range(1, N)]


#: One cell on Summary each: (name, builder).
_SCALARS = [
    ("sum", lambda m: mo.SUM(m.data.revenue)),
    ("min", lambda m: mo.MIN(m.data.revenue)),
    ("max", lambda m: mo.MAX(m.data.revenue)),
    ("average", lambda m: mo.AVERAGE(m.data.revenue)),
    ("count", lambda m: mo.COUNT(m.data.revenue)),
    ("counta", lambda m: mo.COUNTA(m.data.region)),
    ("product", lambda m: mo.PRODUCT(m.data.growth)),
    ("median", lambda m: mo.MEDIAN(m.data.revenue)),
    ("stdev_s", lambda m: mo.STDEV_S(m.data.revenue)),
    ("stdev_p", lambda m: mo.STDEV_P(m.data.revenue)),
    ("var_s", lambda m: mo.VAR_S(m.data.revenue)),
    ("var_p", lambda m: mo.VAR_P(m.data.revenue)),
    ("npv", lambda m: mo.NPV(0.01, m.data.flow)),
    ("countif", lambda m: mo.COUNTIF(m.data.region, "North")),
    ("countifs", lambda m: mo.COUNTIFS(m.data.region, "North", m.data.revenue, ">1100")),
    ("sumif", lambda m: mo.SUMIF(m.data.region, "North", m.data.revenue)),
    ("sumif_self", lambda m: mo.SUMIF(m.data.revenue, ">1100")),
    ("sumifs", lambda m: mo.SUMIFS(m.data.revenue, m.data.region, "North")),
    ("averageif", lambda m: mo.AVERAGEIF(m.data.region, "North", m.data.revenue)),
    ("averageifs", lambda m: mo.AVERAGEIFS(m.data.revenue, m.data.region, "North")),
    ("maxifs", lambda m: mo.MAXIFS(m.data.revenue, m.data.region, "North")),
    ("minifs", lambda m: mo.MINIFS(m.data.revenue, m.data.region, "North")),
    ("match_max", lambda m: mo.MATCH(mo.MAX(m.data.customers), m.data.customers, 0)),
    ("match_key", lambda m: mo.MATCH(m.inputs.key, m.data.customers, 0)),
    ("index", lambda m: mo.INDEX(m.data.revenue, 3)),
    ("index_last", lambda m: mo.INDEX(m.data.revenue, mo.COUNT(m.data.revenue))),
    ("index_match", lambda m: mo.INDEX(m.data.revenue,
                                       mo.MATCH(mo.MAX(m.data.customers), m.data.customers, 0))),
    ("xlookup", lambda m: mo.XLOOKUP(m.inputs.key, m.data.customers, m.data.revenue)),
    ("hlookup", lambda m: mo.HLOOKUP(m.inputs.key, [m.data.customers, m.data.revenue], 2, False)),
    ("large", lambda m: mo.LARGE(m.data.revenue, 2)),
    ("small", lambda m: mo.SMALL(m.data.revenue, 2)),
    ("percentile_inc", lambda m: mo.PERCENTILE_INC(m.data.revenue, 0.3)),
    ("quartile_inc", lambda m: mo.QUARTILE_INC(m.data.revenue, 1)),
    ("rank_eq", lambda m: mo.RANK_EQ(m.inputs.score, m.data.revenue)),
    ("irr", lambda m: mo.IRR(m.data.flow)),
    ("last", lambda m: m.data.revenue[-1] + 0),
    ("sum_slice_cross", lambda m: mo.SUM(m.data.revenue[0:14])),
    ("sum_slice_in", lambda m: mo.SUM(m.data.revenue[12:24])),
    ("match_slice_in", lambda m: mo.MATCH(mo.MAX(m.data.customers[12:24]),
                                          m.data.customers[12:24], 0)),
    ("share", lambda m: mo.SUMIF(m.data.region, "North", m.data.revenue) / mo.SUM(m.data.revenue)),
    ("half_sum", lambda m: mo.SUM(m.data.revenue) / 2),
    ("stdev", lambda m: mo.STDEV(m.data.revenue)),
    ("stdevp", lambda m: mo.STDEVP(m.data.revenue)),
    ("var", lambda m: mo.VAR(m.data.revenue)),
    ("varp", lambda m: mo.VARP(m.data.revenue)),
    ("sumproduct", lambda m: mo.SUM(m.data.revenue * m.data.growth)),
    ("textjoin", lambda m: mo.TEXTJOIN(", ", True, m.data.region)),
]

#: One call per period on Summary: a series among the call's values.
_PER_PERIOD = [
    ("rank_pp", lambda m: mo.RANK_EQ(m.data.revenue, m.data.revenue)),
    ("match_pp", lambda m: mo.MATCH(m.data.customers, m.data.customers, 0)),
    ("large_pp", lambda m: mo.LARGE(m.data.revenue, m.data.k)),
    ("xlookup_pp", lambda m: mo.XLOOKUP(m.data.customers, m.data.customers, m.data.revenue)),
    ("sum_two_pp", lambda m: mo.SUM(m.data.revenue, m.data.customers)),
    ("max_floor_pp", lambda m: mo.MAX(m.data.flow, 0)),
]

#: On the rows' own sheet, Data.
_SAME_SHEET = [
    ("d_sum", lambda m: mo.SUM(m.data.revenue)),
    ("d_match_max", lambda m: mo.MATCH(mo.MAX(m.data.customers), m.data.customers, 0)),
    ("d_large", lambda m: mo.LARGE(m.data.revenue, 1)),
    ("d_countif", lambda m: mo.COUNTIF(m.data.region, "North")),
]


def _matrix(totals=None):
    """Every function over 24 monthly rows on Data, its result on Summary
    (and a few on Data itself), under ``totals`` total columns."""
    view = {"default_excel_view": mo.ExcelView(timeline={"totals": totals})} if totals else {}
    m = mo.Model("m", display_name="M", default_grain="month", default_start="2026-01",
                 default_periods=N, **view)
    with m:
        m.inputs = mo.MultiVariable("Inputs", excel_props={"tab": True})
        m.inputs.key = mo.Variable(_customers()[5], display_name="Key")
        m.inputs.score = mo.Variable(_revenue()[5], display_name="Score")
        m.data = mo.MultiVariable("Data", excel_props={"tab": True})
        with m.data as d:
            d.customers = mo.Variable(_customers(), display_name="Customers",
                                      regrain=mo.up("last"))
            d.revenue = mo.Variable(_revenue(), display_name="Revenue", regrain=mo.up("sum"))
            d.region = mo.Variable(_regions(), display_name="Region")
            d.growth = mo.Variable(_growth(), display_name="Growth", regrain=mo.up("last"))
            d.flow = mo.Variable(_flow(), display_name="Flow", regrain=mo.up("sum"))
            d.k = mo.Variable([float(1 + i % 5) for i in range(N)], display_name="K",
                              regrain=mo.up("last"))
        m.s = mo.MultiVariable("Summary", excel_props={"tab": True})
        for name, build in _SCALARS + _PER_PERIOD:
            setattr(m.s, name, build(m).set_display_name(name))
        for name, build in _SAME_SHEET:
            setattr(m.data, name, build(m).set_display_name(name))
    return m


def _written(m):
    """``{(sheet, label): [cell, …]}`` — what the written book holds on
    each line's row, blank cells left out."""
    path = os.path.join(tempfile.mkdtemp(), "book.xlsx")
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        m.to_excel(path)
    book = load_workbook(path)
    out = {}
    for ws in book.worksheets:
        for row in ws.iter_rows():
            if row[0].value:
                out[(ws.title, str(row[0].value))] = [c.value for c in row[1:]
                                                      if c.value is not None]
    return out


def _columns(first, n):
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.cell import column_index_from_string

    start = column_index_from_string(first)
    return [get_column_letter(start + i) for i in range(n)]


# ══════════════════════════════════════════════════════════════════════
# The control: without total columns every row is one range, as before
# ══════════════════════════════════════════════════════════════════════


#: The book without total columns, cell for cell.
_CONTROL = {
    "sum": "=SUM(Data!C3:Z3)",
    "min": "=MIN(Data!C3:Z3)",
    "max": "=MAX(Data!C3:Z3)",
    "average": "=AVERAGE(Data!C3:Z3)",
    "count": "=COUNT(Data!C3:Z3)",
    "counta": "=COUNTA(Data!C4:Z4)",
    "product": "=PRODUCT(Data!C5:Z5)",
    "median": "=MEDIAN(Data!C3:Z3)",
    "stdev_s": "=_xlfn.STDEV.S(Data!C3:Z3)",
    "stdev_p": "=_xlfn.STDEV.P(Data!C3:Z3)",
    "var_s": "=_xlfn.VAR.S(Data!C3:Z3)",
    "var_p": "=_xlfn.VAR.P(Data!C3:Z3)",
    "npv": "=NPV(0.01, Data!C6:Z6)",
    "countif": '=COUNTIF(Data!C4:Z4, "North")',
    "countifs": '=COUNTIFS(Data!C4:Z4, "North", Data!C3:Z3, ">1100")',
    "sumif": '=SUMIF(Data!C4:Z4, "North", Data!C3:Z3)',
    "sumif_self": '=SUMIF(Data!C3:Z3, ">1100")',
    "sumifs": '=SUMIFS(Data!C3:Z3, Data!C4:Z4, "North")',
    "averageif": '=AVERAGEIF(Data!C4:Z4, "North", Data!C3:Z3)',
    "averageifs": '=AVERAGEIFS(Data!C3:Z3, Data!C4:Z4, "North")',
    "maxifs": '=_xlfn.MAXIFS(Data!C3:Z3, Data!C4:Z4, "North")',
    "minifs": '=_xlfn.MINIFS(Data!C3:Z3, Data!C4:Z4, "North")',
    "match_max": "=MATCH(MAX(Data!C2:Z2), Data!C2:Z2, 0)",
    "match_key": "=MATCH(Inputs!B1, Data!C2:Z2, 0)",
    "index": "=INDEX(Data!C3:Z3, 3)",
    "index_last": "=INDEX(Data!C3:Z3, COUNT(Data!C3:Z3))",
    "index_match": "=INDEX(Data!C3:Z3, MATCH(MAX(Data!C2:Z2), Data!C2:Z2, 0))",
    "xlookup": "=_xlfn.XLOOKUP(Inputs!B1, Data!C2:Z2, Data!C3:Z3)",
    "hlookup": "=HLOOKUP(Inputs!B1, Data!C2:Z3, 2, FALSE)",
    "large": "=LARGE(Data!C3:Z3, 2)",
    "small": "=SMALL(Data!C3:Z3, 2)",
    "percentile_inc": "=_xlfn.PERCENTILE.INC(Data!C3:Z3, 0.3)",
    "quartile_inc": "=_xlfn.QUARTILE.INC(Data!C3:Z3, 1)",
    "rank_eq": "=_xlfn.RANK.EQ(Inputs!B2, Data!C3:Z3)",
    "irr": "=IRR(Data!C6:Z6, 0.1)",
    "last": "=Data!Z3 + 0",
    "sum_slice_cross": "=SUM(Data!C3:P3)",
    "sum_slice_in": "=SUM(Data!O3:Z3)",
    "match_slice_in": "=MATCH(MAX(Data!O2:Z2), Data!O2:Z2, 0)",
    "share": '=SUMIF(Data!C4:Z4, "North", Data!C3:Z3) / SUM(Data!C3:Z3)',
    "half_sum": "=SUM(Data!C3:Z3) / 2",
    "stdev": "=STDEV(Data!C3:Z3)",
    "stdevp": "=STDEVP(Data!C3:Z3)",
    "var": "=VAR(Data!C3:Z3)",
    "varp": "=VARP(Data!C3:Z3)",
    "sumproduct": "=SUMPRODUCT(Data!C3:Z3, Data!C5:Z5)",
    "textjoin": '=_xlfn.TEXTJOIN(", ", TRUE, Data!C4:Z4)',
}

#: The calls one per period, by the column of their period.
_CONTROL_PER_PERIOD = {
    "rank_pp": "=_xlfn.RANK.EQ(Data!{c}3, Data!C3:Z3)",
    "match_pp": "=MATCH(Data!{c}2, Data!C2:Z2, 0)",
    "large_pp": "=LARGE(Data!C3:Z3, Data!{c}7)",
    "xlookup_pp": "=_xlfn.XLOOKUP(Data!{c}2, Data!C2:Z2, Data!C3:Z3)",
    "sum_two_pp": "=SUM(Data!{c}3, Data!{c}2)",
    "max_floor_pp": "=MAX(Data!{c}6, 0)",
}

_CONTROL_SAME_SHEET = {
    "d_sum": "=SUM(C3:Z3)",
    "d_match_max": "=MATCH(MAX(C2:Z2), C2:Z2, 0)",
    "d_large": "=LARGE(C3:Z3, 1)",
    "d_countif": '=COUNTIF(C4:Z4, "North")',
}


class TestTheBookWithoutTotalsIsUnchanged:
    def test_every_function_reads_one_range(self):
        m = _matrix()
        written = _written(m)
        for name, text in _CONTROL.items():
            assert written[("Summary", name)] == [text], name
        for name, text in _CONTROL_PER_PERIOD.items():
            assert written[("Summary", name)] == [text.format(c=c) for c in _columns("C", N)], name
        for name, text in _CONTROL_SAME_SHEET.items():
            assert written[("Data", name)] == [text], name
        assert check_compat(m, "excel") == []


# ══════════════════════════════════════════════════════════════════════
# The runs of a row
# ══════════════════════════════════════════════════════════════════════


class TestTheRunsOfARow:
    @pytest.mark.parametrize("refs, spans", [
        ([], []),
        (["B3"], [(0, 1)]),
        (["B3", "C3", "D3"], [(0, 3)]),
        (["C3", "D3", "E3", "G3", "H3"], [(0, 3), (3, 5)]),
        (["B3", "B4", "B5"], [(0, 3)]),
        (["B3", "B4", "C4"], [(0, 2), (2, 3)]),
        (["D3", "C3"], [(0, 1), (1, 2)]),
        (["C3", "E3", "G3"], [(0, 1), (1, 2), (2, 3)]),
        (["$B$3", "$C$3"], [(0, 2)]),
        (["Data!B3", "Data!C3", "Data!E3"], [(0, 2), (2, 3)]),
        (["?", "C3"], [(0, 2)]),
    ])
    def test_a_run_is_side_by_side_in_one_direction(self, refs, spans):
        assert _run_spans(refs) == spans

    @pytest.mark.parametrize("refs", [
        ["B3", "C3", "D3"], ["B3"], ["B3", "B5"], ["D3", "C3"], ["B3", "B4", "C4"], ["?", "C3"],
    ])
    def test_one_run_is_one_range(self, refs):
        assert (len(_run_spans(refs)) == 1) is _contiguous_run(refs)

    def test_rows_read_together_break_where_any_of_them_breaks(self):
        # By position: one row breaks after its third cell, the other (a
        # sheet with a start of its own) after its second and its fifth.
        a = ["B3", "C3", "D3", "F3", "G3", "H3"]
        b = ["B4", "C4", "E4", "F4", "G4", "I4"]
        assert _united_spans([a]) == [(0, 3), (3, 6)]
        assert _united_spans([a, b]) == [(0, 2), (2, 3), (3, 5), (5, 6)]


# ══════════════════════════════════════════════════════════════════════
# One argument per run
# ══════════════════════════════════════════════════════════════════════


#: A Data row's 24 months under quarter and year totals: eight runs.
_RUNS = ("C{r}:E{r}", "G{r}:I{r}", "K{r}:M{r}", "O{r}:Q{r}",
         "T{r}:V{r}", "X{r}:Z{r}", "AB{r}:AD{r}", "AF{r}:AH{r}")

#: The rows of Data under quarter and year totals (three header rows,
#: then customers on row 4).
_REVENUE, _REGION, _GROWTH, _FLOW = 5, 6, 7, 8


def _runs(row, sheet="Data!"):
    return [sheet + run.format(r=row) for run in _RUNS]


@pytest.fixture(scope="module")
def interleaved():
    m = _matrix(["quarter", "year"])
    return m, _written(m), check_compat(m, "excel")


class TestARowIsWrittenAsItsRuns:
    @pytest.mark.parametrize("name, excel, row", [
        ("sum", "SUM", _REVENUE), ("min", "MIN", _REVENUE), ("max", "MAX", _REVENUE),
        ("average", "AVERAGE", _REVENUE), ("count", "COUNT", _REVENUE),
        ("counta", "COUNTA", _REGION), ("product", "PRODUCT", _GROWTH),
        ("median", "MEDIAN", _REVENUE), ("stdev_s", "_xlfn.STDEV.S", _REVENUE),
        ("stdev_p", "_xlfn.STDEV.P", _REVENUE), ("var_s", "_xlfn.VAR.S", _REVENUE),
        ("var_p", "_xlfn.VAR.P", _REVENUE), ("stdev", "STDEV", _REVENUE),
        ("stdevp", "STDEVP", _REVENUE), ("var", "VAR", _REVENUE), ("varp", "VARP", _REVENUE),
    ])
    def test_one_argument_per_run(self, interleaved, name, excel, row):
        _m, written, _issues = interleaved
        assert written[("Summary", name)] == [f"={excel}({', '.join(_runs(row))})"]

    def test_npv_keeps_its_rate_first(self, interleaved):
        _m, written, _issues = interleaved
        assert written[("Summary", "npv")] == [f"=NPV(0.01, {', '.join(_runs(_FLOW))})"]

    def test_on_the_rows_own_sheet_no_run_names_a_sheet(self, interleaved):
        _m, written, _issues = interleaved
        assert written[("Data", "d_sum")] == [f"=SUM({', '.join(_runs(_REVENUE, ''))})"]

    def test_a_window_is_its_own_runs(self, interleaved):
        _m, written, _issues = interleaved
        # January 2026 to February 2027: four quarters and two months.
        assert written[("Summary", "sum_slice_cross")] == [
            "=SUM(Data!C5:E5, Data!G5:I5, Data!K5:M5, Data!O5:Q5, Data!T5:U5)"]
        assert written[("Summary", "sum_slice_in")] == [f"=SUM({', '.join(_runs(_REVENUE)[4:])})"]

    def test_inside_a_larger_formula(self, interleaved):
        _m, written, _issues = interleaved
        assert written[("Summary", "half_sum")] == [f"=SUM({', '.join(_runs(_REVENUE))}) / 2"]

    def test_a_sum_of_a_product_is_one_sumproduct_per_run(self, interleaved):
        _m, written, _issues = interleaved
        pairs = [f"SUMPRODUCT({a}, {b})" for a, b in zip(_runs(_REVENUE), _runs(_GROWTH))]
        assert written[("Summary", "sumproduct")] == [f"=SUM({', '.join(pairs)})"]

    def test_a_sum_of_a_product_has_its_cells(self, tmp_path):
        # Its rows are laid out: no warning that they have no cells.
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=6,
                     default_excel_view=mo.ExcelView(timeline={"totals": ["quarter"]}))
        with m:
            m.d = mo.MultiVariable("Data", excel_props={"tab": True})
            m.d.p = mo.Variable([1.0, 2, 3, 4, 5, 6], display_name="P", regrain=mo.up("mean"))
            m.d.v = mo.Variable([10.0] * 6, display_name="V", regrain=mo.up("sum"))
            m.s = mo.MultiVariable("S", excel_props={"tab": True})
            m.s.rev = mo.SUM(m.d.p * m.d.v).set_display_name("rev")
        with warnings.catch_warnings():
            warnings.simplefilter("error", mo.CrossScopeReferenceWarning)
            m.to_excel(tmp_path / "book.xlsx")
            assert check_compat(m, "excel") == []
        assert load_workbook(tmp_path / "book.xlsx")["S"]["B1"].value == (
            "=SUM(SUMPRODUCT(Data!B3:D3, Data!B4:D4), SUMPRODUCT(Data!F3:H3, Data!F4:H4))")

    def test_a_running_total_inside_a_formula_sums_its_runs(self):
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=6,
                     default_excel_view=mo.ExcelView(timeline={"totals": ["quarter"]}))
        with m:
            m.a = mo.MultiVariable("A", excel_props={"tab": True})
            with m.a as a:
                a.inc = mo.Variable([100.0] * 6, display_name="In", regrain=mo.up("sum"))
                a.out = mo.Variable([60.0] * 6, display_name="Out", regrain=mo.up("sum"))
                a.cash = (mo.cumsum(a.inc) - mo.cumsum(a.out)).set_display_name("Cash")
        cash = _written(m)[("A", "Cash")]
        # January to March is one run; May steps over the Q1 total at E.
        assert cash[2] == "=SUM(B3:D3) - SUM(B4:D4)"
        assert cash[5] == "=SUM(B3:D3, F3:G3) - SUM(B4:D4, F4:G4)"

    def test_a_quoted_sheet_names_itself_in_front_of_every_run(self):
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=N,
                     default_excel_view=mo.ExcelView(timeline={"totals": ["quarter", "year"]}))
        with m:
            m.d = mo.MultiVariable("Bob's data", excel_props={"tab": True})
            m.d.x = mo.Variable(_revenue(), display_name="X", regrain=mo.up("sum"))
            m.d.r = mo.Variable(_regions(), display_name="R")
            m.s = mo.MultiVariable("S", excel_props={"tab": True})
            m.s.total = (mo.SUM(m.d.x) / 2).set_display_name("total")
            m.s.cnt = mo.COUNTIF(m.d.r, "North").set_display_name("cnt")
            m.s.mean = mo.AVERAGEIF(m.d.r, "North", m.d.x).set_display_name("mean")
        # No column of constants on Bob's data: its months start at B.
        quoted = [f"'Bob''s data'!{run}" for run in (
            "B{r}:D{r}", "F{r}:H{r}", "J{r}:L{r}", "N{r}:P{r}",
            "S{r}:U{r}", "W{r}:Y{r}", "AA{r}:AC{r}", "AE{r}:AG{r}")]
        x, r = [run.format(r=4) for run in quoted], [run.format(r=5) for run in quoted]
        written = _written(m)
        assert written[("S", "total")] == [f"=SUM({', '.join(x)}) / 2"]
        assert written[("S", "cnt")] == [
            "=SUM({})".format(", ".join(f'COUNTIF({c}, "North")' for c in r))]
        numbers = '">=-9.99999999999999E+307"'
        assert written[("S", "mean")] == ["=SUM({}) / SUM({})".format(
            ", ".join(f'SUMIF({c}, "North", {v})' for c, v in zip(r, x)),
            ", ".join(f'COUNTIFS({v}, {numbers}, {c}, "North")' for c, v in zip(r, x)))]

    def test_check_compat_lists_none_of_them(self, interleaved):
        m, _written, issues = interleaved
        listed = {i.qpath for i in issues}
        live = [m.s.sum, m.s.median, m.s.npv, m.s.countif, m.s.sumifs, m.s.averageif,
                m.s.averageifs, m.s.sumproduct, m.s.share, m.s.half_sum, m.s.stdev,
                m.data.d_sum, m.data.d_countif]
        assert not listed & {line.id for line in live}


# ══════════════════════════════════════════════════════════════════════
# A conditional count or sum is one call per run
# ══════════════════════════════════════════════════════════════════════


class TestTheCriteriaAreOneCallPerRun:
    def test_a_conditional_count_or_sum_adds_one_call_per_run(self, interleaved):
        _m, written, _issues = interleaved
        region, revenue = _runs(_REGION), _runs(_REVENUE)
        calls = {
            "countif": [f'COUNTIF({c}, "North")' for c in region],
            "countifs": [f'COUNTIFS({c}, "North", {r}, ">1100")' for c, r in zip(region, revenue)],
            "sumif": [f'SUMIF({c}, "North", {r})' for c, r in zip(region, revenue)],
            "sumif_self": [f'SUMIF({r}, ">1100")' for r in revenue],
            "sumifs": [f'SUMIFS({r}, {c}, "North")' for c, r in zip(region, revenue)],
        }
        for name, each in calls.items():
            assert written[("Summary", name)] == [f"=SUM({', '.join(each)})"], name
        own = [f'COUNTIF({c}, "North")' for c in _runs(_REGION, "")]
        assert written[("Data", "d_countif")] == [f"=SUM({', '.join(own)})"]

    def test_a_conditional_mean_is_the_sum_over_the_count_of_its_numbers(self, interleaved):
        _m, written, _issues = interleaved
        numbers = '">=-9.99999999999999E+307"'
        pairs = list(zip(_runs(_REGION), _runs(_REVENUE)))
        counts = ", ".join(f'COUNTIFS({r}, {numbers}, {c}, "North")' for c, r in pairs)
        sumif = ", ".join(f'SUMIF({c}, "North", {r})' for c, r in pairs)
        sumifs = ", ".join(f'SUMIFS({r}, {c}, "North")' for c, r in pairs)
        assert written[("Summary", "averageif")] == [f"=SUM({sumif}) / SUM({counts})"]
        assert written[("Summary", "averageifs")] == [f"=SUM({sumifs}) / SUM({counts})"]

    def test_inside_a_larger_formula_each_stays_one_operand(self, interleaved):
        _m, written, _issues = interleaved
        sumif = ", ".join(f'SUMIF({c}, "North", {r})'
                          for c, r in zip(_runs(_REGION), _runs(_REVENUE)))
        assert written[("Summary", "share")] == [
            f"=SUM({sumif}) / SUM({', '.join(_runs(_REVENUE))})"]

    def test_a_mean_inside_a_larger_formula_keeps_its_brackets(self):
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=6,
                     default_excel_view=mo.ExcelView(timeline={"totals": ["quarter"]}))
        with m:
            m.d = mo.MultiVariable("Data", excel_props={"tab": True})
            m.d.r = mo.Variable(["a", "b", "a", "a", "b", "a"], display_name="R")
            m.d.x = mo.Variable([1.0, 2, 3, 4, 5, 6], display_name="X", regrain=mo.up("sum"))
            m.s = mo.MultiVariable("S", excel_props={"tab": True})
            m.s.ratio = (m.d.x[0] / mo.AVERAGEIF(m.d.r, "a", m.d.x)).set_display_name("ratio")
        (formula,) = _written(m)[("S", "ratio")]
        assert formula.startswith("=Data!B4 / (SUM(SUMIF(Data!B3:D3, \"a\", Data!B4:D4), ")
        assert formula.endswith('COUNTIFS(Data!F4:H4, ">=-9.99999999999999E+307", '
                                'Data!F3:H3, "a")))')

    def test_the_criteria_is_written_once(self):
        # A criteria no cell holds — a line outside the model — is written
        # as its value: once, and every call reads that text.
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=6,
                     default_excel_view=mo.ExcelView(timeline={"totals": ["quarter"]}))
        outside = mo.Variable("a")
        with m:
            m.d = mo.MultiVariable("Data", excel_props={"tab": True})
            m.d.r = mo.Variable(["a", "b", "a", "a", "b", "a"], display_name="R")
            m.s = mo.MultiVariable("S", excel_props={"tab": True})
            m.s.cnt = mo.COUNTIF(m.d.r, outside).set_display_name("cnt")
        with pytest.warns(mo.CrossScopeReferenceWarning):
            issues = check_compat(m, "excel")
        assert [(i.kind, i.qpath, i.literal) for i in issues] == [("inlined", m.s.cnt.id, '"a"')]

    def test_rows_on_sheets_with_their_own_start_break_together(self):
        # The regions run from January, the amounts from February: the
        # sheets' quarters close after other periods. Each call reads the
        # same periods of both rows — cut where either breaks.
        n = 12
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=n,
                     default_excel_view=mo.ExcelView(timeline={"totals": ["quarter"]}))
        with m:
            m.a = mo.MultiVariable("A", excel_props={"tab": True})
            m.a.region = mo.Variable(["N" if i % 3 else "S" for i in range(n)],
                                     display_name="Region")
            m.b = mo.MultiVariable("B", excel_props={"tab": True}, default_grain="month",
                                   default_start="2026-02", default_periods=n)
            m.b.amt = mo.Variable([float(i + 1) for i in range(n)], display_name="Amt",
                                  regrain=mo.up("sum"))
            m.s = mo.MultiVariable("S", excel_props={"tab": True})
            m.s.cross = mo.SUMIF(m.a.region, "N", m.b.amt).set_display_name("cross")
        assert m.s.cross.value == 56.0
        assert _written(m)[("S", "cross")] == [
            '=SUM(SUMIF(A!B3:C3, "N", B!B3:C3), SUMIF(A!D3, "N", B!E3), '
            'SUMIF(A!F3:G3, "N", B!F3:G3), SUMIF(A!H3, "N", B!I3), '
            'SUMIF(A!J3:K3, "N", B!J3:K3), SUMIF(A!L3, "N", B!M3), '
            'SUMIF(A!N3:O3, "N", B!N3:O3), SUMIF(A!P3, "N", B!Q3))']

    def test_the_largest_or_smallest_matched_keeps_one_range(self, interleaved):
        # MAXIFS over a run with no match is 0, so the largest of one call
        # per run would be wrong for negative amounts.
        m, written, issues = interleaved
        for line, func in ((m.s.maxifs, "MAXIFS"), (m.s.minifs, "MINIFS")):
            assert written[("Summary", line.display_name)] == [line.value]
            (issue,) = [i for i in issues if i.qpath == line.id]
            assert issue.func == func and "'totals_at': 'end'" in str(issue)


# ══════════════════════════════════════════════════════════════════════
# The readers of one ordered range keep the value, and say why
# ══════════════════════════════════════════════════════════════════════


#: Each reads one ordered range: over a row that is not one run, the value.
_READERS = ("maxifs", "minifs", "match_max", "match_key", "index", "index_last", "index_match",
            "xlookup", "hlookup", "large", "small", "percentile_inc", "quartile_inc", "rank_eq",
            "irr", "textjoin", "match_slice_in")


class TestTheReadersOfOneRangeKeepTheValue:
    def test_the_cell_holds_the_engines_value(self, interleaved):
        m, written, _issues = interleaved
        for name in _READERS:
            value = getattr(m.s, name).value
            assert written[("Summary", name)] == [pytest.approx(value)], name
        assert written[("Data", "d_match_max")] == [m.data.d_match_max.value]
        assert written[("Data", "d_large")] == [m.data.d_large.value]

    def test_check_compat_names_the_call_and_the_way_out(self, interleaved):
        m, _written, issues = interleaved
        (issue,) = [i for i in issues if i.qpath == m.s.match_max.id]
        assert issue.func == "MATCH"
        assert str(issue) == (
            f"{m.s.match_max.id}[0]: written as its value — no formula stands for "
            "MATCH((MAX(customers)), customers, 0); customers is read as one range, and the "
            "quarter and year total columns stand between its periods: put them after the last "
            "period, timeline={'totals': ['quarter', 'year'], 'totals_at': 'end'}")
        listed = {i.qpath for i in issues}
        assert {getattr(m.s, name).id for name in _READERS} <= listed

    def test_index_may_read_its_period_by_position(self, interleaved):
        # A literal position is that period's; the line's COUNT the last;
        # any other index names no position.
        m, _written, issues = interleaved
        hints = {i.qpath: i.hint for i in issues}
        assert hints[m.s.index.id].endswith(
            "'totals_at': 'end'}, or read that period by position, revenue[2]")
        assert hints[m.s.index_last.id].endswith(
            "'totals_at': 'end'}, or read the last period by position, revenue[-1]")
        for line in (m.s.index_match, m.s.match_max, m.s.large):
            assert "by position" not in hints[line.id]

    def test_a_block_of_rows_names_its_rows(self, interleaved):
        m, _written, issues = interleaved
        (issue,) = [i for i in issues if i.qpath == m.s.hlookup.id]
        assert issue.hint == (
            "customers and revenue are read as one range, and the quarter and year total "
            "columns stand between their periods: put them after the last period, "
            "timeline={'totals': ['quarter', 'year'], 'totals_at': 'end'}")

    @pytest.mark.parametrize("at", ["between", "end"])
    def test_a_block_of_rows_apart_has_no_way_out(self, at):
        # A line stands between the two rows: no range holds the block in
        # either placement, so the total columns are not the way out.
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=6,
                     default_excel_view=mo.ExcelView(
                         timeline={"totals": ["quarter"], "totals_at": at}))
        with m:
            m.a = mo.MultiVariable("A", excel_props={"tab": True})
            m.a.x = mo.Variable([1.0, 2, 3, 4, 5, 6], display_name="X", regrain=mo.up("last"))
            m.a.gap = mo.Variable(["n", "s", "e", "n", "s", "e"], display_name="Gap")
            m.a.y = mo.Variable([10.0, 20, 30, 40, 50, 60], display_name="Y",
                                regrain=mo.up("sum"))
            m.s = mo.MultiVariable("S", excel_props={"tab": True})
            m.s.h = mo.HLOOKUP(3.0, [m.a.x, m.a.y], 2, False).set_display_name("h")
            m.s.i = mo.INDEX([m.a.x, m.a.y], 2, 3).set_display_name("i")
        written = _written(m)
        assert written[("S", "h")] == [30.0] and written[("S", "i")] == [30.0]
        issues = {i.qpath: i for i in check_compat(m, "excel")}
        assert set(issues) == {m.s.h.id, m.s.i.id}
        assert str(issues[m.s.h.id]) == (
            f"{m.s.h.id}[0]: written as its value — no formula stands for [x, y]")
        assert issues[m.s.i.id].hint is None

    def test_the_last_period_by_position_stays_live(self, interleaved):
        _m, written, _issues = interleaved
        assert written[("Summary", "last")] == [f"=Data!AH{_REVENUE} + 0"]

    def test_a_call_per_period_names_the_way_out_in_every_period(self, interleaved):
        m, written, issues = interleaved
        ranks = [i for i in issues if i.qpath == m.s.rank_pp.id]
        assert [i.period for i in ranks] == list(range(N))
        assert all(i.func == "RANK_EQ" and "'totals_at': 'end'" in i.hint for i in ranks)
        # (Its own total columns hold the projection's #VALUE!: a rank has no total.)
        months = [c for c in written[("Summary", "rank_pp")] if c != "#VALUE!"]
        assert months == m.s.rank_pp.value

    def test_cells_apart_for_another_reason_get_no_totals_at(self):
        # A compare layout puts a line's periods a group of columns apart,
        # and a stepped window skips periods of its own: no total column
        # stands between them, so no view of totals is the way out.
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=6,
                     tracks=mo.Tracks("plan", "actual"),
                     default_excel_view=mo.ExcelView(tracks="compare"))
        with m:
            m.d = mo.MultiVariable("Data", excel_props={"tab": True})
            m.d.rev = mo.Variable(plan=[1.0, 2, 3, 4, 5, 6], actual=[1.5, 2, 3, None, None, None],
                                  display_name="Rev", regrain=mo.up("sum"))
            m.d.flat = mo.Variable([3.0, 1, 4, 1, 5, 9], display_name="Flat",
                                   regrain=mo.up("sum"))
            m.s = mo.MultiVariable("S", excel_props={"tab": True})
            m.s.top = mo.MATCH(mo.MAX(m.d.flat), m.d.flat, 0).set_display_name("top")
            m.s.odd = mo.MATCH(mo.MAX(m.d.flat[0:6:2]), m.d.flat[0:6:2], 0).set_display_name("odd")
        hints = {i.qpath: i.hint for i in check_compat(m, "excel")}
        assert hints == {m.s.top.id: "flat is read as one range, and its cells are not one run",
                         m.s.odd.id: "flat is read as one range, and its cells are not one run"}

    def test_a_hint_is_printed_only_where_there_is_one(self):
        bare = CompatIssue("MATCH", "value", frozenset(), "excel", qpath="m.x", period=0,
                           inlined="MATCH(a, b, 0)")
        assert str(bare) == "m.x[0]: written as its value — no formula stands for MATCH(a, b, 0)"
        hinted = CompatIssue("MATCH", "value", frozenset(), "excel", qpath="m.x", period=0,
                             inlined="MATCH(a, b, 0)", hint="b is read as one range")
        assert str(hinted) == (
            "m.x[0]: written as its value — no formula stands for MATCH(a, b, 0); "
            "b is read as one range")


# ══════════════════════════════════════════════════════════════════════
# The limits of one call
# ══════════════════════════════════════════════════════════════════════


def _long(n, totals=None, sheet="Data"):
    view = {"default_excel_view": mo.ExcelView(timeline={"totals": totals})} if totals else {}
    m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=n, **view)
    with m:
        m.d = mo.MultiVariable(sheet, excel_props={"tab": True})
        m.d.x = mo.Variable([1.0] * n, display_name="X", regrain=mo.up("sum"))
        m.d.r = mo.Variable(["a", "b"] * (n // 2), display_name="R")
        m.s = mo.MultiVariable("S", excel_props={"tab": True})
    return m


class TestTheLimitsOfOneCall:
    def test_more_runs_than_a_call_takes_are_the_value(self):
        # 800 months under quarter totals: 267 runs, past Excel's 255.
        m = _long(800, ["quarter"])
        with m.s as s:
            s.total = mo.SUM(m.d.x).set_display_name("total")
            s.cnt = mo.COUNTIF(m.d.r, "a").set_display_name("cnt")
            s.avg = mo.AVERAGEIF(m.d.r, "a", m.d.x).set_display_name("avg")
            s.sp = mo.SUM(m.d.x * m.d.x).set_display_name("sp")
        written = _written(m)
        for line in (s.total, s.cnt, s.avg, s.sp):
            assert written[("S", line.display_name)] == [line.value]
        issues = {i.qpath: i for i in check_compat(m, "excel")}
        assert set(issues) == {s.total.id, s.cnt.id, s.avg.id, s.sp.id}
        assert str(issues[s.total.id]) == (
            f"{s.total.id}[0]: written as its value — no formula stands for SUM(x); written as "
            "its runs, the call takes 267 arguments, past the 255 one call takes; the quarter "
            "total columns stand between the periods of x: put them after the last period, "
            "timeline={'totals': ['quarter'], 'totals_at': 'end'}")
        assert "the periods of r and x:" in issues[s.avg.id].hint
        assert "the periods of x:" in issues[s.sp.id].hint

    def test_sheets_with_totals_of_their_own_are_each_named(self):
        # The regions under quarter totals, the amounts under year totals:
        # 'totals_at' alone reaches both sheets, each keeping its own kinds.
        n = 800
        quarters = mo.ExcelView(timeline={"totals": ["quarter"]})
        years = mo.ExcelView(timeline={"totals": ["year"]})
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=n)
        with m:
            m.a = mo.MultiVariable("A", excel_props={"tab": True}, default_excel_view=quarters)
            m.a.r = mo.Variable(["a", "b"] * (n // 2), display_name="R")
            m.b = mo.MultiVariable("B", excel_props={"tab": True}, default_excel_view=years)
            m.b.x = mo.Variable([1.0] * n, display_name="X", regrain=mo.up("sum"))
            m.s = mo.MultiVariable("S", excel_props={"tab": True})
            m.s.si = mo.SUMIF(m.a.r, "a", m.b.x).set_display_name("si")
        (issue,) = check_compat(m, "excel")
        assert issue.hint == (
            "written as its runs, the call takes 267 arguments, past the 255 one call takes; "
            "the quarter total columns stand between the periods of r, the year total columns "
            "between those of x: put them after the last period, timeline={'totals_at': 'end'}")

    def test_the_call_s_other_arguments_count(self):
        # Every other month is a run of one cell: 255 runs are a SUM's
        # last argument; NPV's rate makes them 256.
        m = _long(520)
        with m.s as s:
            s.most = mo.SUM(m.d.x[0:510:2]).set_display_name("most")
            s.past = mo.SUM(m.d.x[0:512:2]).set_display_name("past")
            s.npv = mo.NPV(0.01, m.d.x[0:510:2]).set_display_name("npv")
        written = _written(m)
        assert written[("S", "most")][0].count(",") == 254
        assert written[("S", "past")] == [s.past.value]
        assert written[("S", "npv")] == [pytest.approx(s.npv.value)]
        hints = {i.qpath: i.hint for i in check_compat(m, "excel")}
        past = "written as its runs, the call takes 256 arguments, past the 255 one call takes"
        assert hints == {s.past.id: past, s.npv.id: past}

    def test_runs_longer_than_a_formula_holds_are_the_value(self):
        # A sheet named with 31 characters, in front of each of 200 runs.
        m = _long(600, ["quarter"], sheet="A sheet with a thirty-one chars")
        with m.s as s:
            s.total = mo.SUM(m.d.x).set_display_name("total")
        assert _written(m)[("S", "total")] == [600]
        (issue,) = check_compat(m, "excel")
        assert issue.hint.startswith(
            "written as its runs, the formula is 8,640 characters, past the 8,192 one formula "
            "holds; the quarter total columns stand between the periods of x")


# ══════════════════════════════════════════════════════════════════════
# A total column reads a row read as a range at its months
# ══════════════════════════════════════════════════════════════════════


def _buckets(n=24, kinds=("year",)):
    """Formula lines that read their sheet's row ``N`` whole, under total
    columns of ``kinds``."""
    m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=n,
                 default_excel_view=mo.ExcelView(timeline={"totals": list(kinds)}))
    with m:
        m.c = mo.MultiVariable("C", excel_props={"tab": True})
        with m.c as c:
            c.n = mo.Variable([1.0 + i % 5 for i in range(12)] + [100.0 + i for i in range(n - 12)],
                              display_name="N", regrain=mo.up("sum"))
            c.mx = (c.n / mo.MAX(c.n)).set_display_name("Mx")
            c.cnt = (c.n * 0 + mo.COUNTIF(c.n, ">50")).set_display_name("Cnt")
            c.dev = (c.n - mo.LARGE(c.n, 2)).set_display_name("Dev")
            c.gap = (c.n - mo.AVERAGE(c.n)).set_display_name("Gap")
            c.mid = (c.n - mo.MEDIAN(c.n)).set_display_name("Mid")
            c.share = (c.n / mo.SUM(c.n * c.n)).set_display_name("Share")
    return m


class TestTheTotalColumnsReadTheMonths:
    def test_a_row_read_as_a_range_is_its_months(self):
        # Never MAX(N3, AA3) over the year totals: 0.01 for an engine's 0.12.
        rows = _written(_buckets())
        assert rows[("C", "Mx")][12] == "=N3 / MAX(B3:M3, O3:Z3)"
        assert rows[("C", "Mx")][25] == "=AA3 / MAX(B3:M3, O3:Z3)"
        assert rows[("C", "Cnt")][12] == (
            '=N3 * 0 + SUM(COUNTIF(B3:M3, ">50"), COUNTIF(O3:Z3, ">50"))')

    def test_a_reader_of_one_range_is_the_projections_value(self):
        m = _buckets()
        rows = _written(m)
        assert [rows[("C", "Dev")][i] for i in (12, 25)] == m.at("year").c.dev.value

    def test_one_year_reads_one_range(self):
        rows = _written(_buckets(12))
        assert rows[("C", "Mx")][12] == "=N3 / MAX(B3:M3)"
        assert rows[("C", "Cnt")][12] == '=N3 * 0 + COUNTIF(B3:M3, ">50")'

    def test_a_tracked_line_s_total_reads_the_months_too(self):
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=6,
                     tracks=mo.Tracks("plan", "actual"),
                     default_excel_view=mo.ExcelView(timeline={"totals": ["quarter"]},
                                                     tracks="rows"))
        with m:
            m.c = mo.MultiVariable("C", excel_props={"tab": True})
            with m.c as c:
                c.n = mo.Variable(plan=[1.0, 2, 3, 40, 50, 60],
                                  actual=[1.5, 2, 3, 41, None, None],
                                  display_name="N", regrain=mo.up("sum"))
                c.flat = mo.Variable([1.0, 2, 3, 40, 50, 60], display_name="Flat",
                                     regrain=mo.up("sum"))
                c.mx = (c.n / mo.MAX(c.flat)).set_display_name("Mx")
                c.lg = (c.n - mo.LARGE(c.flat, 2)).set_display_name("Lg")
        rows = _written(m)
        # E and I are the quarter totals; Flat's months are B:D and F:H.
        assert rows[("C", "Mx")][3] == "=E3 / MAX(B5:D5, F5:H5)"
        assert rows[("C", "Mx · actual")][7] == "=I4 / MAX(B5:D5, F5:H5)"
        quarters = m.at("quarter").c.lg.value
        assert [rows[("C", "Lg")][i] for i in (3, 7)] == quarters["plan"]
        assert [rows[("C", "Lg · actual")][i] for i in (3, 7)] == quarters["actual"]


# ══════════════════════════════════════════════════════════════════════
# LibreOffice recalculates the runs to the engine's numbers
# ══════════════════════════════════════════════════════════════════════


def _soffice():
    for candidate in ("soffice", "libreoffice"):
        found = shutil.which(candidate)
        if found:
            return found
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


def _recalculated(model, tmp_path):
    """The book after LibreOffice recalculates it (conftest seeds the
    profile to recalculate on load)."""
    soffice = _soffice()
    if soffice is None:
        pytest.skip("LibreOffice not installed")
    path = tmp_path / "book.xlsx"
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        model.to_excel(path)
    subprocess.run(
        [soffice, f"-env:UserInstallation={(tmp_path / 'prof').as_uri()}",
         "--headless", "--convert-to", "xlsx", "--outdir",
         str(tmp_path / "calc"), str(path)],
        check=True, capture_output=True,
    )
    return load_workbook(tmp_path / "calc" / "book.xlsx", data_only=True)


def _row(ws, label):
    return next(row for row in ws.iter_rows() if row[0].value == label)


def _value(ws, label):
    """A scalar line's one cell."""
    return next(c.value for c in _row(ws, label)[1:] if c.value is not None)


def _under(ws, label, header):
    """A line's cells in the columns whose caption (over the header rows)
    passes ``header``: the periods', or a kind of total's."""
    captions = {c: next((ws.cell(r, c).value for r in range(1, 4) if ws.cell(r, c).value), None)
                for c in range(2, ws.max_column + 1)}
    return [cell.value for cell in _row(ws, label)[1:] if header(captions[cell.column])]


def _month(caption):
    return isinstance(caption, str) and not caption.startswith("Total")


def _serial(value):
    """A date as the cell holds it: the recalculated cell is a number."""
    if isinstance(value, dt.datetime):
        return (value - dt.datetime(1899, 12, 30)) / dt.timedelta(days=1)
    if isinstance(value, dt.date):
        return (value - dt.date(1899, 12, 30)).days
    return value


#: A criteria row with blank cells at the ends of its runs and inside one,
#: text and booleans; the averaged rows with a blank, a text, a number
#: written as text, a negative number — and dates.
_CRIT = [None, "North", None, "South", None, "North", True, "North", False, "South", "North", None,
         None, "North", "text", "North", "South", "North", "North", "South", "North", "East",
         "North", "North"]
_AVG = [1.0, 10.0, 2.0, 7.0, 3.0, None, 8.0, "n/a", 9.0, 11.0, "12", 4.0,
        5.0, 20.0, 13.0, 30.0, 14.0, 40.0, -5.0, 15.0, 50.0, 16.0, 60.0, 70.0]


def _means(avg, kinds=("quarter", "year")):
    """Conditional means over 24 months under total columns of ``kinds`` —
    and the same rows on a sheet without them, read as one range each."""
    m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=N,
                 default_excel_view=mo.ExcelView(timeline={"totals": list(kinds)}))
    with m:
        for name, view in (("d", None), ("p", mo.ExcelView(timeline={"totals": []}))):
            sheet = mo.MultiVariable("Data" if view is None else "Plain",
                                     excel_props={"tab": True}, default_excel_view=view)
            setattr(m, name, sheet)
            sheet.crit = mo.Variable(_CRIT, display_name="Crit")
            sheet.avg = mo.Variable(avg, display_name="Avg")
            sheet.tag = mo.Variable([f"x{i}" if i % 4 else "y" for i in range(N)],
                                    display_name="Tag")
            sheet.when = mo.Variable([dt.date(2026, 1, 1) + dt.timedelta(days=17 * i)
                                      if i % 5 else None for i in range(N)], display_name="When")
        m.s = mo.MultiVariable("Summary", excel_props={"tab": True})
        for name, d in (("", m.d), ("plain_", m.p)):
            with m.s as s:
                setattr(s, f"{name}north", mo.AVERAGEIF(d.crit, "North", d.avg)
                        .set_display_name(f"{name}north"))
                setattr(s, f"{name}not_south", mo.AVERAGEIF(d.crit, "<>South", d.avg)
                        .set_display_name(f"{name}not_south"))
                setattr(s, f"{name}blank", mo.AVERAGEIF(d.crit, "", d.avg)
                        .set_display_name(f"{name}blank"))
                setattr(s, f"{name}own", mo.AVERAGEIF(d.avg, ">0").set_display_name(f"{name}own"))
                setattr(s, f"{name}two", mo.AVERAGEIFS(d.avg, d.crit, "<>South", d.tag, "x*")
                        .set_display_name(f"{name}two"))
                setattr(s, f"{name}dates", mo.AVERAGEIF(d.crit, "<>South", d.when)
                        .set_display_name(f"{name}dates"))
                setattr(s, f"{name}count", mo.COUNTIF(d.crit, "<>South")
                        .set_display_name(f"{name}count"))
                setattr(s, f"{name}counts", mo.COUNTIFS(d.tag, "x*", d.crit, "<>South")
                        .set_display_name(f"{name}counts"))
                setattr(s, f"{name}sumif", mo.SUMIF(d.crit, "", d.avg)
                        .set_display_name(f"{name}sumif"))
    return m


_MEANS = ("north", "not_south", "blank", "own", "two", "dates", "count", "counts", "sumif")


@pytest.mark.slow
class TestLibreOfficeAgrees:
    def test_the_runs_recalculate_to_the_engines_values(self, tmp_path):
        m = _matrix(["quarter", "year"])
        book = _recalculated(m, tmp_path)
        for sheet, mv, names in (("Summary", m.s, [name for name, _ in _SCALARS]),
                                 ("Data", m.data, [name for name, _ in _SAME_SHEET])):
            for name in names:
                got = _value(book[sheet], name)
                assert got == pytest.approx(getattr(mv, name).value), (name, got)
        for name, _build in _PER_PERIOD:
            got = _under(book["Summary"], name, _month)
            assert got == pytest.approx(getattr(m.s, name).value), name

    def test_a_conditional_mean_counts_what_it_averages(self, tmp_path):
        # A blank, a text and a number written as text are not averaged; a
        # blank criteria cell at the end of a run meets '<>South' and ''.
        m = _means(_AVG)
        ws = _recalculated(m, tmp_path)["Summary"]
        for name in _MEANS:
            got = _serial(_value(ws, name))
            assert got == pytest.approx(_serial(getattr(m.s, name).value)), (name, got)
            # The runs read as the one range does, cell for cell.
            assert got == pytest.approx(_serial(_value(ws, f"plain_{name}"))), name

    def test_a_boolean_averaged_reads_as_the_one_range_does(self, tmp_path):
        # LibreOffice counts TRUE and FALSE as 1 and 0 in its own AVERAGEIF
        # too: the runs read the cells as the one range does.
        avg = list(_AVG)
        avg[1], avg[15] = True, False
        ws = _recalculated(_means(avg), tmp_path)["Summary"]
        for name in ("north", "not_south", "own", "two"):
            assert _value(ws, name) == pytest.approx(_value(ws, f"plain_{name}")), name

    @pytest.mark.xfail(strict=True, reason=(
        "LibreOffice reads a boolean cell as a number, in AVERAGEIF over one range too; "
        "Excel skips it, and the engine keeps Excel's rule"))
    def test_where_libreoffice_departs_from_excel(self, tmp_path):
        avg = list(_AVG)
        avg[1], avg[15] = True, False
        m = _means(avg)
        ws = _recalculated(m, tmp_path)["Summary"]
        assert _value(ws, "north") == pytest.approx(m.s.north.value)

    def test_the_total_columns_recalculate_to_the_projection(self, tmp_path):
        m = _buckets(kinds=("quarter", "year"))
        ws = _recalculated(m, tmp_path)["C"]
        for kind, word in (("quarter", "Total Q"), ("year", "Total 20")):
            projected = m.at(kind).c
            for label, attr in (("Mx", "mx"), ("Cnt", "cnt"), ("Dev", "dev"), ("Gap", "gap"),
                                ("Mid", "mid"), ("Share", "share")):
                got = _under(ws, label, lambda caption: str(caption).startswith(word))
                assert got == pytest.approx(getattr(projected, attr).value), (kind, label, got)
