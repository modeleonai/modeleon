# SPDX-License-Identifier: Apache-2.0
"""A book carries no green triangles of its writer's own making.

Excel checks a sheet as it opens it and marks a suspected mistake with a green
triangle in the cell's corner, which a reader takes for a note that is not
there. Text that reads as a number (a line's number "1.1", a year label) and
the formulas the writer put in, which differ from their neighbours by design
(a quarter's total between its months), are recorded in the sheet as checked.
A typed value, a label and what a reader types later keep Excel's checks.
"""

import re
import warnings
import zipfile

import openpyxl

import modeleon as mo
from modeleon.compile.excel.cached_values import _ranges, _with_ignored_errors

LEAD = {"article": "No", "label": "Line item", "numbering": "section"}


def _model():
    m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=6)
    with m:
        m.default_excel_view = mo.ExcelView(meta=LEAD, timeline={"totals": ["quarter"]})
        m.ops = mo.MultiVariable("Operations")
        with m.ops as o:
            o.lease = mo.MultiVariable("Lease")
            with o.lease as s:
                s.volume = mo.Variable([1.0, 2.0, 3.0, 4.0, 5.0, 6.0], display_name="Volume",
                                       regrain=mo.up("sum"))
                s.revenue = mo.Variable(s.volume * 2, display_name="Revenue")
    return m


def _sheet_xml(model, tmp_path, sheet="Operations"):
    path = tmp_path / "book.xlsx"
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        model.to_excel(str(path))
    book = openpyxl.load_workbook(str(path))           # still a book Excel can read
    index = book.sheetnames.index(sheet) + 1
    return book[sheet], zipfile.ZipFile(path).read(f"xl/worksheets/sheet{index}.xml").decode()


def _checked(xml):
    """``{rule attribute: set of cells}`` the sheet records as checked."""
    out: dict[str, set] = {}
    for sqref, attrs in re.findall(r'<ignoredError sqref="([^"]+)"([^>]*)/>', xml):
        cells = set()
        for part in sqref.split():
            a, _, b = part.partition(":")
            (c1, r1), (c2, r2) = (re.match(r"([A-Z]+)(\d+)", x).groups() for x in (a, b or a))
            for r in range(int(r1), int(r2) + 1):
                for c in range(openpyxl.utils.column_index_from_string(c1),
                               openpyxl.utils.column_index_from_string(c2) + 1):
                    cells.add(f"{openpyxl.utils.get_column_letter(c)}{r}")
        for rule in re.findall(r'(\w+)="1"', attrs):
            out.setdefault(rule, set()).update(cells)
    return out


def test_line_numbers_and_the_writers_formulas_are_recorded_as_checked(tmp_path):
    ws, xml = _sheet_xml(_model(), tmp_path)
    checked = _checked(xml)
    numbers = {c.coordinate for c in ws["A"] if isinstance(c.value, str) and c.value[:1].isdigit()}
    formulas = {c.coordinate for row in ws.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=")}
    assert numbers and formulas
    assert numbers <= checked["numberStoredAsText"]
    assert formulas <= checked["formula"] and formulas <= checked["formulaRange"]
    # One entry per cell: Excel keeps a single entry for a cell, so the
    # text and the formula entries never meet.
    assert not checked["numberStoredAsText"] & checked["formula"]


def test_a_typed_value_and_a_label_keep_excels_checks(tmp_path):
    ws, xml = _sheet_xml(_model(), tmp_path)
    checked = set().union(*_checked(xml).values())
    typed = {c.coordinate for row in ws.iter_rows() for c in row
             if isinstance(c.value, (int, float)) and not isinstance(c.value, bool)}
    labels = {c.coordinate for c in ws["B"] if c.value in ("Volume", "Revenue")}
    assert typed and labels
    assert not typed & checked and not labels & checked
    # A real error stays marked: no entry silences the check for one.
    assert "evalError" not in xml


def test_the_record_stands_where_a_worksheet_keeps_it():
    sheet = ('<worksheet><sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>1.1</t></is></c>'
             '<c r="B1" t="inlineStr"><is><t>Revenue</t></is></c><c r="C1"><f>B2*2</f><v>4</v></c>'
             '</row></sheetData><pageMargins left="0.75"/>'
             '<tableParts count="1"><tablePart/></tableParts></worksheet>')
    out = _with_ignored_errors(sheet)
    assert out.index("<pageMargins") < out.index("<ignoredErrors>") < out.index("<tableParts")
    assert '<ignoredError sqref="A1" numberStoredAsText="1" twoDigitTextYear="1"/>' in out
    assert '<ignoredError sqref="C1" formula="1" formulaRange="1"/>' in out
    plain = ('<worksheet><sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>Revenue</t></is></c>'
             '<c r="B1"><v>4</v></c></row></sheetData><pageMargins left="0.75"/></worksheet>')
    assert _with_ignored_errors(plain) == plain


def test_cells_are_named_as_few_ranges():
    cells = [("A", 3), ("A", 4), ("A", 5), ("B", 3), ("B", 4), ("B", 5),
             ("C", 1), ("D", 1), ("E", 1), ("A", 7)]
    assert _ranges(cells) == ["C1:E1", "A3:B5", "A7"]
