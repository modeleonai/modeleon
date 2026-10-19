# SPDX-License-Identifier: Apache-2.0
"""A lag's fill is ONE operand of the formula around it.

``x - mo.lag(y, fill=a + b - c)`` writes the fill in the periods the
shift exposes. A fill that is a number or a named cell is one token; an
unnamed compound fill is inlined, and used to lose its brackets —
``=B2 - I!B1 + I!B2 - I!B3`` where the model says ``200 - (120 + 90 -
50)``. The cached value hid it until a recalculation. The brackets now
come with the fill; the other periods and the simple fills render
exactly as before.
"""

import shutil
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook

import modeleon as mo


def _rows(ws):
    out = {}
    for row in ws.iter_rows():
        label = row[0].value
        if label:
            out[str(label)] = [c.value for c in row[1:]]
    return out


def _book(model, tmp_path, sheet):
    path = tmp_path / "book.xlsx"
    model.to_excel(path)
    return _rows(load_workbook(path)[sheet])


def _model():
    m = mo.Model("m", default_grain="year", default_start="2026", default_periods=3)
    with m:
        m.i = mo.MultiVariable("I")
        m.i.a = mo.Variable(120.0, display_name="a")
        m.i.b = mo.Variable(90.0, display_name="b")
        m.i.c = mo.Variable(50.0, display_name="c")
        m.i.k = mo.Variable(7.0, display_name="k")
        m.x = mo.MultiVariable("X")
        m.x.nwc = mo.Variable([200.0, 210.0, 220.0], display_name="nwc")
    return m


class TestACompoundFillKeepsItsBrackets:
    @pytest.mark.parametrize("build, formula", [
        (lambda x, i: x.nwc - mo.lag(x.nwc, fill=i.a + i.b - i.c),
         "=B2 - (I!B1 + I!B2 - I!B3)"),
        (lambda x, i: x.nwc * mo.lag(x.nwc, fill=i.a + i.b),
         "=B2 * (I!B1 + I!B2)"),
        (lambda x, i: x.nwc / mo.lag(x.nwc, fill=i.a * i.b),
         "=B2 / (I!B1 * I!B2)"),
        (lambda x, i: x.nwc ** mo.lag(x.nwc, fill=i.a - i.b),
         "=B2 ^ (I!B1 - I!B2)"),
        (lambda x, i: mo.lag(x.nwc, fill=i.a - i.b) - x.nwc,
         "=(I!B1 - I!B2) - B2"),
        (lambda x, i: -mo.lag(x.nwc, fill=i.a - i.b),
         "=-(I!B1 - I!B2)"),
        (lambda x, i: x.nwc * mo.lag(x.nwc, fill=i.a > i.b),
         "=B2 * (I!B1 > I!B2)"),
    ])
    def test_under_each_operator(self, tmp_path, build, formula):
        m = _model()
        with m:
            m.x.d = build(m.x, m.i).set_display_name("d")
        d = _book(m, tmp_path, "X")["d"]
        assert d[0] == formula, d
        # The periods the shift does not expose read the previous cell.
        assert "I!" not in str(d[1]) and "I!" not in str(d[2]), d

    def test_wherever_the_fill_is_written(self, tmp_path):
        m = _model()
        with m:
            # A lead exposes the LAST period.
            m.x.lead = (m.x.nwc - mo.lag(m.x.nwc, -1, fill=m.i.a + m.i.b)).set_display_name("lead")
            # Two periods back exposes two periods.
            m.x.two = (m.x.nwc - mo.lag(m.x.nwc, 2, fill=m.i.a + m.i.b)).set_display_name("two")
            # A lag of an unnamed line writes the fill the same way.
            m.x.unlaid = (m.x.nwc - mo.lag(m.x.nwc * 2.0, fill=m.i.a + m.i.b)).set_display_name("u")
        x = _book(m, tmp_path, "X")
        assert x["lead"] == ["=B2 - C2", "=C2 - D2", "=D2 - (I!B1 + I!B2)"], x["lead"]
        assert x["two"] == ["=B2 - (I!B1 + I!B2)", "=C2 - (I!B1 + I!B2)", "=D2 - B2"], x["two"]
        assert x["u"] == ["=B2 - (I!B1 + I!B2)", "=C2 - (B2 * 2.0)", "=D2 - (C2 * 2.0)"], x["u"]


class TestEverythingElseIsUnchanged:
    def test_simple_fills_and_the_whole_cell(self, tmp_path):
        m = _model()
        with m:
            m.x.lit = (m.x.nwc - mo.lag(m.x.nwc, fill=5.0)).set_display_name("lit")
            m.x.ref = (m.x.nwc - mo.lag(m.x.nwc, fill=m.i.k)).set_display_name("ref")
            m.x.fn = (m.x.nwc - mo.lag(m.x.nwc, fill=mo.ABS(m.i.a - m.i.b))).set_display_name("fn")
            m.x.whole = mo.lag(m.x.nwc, fill=m.i.a + m.i.b - m.i.c).set_display_name("whole")
        x = _book(m, tmp_path, "X")
        assert x["lit"] == ["=B2 - 5.0", "=C2 - B2", "=D2 - C2"], x["lit"]
        assert x["ref"] == ["=B2 - I!B4", "=C2 - B2", "=D2 - C2"], x["ref"]
        assert x["fn"] == ["=B2 - ABS(I!B1 - I!B2)", "=C2 - B2", "=D2 - C2"], x["fn"]
        # The whole cell is the fill: no brackets around a whole formula.
        assert x["whole"] == ["=I!B1 + I!B2 - I!B3", "=B2", "=C2"], x["whole"]


def _soffice():
    for candidate in ("soffice", "libreoffice"):
        found = shutil.which(candidate)
        if found:
            return found
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


@pytest.mark.slow
def test_the_workbook_recalculates_to_the_python_values(tmp_path):
    soffice = _soffice()
    if soffice is None:
        pytest.skip("LibreOffice not installed")
    m = _model()
    with m:
        m.x.ex = mo.Variable([2.0, 3.0, 4.0], display_name="ex")
        m.x.minus = (m.x.nwc - mo.lag(m.x.nwc, fill=m.i.a + m.i.b - m.i.c)).set_display_name("minus")
        m.x.times = (m.x.nwc * mo.lag(m.x.nwc, fill=m.i.a - m.i.b)).set_display_name("times")
        m.x.over = (m.x.nwc / mo.lag(m.x.nwc, fill=m.i.a - m.i.b)).set_display_name("over")
        m.x.power = (m.x.ex ** mo.lag(m.x.ex, fill=m.i.a / m.i.b)).set_display_name("power")
    path = tmp_path / "fill.xlsx"
    m.to_excel(path)
    subprocess.run(
        [soffice, f"-env:UserInstallation={(tmp_path / 'prof').as_uri()}",
         "--headless", "--convert-to", "xlsx", "--outdir",
         str(tmp_path / "calc"), str(path)],
        check=True, capture_output=True,
    )
    x = _rows(load_workbook(tmp_path / "calc" / "fill.xlsx", data_only=True)["X"])
    assert m.x.minus.value == [40.0, 10.0, 10.0]
    for name in ("minus", "times", "over", "power"):
        got = [round(v, 6) for v in x[name][:3]]
        assert got == [round(v, 6) for v in getattr(m.x, name).value], name
