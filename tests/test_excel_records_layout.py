"""``ExcelView(orient='records')`` — records laid out as a table.

The default stacked layout gives every child MV a header row and a
block of its own. That is right for sections of a model and wrong for a
LIST: 29 employees of five fields each become 29 headers and ~150 rows
repeating the same five labels, and no one writes that by hand.

Records orientation says the shape out loud, and the fields split by
whether they carry a period axis:

* scalar fields → a table, fields across, records down;
* per-period fields → one block each, titled once, records under it,
  values landing in the sheet's own period columns.

Both halves keep real addresses, so formulas stay live and the grid and
the .xlsx are one picture.

The form is DECLARED, never assumed: a container prints as stacked
sections unless its own view says otherwise. Nothing cascades it — a
model that declared it would turn every section into a record table,
and a section whose children are Variables has no records to lay out,
so its rows would silently vanish. A shape this violent is stated
where it applies.
"""

import openpyxl
import pytest

import modeleon as mo


def _grid(model, sheet, cols=6):
    import os
    import tempfile

    path = os.path.join(tempfile.mkdtemp(), "t.xlsx")
    model.to_excel(path)
    ws = openpyxl.load_workbook(path)[sheet]
    return [
        [ws.cell(r, c).value for c in range(1, cols + 1)]
        for r in range(1, ws.max_row + 1)
    ]


class Сотрудник(mo.MultiVariableClass):
    def compute(self, оклад):
        self.оклад = mo.Variable(оклад, display_name="Оклад")
        self.налоги = mo.Variable(оклад * 0.1, display_name="Налоги")


def _register():
    m = mo.Model("m", display_name="M")
    with m:
        m.штат = mo.MultiVariable("Штат", excel_props={"tab": True})
        m.штат.default_excel_view = mo.ExcelView(orient="records")
        with m.штат as c:
            c.s1 = Сотрудник(оклад=100, display_name="Сотрудник №1")
            c.s2 = Сотрудник(оклад=200, display_name="Сотрудник №2")
    return m


class TestScalarRecords:
    def test_fields_across_records_down(self):
        g = _grid(_register(), "Штат")
        assert g[0][:3] == [None, "Оклад", "Налоги"]
        assert g[1][:3] == ["Сотрудник №1", 100, 10]
        assert g[2][:3] == ["Сотрудник №2", 200, 20]
        assert len(g) == 3  # no per-record sections

    def test_a_field_only_some_records_have_keeps_its_column(self):
        # First-appearance order, and a record that lacks a field
        # leaves the cell empty rather than shifting the row — the
        # failure that turns a table into misaligned nonsense.
        class Свой(mo.MultiVariableClass):
            def compute(self, оклад, бонус=None):
                self.оклад = mo.Variable(оклад, display_name="Оклад")
                if бонус is not None:
                    self.бонус = mo.Variable(бонус, display_name="Бонус")

        m = mo.Model("m", display_name="M")
        with m:
            m.штат = mo.MultiVariable(
                "Штат", excel_props={"tab": True, "orient": "records"}
            )
            with m.штат as c:
                c.a = Свой(оклад=1, display_name="A")
                c.b = Свой(оклад=2, бонус=9, display_name="B")
                c.c = Свой(оклад=3, display_name="C")
        g = _grid(m, "Штат")
        assert g[0][:3] == [None, "Оклад", "Бонус"]
        assert g[1][:3] == ["A", 1, None]
        assert g[2][:3] == ["B", 2, 9]
        assert g[3][:3] == ["C", 3, None]

    def test_the_first_record_names_the_column(self):
        # Records of one class usually agree on a field's display name.
        # When they don't, the header must be the first record's — a
        # stable rule — not "whoever the walk wrote last".
        class Свой(mo.MultiVariableClass):
            def compute(self, оклад, подпись):
                self.оклад = mo.Variable(оклад, display_name=подпись)

        m = mo.Model("m", display_name="M")
        with m:
            m.штат = mo.MultiVariable(
                "Штат", excel_props={"tab": True, "orient": "records"}
            )
            with m.штат as c:
                c.a = Свой(оклад=1, подпись="Оклад", display_name="A")
                c.b = Свой(оклад=2, подпись="Зарплата", display_name="B")
        g = _grid(m, "Штат")
        assert g[0][1] == "Оклад"
        assert [row[1] for row in g[1:3]] == [1, 2]

    def test_no_period_header_over_a_scalar_register(self):
        # The sheet's columns are FIELDS here. A period header would
        # name column B «Jan 2026» when B is «Оклад» — the same lie
        # ``orient='columns'`` refuses, and it must not cost a
        # reserved blank row either.
        m = _register()
        m.default_start, m.default_grain, m.default_periods = (
            "2026-01",
            "month",
            3,
        )
        g = _grid(m, "Штат")
        assert g[0][:2] == [None, "Оклад"]


class TestPerPeriodRecords:
    """A series cannot be a column of a record table: that table's
    columns are fields and a series' columns are periods."""

    def _model(self):
        m = mo.Model("m", display_name="M")
        with m:
            m.ставка = mo.Variable(0.5, display_name="Ставка")

            class Проект(mo.MultiVariableClass):
                def compute(self, цена):
                    self.цена = mo.Variable(цена, display_name="Цена")
                    self.профиль = mo.Variable(
                        [0.2, 0.3, 0.5], display_name="Профиль"
                    )
                    self.выручка = self.профиль * self.цена
                    self.доля = self.выручка * m.ставка

            m.проекты = mo.MultiVariable(
                "Проекты", excel_props={"tab": True, "orient": "records"}
            )
            with m.проекты as c:
                c.p1 = Проект(цена=100, display_name="Проект №1")
                c.p2 = Проект(цена=200, display_name="Проект №2")
        return m

    def test_one_block_per_field_records_under_it(self):
        g = _grid(self._model(), "Проекты")
        labels = [row[0] for row in g]
        # Scalar table first, then a block per series field.
        assert labels[:3] == [None, "Проект №1", "Проект №2"]
        assert [x for x in labels if x in ("Профиль", "Выручка", "Доля")] == [
            "Профиль",
            "Выручка",
            "Доля",
        ]
        i = labels.index("Профиль")
        assert labels[i + 1 : i + 3] == ["Проект №1", "Проект №2"]
        # Period values start PAST the fields table (one scalar field
        # here → column C onward), so the month masthead never paints
        # over «Цена».
        assert g[i + 1][2:5] == pytest.approx([0.2, 0.3, 0.5])
        assert g[i + 2][2:5] == pytest.approx([0.2, 0.3, 0.5])
        # And the computed block reads ACROSS the blocks: profile row ×
        # the record's own price cell in the table above.
        j = labels.index("Выручка")
        assert g[j + 1][2] == f"=C{i + 2} * B2"

    def test_formulas_stay_live_across_the_blocks(self):
        # The whole reason a collection cannot be flattened to a value
        # table: «доля» references «выручка» in another block and the
        # model's rate on another sheet, and both must survive as cell
        # references, not as numbers.
        g = _grid(self._model(), "Проекты")
        labels = [row[0] for row in g]
        i = labels.index("Доля")
        first = g[i + 1][2]
        assert isinstance(first, str) and first.startswith("=")
        assert "M!" in first  # the model's rate, by reference

    def test_months_never_paint_over_the_fields_table(self):
        """The defect this geometry exists to prevent: «Город» under
        «Feb 2026». The month masthead is read back from the period
        blocks' value addresses; those must start past the widest
        scalar column, leaving row 1 over the table empty."""
        m = self._model()
        m.default_start, m.default_grain, m.default_periods = (
            "2026-01", "month", 3,
        )
        import tempfile, pathlib as _pl
        with tempfile.TemporaryDirectory() as d:
            path = _pl.Path(d) / "m.xlsx"
            m.to_excel(str(path))
            ws = openpyxl.load_workbook(str(path))["Проекты"]
            # Column B is «Цена» — the field header owns it, and no
            # month label may stand above it.
            assert ws["B1"].value is None
            # The first month lands past the table, over the blocks.
            row1 = [ws.cell(1, c).value for c in range(1, 8)]
            months = [x for x in row1 if isinstance(x, str) and "2026" in x]
            assert months, row1
            first_month_col = row1.index(months[0]) + 1
            assert first_month_col >= 3

    def test_a_field_named_once_per_block_not_once_per_record(self):
        g = _grid(self._model(), "Проекты")
        titles = [row[0] for row in g]
        assert titles.count("Выручка") == 1

    def test_records_name_themselves_in_every_block(self):
        # Rows without an identity are the defect this shape exists to
        # fix: 29 salaries with nothing saying whose.
        g = _grid(self._model(), "Проекты")
        assert [row[0] for row in g].count("Проект №1") == 4  # table + 3


class TestEdges:
    def test_a_collection_with_no_records_lays_out_nothing(self):
        m = mo.Model("m", display_name="M")
        with m:
            m.пусто = mo.MultiVariable(
                "Пусто", excel_props={"tab": True, "orient": "records"}
            )
            m.x = mo.Variable(1, display_name="X")
        # No crash, and the model's own rows are untouched.
        assert _grid(m, "M")[0][:2] == ["X", 1]

    def test_records_with_no_variable_fields_lay_out_nothing(self):
        # A record built purely from constructor kwargs has no cells to
        # place — the layout must not invent rows for it.
        class Пустой(mo.MultiVariableClass):
            def compute(self, что):
                self._что = что

        m = mo.Model("m", display_name="M")
        with m:
            m.x = mo.Variable(1, display_name="X")
            m.список = mo.MultiVariable(
                "Список", excel_props={"tab": True, "orient": "records"}
            )
            with m.список as c:
                c.a = Пустой(что=1, display_name="A")
        assert _grid(m, "Список") in ([], [[None] * 6])


class TestTheFormIsDeclared:
    def test_a_container_without_a_view_stays_stacked(self):
        # The default every collection prints as: a header per record,
        # its fields beneath. Right for records that carry per-period
        # lines, where the table's own repetition is worse than what it
        # replaces.
        m = mo.Model("m", display_name="M")
        with m:
            m.штат = mo.MultiVariable("Штат", excel_props={"tab": True})
            with m.штат as c:
                c.s1 = Сотрудник(оклад=100, display_name="Сотрудник №1")
        g = _grid(m, "Штат")
        assert [row[0] for row in g][:3] == ["Сотрудник №1", "Оклад", "Налоги"]

    def test_the_legacy_excel_props_spelling_still_wins(self):
        # Models in the wild carry ``excel_props={'orient': …}``; they
        # must stay byte-identical.
        m = mo.Model("m", display_name="M")
        with m:
            m.штат = mo.MultiVariable(
                "Штат", excel_props={"tab": True, "orient": "records"}
            )
            with m.штат as c:
                c.s1 = Сотрудник(оклад=100, display_name="Сотрудник №1")
        assert _grid(m, "Штат")[0][:2] == [None, "Оклад"]

    def test_an_ancestor_cannot_impose_the_shape(self):
        # ``resolve_excel_view`` cascades ink down the tree. Geometry
        # must not ride it: a model-level ``orient='records'`` would
        # reach a section whose children are Variables — no records to
        # lay out — and its rows would disappear from the sheet.
        m = mo.Model("m", display_name="M")
        m.default_excel_view = mo.ExcelView(orient="records")
        with m:
            m.pnl = mo.MultiVariable("PnL", excel_props={"tab": True})
            with m.pnl as p:
                p.revenue = mo.Variable(100, display_name="Revenue")
                p.cogs = mo.Variable(60, display_name="COGS")
        g = _grid(m, "PnL")
        assert [row[0] for row in g] == ["Revenue", "COGS"]
