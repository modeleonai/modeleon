# SPDX-License-Identifier: Apache-2.0
"""Tests for ``mo.lag`` — period shift with Excel cell-back-references."""

import pytest
from openpyxl import load_workbook

import modeleon as mo


class TestLagValues:
    def test_basic_lag_shifts_and_fills(self):
        x = mo.Variable([10.0, 20.0, 30.0, 40.0])
        result = mo.lag(x)
        assert result._value == [0.0, 10.0, 20.0, 30.0]
        assert result.var_type == "list"

    def test_lag_two_periods(self):
        x = mo.Variable([1.0, 2.0, 3.0, 4.0])
        assert mo.lag(x, 2)._value == [0.0, 0.0, 1.0, 2.0]

    def test_lead_negative_periods(self):
        x = mo.Variable([1.0, 2.0, 3.0])
        assert mo.lag(x, -1)._value == [2.0, 3.0, 0.0]

    def test_numeric_fill(self):
        x = mo.Variable([5.0, 6.0])
        assert mo.lag(x, fill=99.0)._value == [99.0, 5.0]

    def test_variable_fill_uses_its_value_and_dependency(self):
        opening = mo.Variable(500.0)
        x = mo.Variable([10.0, 20.0, 30.0])
        result = mo.lag(x, fill=opening)
        assert result._value == [500.0, 10.0, 20.0]
        refs = list(result._expr.iter_refs())
        assert any(r is opening for r in refs)
        assert any(r is x for r in refs)

    def test_delta_idiom(self):
        nwc = mo.Variable([100.0, 130.0, 120.0])
        delta = nwc - mo.lag(nwc)
        assert delta._value == [100.0, 30.0, -10.0]

    def test_opening_balance_idiom(self):
        closing = mo.Variable([90.0, 80.0, 70.0])
        opening = mo.lag(closing, fill=100.0)
        assert opening._value == [100.0, 90.0, 80.0]

    def test_scalar_source_rejected(self):
        with pytest.raises(TypeError, match="list-valued"):
            mo.lag(mo.Variable(5.0))

    def test_list_fill_rejected(self):
        x = mo.Variable([1.0, 2.0])
        with pytest.raises(TypeError, match="SCALAR"):
            mo.lag(x, fill=mo.Variable([1.0, 2.0]))

    def test_lag_longer_than_series(self):
        x = mo.Variable([1.0, 2.0])
        assert mo.lag(x, 5)._value == [0.0, 0.0]


class TestLagExcel:
    """The emitted workbook: each lag cell references the source row's
    cell one column back; the exposed head cell holds the fill — a
    literal, or a REFERENCE to the fill Variable's own cell."""

    def _formulas(self, path, sheet):
        ws = load_workbook(path)[sheet]
        return {
            cell.coordinate: cell.value
            for row in ws.iter_rows()
            for cell in row
            if cell.value is not None
        }

    def test_lag_emits_back_references(self, tmp_path):
        model = mo.MultiVariable("m")
        model.s = mo.MultiVariable("S", excel_props={"tab": True})
        with model.s as s:
            s.flow = mo.Variable([10.0, 20.0, 30.0], display_name="Flow")
            s.prior = mo.lag(s.flow, display_name="Prior")

        path = tmp_path / "lag.xlsx"
        model.to_excel(path)
        cells = self._formulas(path, "S")

        flow_row = next(
            coord for coord, v in cells.items() if v == "Flow"
        )[1:]
        prior_row = next(
            coord for coord, v in cells.items() if v == "Prior"
        )[1:]
        # Head cell: fill literal (formula-path "=0"). Later cells:
        # previous column of the source row.
        assert cells[f"B{prior_row}"] == "=0"
        assert cells[f"C{prior_row}"] == f"=B{flow_row}"
        assert cells[f"D{prior_row}"] == f"=C{flow_row}"

    def test_variable_fill_emits_cell_reference(self, tmp_path):
        model = mo.MultiVariable("m")
        model.s = mo.MultiVariable("S", excel_props={"tab": True})
        with model.s as s:
            s.opening = mo.Variable(500.0, display_name="Opening")
            s.closing = mo.Variable([90.0, 80.0], display_name="Closing")
            s.prior = mo.lag(s.closing, fill=s.opening, display_name="Prior")

        path = tmp_path / "lag_fill.xlsx"
        model.to_excel(path)
        cells = self._formulas(path, "S")

        opening_row = next(
            coord for coord, v in cells.items() if v == "Opening"
        )[1:]
        prior_row = next(
            coord for coord, v in cells.items() if v == "Prior"
        )[1:]
        # The head cell REFERENCES the opening cell — the dependency
        # stays live in the workbook, not a frozen literal.
        assert cells[f"B{prior_row}"] == f"=B{opening_row}"
