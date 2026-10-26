# SPDX-License-Identifier: Apache-2.0
"""The workbook writes ``DATEDIF``, ``WEEKDAY``, ``WEEKNUM``,
``NETWORKDAYS`` and ``WORKDAY`` as the live calls they are, and
recalculates to the engine's values.

A scalar call reads its cells (``=DATEDIF(B2, B3, "M")``); a call over
a series is one call per period, each reading this period's cell
(``=WEEKDAY(C2, 2)``); the holidays of ``NETWORKDAYS`` / ``WORKDAY`` are
the one operand written as a range, the same in every period
(``=NETWORKDAYS(I!B2, C2, I!C4:E4)``); a ``WORKDAY`` cell shows as a
date. All five names are classic, stored bare (no ``_xlfn.`` prefix).
The slow test lets LibreOffice recalculate a workbook that holds every
function under test and compares; the three documented differences
between Excel and LibreOffice — ``DATEDIF``'s ``"YD"`` across a 29
February, ``WEEKNUM`` in the last days of a year, and the error a
``DATEDIF`` with the start after the end shows — are each pinned as an
expected failure.
"""

from __future__ import annotations

import shutil
import subprocess
import warnings
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

import modeleon as mo
from modeleon.compile.excel.renderer import CrossScopeReferenceWarning
from modeleon.functions.dates import DATEDIF, NETWORKDAYS, WEEKDAY, WEEKNUM, WORKDAY


def _cells(ws):
    """``{label: [cells right of the label]}`` of one sheet."""
    out = {}
    for row in ws.iter_rows():
        label = row[0].value
        if label:
            out[str(label)] = [c for c in row[1:] if c.value is not None]
    return out


def _values(ws):
    return {label: [c.value for c in cells] for label, cells in _cells(ws).items()}


def _model():
    m = mo.Model("m", default_grain="month", default_start="2024-01", default_periods=3)
    with m:
        m.i = mo.MultiVariable("I")
        m.i.start = mo.Variable(date(2024, 1, 15), display_name="start")
        m.i.end = mo.Variable(date(2024, 3, 20), display_name="end")
        m.i.hol = mo.Variable(
            [date(2024, 1, 1), date(2024, 2, 5), date(2024, 3, 15)], display_name="hol",
        )
        m.x = mo.MultiVariable("X")
        m.x.ends = mo.Variable(
            [date(2024, 1, 31), date(2024, 2, 29), date(2024, 3, 31)], display_name="ends",
        )
    return m


def _book(model, tmp_path):
    path = tmp_path / "book.xlsx"
    model.to_excel(str(path))
    return load_workbook(str(path))


class TestTheWrittenFormulas:
    def test_scalar_calls_read_their_cells(self, tmp_path):
        m = _model()
        with m:
            m.i.months = DATEDIF(m.i.start, m.i.end, "M").set_display_name("months")
            m.i.wd = WEEKDAY(m.i.start).set_display_name("wd")
            m.i.wd2 = WEEKDAY(m.i.start, 2).set_display_name("wd2")
            m.i.wk = WEEKNUM(m.i.end, 21).set_display_name("wk")
            m.i.n = NETWORKDAYS(m.i.start, m.i.end, m.i.hol).set_display_name("n")
            m.i.n0 = NETWORKDAYS(m.i.start, m.i.end).set_display_name("n0")
            m.i.n1 = NETWORKDAYS(m.i.start, m.i.end, m.i.start).set_display_name("n1")
            m.i.w = WORKDAY(m.i.start, 10, m.i.hol).set_display_name("w")
            m.i.w1 = WORKDAY(m.i.start, 10, date(2024, 1, 16)).set_display_name("w1")
        i = _values(_book(m, tmp_path)["I"])
        assert i["start"] == [datetime(2024, 1, 15)] and i["end"] == [datetime(2024, 3, 20)]
        assert i["hol"] == [datetime(2024, 1, 1), datetime(2024, 2, 5), datetime(2024, 3, 15)]
        assert i["months"] == ['=DATEDIF(B2, B3, "M")']
        assert i["wd"] == ["=WEEKDAY(B2)"]
        assert i["wd2"] == ["=WEEKDAY(B2, 2)"]
        assert i["wk"] == ["=WEEKNUM(B3, 21)"]
        assert i["n"] == ["=NETWORKDAYS(B2, B3, C4:E4)"]
        assert i["n0"] == ["=NETWORKDAYS(B2, B3)"]
        assert i["n1"] == ["=NETWORKDAYS(B2, B3, B2)"]
        assert i["w"] == ["=WORKDAY(B2, 10, C4:E4)"]
        assert i["w1"] == ["=WORKDAY(B2, 10, DATE(2024, 1, 16))"]

    def test_a_call_over_a_series_is_one_call_per_period(self, tmp_path):
        m = _model()
        with m:
            m.x.d = DATEDIF(m.i.start, m.x.ends, "D").set_display_name("d")
            m.x.wd = WEEKDAY(m.x.ends, 2).set_display_name("wd")
            m.x.wk = WEEKNUM(m.x.ends).set_display_name("wk")
            m.x.n = NETWORKDAYS(m.i.start, m.x.ends, m.i.hol).set_display_name("n")
            m.x.w = WORKDAY(m.x.ends, -5, m.i.hol).set_display_name("w")
        x = _values(_book(m, tmp_path)["X"])
        assert x["d"] == ['=DATEDIF(I!B2, B2, "D")', '=DATEDIF(I!B2, C2, "D")',
                          '=DATEDIF(I!B2, D2, "D")']
        assert x["wd"] == ["=WEEKDAY(B2, 2)", "=WEEKDAY(C2, 2)", "=WEEKDAY(D2, 2)"]
        assert x["wk"] == ["=WEEKNUM(B2)", "=WEEKNUM(C2)", "=WEEKNUM(D2)"]
        assert x["n"] == ["=NETWORKDAYS(I!B2, B2, I!C4:E4)", "=NETWORKDAYS(I!B2, C2, I!C4:E4)",
                          "=NETWORKDAYS(I!B2, D2, I!C4:E4)"]
        assert x["w"] == ["=WORKDAY(B2, -5, I!C4:E4)", "=WORKDAY(C2, -5, I!C4:E4)",
                          "=WORKDAY(D2, -5, I!C4:E4)"]

    def test_a_workday_cell_shows_as_a_date_and_a_count_as_a_number(self, tmp_path):
        m = _model()
        with m:
            m.i.w = WORKDAY(m.i.start, 10, m.i.hol).set_display_name("w")
            m.i.n = NETWORKDAYS(m.i.start, m.i.end, m.i.hol).set_display_name("n")
            m.x.w = WORKDAY(m.x.ends, -5).set_display_name("w")
            m.x.d = DATEDIF(m.i.start, m.x.ends, "D").set_display_name("d")
        book = _book(m, tmp_path)
        i, x = _cells(book["I"]), _cells(book["X"])
        assert [c.number_format for c in i["w"]] == ["yyyy-mm-dd"]
        assert [c.number_format for c in x["w"]] == ["yyyy-mm-dd"] * 3
        assert "yy" not in i["n"][0].number_format
        assert all("yy" not in c.number_format for c in x["d"])

    def test_the_names_are_stored_bare(self, tmp_path):
        m = _model()
        with m:
            m.i.a = DATEDIF(m.i.start, m.i.end, "Y").set_display_name("a")
            m.i.b = WEEKDAY(m.i.start).set_display_name("b")
            m.i.c = WEEKNUM(m.i.start).set_display_name("c")
            m.i.d = NETWORKDAYS(m.i.start, m.i.end).set_display_name("d")
            m.i.e = WORKDAY(m.i.start, 1).set_display_name("e")
        i = _values(_book(m, tmp_path)["I"])
        for label in "abcde":
            (formula,) = i[label]
            assert formula.startswith("=") and "_xlfn." not in formula, formula

    def test_a_shorter_holidays_line_is_written_as_its_own_cells(self, tmp_path):
        m = _model()
        with m:
            m.i.hol2 = mo.Variable([date(2024, 1, 1), date(2024, 2, 5)], extend=mo.none(),
                                   display_name="hol2")
            m.i.n2 = NETWORKDAYS(m.i.start, m.i.end, m.i.hol2).set_display_name("n2")
            m.x.n2 = NETWORKDAYS(m.i.start, m.x.ends, m.i.hol2).set_display_name("n2")
        book = _book(m, tmp_path)
        i, x = _values(book["I"]), _values(book["X"])
        assert i["hol2"] == [datetime(2024, 1, 1), datetime(2024, 2, 5)]
        assert i["n2"] == ["=NETWORKDAYS(B2, B3, C5:D5)"]
        assert x["n2"] == ["=NETWORKDAYS(I!B2, B2, I!C5:D5)", "=NETWORKDAYS(I!B2, C2, I!C5:D5)",
                           "=NETWORKDAYS(I!B2, D2, I!C5:D5)"]

    def test_a_plain_list_of_holidays_has_no_cells_so_the_value_is_written(self, tmp_path):
        m = _model()
        with m:
            m.x.n = NETWORKDAYS(m.i.start, m.x.ends, [date(2024, 1, 16)]).set_display_name("n")
        assert m.x.n.value == [12, 33, 54]
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter("always")
            x = _values(_book(m, tmp_path)["X"])
        assert x["n"] == [12, 33, 54]
        assert any(issubclass(w.category, CrossScopeReferenceWarning)
                   and "NETWORKDAYS" in str(w.message) for w in caught)


# ─── LibreOffice recalculates the workbook ────────────────────────────


def _soffice():
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


def _recalculated(model, tmp_path):
    """Write ``model``, let LibreOffice calculate every formula, and
    return the recalculated workbook's stored values."""
    soffice = _soffice()
    if soffice is None:
        pytest.skip("LibreOffice is not installed")
    src = tmp_path / "book.xlsx"
    model.to_excel(str(src))
    outdir = tmp_path / "calc"
    profile = (tmp_path / "lo-profile").as_uri()   # private: parallel runs don't clash
    subprocess.run(
        [soffice, f"-env:UserInstallation={profile}", "--headless",
         "--convert-to", "xlsx", "--outdir", str(outdir), str(src)],
        check=True, capture_output=True, timeout=180,
    )
    return load_workbook(str(outdir / "book.xlsx"), data_only=True)


def _as_date(value):
    return value.date() if isinstance(value, datetime) else value


def _parity_model():
    m = _model()
    with m:
        m.i.s = mo.Variable(date(2001, 6, 1), display_name="s")
        m.i.e = mo.Variable(date(2002, 8, 15), display_name="e")
        for unit in ("Y", "M", "D", "MD", "YM", "YD"):
            setattr(m.i, f"dif_{unit.lower()}",
                    DATEDIF(m.i.s, m.i.e, unit).set_display_name(f"dif_{unit}"))
        m.i.jan31 = mo.Variable(date(2016, 1, 31), display_name="jan31")
        m.i.mar1 = mo.Variable(date(2016, 3, 1), display_name="mar1")
        m.i.md_neg = DATEDIF(m.i.jan31, m.i.mar1, "MD").set_display_name("md_neg")
        m.i.months = DATEDIF(m.i.start, m.i.end, "M").set_display_name("months")
        m.i.wd = WEEKDAY(m.i.start).set_display_name("wd")
        m.i.wd3 = WEEKDAY(m.i.start, 3).set_display_name("wd3")
        m.i.wd16 = WEEKDAY(m.i.end, 16).set_display_name("wd16")
        m.i.wk = WEEKNUM(m.i.end, 21).set_display_name("wk")
        m.i.wk2 = WEEKNUM(m.i.end, 2).set_display_name("wk2")
        m.i.n = NETWORKDAYS(m.i.start, m.i.end, m.i.hol).set_display_name("n")
        m.i.n0 = NETWORKDAYS(m.i.start, m.i.end).set_display_name("n0")
        m.i.nrev = NETWORKDAYS(m.i.end, m.i.start, m.i.hol).set_display_name("nrev")
        m.i.hol2 = mo.Variable([date(2024, 1, 1), date(2024, 2, 5)], extend=mo.none(),
                               display_name="hol2")
        m.i.n2 = NETWORKDAYS(m.i.start, m.i.end, m.i.hol2).set_display_name("n2")
        m.i.w = WORKDAY(m.i.start, 10, m.i.hol).set_display_name("w")
        m.i.w0 = WORKDAY(m.i.start, 0).set_display_name("w0")
        m.i.wback = WORKDAY(m.i.end, -12, m.i.hol).set_display_name("wback")
        m.x.d = DATEDIF(m.i.start, m.x.ends, "D").set_display_name("d")
        # No whole year from s2 to the ends: Excel's and LibreOffice's "YD"
        # count the same window; the one that straddles a 29 February is
        # the expected failure below.
        m.i.s2 = mo.Variable(date(2023, 6, 1), display_name="s2")
        m.x.yd = DATEDIF(m.i.s2, m.x.ends, "YD").set_display_name("yd")
        m.x.wd = WEEKDAY(m.x.ends, 2).set_display_name("wd")
        m.x.wd14 = WEEKDAY(m.x.ends, 14).set_display_name("wd14")
        m.x.wk = WEEKNUM(m.x.ends).set_display_name("wk")
        m.x.wk11 = WEEKNUM(m.x.ends, 11).set_display_name("wk11")
        m.x.wk17 = WEEKNUM(m.x.ends, 17).set_display_name("wk17")
        m.x.n = NETWORKDAYS(m.i.start, m.x.ends, m.i.hol).set_display_name("n")
        m.x.w = WORKDAY(m.x.ends, -5, m.i.hol).set_display_name("w")
    return m


@pytest.mark.slow
def test_the_workbook_recalculates_to_the_python_values(tmp_path):
    m = _parity_model()
    # The engine's own numbers first, so a parity failure is LibreOffice's.
    assert [getattr(m.i, f"dif_{u}").value for u in ("y", "m", "d", "md", "ym", "yd")] == \
        [1, 14, 440, 14, 2, 75]
    assert m.i.md_neg.value == -1
    assert (m.i.n.value, m.i.n0.value, m.i.nrev.value, m.i.n2.value) == (46, 48, -46, 47)
    assert (m.i.w.value, m.i.w0.value, m.i.wback.value) == \
        (date(2024, 1, 29), date(2024, 1, 15), date(2024, 3, 1))
    assert m.x.n.value == [13, 33, 53]
    assert m.x.yd.value == [244, 273, 304]
    assert (m.x.wd14.value, m.x.wk11.value, m.x.wk17.value) == ([7, 1, 4], [5, 9, 13], [5, 9, 14])
    book = _recalculated(m, tmp_path)
    i, x = _values(book["I"]), _values(book["X"])
    for label in ("dif_Y", "dif_M", "dif_D", "dif_MD", "dif_YM", "dif_YD", "md_neg", "months",
                  "wd", "wd3", "wd16", "wk", "wk2", "n", "n0", "nrev", "n2"):
        assert i[label] == [getattr(m.i, label.lower()).value], label
    for label in ("w", "w0", "wback"):
        assert [_as_date(v) for v in i[label]] == [getattr(m.i, label).value], label
    for label in ("d", "yd", "wd", "wd14", "wk", "wk11", "wk17", "n"):
        assert x[label] == getattr(m.x, label).value, label
    assert [_as_date(v) for v in x["w"]] == m.x.w.value


@pytest.mark.slow
@pytest.mark.xfail(
    strict=True,
    reason="LibreOffice's DATEDIF \"YD\" places the start's month and day in the end's "
           "year and counts 324 days here; Excel places the end's month and day in the "
           "start's year and returns 323, which the engine computes (the difference is "
           "noted in LibreOffice's own source, sc/source/core/tool/interpr2.cxx).",
)
def test_datedif_yd_across_a_leap_day_is_excel_s_not_libreoffice_s(tmp_path):
    m = mo.Model("m")
    m.s = mo.Variable(date(1988, 6, 22), display_name="s")
    m.e = mo.Variable(date(2012, 5, 11), display_name="e")
    m.yd = DATEDIF(m.s, m.e, "YD").set_display_name("yd")
    assert m.yd.value == 323
    ws = _recalculated(m, tmp_path).active
    assert _values(ws)["yd"] == [323]


@pytest.mark.slow
@pytest.mark.xfail(
    strict=True,
    reason="LibreOffice's WEEKNUM in every mode but 21 (tools::Date::GetWeekOfYear with one "
           "minimum day in a week) numbers a date in the last partial week of a year, the one "
           "that runs into the next year, week 1 of that next year, and turns a week 54 into 1; "
           "Excel numbers the week holding 1 January 1 and lets the last days of a year reach 53 "
           "or 54 (31 December 2020 is 53, 31 December 2000 is 54), which the engine computes.",
)
def test_weeknum_s_last_days_of_a_year_are_excel_s_not_libreoffice_s(tmp_path):
    m = mo.Model("m")
    m.d = mo.Variable(date(2020, 12, 31), display_name="d")
    m.e = mo.Variable(date(2000, 12, 31), display_name="e")
    m.wk = WEEKNUM(m.d).set_display_name("wk")
    m.wk2 = WEEKNUM(m.d, 2).set_display_name("wk2")
    m.wk54 = WEEKNUM(m.e).set_display_name("wk54")
    assert (m.wk.value, m.wk2.value, m.wk54.value) == (53, 53, 54)
    ws = _recalculated(m, tmp_path).active
    values = _values(ws)
    assert (values["wk"], values["wk2"], values["wk54"]) == ([53], [53], [54])


@pytest.mark.slow
@pytest.mark.xfail(
    strict=True,
    reason="A DATEDIF with the start after the end is #NUM! in Excel, which the engine "
           "computes; LibreOffice recalculates it to Err:502 and saves the cell as #VALUE!.",
)
def test_datedif_with_the_start_after_the_end_is_num_in_excel_not_value(tmp_path):
    m = mo.Model("m")
    m.s = mo.Variable(date(2002, 8, 15), display_name="s")
    m.e = mo.Variable(date(2001, 6, 1), display_name="e")
    m.d = DATEDIF(m.s, m.e, "D").set_display_name("d")
    assert m.d.value == "#NUM!"
    ws = _recalculated(m, tmp_path).active
    assert _values(ws)["d"] == ["#NUM!"]
