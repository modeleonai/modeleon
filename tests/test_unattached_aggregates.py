# SPDX-License-Identifier: Apache-2.0
"""An aggregate over a list the workbook never lays out.

``x = mo.Variable([10, 20, 30])`` that is never attached to the model has
no cells, so ``mo.SUM(x) + a`` has no range for the ``SUM`` to read.
Rendering it period by period would sum ONE element per column —
``=SUM(10) + B1``, ``=SUM(20) + C1`` — and the workbook would show
11, 22, 33 where the model says 61, 62, 63.

The honest rendering writes the aggregate's computed value (one number,
the same in every period) in place of the call, keeps the rest of the
formula live, and says so with a :class:`~modeleon.CrossScopeReferenceWarning`.
An attached list keeps its live range formula.
"""

from __future__ import annotations

import warnings

import pytest

import modeleon as mo


def _soffice() -> str | None:
    """Path to the LibreOffice command line, or None when not installed."""
    import shutil
    from pathlib import Path

    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


def _recalculated(xlsx, workdir) -> dict:
    """Let LibreOffice calculate every formula of the workbook and return
    the first sheet as ``{label in column A: [values of the other cells]}``."""
    import subprocess

    import openpyxl

    soffice = _soffice()
    if soffice is None:
        pytest.skip("LibreOffice is not installed")
    out_dir = workdir / "recalculated"
    out_dir.mkdir()
    profile = (workdir / "lo-profile").as_uri()  # private profile: parallel runs don't clash
    subprocess.run(
        [soffice, f"-env:UserInstallation={profile}", "--headless",
         "--convert-to", "xlsx", "--outdir", str(out_dir), str(xlsx)],
        check=True, capture_output=True, timeout=120,
    )
    ws = openpyxl.load_workbook(out_dir / xlsx.name, data_only=True).active
    return {row[0].value: [c.value for c in row[1:]] for row in ws.iter_rows()}


def _written(xlsx) -> dict:
    """The first sheet as written: ``{label: [cells after the label]}``."""
    import openpyxl

    ws = openpyxl.load_workbook(xlsx).active
    return {row[0].value: row[1:] for row in ws.iter_rows()}


def _emit(m: mo.Model, path):
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        m.to_excel(str(path))
    return [w for w in caught
            if issubclass(w.category, mo.CrossScopeReferenceWarning)]


# (name, build the aggregate from the unattached list, the list's values)
_AGGREGATES = [
    ("SUM", lambda x: mo.SUM(x), [10, 20, 30]),
    ("AVERAGE", lambda x: mo.AVERAGE(x), [10, 20, 30]),
    ("MIN", lambda x: mo.MIN(x), [10, 20, 30]),
    ("MAX", lambda x: mo.MAX(x), [10, 20, 30]),
    ("NPV", lambda x: mo.NPV(0.1, x), [10, 20, 30]),
    ("IRR", lambda x: mo.IRR(x), [-100, 60, 70]),
]


def _model(build, values, *, attached: bool):
    m = mo.Model("M")
    m.a = mo.Variable([1, 2, 3], display_name="A")
    x = mo.Variable(values, display_name="X")
    if attached:
        m.x = x
    m.u = build(x) + m.a
    m.u._display_name = "U"
    return m, x


class TestUnattachedListInsideAFormula:
    @pytest.mark.parametrize("name, build, values", _AGGREGATES,
                             ids=[a[0] for a in _AGGREGATES])
    def test_the_workbook_computes_the_python_values(self, tmp_path, name,
                                                     build, values):
        m, _x = _model(build, values, attached=False)
        agg = build(mo.Variable(values))._value
        assert m.u._value == pytest.approx([agg + a for a in (1, 2, 3)])

        out = tmp_path / "u.xlsx"
        _emit(m, out)
        calc = _recalculated(out, tmp_path)
        assert calc["U"][:3] == pytest.approx(m.u._value, rel=1e-12)

    @pytest.mark.parametrize("name, build, values", _AGGREGATES,
                             ids=[a[0] for a in _AGGREGATES])
    def test_the_aggregate_is_one_number_and_the_rest_stays_live(
            self, tmp_path, name, build, values):
        m, _x = _model(build, values, attached=False)
        out = tmp_path / "u.xlsx"
        warned = _emit(m, out)
        rows = _written(out)
        a_cells = [c.coordinate for c in rows["A"][:3]]
        u_cells = [c.value for c in rows["U"][:3]]
        for formula, a_cell in zip(u_cells, a_cells):
            assert isinstance(formula, str) and formula.startswith("="), formula
            assert a_cell in formula, formula          # `+ a` is still a reference
            assert f"{name}(" not in formula, formula  # no call over one element
        # the aggregate part is the same number in every period
        constants = {f.replace(a, "") for f, a in zip(u_cells, a_cells)}
        assert len(constants) == 1, u_cells
        # and the file says why it is a number, naming the list's function
        assert any(name in str(w.message) for w in warned), [
            str(w.message) for w in warned
        ]


class TestAttachedListKeepsTheRange:
    @pytest.mark.parametrize("name, build, values", _AGGREGATES,
                             ids=[a[0] for a in _AGGREGATES])
    def test_live_range_formula(self, tmp_path, name, build, values):
        m, _x = _model(build, values, attached=True)
        out = tmp_path / "u.xlsx"
        warned = _emit(m, out)
        assert warned == []
        rows = _written(out)
        x_row = rows["X"]
        x_range = f"{x_row[0].coordinate}:{x_row[2].coordinate}"
        for formula, a in zip([c.value for c in rows["U"][:3]], rows["A"][:3]):
            assert f"{name}(" in formula and x_range in formula, formula
            assert a.coordinate in formula, formula

        calc = _recalculated(out, tmp_path)
        assert calc["U"][:3] == pytest.approx(m.u._value, rel=1e-6)

    def test_a_window_of_an_attached_row_keeps_its_range(self, tmp_path):
        m = mo.Model("M")
        m.a = mo.Variable([1, 2, 3], display_name="A")
        m.u = mo.SUM(m.a[0:2]) + m.a
        m.u._display_name = "U"
        out = tmp_path / "u.xlsx"
        assert _emit(m, out) == []
        rows = _written(out)
        window = f"{rows['A'][0].coordinate}:{rows['A'][1].coordinate}"
        for formula, a in zip([c.value for c in rows["U"][:3]], rows["A"][:3]):
            assert formula == f"=SUM({window}) + {a.coordinate}"


class TestOtherListsWithNoCells:
    """Every list an aggregate can reduce that has no row in the book —
    an intermediate expression, a literal list, a slice of an unattached
    list, one list among several arguments — follows the same rule."""

    @pytest.mark.parametrize("build", [
        lambda m, x: mo.SUM(m.a * 2) + m.a,
        lambda m, x: mo.SUM([10, 20, 30]) + m.a,
        lambda m, x: mo.SUM(x[0:2]) + m.a,
        lambda m, x: mo.SUM(x, m.a) + m.a,
    ], ids=["intermediate", "literal", "slice", "mixed"])
    def test_the_workbook_computes_the_python_values(self, tmp_path, build):
        m = mo.Model("M")
        m.a = mo.Variable([1, 2, 3], display_name="A")
        x = mo.Variable([10, 20, 30])
        m.u = build(m, x)
        m.u._display_name = "U"
        out = tmp_path / "u.xlsx"
        _emit(m, out)
        calc = _recalculated(out, tmp_path)
        assert calc["U"][:3] == pytest.approx(m.u._value, rel=1e-12)

    def test_a_standalone_aggregate_is_its_value(self, tmp_path):
        m = mo.Model("M")
        m.s = mo.SUM(mo.Variable([10, 20, 30]))
        m.s._display_name = "S"
        m.lit = mo.SUM([10, 20, 30])
        m.lit._display_name = "Literal"
        out = tmp_path / "s.xlsx"
        _emit(m, out)
        rows = _written(out)
        assert rows["S"][0].value == 60
        assert rows["Literal"][0].value == 60


class TestScalarOperandsAreUnchanged:
    def test_an_unattached_scalar_is_inlined_inside_the_live_call(self, tmp_path):
        m = mo.Model("M")
        m.a = mo.Variable([1, 2, 3], display_name="A")
        s = mo.Variable(5)
        m.u = mo.SUM(s, m.a) + m.a
        m.u._display_name = "U"
        out = tmp_path / "u.xlsx"
        warned = _emit(m, out)
        assert warned
        rows = _written(out)
        a_range = f"{rows['A'][0].coordinate}:{rows['A'][2].coordinate}"
        for formula, a in zip([c.value for c in rows["U"][:3]], rows["A"][:3]):
            assert formula == f"=SUM(5, {a_range}) + {a.coordinate}"


class TestTrackedLines:
    def test_a_track_coordinate_reduces_the_whole_list(self, tmp_path):
        # In the blend view each tracked line is one row of its live
        # values; the SUM inside the plan coordinate is still one number.
        from modeleon.compile.excel.view import ExcelView

        m = mo.Model("m", display_name="M",
                     tracks=mo.Tracks("plan", "actual",
                                      blend=mo.blend(given="actual", follow="plan",
                                                     until="2026-02")),
                     default_grain="month", default_start="2026-01",
                     default_periods=4)
        m.default_excel_view = ExcelView(tracks="blend")
        m.s = mo.MultiVariable("S", excel_props={"tab": True})
        x = mo.Variable([10.0, 20.0, 30.0, 40.0])
        with m.s as s:
            s.d = mo.Variable(plan=[1.0, 2.0, 3.0, 4.0],
                              actual=[10.0, 20.0, 30.0, 40.0], display_name="D")
            s.t = mo.Variable(plan=mo.SUM(x), display_name="T")
            s.u = mo.Variable(plan=mo.SUM(x) + s.d, display_name="U")
        out = tmp_path / "t.xlsx"
        _emit(m, out)
        calc = _recalculated(out, tmp_path)
        assert calc["T"][:4] == pytest.approx(s.t._value["live"])
        assert calc["U"][:4] == pytest.approx(s.u._value["live"])
