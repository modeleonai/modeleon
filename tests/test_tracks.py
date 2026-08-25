# SPDX-License-Identifier: Apache-2.0
"""The track value layer (§14.2.2, §16, §17) — P1 core.

A track-axised Variable stores role → time-series tracks. One formula
broadcasts across coordinates; the mismatch law is loud; time-coupled
functions lift per track; slices are first-class AST nodes.
"""

from __future__ import annotations

import re

import pytest

import modeleon as mo
from modeleon.core.expr import Restrict
from modeleon.core.tracks import TrackValues


def _model(periods: int = 4) -> "mo.Model":
    return mo.Model('m', default_grain='month', default_start='2025-01',
                    default_periods=periods)


class TestConstruction:
    def test_tracked_variable(self):
        m = _model()
        with m:
            m.x = mo.Variable(tracks={'план': [1.0] * 4, 'факт': [2.0] * 4})
        assert isinstance(m.x._value, TrackValues)
        assert set(m.x._value.roles) == {'план', 'факт'}
        assert m.x.var_type == 'list'

    def test_time_resolves_from_ambient(self):
        m = _model()
        with m:
            m.x = mo.Variable(tracks={'план': [1.0] * 4})
        assert m.x.time is not None and m.x.time.grain == 'month'

    def test_extent_law_applies_per_track(self):
        m = _model(24)
        with pytest.raises(ValueError, match="24-period window"):
            with m:
                m.x = mo.Variable(tracks={'план': [1.0, 2.0]})

    def test_ragged_tracks_rejected(self):
        with pytest.raises(ValueError, match="one time length"):
            TrackValues({'план': [1.0, 2.0], 'факт': [1.0]})

    def test_nested_tracks_rejected(self):
        with pytest.raises(TypeError, match="second axis"):
            TrackValues({'план': {'вложенный': [1.0]}})

    def test_exclusive_with_value_and_indexed_by(self):
        with pytest.raises(ValueError, match="exclusive"):
            mo.Variable([1.0], tracks={'план': [1.0]})
        оси = mo.Variable(['a', 'b'])
        with pytest.raises(ValueError, match="axis budget"):
            mo.Variable(tracks={'план': [1.0, 2.0]}, indexed_by=[оси])


class TestBroadcast:
    def _rev(self, m):
        with m:
            m.выручка = mo.Variable(tracks={
                'план': [100.0, 110.0, 120.0, 130.0],
                'факт': [95.0, 118.0, 0.0, 0.0],
            })
        return m.выручка

    def test_tracked_minus_plain_broadcasts_into_every_role(self):
        m = _model()
        rev = self._rev(m)
        with m:
            m.косты = mo.Variable([50.0] * 4)
            m.маржа = rev - m.косты
        v = m.маржа._value
        assert isinstance(v, TrackValues)
        assert v['план'] == [50.0, 60.0, 70.0, 80.0]
        assert v['факт'] == [45.0, 68.0, -50.0, -50.0]

    def test_tracked_times_scalar(self):
        m = _model()
        rev = self._rev(m)
        with m:
            m.удвоено = rev * 2
        assert m.удвоено._value['план'] == [200.0, 220.0, 240.0, 260.0]

    def test_same_roles_zip_per_track(self):
        m = _model()
        with m:
            m.a = mo.Variable(tracks={'план': [1.0] * 4, 'факт': [2.0] * 4})
            m.b = mo.Variable(tracks={'план': [10.0] * 4, 'факт': [20.0] * 4})
            m.c = m.a + m.b
        assert m.c._value['план'] == [11.0] * 4
        assert m.c._value['факт'] == [22.0] * 4

    def test_mismatched_role_sets_die_loudly(self):
        m = _model()
        with m:
            m.a = mo.Variable(tracks={'план': [1.0] * 4, 'факт': [2.0] * 4})
            m.b = mo.Variable(tracks={'план': [1.0] * 4, 'база': [9.0] * 4})
        with pytest.raises(ValueError, match="never silently intersects"):
            _ = m.a + m.b

    def test_comparison_lifts(self):
        m = _model()
        rev = self._rev(m)
        флаг = rev > 100
        assert флаг._value['план'] == [False, True, True, True]


class TestSlice:
    def test_at_basis_returns_the_track(self):
        m = _model()
        with m:
            m.x = mo.Variable(tracks={'план': [1.0, 2.0, 3.0, 4.0],
                                      'факт': [9.0] * 4})
        срез = m.x.at(track='план')
        assert срез._value == [1.0, 2.0, 3.0, 4.0]
        assert isinstance(срез._expr, Restrict)
        assert срез._expr.label == 'план'

    def test_variance_is_one_formula(self):
        m = _model()
        with m:
            m.x = mo.Variable(tracks={'план': [100.0] * 4,
                                      'факт': [95.0, 118.0, 0.0, 0.0]})
            m.отклонение = m.x.at(track='факт') - m.x.at(track='план')
        assert m.отклонение._value == [-5.0, 18.0, -100.0, -100.0]

    def test_unknown_role_dies_with_the_declared_list(self):
        x = mo.Variable(tracks={'план': [1.0, 2.0]})
        with pytest.raises(ValueError, match="план"):
            x.at(track='факт')

    def test_at_basis_on_plain_variable_dies(self):
        with pytest.raises(ValueError, match="no track coordinates"):
            mo.Variable([1.0, 2.0]).at(track='план')

    def test_at_needs_something(self):
        with pytest.raises(TypeError, match="grain"):
            mo.Variable(tracks={'план': [1.0]}).at()


class TestLifting:
    def _rev(self, m):
        with m:
            m.выручка = mo.Variable(tracks={
                'план': [100.0, 110.0, 120.0, 130.0],
                'факт': [95.0, 118.0, 0.0, 0.0],
            })
        return m.выручка

    def test_lag_shifts_each_track(self):
        m = _model()
        rev = self._rev(m)
        лаг = mo.lag(rev)
        assert лаг._value['план'] == [0.0, 100.0, 110.0, 120.0]
        assert лаг._value['факт'] == [0.0, 95.0, 118.0, 0.0]

    def test_cumsum_accumulates_each_track(self):
        m = _model()
        rev = self._rev(m)
        acc = mo.cumsum(rev)
        assert acc._value['план'] == [100.0, 210.0, 330.0, 460.0]
        assert acc._value['факт'] == [95.0, 213.0, 213.0, 213.0]

    def test_sum_reduces_each_track(self):
        m = _model()
        rev = self._rev(m)
        итог = mo.SUM(rev)
        assert итог._value['план'] == 460.0
        assert итог._value['факт'] == 213.0

    def test_if_lifts_across_roles(self):
        m = _model()
        rev = self._rev(m)
        флаг = mo.IF(rev > 100, 1.0, 0.0)
        assert флаг._value['план'] == [0.0, 1.0, 1.0, 1.0]
        assert флаг._value['факт'] == [0.0, 1.0, 0.0, 0.0]

    def test_recurrence_with_tracked_driver_lifts(self):
        # Был teaching-отказ; теперь rank-lifting — см. TestRecurrenceLift.
        m = _model()
        rev = self._rev(m)
        r = mo.recurrence(start=0.0, formula="{prev} + {r}",
                          variables={"r": rev}, periods=4)
        assert r._value['план'] == [0.0, 110.0, 230.0, 360.0]
        assert r._value['факт'] == [0.0, 118.0, 118.0, 118.0]


class TestGates:
    def _tracked(self):
        m = _model()
        with m:
            m.x = mo.Variable(tracks={'план': [1.0] * 4, 'факт': [2.0] * 4},
                              display_name='X')
        return m.x

    def test_whole_cube_grain_projects_per_track(self):
        """P2.5: at(grain) on a tracked variable re-grains EACH track by
        the line's rule — the axis passes through (was a refusal)."""
        m = _model(12)
        with m:
            m.x = mo.Variable(
                tracks={'план': [1.0] * 12, 'факт': [2.0] * 12},
                regrain=mo.up('sum'),
            )
        q = m.x.at('quarter')
        assert q._value['план'] == [3.0] * 4
        assert q._value['факт'] == [6.0] * 4

    def test_whole_cube_without_rule_still_teaches(self):
        with pytest.raises(ValueError, match="re-grain rule"):
            self._tracked().at('quarter')

    def test_slice_then_grain_works(self):
        m = _model(12)
        with m:
            m.x = mo.Variable(
                tracks={'план': [1.0] * 12, 'факт': [2.0] * 12},
                regrain=mo.up('sum'),
            )
        q = m.x.at('quarter', track='план')
        assert q._value == [3.0, 3.0, 3.0, 3.0]

    def test_emission_expands_tracks_into_rows(self):
        """P3: the xlsx writer expands a tracked line — the display
        default keeps the name (BLEND-sheet law), other tracks follow
        as labeled value rows (was a refusal)."""
        import os
        import tempfile
        from openpyxl import load_workbook
        m = _model()
        with m:
            m.x = mo.Variable(tracks={'план': [1.0] * 4, 'факт': [2.0] * 4},
                              display_name='Икс')
        p = os.path.join(tempfile.mkdtemp(), 'т.xlsx')
        m.to_excel(p)
        labels = [r[0] for r in load_workbook(p).worksheets[0]
                  .iter_rows(min_row=2, max_row=3, max_col=1,
                             values_only=True)]
        assert labels[0] == 'Икс'          # default row keeps the name
        assert 'факт' in (labels[1] or '')  # the other track, labeled

    def test_model_projection_projects_per_track(self):
        from modeleon.core.projection import project_variable
        m = _model(12)
        with m:
            m.x = mo.Variable(
                tracks={'план': [1.0] * 12, 'факт': [2.0] * 12},
                regrain=mo.up('sum'),
            )
        p = project_variable(m.x, 'quarter', {})
        assert p._value['план'] == [3.0] * 4
        assert p._value['факт'] == [6.0] * 4


class TestRecurrenceLift:
    """Recurrence over tracks = N independent chains (§15.3.5).

    План-остаток катится от план-потоков, факт-остаток от фактических.
    (Перепривязанная live-цепочка — это третья вселенная закона сшивки,
    §16.9 — P2, не эта функция.)
    """

    def test_tracked_driver_rolls_per_track(self):
        m = _model()
        with m:
            m.поток = mo.Variable(tracks={
                'план': [100.0, 100.0, 100.0, 100.0],
                'факт': [95.0, 120.0, 0.0, 0.0],
            })
            m.остаток = mo.recurrence(
                start=1000.0, formula="{prev} + {п}",
                variables={"п": m.поток},
            )
        assert m.остаток._value['план'] == [1000.0, 1100.0, 1200.0, 1300.0]
        assert m.остаток._value['факт'] == [1000.0, 1120.0, 1120.0, 1120.0]

    def test_tracked_start_spawns_chains(self):
        m = _model(3)
        with m:
            m.старт = mo.Variable(tracks={'план': 10.0, 'факт': 20.0})
            m.ряд = mo.recurrence(start=m.старт, formula="{prev} * 2",
                                  periods=3)
        assert m.ряд._value['план'] == [10.0, 20.0, 40.0]
        assert m.ряд._value['факт'] == [20.0, 40.0, 80.0]

    def test_mismatched_role_sets_die(self):
        m = _model()
        with m:
            m.а = mo.Variable(tracks={'план': [1.0] * 4, 'факт': [2.0] * 4})
            m.б = mo.Variable(tracks={'план': [1.0] * 4, 'база': [9.0] * 4})
        with pytest.raises(ValueError, match="differ"):
            mo.recurrence(start=0.0, formula="{prev} + {a} + {b}",
                          variables={"a": m.а, "b": m.б})

    def test_lambda_mode_lifts_too(self):
        m = _model(3)
        with m:
            m.старт = mo.Variable(tracks={'план': 1.0, 'факт': 2.0})
            m.р = mo.recurrence(start=m.старт,
                                formula=lambda prev, t: prev + 1, periods=3)
        assert m.р._value['план'] == [1.0, 2.0, 3.0]
        assert m.р._value['факт'] == [2.0, 3.0, 4.0]


class TestCyrillicPlaceholders:
    def test_cyrillic_template_names_substitute(self):
        # Была ASCII-only регулярка: «{п}» не подставлялся и парсился
        # как set-литерал. Продукт пишет модели по-русски.
        r = mo.recurrence(start=0.0, formula="{prev} + {доход}",
                          variables={"доход": mo.Variable([1.0, 2.0, 3.0])})
        assert r._value == [0.0, 2.0, 5.0]


class TestDateAlignment:
    """The date-aligned zip (§16.5): operands with different windows
    align by CALENDAR; outside its own extent each side follows its
    declared extend rule — never a positional zip, never a guess."""

    def _m(self):
        return mo.Model('m', default_grain='month',
                        default_start='2025-01', default_periods=6)

    def test_flow_extends_with_zeros(self):
        m = self._m()
        with m:
            m.выручка = mo.Variable([100.0] * 6)
            m.бонус = mo.Variable([10.0] * 3, start='2025-04',
                                  grain='month', extend=mo.zero())
            m.итого = m.выручка + m.бонус
        assert m.итого._value == [100.0, 100.0, 100.0, 110.0, 110.0, 110.0]

    def test_rate_holds_forward_and_backward(self):
        m = self._m()
        with m:
            m.база = mo.Variable([100.0] * 6)
            m.ставка = mo.Variable([0.10, 0.12], start='2025-05',
                                   grain='month', extend=mo.hold())
            m.налог = m.база * m.ставка
        # назад — первым значением, вперёд — последним.
        assert m.налог._value == [10.0, 10.0, 10.0, 10.0, 10.0, 12.0]

    def test_result_is_located_at_the_union(self):
        m = self._m()
        with m:
            m.a = mo.Variable([1.0] * 6)
            m.b = mo.Variable([1.0] * 2, start='2025-05', grain='month',
                              extend=mo.zero())
            m.c = m.a + m.b
        assert m.c.time.start == '2025-01'
        assert len(m.c._value) == 6

    def test_undeclared_continuation_teaches(self):
        m = self._m()
        with pytest.raises(ValueError, match="Say what continues it"):
            with m:
                m.a = mo.Variable([1.0] * 6)
                m.b = mo.Variable([5.0, 5.0], start='2025-03',
                                  grain='month')
                m.x = m.a + m.b

    def test_same_window_stays_positional(self):
        # Быстрый путь не трогается — байт-в-байт как раньше.
        m = self._m()
        with m:
            m.a = mo.Variable([1.0] * 6)
            m.b = mo.Variable([2.0] * 6)
            m.c = m.a + m.b
        assert m.c._value == [3.0] * 6

    def test_grain_mismatch_refuses(self):
        a = mo.Variable([1.0] * 6, start='2025-01', grain='month')
        b = mo.Variable([1.0, 1.0], start='2025-01', grain='quarter')
        with pytest.raises(ValueError, match="different grains"):
            _ = a + b


class TestAggregateLifts:
    def test_max_min_average_reduce_per_track(self):
        m = _model()
        with m:
            m.x = mo.Variable(tracks={'план': [1.0, 5.0, 3.0, 2.0],
                                      'факт': [9.0, 0.0, 0.0, 0.0]})
        assert dict(mo.MAX(m.x)._value.items()) == {'план': 5.0, 'факт': 9.0}
        assert dict(mo.MIN(m.x)._value.items()) == {'план': 1.0, 'факт': 0.0}
        assert dict(mo.AVERAGE(m.x)._value.items()) == {
            'план': 2.75, 'факт': 2.25,
        }


# ─── P2: mo.Tracks declaration + role-kwarg authoring (§16.2/§16.10) ──


def _declared(periods: int = 6, **decl_labels) -> "mo.Model":
    labels = decl_labels or {'actual': 'факт', 'plan': 'бюджет'}
    return mo.Model('м', tracks=mo.Tracks(**labels),
                    default_grain='month', default_start='2025-01',
                    default_periods=periods)


class TestTracksDeclaration:
    def test_declaration_on_model(self):
        m = _declared()
        assert m.tracks.names == ('actual', 'plan')
        assert m.tracks.label('plan') == 'бюджет'

    def test_names_are_user_content(self):
        """Any words, any language, any count — the engine attaches no
        semantics to a track name (roles are a later, explicit layer)."""
        d = mo.Tracks('факт', 'бюджет_утв', 'бюджет_корр')
        assert d.names == ('факт', 'бюджет_утв', 'бюджет_корр')
        assert d.label('факт') == 'факт'   # positional: name is the label

    def test_labeled_spelling(self):
        d = mo.Tracks('факт', бюджет='Бюджет 2026')
        assert d.names == ('факт', 'бюджет')
        assert d.label('бюджет') == 'Бюджет 2026'

    def test_single_track_is_legal(self):
        m = mo.Model('м', tracks=mo.Tracks('план'),
                     default_grain='month', default_start='2025-01',
                     default_periods=3)
        with m:
            m.x = mo.Variable(план=[1, 2, 3])
        assert m.x._value['план'] == [1, 2, 3]

    def test_empty_declaration_refused(self):
        with pytest.raises(TypeError, match="name them"):
            mo.Tracks()

    def test_declaration_must_be_tracks(self):
        with pytest.raises(TypeError, match="mo.Tracks"):
            mo.Model('t', tracks={'факт': 'ф'})

    def test_labels_are_nonempty_strings(self):
        with pytest.raises(TypeError, match="non-empty"):
            mo.Tracks(факт='')


class TestRoleKwargAuthoring:
    def test_role_lists_materialize(self):
        m = _declared()
        with m:
            m.выручка = mo.Variable(actual=[1] * 6, plan=[10] * 6)
        v = m.выручка._value
        assert isinstance(v, TrackValues)
        assert v.roles == ('actual', 'plan')
        assert v['plan'] == [10] * 6
        assert m.выручка.var_type == 'list'

    def test_arithmetic_broadcasts(self):
        m = _declared()
        with m:
            m.a = mo.Variable(actual=[1] * 6, plan=[10] * 6)
            m.b = m.a * 2
        assert m.b._value['actual'] == [2] * 6
        assert m.b._value['plan'] == [20] * 6

    def test_pin_spelling_shared_formula(self):
        """§16.10: positional shared expression + факт override."""
        m = _declared()
        with m:
            m.цена = mo.Variable(10)
            m.штук = mo.Variable([1, 1, 2, 2, 3, 3])
            m.доход = mo.Variable(m.цена * m.штук,
                                  actual=[9, 11, 22, 18, 33, 27])
        d = m.доход._value
        assert d['plan'] == [10, 10, 20, 20, 30, 30]
        assert d['actual'] == [9, 11, 22, 18, 33, 27]
        # the shared expression stays the canonical formula
        assert m.доход._expr is not None

    def test_coordinate_aligned_pull(self):
        """A tracked operand contributes its own same-role track."""
        m = _declared()
        with m:
            m.a = mo.Variable(actual=[1.0] * 6, plan=[10.0] * 6)
            m.b = mo.Variable(plan=m.a * 1.1, actual=m.a * 1.0)
        assert abs(m.b._value['plan'][0] - 11.0) < 1e-9
        assert m.b._value['actual'][0] == 1.0

    def test_variable_operand_is_dep_edge(self):
        m = _declared()
        with m:
            m.источник = mo.Variable([5] * 6)
            m.привязка = mo.Variable(plan=[1] * 6, actual=m.источник)
        assert m.источник in m.привязка._dependency_refs

    def test_slice_after_materialization(self):
        m = _declared()
        with m:
            m.x = mo.Variable(actual=[1] * 6, plan=[2] * 6)
        assert m.x.at(track='plan')._value == [2] * 6


class TestRoleKwargGates:
    def test_no_declaration_teaches(self):
        m = mo.Model('без', default_grain='month',
                     default_start='2025-01', default_periods=3)
        with pytest.raises(ValueError, match="mo.Tracks"):
            with m:
                m.x = mo.Variable(plan=[1, 2, 3], actual=[1, 2, 3])

    def test_undeclared_role_refused(self):
        m = _declared()
        with pytest.raises(ValueError, match="forecast"):
            with m:
                m.y = mo.Variable(actual=[1] * 6, plan=[1] * 6,
                                  forecast=[1] * 6)

    def test_subset_materializes_incrementally(self):
        """A track subset is legal — incremental authoring accumulates
        tracks one statement at a time; the mismatch law polices any
        cross-track arithmetic against an incomplete set."""
        m = _declared()
        with m:
            m.z = mo.Variable(actual=[1.0] * 6)
        assert m.z._value.roles == ('actual',)
        m.z.plan = [2.0] * 6
        assert m.z._value['plan'] == [2.0] * 6

    def test_roles_exclusive_with_formula(self):
        with pytest.raises(ValueError, match="positionally"):
            mo.Variable(formula='a+b', actual=[1])

    def test_roles_exclusive_with_tracks(self):
        with pytest.raises(ValueError, match="two spellings"):
            mo.Variable(tracks={'факт': [1]}, actual=[1])

    def test_roles_exclusive_with_indexed_by(self):
        axis = mo.Variable(['a', 'b'])
        with pytest.raises(ValueError, match="axis budget"):
            mo.Variable(indexed_by=[axis], actual=[1, 2])

    def test_track_window_extent_still_applies(self):
        m = _declared(periods=6)
        with pytest.raises(ValueError, match="cover the"):
            with m:
                m.short = mo.Variable(actual=[1, 2, 3], plan=[1, 2, 3])


class TestRoleKwargLifecycle:
    """The adversarial-review cluster (2026-08-05): lifecycle traps in
    adoption-time materialization — each of these reproduced a real
    silent-wrong-value or crash before the fixes."""

    def test_two_list_operands_no_crash_both_edges(self):
        """Membership via ``in`` routed through the overloaded __eq__ —
        list-valued operands crashed, equal scalars dropped the edge."""
        m = _declared()
        with m:
            m.бюджет = mo.Variable([1.0] * 6)
            m.леджер = mo.Variable([2.0] * 6)
            m.выручка = mo.Variable(plan=m.бюджет, actual=m.леджер)
        refs = m.выручка._dependency_refs
        assert any(r is m.бюджет for r in refs)
        assert any(r is m.леджер for r in refs)

    def test_equal_scalar_operands_keep_both_edges(self):
        m = _declared()
        with m:
            m.a = mo.Variable(5.0)
            m.b = mo.Variable(5.0)
            m.c = mo.Variable(plan=m.a, actual=m.b)
        refs = m.c._dependency_refs
        assert any(r is m.a for r in refs) and any(r is m.b for r in refs)

    def test_tracked_shared_expression_honors_overrides(self):
        """A tracked shared expr made _value TrackValues at construction
        — the old type-based guard skipped materialization and silently
        discarded the overrides."""
        m = _declared()
        with m:
            m.цена = mo.Variable(actual=[10.0] * 6, plan=[10.0] * 6)
            m.штук = mo.Variable([1.0, 1.0, 2.0, 2.0, 3.0, 3.0])
            m.доход = mo.Variable(m.цена * m.штук,
                                  actual=[9.0, 11.0, 22.0, 18.0, 33.0, 27.0])
        d = m.доход._value
        assert d['actual'] == [9.0, 11.0, 22.0, 18.0, 33.0, 27.0]
        assert d['plan'] == [10.0, 10.0, 20.0, 20.0, 30.0, 30.0]

    def test_tracked_shared_variable_does_not_alias(self):
        m = _declared()
        with m:
            m.бюджет = mo.Variable(actual=[1.0] * 6, plan=[2.0] * 6)
            m.версия = mo.Variable(m.бюджет, actual=[9.0] * 6)
        assert m.версия._value['actual'] == [9.0] * 6
        assert m.версия._value['plan'] == [2.0] * 6
        assert m.версия._value is not m.бюджет._value

    def test_consumer_before_operand_now_works(self):
        """Pure track kwargs materialize AT CONSTRUCTION (from their own
        names), so a floating operand already carries its coordinates —
        the class-body pattern works and adoption order stops mattering
        for this shape."""
        m = _declared()
        a = mo.Variable(actual=[1.0] * 6, plan=[2.0] * 6)
        b = mo.Variable(plan=a, actual=[3.0] * 6)
        m.b = b
        assert m.b._value['plan'] == [2.0] * 6
        assert m.b._value['actual'] == [3.0] * 6

    def test_unadopted_schedule_operand_is_loud(self):
        m = _declared(periods=3)
        with pytest.raises(ValueError, match="no value yet"):
            with m:
                m.v = mo.Variable(actual=[99.0, 98.0, 97.0],
                                  plan=mo.schedule({'2025-01': 100}))

    def test_floating_subtree_with_local_window_defers(self):
        """The no-declaration error fired mid-construction on a floating
        subtree whose own window resolved; now it defers to the mount."""
        pnl = mo.MultiVariable('pnl', default_grain='month',
                               default_start='2025-01', default_periods=6)
        pnl.rev = mo.Variable(actual=[1.0] * 6, plan=[2.0] * 6)
        m = mo.Model('м', tracks=mo.Tracks(actual='факт', plan='бюджет'))
        m.pnl = pnl
        assert isinstance(m.pnl.rev._value, TrackValues)
        assert m.pnl.rev._value['plan'] == [2.0] * 6

    def test_factory_subtree_with_local_window_defers(self):
        pnl = mo.MultiVariable(
            'pnl2', default_grain='month', default_start='2025-01',
            default_periods=6,
            rev=mo.Variable(actual=[1.0] * 6, plan=[2.0] * 6))
        m = mo.Model('м', tracks=mo.Tracks(actual='факт', plan='бюджет'))
        m.pnl2 = pnl
        assert isinstance(m.pnl2.rev._value, TrackValues)

    def test_copy_carries_role_stash(self):
        """Clone-on-ownership-change stripped _role_kwargs — the copy
        landed valueless with no error."""
        tmpl = mo.MultiVariable('tmpl')
        tmpl.x = mo.Variable(actual=[1.0, 2.0, 3.0], plan=[4.0, 5.0, 6.0])
        m = _declared(periods=3)
        m.x = tmpl.x
        assert isinstance(m.x._value, TrackValues)
        assert m.x._value['plan'] == [4.0, 5.0, 6.0]

    def test_schedule_as_shared_value(self):
        """The fall-through's stated purpose was unreachable: the wrapper
        never inherited _schedule from the positional Variable."""
        m = _declared(periods=3)
        with m:
            m.rate = mo.Variable(
                mo.schedule({'2025-01': 100.0, '2025-03': 120.0}),
                actual=[99.0, 98.0, 97.0])
        r = m.rate._value
        assert r['actual'] == [99.0, 98.0, 97.0]
        assert r['plan'] == [100.0, 100.0, 120.0]

    def test_plain_schedule_wrapper_materializes(self):
        m = _declared(periods=3)
        with m:
            m.ставка = mo.Variable(mo.schedule({'2025-01': 0.12}))
        assert m.ставка._value == [0.12, 0.12, 0.12]

    def test_windowless_model_scalar_roles_materialize(self):
        """The window gate swallowed role materialization; scalar tracks
        need no window (time lives inside a track)."""
        m = mo.Model('м', tracks=mo.Tracks(actual='факт', plan='бюджет'))
        with m:
            m.kpi = mo.Variable(actual=100.0, plan=120.0)
        assert isinstance(m.kpi._value, TrackValues)
        assert m.kpi.at(track='plan')._value == 120.0

    def test_eager_compute_inside_a_class_body_works(self):
        """The instance pattern (§16.11): a tracked field is used by the
        very next line, long before the instance joins a model. Pure
        kwargs materialize at construction, so the arithmetic computes
        per track instead of dying on a pending operand."""
        pnl = mo.MultiVariable('pnl3')
        pnl.rev = mo.Variable(actual=[1.0] * 6, plan=[2.0] * 6)
        pnl.margin = pnl.rev * 0.5
        assert pnl.margin._value['actual'] == [0.5] * 6
        assert pnl.margin._value['plan'] == [1.0] * 6

    def test_pin_spelling_still_defers_and_guards(self):
        """The pin spelling needs the declared vocabulary to fill the
        tracks the shared expression owes — so it still defers, and
        eager arithmetic on it still teaches."""
        pnl = mo.MultiVariable('pnl4')
        pnl.rev = mo.Variable([1.0] * 6, actual=[2.0] * 6)
        with pytest.raises(ValueError, match="materialized"):
            pnl.margin = pnl.rev * 0.5

    def test_at_track_works_on_a_floating_variable(self):
        v = mo.Variable(actual=[1.0] * 3, plan=[2.0] * 3)
        assert v.at(track='actual')._value == [1.0] * 3

    def test_at_track_on_a_pending_pin_teaches(self):
        v = mo.Variable([9.0] * 3, actual=[1.0] * 3)
        with pytest.raises(ValueError, match="materialized"):
            v.at(track='actual')

    def test_crystallization_refire_is_idempotent(self):
        m = _declared()
        sub = mo.MultiVariable('sub')
        sub.v = mo.Variable(actual=[1.0] * 6, plan=[2.0] * 6)
        m.sub = sub
        assert m.sub.v._value['actual'] == [1.0] * 6

    def test_none_track_rejected_by_value_layer(self):
        with pytest.raises(TypeError, match="None"):
            TrackValues({'факт': [1.0], 'план': None})


class TestDottedCoordinateRead:
    """§0's promise — ``выручка.факт`` reads a coordinate; sugar over
    ``.at(track=...)``, real attributes always win (getattr fires only
    on lookup miss)."""

    def test_dotted_read_equals_at(self):
        m = _declared()
        with m:
            m.выручка = mo.Variable(actual=[1.0] * 6, plan=[2.0] * 6)
        assert m.выручка.actual._value == m.выручка.at(track='actual')._value

    def test_dotted_read_in_formula(self):
        m = _declared()
        with m:
            m.a = mo.Variable(actual=[10.0] * 6, plan=[7.0] * 6)
            m.d = m.a.actual - m.a.plan
        assert m.d._value == [3.0] * 6

    def test_real_attributes_win(self):
        m = _declared()
        with m:
            m.x = mo.Variable(actual=[1.0] * 6, plan=[1.0] * 6)
        assert not isinstance(m.x.value, mo.Variable)   # the property, not a track
        assert m.x.formula is None or isinstance(m.x.formula, str)

    def test_miss_teaches_track_names(self):
        m = _declared()
        with m:
            m.x = mo.Variable(actual=[1.0] * 6, plan=[1.0] * 6)
        with pytest.raises(AttributeError, match="actual, plan"):
            m.x.forecast

    def test_plain_variable_unchanged(self):
        v = mo.Variable([1.0, 2.0])
        with pytest.raises(AttributeError):
            v.несуществующее


class TestDottedTrackWrite:
    """§16.10 MV-style incremental write — ``выручка.факт = [...]`` as
    its own statement, after the definition, without rewriting it."""

    def test_incremental_authoring(self):
        m = _declared()
        with m:
            m.выручка = mo.Variable()
        m.выручка.actual = [1.0] * 6
        m.выручка.plan = [2.0] * 6
        assert m.выручка._value['actual'] == [1.0] * 6
        assert m.выручка._value['plan'] == [2.0] * 6
        assert m.выручка.var_type == 'list'

    def test_pin_over_shared_formula(self):
        """The January door: the line has ONE formula; the actual lands
        later by dot — the formula keeps computing only the un-pinned
        tracks (given beats derived, §16.10)."""
        m = _declared()
        with m:
            m.цена = mo.Variable([10.0] * 6)
            m.штук = mo.Variable([2.0] * 6)
            m.доход = mo.Variable(m.цена * m.штук)
        m.доход.actual = [19.0, 21.0, 20.0, 20.0, 20.0, 18.0]
        v = m.доход._value
        assert v['plan'] == [20.0] * 6                    # formula-derived
        assert v['actual'] == [19.0, 21.0, 20.0, 20.0, 20.0, 18.0]
        assert m.доход.formula == 'цена * штук'           # ONE formula, intact

    def test_undeclared_name_teaches(self):
        m = _declared()
        with m:
            m.x = mo.Variable(actual=[1.0] * 6, plan=[1.0] * 6)
        with pytest.raises(ValueError, match="declared track"):
            m.x.буджет = [1.0] * 6      # typo of бюджет

    def test_wrong_length_refused(self):
        m = _declared()
        with m:
            m.x = mo.Variable(actual=[1.0] * 6, plan=[1.0] * 6)
        with pytest.raises(ValueError, match="6-period"):
            m.x.plan = [1.0, 2.0]

    def test_variable_operand_becomes_dep_edge(self):
        m = _declared()
        with m:
            m.леджер = mo.Variable([5.0] * 6)
            m.x = mo.Variable(plan=[1.0] * 6)
        m.x.actual = m.леджер
        assert any(r is m.леджер for r in m.x._dependency_refs)
        assert m.x._value['actual'] == [5.0] * 6

    def test_floating_binding_stashes_until_adoption(self):
        v = mo.Variable()
        v.actual = [1.0, 2.0, 3.0]
        v.plan = [4.0, 5.0, 6.0]
        m = _declared(periods=3)
        m.v = v
        assert m.v._value['actual'] == [1.0, 2.0, 3.0]
        assert m.v._value['plan'] == [4.0, 5.0, 6.0]

    def test_engine_attrs_untouched(self):
        v = mo.Variable([1.0])
        v.var_type = 'list'        # plain engine attr — normal path
        v.formula = 'a + b'        # property setter — normal path
        assert v.var_type == 'list'



class TestBlendLiveUniverse:
    """§16.9 — the live universe. Data lines splice givens; memoryless
    formulas coincide with the splice via broadcast; stateful chains
    re-anchor because they roll over live operand flows. План is never
    rewritten."""

    def _m(self, until='2025-02', periods=4):
        return mo.Model('м',
            tracks=mo.Tracks('факт', 'бюджет',
                             blend=mo.blend(given='факт', follow='бюджет',
                                            until=until)),
            default_grain='month', default_start='2025-01',
            default_periods=periods)

    def _flows(self, m):
        with m:
            m.поток = mo.Variable(факт=[8.0, 7.0, 0.0, 0.0],
                                  бюджет=[10.0] * 4)
        return m.поток

    def test_data_line_splices_givens(self):
        m = self._m()
        поток = self._flows(m)
        v = поток._value
        assert v['live'] == [8.0, 7.0, 10.0, 10.0]
        assert v['факт'] == [8.0, 7.0, 0.0, 0.0]      # untouched
        assert v['бюджет'] == [10.0] * 4              # план inviolable

    def test_memoryless_formula_coincides_with_splice(self):
        m = self._m()
        поток = self._flows(m)
        with m:
            m.двойной = поток * 2
        assert m.двойной._value['live'] == [16.0, 14.0, 20.0, 20.0]

    def test_recurrence_re_anchors(self):
        """live[t>close] continues LIVE's own past, not план's — the
        whole reason blend is a universe, not an output splice."""
        m = self._m()
        поток = self._flows(m)
        with m:
            m.остаток = mo.recurrence(start=0.0, formula="{prev} + {п}",
                                      variables={"п": поток}, periods=4)
        o = m.остаток._value
        assert o['бюджет'] == [0.0, 10.0, 20.0, 30.0]  # план from план
        assert o['факт'] == [0.0, 7.0, 7.0, 7.0]
        assert o['live'] == [0.0, 7.0, 17.0, 27.0]     # 17 ≠ план's 20
        assert o['live'][2] == o['live'][1] + 10.0     # re-anchored

    def test_cumsum_re_anchors(self):
        m = self._m()
        поток = self._flows(m)
        acc = mo.cumsum(поток)
        assert acc._value['live'] == [8.0, 15.0, 25.0, 35.0]

    def test_moving_the_boundary_re_splices(self):
        m = self._m(until='2025-01')
        поток = self._flows(m)
        assert поток._value['live'] == [8.0, 10.0, 10.0, 10.0]

    def test_missing_given_falls_back_to_follow(self):
        m = self._m()
        with m:
            m.аренда = mo.Variable(бюджет=[50.0] * 4)
        assert m.аренда._value['live'] == [50.0] * 4

    def test_dotted_and_at_slices_reach_live(self):
        m = self._m()
        поток = self._flows(m)
        assert поток.live._value == [8.0, 7.0, 10.0, 10.0]
        assert поток.at(track='live')._value == [8.0, 7.0, 10.0, 10.0]

    def test_dotted_binding_re_splices(self):
        m = self._m()
        with m:
            m.поток = mo.Variable(бюджет=[10.0] * 4)
        assert m.поток._value['live'] == [10.0] * 4
        m.поток.факт = [8.0, 7.0, 0.0, 0.0]
        assert m.поток._value['live'] == [8.0, 7.0, 10.0, 10.0]

    def test_blend_validates_declared_tracks(self):
        with pytest.raises(ValueError, match="not\\s+declared"):
            mo.Tracks('факт', blend=mo.blend(given='факт', follow='план',
                                             until='x'))
        with pytest.raises(ValueError, match="same track"):
            mo.blend(given='факт', follow='факт', until='x')
        with pytest.raises(ValueError, match="collides"):
            mo.Tracks('факт', 'бюджет', 'live',
                      blend=mo.blend(given='факт', follow='бюджет',
                                     until='x'))

    def test_no_blend_no_live(self):
        m = mo.Model('м', tracks=mo.Tracks('факт', 'бюджет'),
                     default_grain='month', default_start='2025-01',
                     default_periods=4)
        with m:
            m.x = mo.Variable(факт=[1.0] * 4, бюджет=[2.0] * 4)
        assert 'live' not in m.x._value.roles


class TestRoleBoundCompoundLists:
    """A role kwarg bound to a LIST of Variables lowers its elements.

    The positional path lowers compound lists through the formula
    builder; the role path used to store raw Variable objects inside
    TrackValues — and from there they leaked onto the wire as
    serialized reprs. Elements must lower to values with dependency
    edges, exactly like a whole-Variable operand.
    """

    def _m(self):
        return mo.Model('м',
            tracks=mo.Tracks('план', 'прогноз', 'факт',
                             blend=mo.blend(given='факт', follow='прогноз')),
            default_grain='month', default_start='2026-01',
            default_periods=3)

    def test_variable_elements_lower_to_values(self):
        m = self._m()
        with m:
            m.реестр = mo.MultiVariable('Реестр')
            with m.реестр as р:
                р.акт = mo.Variable(10.0)
            ряд = [mo.Variable(0.0), р.акт, mo.Variable(0.0)]
            m.x = mo.Variable(план=ряд, прогноз=ряд, факт=[None] * 3)
        v = m.x._value
        for role in ('план', 'прогноз'):
            assert v[role] == [0.0, 10.0, 0.0]
            assert all(not isinstance(c, mo.Variable) for c in v[role])
        assert v['факт'] == [None, None, None]

    def test_lowered_elements_keep_dependency_edges(self):
        m = self._m()
        with m:
            m.реестр = mo.MultiVariable('Реестр')
            with m.реестр as р:
                р.акт = mo.Variable(10.0)
            m.x = mo.Variable(план=[р.акт, 0.0, 0.0], факт=[None] * 3)
        assert any(d is m.реестр.акт for d in m.x._dependency_refs)

    def test_valueless_element_is_loud(self):
        m = self._m()
        with m:
            призрак = mo.Variable(formula='deferred')
            призрак._value = None
            with pytest.raises(ValueError):
                m.x = mo.Variable(план=[призрак, 0.0, 0.0])


class TestBlendWithoutBoundary:
    """``mo.blend`` without ``until`` — the register shape.

    A boundary is a promise that the given track is complete up to a
    date. A register typed as events arrive makes no such promise:
    facts land irregularly, out of order, and sometimes not at all.
    The boundary-less spec says the rule directly — given where it
    carries a value, follow everywhere else — with no date to keep
    current.
    """

    def _m(self, periods=6):
        return mo.Model('м',
            tracks=mo.Tracks('план', 'прогноз', 'факт',
                             blend=mo.blend(given='факт', follow='прогноз')),
            default_grain='month', default_start='2026-01',
            default_periods=periods)

    def test_given_wins_wherever_it_has_a_value(self):
        m = self._m()
        with m:
            m.поток = mo.Variable(план=[10.0] * 6, прогноз=[9.0] * 6,
                                  факт=[7.0, None, 0.0, None, 33.0, None])
        v = m.поток._value
        # 0.0 is a FACT ("nothing happened"), not a hole — it must win
        # over the forecast; only None yields.
        assert v['live'] == [7.0, 9.0, 0.0, 9.0, 33.0, 9.0]
        assert v['факт'] == [7.0, None, 0.0, None, 33.0, None]
        assert v['план'] == [10.0] * 6        # план inviolable

    def test_a_late_period_fact_is_not_overruled(self):
        """The difference a boundary makes: past it, the bounded form
        prefers ``follow`` and a typed fact loses. Without a boundary
        the fact stands wherever the author put it."""
        bounded = mo.Model('b',
            tracks=mo.Tracks('прогноз', 'факт',
                             blend=mo.blend(given='факт', follow='прогноз',
                                            until='2026-01')),
            default_grain='month', default_start='2026-01',
            default_periods=3)
        with bounded:
            bounded.x = mo.Variable(прогноз=[9.0] * 3,
                                    факт=[None, 5.0, None])
        assert bounded.x._value['live'] == [9.0, 9.0, 9.0]   # fact lost

        free = self._m(periods=3)
        with free:
            free.x = mo.Variable(план=[9.0] * 3, прогноз=[9.0] * 3,
                                 факт=[None, 5.0, None])
        assert free.x._value['live'] == [9.0, 5.0, 9.0]      # fact stands

    def test_empty_given_leaves_follow_intact(self):
        m = self._m(periods=3)
        with m:
            m.x = mo.Variable(план=[1.0] * 3, прогноз=[2.0] * 3,
                              факт=[None, None, None])
        assert m.x._value['live'] == [2.0, 2.0, 2.0]

    def test_formulas_carry_the_boundary_less_live(self):
        m = self._m(periods=3)
        with m:
            m.a = mo.Variable(план=[10.0] * 3, прогноз=[9.0] * 3,
                              факт=[7.0, None, None])
            m.b = mo.Variable(план=[1.0] * 3, прогноз=[1.0] * 3,
                              факт=[2.0, None, None])
            m.s = mo.Variable(m.a + m.b)
        assert m.s._value['live'] == [9.0, 10.0, 10.0]

    def test_repr_omits_the_absent_boundary(self):
        spec = mo.blend(given='факт', follow='прогноз')
        assert 'until' not in repr(spec)
        assert repr(mo.blend(given='ф', follow='п', until='2026-01')).count(
            'until') == 1

    def test_a_blank_boundary_is_still_a_mistake(self):
        # Omitting the boundary is a choice; spelling it empty is a typo.
        with pytest.raises(TypeError):
            mo.blend(given='факт', follow='прогноз', until='')


class TestBlendReviewHardening:
    """The 2026-08-06 adversarial-review cluster — 10 confirmed defects
    in the first blend cut, each reproduced before its fix."""

    def _m(self, until='2025-02', periods=4, name=None):
        kw = dict(given='факт', follow='бюджет', until=until)
        if name:
            kw['name'] = name
        return mo.Model('м',
            tracks=mo.Tracks('факт', 'бюджет', blend=mo.blend(**kw)),
            default_grain='month', default_start='2025-01',
            default_periods=periods)

    def test_until_compares_parsed_anchors_not_strings(self):
        """Lexicographic '2025-12' <= '2025-2' silently misrouted the
        whole splice. Boundaries now compare PARSED anchors: a lenient
        spelling splices correctly; a malformed one teaches."""
        m = self._m(until='2025-2')     # unpadded February — still February
        with m:
            m.x = mo.Variable(факт=[8.0, 7.0, 0.0, 0.0],
                              бюджет=[10.0] * 4)
        assert m.x._value['live'] == [8.0, 7.0, 10.0, 10.0]

        m2 = self._m(until='какая-то дата')
        with pytest.raises(ValueError, match="period label"):
            with m2:
                m2.x = mo.Variable(факт=[1.0] * 4, бюджет=[2.0] * 4)

    def test_scalar_lines_get_live(self):
        """Scalar-per-track assumptions skipped synthesis and broke
        every multiplication via the mismatch law."""
        m = self._m()
        with m:
            m.объём = mo.Variable(факт=[8.0, 7.0, 0.0, 0.0],
                                  бюджет=[10.0] * 4)
            m.цена = mo.Variable(факт=100.0, бюджет=90.0)
            m.выручка = m.объём * m.цена
        assert m.цена._value['live'] == [100.0, 100.0, 90.0, 90.0]
        assert m.выручка._value['live'] == [800.0, 700.0, 900.0, 900.0]

    def test_length_one_track_splices_as_constant(self):
        m = self._m()
        with m:
            m.y = mo.Variable(факт=[8.0], бюджет=[10.0])
        assert m.y._value['live'] == [8.0, 8.0, 10.0, 10.0]

    def test_ragged_track_still_teaches(self):
        """Synthesis ran before the extent law and крашился raw
        IndexError; the teaching error must win."""
        m = self._m()
        with pytest.raises(ValueError, match="4-period window"):
            with m:
                m.z = mo.Variable(факт=[8.0, 7.0], бюджет=[10.0, 10.0])

    def test_tracks_dict_authoring_live_is_loud(self):
        m = self._m()
        with pytest.raises(ValueError, match="synthesizes"):
            with m:
                m.x = mo.Variable(tracks={'live': [1.0] * 4,
                                          'бюджет': [1.0] * 4})

    def test_kwarg_authoring_live_names_the_blend(self):
        m = self._m()
        with pytest.raises(ValueError, match="synthesizes"):
            with m:
                m.x = mo.Variable(live=[1.0] * 4, бюджет=[1.0] * 4)

    def test_dotted_write_to_live_teaches_sources(self):
        m = self._m()
        with m:
            m.x = mo.Variable(факт=[1.0] * 4, бюджет=[2.0] * 4)
        with pytest.raises(ValueError, match="SYNTHESIZED"):
            m.x.live = [9.0] * 4

    def test_blend_without_window_refused_at_declaration(self):
        with pytest.raises(ValueError, match="window"):
            mo.Model('м', tracks=mo.Tracks(
                'факт', 'бюджет',
                blend=mo.blend(given='факт', follow='бюджет',
                               until='2025-02')))

    def test_pin_on_derived_line_continues_inherited_live(self):
        """A binding on a formula line must NOT replace the inherited
        live tail with the follow track (the output splice §16.9
        exists to kill)."""
        m = self._m()
        with m:
            m.поток = mo.Variable(факт=[8.0, 7.0, 0.0, 0.0],
                                  бюджет=[10.0] * 4)
            m.двойной = mo.Variable(m.поток * 2)
        inherited_tail = m.двойной._value['live'][2:]
        m.двойной.факт = [16.0, 15.0, 0.0, 0.0]
        v = m.двойной._value
        assert v['live'][:2] == [16.0, 15.0]          # pinned data
        assert v['live'][2:] == inherited_tail        # inherited, not бюджет

    def test_copy_into_blendless_context_sheds_live(self):
        m = self._m()
        with m:
            m.x = mo.Variable(факт=[1.0] * 4, бюджет=[2.0] * 4)
        assert 'live' in m.x._value.roles
        plain = mo.Model('без', tracks=mo.Tracks('факт', 'бюджет'),
                         default_grain='month', default_start='2025-01',
                         default_periods=4)
        plain.x = m.x                     # ownership change → copy
        assert 'live' not in plain.x._value.roles

    def test_none_holes_fall_to_follow(self):
        m = self._m()
        with m:
            m.x = mo.Variable(факт=[8.0, None, 0.0, 0.0],
                              бюджет=[10.0] * 4)
        assert m.x._value['live'][1] == 10.0

    def test_own_located_lines_skip_synthesis(self):
        m = self._m()
        with m:
            m.кв = mo.Variable(tracks={'факт': [1.0], 'бюджет': [2.0]},
                               start='2025-01', grain='year')
        assert 'live' not in m.кв._value.roles


class TestExcelViewTracksParam:
    """The view owns the PRINTED form (§ExcelView cascade): tracks=
    'blend' prints the clean board book (display-default rows only),
    'rows' adds labeled per-track rows, 'compare' refuses until column
    groups land; grain= is the view's lens declaration."""

    def _m(self, view=None):
        m = mo.Model('м',
            tracks=mo.Tracks('факт', 'бюджет',
                             blend=mo.blend(given='факт', follow='бюджет',
                                            until='2025-02')),
            default_grain='month', default_start='2025-01',
            default_periods=4)
        with m:
            m.поток = mo.Variable(факт=[8.0, 7.0, 0.0, 0.0],
                                  бюджет=[10.0] * 4)
        if view is not None:
            m.default_excel_view = view
        return m

    def _labels(self, m):
        import os
        import tempfile
        from openpyxl import load_workbook
        p = os.path.join(tempfile.mkdtemp(), 'т.xlsx')
        m.to_excel(p)
        return [r[0] for r in load_workbook(p).worksheets[0]
                .iter_rows(min_row=2, max_row=4, max_col=1,
                           values_only=True)]

    def test_rows_is_the_default(self):
        assert self._labels(self._m()) == [
            'Поток', 'поток · факт', 'поток · бюджет',
        ]

    def test_blend_prints_default_rows_only(self):
        labels = self._labels(self._m(mo.ExcelView(tracks='blend')))
        assert labels[0] == 'Поток' and labels[1] is None

    def test_compare_draws_period_major_column_groups(self):
        # The IDE lens's own drawing, now in the file: each period is
        # a GROUP of one column per track, the blend column a LIVE
        # splice over its group neighbours, «откл» a live difference.
        import os
        import tempfile

        from openpyxl import load_workbook

        p = os.path.join(tempfile.mkdtemp(), 'т.xlsx')
        self._m(mo.ExcelView(tracks='compare')).to_excel(p)
        ws = load_workbook(p).worksheets[0]
        # One row — the line keeps its name; no «· факт» sub-rows.
        labels = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]
        assert 'Поток' in labels
        assert not any('·' in str(x) for x in labels if x)
        row = labels.index('Поток') + 1
        # Track words under the period labels, one group per month:
        # given, follow, blend (declared order — no selection given).
        assert [ws.cell(2, c).value for c in range(2, 6)] == [
            'факт', 'бюджет', 'live', 'откл',
        ]
        assert ws.cell(1, 2).value is not None      # period label
        # The blend column splices its own GROUP NEIGHBOURS, live.
        assert ws.cell(row, 4).value == '=IF(B3="",C3,B3)'
        # «откл» is a live difference, not a baked number.
        assert ws.cell(row, 5).value == '=B3 - C3'
        # Numbers land per track: факт 8, бюджет 10.
        assert ws.cell(row, 2).value == 8
        assert ws.cell(row, 3).value == 10
        # The next month's group starts one stride (4) over.
        assert ws.cell(row, 6).value == 7

    def test_vocabularies(self):
        with pytest.raises(ValueError, match="blend"):
            mo.ExcelView(tracks='чушь')
        with pytest.raises(ValueError, match="grain"):
            mo.ExcelView(grain='decade')

    def test_resolved_view_carries_fields(self):
        from modeleon.compile.excel.view import resolve_excel_view
        m = self._m(mo.ExcelView(tracks='blend', grain='quarter'))
        r = resolve_excel_view(m.поток)
        assert r.tracks == 'blend'
        assert r.grain == 'quarter'


class TestHolesArePartialKnowledge:
    """A hole is ABSENCE, not failure (§16.9): an un-entered fact month
    propagates as a hole, so formulas and chains over it go quiet
    (blank) instead of painting the future red. The January case —
    facts arrive monthly, and the un-entered rest is the normal state,
    not an error state."""

    def test_chain_over_none_holes(self):
        m = mo.Model('м',
            tracks=mo.Tracks('факт', 'бюджет',
                             blend=mo.blend(given='факт', follow='бюджет',
                                            until='2025-02')),
            default_grain='month', default_start='2025-01',
            default_periods=4)
        with m:
            m.поток = mo.Variable(факт=[8.0, 7.0, None, None],
                                  бюджет=[10.0] * 4)
            m.остаток = mo.recurrence(start=0.0, formula="{prev} + {п}",
                                      variables={"п": m.поток}, periods=4)
        o = m.остаток._value
        assert o['факт'] == [0.0, 7.0, None, None]      # quiet, not red
        assert o['бюджет'] == [0.0, 10.0, 20.0, 30.0]
        assert o['live'] == [0.0, 7.0, 17.0, 27.0]      # re-anchored, alive

    def test_formula_over_holes_is_blank(self):
        m = mo.Model('м',
            tracks=mo.Tracks('факт', 'бюджет',
                             blend=mo.blend(given='факт', follow='бюджет',
                                            until='2025-02')),
            default_grain='month', default_start='2025-01',
            default_periods=4)
        with m:
            m.выручка = mo.Variable(факт=[8.0, 7.0, None, None],
                                    бюджет=[10.0] * 4)
            m.расходы = mo.Variable(факт=[6.0, 5.0, None, None],
                                    бюджет=[7.0] * 4)
            m.прибыль = mo.Variable(m.выручка - m.расходы)
        p = m.прибыль._value
        assert p['факт'] == [2.0, 2.0, None, None]
        assert p['бюджет'] == [3.0] * 4
        assert p['live'] == [2.0, 2.0, 3.0, 3.0]

    def test_genuine_errors_still_absorb(self):
        """Absence is quiet; failure stays loud."""
        a = mo.Variable([1.0, 2.0])
        b = mo.Variable([0.0, 2.0])
        assert (a / b)._value[0] == '#DIV/0!'


class TestTracksInsideInstances:
    """§16.11 — the Arrual shape: a Collection of projects, each with a
    tracked field authored by constructor params. Pure track kwargs
    materialize PROVISIONALLY at construction (from their own names),
    so the class body's next line computes; adoption validates the
    names against the declaration, attaches the blend, and memoryless
    derived lines get their live as the splice of their own tracks
    (splice-equal for per-period formulas)."""

    def _model(self):
        class Проект(mo.MultiVariableClass):
            def compute(self, ставка=100.0, часы_план=None,
                        часы_факт=None):
                self.часы = mo.Variable(план=часы_план, факт=часы_факт)
                self.выручка = self.часы * ставка

        m = mo.Model('м',
            tracks=mo.Tracks('факт', 'план',
                             blend=mo.blend(given='факт', follow='план',
                                            until='2025-02')),
            default_grain='month', default_start='2025-01',
            default_periods=4)
        with m:
            m.альфа = Проект(ставка=100.0, часы_план=[10.0] * 4,
                             часы_факт=[9.0, 11.0, None, None])
            m.бета = Проект(ставка=150.0, часы_план=[20.0] * 4,
                            часы_факт=[18.0, 21.0, None, None])
            m.портфель = m.альфа.выручка + m.бета.выручка
        return m

    def test_tracked_field_in_class_body(self):
        m = self._model()
        v = m.альфа.часы._value
        assert v['план'] == [10.0] * 4
        assert v['факт'] == [9.0, 11.0, None, None]
        assert v['live'] == [9.0, 11.0, 10.0, 10.0]

    def test_derived_line_in_class_body_gets_spliced_live(self):
        m = self._model()
        v = m.альфа.выручка._value
        assert v['live'] == [900.0, 1100.0, 1000.0, 1000.0]

    def test_portfolio_aggregation_across_instances(self):
        m = self._model()
        p = m.портфель._value
        assert p['план'] == [4000.0] * 4
        assert p['факт'][:2] == [3600.0, 4250.0]
        assert p['live'] == [3600.0, 4250.0, 4000.0, 4000.0]

    def test_typo_in_class_body_still_teaches_at_adoption(self):
        class Кривой(mo.MultiVariableClass):
            def compute(self):
                self.x = mo.Variable(фактт=[1.0] * 4)

        m = mo.Model('м', tracks=mo.Tracks('факт', 'план'),
                     default_grain='month', default_start='2025-01',
                     default_periods=4)
        with pytest.raises(ValueError, match="фактт"):
            with m:
                m.к = Кривой()

    def test_stateful_derived_in_class_body_has_no_live(self):
        """A recurrence inside a class body rolled before its driver
        had live — a splice would be the §16.9 bug, so live is honestly
        ABSENT and the mismatch law polices any cross-use."""
        class Накопитель(mo.MultiVariableClass):
            def compute(self, поток_план=None, поток_факт=None):
                self.поток = mo.Variable(план=поток_план, факт=поток_факт)
                self.остаток = mo.cumsum(self.поток)

        m = mo.Model('м',
            tracks=mo.Tracks('факт', 'план',
                             blend=mo.blend(given='факт', follow='план',
                                            until='2025-02')),
            default_grain='month', default_start='2025-01',
            default_periods=4)
        with m:
            m.n = Накопитель(поток_план=[10.0] * 4,
                             поток_факт=[9.0, 11.0, None, None])
        assert 'live' not in m.n.остаток._value.roles


class TestScalarBlend:
    """Non-temporal lines carry the axis too — without a timeline."""

    def _model(self):
        m = mo.Model(
            'm', display_name='M', default_grain='month',
            default_start='2026-01', default_periods=6,
            tracks=mo.Tracks(
                'план', 'факт',
                blend=mo.blend(given='факт', follow='план',
                               until='2026-03', name='прогноз')))
        return m

    def test_a_scalar_line_blends_to_a_scalar(self):
        # «Сумма договора»: план 50 000, факт (доп.соглашение) 52 000.
        # The time splice used to spread the constant into a full
        # 6-period series — a record field quietly turned temporal.
        m = self._model()
        with m:
            m.p = mo.MultiVariable('P', excel_props={'tab': True})
            m.p._records_container = True  # the Collection contract
            with m.p as p:
                p.сумма = mo.Variable(план=50000, факт=52000,
                                      display_name='Сумма')
        tv = m.p.сумма._value
        assert tv['прогноз'] == 52000
        assert not isinstance(tv['прогноз'], list)

    def test_fact_absent_the_scalar_forecast_is_the_plan(self):
        m = self._model()
        with m:
            m.p = mo.MultiVariable('P', excel_props={'tab': True})
            m.p._records_container = True
            with m.p as p:
                p.срок = mo.Variable(план=12, display_name='Срок')
        assert m.p.срок._value['прогноз'] == 12

    def test_temporal_lines_still_splice_per_period(self):
        m = self._model()
        with m:
            m.p = mo.MultiVariable('P', excel_props={'tab': True})
            with m.p as p:
                p.ряд = mo.Variable(план=[1.0] * 6, display_name='Ряд')
        assert len(m.p.ряд._value['прогноз']) == 6


class TestBlendRowIsAlive:
    """The blended row references its subrows — the file stays live."""

    def _book(self, model):
        import os
        import tempfile

        import openpyxl

        path = os.path.join(tempfile.mkdtemp(), "b.xlsx")
        model.to_excel(path)
        return openpyxl.load_workbook(path)

    def _model(self, **view_kw):
        m = mo.Model(
            'm', display_name='M', default_grain='month',
            default_start='2026-01', default_periods=4,
            tracks=mo.Tracks(
                'план', 'факт',
                blend=mo.blend(given='факт', follow='план',
                               until='2026-02')),
            **view_kw)
        with m:
            m.p = mo.MultiVariable('P', excel_props={'tab': True})
            with m.p as p:
                p.выручка = mo.Variable(
                    план=[10.0] * 4, факт=[12.0, 11.0, None, None],
                    display_name='Выручка')
                p.налог = p.выручка * 0.1
        return m

    def test_the_splice_is_a_formula_over_its_subrows(self):
        ws = self._book(self._model())['P']
        # Blend row 2; план row 3; факт row 4. Before the boundary the
        # FACT wins with a fallback; after it the PLAN does.
        assert ws['B2'].value == '=IF(B4="",B3,B4)'
        assert ws['C2'].value == '=IF(C4="",C3,C4)'
        assert ws['D2'].value == '=IF(D3="",D4,D3)'
        assert ws['E2'].value == '=IF(E3="",E4,E3)'

    def test_a_derived_line_keeps_its_own_formula(self):
        # налог = выручка × 0.1 — it references the BLEND row (the live
        # universe). Replaying the splice on an output would be wrong.
        ws = self._book(self._model())['P']
        assert ws['B5'].value == '=B2 * 0.1'

    def test_a_computed_plan_with_authored_facts_still_splices(self):
        # The customer shape: план is a FORMULA (a sum over a dated
        # register), факт and прогноз are authored series. The line has
        # both an expression and role kwargs — the engine calls it a
        # DATA line and splices it, so the workbook must splice it too.
        # Left with its base expression the file showed ПЛАН where the
        # app showed the blend (акт перенесён с сентября на октябрь).
        m = mo.Model(
            'm', display_name='M', default_grain='month',
            default_start='2026-01', default_periods=4,
            tracks=mo.Tracks(
                'план', 'прогноз', 'факт',
                blend=mo.blend(given='факт', follow='прогноз')),
        )
        with m:
            m.p = mo.MultiVariable('P', excel_props={'tab': True})
            with m.p as p:
                p.акты = mo.Variable([0.0, 20.0, 0.0, 0.0],
                                     display_name='Акты')
                p.выручка = mo.Variable(
                    p.акты * 1.0,
                    прогноз=[0.0, 0.0, 20.0, 0.0],
                    факт=[None, None, None, None],
                    display_name='Выручка')
        ws = self._book(m)['P']
        rows = {ws.cell(r, 1).value: r for r in range(1, ws.max_row + 1)}
        blend = rows['Выручка']
        given = rows['Выручка · факт']
        follow = rows['Выручка · прогноз']
        # Boundary-less: given wins wherever it carries a value.
        assert ws.cell(blend, 2).value == (
            f'=IF(B{given}="",B{follow},B{given})')
        # …and the splice, not the план formula: september is the
        # прогноз row's zero, october its 20.
        assert ws.cell(blend, 4).value == (
            f'=IF(D{given}="",D{follow},D{given})')

    def test_the_settled_ink_is_a_rule_not_a_stamp(self):
        # The blended cell goes blue WHILE a fact stands behind it —
        # as a conditional rule over the fact cell, so typing a fact
        # into the downloaded file colours its own month. A stamped
        # font would freeze provenance at download time, the same
        # dead-values habit the splice formulas exist to kill.
        ws = self._book(self._model())['P']
        rules = ws.conditional_formatting
        by_range = {}
        for rng in rules:
            for rule in rng.rules:
                by_range[str(rng.sqref)] = rule
        # Blend row 2, fact row 4: each month keys off ITS fact cell.
        assert by_range['B2'].formula == ['B4<>""']
        assert by_range['C2'].formula == ['C4<>""']
        # Font only — a fill would cut holes in the house style's
        # section bands.
        assert by_range['B2'].dxf.font is not None
        assert by_range['B2'].dxf.fill is None

    def test_the_blend_only_book_bakes_nothing_extra(self):
        # tracks='blend' emits the default row ALONE — no subrows to
        # reference, so the row keeps its computed values.
        m = self._model(
            default_excel_view=mo.ExcelView(tracks='blend'))
        ws = self._book(m)['P']
        assert ws['B2'].value == 12.0
        labels = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]
        assert not any('·' in str(x) for x in labels if x)  # no subrows


class TestTrackSubrowsCarryTheirOwnFormula:
    """A DERIVED track explains itself in ITS OWN coordinates.

    The blend row has always carried the line's formula over its
    operands' blend rows. The other subrows got numbers, so «Итого ·
    план» sat frozen under a row that recomputed — change January's
    план in the downloaded file and the план total did not move. The
    engine already knows the answer (a derived track's value IS the
    line's expression over the operands' same-track values), so the
    same expression through a per-track address book is the honest
    formula.
    """

    def _book(self, tmp_path, **kw):
        import openpyxl

        m = mo.Model(
            "m", display_name="M",
            tracks=mo.Tracks("план", "факт",
                             blend=mo.blend(given="факт", follow="план")),
            default_grain="month", default_start="2026-01",
            default_periods=2,
            default_excel_view=mo.ExcelView(tracks="rows"),
        )
        with m:
            m.s = mo.MultiVariable("S", excel_props={"tab": True})
            with m.s as s:
                s.a = mo.Variable([10.0, 20.0], display_name="А",
                                  факт=[11.0, None])
                s.ставка = mo.Variable(0.5, display_name="Ставка")
                s.derived = mo.Variable(s.a * s.ставка, display_name="Д")
                s.mixed = mo.Variable(s.a + s.a, display_name="М",
                                      факт=[99.0, 98.0])
        path = tmp_path / "b.xlsx"
        m.to_excel(str(path))
        ws = openpyxl.load_workbook(str(path))["S"]
        return {
            (ws.cell(row=r, column=1).value
             or ws.cell(row=r, column=2).value): [
                ws.cell(row=r, column=c).value for c in (3, 4)
            ]
            for r in range(1, 40)
            if (ws.cell(row=r, column=1).value
                or ws.cell(row=r, column=2).value)
        }

    def test_a_derived_track_reads_its_own_track(self, tmp_path):
        rows = self._book(tmp_path)
        план = rows["Д · план"][0]
        факт = rows["Д · факт"][0]
        assert isinstance(план, str) and план.startswith("=")
        assert isinstance(факт, str) and факт.startswith("=")
        # …and they point at DIFFERENT rows — each its own track's.
        assert план != факт

    def test_an_untracked_operand_keeps_its_single_cell(self, tmp_path):
        rows = self._book(tmp_path)
        # «Ставка» has no tracks, so every track's formula multiplies by
        # the same cell — a constant reads the same from anywhere.
        план, факт = rows["Д · план"][0], rows["Д · факт"][0]
        общий = set(re.findall(r"B\d+", план)) & set(re.findall(r"B\d+", факт))
        assert общий, (план, факт)

    def test_an_authored_track_stays_data(self, tmp_path):
        rows = self._book(tmp_path)
        # «А» typed both tracks; «М» typed only факт. Typed series are
        # data — a formula there would overwrite what the user entered.
        assert rows["А · план"][0] == 10.0
        assert rows["А · факт"][0] == 11.0
        assert rows["М · факт"][0] == 99.0
        # …while М's план is derived and does get one.
        assert str(rows["М · план"][0]).startswith("=")

    def test_the_blend_row_is_left_to_its_own_writer(self, tmp_path):
        rows = self._book(tmp_path)
        # A spliced line's blend still asks its subrows; a derived
        # line's blend still carries the shared formula.
        assert str(rows["А"][0]).startswith("=IF(")
        assert str(rows["Д"][0]).startswith("=")


class TestLoopBuiltIntermediatesReferenceTheirRows:
    """A loop-built model computes intermediates as per-month
    expressions in local Python lists, then materializes the same
    lists as rows. The row's wrapper shares the ORIGINAL expression
    tree with every downstream consumer — identity is the only join —
    and without it each consumer unrolled the intermediates wholesale:
    a 1040-char payroll cell repeating the ОПВ cap four times, one row
    under the ОПВ row that holds it.
    """

    def _book(self, tmp_path):
        import openpyxl

        m = mo.Model(
            "m", display_name="M",
            tracks=mo.Tracks("план", "факт",
                             blend=mo.blend(given="факт", follow="план")),
            default_grain="month", default_start="2026-01",
            default_periods=2,
        )
        with m:
            m.s = mo.MultiVariable("S", excel_props={"tab": True})
            with m.s as s:
                s.ставка = mo.Variable(0.1, display_name="Ставка")
                оклады = [mo.Variable(100.0), mo.Variable(200.0)]
                налог = [о * s.ставка for о in оклады]
                чистыми = [о - н for о, н in zip(оклады, налог)]
                s.оклад = mo.Variable(оклады, display_name="Оклад",
                                      факт=[None, None])
                s.налог = mo.Variable(налог, display_name="Налог",
                                      факт=[None, None])
                s.чистыми = mo.Variable(чистыми, display_name="Чистыми",
                                        факт=[None, None])
        path = tmp_path / "b.xlsx"
        m.to_excel(str(path))
        ws = openpyxl.load_workbook(str(path))["S"]
        return {
            (ws.cell(row=r, column=1).value
             or ws.cell(row=r, column=2).value): [
                ws.cell(row=r, column=c).value for c in (3, 4)
            ]
            for r in range(1, 30)
            if (ws.cell(row=r, column=1).value
                or ws.cell(row=r, column=2).value)
        }

    def test_the_consumer_references_the_intermediate_rows(self, tmp_path):
        rows = self._book(tmp_path)
        налог = str(rows["Налог · план"][0])
        чистыми = str(rows["Чистыми · план"][0])
        # Налог references the оклад subrow's cell, not an unrolled
        # copy of its expression…
        assert налог.startswith("=") and "*" in налог
        # …and Чистыми references BOTH rows by cell: it must contain
        # cell refs and no second copy of the multiplication.
        assert чистыми.startswith("=")
        assert "*" not in чистыми, чистыми
    def test_lengths_stay_flat(self, tmp_path):
        rows = self._book(tmp_path)
        # The unrolled form grew with every hop; the referenced form
        # stays a handful of cells.
        assert len(str(rows["Чистыми · план"][0])) < 30
