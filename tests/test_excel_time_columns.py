"""A period header names period columns — and nothing else.

Two rules, one idea: the time axis is a property of the LINE, not of
the model. A model declaring a window doesn't make every sheet a
time series, and it doesn't make a constant a January figure.

  1. No period header on a sheet that carries no per-period line.
     A sheet of constants (a price list, say) must not get
     «Jan 2026» over column B — its values would read as January
     figures.
  2. On a sheet that carries BOTH, constants get a column of their
     own, left of the periods, and the header starts after it.

The addresses move with the rule, so formulas follow for free — that
is what makes this a layout change rather than a cosmetic one.
"""

import os
import tempfile

import openpyxl
import pytest

import modeleon as mo


def _sheet(model, name):
    path = os.path.join(tempfile.mkdtemp(), "t.xlsx")
    model.to_excel(path)
    return openpyxl.load_workbook(path)[name]


def _grid(ws, cols=6):
    return [
        [ws.cell(r, c).value for c in range(1, cols + 1)]
        for r in range(1, ws.max_row + 1)
    ]


def _windowed():
    m = mo.Model("m", display_name="M")
    m.default_start, m.default_grain, m.default_periods = "2026-01", "month", 3
    return m


class TestTheHeaderFollowsContent:
    def test_a_register_of_constants_gets_no_period_header(self):
        m = _windowed()
        with m:
            m.roles = mo.MultiVariable("Roles", excel_props={"tab": True})
            with m.roles as d:
                d.вилка = mo.Variable(500, display_name="Вилка")
                d.разряд = mo.Variable(3, display_name="Разряд")
            # A second sheet DOES carry periods — the window is real,
            # so this pins content and not "the model has no window".
            m.план = mo.MultiVariable("План", excel_props={"tab": True})
            with m.план as p:
                p.выручка = mo.Variable([1, 2, 3], display_name="Выручка")
        g = _grid(_sheet(m, "Roles"))
        assert g[0][:2] == ["Вилка", 500]
        assert not any("2026" in str(c) for row in g for c in row if c)

    def test_the_sheet_that_has_periods_keeps_its_header(self):
        m = _windowed()
        with m:
            m.план = mo.MultiVariable("План", excel_props={"tab": True})
            with m.план as p:
                p.выручка = mo.Variable([1, 2, 3], display_name="Выручка")
        assert _grid(_sheet(m, "План"))[0][1:4] == [
            "Jan 2026",
            "Feb 2026",
            "Mar 2026",
        ]

    def test_two_sheets_are_judged_separately(self):
        # Content belongs to the sheet that prints it: one tab having
        # periods must not put a header over another that hasn't.
        m = _windowed()
        with m:
            m.курсы = mo.MultiVariable("Курсы", excel_props={"tab": True})
            with m.курсы as c:
                c.usd = mo.Variable(450, display_name="USD")
            m.план = mo.MultiVariable("План", excel_props={"tab": True})
            with m.план as p:
                p.ряд = mo.Variable([1, 2, 3], display_name="Ряд")
        assert not any(
            "2026" in str(c)
            for row in _grid(_sheet(m, "Курсы"))
            for c in row
            if c
        )
        assert _grid(_sheet(m, "План"))[0][1] == "Jan 2026"


class TestConstantsKeepTheirOwnColumn:
    def _mixed(self):
        m = _windowed()
        with m:
            m.план = mo.MultiVariable("План", excel_props={"tab": True})
            with m.план as p:
                p.ставка = mo.Variable(0.12, display_name="Ставка")
                p.выручка = mo.Variable([100, 110, 120], display_name="Выручка")
                p.ндс = p.выручка * p.ставка
        return m

    def test_a_constant_is_not_a_january_figure(self):
        ws = _sheet(self._mixed(), "План")
        g = _grid(ws)
        # Periods start at C; B is the constants column.
        assert g[0][1] is None
        assert g[0][2:5] == ["Jan 2026", "Feb 2026", "Mar 2026"]
        assert g[1][:3] == ["Ставка", 0.12, None]
        assert g[2][1:5] == ["Выручка", 100, 110, 120][1:] or g[2][2:5] == [
            100,
            110,
            120,
        ]

    def test_formulas_follow_the_moved_cells(self):
        # The point of moving ADDRESSES rather than painting: a
        # reference to the constant still lands on it.
        g = _grid(_sheet(self._mixed(), "План"))
        ндс = next(row for row in g if row[0] == "Ндс")
        assert ндс[2] == "=C3 * B2"

    def test_the_pane_freezes_left_of_the_periods(self):
        # The constants column must stay in view while the periods
        # scroll — it is a lead column in everything but name.
        assert _sheet(self._mixed(), "План").freeze_panes == "C2"

    def test_no_empty_column_when_the_sheet_is_all_periods(self):
        m = _windowed()
        with m:
            m.план = mo.MultiVariable("План", excel_props={"tab": True})
            with m.план as p:
                p.выручка = mo.Variable([1, 2, 3], display_name="Выручка")
        g = _grid(_sheet(m, "План"))
        assert g[0][1] == "Jan 2026"
        assert g[1][:4] == ["Выручка", 1, 2, 3]

    def test_lead_metadata_columns_still_come_first(self):
        m = _windowed()
        m.default_excel_view = mo.ExcelView(
            meta={"article": True, "unit": "Ед. изм.", "label": "Наименование"}
        )
        with m:
            m.план = mo.MultiVariable("План", excel_props={"tab": True})
            with m.план as p:
                p.ставка = mo.Variable(0.12, display_name="Ставка")
                p.выручка = mo.Variable([1, 2, 3], display_name="Выручка")
        g = _grid(_sheet(m, "План"), cols=8)
        # A = №, B = name, C = unit, D = constants, E.. = periods.
        assert g[0][1] == "Наименование" and g[0][2] == "Ед. изм."
        assert g[0][4] == "Jan 2026"
        ставка = next(row for row in g if row[1] == "Ставка")
        assert ставка[3] == pytest.approx(0.12)


class TestProjectedFieldColumns:
    """Any field can ride its own lead column, by analogy with the
    unit: ``meta={'fields': [('description', 'Комментарий')]}``.

    A list of PAIRS rather than free-standing keys because a view is a
    real MV tree — a field named ``description`` would collide with
    ``MultiVariable.description`` the moment it became an attribute of
    the meta node.
    """

    def _model(self, fields):
        m = _windowed()
        m.default_excel_view = mo.ExcelView(
            meta={"unit": "Ед. изм.", "label": "Наименование", "fields": fields}
        )
        with m:
            m.план = mo.MultiVariable("План", excel_props={"tab": True})
            with m.план as p:
                p.ставка = mo.Variable(
                    0.12,
                    display_name="Ставка",
                    unit="%",
                    description="по НК РК",
                )
                p.выручка = mo.Variable(
                    [1, 2, 3], display_name="Выручка", unit="₸"
                )
        return m

    def test_a_field_gets_a_column_and_a_caption(self):
        g = _grid(_sheet(self._model([("description", "Комментарий")]), "План"), 8)
        # A = name, B = unit, C = the projected field, D = constants,
        # E.. = periods.
        assert g[0][:5] == ["Наименование", "Ед. изм.", "Комментарий", None, "Jan 2026"]
        assert g[1][:4] == ["Ставка", "%", "по НК РК", 0.12]

    def test_a_row_without_the_field_leaves_it_empty(self):
        # A projected column reports what the model holds. It never
        # invents a value, and never borrows the row above's.
        g = _grid(_sheet(self._model([("description", "Комментарий")]), "План"), 8)
        выручка = next(row for row in g if row[0] == "Выручка")
        assert выручка[2] is None

    def test_several_fields_keep_declaration_order(self):
        m = _windowed()
        m.default_excel_view = mo.ExcelView(
            meta={
                "fields": [
                    ("description", "Комментарий"),
                    ("number_format", "Формат"),
                ]
            }
        )
        with m:
            m.план = mo.MultiVariable("План", excel_props={"tab": True})
            with m.план as p:
                p.ставка = mo.Variable(
                    0.12,
                    display_name="Ставка",
                    description="по НК РК",
                    excel_props={"number_format": "0.0%"},
                )
                p.ряд = mo.Variable([1, 2, 3], display_name="Ряд")
        g = _grid(_sheet(m, "План"), 8)
        assert g[0][1:3] == ["Комментарий", "Формат"]
        # The second one reads from excel_props — a field lives
        # wherever the model honestly keeps it.
        assert g[1][1:3] == ["по НК РК", "0.0%"]

    def test_an_unnamed_field_column_still_carries_its_values(self):
        # ``True`` instead of a caption: the column renders headerless,
        # the same courtesy the unit column has always had.
        g = _grid(_sheet(self._model([("description", True)]), "План"), 8)
        assert g[0][2] is None
        assert g[1][2] == "по НК РК"


class TestConstantsCaption:
    """The constants column can be HEADED — ``meta={'constants':
    'Значение'}`` — the book form's own word for it. Without a
    header the column reads as the sheet having slipped against its
    own timeline.
    """

    def _model(self, **meta):
        m = _windowed()
        if meta:
            m.default_excel_view = mo.ExcelView(meta=meta)
        with m:
            m.план = mo.MultiVariable("План", excel_props={"tab": True})
            with m.план as p:
                p.ставка = mo.Variable(0.12, display_name="Ставка")
                p.выручка = mo.Variable([1, 2, 3], display_name="Выручка")
        return m

    def test_the_caption_heads_the_column(self):
        g = _grid(_sheet(self._model(constants="Значение"), "План"))
        assert g[0][1:3] == ["Значение", "Jan 2026"]
        assert g[1][:2] == ["Ставка", 0.12]

    def test_undeclared_stays_headerless(self):
        # Same rule as the unit column always had: the flag is the
        # caption, and nothing invents wording the model didn't say.
        g = _grid(_sheet(self._model(), "План"))
        assert g[0][1] is None
        assert g[0][2] == "Jan 2026"

    def test_no_caption_without_the_column(self):
        # All-periodic sheet: no constants column is reserved, so the
        # declared caption must NOT claim the label column.
        m = _windowed()
        m.default_excel_view = mo.ExcelView(meta={"constants": "Значение"})
        with m:
            m.план = mo.MultiVariable("План", excel_props={"tab": True})
            with m.план as p:
                p.выручка = mo.Variable([1, 2, 3], display_name="Выручка")
        g = _grid(_sheet(m, "План"))
        assert g[0][1] == "Jan 2026"
        assert not any("Значение" in str(c) for row in g for c in row if c)


class TestALineOnItsOwnWindow:
    """A line declared with its own ``start=`` / ``grain=`` that differs
    from its sheet's window cannot be written yet. Every series is
    written from the sheet's first period column, so such a line would
    sit under the wrong dates, and formulas reading it would pair the
    wrong periods — the file would compute other numbers than Python.
    The export refuses before writing anything.
    """

    def _refused(self, model):
        path = os.path.join(tempfile.mkdtemp(), "t.xlsx")
        with pytest.raises(ValueError) as err:
            model.to_excel(path)
        assert not os.path.exists(path)
        return str(err.value)

    def _later_start(self):
        m = _windowed()
        with m:
            m.план = mo.MultiVariable("План", excel_props={"tab": True})
            with m.план as p:
                p.выручка = mo.Variable([1, 2, 3], display_name="Выручка")
                p.бонус = mo.Variable(
                    [10, 20], start="2026-02", grain="month",
                    display_name="Бонус",
                )
        return m

    def test_a_later_start_is_refused(self):
        self._refused(self._later_start())

    def test_a_different_grain_is_refused(self):
        m = _windowed()
        with m:
            m.план = mo.MultiVariable("План", excel_props={"tab": True})
            with m.план as p:
                p.выручка = mo.Variable([1, 2, 3], display_name="Выручка")
                p.квартал = mo.Variable(
                    [30, 40], start="2026-Q1", grain="quarter",
                    display_name="Квартал",
                )
        msg = self._refused(m)
        assert "'Квартал'" in msg and "quarter" in msg

    def test_the_message_names_the_line_both_windows_and_the_way_out(self):
        msg = self._refused(self._later_start())
        assert "'Бонус'" in msg and "m.план.бонус" in msg
        assert "2026-02" in msg and "month" in msg
        assert "'План'" in msg and "2026-01" in msg
        assert "start=" in msg and "extend=" in msg

    def test_an_own_window_equal_to_the_sheets_exports(self):
        m = _windowed()
        with m:
            m.план = mo.MultiVariable("План", excel_props={"tab": True})
            with m.план as p:
                p.ряд = mo.Variable(
                    [1, 2, 3], start="2026-01", grain="month",
                    display_name="Ряд",
                )
        g = _grid(_sheet(m, "План"))
        assert g[0][1:4] == ["Jan 2026", "Feb 2026", "Mar 2026"]
        assert g[1][:4] == ["Ряд", 1, 2, 3]

    def test_the_same_month_spelled_without_a_leading_zero_exports(self):
        # '2026-1' and '2026-01' are the same month: the line is on
        # its sheet's window, whatever the spelling.
        m = _windowed()
        with m:
            m.план = mo.MultiVariable("План", excel_props={"tab": True})
            with m.план as p:
                p.ряд = mo.Variable(
                    [1, 2, 3], start="2026-1", grain="month",
                    display_name="Ряд",
                )
        assert _grid(_sheet(m, "План"))[1][:4] == ["Ряд", 1, 2, 3]

    def test_an_own_window_that_ends_early_is_refused(self):
        # Same start and grain, fewer periods: Python continues the
        # series by its extend rule (20 held into March), while a
        # formula in the book would read an empty or wrong cell there.
        m = _windowed()
        with m:
            m.план = mo.MultiVariable("План", excel_props={"tab": True})
            with m.план as p:
                p.выручка = mo.Variable([1, 2, 3], display_name="Выручка")
                p.ставка = mo.Variable(
                    [10, 20], start="2026-01", grain="month",
                    extend=mo.hold(), display_name="Ставка",
                )
                p.итого = p.выручка + p.ставка
        assert m.план.итого.value == [11, 22, 23]
        msg = self._refused(m)
        assert "'Ставка'" in msg and "2 period(s)" in msg
        assert "3 period(s)" in msg

    def test_extend_and_schedule_on_the_models_window_still_export(self):
        m = _windowed()
        with m:
            m.план = mo.MultiVariable("План", excel_props={"tab": True})
            with m.план as p:
                p.выручка = mo.Variable([1, 2, 3], display_name="Выручка")
                p.бонус = mo.Variable(
                    [10, 20], extend=mo.zero(), display_name="Бонус"
                )
                p.ставка = mo.schedule({"2026-01": 0.1, "2026-03": 0.2})
                p.итого = p.выручка + p.бонус
        g = _grid(_sheet(m, "План"))
        assert next(r for r in g if r[0] == "Бонус")[1:4] == [10, 20, 0]
        assert next(r for r in g if r[0] == "Ставка")[1:4] == [0.1, 0.1, 0.2]
        assert next(r for r in g if r[0] == "Итого")[1] == "=B2 + B3"

    def test_a_model_without_a_window_is_unaffected(self):
        m = mo.Model("m", display_name="M")
        with m:
            m.план = mo.MultiVariable("План", excel_props={"tab": True})
            with m.план as p:
                p.бонус = mo.Variable(
                    [10, 20], start="2026-02", grain="month",
                    display_name="Бонус",
                )
        assert _grid(_sheet(m, "План"))[0][:3] == ["Бонус", 10, 20]

    def test_a_line_derived_from_it_is_not_written_either(self):
        m = _windowed()
        with m:
            m.план = mo.MultiVariable("План", excel_props={"tab": True})
            with m.план as p:
                p.выручка = mo.Variable([1, 2, 3], display_name="Выручка")
                p.бонус = mo.Variable(
                    [10, 20], start="2026-02", grain="month",
                    extend=mo.zero(), display_name="Бонус",
                )
                p.итого = p.выручка + p.бонус
        # The sum itself runs on the sheet's window — it is its operand
        # that would be misplaced, and the whole book with it.
        assert m.план.итого.time.start == "2026-01"
        assert m.план.итого.value == [1, 12, 23]
        self._refused(m)

    def test_a_derived_line_reading_an_unattached_one_is_refused(self):
        # The operand lives outside the model, so it has no cells of
        # its own: its values would be inlined into the sum from the
        # first period — under the wrong dates.
        бонус = mo.Variable(
            [10, 20], start="2026-02", grain="month",
            extend=mo.zero(), display_name="Бонус",
        )
        m = _windowed()
        with m:
            m.план = mo.MultiVariable("План", excel_props={"tab": True})
            with m.план as p:
                p.выручка = mo.Variable([1, 2, 3], display_name="Выручка")
                p.итого = p.выручка + бонус
        assert m.план.итого.value == [1, 12, 23]
        msg = self._refused(m)
        assert "'Итого'" in msg and "'Бонус'" in msg
