"""A track authored as a list of REFERENCES keeps them as formulas.

``forecast=[known, None, None, solved_cell]`` is a sparse row of links — each
milestone cell points at the cell that holds the number. The role stash
keeps the original list, so the track's expression is a ListExpr: a
reference per period where the list holds a Variable, a literal (or a
blank) elsewhere.
"""

import os
import tempfile

import openpyxl

import modeleon as mo
from modeleon.compile.excel.view import ExcelView
from modeleon.core.expr import ListExpr


def _book(model, name):
    path = os.path.join(tempfile.mkdtemp(), "t.xlsx")
    model.to_excel(path)
    return openpyxl.load_workbook(path)[name]


def _rows(ws):
    return {ws.cell(r, 1).value: r for r in range(1, ws.max_row + 1)}


def test_sparse_reference_list_emits_links():
    m = mo.Model("m", display_name="M",
                 tracks=mo.Tracks("forecast", "actual"),
                 default_grain="month", default_start="2026-01",
                 default_periods=4)
    m.default_excel_view = ExcelView(tracks="blend")
    m.s = mo.MultiVariable("S", excel_props={"tab": True})
    with m.s as s:
        s.known = mo.Variable(280.0, display_name="Known")
        s.solved = mo.Variable(1000.0, display_name="Solved")
        s.points = mo.Variable(forecast=[s.known, None, s.solved * 2, s.solved],
                               display_name="Points")
    assert isinstance(m.s.points.track_expr("forecast"), ListExpr)
    ws = _book(m, "S")
    rows = _rows(ws)
    r = rows["Points"]
    cells = [ws.cell(r, c).value for c in range(2, ws.max_column + 1)]
    vals = [v for v in cells if v is not None]   # blanks and the constants column drop out
    assert vals[0] == f"=B{rows['Known']}", vals
    assert vals[1] == f"=B{rows['Solved']} * 2", vals
    assert vals[2] == f"=B{rows['Solved']}", vals


def test_all_literal_list_stays_data():
    m = mo.Model("m", display_name="M", tracks=mo.Tracks("forecast", "actual"),
                 default_grain="month", default_start="2026-01",
                 default_periods=3)
    m.s = mo.MultiVariable("S", excel_props={"tab": True})
    with m.s as s:
        s.points = mo.Variable(forecast=[1.0, None, 3.0], display_name="Points")
    assert m.s.points.track_expr("forecast") is None


class TestATrackedRowReadByPosition:
    """``row[-1]`` / ``row[0]`` of a plan/actual row is that period of every
    track, one live cell per track (before, ``[-1]`` raised "Index -1 on
    a scalar Variable … use var.value", and ``[0]`` copied the whole row as
    numbers)."""

    @staticmethod
    def _model():
        m = mo.Model("m", display_name="M", tracks=mo.Tracks("plan", "actual"),
                     default_grain="month", default_start="2026-01", default_periods=3)
        with m:
            m.inputs = mo.MultiVariable("Inputs")
            with m.inputs as a:
                a.opening = mo.Variable(1000.0, display_name="Opening")
            m.pnl = mo.MultiVariable("P&L")
            with m.pnl as p:
                p.revenue = mo.Variable(plan=[100.0, 100.0, 100.0], actual=[110.0, 190.0, None],
                                        display_name="Revenue", regrain=mo.up("sum"))
                p.cash = mo.recurrence(start=a.opening, formula="{prev} + {r}", r=p.revenue,
                                       display_name="Cash").set_regrain(mo.frozen(values="last"))
                p.rent = mo.Variable(plan=1500.0, actual=[1500.0, 1500.0, None], display_name="Rent")
            m.summary = mo.MultiVariable("Summary")
            with m.summary as s:
                s.end_cash = m.pnl.cash[-1].set_display_name("End cash")
                s.first_revenue = m.pnl.revenue[0].set_display_name("First revenue")
                s.end_rent = m.pnl.rent[-1].set_display_name("End rent")
        return m

    def test_each_track_reads_its_own_period(self):
        m = self._model()
        cash = m.summary.end_cash.value
        # the recurrence starts at the opening: plan 1000 → 1100 → 1200,
        # actual 1000 → 1190 → 1190 (March not entered yet counts 0)
        assert cash["plan"] == 1200.0 and cash["actual"] == 1190.0
        assert m.summary.first_revenue.value["plan"] == 100.0
        assert m.summary.first_revenue.value["actual"] == 110.0
        rent = m.summary.end_rent.value
        # a scalar track is the same in every period; an actual not entered
        # yet reads as a blank cell does in a formula: 0
        assert rent["plan"] == 1500.0 and rent["actual"] == 0

    def test_the_book_refers_to_each_track_s_cell(self):
        m = self._model()
        pnl, summary = _book(m, "P&L"), _book(m, "Summary")
        prow = _rows(pnl)
        srow = _rows(summary)
        last_col = openpyxl.utils.get_column_letter(pnl.max_column)
        assert summary.cell(srow["End cash"], 2).value == f"='P&L'!{last_col}{prow['Cash']}"
        assert summary.cell(srow["End cash · actual"], 2).value == (
            f"='P&L'!{last_col}{prow['Cash · actual']}")
        assert summary.cell(srow["First revenue · actual"], 2).value == (
            f"='P&L'!C{prow['Revenue · actual']}")

    def test_a_period_past_the_window_is_an_index_error(self):
        m = self._model()
        try:
            m.pnl.revenue[5]
        except IndexError:
            pass
        else:
            raise AssertionError("a period past the window must not read")

    def test_a_scalar_row_says_to_read_it_by_name(self):
        m = self._model()
        try:
            m.inputs.opening[-1]
        except IndexError as exc:
            assert "read it as it is, by its name" in str(exc)
            assert "var.value" not in str(exc)
        else:
            raise AssertionError("an index past [0] on a scalar must refuse")
