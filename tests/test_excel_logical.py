# SPDX-License-Identifier: Apache-2.0
"""The workbook writes ``IFERROR``, ``ISERROR``, ``ISNUMBER``, ``XOR``,
``IFS`` and ``SWITCH`` as the spreadsheet's own functions — one call per
period over that period's cells, the names Excel added after 2007 stored
with the ``_xlfn.`` prefix — and, recalculated by LibreOffice, computes
the engine's values.
"""

from __future__ import annotations

import datetime as dt
import shutil
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook

import modeleon as mo
from modeleon.functions.logical import IFERROR, IFS, ISERROR, ISNUMBER, SWITCH, XOR


def _rows(ws):
    out = {}
    for row in ws.iter_rows():
        label = row[0].value
        if label:
            out[str(label)] = [c.value for c in row[1:]]
    return out


def _model():
    m = mo.Model("m", default_grain="year", default_start="2026", default_periods=3)
    with m:
        m.i = mo.MultiVariable("I")
        m.i.k = mo.Variable(7.0, display_name="k")
        m.i.d0 = mo.Variable(dt.date(2026, 1, 1), display_name="d0")
        m.x = mo.MultiVariable("X")
        m.x.a = mo.Variable([3.0, 8.0, 5.0], display_name="a")
        m.x.z = mo.Variable([1.0, 0.0, 2.0], display_name="z")
        m.x.q = (m.x.a / m.x.z).set_display_name("q")
        m.x.blank = mo.Variable([None, 5.0, None], display_name="blank")
        m.x.words = mo.Variable(["abc", "", "zz"], display_name="words")
        m.x.names = mo.Variable(["Base", "bull", "x"], display_name="names")
        m.x.flags = mo.Variable([True, False, True], display_name="flags")
        m.x.nums = mo.Variable([0, 2, -1], display_name="nums")
    return m


def _add(container, **rows):
    for name, var in rows.items():
        setattr(container, name, var.set_display_name(name))


class TestTheWorkbook:
    def test_each_function_is_written_as_the_spreadsheet_s_own(self, tmp_path):
        m = _model()
        with m:
            _add(
                m.x,
                iferr=IFERROR(m.x.q, 0.0),
                iferr_txt=IFERROR(m.x.q, "n/a"),
                iferr_row=IFERROR(m.x.q, [9.0, 8.0, 7.0]),
                iserr=ISERROR(m.x.q),
                isnum=ISNUMBER(m.x.a),
                xor=XOR(m.x.a > 4.0, True),
                ifs=IFS(m.x.a > 7.0, "big", m.x.a > 4.0, "mid", True, "small"),
                sw=SWITCH(m.x.names, "base", 0.05, "bull", 0.08, 0.03),
                nested=IFERROR(IFS(m.x.a > 100.0, 1.0), -1.0),
            )
            _add(
                m.i,
                ifs_k=IFS(m.i.k > 5.0, "hi", True, "lo"),
                sw_date=SWITCH(m.i.d0, 46023, "yes", "no"),
                isnum_date=ISNUMBER(m.i.d0),
            )
        path = tmp_path / "logical.xlsx"
        m.to_excel(path)
        book = load_workbook(path)
        x = _rows(book["X"])
        i = _rows(book["I"])
        assert x["iferr"] == ["=IFERROR(B4, 0.0)", "=IFERROR(C4, 0.0)", "=IFERROR(D4, 0.0)"]
        assert x["iferr_txt"][0] == '=IFERROR(B4, "n/a")'
        # A per-period list literal is this period's number.
        assert x["iferr_row"] == ["=IFERROR(B4, 9.0)", "=IFERROR(C4, 8.0)", "=IFERROR(D4, 7.0)"]
        assert x["iserr"] == ["=ISERROR(B4)", "=ISERROR(C4)", "=ISERROR(D4)"]
        assert x["isnum"][0] == "=ISNUMBER(B2)"
        assert x["xor"] == ["=_xlfn.XOR(B2 > 4.0, TRUE)", "=_xlfn.XOR(C2 > 4.0, TRUE)",
                            "=_xlfn.XOR(D2 > 4.0, TRUE)"]
        assert x["ifs"][0] == '=_xlfn.IFS(B2 > 7.0, "big", B2 > 4.0, "mid", TRUE, "small")'
        assert x["ifs"][2] == '=_xlfn.IFS(D2 > 7.0, "big", D2 > 4.0, "mid", TRUE, "small")'
        assert x["sw"][0] == '=_xlfn.SWITCH(B7, "base", 0.05, "bull", 0.08, 0.03)'
        assert x["sw"][1] == '=_xlfn.SWITCH(C7, "base", 0.05, "bull", 0.08, 0.03)'
        assert x["nested"][0] == "=IFERROR(_xlfn.IFS(B2 > 100.0, 1.0), -1.0)"
        assert i["ifs_k"][0] == '=_xlfn.IFS(B1 > 5.0, "hi", TRUE, "lo")'
        assert i["sw_date"][0] == '=_xlfn.SWITCH(B2, 46023, "yes", "no")'
        assert i["isnum_date"][0] == "=ISNUMBER(B2)"

    def test_the_cells_carry_the_engine_s_values(self, tmp_path):
        m = _model()
        with m:
            _add(
                m.x,
                iferr=IFERROR(m.x.q, 0.0),
                iserr=ISERROR(m.x.q),
                xor=XOR(m.x.a > 4.0, True),
                ifs=IFS(m.x.a > 7.0, "big", m.x.a > 4.0, "mid", True, "small"),
                ifs_na=IFS(m.x.a > 100.0, 1.0),
                sw=SWITCH(m.x.names, "base", 0.05, "bull", 0.08, 0.03),
            )
        path = tmp_path / "logical.xlsx"
        m.to_excel(path)
        x = _rows(load_workbook(path, data_only=True)["X"])
        assert x["iferr"] == [3.0, 0.0, 2.5] == m.x.iferr.value
        assert x["iserr"] == [False, True, False] == m.x.iserr.value
        assert x["xor"] == [True, False, False] == m.x.xor.value
        assert x["ifs"] == ["small", "big", "mid"] == m.x.ifs.value
        assert x["ifs_na"] == ["#N/A"] * 3 == m.x.ifs_na.value
        assert x["sw"] == [0.05, 0.08, 0.03] == m.x.sw.value


def _soffice():
    for candidate in ("soffice", "libreoffice"):
        found = shutil.which(candidate)
        if found:
            return found
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


def _agree(got, want) -> bool:
    if got is None and want == "":
        return True  # a formula's empty text reads back as an empty cell
    if isinstance(want, bool) or isinstance(got, bool):
        return got is want
    if isinstance(want, (int, float)) and isinstance(got, (int, float)):
        return round(float(got), 9) == round(float(want), 9)
    return got == want


@pytest.mark.slow
def test_the_workbook_recalculates_to_the_python_values(tmp_path):
    """Every function of the module, in the corners Excel and LibreOffice
    read alike: errors caught and named, blanks and text in conditions
    (the text TRUE and FALSE among them), a lazy branch, the ``#N/A`` of
    nothing matching, computed numbers matched to 15 significant digits,
    text matched regardless of case, a date by its serial number, an
    empty text typed into a line (an empty cell) and a formula's (text),
    a per-period and a scalar call."""
    soffice = _soffice()
    if soffice is None:
        pytest.skip("LibreOffice not installed")
    m = _model()
    with m:
        m.x.f = mo.Variable([0.1, 0.2, 0.3], display_name="f")
        m.x.tf = mo.Variable(["TRUE", "FALSE", "true"], display_name="tf")
        m.x.gaps = mo.Variable([True, "", False], display_name="gaps")
        m.x.fempty = IFS(m.x.a > 100.0, "x", True, "").set_display_name("fempty")
        m.i.big = mo.Variable(4503599627370496.0, display_name="big")
        _add(
            m.x,
            iferr=IFERROR(m.x.q, 0.0),
            iferr_txt=IFERROR(m.x.q, "n/a"),
            iferr_blank=IFERROR(m.x.blank, "x"),
            iferr_typed=IFERROR(m.x.words, "x"),
            iferr_fempty=IFERROR(m.x.fempty, "x"),
            iferr_fallback_blank=IFERROR(m.x.q, m.x.blank),
            iferr_nested=IFERROR(IFS(m.x.a > 100.0, 1.0), -1.0),
            iserr=ISERROR(m.x.q),
            iserr_blank=ISERROR(m.x.blank),
            isnum_q=ISNUMBER(m.x.q),
            isnum_words=ISNUMBER(m.x.words),
            isnum_blank=ISNUMBER(m.x.blank),
            xor_tf=XOR(m.x.a > 4.0, True),
            xor_three=XOR(m.x.a > 4.0, m.x.z > 0.0, True),
            xor_blank=XOR(m.x.blank, True),
            xor_gaps=XOR(m.x.gaps, True),
            xor_fempty=XOR(m.x.fempty, True),
            xor_num=XOR(m.x.nums),
            xor_err=XOR(m.x.q > 1.0, m.x.flags),
            xor_all_blank=XOR(m.x.blank),
            ifs_band=IFS(m.x.a > 7.0, "big", m.x.a > 4.0, "mid", True, "small"),
            ifs_na=IFS(m.x.a > 100.0, 1.0),
            ifs_lazy=IFS(m.x.z == 0, 0.0, True, m.x.q),
            ifs_err=IFS(m.x.q > 1.0, 1.0, True, 2.0),
            ifs_cond_after=IFS(True, 1.0, m.x.q > 1.0, 2.0),
            ifs_num=IFS(m.x.nums, "yes", True, "no"),
            ifs_words=IFS(m.x.words, 1.0, True, 2.0),
            ifs_tf=IFS(m.x.tf, 1.0, True, 2.0),
            ifs_fempty=IFS(m.x.fempty, 1.0, True, 2.0),
            ifs_blank=IFS(m.x.blank, 1.0, True, 2.0),
            ifs_pick_blank=IFS(True, m.x.blank),
            ifs_pick_typed=IFS(True, m.x.words),
            ifs_flags=IFS(m.x.flags, m.x.a, True, 0.0),
            sw_txt=SWITCH(m.x.names, "base", 0.05, "BULL", 0.08, 0.03),
            sw_na=SWITCH(m.x.a, 3, "three"),
            sw_num=SWITCH(m.x.a, 3, "three", 8, "eight", "other"),
            sw_blank_txt=SWITCH(m.x.blank, "", "empty", "other"),
            sw_typed=SWITCH(m.x.words, "", "empty", "other"),
            sw_fempty=SWITCH(m.x.fempty, "", "empty", "other"),
            sw_fempty0=SWITCH(m.x.fempty, 0, "zero", "other"),
            sw_money=SWITCH(m.x.f * 3, 0.3, "p3", 0.6, "p6", 0.9, "p9", "other"),
            sw_rate=SWITCH(m.x.f + 0.2, 0.3, "a", 0.4, "b", 0.5, "c", "other"),
            sw_apart=SWITCH(m.x.f + 1e-12, 0.1, "eq", 0.2, "eq", 0.3, "eq", "apart"),
            sw_err=SWITCH(m.x.q, 3, "three", "other"),
            sw_pick_blank=SWITCH(m.x.a, 3, m.x.blank, 8, m.x.blank, "other"),
            sw_series=SWITCH(2, 1, m.x.a, 2, m.x.z),
        )
        _add(
            m.i,
            ifs_k=IFS(m.i.k > 5.0, "hi", True, "lo"),
            sw_k=SWITCH(m.i.k, 7, "seven", 8, "eight"),
            sw_date=SWITCH(m.i.d0, 46023, "yes", "no"),
            sw_whole=SWITCH(m.i.big, 4503599627370497, "eq", "ne"),
            isnum_date=ISNUMBER(m.i.d0),
            ifs_true_lit=IFS("TRUE", 1.0, True, 2.0),
            ifs_false_lit=IFS("False", 1.0, True, 2.0),
        )
    assert m.x.ifs_words.value == ["#VALUE!", 2.0, "#VALUE!"]
    assert m.x.sw_err.value == ["three", "#DIV/0!", "other"]
    assert m.x.ifs_tf.value == [1.0, 2.0, 1.0] and m.x.ifs_fempty.value == ["#VALUE!"] * 3
    assert m.x.iferr_typed.value == ["abc", 0, "zz"] == m.x.ifs_pick_typed.value
    assert m.x.iferr_fempty.value == ["", "", ""] and m.x.xor_gaps.value == [False, True, True]
    assert m.x.sw_money.value == ["p3", "p6", "p9"] and m.x.sw_rate.value == ["a", "b", "c"]
    assert m.x.sw_apart.value == ["apart"] * 3 and m.i.sw_whole.value == "ne"
    assert m.x.sw_typed.value == ["other", "empty", "other"]
    assert m.x.sw_fempty0.value == ["other"] * 3
    path = tmp_path / "logical.xlsx"
    m.to_excel(path)
    subprocess.run(
        [soffice, f"-env:UserInstallation={(tmp_path / 'prof').as_uri()}",
         "--headless", "--convert-to", "xlsx", "--outdir",
         str(tmp_path / "calc"), str(path)],
        check=True, capture_output=True,
    )
    book = load_workbook(tmp_path / "calc" / "logical.xlsx", data_only=True)
    for sheet, container in (("X", m.x), ("I", m.i)):
        rows = _rows(book[sheet])
        for name, var in container._components.items():
            if not isinstance(var, mo.Variable) or var._expr is None:
                continue
            want = var.value if isinstance(var.value, list) else [var.value]
            got = rows[var.display_name][:len(want)]
            assert len(got) == len(want) and all(
                _agree(g, w) for g, w in zip(got, want)
            ), f"{sheet}.{name}: LibreOffice {got!r}, the engine {want!r}"
