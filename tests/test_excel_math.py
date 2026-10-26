# SPDX-License-Identifier: Apache-2.0
"""The rounding family, the powers and logarithms, ``SIGN`` and
``PRODUCT`` in the workbook.

An element-wise function writes one call per period cell — ``=ROUNDUP(B2,
1)``, ``=CEILING(B2, B3)``, ``=LOG(B4)`` with the base left out when none
was given — a scalar one cell, and ``PRODUCT`` writes a row as a range,
``=PRODUCT(X!B4:G4)``. Every name is the classic one, stored without a
prefix, and every call renders natively. The slow test recalculates a
workbook with every function in LibreOffice and reads the engine's
numbers back, a whole number to the unit; the cells where LibreOffice's
answer is not the engine's are pinned, and the parity case over them is
an expected failure.
"""

import datetime as dt
import shutil
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook

import modeleon as mo
from modeleon.functions._helpers import is_error
from modeleon.functions.mathfn import (
    CEILING,
    EXP,
    FLOOR,
    LN,
    LOG,
    LOG10,
    POWER,
    PRODUCT,
    ROUNDDOWN,
    ROUNDUP,
    SIGN,
    SQRT,
    TRUNC,
)

NAMES = (
    "ROUNDUP", "ROUNDDOWN", "TRUNC", "CEILING", "FLOOR", "POWER", "SQRT", "EXP",
    "LN", "LOG", "LOG10", "SIGN", "PRODUCT",
)


def _rows(ws):
    out = {}
    for row in ws.iter_rows():
        label = row[0].value
        if label:
            out[str(label)] = [c.value for c in row[1:]]
    return out


def _book(model, tmp_path, sheet):
    path = tmp_path / "book.xlsx"
    model.to_excel(path)
    return _rows(load_workbook(path)[sheet])


def _model():
    m = mo.Model("m", default_grain="year", default_start="2026", default_periods=6)
    with m:
        m.i = mo.MultiVariable("I")
        m.i.p = mo.Variable(2.5, display_name="p")
        m.i.q = mo.Variable(-2.0, display_name="q")
        m.i.z = mo.Variable(0.0, display_name="z")
        m.i.eight = mo.Variable(8.0, display_name="eight")
        m.i.m8 = mo.Variable(-8.0, display_name="m8")
        m.i.t = mo.Variable(True, display_name="t")
        m.i.four = mo.Variable("4", display_name="four")
        m.i.word = mo.Variable("abc", display_name="word")
        m.x = mo.MultiVariable("X")
        m.x.a = mo.Variable([2.5, -2.5, 0.0, 3.14159, 1234.5678, -0.125], display_name="a")
        m.x.b = mo.Variable([0.1, 2.0, 0.5, 100.0, 3.0, 0.25], display_name="b")
        m.x.c = mo.Variable([16.0, 2.0, 1.0, 100000.0, 0.5, 1000.0], display_name="c")
        m.x.d = mo.Variable([0.0, 1.0, -1.0, 2.5, 10.0, -20.0], display_name="d")
    return m


class TestEachFunctionWritesOneCallPerPeriod:
    @pytest.mark.parametrize("build, first, second", [
        (lambda x, i: ROUNDUP(x.a, 1), "=ROUNDUP(B2, 1)", "=ROUNDUP(C2, 1)"),
        (lambda x, i: ROUNDDOWN(x.a, 1), "=ROUNDDOWN(B2, 1)", "=ROUNDDOWN(C2, 1)"),
        (lambda x, i: TRUNC(x.a), "=TRUNC(B2)", "=TRUNC(C2)"),
        (lambda x, i: TRUNC(x.a, 2), "=TRUNC(B2, 2)", "=TRUNC(C2, 2)"),
        (lambda x, i: CEILING(x.a, 0.5), "=CEILING(B2, 0.5)", "=CEILING(C2, 0.5)"),
        (lambda x, i: CEILING(x.a, x.b), "=CEILING(B2, B3)", "=CEILING(C2, C3)"),
        (lambda x, i: FLOOR(x.a, 0.5), "=FLOOR(B2, 0.5)", "=FLOOR(C2, 0.5)"),
        (lambda x, i: FLOOR(x.a, i.p), "=FLOOR(B2, I!B1)", "=FLOOR(C2, I!B1)"),
        (lambda x, i: POWER(x.a, 2), "=POWER(B2, 2)", "=POWER(C2, 2)"),
        (lambda x, i: POWER(x.c, x.d), "=POWER(B4, B5)", "=POWER(C4, C5)"),
        (lambda x, i: SQRT(x.c), "=SQRT(B4)", "=SQRT(C4)"),
        (lambda x, i: EXP(x.d), "=EXP(B5)", "=EXP(C5)"),
        (lambda x, i: LN(x.c), "=LN(B4)", "=LN(C4)"),
        (lambda x, i: LOG(x.c), "=LOG(B4)", "=LOG(C4)"),
        (lambda x, i: LOG(x.c, 2), "=LOG(B4, 2)", "=LOG(C4, 2)"),
        (lambda x, i: LOG(x.c, i.eight), "=LOG(B4, I!B4)", "=LOG(C4, I!B4)"),
        (lambda x, i: LOG10(x.c), "=LOG10(B4)", "=LOG10(C4)"),
        (lambda x, i: SIGN(x.a), "=SIGN(B2)", "=SIGN(C2)"),
    ])
    def test_the_call_over_this_period_s_cells(self, tmp_path, build, first, second):
        m = _model()
        with m:
            m.x.r = build(m.x, m.i).set_display_name("r")
        r = _book(m, tmp_path, "X")["r"][:6]
        assert r[:2] == [first, second], r
        assert all(isinstance(cell, str) and cell.startswith("=") for cell in r), r

    def test_a_scalar_is_one_cell(self, tmp_path):
        m = _model()
        with m:
            m.i.ru = ROUNDUP(m.i.p, 0).set_display_name("ru")
            m.i.pw = POWER(m.i.q, 0.5).set_display_name("pw")
            m.i.lg = LOG(m.i.eight, 2).set_display_name("lg")
            m.i.ce = CEILING(m.i.p, m.i.q).set_display_name("ce")
            m.i.sq = SQRT(m.i.eight * 2).set_display_name("sq")
        i = _book(m, tmp_path, "I")
        assert i["ru"][0] == "=ROUNDUP(B1, 0)", i["ru"]
        assert i["pw"][0] == "=POWER(B2, 0.5)", i["pw"]
        assert i["lg"][0] == "=LOG(B4, 2)", i["lg"]
        assert i["ce"][0] == "=CEILING(B1, B2)", i["ce"]
        assert i["sq"][0] == "=SQRT(B4 * 2)", i["sq"]


class TestProductWritesARange:
    def test_a_row_is_a_range_and_a_scalar_a_cell(self, tmp_path):
        m = _model()
        with m:
            m.i.row = PRODUCT(m.x.c).set_display_name("row")
            m.i.cells = PRODUCT(m.i.p, m.i.q, 3).set_display_name("cells")
            m.i.rows = (PRODUCT(m.x.b) * PRODUCT(m.x.d)).set_display_name("rows")
        i = _book(m, tmp_path, "I")
        assert i["row"][0] == "=PRODUCT(X!B4:G4)", i["row"]
        assert i["cells"][0] == "=PRODUCT(B1, B2, 3)", i["cells"]
        assert i["rows"][0] == "=PRODUCT(X!B3:G3) * PRODUCT(X!B5:G5)", i["rows"]

    def test_several_series_are_one_call_per_period(self, tmp_path):
        # One call per period, as ``b * d`` and ``c * 2`` are.
        m = _model()
        with m:
            m.x.rows = PRODUCT(m.x.b, m.x.d).set_display_name("rows")
            m.x.mixed = PRODUCT(m.x.c, 2).set_display_name("mixed")
        x = _book(m, tmp_path, "X")
        assert x["rows"][:2] == ["=PRODUCT(B3, B5)", "=PRODUCT(C3, C5)"], x["rows"]
        assert x["mixed"][:2] == ["=PRODUCT(B4, 2)", "=PRODUCT(C4, 2)"], x["mixed"]
        assert m.x.rows.value == [0.0, 2.0, -0.5, 250.0, 30.0, -5.0]


def _every_function(m):
    """Every function above in ``m``, label -> line: the series on
    X, the scalars and the reductions on I."""
    with m:
        m.x.tenth = mo.Variable([0.1, 4.35, 1.1, 2.675, 0.7, 0.3], display_name="tenth")
        m.x.mult = mo.Variable([3, 100, 1, 100, 3, 1], display_name="mult")
        m.x.blank = mo.Variable([None] * 6, display_name="blank")
        # Held as 0.30000000000000004, 434.99999999999994, …; shown as 0.3, 435, ….
        m.x.held = (m.x.tenth * m.x.mult).set_display_name("held")
        m.x.e = (m.x.b / m.x.d).set_display_name("e")        # #DIV/0! in the first period
        m.x.when = mo.EOMONTH(mo.time.start, 0).set_display_name("when")     # a date row
        m.i.day = mo.Variable(dt.date(2024, 2, 29), display_name="day")
        m.i.noon = mo.Variable(dt.datetime(2024, 2, 29, 12, 0), display_name="noon")
        m.i.big = mo.Variable(1e300, display_name="big")
        m.i.tiny = mo.Variable(1e-300, display_name="tiny")
        m.i.whole16 = mo.Variable(1000000000000001, display_name="whole16")
        # A cell is written to 15 digits; the half is computed in the sheet.
        m.i.half16 = (m.i.whole16 - 0.5).set_display_name("half16")
        lines = {
            "roundup": ROUNDUP(m.x.a, 1),
            "rounddown": ROUNDDOWN(m.x.a, 1),
            "trunc": TRUNC(m.x.a),
            "trunc2": TRUNC(m.x.a, 2),
            "ceil": CEILING(m.x.a, 0.5),
            "floor": FLOOR(m.x.a, 0.5),
            "ceil_b": CEILING(m.x.a, m.x.b),
            "floor_b": FLOOR(m.x.a, m.x.b),
            "power": POWER(m.x.a, 2),
            "power_d": POWER(m.x.c, m.x.d),
            "sqrt": SQRT(m.x.c),
            "exp": EXP(m.x.d),
            "ln": LN(m.x.c),
            "log": LOG(m.x.c),
            "log2": LOG(m.x.c, 2),
            "log10": LOG10(m.x.c),
            "sign": SIGN(m.x.a),
            "ru_shown": ROUNDUP(m.x.held, 1),
            "rd_shown": ROUNDDOWN(m.x.held, 0),
            "tr_shown": TRUNC(m.x.held, 1),
            "ce_shown": CEILING(m.x.held, 0.1),
            "fl_shown": FLOOR(m.x.held, 0.1),
            "sqrt_a": SQRT(m.x.a),
            "ln_d": LN(m.x.d),
            "exp_big": EXP(m.x.c),
            "ru_err": ROUNDUP(m.x.e, 1),
            "sq_err": SQRT(m.x.e),
            "pw_err": POWER(m.x.e, 2),
            "c0": CEILING(m.i.p, m.i.z),
            "f0": FLOOR(m.i.p, m.i.z),
            "p00": POWER(m.i.z, m.i.z),
            "p0m1": POWER(m.i.z, -1),
            "pneg": POWER(m.i.q, 0.5),
            "l1": LOG(m.i.eight, 1),
            "lneg": LOG(m.i.eight, m.i.q),
            "c_nn": CEILING(m.i.q * 1.25, m.i.q),
            "f_nn": FLOOR(m.i.q * 1.25, m.i.q),
            "c_pn": CEILING(m.i.p, m.i.q),
            "f_pn": FLOOR(m.i.p, m.i.q),
            "c_np": CEILING(m.i.q * 1.25, 2),
            "f_np": FLOOR(m.i.q * 1.25, 2),
            "ru_s": ROUNDUP(m.i.p, 0),
            "sq_s": SQRT(m.i.eight),
            "lg_s": LOG(m.i.eight, 2),
            "odd_root": POWER(m.i.m8, 1 / 3),
            "odd_fifth": POWER(m.i.m8 * 4, 0.2),
            "odd_neg": POWER(m.i.m8, -1 / 3),
            "even_root": POWER(m.i.m8, 0.4),
            "underflow": POWER(m.i.eight, -400),
            "ru_big": ROUNDUP(m.i.p, -320),
            "rd_big": ROUNDDOWN(m.i.p, -320),
            "t_ru": ROUNDUP(m.i.t, 0),
            "t_pw": POWER(m.i.t, 2),
            "four_ru": ROUNDUP(m.i.four, 0),
            "four_sq": SQRT(m.i.four),
            "four_pw": POWER(2, m.i.four),
            "word_ru": ROUNDUP(m.i.word, 0),
            "word_sg": SIGN(m.i.word),
            "prod_tb": PRODUCT(m.i.t, m.i.p),
            "prod_blank": PRODUCT(m.x.blank),
            "prod_c": PRODUCT(m.x.c),
            "prod_pq": PRODUCT(m.i.p, m.i.q, 3),
            "prod_rows": PRODUCT(m.x.b, m.x.d),
            "fl_date": FLOOR(m.x.when, 7),
            "ce_date": CEILING(m.x.when, 7),
            "ru_date": ROUNDUP(m.x.when, -1),
            "sg_date": SIGN(m.x.when),
            "tr_noon": TRUNC(m.i.noon),
            "sq_day": SQRT(m.i.day),
            "prod_day": PRODUCT(m.i.day, 2),
            "f00": FLOOR(m.i.z, m.i.z),
            "c_past": CEILING(m.i.big, m.i.tiny),
            "f_past": FLOOR(m.i.big, m.i.tiny),
            "c_below": CEILING(m.i.tiny, m.i.big),
            "c_whole16": CEILING(m.i.whole16, 1),
            "c_half16": CEILING(m.i.half16, 1),
            "f_half16": FLOOR(m.i.half16, 1),
        }
        for label, line in lines.items():
            sheet = m.x if line.var_type == 'list' else m.i
            setattr(sheet, label, line.set_display_name(label))
    return lines


def test_every_name_is_classic_and_renders_natively(tmp_path):
    m = _model()
    lines = _every_function(m)
    assert mo.check_compat(m, "excel") == []
    path = tmp_path / "book.xlsx"
    m.to_excel(path)
    formulas = [
        cell.value
        for ws in load_workbook(path).worksheets
        for row in ws.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    ]
    assert len(formulas) >= len(lines)
    assert not [f for f in formulas if "_xlfn." in f]
    for name in NAMES:
        assert any(f.startswith(f"={name}(") for f in formulas), name


# ─── LibreOffice recalculates the workbook ──────────────────────────


def _soffice():
    for candidate in ("soffice", "libreoffice"):
        found = shutil.which(candidate)
        if found:
            return found
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


def _recalculated(m, tmp_path):
    """The workbook of ``m`` recalculated by LibreOffice: ``(X rows, I rows)``
    of values."""
    soffice = _soffice()
    if soffice is None:
        pytest.skip("LibreOffice not installed")
    path = tmp_path / "math.xlsx"
    m.to_excel(path)
    subprocess.run(
        [soffice, f"-env:UserInstallation={(tmp_path / 'prof').as_uri()}",
         "--headless", "--convert-to", "xlsx", "--outdir",
         str(tmp_path / "calc"), str(path)],
        check=True, capture_output=True, timeout=300,
    )
    book = load_workbook(tmp_path / "calc" / "math.xlsx", data_only=True)
    return _rows(book["X"]), _rows(book["I"])


def _number(cell):
    """A cell's number: one LibreOffice formats as a date (it gives a
    product of a date and a number the date's format) is read back as the
    serial it holds."""
    if isinstance(cell, dt.datetime):
        days = (cell - dt.datetime(1899, 12, 30)) / dt.timedelta(days=1)
        return int(days) if days.is_integer() else days
    return cell


def _cells(lines, x, i):
    """Every function cell as ``((label, period), engine value, LibreOffice
    value)``."""
    for label, line in lines.items():
        if line.var_type == 'list':
            for period, (ours, theirs) in enumerate(zip(line.value, x[label][:6])):
                yield (label, period), ours, _number(theirs)
        else:
            yield (label, 0), line.value, _number(i[label][0])


#: Where LibreOffice's answer is not the engine's: what LibreOffice 25.2
#: writes instead, pinned by cell. Most are where it departs from Excel,
#: whose answers (Microsoft's documentation, its known behaviour) the
#: engine computes: LibreOffice's "invalid argument" error is written as
#: #VALUE! where Excel says #NUM! (a root or a logarithm outside its
#: domain, a positive number with a negative significance) or #DIV/0!
#: (FLOOR of a number by 0, a logarithm to base 1); 0 ^ -1 is its #NUM!
#: where Excel divides by zero; it defines 0 ^ 0 as 1 where Excel says
#: #NUM!. Two are the engine's own choice where Excel's answer is neither
#: documented nor checked: a power too small for a cell is 0 (as EXP(-1000)
#: is 0 in both programs) where LibreOffice says #NUM!, and ROUNDUP to a
#: multiple no cell can hold (1e320) is #NUM! where LibreOffice says 0.
_LIBREOFFICE_INSTEAD = {
    ("sqrt_a", 1): "#VALUE!", ("sqrt_a", 5): "#VALUE!",
    ("ln_d", 0): "#VALUE!", ("ln_d", 2): "#VALUE!", ("ln_d", 5): "#VALUE!",
    ("sq_err", 2): "#VALUE!", ("sq_err", 5): "#VALUE!",
    ("c_pn", 0): "#VALUE!", ("f_pn", 0): "#VALUE!", ("lneg", 0): "#VALUE!",
    ("f0", 0): "#VALUE!", ("l1", 0): "#VALUE!",
    ("p0m1", 0): "#NUM!",
    ("p00", 0): 1,
    ("underflow", 0): "#NUM!",
    ("ru_big", 0): 0,
}


@pytest.mark.slow
def test_the_workbook_recalculates_to_the_python_values(tmp_path):
    m = _model()
    lines = _every_function(m)
    x, i = _recalculated(m, tmp_path)
    for cell, ours, theirs in _cells(lines, x, i):
        expected = _LIBREOFFICE_INSTEAD.get(cell, ours)
        if is_error(expected):
            assert theirs == expected, (cell, ours, theirs)
        elif float(expected).is_integer() and abs(expected) < 2 ** 53:
            assert theirs == expected, (cell, ours, theirs)        # a whole number to the unit
        else:
            assert theirs == pytest.approx(expected, rel=1e-9, abs=1e-12), (cell, ours, theirs)
    # PRODUCT over two rows is one call per period: 33 rows.
    assert sum(1 for _ in _cells(lines, x, i)) == 33 * 6 + 44


@pytest.mark.slow
@pytest.mark.xfail(strict=True, reason=(
    "LibreOffice names a domain error differently from Excel (#VALUE! for its "
    "'invalid argument' where Excel says #NUM! or #DIV/0!) and defines 0 ^ 0 as 1, "
    "where the engine computes Excel's answers; where Excel's answer is unchecked, the "
    "engine calls a power too small for a cell 0 and a ROUNDUP past the largest number "
    "#NUM!, LibreOffice #NUM! and 0 — see _LIBREOFFICE_INSTEAD."
))
def test_the_errors_libreoffice_names_differently(tmp_path):
    m = _model()
    lines = _every_function(m)
    x, i = _recalculated(m, tmp_path)
    for cell, ours, theirs in _cells(lines, x, i):
        if cell in _LIBREOFFICE_INSTEAD:
            assert theirs == ours, (cell, ours, theirs)
