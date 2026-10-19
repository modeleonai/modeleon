# SPDX-License-Identifier: Apache-2.0
"""``ExcelView(meta={'unit': …})`` has three states, told apart once.

A caption (or ``True``) puts the units in a column of their own — and,
under the captions of a records table or a board, in a unit row. Unset
suffixes each unit into its row label and table caption ("Revenue
(USD m)"). ``False`` shows the unit text NOWHERE: no column, no suffix,
no unit row — in the workbook and the notebook repr alike — while the
number formats keyed by unit still apply: the unit stays on the line,
only its text is hidden. The decision is read from one resolved value
(``ResolvedView.meta_unit_col``, carried in the layout's meta triple)
through one predicate, ``units_hidden``.
"""

import re

from openpyxl import load_workbook

import modeleon as mo
from modeleon.compile.excel.layout import LayoutEngine
from modeleon.compile.excel.view import units_hidden

COLUMN = {"unit": "Unit"}
SUFFIX: dict = {}
HIDDEN = {"unit": False}

FORMATS = {"%": "0.0%", "USD m": "#,##0.0"}


class Comp(mo.MultiVariableClass):
    def compute(self, ev, ebitda):
        self.ev = mo.Variable(ev, display_name="EV", unit="USD m")
        self.ebitda = mo.Variable(ebitda, display_name="EBITDA", unit="USD m")
        self.deals = mo.Variable(ebitda * 2, display_name="Deals")


def _book(meta, **view):
    m = mo.Model("m", display_name="M", default_grain="year", default_start="2026",
                 default_periods=2)
    m.default_excel_view = mo.ExcelView(meta=meta, formats=FORMATS, **view)
    with m:
        m.pnl = mo.MultiVariable("PnL", excel_props={"tab": True})
        with m.pnl as p:
            p.revenue = mo.Variable([100, 110], display_name="Revenue", unit="USD m")
            p.margin = mo.Variable(0.2, display_name="Margin", unit="%")
            p.heads = mo.Variable([5, 6], display_name="Heads")
        m.comps = mo.MultiVariable("Comps", excel_props={"tab": True, "orient": "records"})
        with m.comps as c:
            c.a = Comp(ev=1200, ebitda=100, display_name="Alpha")
            c.b = Comp(ev=900, ebitda=90, display_name="Beta")
        m.kpis = mo.MultiVariable("KPIs", excel_props={"tab": True, "orient": "columns"})
        with m.kpis as k:
            k.irr = mo.Variable(0.2, display_name="IRR", unit="%")
            k.moic = mo.Variable(2.5, display_name="MOIC", unit="x")
            k.deals = mo.Variable(3, display_name="Deals")
    return m


def _sheets(model, tmp_path):
    path = tmp_path / "t.xlsx"
    model.to_excel(str(path))
    return load_workbook(str(path))


def _rows(ws, n, cols="ABCDE"):
    return [[ws[f"{c}{r}"].value for c in cols] for r in range(1, n + 1)]


def _html_rows(html):
    out = []
    for tr in re.findall(r"<tr>(.*?)</tr>", html, re.S):
        cells = re.findall(r"<t[hd][^>]*>(.*?)</t[hd]>", tr, re.S)
        out.append([re.sub(r"<[^>]+>", "", c).strip() for c in cells])
    return out


def _texts(ws):
    return {str(c.value) for row in ws.iter_rows() for c in row if isinstance(c.value, str)}


class TestRows:
    def test_a_caption_puts_the_unit_in_a_column(self, tmp_path):
        ws = _sheets(_book(COLUMN), tmp_path)["PnL"]
        assert _rows(ws, 4) == [
            [None, "Unit", None, "2026", "2027"],
            ["Revenue", "USD m", None, 100, 110],
            ["Margin", "%", 0.2, None, None],
            ["Heads", None, None, 5, 6],
        ]

    def test_unset_suffixes_the_unit_into_the_label(self, tmp_path):
        ws = _sheets(_book(SUFFIX), tmp_path)["PnL"]
        assert _rows(ws, 4, "ABCD") == [
            [None, None, "2026", "2027"],
            ["Revenue (USD m)", None, 100, 110],
            ["Margin (%)", 0.2, None, None],
            ["Heads", None, 5, 6],
        ]

    def test_false_shows_the_unit_nowhere(self, tmp_path):
        ws = _sheets(_book(HIDDEN), tmp_path)["PnL"]
        assert _rows(ws, 4, "ABCD") == [
            [None, None, "2026", "2027"],
            ["Revenue", None, 100, 110],
            ["Margin", 0.2, None, None],
            ["Heads", None, 5, 6],
        ]
        assert not {"USD m", "%"} & _texts(ws)

    def test_the_number_format_keyed_by_the_unit_still_applies(self, tmp_path):
        ws = _sheets(_book(HIDDEN), tmp_path)["PnL"]
        assert ws["B3"].number_format == "0.0%"
        assert ws["C2"].number_format == "#,##0.0"


class TestARecordsTable:
    def test_a_caption_puts_the_units_in_a_row_under_the_captions(self, tmp_path):
        ws = _sheets(_book(COLUMN), tmp_path)["Comps"]
        assert _rows(ws, 4, "ABCDEF") == [
            [None, None, "EV", "EBITDA", "Deals", None],
            [None, None, "USD m", "USD m", None, None],
            ["Alpha", None, 1200, 100, 200, None],
            ["Beta", None, 900, 90, 180, None],
        ]

    def test_unset_suffixes_the_unit_into_the_caption(self, tmp_path):
        ws = _sheets(_book(SUFFIX), tmp_path)["Comps"]
        assert _rows(ws, 3, "ABCD") == [
            [None, "EV (USD m)", "EBITDA (USD m)", "Deals"],
            ["Alpha", 1200, 100, 200],
            ["Beta", 900, 90, 180],
        ]

    def test_false_shows_the_unit_nowhere(self, tmp_path):
        ws = _sheets(_book(HIDDEN), tmp_path)["Comps"]
        assert _rows(ws, 3, "ABCD") == [
            [None, "EV", "EBITDA", "Deals"],
            ["Alpha", 1200, 100, 200],
            ["Beta", 900, 90, 180],
        ]
        assert "USD m" not in _texts(ws)
        assert ws["B2"].number_format == "#,##0.0"


class TestAScalarsAcrossBoard:
    def test_a_caption_puts_the_units_in_a_row_under_the_captions(self, tmp_path):
        ws = _sheets(_book(COLUMN), tmp_path)["KPIs"]
        assert _rows(ws, 4, "ABCDE")[1:] == [
            [None, None, "IRR", "MOIC", "Deals"],
            [None, None, "%", "x", None],
            [None, None, 0.2, 2.5, 3],
        ]

    def test_unset_suffixes_the_unit_into_the_caption(self, tmp_path):
        ws = _sheets(_book(SUFFIX), tmp_path)["KPIs"]
        assert _rows(ws, 3, "ABCD")[1:] == [
            [None, "IRR (%)", "MOIC (x)", "Deals"],
            [None, 0.2, 2.5, 3],
        ]

    def test_false_shows_the_unit_nowhere(self, tmp_path):
        ws = _sheets(_book(HIDDEN), tmp_path)["KPIs"]
        assert _rows(ws, 3, "ABCD")[1:] == [
            [None, "IRR", "MOIC", "Deals"],
            [None, 0.2, 2.5, 3],
        ]
        assert not {"%", "x"} & _texts(ws)
        assert ws["B3"].number_format == "0.0%"


class TestTheNotebookRepr:
    @staticmethod
    def _label(meta, name):
        rows = _html_rows(_book(meta).pnl._repr_html_())
        return next(r[1] for r in rows if r[1].startswith(name))

    def test_unset_suffixes_the_unit_into_the_label(self):
        assert self._label(SUFFIX, "Revenue") == "Revenue (USD m)"
        assert self._label(SUFFIX, "Margin") == "Margin (%)"

    def test_a_caption_suffixes_too_the_repr_has_no_unit_column(self):
        assert self._label(COLUMN, "Revenue") == "Revenue (USD m)"

    def test_false_shows_the_unit_nowhere(self):
        assert self._label(HIDDEN, "Revenue") == "Revenue"
        assert self._label(HIDDEN, "Margin") == "Margin"
        assert self._label(HIDDEN, "Heads") == "Heads"


class TestTheCascade:
    def test_a_sheet_hides_its_units_under_a_book_with_a_unit_column(self, tmp_path):
        m = _book(COLUMN)
        m.kpis.default_excel_view = mo.ExcelView(meta=HIDDEN)
        book = _sheets(m, tmp_path)
        assert _rows(book["KPIs"], 3, "ABCD")[1:] == [
            [None, "IRR", "MOIC", "Deals"],
            [None, 0.2, 2.5, 3],
        ]
        # The other sheets keep the column and the unit row.
        assert _rows(book["Comps"], 2, "ABCD")[1] == [None, None, "USD m", "USD m"]
        assert _rows(book["PnL"], 2, "AB")[1] == ["Revenue", "USD m"]

    def test_the_layout_carries_the_three_states_apart(self):
        m = _book(COLUMN)
        m.kpis.default_excel_view = mo.ExcelView(meta=HIDDEN)
        m.comps.default_excel_view = mo.ExcelView(meta={"unit": None})
        engine = LayoutEngine([m.pnl, m.comps, m.kpis], flatten_nested_sheets=True,
                              header_rows=True)
        engine.compute_addresses()
        assert engine.meta_by_sheet["PnL"][1] == "Unit"
        assert engine.meta_by_sheet["Comps"][1] == "Unit"    # None inherits
        assert engine.meta_by_sheet["KPIs"][1] is False
        plain = LayoutEngine([_book(SUFFIX).pnl], flatten_nested_sheets=True, header_rows=True)
        plain.compute_addresses()
        assert plain.meta_by_sheet["PnL"][1] is None


class TestATransposedSheet:
    @staticmethod
    def _ws(tmp_path, meta):
        m = mo.Model("m", display_name="M")
        m.default_excel_view = mo.ExcelView(orient="down", meta=meta)
        m.price = mo.Variable([10, 11], display_name="Price", unit="USD")
        m.volume = mo.Variable([100, 120], display_name="Volume")
        book = _sheets(m, tmp_path)
        return book[book.sheetnames[0]]

    def test_a_caption_reserves_no_column_down_the_unit_is_suffixed(self, tmp_path):
        ws = self._ws(tmp_path, COLUMN)
        assert [ws["A1"].value, ws["B1"].value] == ["Price (USD)", "Volume"]

    def test_false_shows_the_unit_nowhere(self, tmp_path):
        ws = self._ws(tmp_path, HIDDEN)
        assert [ws["A1"].value, ws["B1"].value] == ["Price", "Volume"]


def test_units_hidden_is_the_one_predicate():
    assert units_hidden(False) is True
    for shown in (None, True, "Unit", ""):
        assert units_hidden(shown) is False
