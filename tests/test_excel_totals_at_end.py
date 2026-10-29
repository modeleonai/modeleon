# SPDX-License-Identifier: Apache-2.0
"""Total columns after the last period — ``timeline={'totals': [...],
'totals_at': 'end'}``.

Between the periods ("Jan Feb Mar (Q1) Apr …") a line's period cells are
no run, so a function that reads a whole row as one range — ``MATCH``,
``INDEX``, ``LARGE``, ``COUNTIF`` … — cannot be written over it. With
every total after the last period ("Jan … Dec (Q1) … (Q4) (year)") every
line's periods stay one run and those functions stay live; the totals keep
their own formulas, the same text at other columns, and the header names
each total in full ("Total Q1 2026"), no tier above it saying its year.
The kinds stand finest first whatever order they are declared in, in both
placements. Unset, the totals stay between the periods.
"""

import os
import re
import tempfile
import warnings

import openpyxl
import pytest
from openpyxl.utils import column_index_from_string, get_column_letter

import modeleon as mo
from modeleon.compile.compat import check_compat
from modeleon.compile.excel.totals import NATIVE, build_totals_plan, header_rows
from modeleon.compile.excel.view import resolve_excel_view

#: 24 months of a count: the peak (26) in the 22nd month, where the
#: readers below point.
COUNTS = [12.0, 17.0, 9.0, 21.0, 14.0, 8.0, 19.0, 23.0, 11.0, 16.0, 7.0, 20.0,
          13.0, 25.0, 10.0, 18.0, 22.0, 6.0, 15.0, 24.0, 5.0, 26.0, 4.0, 3.0]


def _book(model):
    path = os.path.join(tempfile.mkdtemp(), "t.xlsx")
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        model.to_excel(path)
    return openpyxl.load_workbook(path)


def _timeline(kinds, at):
    out = {"totals": list(kinds)}
    if at is not None:
        out["totals_at"] = at
    return out


def _merges(ws):
    return sorted(str(r) for r in ws.merged_cells.ranges)


def _overlapping(ws):
    """Every pair of merged ranges that share a cell — a file Excel repairs."""
    ranges = list(ws.merged_cells.ranges)
    return [(str(a), str(b)) for i, a in enumerate(ranges) for b in ranges[i + 1:]
            if not (a.max_row < b.min_row or b.max_row < a.min_row
                    or a.max_col < b.min_col or b.max_col < a.min_col)]


def _level(ws, letter):
    dim = ws.column_dimensions.get(letter)
    return dim.outline_level if dim is not None else 0


def _readers(at="end", kinds=("quarter", "year")):
    """A monthly count on Data and every way of reading the whole row: on
    another sheet (Summary) and on its own (Data's constants column)."""
    m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=24,
                 default_excel_view=mo.ExcelView(timeline=_timeline(kinds, at)))
    with m:
        m.data = mo.MultiVariable("Data")
        with m.data as d:
            d.customers = mo.Variable(COUNTS, display_name="Customers", regrain=mo.up("sum"))
            d.price = mo.Variable([10.0 + i % 4 for i in range(24)], display_name="Price",
                                  regrain=mo.up("mean"))
            d.revenue = d.customers * d.price
            d.peak_month = mo.MATCH(mo.MAX(d.customers), d.customers, 0)
        m.summary = mo.MultiVariable("Summary")
        with m.summary as s:
            c = m.data.customers
            s.total = mo.SUM(c)
            s.peak = mo.MAX(c)
            s.peak_month = mo.MATCH(mo.MAX(c), c, 0)
            s.third = mo.INDEX(c, 3)
            s.big = mo.COUNTIF(c, ">15")
            s.big_sum = mo.SUMIFS(c, c, ">15")
            s.second = mo.LARGE(c, 2)
            s.first_half = mo.SUM(c[0:6])
            s.last = c[-1]
    return m


# ═════════════════════════════════════════════════════════════════════
# The plan
# ═════════════════════════════════════════════════════════════════════


class TestThePlan:
    def test_a_positional_call_still_places_the_totals_between(self):
        plan = build_totals_plan(("quarter", "year"), "month", "2026-01", 6, 3)
        assert plan.at == "between"
        assert plan.columns == [
            (NATIVE, 0), (NATIVE, 1), (NATIVE, 2), ("quarter", 0),
            (NATIVE, 3), (NATIVE, 4), (NATIVE, 5), ("quarter", 1), ("year", 0),
        ]
        # The stride, positional too: a group of two columns per key.
        strided = build_totals_plan(("quarter",), "month", "2026-01", 6, 3, 2)
        assert strided.columns[3] == ("quarter", 0)
        assert [strided.col_of[k] for k in strided.columns] == list(range(3, 19, 2))

    def test_at_the_end_the_periods_come_first_then_a_block_per_kind(self):
        plan = build_totals_plan(("quarter", "year"), "month", "2026-01", 24, 3, at="end")
        assert plan.at == "end"
        assert plan.columns == (
            [(NATIVE, i) for i in range(24)]
            + [("quarter", bi) for bi in range(8)]
            + [("year", 0), ("year", 1)]
        )
        assert [plan.col_of[k] for k in plan.columns] == list(range(3, 37))
        assert plan.depth == 3 and plan.n == 24

    def test_the_end_block_keeps_the_stride(self):
        plan = build_totals_plan(("quarter",), "month", "2026-01", 6, 2, stride=4, at="end")
        assert [plan.native_col(i) for i in range(6)] == [2, 6, 10, 14, 18, 22]
        assert [plan.col_of[("quarter", bi)] for bi in range(2)] == [26, 30]

    @pytest.mark.parametrize("at", ["between", "end"])
    def test_the_kinds_stand_finest_first_whatever_their_order(self, at):
        # Declared year-first, the year's total used to stand before Q4's,
        # inside Q4's span: the header's merged cells overlapped.
        fine = build_totals_plan(("quarter", "year"), "month", "2026-01", 24, 3, at=at)
        coarse = build_totals_plan(("year", "quarter"), "month", "2026-01", 24, 3, at=at)
        assert coarse.columns == fine.columns
        assert coarse.col_of == fine.col_of
        assert coarse.kinds == ("year", "quarter")    # as declared
        if at == "between":
            assert fine.columns[15:17] == [("quarter", 3), ("year", 0)]

    def test_the_header_at_the_end(self):
        plan = build_totals_plan(("quarter", "year"), "month", "2026-01", 24, 3, at="end")
        labels = [f"{mon} {year}" for year in (2026, 2027)
                  for mon in ("Jan", "Feb", "Mar", "Apr", "May", "Jun",
                              "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")]
        years, quarters, months = header_rows(plan, labels)
        # The tiers span their own months only…
        assert [(c["col"], c["colspan"], c["label"]) for c in years["cells"]
                if "rowspan" not in c] == [(3, 12, "2026"), (15, 12, "2027")]
        assert [(c["col"], c["colspan"], c["label"]) for c in quarters["cells"]
                if "rowspan" not in c] == [
            (3 + 3 * q, 3, f"Q{q % 4 + 1}") for q in range(8)]
        # …and a total, with no tier over it, names its bucket in full,
        # from its own tier down through the month row.
        assert [(c["col"], c["rowspan"], c["label"]) for c in years["cells"]
                if "rowspan" in c] == [(35, 3, "Total 2026"), (36, 3, "Total 2027")]
        assert [(c["col"], c["rowspan"], c["label"]) for c in quarters["cells"]
                if "rowspan" in c] == [
            (27 + q, 2, f"Total Q{q % 4 + 1} {2026 + q // 4}") for q in range(8)]
        assert [c["label"] for c in months["cells"]] == [lbl.split(" ")[0] for lbl in labels]


# ═════════════════════════════════════════════════════════════════════
# The book
# ═════════════════════════════════════════════════════════════════════


class TestTheBook:
    def test_the_header_at_the_end(self):
        ws = _book(_readers())["Data"]
        assert _overlapping(ws) == []
        merges = _merges(ws)
        for span in ("C1:N1", "O1:Z1", "C2:E2", "F2:H2", "X2:Z2", "AA2:AA3", "AH2:AH3",
                     "AI1:AI3", "AJ1:AJ3"):
            assert span in merges, span
        assert ws["AA2"].value == "Total Q1 2026"
        assert ws["AH2"].value == "Total Q4 2027"
        assert ws["AI1"].value == "Total 2026" and ws["AJ1"].value == "Total 2027"
        # The month row carries the months and nothing else.
        assert [ws.cell(3, c).value for c in range(3, 27)] == [
            "Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov",
            "Dec"] * 2
        assert ws["AA3"].value is None
        assert ws.freeze_panes == "C4"

    @pytest.mark.parametrize("at", [None, "end"])
    def test_a_coarse_first_declaration_lays_out_as_a_fine_first_one(self, at):
        fine = _book(_readers(at, ("quarter", "year")))["Data"]
        coarse = _book(_readers(at, ("year", "quarter")))["Data"]
        assert _overlapping(coarse) == []
        assert _merges(coarse) == _merges(fine)
        assert [[c.value for c in row] for row in coarse.iter_rows()] == [
            [c.value for c in row] for row in fine.iter_rows()]

    def test_the_totals_keep_their_formulas(self):
        """Every bucket cell is written as it is between the periods: the
        same text, its cells at other columns."""

        def model(at):
            m = mo.Model("m", default_grain="month", default_start="2026-01",
                         default_periods=6,
                         default_excel_view=mo.ExcelView(
                             timeline=_timeline(("quarter", "year"), at)))
            with m:
                m.plan = mo.MultiVariable("Plan")
                with m.plan as p:
                    p.rate = mo.Variable(0.1, display_name="Rate")
                    p.revenue = mo.Variable([10.0, 20.0, 30.0, 40.0, 50.0, 60.0],
                                            display_name="Revenue", regrain=mo.up("sum"))
                    p.price = mo.Variable([5.0, 6.0, 7.0, 8.0, 9.0, 10.0],
                                          display_name="Price", regrain=mo.up("mean"))
                    p.tax = p.revenue * p.rate
                    p.profit = mo.Variable([4.0, 3.0, 9.0, 8.0, 15.0, 12.0],
                                           display_name="Profit", regrain=mo.up("sum"))
                    p.margin = (p.profit / p.revenue).set_display_name("Margin").set_regrain(
                        mo.ratio("profit", "revenue"))
                    p.cash = mo.Variable([1.0, 2.0, 3.0, 4.0, 5.0, 6.0], display_name="Cash",
                                         regrain=mo.up("last"))
            return m

        def bucket_formulas(at):
            ws = _book(model(at))["Plan"]
            plan = build_totals_plan(("quarter", "year"), "month", "2026-01", 6, 3,
                                     at=at or "between")
            key_of = {col: key for key, col in plan.col_of.items()}

            def keyed(formula):
                # Each cell of the period area by its plan key, so the two
                # books' texts compare apart from where the columns stand.
                def name(found):
                    key = key_of.get(column_index_from_string(found.group(1)))
                    return f"<{key[0]} {key[1]}>{found.group(2)}" if key else found.group(0)
                return re.sub(r"(?<![A-Z])([A-Z]{1,3})(\d+)(?![\d(])", name, formula)

            return {
                (ws.cell(row, 1).value, key): keyed(ws.cell(row, col).value)
                for key, col in plan.col_of.items() if key[0] != NATIVE
                for row in range(4, ws.max_row + 1)
                if isinstance(ws.cell(row, col).value, str)
            }

        end, between = bucket_formulas("end"), bucket_formulas(None)
        assert end == between
        lines = {line for line, _key in end}
        assert lines == {"Revenue", "Price", "Tax", "Profit", "Margin", "Cash"}
        assert end[("Revenue", ("year", 0))] == (
            "=SUM(<period 0>5,<period 1>5,<period 2>5,<period 3>5,<period 4>5,<period 5>5)")
        assert end[("Tax", ("quarter", 1))] == "=<quarter 1>5 * B4"

    def test_the_outline_folds_the_months_then_the_quarters(self):
        """The same levels as between the periods: a month two buckets deep,
        a quarter's total one, a year's none. At the end they make the
        months one group and the quarters the next."""
        end = _book(_readers("end"))["Data"]
        assert {_level(end, get_column_letter(c)) for c in range(3, 27)} == {2}
        assert {_level(end, get_column_letter(c)) for c in range(27, 35)} == {1}
        assert _level(end, "AI") == 0 and _level(end, "AJ") == 0
        assert end.sheet_properties.outlinePr.summaryRight is True
        between = _book(_readers(None))["Data"]
        assert end.freeze_panes == between.freeze_panes == "C4"
        assert end["C3"].value == between["C3"].value == "Jan"
        assert end["C4"].value == between["C4"].value == COUNTS[0]


# ═════════════════════════════════════════════════════════════════════
# A row read whole
# ═════════════════════════════════════════════════════════════════════


class TestARowStaysOneRange:
    def test_every_reader_of_a_whole_row_is_a_live_range(self):
        m = _readers()
        wb = _book(m)
        summary = {wb["Summary"].cell(r, 1).value: wb["Summary"].cell(r, 2).value
                   for r in range(1, wb["Summary"].max_row + 1)}
        assert summary == {
            "Total": "=SUM(Data!C4:Z4)",
            "Peak": "=MAX(Data!C4:Z4)",
            "Peak Month": "=MATCH(MAX(Data!C4:Z4), Data!C4:Z4, 0)",
            "Third": "=INDEX(Data!C4:Z4, 3)",
            "Big": '=COUNTIF(Data!C4:Z4, ">15")',
            "Big Sum": '=SUMIFS(Data!C4:Z4, Data!C4:Z4, ">15")',
            "Second": "=LARGE(Data!C4:Z4, 2)",
            "First Half": "=SUM(Data!C4:H4)",
            "Last": "=Data!Z4",
        }
        # On the row's own sheet too.
        assert wb["Data"]["B7"].value == "=MATCH(MAX(C4:Z4), C4:Z4, 0)"
        assert check_compat(m, "excel") == []

    @pytest.mark.parametrize("at", [None, "between"])
    def test_between_the_periods_a_row_is_no_run(self, at):
        # The default does not move: the year's total stands after
        # December's quarter, and a reader of one range keeps its value.
        m = _readers(at)
        wb = _book(m)
        assert wb["Data"]["R2"].value == "Total Q4" and wb["Data"]["S1"].value == "Total 2026"
        assert wb["Summary"]["B3"].value == m.summary.peak_month.value == 22
        assert "m.summary.peak_month" in {i.qpath for i in check_compat(m, "excel")}


# ═════════════════════════════════════════════════════════════════════
# The book's other shapes
# ═════════════════════════════════════════════════════════════════════


class TestTheShapes:
    def test_a_header_anchor_with_an_opening_column(self):
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=15,
                     default_excel_view=mo.ExcelView(timeline={
                         "totals": ["quarter", "year"], "totals_at": "end",
                         "opening": "Pre-start"}))
        with m:
            m.time = mo.MultiVariable("Time")
            with m.time as t:
                t.start = mo.time.start
                t.end = mo.time.end
                t.days = mo.time.days
            m.default_header = m.time
            m.ops = mo.MultiVariable("Ops")
            with m.ops as o:
                o.flow = mo.Variable([float(i) for i in range(15)], display_name="Flow",
                                     regrain=mo.up("sum"))
                o.where = mo.MATCH(mo.MAX(o.flow), o.flow, 0)
        wb = _book(m)
        time, ops = wb["Time"], wb["Ops"]
        assert _overlapping(time) == [] and _overlapping(ops) == []
        # The chain off the opening column runs through the fifteen months,
        # the totals after them.
        assert time["B3"].value == "Pre-start"
        assert time["C4"].value == "=B5+1" and time["Q4"].value == "=P5+1"
        assert time["C5"].value == "=EOMONTH(C4,0)"
        assert time["R2"].value == "Total Q1 2026" and time["W1"].value == "Total 2026"
        # The mirrored rows read the anchor's month at the same place.
        assert ops["D4"].value == "=Time!C4" and ops["R4"].value == "=Time!Q4"
        assert ops["C5"].value == "=Time!B5"
        rows = {ops.cell(r, 1).value: r for r in range(1, ops.max_row + 1)}
        flow = rows["Flow"]
        assert ops.cell(flow, 19).value == f"=SUM(D{flow},E{flow},F{flow})"   # S: Q1
        assert ops.cell(rows["Where"], 2).value == (
            f"=MATCH(MAX(D{flow}:R{flow}), D{flow}:R{flow}, 0)")
        assert check_compat(m, "excel") == []

    def test_a_compare_book_puts_its_groups_of_totals_after_the_last_period(self):
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=6,
                     tracks=mo.Tracks("plan", "actual",
                                      blend=mo.blend(given="actual", follow="plan",
                                                     until="2026-02", name="live")),
                     default_excel_view=mo.ExcelView(
                         tracks="compare", timeline={"totals": ["quarter"], "totals_at": "end"}))
        with m:
            m.p = mo.MultiVariable("P", excel_props={"tab": True})
            with m.p as p:
                p.flow = mo.Variable(plan=[10.0] * 6,
                                     actual=[12.0, 11.0, None, None, None, None],
                                     display_name="Flow", regrain=mo.up("sum"))
        ws = _book(m)["P"]
        assert _overlapping(ws) == []
        # A period is a group of four columns, and so is each total.
        assert "B1:M1" in _merges(ws) and "N1:Y1" in _merges(ws)
        assert "Z1:AC2" in _merges(ws) and "AD1:AG2" in _merges(ws)
        assert ws["Z1"].value == "Total Q1 2026" and ws["AD1"].value == "Total Q2 2026"
        # The word row runs over the total groups too.
        assert [ws.cell(3, c).value for c in range(26, 30)] == ["actual", "plan", "live", "Var"]
        assert ws["Z4"].value == "=SUM(B4,F4,J4)" and ws["AA4"].value == "=SUM(C4,G4,K4)"

    def test_a_day_sheet_totals_its_months_at_the_end(self):
        m = mo.Model("m", default_grain="day", default_start="2026-01-30", default_periods=5,
                     default_excel_view=mo.ExcelView(
                         timeline={"totals": ["month"], "totals_at": "end"}))
        with m:
            m.flow = mo.Variable([1.0, 2.0, 3.0, 4.0, 5.0], regrain=mo.up("sum"),
                                 display_name="Flow")
            m.top = mo.LARGE(m.flow, 2)
        ws = _book(m).worksheets[0]
        assert _merges(ws) == ["C1:D1", "E1:G1", "H1:H2", "I1:I2"]
        assert [ws["C1"].value, ws["E1"].value] == ["Jan", "Feb"]
        assert [ws["H1"].value, ws["I1"].value] == ["Total Jan 2026", "Total Feb 2026"]
        assert ws["H3"].value == "=SUM(C3,D3)" and ws["I3"].value == "=SUM(E3,F3,G3)"
        assert ws["B4"].value == "=LARGE(C3:G3, 2)"

    def test_a_quarterly_sheet_with_a_title_row(self):
        m = mo.Model("m", default_grain="quarter", default_start="2026-Q1", default_periods=8,
                     default_excel_view=mo.ExcelView(
                         timeline={"totals": ["year"], "totals_at": "end"},
                         sheet={"title": True}))
        with m:
            m.s = mo.MultiVariable("S", description="quarterly book")
            with m.s as s:
                s.flow = mo.Variable([float(i) for i in range(8)], display_name="Flow",
                                     regrain=mo.up("sum"))
                s.big = mo.COUNTIF(s.flow, ">3")
        ws = _book(m)["S"]
        assert ws["A1"].value == "S"
        assert _merges(ws) == ["C2:F2", "G2:J2", "K2:K3", "L2:L3"]
        assert [ws["K2"].value, ws["L2"].value] == ["Total 2026", "Total 2027"]
        assert ws["K4"].value == "=SUM(C4,D4,E4,F4)" and ws["L4"].value == "=SUM(G4,H4,I4,J4)"
        assert ws["B5"].value == '=COUNTIF(C4:J4, ">3")'
        assert ws.freeze_panes == "C4"


# ═════════════════════════════════════════════════════════════════════
# The view
# ═════════════════════════════════════════════════════════════════════


def _two_sheets(model_at="end", a_view=None, b_view=None):
    """Two monthly sheets of 15 months under a model whose view puts the
    totals at ``model_at``; each sheet may hold a view of its own."""
    m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=15,
                 default_excel_view=mo.ExcelView(
                     timeline=_timeline(("quarter", "year"), model_at)))
    with m:
        m.a = mo.MultiVariable("A")
        with m.a as a:
            a.x = mo.Variable([1.0] * 15, display_name="X", regrain=mo.up("sum"))
        m.b = mo.MultiVariable("B")
        with m.b as b:
            b.y = mo.Variable([2.0] * 15, display_name="Y", regrain=mo.up("sum"))
    if a_view is not None:
        m.a.default_excel_view = a_view
    if b_view is not None:
        m.b.default_excel_view = b_view
    return m


def _year_total_col(ws):
    """The column of "Total 2026" on a sheet's top row."""
    return next(c for c in range(1, ws.max_column + 1) if ws.cell(1, c).value == "Total 2026")


class TestTheView:
    def test_a_place_the_book_does_not_know_is_refused_by_name(self):
        with pytest.raises(ValueError, match=(
                r"ExcelView timeline totals_at may name between, end; got 'middle'\. "
                r"'between' stands each total column after its own periods, 'end' "
                r"stands every one after the last period\.")):
            mo.ExcelView(timeline={"totals": ["year"], "totals_at": "middle"})

    def test_a_view_edited_after_it_was_built_is_refused_too(self):
        view = mo.ExcelView(timeline={"totals": ["year"]})
        view.timeline = {"totals": ["year"], "totals_at": "End"}
        with pytest.raises(ValueError, match="got 'End'"):
            view.snapshot()
        # The form the default view is written in: a timeline of lines.
        view = mo.ExcelView(timeline=mo.MultiVariable(
            "timeline", totals=mo.Variable(["year"]), totals_at=mo.Variable("after")))
        with pytest.raises(ValueError, match="got 'after'"):
            view.snapshot()

    def test_unset_is_between(self):
        assert mo.ExcelView(timeline={"totals": ["year"]}).snapshot().time_totals_at is None
        assert mo.ExcelView(timeline={"totals_at": "end"}).snapshot().time_totals_at == "end"
        ws = _book(_two_sheets(None))["A"]
        assert _year_total_col(ws) == 2 + 12 + 4    # after December and four quarters

    def test_a_nearer_view_inherits_the_end_and_between_takes_it_back(self):
        m = _two_sheets(
            "end",
            a_view=mo.ExcelView(timeline={"totals": ["year"]}),
            b_view=mo.ExcelView(timeline={"totals": ["year"], "totals_at": "between"}),
        )
        a, b = resolve_excel_view(m.a.x), resolve_excel_view(m.b.y)
        assert (a.time_totals, a.time_totals_at) == (("year",), "end")
        assert (b.time_totals, b.time_totals_at) == (("year",), "between")
        wb = _book(m)
        assert _year_total_col(wb["A"]) == 2 + 15        # after the fifteenth month
        assert _year_total_col(wb["B"]) == 2 + 12        # right after December

    def test_extend_carries_it(self):
        view = mo.ExcelView(timeline={"totals": ["year"], "totals_at": "end"})
        assert view.extend(font={"name": "Arial"}).snapshot().time_totals_at == "end"

    def test_a_named_view_puts_the_totals_at_the_end_once_applied(self):
        m = _two_sheets(None)
        m.board = mo.ExcelView("Board", timeline={"totals": ["quarter", "year"],
                                                  "totals_at": "end"})
        assert _year_total_col(_book(m)["A"]) == 2 + 12 + 4
        with mo.applied_view(m.board):
            assert _year_total_col(_book(m)["A"]) == 2 + 15 + 5

    def test_the_house_view_and_the_default_view_carry_it(self, monkeypatch):
        m = _two_sheets(None)
        with mo.house_view(mo.ExcelView(timeline={"totals_at": "end"})):
            assert _year_total_col(_book(m)["A"]) == 2 + 15 + 5
        floor = mo.default_excel_view.extend(
            timeline={"header": True, "label_format": "finance", "totals_at": "end"})
        monkeypatch.setattr(mo, "default_excel_view", floor)
        assert _year_total_col(_book(m)["A"]) == 2 + 15 + 5
        # The model's own view takes it back.
        m.default_excel_view = mo.ExcelView(timeline=_timeline(("quarter", "year"), "between"))
        assert _year_total_col(_book(m)["A"]) == 2 + 12 + 4


# ═════════════════════════════════════════════════════════════════════
# Recalculated
# ═════════════════════════════════════════════════════════════════════


def _recalculated(model, tmp_path):
    import shutil
    import subprocess
    from pathlib import Path

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    if soffice is None and mac.exists():
        soffice = str(mac)
    if soffice is None:
        pytest.skip("LibreOffice is not installed")
    src = tmp_path / "book.xlsx"
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        model.to_excel(str(src))
    outdir = tmp_path / "calc"
    profile = (tmp_path / "lo-profile").as_uri()  # private per run
    subprocess.run(
        [soffice, f"-env:UserInstallation={profile}", "--headless",
         "--convert-to", "xlsx", "--outdir", str(outdir), str(src)],
        check=True, capture_output=True, timeout=120,
    )
    return openpyxl.load_workbook(str(outdir / "book.xlsx"), data_only=True)


@pytest.mark.slow
def test_the_end_book_recalculates_to_the_engine(tmp_path):
    """Recalculated by a spreadsheet program, every total of the end book
    is the engine's re-grained number, and every whole-row reader the
    engine's own."""
    m = _readers()
    wb = _recalculated(m, tmp_path)
    data, summary = wb["Data"], wb["Summary"]
    plan = build_totals_plan(("quarter", "year"), "month", "2026-01", 24, 3, at="end")
    rows = {data.cell(r, 1).value: r for r in range(4, data.max_row + 1)}
    checked = 0
    for kind in ("quarter", "year"):
        coarse = m.at(kind).data
        for label, line in (("Customers", coarse.customers), ("Price", coarse.price),
                            ("Revenue", coarse.revenue)):
            for bi, want in enumerate(line.value):
                got = data.cell(rows[label], plan.col_of[(kind, bi)]).value
                assert got == pytest.approx(want, rel=1e-12), (label, kind, bi)
                checked += 1
    assert checked == 3 * (8 + 2)
    readers = {summary.cell(r, 1).value: summary.cell(r, 2).value
               for r in range(1, summary.max_row + 1)}
    s = m.summary
    assert readers == {
        "Total": pytest.approx(s.total.value), "Peak": pytest.approx(s.peak.value),
        "Peak Month": s.peak_month.value, "Third": pytest.approx(s.third.value),
        "Big": s.big.value, "Big Sum": pytest.approx(s.big_sum.value),
        "Second": pytest.approx(s.second.value), "First Half": pytest.approx(s.first_half.value),
        "Last": pytest.approx(s.last.value),
    }
    assert (s.peak_month.value, s.third.value, s.big.value, s.second.value) == (22, 9.0, 11, 25.0)
    assert data["B7"].value == m.data.peak_month.value == 22
