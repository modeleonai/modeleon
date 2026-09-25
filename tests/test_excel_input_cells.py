# SPDX-License-Identifier: Apache-2.0
"""Input cells that hold dates and durations.

A date typed into a spreadsheet is a number (days since the sheet's
epoch) that is merely DISPLAYED as a date. Written as the text
``'2024-03-31'`` it still looks right, and ``+ 1`` even happens to work,
but a comparison does not: text sorts above every number, so
``start < EDATE(start, 1)`` is false in the workbook while Python says
true. A duration (``timedelta``) has no spreadsheet type of its own; its
natural form is a count of days, which is also what a formula adds to a
date. These tests pin both, and recalculate the written workbook to show
the file computes what Python computed.
"""

from __future__ import annotations

import subprocess
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import openpyxl
import pytest

import modeleon as mo


def _cells(model, tmp_path, sheet=None):
    """Write ``model`` and return ``{label: [value cells]}`` of one sheet."""
    out = tmp_path / "book.xlsx"
    model.to_excel(str(out))
    wb = openpyxl.load_workbook(str(out))
    ws = wb[sheet] if sheet else wb.active
    return {
        row[0].value: [c for c in row[1:] if c.value is not None]
        for row in ws.iter_rows()
        if row[0].value is not None
    }


def _soffice() -> str | None:
    """Path to the LibreOffice command line, or None when not installed."""
    import shutil

    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


def _recalculated(model, tmp_path) -> openpyxl.Workbook:
    """Write ``model``, let LibreOffice calculate every formula, and
    return the recalculated workbook's stored values (every sheet)."""
    soffice = _soffice()
    if soffice is None:
        pytest.skip("LibreOffice is not installed")
    src = tmp_path / "book.xlsx"
    model.to_excel(str(src))
    outdir = tmp_path / "calc"
    profile = (tmp_path / "lo-profile").as_uri()  # private: parallel runs don't clash
    subprocess.run(
        [soffice, f"-env:UserInstallation={profile}", "--headless",
         "--convert-to", "xlsx", "--outdir", str(outdir), str(src)],
        check=True, capture_output=True, timeout=120,
    )
    return openpyxl.load_workbook(str(outdir / "book.xlsx"), data_only=True)


def _row(ws, label):
    """The value cells right of ``label`` in column A."""
    for row in ws.iter_rows():
        if row[0].value == label:
            return [c.value for c in row[1:] if c.value is not None]
    raise KeyError(label)


def _as_date(value):
    """A recalculated cell's date, whether it came back as a date or a
    datetime at midnight."""
    return value.date() if isinstance(value, datetime) else value


class TestDateInputs:
    def test_a_date_is_a_real_date_cell(self, tmp_path):
        m = mo.Model("Dates")
        m.start = mo.Variable(date(2024, 1, 15), display_name="Start")
        (cell,) = _cells(m, tmp_path)["Start"]
        # Stored as a serial number with a date format — openpyxl reads
        # such a cell back as a date; a text cell would come back as 's'.
        assert cell.data_type == "d"
        assert cell.is_date
        assert cell.value == datetime(2024, 1, 15)
        assert cell.number_format == "yyyy-mm-dd"

    def test_a_datetime_keeps_its_time_of_day(self, tmp_path):
        m = mo.Model("Dates")
        m.stamp = mo.Variable(datetime(2024, 1, 15, 18, 30), display_name="Stamp")
        (cell,) = _cells(m, tmp_path)["Stamp"]
        assert cell.data_type == "d"
        assert cell.value == datetime(2024, 1, 15, 18, 30)
        # The display shows the clock too, or 18:30 would silently read
        # as the start of the day.
        assert cell.number_format.startswith("yyyy-mm-dd")
        assert "hh:mm" in cell.number_format

    def test_a_datetime_at_midnight_shows_as_a_date(self, tmp_path):
        m = mo.Model("Dates")
        m.day = mo.Variable(datetime(2024, 1, 15), display_name="Day")
        (cell,) = _cells(m, tmp_path)["Day"]
        assert cell.data_type == "d"
        assert cell.number_format == "yyyy-mm-dd"

    def test_a_list_of_dates_is_one_date_cell_per_period(self, tmp_path):
        ends = [date(2024, 1, 31), date(2024, 2, 29), date(2024, 3, 31)]
        m = mo.Model("Dates")
        m.ends = mo.Variable(ends, display_name="Ends")
        cells = _cells(m, tmp_path)["Ends"]
        assert [c.data_type for c in cells] == ["d", "d", "d"]
        assert [c.value.date() for c in cells] == ends
        assert {c.number_format for c in cells} == {"yyyy-mm-dd"}

    def test_an_explicit_number_format_still_wins(self, tmp_path):
        m = mo.Model("Dates")
        m.start = mo.Variable(
            date(2024, 1, 15), display_name="Start",
            excel_props={"number_format": "mmm yyyy"},
        )
        (cell,) = _cells(m, tmp_path)["Start"]
        assert cell.data_type == "d"
        assert cell.number_format == "mmm yyyy"

    def test_a_sections_number_format_does_not_turn_a_date_into_a_number(self, tmp_path):
        # A section's number format speaks for its numbers. A date line
        # inside it keeps a date format, or 2024-01-15 would read 45,306.00;
        # a date format the section states itself is still honoured.
        m = mo.Model("Dates")
        with m:
            m.sec = mo.MultiVariable(
                "Sec", excel_props={"tab": True, "number_format": "#,##0.00"},
            )
            m.sec.start = mo.Variable(date(2024, 1, 15), display_name="Start")
            m.sec.next = mo.EDATE(m.sec.start, 1)
            m.sec.next._display_name = "Next"
            m.sec.amount = mo.Variable(1250.5, display_name="Amount")
            m.cal = mo.MultiVariable(
                "Cal", excel_props={"tab": True, "number_format": "mmm yyyy"},
            )
            m.cal.month = mo.Variable(date(2024, 1, 31), display_name="Month")
        sec = _cells(m, tmp_path, "Sec")
        assert [c.number_format for c in sec["Start"]] == ["yyyy-mm-dd"]
        assert [c.number_format for c in sec["Next"]] == ["yyyy-mm-dd"]
        assert [c.number_format for c in sec["Amount"]] == ["#,##0.00"]
        cal = _cells(m, tmp_path, "Cal")
        assert [c.number_format for c in cal["Month"]] == ["mmm yyyy"]

    def test_a_zoned_datetime_keeps_its_clock_reading(self, tmp_path):
        # A spreadsheet date has no time zone; the cell shows the time
        # the value states rather than refusing to write the workbook.
        m = mo.Model("Dates")
        m.stamp = mo.Variable(
            datetime(2024, 1, 15, 18, 0, tzinfo=timezone.utc), display_name="Stamp",
        )
        (cell,) = _cells(m, tmp_path)["Stamp"]
        assert cell.data_type == "d"
        assert cell.value == datetime(2024, 1, 15, 18, 0)

    def test_a_date_on_a_downward_sheet(self, tmp_path):
        m = mo.Model("Dates")
        with m:
            m.inputs = mo.MultiVariable("Inputs", excel_props={"tab": True})
            m.inputs.default_excel_view = mo.ExcelView(orient="down")
            m.inputs.ends = mo.Variable(
                [date(2024, 1, 31), date(2024, 2, 29)], display_name="Ends",
            )
        out = tmp_path / "book.xlsx"
        m.to_excel(str(out))
        ws = openpyxl.load_workbook(str(out))["Inputs"]
        dated = [c for row in ws.iter_rows() for c in row if c.is_date]
        assert [c.value.date() for c in dated] == [date(2024, 1, 31), date(2024, 2, 29)]
        assert {c.data_type for c in dated} == {"d"}


class TestComputedDates:
    def test_a_formula_that_computes_a_date_shows_as_a_date(self, tmp_path):
        # IF / CHOOSE / DATE results carry no date type of their own; the
        # values decide, or the cell would show a bare day number.
        m = mo.Model("Dates")
        m.start = mo.Variable(date(2024, 1, 15), display_name="Start")
        m.end = mo.Variable(date(2024, 3, 1), display_name="End")
        m.pick = mo.IF(m.start < m.end, m.start, m.end)
        m.pick._display_name = "Pick"
        m.first = mo.CHOOSE(1, m.start, m.end)
        m.first._display_name = "First"
        m.due = mo.Variable(mo.DATE(2024, 5, 1), display_name="Due")
        cells = _cells(m, tmp_path)
        for label in ("Pick", "First", "Due"):
            (cell,) = cells[label]
            assert cell.value.startswith("="), label
            assert cell.number_format == "yyyy-mm-dd", label

    def test_days_between_two_dates_stay_a_number(self, tmp_path):
        m = mo.Model("Dates")
        m.start = mo.Variable(date(2024, 1, 15), display_name="Start")
        m.end = mo.Variable(date(2024, 3, 1), display_name="End")
        m.span = m.end - m.start
        m.span._display_name = "Span"
        (cell,) = _cells(m, tmp_path)["Span"]
        assert cell.value.startswith("=")
        assert "yy" not in cell.number_format


class TestDurationInputs:
    def test_a_timedelta_is_its_count_of_days(self, tmp_path):
        m = mo.Model("Durations")
        m.gap = mo.Variable(timedelta(days=5), display_name="Gap")
        m.half = mo.Variable(timedelta(hours=12), display_name="Half")
        cells = _cells(m, tmp_path)
        (gap,) = cells["Gap"]
        (half,) = cells["Half"]
        assert (gap.data_type, gap.value) == ("n", 5)
        assert (half.data_type, half.value) == ("n", 0.5)
        assert not gap.is_date and not half.is_date

    def test_a_list_of_timedeltas(self, tmp_path):
        m = mo.Model("Durations")
        m.gaps = mo.Variable(
            [timedelta(days=1), timedelta(days=2, hours=6)], display_name="Gaps",
        )
        cells = _cells(m, tmp_path)["Gaps"]
        assert [(c.data_type, c.value) for c in cells] == [("n", 1), ("n", 2.25)]


class TestOtherInputsUnchanged:
    def test_numbers_text_and_booleans(self, tmp_path):
        m = mo.Model("Plain")
        m.n = mo.Variable(7, display_name="N")
        m.x = mo.Variable(2.5, display_name="X")
        m.s = mo.Variable("north", display_name="S")
        m.b = mo.Variable(True, display_name="B")
        m.gaps = mo.Variable([1, None, 3], display_name="Gaps")
        cells = _cells(m, tmp_path)
        assert [(c.data_type, c.value) for c in cells["N"]] == [("n", 7)]
        assert [(c.data_type, c.value) for c in cells["X"]] == [("n", 2.5)]
        assert [(c.data_type, c.value) for c in cells["S"]] == [("s", "north")]
        assert [(c.data_type, c.value) for c in cells["B"]] == [("b", True)]
        assert [c.value for c in cells["Gaps"]] == [1, 3]


class TestWorkbookComputesWhatPythonComputes:
    """Recalculate the written file and compare with the Python values."""

    def test_comparing_a_date_input(self, tmp_path):
        m = mo.Model("Dates")
        m.start = mo.Variable(date(2024, 1, 15), display_name="Start")
        m.early = mo.IF(m.start < mo.EDATE(m.start, 1), 1, 0)
        m.early._display_name = "Early"
        assert m.early._value == 1
        ws = _recalculated(m, tmp_path).active
        assert _row(ws, "Early") == [1]

    def test_adding_a_timedelta_input_to_a_date(self, tmp_path):
        m = mo.Model("Dates")
        m.start = mo.Variable(date(2024, 1, 15), display_name="Start")
        m.gap = mo.Variable(timedelta(days=5), display_name="Gap")
        m.half = mo.Variable(timedelta(hours=12), display_name="Half")
        m.later = m.start + m.gap
        m.later._display_name = "Later"
        m.stamp = m.start + m.half
        m.stamp._display_name = "Stamp"
        assert m.later._value == date(2024, 1, 20)
        assert m.stamp._value == datetime(2024, 1, 15, 12, 0)
        ws = _recalculated(m, tmp_path).active
        assert [_as_date(v) for v in _row(ws, "Later")] == [date(2024, 1, 20)]
        assert _row(ws, "Stamp") == [datetime(2024, 1, 15, 12, 0)]

    def test_per_period_dates_compared_across_sheets(self, tmp_path):
        m = mo.Model("Dates")
        with m:
            m.out = mo.MultiVariable("Out", excel_props={"tab": True})
            m.inputs = mo.MultiVariable("Inputs", excel_props={"tab": True})
            m.inputs.default_excel_view = mo.ExcelView(orient="down")
            m.inputs.cutoff = mo.Variable(date(2024, 2, 15), display_name="Cutoff")
            m.inputs.ends = mo.Variable(
                [date(2024, 1, 31), date(2024, 2, 29), date(2024, 3, 31)],
                display_name="Ends",
            )
            # Against a COMPUTED date: two text dates would compare in
            # the right order by accident, text against a number never does.
            m.out.after = mo.IF(m.inputs.ends > mo.EOMONTH(m.inputs.cutoff, 0), 1, 0)
            m.out.after._display_name = "After"
        assert m.out.after._value == [0, 0, 1]
        wb = _recalculated(m, tmp_path)
        assert _row(wb["Out"], "After") == [0, 0, 1]
