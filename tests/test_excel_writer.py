# SPDX-License-Identifier: Apache-2.0
"""Tests for excel_writer — the explicit-root emission path.

Every test builds an explicit root MultiVariable and calls ``.to_excel()``
on it. Components attach via ``parent.child = ...`` only.
"""

import pytest
from openpyxl import load_workbook

import modeleon as mo
from modeleon.compile.excel.layout import sheet_name_for


class TestSheetNameFor:
    """``sheet_name_for`` is the canonical sheet-name derivation used by
    LayoutEngine, _build_var_to_sheet, the writer, and pro's bridge.
    Drift between any two of those callers produced the
    case-mismatched-qualifier bug; this test pins the priority chain.

    Effective chain (since ``MultiVariableBase.sheet_name`` is a
    property that returns ``display_name`` for sheet-roled MVs and
    None otherwise):

      1. sheet_name (for sheet-roled MVs → display_name)
      2. display_name
      3. id

    Same string either way for sheet MVs; the test pins that the
    helper returns the user-visible label.
    """

    def test_model_returns_display_name(self):
        # ``mo.Model("forecast")`` title-cases display_name to
        # "Forecast" — that's the Excel tab name, and the value
        # downstream layers (translator's var_to_sheet, the writer
        # itself) must agree on.
        m = mo.Model("forecast")
        assert m.display_name == "Forecast"
        assert sheet_name_for(m) == "Forecast"

    def test_tab_marked_multivariable_returns_display_name(self):
        # Sheet-roled MultiVariable with explicit display_name.
        mv = mo.MultiVariable("Income Statement", excel_props={'tab': True})
        mv.revenue = mo.Variable(1_000_000)
        assert sheet_name_for(mv) == "Income Statement"

    def test_unanchored_mv_falls_through_to_id_or_display(self):
        # Non-sheet MV has ``sheet_name`` = None; the helper then
        # uses display_name. We don't pin the exact case (humanize
        # rules) — just that the helper returns a non-empty string
        # that's NOT None.
        mv = mo.MultiVariable("placeholder")
        result = sheet_name_for(mv)
        assert isinstance(result, str)
        assert result


class TestToExcel:
    def test_writes_file(self, tmp_path):
        model = mo.MultiVariable("Writes")
        model.test = mo.MultiVariable("Test", excel_props={'tab': True})
        model.test.revenue = mo.Variable(1_000_000, display_name="Revenue")

        model.to_excel(tmp_path / "out.xlsx")

        path = tmp_path / "out.xlsx"
        assert path.exists()

    def test_empty_model_raises(self, tmp_path):
        model = mo.MultiVariable("Empty")
        with pytest.raises(ValueError, match="Nothing to emit"):
            model.to_excel(tmp_path / "out.xlsx")

    def test_creates_tab_per_sheet(self, tmp_path):
        model = mo.MultiVariable("Tabs")
        model.assumptions = mo.MultiVariable(
            "Assumptions", rate=mo.Variable(0.25), excel_props={'tab': True},
        )
        model.results = mo.MultiVariable(
            "Results", profit=mo.Variable(100_000), excel_props={'tab': True},
        )

        model.to_excel(tmp_path / "out.xlsx")

        path = tmp_path / "out.xlsx"
        wb = load_workbook(path)
        assert set(wb.sheetnames) == {"Assumptions", "Results"}

    def test_input_variable_writes_value(self, tmp_path):
        model = mo.MultiVariable("Inputs")
        model.test = mo.MultiVariable("Test", excel_props={'tab': True})
        model.test.revenue = mo.Variable(1_000_000, display_name="Revenue")

        model.to_excel(tmp_path / "out.xlsx")

        path = tmp_path / "out.xlsx"
        wb = load_workbook(path)
        ws = wb["Test"]

        values = [
            (cell.coordinate, cell.value)
            for row in ws.iter_rows()
            for cell in row
            if cell.value is not None
        ]
        labels = [v for _, v in values if isinstance(v, str)]
        numbers = [v for _, v in values if isinstance(v, (int, float))]

        assert "Revenue" in labels
        assert 1_000_000 in numbers

    def test_formula_variable_writes_excel_formula(self, tmp_path):
        model = mo.MultiVariable("Formulas")
        model.test = mo.MultiVariable("Test", excel_props={'tab': True})
        model.test.revenue = mo.Variable(1_000_000, display_name="Revenue")
        model.test.cogs_pct = mo.Variable(0.6, display_name="COGS %")
        model.test.cogs = model.test.revenue * model.test.cogs_pct

        model.to_excel(tmp_path / "out.xlsx")

        path = tmp_path / "out.xlsx"
        wb = load_workbook(path)
        ws = wb["Test"]

        formulas = []
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas.append(cell.value)

        assert len(formulas) >= 1
        for f in formulas:
            assert "revenue" not in f.lower() or "*" in f

    def test_unsafe_sheet_name_sanitized(self, tmp_path):
        model = mo.MultiVariable("Unsafe")
        model.unsafe = mo.MultiVariable(
            "Tab:With/Slash", x=mo.Variable(1), excel_props={'tab': True},
        )

        model.to_excel(tmp_path / "out.xlsx")

        path = tmp_path / "out.xlsx"
        wb = load_workbook(path)
        tab = wb.sheetnames[0]
        assert ":" not in tab
        assert "/" not in tab


class TestPositionalSheetSectionRule:
    """`to_excel` matches the HTML repr's positional rule: the root is
    the workbook (never its own tab), first-depth sub-MVs become tabs,
    deeper sub-MVs render as sections inside their containing tab —
    regardless of an inner MV's ``_is_sheet`` flag."""

    def test_root_is_not_a_tab_even_if_marked(self, tmp_path):
        # Even though `t` is marked `tab=True`, it's the workbook
        # itself; the only tab should come from its first-depth child.
        acme = mo.Model("acme")
        acme.revenue = mo.Variable(1_000_000)
        t = mo.Model("t", excel_props={"tab": True})
        t.h = acme

        t.to_excel(tmp_path / "t.xlsx")

        wb = load_workbook(tmp_path / "t.xlsx")
        assert wb.sheetnames == ["Acme"]

    def test_nested_sheet_marker_renders_as_section(self, tmp_path):
        # `pnl._is_sheet=True`, but it's depth 2 from the root passed
        # to `to_excel`, so it renders as a section in `Acme`, not a
        # separate tab. Section header lands at row 1; the variable
        # rows + cogs formula reference the flat layout.
        acme = mo.Model("acme")
        with mo.MultiVariable(excel_props={"tab": True}) as pnl:
            pnl.revenue = mo.Variable(1_000_000, display_name="Revenue")
            pnl.cogs = (pnl.revenue * mo.Variable(0.6)).set_display_name("Cogs")
        acme.pnl = pnl
        t = mo.Model("t")
        t.h = acme

        t.to_excel(tmp_path / "t.xlsx")

        wb = load_workbook(tmp_path / "t.xlsx")
        assert wb.sheetnames == ["Acme"]
        ws = wb["Acme"]
        cells = {
            cell.coordinate: cell.value
            for row in ws.iter_rows()
            for cell in row
            if cell.value is not None
        }
        assert cells == {
            "A1": "Pnl",  # nested sheet flattens to a section header
            "A2": "Revenue", "B2": 1_000_000,
            "A3": "Cogs",    "B3": "=B2 * 0.6",
        }

    def test_orphan_variables_get_synthesized_overview_tab(self, tmp_path):
        # When the root has no sub-MVs, only Variables, those orphans
        # land in a synthesized overview tab named after the root.
        pnl = mo.MultiVariable("PnL")
        pnl.revenue = mo.Variable(100, display_name="Revenue")
        pnl.cogs = (pnl.revenue * mo.Variable(0.6)).set_display_name("Cogs")

        pnl.to_excel(tmp_path / "p.xlsx")

        wb = load_workbook(tmp_path / "p.xlsx")
        assert wb.sheetnames == ["PnL"]


class TestVariableToExcel:
    """A lone ``Variable`` exports just like its HTML repr — one row
    on a tab named after the Variable. The pseudo-sheet wrapper used
    for layout is the same one ``variable_html`` uses, so file output
    and notebook output agree on tab name + cell shape."""

    def test_named_variable_writes_one_row(self, tmp_path):
        v = mo.Variable(100, display_name="Revenue")

        v.to_excel(tmp_path / "v.xlsx")

        wb = load_workbook(tmp_path / "v.xlsx")
        assert wb.sheetnames == ["Revenue"]
        ws = wb["Revenue"]
        cells = {
            cell.coordinate: cell.value
            for row in ws.iter_rows()
            for cell in row
            if cell.value is not None
        }
        assert cells == {"A1": "Revenue", "B1": 100}

    def test_recurrence_writes_period_chain(self, tmp_path):
        r = mo.recurrence(
            start=1000, formula="{prev}*1.05", periods=4, display_name="Users",
        )

        r.to_excel(tmp_path / "r.xlsx")

        wb = load_workbook(tmp_path / "r.xlsx")
        assert wb.sheetnames == ["Users"]
        ws = wb["Users"]
        formulas = {
            cell.coordinate: cell.value
            for row in ws.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        }
        # Period 0 = literal start; later periods chain via {prev}.
        assert formulas == {
            "B1": "=1000",
            "C1": "=B1*1.05",
            "D1": "=C1*1.05",
            "E1": "=D1*1.05",
        }


class TestExcelLayoutOverrides:
    """The ``excel_layout`` kwarg lets pro (or any caller) override
    engine defaults via a duck-typed object — engine reads ``.sheet``
    and ``.format`` off it without importing pro types. These tests
    use ``SimpleNamespace`` stubs so the engine suite stays free of
    pro dependencies.
    """

    def test_layout_sheet_overrides_display_name(self, tmp_path):
        from types import SimpleNamespace

        # Without the override, ``mo.Model("forecast")`` lands on a
        # sheet named "Forecast" (display_name title-case). With
        # ``excel_layout.sheet="My P&L"``, the override wins.
        layout = SimpleNamespace(sheet="My P&L", format=None)
        m = mo.Model("forecast", excel_layout=layout)
        m.revenue = mo.Variable(1_000_000)

        m.to_excel(tmp_path / "sheet_override.xlsx")

        wb = load_workbook(tmp_path / "sheet_override.xlsx")
        assert wb.sheetnames == ["My P&L"]

    def test_layout_format_applies_to_variable_cells(self, tmp_path):
        from types import SimpleNamespace

        layout = SimpleNamespace(sheet=None, format="0.0%")
        m = mo.Model("metrics")
        m.tax_rate = mo.Variable(0.25, excel_layout=layout)

        m.to_excel(tmp_path / "format_override.xlsx")

        wb = load_workbook(tmp_path / "format_override.xlsx")
        ws = wb["Metrics"]
        # Find the value cell (column B in the standard scalar layout).
        value_cell = None
        for row in ws.iter_rows():
            for cell in row:
                if cell.value == 0.25:
                    value_cell = cell
                    break
        assert value_cell is not None, "tax_rate value cell not found"
        assert value_cell.number_format == "0.0%"

    def test_layout_format_wins_over_excel_props(self, tmp_path):
        """Priority: ``excel_layout.format`` beats ``excel_props
        ['number_format']``. The structured layout system is the
        primary surface; the flag-dict is legacy."""
        from types import SimpleNamespace

        layout = SimpleNamespace(sheet=None, format="0.00%")
        m = mo.Model("priority")
        m.r = mo.Variable(
            0.5,
            excel_props={"number_format": "#,##0"},
            excel_layout=layout,
        )

        m.to_excel(tmp_path / "priority.xlsx")

        wb = load_workbook(tmp_path / "priority.xlsx")
        ws = wb["Priority"]
        value_cell = next(
            (c for row in ws.iter_rows() for c in row if c.value == 0.5),
            None,
        )
        assert value_cell is not None
        # The layout's format wins, not the props dict's.
        assert value_cell.number_format == "0.00%"

    def test_no_layout_falls_back_to_existing_behavior(self, tmp_path):
        """Sanity: a model with no ``excel_layout`` set anywhere keeps
        the existing sheet-name + format heuristics. Regression guard
        — the override path mustn't alter behavior for callers that
        haven't opted in."""
        m = mo.Model("baseline")
        m.x = mo.Variable(42, excel_props={"number_format": "#,##0"})

        m.to_excel(tmp_path / "baseline.xlsx")

        wb = load_workbook(tmp_path / "baseline.xlsx")
        assert wb.sheetnames == ["Baseline"]
        ws = wb["Baseline"]
        value_cell = next(
            (c for row in ws.iter_rows() for c in row if c.value == 42),
            None,
        )
        assert value_cell is not None
        # Legacy excel_props path still applies when no layout is set.
        assert value_cell.number_format == "#,##0"


class TestExcelStyleProps:
    """Pin the ``_excel_props`` styling vocabulary on the emitted .xlsx.

    The writer consumes typography (``bold`` / ``italic`` / ``font_*``),
    fill (``bg``), borders (``border_*``), and alignment
    (``text_align`` / ``indent``). Each test asserts one prop class
    lands on the value cell (and label, where styling is row-wide).
    """

    def test_bold_applies_to_value_and_label(self, tmp_path):
        m = mo.Model("style")
        m.x = mo.Variable(10, display_name="X", excel_props={"bold": True})
        m.to_excel(tmp_path / "bold.xlsx")

        wb = load_workbook(tmp_path / "bold.xlsx")
        ws = wb["Style"]
        value_cell = next(
            (c for row in ws.iter_rows() for c in row if c.value == 10),
            None,
        )
        label_cell = next(
            (c for row in ws.iter_rows() for c in row if c.value == "X"),
            None,
        )
        assert value_cell is not None and value_cell.font.bold is True
        assert label_cell is not None and label_cell.font.bold is True

    def test_font_color_and_size(self, tmp_path):
        m = mo.Model("style")
        m.x = mo.Variable(
            7,
            excel_props={"font_color": "#FF0000", "font_size": 14},
        )
        m.to_excel(tmp_path / "font.xlsx")

        wb = load_workbook(tmp_path / "font.xlsx")
        ws = wb["Style"]
        cell = next(
            (c for row in ws.iter_rows() for c in row if c.value == 7),
            None,
        )
        assert cell is not None
        # openpyxl returns colors uppercased + may prepend the alpha
        # channel; assert membership rather than strict equality.
        assert cell.font.color is not None
        assert "FF0000" in str(cell.font.color.rgb).upper()
        assert cell.font.size == 14.0

    def test_bg_fill(self, tmp_path):
        m = mo.Model("style")
        m.x = mo.Variable(5, excel_props={"bg": "#FFFBE6"})
        m.to_excel(tmp_path / "fill.xlsx")

        wb = load_workbook(tmp_path / "fill.xlsx")
        ws = wb["Style"]
        cell = next(
            (c for row in ws.iter_rows() for c in row if c.value == 5),
            None,
        )
        assert cell is not None
        assert cell.fill.fill_type == "solid"
        assert "FFFBE6" in str(cell.fill.start_color.rgb).upper()

    def test_border_bool_default_thin(self, tmp_path):
        m = mo.Model("style")
        m.x = mo.Variable(
            1, excel_props={"border_top": True, "border_bottom": True}
        )
        m.to_excel(tmp_path / "borders.xlsx")

        wb = load_workbook(tmp_path / "borders.xlsx")
        ws = wb["Style"]
        cell = next(
            (c for row in ws.iter_rows() for c in row if c.value == 1),
            None,
        )
        assert cell is not None
        # Truthy bool defaults to "thin" — matches Excel's default
        # border weight when a user clicks the border button.
        assert cell.border.top.style == "thin"
        assert cell.border.bottom.style == "thin"
        # Untouched sides stay absent — openpyxl leaves the Side
        # slot itself ``None`` rather than a ``Side(style=None)``.
        assert cell.border.left is None or cell.border.left.style is None
        assert cell.border.right is None or cell.border.right.style is None

    def test_border_named_style(self, tmp_path):
        m = mo.Model("style")
        m.x = mo.Variable(2, excel_props={"border_top": "medium"})
        m.to_excel(tmp_path / "borders2.xlsx")

        wb = load_workbook(tmp_path / "borders2.xlsx")
        ws = wb["Style"]
        cell = next(
            (c for row in ws.iter_rows() for c in row if c.value == 2),
            None,
        )
        assert cell is not None and cell.border.top.style == "medium"

    def test_alignment(self, tmp_path):
        m = mo.Model("style")
        m.x = mo.Variable(
            3, excel_props={"text_align": "right", "indent": 2}
        )
        m.to_excel(tmp_path / "align.xlsx")

        wb = load_workbook(tmp_path / "align.xlsx")
        ws = wb["Style"]
        cell = next(
            (c for row in ws.iter_rows() for c in row if c.value == 3),
            None,
        )
        assert cell is not None
        assert cell.alignment.horizontal == "right"
        assert cell.alignment.indent == 2

    def test_unset_props_leave_defaults(self, tmp_path):
        # When no styling props are set, the cell's defaults apply
        # (no bold, no fill, no borders).
        m = mo.Model("plain")
        m.x = mo.Variable(8)
        m.to_excel(tmp_path / "plain.xlsx")

        wb = load_workbook(tmp_path / "plain.xlsx")
        ws = wb["Plain"]
        cell = next(
            (c for row in ws.iter_rows() for c in row if c.value == 8),
            None,
        )
        assert cell is not None
        assert cell.font.bold in (False, None)
        # openpyxl returns "none" fill_type for un-styled cells.
        assert cell.fill.fill_type in (None, "none")


class TestMultiVariableHeaderStyling:
    """Pin ``_excel_props`` styling on nested-MV section header cells.

    An MV with ``display_name`` produces a section header row in the
    flatten-nested-sheets layout. Before this support, the header cell
    landed bare — users (and the AI authoring loop) had to workaround
    by declaring a phantom ``mo.Variable(0.0, ...)`` row purely to
    get a styled label. With MV-level ``excel_props`` honoured, the
    natural shape works.
    """

    def test_mv_header_bold_and_bg(self, tmp_path):
        # MV section headers only appear when the MV is NESTED — at
        # first depth it becomes its own tab (and the display_name is
        # the tab title, not a header cell). So set up two layers.
        m = mo.Model("book")
        m.statements = mo.MultiVariable("Statements")
        m.statements.balance_sheet = mo.MultiVariable(
            "Balance Sheet",
            excel_props={"bold": True, "bg": "#0D1F3C", "font_color": "#FFFFFF"},
        )
        m.statements.balance_sheet.cash = mo.Variable(50_000, display_name="Cash")

        m.to_excel(tmp_path / "mv_header.xlsx")

        wb = load_workbook(tmp_path / "mv_header.xlsx")
        ws = wb["Statements"]
        header_cell = next(
            (
                c
                for row in ws.iter_rows()
                for c in row
                if c.value == "Balance Sheet"
            ),
            None,
        )
        assert header_cell is not None, "MV header cell not found"
        assert header_cell.font.bold is True
        assert "0D1F3C" in str(header_cell.fill.start_color.rgb).upper()
        assert "FFFFFF" in str(header_cell.font.color.rgb).upper()

    def test_mv_header_border(self, tmp_path):
        m = mo.Model("book")
        m.bs = mo.MultiVariable("Statements")
        m.bs.assets = mo.MultiVariable(
            "ASSETS",
            excel_props={"border_bottom": "medium"},
        )
        m.bs.assets.cash = mo.Variable(1, display_name="Cash")

        m.to_excel(tmp_path / "mv_border.xlsx")

        wb = load_workbook(tmp_path / "mv_border.xlsx")
        ws = wb["Statements"]
        header_cell = next(
            (
                c
                for row in ws.iter_rows()
                for c in row
                if c.value == "ASSETS"
            ),
            None,
        )
        assert header_cell is not None
        assert header_cell.border.bottom.style == "medium"

    def test_same_named_sections_on_different_sheets_dont_collide(self, tmp_path):
        # Regression: ``section_header_rows`` was keyed by the short name,
        # so two sections both named ``dates`` (control.dates and
        # timing.flags.dates) collided — the writer painted one sheet's
        # band on the *other* section's row, bleeding the band onto an
        # unrelated data cell. Keying by the fully-qualified id fixes it.
        m = mo.Model("book")

        m.control = mo.Model("control", display_name="Control")
        m.control.dates = mo.MultiVariable("Dates", excel_props={"bg": "#E9BA0D"})
        m.control.dates.valuation = mo.Variable(1, display_name="Valuation")
        m.control.sensitivity = mo.MultiVariable("Sensitivity")
        m.control.sensitivity.a = mo.Variable(1, display_name="A")
        m.control.sensitivity.b = mo.Variable(2, display_name="B")
        m.control.sensitivity.raw = mo.Variable(3, display_name="Raw")

        m.timing = mo.Model("timing", display_name="Timing")
        m.timing.flags = mo.MultiVariable("Flags")
        m.timing.flags.dates = mo.MultiVariable("Dates", excel_props={"bg": "#E9BA0D"})
        m.timing.flags.dates.x = mo.Variable(1, display_name="X")

        m.to_excel(tmp_path / "collide.xlsx")
        wb = load_workbook(tmp_path / "collide.xlsx")
        ws = wb["Control"]

        def band(label):
            cell = next(
                c for row in ws.iter_rows() for c in row if c.value == label
            )
            f = ws.cell(cell.row, 1).fill
            return f.fgColor.rgb if f and f.patternType else None

        # Only the real section header carries the band; data rows stay clean.
        assert "E9BA0D" in str(band("Dates")).upper()
        for data_label in ("A", "B", "Raw", "Valuation"):
            got = band(data_label)
            assert got is None or "E9BA0D" not in str(got).upper(), (
                f"{data_label!r} wrongly got the section band"
            )

    def test_mv_header_unstyled_when_no_props(self, tmp_path):
        # Regression: MV without ``excel_props`` still emits its header
        # but applies no extra styling — should remain backwards-compat
        # with existing models.
        m = mo.Model("book")
        m.bs = mo.MultiVariable("Statements")
        m.bs.section = mo.MultiVariable("Section")
        m.bs.section.x = mo.Variable(1)

        m.to_excel(tmp_path / "mv_plain.xlsx")

        wb = load_workbook(tmp_path / "mv_plain.xlsx")
        ws = wb["Statements"]
        header = next(
            (
                c
                for row in ws.iter_rows()
                for c in row
                if c.value == "Section"
            ),
            None,
        )
        assert header is not None
        assert header.font.bold in (False, None)
        assert header.fill.fill_type in (None, "none")


class TestFormatByType:
    """``format_by_type`` on an MV cascades a cell-type colour convention
    to every descendant value cell."""

    def _color(self, ws, label):
        # Return the explicit RGB font colour, or None for an uncoloured cell
        # (no colour, or the default ``theme`` text colour — ``.rgb`` errors
        # on theme colours, so guard on ``type``).
        for row in ws.iter_rows():
            if row and row[0].value == label:
                c = row[1].font.color
                if c is None or c.type != "rgb":
                    return None
                return c.rgb
        return None

    def test_default_palette_input_formula_reference(self, tmp_path):
        m = mo.Model("m", excel_props={"format_by_type": True})
        with m:
            m.ctrl = mo.Model("Control", display_name="Control")
            with m.ctrl:
                m.ctrl.rate = mo.Variable(0.1, display_name="Rate")
            m.calc = mo.Model("Calc", display_name="Calc")
            with m.calc:
                m.calc.base = mo.Variable(100, display_name="Base")
                m.calc.local = (m.calc.base * 2).set_display_name("Local")
                m.calc.xref = (m.calc.base * m.ctrl.rate).set_display_name("Xref")

        m.to_excel(tmp_path / "fbt.xlsx")
        ws = load_workbook(tmp_path / "fbt.xlsx")["Calc"]
        assert self._color(ws, "Base") == "FF0563C1"        # input → blue
        assert self._color(ws, "Local") is None             # same-sheet formula → uncoloured
        assert self._color(ws, "Xref") == "FF2E7D32"        # cross-sheet → green

    def test_reference_detected_through_nested_intermediates(self, tmp_path):
        # The cross-sheet cell sits below an operator/function intermediate
        # (``IF(end > control.date, …)``) — detection must recurse through
        # the floating ``end > control.date`` Variable to find it.
        m = mo.Model("m", excel_props={"format_by_type": True})
        with m:
            m.ctrl = mo.Model("Control", display_name="Control")
            with m.ctrl:
                m.ctrl.thresh = mo.Variable(50, display_name="Thresh")
            m.calc = mo.Model("Calc", display_name="Calc")
            with m.calc:
                m.calc.x = mo.Variable(100, display_name="X")
                m.calc.flag = mo.IF(m.calc.x > m.ctrl.thresh, 1, 0).set_display_name("Flag")

        m.to_excel(tmp_path / "nest.xlsx")
        ws = load_workbook(tmp_path / "nest.xlsx")["Calc"]
        assert self._color(ws, "Flag") == "FF2E7D32"        # cross-sheet → green

    def test_custom_palette_and_explicit_color_wins(self, tmp_path):
        m = mo.Model("m", excel_props={"format_by_type": {"input": "#B08C0A"}})
        with m:
            m.s = mo.MultiVariable("S", excel_props={"tab": True})
            m.s.a = mo.Variable(5, display_name="A")
            m.s.b = mo.Variable(6, display_name="B", excel_props={"font_color": "#FF0000"})

        m.to_excel(tmp_path / "fbt2.xlsx")
        ws = load_workbook(tmp_path / "fbt2.xlsx")["S"]
        assert self._color(ws, "A") == "FFB08C0A"           # custom bronze input
        assert self._color(ws, "B") == "FFFF0000"           # explicit font_color wins


class TestListExprPerCellRendering:
    """A ListExpr is positional — item *i* is the cell content for period
    *i*. Regression: the whole bracketed list used to be stamped into
    every cell (``=[0.0, -'Sheet'!B1, ...]`` — invalid Excel)."""

    def _grid(self, ws):
        return {
            c.coordinate: c.value
            for row in ws.iter_rows() for c in row if c.value is not None
        }

    def test_literal_items_write_raw_values_ref_items_write_formulas(self, tmp_path):
        root = mo.Model("m")
        root.gi = mo.Model("gi", display_name="Inputs")
        root.gi.rate = mo.Variable(53.0, display_name="Rate")
        root.inc = mo.Model("inc", display_name="Statement")
        root.inc.row = mo.Variable(
            [0.0, 0.0, -root.gi.rate, -root.gi.rate],
            var_type="list", display_name="Row",
        )
        root.to_excel(tmp_path / "out.xlsx")
        ws = load_workbook(tmp_path / "out.xlsx")["Statement"]
        grid = self._grid(ws)
        assert grid["B1"] == 0 and grid["C1"] == 0          # authored values, not '=0'
        assert grid["D1"] == "=-Inputs!B1"                  # live cross-sheet ref
        assert grid["E1"] == "=-Inputs!B1"

    def test_plain_ref_items_render_per_cell(self, tmp_path):
        root = mo.Model("m")
        root.gi = mo.Model("gi", display_name="Inputs")
        root.gi.rate = mo.Variable(53.0, display_name="Rate")
        root.s = mo.Model("s", display_name="S")
        root.s.row = mo.Variable(
            [root.gi.rate, root.gi.rate, root.gi.rate],
            var_type="list", display_name="Row",
        )
        root.to_excel(tmp_path / "out.xlsx")
        grid = self._grid(load_workbook(tmp_path / "out.xlsx")["S"])
        assert [grid["B1"], grid["C1"], grid["D1"]] == ["=Inputs!B1"] * 3

    def test_no_cell_ever_contains_bracket_pseudo_formula(self, tmp_path):
        root = mo.Model("m")
        root.gi = mo.Model("gi", display_name="Inputs")
        root.gi.rate = mo.Variable(2.0, display_name="Rate")
        root.s = mo.Model("s", display_name="S")
        root.s.row = mo.Variable(
            [1.0, -root.gi.rate, root.gi.rate], var_type="list", display_name="Row",
        )
        root.s.derived = (root.s.row * 2.0).set_display_name("Derived")
        root.to_excel(tmp_path / "out.xlsx")
        wb = load_workbook(tmp_path / "out.xlsx")
        for sn in wb.sheetnames:
            for row in wb[sn].iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("="):
                        assert "[" not in c.value, f"{sn}!{c.coordinate}: {c.value}"

    def test_derived_row_references_listexpr_row_cells(self, tmp_path):
        root = mo.Model("m")
        root.gi = mo.Model("gi", display_name="Inputs")
        root.gi.rate = mo.Variable(2.0, display_name="Rate")
        root.s = mo.Model("s", display_name="S")
        root.s.row = mo.Variable(
            [0.0, -root.gi.rate, -root.gi.rate], var_type="list", display_name="Row",
        )
        root.s.doubled = (root.s.row * 2.0).set_display_name("Doubled")
        root.to_excel(tmp_path / "out.xlsx")
        grid = self._grid(load_workbook(tmp_path / "out.xlsx")["S"])
        # the derived row references Row's own cells, never the inlined items
        assert grid["B2"] == "=B1 * 2.0"
        assert grid["C2"] == "=C1 * 2.0"

    def test_values_snapshot_unchanged(self):
        gi = mo.Model("gi")
        gi.rate = mo.Variable(53.0)
        row = mo.Variable([0.0, -gi.rate, -gi.rate], var_type="list")
        assert row._value == [0.0, -53.0, -53.0]

    def test_orient_down_transposes_per_item(self, tmp_path):
        root = mo.Model("m")
        root.default_excel_view = mo.ExcelView(orient="down")
        root.rate = mo.Variable(2.0, display_name="Rate")
        root.row = mo.Variable(
            [1.0, -root.rate, -root.rate], var_type="list", display_name="Row",
        )
        root.to_excel(tmp_path / "down.xlsx")
        wb = load_workbook(tmp_path / "down.xlsx")
        grid = self._grid(wb[wb.sheetnames[0]])
        cells = [v for v in grid.values() if isinstance(v, str) and v.startswith("=")]
        assert cells, "expected formula cells in transposed layout"
        assert all("[" not in v for v in cells)

    def test_short_list_beside_longer_sibling_emits_own_cells_only(self, tmp_path):
        # The writer loops over the variable's OWN address list, so a 2-item
        # row beside a 4-period sibling writes exactly its 2 cells — item
        # per cell, no bracket garbage, no spill into the sibling's columns.
        root = mo.Model("m")
        root.rate = mo.Variable(2.0, display_name="Rate")
        root.s = mo.Model("s", display_name="S")
        root.s.wide = mo.Variable([1, 2, 3, 4], var_type="list", display_name="Wide")
        root.s.narrow = mo.Variable(
            [0.0, -root.rate], var_type="list", display_name="Narrow",
        )
        root.to_excel(tmp_path / "clamp.xlsx")
        grid = self._grid(load_workbook(tmp_path / "clamp.xlsx")["S"])
        assert grid["B2"] == 0
        assert grid["C2"] == "=-M!B1"      # rate lives on the root overview sheet
        assert "D2" not in grid and "E2" not in grid


class TestInlinedListExprPrecedence:
    """A floating (unaddressed) list Variable inlined into a larger formula
    renders one positional item per period; the precedence helpers must
    parenthesize whenever ANY period's item binds looser than the parent
    op — a dropped paren is a silently wrong number, worse than the loud
    bracket garbage this replaced."""

    def _grid(self, ws):
        return {
            c.coordinate: c.value
            for row in ws.iter_rows() for c in row if c.value is not None
        }

    def test_mixed_precedence_items_parenthesize_every_period(self, tmp_path):
        m = mo.Model("m")
        m.a = mo.Variable(2.0, display_name="A")
        m.b = mo.Variable(3.0, display_name="B")
        m.c = mo.Variable(4.0, display_name="C")
        m.d = mo.Variable(5.0, display_name="D")
        m.x = mo.Variable([10.0, 20.0], var_type="list", display_name="X")
        floating = mo.Variable([m.a * m.b, m.c + m.d], var_type="list")
        m.y = (m.x * floating).set_display_name("Y")
        m.to_excel(tmp_path / "prec.xlsx")
        wb = load_workbook(tmp_path / "prec.xlsx")
        grid = self._grid(wb[wb.sheetnames[0]])
        # period 1's item is c+d — binds looser than '*', parens REQUIRED
        assert grid["C6"] == "=C5 * (B3 + B4)"
        assert m.y._value == [60.0, 180.0]

    def test_division_by_floating_list_keeps_parens(self, tmp_path):
        m = mo.Model("m")
        m.p = mo.Variable(2.0, display_name="P")
        m.q = mo.Variable(3.0, display_name="Q")
        m.num = mo.Variable(12.0, display_name="Num")
        divisor = mo.Variable([m.p + m.q, m.p + m.q], var_type="list")
        m.res = (m.num / divisor).set_display_name("Res")
        m.to_excel(tmp_path / "div.xlsx")
        wb = load_workbook(tmp_path / "div.xlsx")
        grid = self._grid(wb[wb.sheetnames[0]])
        assert grid["B4"] == "=B3 / (B1 + B2)"
        assert m.res._value[0] == 2.4


class TestSliceEmission:
    """Slices of list Variables: element-wise contexts emit the CURRENT
    period's cell inside the sliced window; aggregate args emit the
    sliced RANGE. Regression: the element-wise case used to stamp one
    identical full-range formula (``=Input!B1:U1``) into every cell —
    #SPILL! chaos in Excel 365, accidental implicit intersection in
    legacy Excel."""

    @staticmethod
    def _grid(ws):
        return {
            c.coordinate: c.value
            for row in ws.iter_rows()
            for c in row
            if c.value is not None
        }

    def _model(self):
        m = mo.Model("m")
        with m:
            m.input_sheet = mo.MultiVariable("Input")
            with m.input_sheet:
                m.input_sheet.pl_cogs = mo.Variable(
                    [float(i) for i in range(26)], var_type="list"
                )
            m.income_statement = mo.MultiVariable("IS")
            with m.income_statement:
                m.income_statement.cogs = m.input_sheet.pl_cogs[:20]
                m.income_statement.h1 = mo.SUM(m.input_sheet.pl_cogs[0:12])
        return m

    def test_elementwise_slice_emits_per_period_cells(self, tmp_path):
        m = self._model()
        m.to_excel(tmp_path / "slice.xlsx")
        wb = load_workbook(tmp_path / "slice.xlsx")
        grid = self._grid(wb["IS"])
        assert grid["B1"] == "=Input!B1"
        assert grid["C1"] == "=Input!C1"
        assert grid["U1"] == "=Input!U1"  # 20th period of the window
        assert m.income_statement.cogs._value == [float(i) for i in range(20)]

    def test_slice_inside_aggregate_emits_the_range(self, tmp_path):
        m = self._model()
        m.to_excel(tmp_path / "sliceagg.xlsx")
        wb = load_workbook(tmp_path / "sliceagg.xlsx")
        grid = self._grid(wb["IS"])
        assert grid["B2"] == "=SUM(Input!B1:M1)"
        assert m.income_statement.h1._value == sum(range(12))
