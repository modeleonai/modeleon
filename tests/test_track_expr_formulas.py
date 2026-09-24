"""Track-authored expressions keep their AST and emit live formulas.

``mo.Variable(plan=<row>)`` used to copy the operand's NUMBERS into the
plan track and forget where they came from — the workbook showed dead
literals for every such row. The role stash still holds the operand,
so the track's expression is recoverable: the emission copies render
it, and ``x.at(track=...)`` resolves to the coordinate's laid-out row.
"""

import os
import tempfile

import openpyxl

import modeleon as mo
from modeleon.compile.excel.view import ExcelView


def _book(model, name):
    path = os.path.join(tempfile.mkdtemp(), "t.xlsx")
    model.to_excel(path)
    return openpyxl.load_workbook(path)[name]


def _row_cells(ws, label):
    """Value cells (col B onward) of the row whose label is ``label``."""
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value == label:
            return [ws.cell(r, c).value for c in range(2, ws.max_column + 1)
                    if ws.cell(r, c).value is not None]
    raise AssertionError(f"row {label!r} not found; labels: "
                         f"{[ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]}")


def _tracked_model(mode):
    m = mo.Model("m", display_name="M",
                 tracks=mo.Tracks("forecast", "plan", "actual",
                                  blend=mo.blend(given="actual", follow="plan",
                                                 until="2026-02")),
                 default_grain="month", default_start="2026-01",
                 default_periods=3)
    m.default_excel_view = ExcelView(tracks=mode)
    m.s = mo.MultiVariable("S", excel_props={"tab": True})
    with m.s as s:
        s.driver = mo.Variable([1.0, 2.0, 3.0], display_name="Driver")
        s.rev = mo.Variable(
            forecast=[10.0, 10.0, 10.0],
            plan=s.driver * 2,
            actual=[9.0, None, None],
            display_name="Rev",
        )
    return m


def test_plan_expression_track_is_a_formula_row():
    ws = _book(_tracked_model("rows"), "S")
    plan_cells = _row_cells(ws, "Rev · plan")
    assert all(isinstance(c, str) and c.startswith("=") for c in plan_cells), plan_cells
    # the formula reads the driver's cell, not a literal
    assert all("2" in c for c in plan_cells)
    # data tracks stay values
    assert all(not (isinstance(c, str) and c.startswith("="))
               for c in _row_cells(ws, "Rev · forecast"))


def test_single_role_recurrence_track_shows_its_formula():
    m = mo.Model("m", display_name="M",
                 tracks=mo.Tracks("forecast", "actual",
                                  blend=mo.blend(given="actual", follow="forecast",
                                                 until="2026-01")),
                 default_grain="month", default_start="2026-01",
                 default_periods=4)
    m.default_excel_view = ExcelView(tracks="blend")
    m.s = mo.MultiVariable("S", excel_props={"tab": True})
    with m.s as s:
        s.g = mo.Variable([1.1, 1.1, 1.2, 1.2], display_name="Pace")
        s.traj = mo.Variable(
            forecast=mo.recurrence(100.0, "{prev} * {g}", g=s.g),
            display_name="Trajectory",
        )
        s.req = mo.Variable(forecast=s.traj / 15, display_name="Required")
    ws = _book(m, "S")
    traj = _row_cells(ws, "Trajectory")
    # period 0 is the seed, later periods chain on the previous cell
    assert isinstance(traj[1], str) and traj[1].startswith("=") and "*" in traj[1]
    req = _row_cells(ws, "Required")
    assert all(isinstance(c, str) and c.startswith("=") and "/ 15" in c for c in req)


def test_restrict_resolves_to_the_coordinate_row():
    m = _tracked_model("rows")
    with m.s as s:
        s.from_plan = mo.Variable(s.rev.at(track="plan") * 10,
                                  display_name="From plan")
        s.from_forecast = mo.Variable(s.rev.at(track="forecast") + 1,
                                      display_name="From forecast")
    ws = _book(m, "S")
    # locate the plan subrow and the forecast subrow
    rows = {ws.cell(r, 1).value: r for r in range(1, ws.max_row + 1)}
    plan_row, forecast_row = rows["Rev · plan"], rows["Rev · forecast"]
    fp = _row_cells(ws, "From plan")
    fl = _row_cells(ws, "From forecast")
    assert fp[0] == f"=B{plan_row} * 10", fp
    assert fl[0] == f"=B{forecast_row} + 1", fl


def test_restrict_on_the_shown_role_hits_the_head_row():
    # A single-role line: .at(track='forecast') IS the line's own row.
    m = mo.Model("m", display_name="M", tracks=mo.Tracks("forecast", "actual"),
                 default_grain="month", default_start="2026-01",
                 default_periods=2)
    m.default_excel_view = ExcelView(tracks="blend")
    m.s = mo.MultiVariable("S", excel_props={"tab": True})
    with m.s as s:
        s.base = mo.Variable(forecast=[5.0, 6.0], display_name="Base")
        s.twice = mo.Variable(forecast=s.base.at(track="forecast") * 2,
                              display_name="Twice")
    ws = _book(m, "S")
    rows = {ws.cell(r, 1).value: r for r in range(1, ws.max_row + 1)}
    assert _row_cells(ws, "Twice")[0] == f"=B{rows['Base']} * 2"


def test_blend_book_inlines_a_missing_coordinate_row():
    # blend mode lays out no per-track subrows: the restricted operand
    # degrades to its value, the rest of the formula stays live.
    m = _tracked_model("blend")
    with m.s as s:
        s.from_forecast = mo.Variable(s.rev.at(track="forecast") + s.driver,
                                      display_name="From forecast")
    ws = _book(m, "S")
    rows = {ws.cell(r, 1).value: r for r in range(1, ws.max_row + 1)}
    cell = _row_cells(ws, "From forecast")[0]
    assert cell == f"=10.0 + B{rows['Driver']}", cell


def test_json_renderer_has_a_restrict_node():
    from modeleon.compile.json.renderer import to_json
    m = _tracked_model("rows")
    node = m.s.rev.at(track="plan")._expr
    out = to_json(node)
    assert out["kind"] == "restrict" and out["label"] == "plan"


def test_untracked_book_unchanged():
    m = mo.Model("m", display_name="M", default_grain="month",
                 default_start="2026-01", default_periods=2)
    m.s = mo.MultiVariable("S", excel_props={"tab": True})
    with m.s as s:
        s.a = mo.Variable([1.0, 2.0], display_name="A")
        s.b = mo.Variable(s.a * 3, display_name="B")
    ws = _book(m, "S")
    rows = {ws.cell(r, 1).value: r for r in range(1, ws.max_row + 1)}
    assert _row_cells(ws, "B")[0] == f"=B{rows['A']} * 3"
