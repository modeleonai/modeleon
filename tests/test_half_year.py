# SPDX-License-Identifier: Apache-2.0
"""The half-year: a ruler of the gregorian calendar.

Six whole months, keyed ``YYYY-Hn``, written ``H1 2026``; a window of
its own, a bucket between quarter and year, a target of ``.at()`` —
every answer derived from the calendar, nothing per-grain added.
"""
from datetime import date

import pytest
from openpyxl import load_workbook

import modeleon as mo
from modeleon.compile.excel.header import spell_anchor, spell_inline
from modeleon.compile.excel.totals import build_totals_plan, display_label, header_rows, short_label
from modeleon.core.time import (
    _period_labels, _target_label, bucket_ranges, period_weights, time_field_series,
    window_start_label,
)


class TestTheRuler:
    def test_labels_in_three_styles(self):
        assert _period_labels('2026-H1', None, 3, 'half', 'finance') == ['H1 2026', 'H2 2026', 'H1 2027']
        assert _period_labels('2026-H2', None, 2, 'half') == ['2026-H2', '2027-H1']
        assert _period_labels('2026-H1', None, 2, 'half', 'compact') == ['H1-26', 'H2-26']

    def test_a_half_window_reads_its_time(self):
        m = mo.Model('m', default_grain='half', default_start='2026-H1', default_periods=4)
        with m:
            m.days = mo.time.days
            m.months = mo.time.months
            m.fee = mo.Variable(100.0) * mo.time.days
        assert m.days.value == [181, 184, 181, 184]
        assert m.months.value == [6, 6, 6, 6]
        assert m.fee.value == [18100.0, 18400.0, 18100.0, 18400.0]
        assert time_field_series('end', '2026-H1', 'half', 2) == [date(2026, 6, 30), date(2026, 12, 31)]

    def test_a_half_window_may_open_inside_its_first_half(self):
        # A valuation on 1 September: the first half is September to
        # December, keeps its label and counts its own days.
        m = mo.Model('m', default_grain='half', default_start=mo.DATE(2026, 9, 1), default_periods=3)
        with m:
            m.days = mo.time.days
        assert m.default_start == '2026-H2'
        assert m.days.value == [122, 181, 184]
        assert window_start_label(date(2026, 7, 1), 'half') == '2026-H2'
        assert window_start_label(date(2026, 2, 1), 'half') == '2026-H1'
        with pytest.raises(ValueError, match="A half window starts on the 1st of a month"):
            window_start_label(date(2026, 2, 15), 'half')

    def test_months_and_quarters_roll_into_halves_and_halves_into_years(self):
        assert bucket_ranges('month', 'half', '2026-11', 8) == [('2026-H2', 0, 2), ('2027-H1', 2, 8)]
        assert bucket_ranges('quarter', 'half', '2026-Q2', 4) == [('2026-H1', 0, 1), ('2026-H2', 1, 3), ('2027-H1', 3, 4)]
        assert bucket_ranges('half', 'year', '2026-H2', 3) == [('2026', 0, 1), ('2027', 1, 3)]
        assert _target_label((2026, 11), 'month', 'half') == '2026-H2'
        assert period_weights('half', '2026-H1', 2) == [181, 184]
        with pytest.raises(ValueError, match="Re-grain coarsens only: half -> quarter"):
            bucket_ranges('half', 'quarter', '2026-H1', 2)

    def test_at_half_sums_a_monthly_line(self):
        m = mo.Model('m', default_grain='month', default_start='2026-01', default_periods=12)
        with m:
            m.flow = mo.Variable([float(i + 1) for i in range(12)], regrain=mo.up('sum'))
            m.level = mo.Variable([10.0] * 12, regrain=mo.up('last'))
        assert m.flow.at('half').value == [21.0, 57.0]
        assert m.flow.at('half').time.start == '2026-H1'
        assert m.level.at('half').value == [10.0, 10.0]


class TestTheBook:
    def test_half_totals_interleave_between_quarters_and_the_year(self):
        plan = build_totals_plan(('quarter', 'half', 'year'), 'month', '2026-01', 12, 3)
        assert plan is not None
        kinds = [k for k, _i in plan.columns]
        # … Jun (Q2) (H1) Jul … Dec (Q4) (H2) (2026): the half closes
        # after the quarter that ends with it, the year after both.
        assert kinds[6:10] == ['period', 'quarter', 'half', 'period']
        assert kinds[-4:] == ['period', 'quarter', 'half', 'year']
        tiers = [row['cells'][0]['kind'] for row in header_rows(plan, _period_labels('2026-01', None, 12, 'month', 'finance'))]
        assert tiers == ['year', 'half', 'quarter', 'period']
        assert display_label('2026-H1', 'half', 'finance') == 'H1 2026'
        assert display_label('2026-H1', 'half', 'iso') == '2026-H1'
        assert short_label('2026-H2', 'half', 'finance', True) == 'H2'
        labels = [c['label'] for row in header_rows(plan, _period_labels('2026-01', None, 12, 'month', 'finance'))
                  for c in row['cells'] if c['kind'] == 'half']
        assert labels == ['H1', 'Total H1', 'H2', 'Total H2']

    def test_a_half_bucket_is_a_live_sum_over_its_months(self, tmp_path):
        m = mo.Model('m', default_grain='month', default_start='2026-01', default_periods=12,
                     default_excel_view=mo.ExcelView(timeline={'totals': ['half']}))
        with m:
            m.flow = mo.Variable([1.0] * 12, regrain=mo.up('sum'), display_name='Flow')
        path = tmp_path / 'half.xlsx'
        m.to_excel(path)
        ws = load_workbook(path)['M']
        row = next(r for r in ws.iter_rows() if r[0].value == 'Flow')
        formulas = [c.value for c in row if isinstance(c.value, str) and c.value.startswith('=')]
        assert formulas == ['=SUM(B3,C3,D3,E3,F3,G3)', '=SUM(I3,J3,K3,L3,M3,N3)']
        # The half is the top tier of this book: 'H1' spans January to June.
        assert [c.value for c in ws[1] if c.value][:1] == ['H1']

    def test_the_header_spells_a_half_window_by_six_months(self):
        assert spell_inline('end', 1, 'half', 'B2') == 'EOMONTH(B2,11)'
        assert spell_inline('months', 0, 'half', 'B2') == '6'
        cell = lambda f, i: None  # noqa: E731
        assert spell_anchor('end', 0, 'half', 'B2', cell) == 'EOMONTH(B2,5)'
        # A short first half: each period still ends with its half.
        assert spell_inline('end', 0, 'half', 'B2', date(2026, 9, 1)) == 'EOMONTH(B2,5-MOD(MONTH(B2)-1,6))'


def test_the_words_a_reader_sees():
    from modeleon.core.calendar import HALF
    assert (HALF.word, HALF.letter, HALF.grammar) == ('Half-yearly', 'H', "'YYYY-Hn'")
