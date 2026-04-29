# SPDX-License-Identifier: Apache-2.0
"""Golden-xlsx regression harness.

Runs each example script in a temp directory and diffs the produced .xlsx
against a committed baseline under ``tests/golden/``. Compares cell values
and formulas (data_type + value), ignores formatting/styles/metadata.

The harness is the regression net for the qualified-path identity
machinery: any change must leave every baseline bit-identical at the
formula/value level. A drift in any cell is a stop-the-line signal.

Bootstrap: if ``tests/golden/<name>.xlsx`` is missing, the test copies the
newly-generated file into ``golden/`` and xfails with a bootstrap message.
Commit the bootstrapped file to freeze the baseline.
"""

from __future__ import annotations

import shutil
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

REPO_ENGINE_DIR = Path(__file__).resolve().parent.parent
EXAMPLES_DIR = REPO_ENGINE_DIR / "examples"
GOLDEN_DIR = Path(__file__).resolve().parent / "golden"

# (example_script, output_filename_relative_to_cwd)
EXAMPLES = [
    ("01_pnl.py", "pnl.xlsx"),
    ("02_multi_sheet.py", "forecast.xlsx"),
]


def _run_example(script_name: str, cwd: Path) -> Path:
    script = EXAMPLES_DIR / script_name
    result = subprocess.run(
        [sys.executable, str(script)],
        cwd=cwd,
        capture_output=True,
        text=True,
        timeout=60,
    )
    if result.returncode != 0:
        raise RuntimeError(
            f"example {script_name} failed (rc={result.returncode}):\n"
            f"stdout:\n{result.stdout}\nstderr:\n{result.stderr}"
        )
    produced = cwd / Path(script_name).stem  # not used — keep explicit output name
    return produced


def _cell_signature(cell) -> tuple[str, object]:
    """Return (data_type, value) — enough to detect formula/value drift."""
    return (cell.data_type, cell.value)


def _workbook_signature(path: Path) -> dict[str, dict[str, tuple[str, object]]]:
    """{sheet_name: {cell_coord: (data_type, value)}}."""
    wb = load_workbook(path, data_only=False)
    out: dict[str, dict[str, tuple[str, object]]] = {}
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        cells: dict[str, tuple[str, object]] = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                cells[cell.coordinate] = _cell_signature(cell)
        out[sheet_name] = cells
    return out


def _diff_signatures(
    golden: dict[str, dict[str, tuple[str, object]]],
    actual: dict[str, dict[str, tuple[str, object]]],
) -> list[str]:
    diffs: list[str] = []
    all_sheets = set(golden) | set(actual)
    for sheet in sorted(all_sheets):
        if sheet not in golden:
            diffs.append(f"[{sheet}] sheet missing in golden")
            continue
        if sheet not in actual:
            diffs.append(f"[{sheet}] sheet missing in actual")
            continue
        g = golden[sheet]
        a = actual[sheet]
        all_cells = set(g) | set(a)
        for coord in sorted(all_cells):
            if coord not in g:
                diffs.append(f"[{sheet}!{coord}] cell absent in golden, actual={a[coord]!r}")
                continue
            if coord not in a:
                diffs.append(f"[{sheet}!{coord}] cell absent in actual, golden={g[coord]!r}")
                continue
            if g[coord] != a[coord]:
                diffs.append(f"[{sheet}!{coord}] golden={g[coord]!r} actual={a[coord]!r}")
    return diffs


@pytest.mark.parametrize("script,output", EXAMPLES, ids=[s for s, _ in EXAMPLES])
def test_example_xlsx_matches_golden(script: str, output: str, tmp_path: Path):
    _run_example(script, cwd=tmp_path)
    produced = tmp_path / output
    assert produced.exists(), f"example did not produce {output}"

    golden = GOLDEN_DIR / output
    if not golden.exists():
        GOLDEN_DIR.mkdir(parents=True, exist_ok=True)
        shutil.copy2(produced, golden)
        pytest.xfail(f"bootstrapped {golden.name} — commit tests/golden/ and re-run")

    golden_sig = _workbook_signature(golden)
    actual_sig = _workbook_signature(produced)
    diffs = _diff_signatures(golden_sig, actual_sig)
    assert not diffs, (
        f"{output} drifted from golden ({len(diffs)} cell(s)):\n  "
        + "\n  ".join(diffs[:20])
        + (f"\n  ... {len(diffs) - 20} more" if len(diffs) > 20 else "")
    )
