# SPDX-License-Identifier: Apache-2.0
"""The look of a hand-built book, declared once on a view.

Every capability here is off until a view asks for it — a model that
declares none writes the book it always wrote. Pinned: the book's font,
the look of a cell by what it is, formats by what a line is, the rows
repeated at the top of every sheet, a title row, the opening column
("Pre-start") the time rows chain off, full-width rows, column widths,
and the house view an application lays under a model without touching it.
"""

import warnings
from datetime import date

import openpyxl
import pytest

import modeleon as mo
from modeleon.compile.excel.view import resolve_excel_view

GREY = "#7F7F7F"


def deal(view=None):
    """Three sheets like a project-finance book: inputs, series, and
    the time anchor the others repeat at their top."""
    m = mo.Model("m", default_grain="year", default_periods=3)
    with m:
        if view is not None:
            m.default_excel_view = view
        m.ti = mo.MultiVariable("Inputs")
        with m.ti as i:
            i.start = mo.Variable(date(2026, 1, 1), display_name="Model start")
            i.rate = mo.Variable(0.05, display_name="Rate")
        m.default_start = m.ti.start
        m.td = mo.MultiVariable("Series")
        with m.td as s:
            s.volume = mo.Variable([1.0, 2.0, 3.0], display_name="Volume",
                                   excel_props={"row_sum": True})
        m.timing = mo.MultiVariable("Timing")
        with m.timing as t:
            t.ruler = mo.MultiVariable("Timeline")
            with t.ruler as r:
                r.index = mo.time.index.set_display_name("Period number")
                r.index.unit = "#"
                r.start = mo.time.start.set_display_name("Period start")
                r.end = mo.time.end.set_display_name("Period end")
                r.end.unit = "date"
            t.flags = mo.MultiVariable("Flags")
            with t.flags as f:
                f.first = mo.IF(m.timing.ruler.index == 1, 1, 0).set_display_name(
                    "First period flag")
        m.default_header = m.timing.ruler
    return m


def book(model, tmp_path):
    path = tmp_path / "book.xlsx"
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        model.to_excel(str(path))
    return openpyxl.load_workbook(str(path))


def cells(ws):
    return {c.coordinate: c.value for row in ws.iter_rows() for c in row
            if c.value is not None}


def column(ws, caption):
    """The column a header caption heads."""
    return next(c.column for row in ws.iter_rows(max_row=3) for c in row
                if c.value == caption)


LEAD = {"article": "No", "label": "Line item", "unit": "Unit", "constants": "Constant",
        "fields": [("row_sum", "Total")]}


# ══════════════════════════════════════════════════════════════════════
# Off by default
# ══════════════════════════════════════════════════════════════════════


def test_a_model_without_a_view_writes_the_book_it_always_did(tmp_path):
    wb = book(deal(), tmp_path)
    assert wb["Timing"]["A1"].font.name == "Calibri"
    assert wb["Timing"].sheet_view.showGridLines is not False
    assert wb["Inputs"]["A1"].value != "INPUTS"


# ══════════════════════════════════════════════════════════════════════
# Font, cell types, formats
# ══════════════════════════════════════════════════════════════════════


def test_the_view_font_is_the_book_font(tmp_path):
    wb = book(deal(mo.ExcelView(font={"name": "Arial", "size": 10}, meta=LEAD)), tmp_path)
    ws = wb["Timing"]
    # The header captions, the section bands and an empty cell alike.
    for ref in ("A1", "B1", "Z99"):
        assert (ws[ref].font.name, ws[ref].font.sz) == ("Arial", 10.0)
    assert wb._named_styles["Normal"].font.name == "Arial"


def test_an_input_wears_its_type_look_under_its_own(tmp_path):
    view = mo.ExcelView(cell_types={"on": True, "colors": {"input": "#0000FF"},
                                    "styles": {"input": {"bg": "#FFF2CC"}}})
    m = deal(view)
    m.ti.rate.set_style(bg="#FFFFFF")
    ws = book(m, tmp_path)["Inputs"]
    start = next(c for row in ws.iter_rows() for c in row if c.value == date(2026, 1, 1)
                 or getattr(c.value, "date", lambda: None)() == date(2026, 1, 1))
    assert start.fill.fgColor.rgb == "FFFFF2CC" and start.font.color.rgb == "FF0000FF"
    rate = next(c for row in ws.iter_rows() for c in row if c.value == 0.05)
    assert rate.fill.fgColor.rgb == "FFFFFFFF"          # the line's own look wins


def test_formats_follow_what_a_line_is(tmp_path):
    view = mo.ExcelView(meta=LEAD, formats={"date": "dd mmm yyyy", "flag": "0", "#": "0",
                                            "number": "#,##0.0"})
    ws = book(deal(view), tmp_path)["Timing"]
    first = column(ws, "Total") + 1          # the Timing sheet has no constants
    by_label = {ws.cell(r, 2).value: ws.cell(r, first) for r in range(2, ws.max_row + 1)}
    assert by_label["Period end"].number_format == "dd mmm yyyy"      # a date
    assert by_label["Period number"].number_format == "0"              # its unit


# ══════════════════════════════════════════════════════════════════════
# The rows at the top of every sheet
# ══════════════════════════════════════════════════════════════════════


def test_header_rows_wear_the_mirror_band_and_their_unit(tmp_path):
    view = mo.ExcelView(meta=LEAD, bands={"mirror": {"italic": True, "font_color": GREY}},
                        formats={"date": "dd mmm yyyy"})
    ws = book(deal(view), tmp_path)["Series"]
    got = cells(ws)
    row = next(r for r in range(1, 10) if got.get(f"B{r}") == "Period end")
    assert got[f"C{row}"] == "date"
    value = ws.cell(row, 6)      # first period after Constant: No, label, unit, Total, Constant
    assert value.font.i and value.font.color.rgb == "FF7F7F7F"
    assert value.number_format == "dd mmm yyyy"      # the format of the row it repeats


def test_same_columns_keeps_the_constant_column_on_every_sheet(tmp_path):
    view = mo.ExcelView(meta=LEAD, sheet={"same_columns": True})
    wb = book(deal(view), tmp_path)
    starts = {name: next(c.column for c in wb[name][1] if c.value == "Constant")
              for name in ("Inputs", "Series", "Timing")}
    assert len(set(starts.values())) == 1


def test_a_sheet_of_constants_heads_its_columns_without_periods(tmp_path):
    view = mo.ExcelView(meta=LEAD, sheet={"same_columns": True, "title": "upper"})
    m = deal(view)
    m.ti.default_header = None                   # constants: no time rows on top
    wb = book(m, tmp_path)
    inputs, series = wb["Inputs"], wb["Series"]
    captions = ["No", "Line item", "Unit", "Total", "Constant"]
    assert [inputs.cell(2, c).value for c in range(1, 6)] == captions
    # The same words in the same places as on a sheet with periods.
    assert [series.cell(2, c).value for c in range(1, 6)] == captions
    # …and no periods: nothing right of the constants column.
    assert all(inputs.cell(2, c).value is None for c in range(6, 10))
    got = cells(inputs)
    first = min(int(ref[1:]) for ref, v in got.items() if v == "Model start")
    assert first > 2                             # the data stands under the captions
    assert got[f"E{first}"] is not None          # in the constants column
    assert inputs.freeze_panes.endswith("3")     # the captions stay in view
    assert inputs["A2"].font.b


def test_without_same_columns_a_sheet_of_constants_keeps_row_1_for_data(tmp_path):
    m = deal(mo.ExcelView(meta=LEAD))
    m.ti.default_header = None
    got = cells(book(m, tmp_path)["Inputs"])
    assert "Line item" not in got.values()
    assert min(int(ref[1:]) for ref, v in got.items() if v == "Model start") == 1


def test_a_scalar_header_row_sits_in_the_constant_column(tmp_path):
    m = deal(mo.ExcelView(meta=LEAD))
    with m:
        m.timing.flags.status = mo.IF(m.ti.rate > 0, "OK", "ERROR").set_display_name("Status")
    m.td.default_header = [m.timing.ruler.end, m.timing.flags.status]
    ws = book(m, tmp_path)["Series"]
    got = cells(ws)
    assert got["E1"] == "Constant"
    status = next(r for r in range(1, 10) if got.get(f"B{r}") == "Status")
    assert str(got[f"E{status}"]).startswith("=Timing!")


def test_the_anchor_can_repeat_its_own_rows_at_its_top(tmp_path):
    view = mo.ExcelView(timeline={"repeat_on_anchor": True}, meta={"unit": True})
    ws = book(deal(view), tmp_path)["Timing"]
    got = cells(ws)                              # label | unit | 2026 …
    labels = [got.get(f"A{r}") for r in range(2, 5)]
    assert labels == ["Period number", "Period start", "Period end"]
    own = {got.get(f"A{r}"): r for r in range(5, 20)}
    assert got["C4"] == f"=C{own['Period end']}"     # a reference to the row below
    # A formula on the anchor still reads the row itself, not its mirror.
    assert f"C{own['Period number']}" in got[f"C{own['First period flag']}"]
    assert ws.freeze_panes == "C5"               # under the repeated rows


# ══════════════════════════════════════════════════════════════════════
# Title row and the opening column
# ══════════════════════════════════════════════════════════════════════


def test_a_title_row_heads_every_sheet(tmp_path):
    view = mo.ExcelView(meta={**LEAD, "fields": [("description", "Note")]},
                        sheet={"title": "upper"},
                        bands={"title_note": {"italic": True, "font_color": GREY}})
    m = deal(view)
    m.ti.set_description("Constant inputs")
    ws = book(m, tmp_path)["Inputs"]
    assert ws["B1"].value == "INPUTS" and ws["B1"].font.b
    assert ws["D1"].value == "Constant inputs" and ws["D1"].font.i     # the Note column
    assert ws["A2"].value == "No" and ws["B2"].value == "Line item"    # header one row down


def test_the_opening_column_holds_the_day_before_the_start(tmp_path):
    view = mo.ExcelView(meta=LEAD, timeline={"opening": "Pre-start"},
                        sheet={"same_columns": True})
    wb = book(deal(view), tmp_path)
    ws = wb["Timing"]
    got = cells(ws)
    # No | Line item | Unit | Total | Constant | Pre-start | 2026 …
    assert got["E1"] == "Constant" and got["F1"] == "Pre-start"
    rows = {got.get(f"B{r}"): r for r in range(2, 20)}
    idx, start, end = rows["Period number"], rows["Period start"], rows["Period end"]
    assert got[f"F{end}"] == "=Inputs!E6-1"        # the day before the model start
    assert got[f"F{idx}"] == 0
    # One formula across the row, the first period included.
    assert got[f"G{start}"] == f"=F{end}+1" and got[f"H{start}"] == f"=G{end}+1"
    assert got[f"G{idx}"] == f"=F{idx}+1"
    assert f"F{start}" not in got                  # a start has no opening value
    series = cells(wb["Series"])
    mirror = next(r for r in range(1, 10) if series.get(f"B{r}") == "Period end")
    assert series[f"F{mirror}"] == f"=Timing!F{end}"
    assert ws.freeze_panes.startswith("G")         # the opening column stays in view


# ══════════════════════════════════════════════════════════════════════
# Full-width rows, widths, gridlines
# ══════════════════════════════════════════════════════════════════════


def test_a_full_row_spans_the_sheet(tmp_path):
    view = mo.ExcelView(meta=LEAD, bands={"section": {"bg": "#D9E2F3"}},
                        formats={"number": "#,##0.0"},
                        sheet={"full_rows": True})
    m = deal(view)
    m.td.volume.set_style(bold=True, border_top=True)
    wb = book(m, tmp_path)
    ws = wb["Series"]
    got = cells(ws)
    row = next(r for r in range(1, 10) if got.get(f"B{r}") == "Volume")
    for col in range(1, ws.max_column + 1):
        assert ws.cell(row, col).border.top.style == "thin"
    total = ws.cell(row, 4)
    assert total.value.startswith("=SUM(") and total.font.b
    assert total.number_format == "#,##0.0"
    timing = wb["Timing"]
    band = next(r for r in range(1, 10) if timing.cell(r, 2).value == "Flags")
    assert timing.cell(band, 1).fill.fgColor.rgb == "FFD9E2F3"     # from the first column


def test_widths_by_role_and_no_gridlines(tmp_path):
    view = mo.ExcelView(meta=LEAD, sheet={
        "gridlines": False,
        "columns": [("label", {"width": 40}), ("unit", {"width": 9, "font_color": GREY}),
                    ("period", {"width": 12})]})
    ws = book(deal(view), tmp_path)["Series"]
    assert ws.sheet_view.showGridLines is False
    assert ws.column_dimensions["B"].width == 40 and ws.column_dimensions["C"].width == 9
    first = column(ws, "Total") + 1              # no constants on this sheet
    for col in range(first, first + 3):
        assert ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width == 12
    unit = next(c for c in ws["C"] if c.value == "date")
    assert unit.font.color.rgb == "FF7F7F7F"


# ══════════════════════════════════════════════════════════════════════
# The house view
# ══════════════════════════════════════════════════════════════════════


def test_a_house_view_dresses_a_model_that_declares_nothing(tmp_path):
    house = mo.ExcelView(font={"name": "Arial", "size": 10}, sheet={"gridlines": False})
    with mo.house_view(house):
        wb = book(deal(), tmp_path)
    assert wb["Timing"]["A1"].font.name == "Arial"
    assert wb["Timing"].sheet_view.showGridLines is False
    # Outside the block the floor is back.
    assert resolve_excel_view(None).sheet_gridlines is None


def test_the_model_s_own_view_wins_over_the_house(tmp_path):
    house = mo.ExcelView(font={"name": "Arial", "size": 10}, meta=LEAD)
    team = mo.ExcelView(sheet={"gridlines": False})
    m = deal(mo.ExcelView(font={"name": "Georgia", "size": 12}))
    with mo.house_view(house, team):
        view = resolve_excel_view(m.ti.rate)
    assert view.base_font == {"name": "Georgia", "size": 12}     # the model
    assert view.sheet_gridlines is False                          # the team
    assert view.meta_fields == (("row_sum", "Total"),)            # the house, kept


def test_a_nested_view_without_fields_keeps_its_parent_s():
    m = deal(mo.ExcelView(meta=LEAD))
    m.td.default_excel_view = mo.ExcelView(font={"name": "Arial"})
    assert resolve_excel_view(m.td.volume).meta_fields == (("row_sum", "Total"),)


def test_house_view_refuses_what_is_not_a_view():
    with pytest.raises(TypeError, match="ExcelView"):
        with mo.house_view({"font": {"name": "Arial"}}):
            pass


def test_a_model_names_the_rows_every_sheet_opens_on(tmp_path):
    """A header LIST on the model: every sheet shows those rows; the
    section of its first time row still anchors the time."""
    m = deal(mo.ExcelView(timeline={"repeat_on_anchor": True}))
    m.default_header = [m.timing.ruler.end, m.timing.flags.first]
    wb = book(m, tmp_path)
    for name in ("Inputs", "Series", "Timing"):
        got = cells(wb[name])
        assert [got.get(f"A{r}") for r in (2, 3)] == ["Period end", "First period flag"]
    timing = cells(wb["Timing"])
    rows = {timing.get(f"A{r}"): r for r in range(4, 20)}
    # The ruler still chains off the start cell: the anchor is its section.
    assert timing[f"C{rows['Period start']}"] == f"=B{rows['Period end (date)']}+1"


# ══════════════════════════════════════════════════════════════════════
# What a line is: its look, its number, its total
# ══════════════════════════════════════════════════════════════════════


def _linked(view):
    m = deal(view)
    with m:
        m.timing.dates = mo.MultiVariable("Key dates")
        with m.timing.dates as d:
            d.start = m.ti.start.copy().set_display_name("Model start")
            d.close = mo.EDATE(m.timing.dates.start, 6).set_display_name("Financial close")
            d.cod = mo.EDATE(m.timing.dates.close, 18).set_display_name(
                "Commercial operation").set_style(result=True)
    return m


def test_only_a_marked_line_wears_the_result_look(tmp_path):
    view = mo.ExcelView(meta=LEAD, bands={"result": {"bold": True, "border_top": True},
                                          "link": {"font_color": GREY}})
    ws = book(_linked(view), tmp_path)["Timing"]
    got = cells(ws)
    rows = {got.get(f"B{r}"): r for r in range(2, 30)}
    assert not ws[f"B{rows['Financial close']}"].font.b          # computed, not marked
    assert ws[f"B{rows['Commercial operation']}"].font.b          # marked
    assert ws[f"B{rows['Commercial operation']}"].border.top.style == "thin"
    assert ws[f"B{rows['Model start']}"].font.color.rgb == "FF7F7F7F"   # a link


def test_a_link_to_a_result_is_no_result_of_its_own(tmp_path):
    # The next block opens by reading the result: the link carries the
    # source's excel_props, mark included, but it is not this block's result.
    view = mo.ExcelView(meta=LEAD, bands={"result": {"bold": True, "border_top": True}})
    m = _linked(view)
    with m:
        m.timing.use = mo.MultiVariable("Uses")
        with m.timing.use as u:
            u.cod = m.timing.dates.cod.copy().set_display_name("COD (link)")
    ws = book(m, tmp_path)["Timing"]
    got = cells(ws)
    rows = {got.get(f"B{r}"): r for r in range(2, 40)}
    assert ws[f"B{rows['Commercial operation']}"].font.b
    assert not ws[f"B{rows['COD (link)']}"].font.b
    assert ws[f"B{rows['COD (link)']}"].border.top.style is None


def test_lines_are_numbered_by_section_and_links_keep_none(tmp_path):
    view = mo.ExcelView(meta={**LEAD, "numbering": "section"})
    got = cells(book(_linked(view), tmp_path)["Timing"])
    numbers = {got[f"B{r}"]: got.get(f"A{r}") for r in range(2, 30) if f"B{r}" in got}
    assert [numbers[s] for s in ("Timeline", "Flags", "Key dates")] == ["1", "2", "3"]
    assert numbers["Period number"] == "1.1" and numbers["Period end"] == "1.3"
    assert numbers["First period flag"] == "2.1"
    assert numbers["Model start"] is None                         # a link has no number
    assert numbers["Financial close"] == "3.1" and numbers["Commercial operation"] == "3.2"


def _site(view):
    """A sheet three levels deep: a site, its blocks, their lines — and a
    line on the sheet itself."""
    m = mo.Model("m", default_grain="year", default_start="2026", default_periods=3)
    with m:
        m.default_excel_view = view
        m.ops = mo.MultiVariable("Operations")
        with m.ops as ops:
            ops.days = mo.Variable(365.0, display_name="Days in year")
            ops.site_a = mo.MultiVariable("Site A")
            with ops.site_a as sa:
                sa.rev = mo.MultiVariable("Sales")
                with sa.rev as r:
                    r.units = mo.Variable([10.0, 20.0, 30.0], display_name="Units")
                    r.price = mo.Variable(5.0, display_name="Price")
                    r.total = (r.units * r.price).set_display_name("Revenue")
                sa.costs = mo.MultiVariable("Operating costs")
                with sa.costs as c:
                    c.rev = sa.rev.total.copy().set_display_name("Revenue (link)")
                    c.opex = (c.rev * 0.1).set_display_name("Opex")
                sa.margin = (sa.rev.total - sa.costs.opex).set_display_name("EBITDA")
            ops.other = mo.MultiVariable("Site B", excel_props={"article": "B"})
            with ops.other as o:
                o.x = mo.Variable([1.0, 2.0, 3.0], display_name="Other line")
    return m


def _numbers(ws):
    """Label → the number in column A, for every row with a label."""
    got = cells(ws)
    return {got[f"B{r}"]: got.get(f"A{r}") for r in range(2, ws.max_row + 1) if f"B{r}" in got}


def test_numbers_follow_the_outline_to_any_depth(tmp_path):
    # A section takes a number like the lines beside it, and its lines are
    # numbered one level below it — at every depth, not just the first.
    view = mo.ExcelView(meta={**LEAD, "numbering": "section"})
    numbers = _numbers(book(_site(view), tmp_path)["Operations"])
    assert numbers["Days in year"] == "1"                 # a line on the sheet itself
    assert numbers["Site A"] == "2"
    assert numbers["Sales"] == "2.1"
    assert [numbers[k] for k in ("Units", "Price", "Revenue")] == [
        "2.1.1", "2.1.2", "2.1.3"]
    assert numbers["Operating costs"] == "2.2"
    assert numbers["Revenue (link)"] is None and numbers["Opex"] == "2.2.1"
    assert numbers["EBITDA"] == "2.3"                     # a line after two sections
    assert numbers["Site B"] == "B" and numbers["Other line"] == "B.1"


def test_a_sections_number_wears_its_heading_look(tmp_path):
    view = mo.ExcelView(meta={**LEAD, "numbering": "section"},
                        bands={"section": {"bold": True, "bg": "#D9E2F3"}})
    ws = book(_site(view), tmp_path)["Operations"]
    row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 2).value == "Site A")
    number = ws.cell(row, 1)
    assert number.value == "2" and number.font.b
    assert number.fill.fgColor.rgb == "FFD9E2F3"


def test_a_row_riding_another_or_a_table_takes_one_number_or_none(tmp_path):
    # A total printed on its section's own row is that row: it takes no
    # number of its own. A table of records is one item, and the rows of
    # a matrix of instances are numbered like lines.
    m = mo.Model("m", default_grain="year", default_start="2026", default_periods=2)
    with m:
        m.default_excel_view = mo.ExcelView(meta={**LEAD, "numbering": "section"})
        m.book = mo.MultiVariable("Book")
        with m.book as b:
            b.charges = mo.MultiVariable("Charges")
            with b.charges as c:
                c.rent = mo.Variable([5.0, 5.0], display_name="Rent")
                c.power = mo.Variable([3.0, 3.0], display_name="Power")
                c.total = mo.Variable(c.rent + c.power, excel_props={"header_row": True})
            b.staff = mo.MultiVariable("Staff", excel_props={"orient": "records"})
            with b.staff as s:
                s.e1 = mo.MultiVariable("Employee 1")
                s.e1.salary = mo.Variable(100.0, display_name="Salary")
                s.e2 = mo.MultiVariable("Employee 2")
                s.e2.salary = mo.Variable(200.0, display_name="Salary")
            b.board = mo.MultiVariable("Board", excel_props={"orient": "columns"})
            with b.board as d:
                d.p1 = mo.MultiVariable("Plant 1")
                d.p1.capacity = mo.Variable(10.0, display_name="Capacity")
                d.p1.capex = mo.Variable(50.0, display_name="Capex")
                d.p2 = mo.MultiVariable("Plant 2")
                d.p2.capacity = mo.Variable(20.0, display_name="Capacity")
                d.p2.capex = mo.Variable(90.0, display_name="Capex")
            b.after = mo.Variable([1.0, 1.0], display_name="After the tables")
    ws = book(m, tmp_path)["Book"]
    numbers = _numbers(ws)
    assert numbers["Charges"] == "1"
    assert numbers["Rent"] == "1.1" and numbers["Power"] == "1.2"
    assert numbers["Staff"] == "2"
    assert numbers["Employee 1"] is None and numbers["Employee 2"] is None
    assert numbers["Board"] == "3"
    assert numbers["Capacity"] == "3.1" and numbers["Capex"] == "3.2"
    assert numbers["After the tables"] == "4"
    # Nothing else in the number column: no field of the table numbered
    # onto its caption row.
    assert sorted(v for v in numbers.values() if v) == sorted(
        ["1", "1.1", "1.2", "2", "3", "3.1", "3.2", "4"])


def test_row_totals_follow_what_a_line_is(tmp_path):
    # A line totals by its re-grain rule (``regrain=mo.up('sum')``); a date
    # line never does.
    view = mo.ExcelView(meta={**LEAD, "row_sum_for": ["sum"]})
    m = deal(view)
    m.td.volume.excel_props.pop("row_sum")
    m.td.volume._regrain = mo.up("sum")
    got = cells(book(m, tmp_path)["Series"])
    rows = {got.get(f"B{r}"): r for r in range(2, 30)}
    assert any(str(v).startswith("=SUM(") for k, v in got.items()
               if k[1:] == str(rows["Volume"]))                        # a sum
    timing = cells(book(m, tmp_path)["Timing"])
    trows = {timing.get(f"B{r}"): r for r in range(2, 30)}
    assert not any(str(v).startswith("=SUM(") for k, v in timing.items()
                   if k[1:] == str(trows["Period end"]))               # a date


def test_a_bar_rows_total_stays_a_number(tmp_path):
    """A row drawn as a bar (``'"X";;'`` marks any positive value) totals
    to a count; the total takes the row's format only when that format
    would show a number."""
    from modeleon.compile.excel.writer import format_shows_a_number

    assert format_shows_a_number('#,##0;[Red](#,##0);')
    assert format_shows_a_number('0.##;-0.##;')
    assert format_shows_a_number("General")
    assert format_shows_a_number('0%')
    assert not format_shows_a_number('"X";;')
    assert not format_shows_a_number('"yes";"no";"-"')

    m = mo.Model("m", default_grain="year", default_start="2026", default_periods=4)
    m.default_excel_view = mo.ExcelView(meta=LEAD)
    m.s = mo.MultiVariable("S")
    with m.s as s:
        s.active = mo.Variable([1, 1, 0, 1], display_name="Active",
                               excel_props={"number_format": '"X";;', "row_sum": True})
        s.cash = mo.Variable([10, 20, 30, 40], display_name="Cash",
                             excel_props={"number_format": "#,##0", "row_sum": True})
    ws = book(m, tmp_path)["S"]
    label, total = column(ws, "Line item"), column(ws, "Total")
    totals = {ws.cell(r, label).value: ws.cell(r, total) for r in range(1, ws.max_row + 1)}
    assert totals["Active"].value.startswith("=SUM(")
    assert totals["Active"].number_format == "General"
    assert totals["Cash"].number_format == "#,##0"


# ══════════════════════════════════════════════════════════════════════
# A named view, applied
# ══════════════════════════════════════════════════════════════════════


def test_an_applied_named_view_dresses_the_book_and_ends_with_the_block(tmp_path):
    m = deal(mo.ExcelView(font={"name": "Arial", "size": 10}, sheet={"gridlines": False}))
    m.board = mo.ExcelView("Board pack", font={"name": "Georgia", "size": 11})
    # A named view on the model changes nothing until applied: the book is the model's own.
    assert book(m, tmp_path)["Series"]["A1"].font.name == "Arial"
    with mo.applied_view(m.board):
        view = resolve_excel_view(m.td.volume)
        assert view.base_font == {"name": "Georgia", "size": 11}     # the applied view
        assert view.sheet_gridlines is False                          # the model's, kept
        wb = book(m, tmp_path)
    assert wb["Series"]["A1"].font.name == "Georgia"
    assert wb["Series"].sheet_view.showGridLines is False
    # Outside the block the model's own look is back.
    assert resolve_excel_view(m.td.volume).base_font == {"name": "Arial", "size": 10}
    assert book(m, tmp_path)["Series"]["A1"].font.name == "Arial"


def test_an_applied_view_lies_over_the_house_and_the_model(tmp_path):
    house = mo.ExcelView(font={"name": "Arial", "size": 10}, meta=LEAD)
    m = deal(mo.ExcelView(font={"name": "Georgia", "size": 12}))
    m.print = mo.ExcelView("Print", font={"name": "Verdana", "size": 9})
    with mo.house_view(house), mo.applied_view(m.print):
        view = resolve_excel_view(m.ti.rate)
    assert view.base_font == {"name": "Verdana", "size": 9}          # the applied view
    assert view.meta_fields == (("row_sum", "Total"),)               # the house, kept


def test_an_applied_view_s_sheet_order_wins(tmp_path):
    m = deal()
    m.series_first = mo.ExcelView("Series first", sheet={"order": ["td", "ti", "timing"]})
    assert book(m, tmp_path).sheetnames == ["Inputs", "Series", "Timing"]
    with mo.applied_view(m.series_first):
        assert book(m, tmp_path).sheetnames == ["Series", "Inputs", "Timing"]


def test_an_applied_view_s_orient_is_not_a_look():
    m = deal()
    m.side = mo.ExcelView("Side", orient="records", font={"name": "Georgia"})
    own = resolve_excel_view(m.td.volume).orient
    with mo.applied_view(m.side):
        view = resolve_excel_view(m.td.volume)
    assert view.base_font == {"name": "Georgia"}
    assert view.orient == own and view.orient != "records"


def test_applied_view_refuses_what_is_not_a_view():
    with pytest.raises(TypeError, match="ExcelView"):
        with mo.applied_view({"font": {"name": "Georgia"}}):
            pass
