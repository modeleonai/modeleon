# SPDX-License-Identifier: Apache-2.0
"""YEARFRAC, DAYS, IRR and XNPV compute what Excel computes.

The expected values in ``data/excel_functions_measured.json`` were
recalculated by Microsoft Excel itself (16.95.1, macOS) from the same
inputs written as native formulas: every YEARFRAC basis over month ends,
29 February, leap years and reversed dates, DAYS over the same pairs,
IRR over ordinary, tiny, huge and several-sign-change flows, and XNPV
including the rates and the early dates Excel refuses.
"""

from __future__ import annotations

import json
import math
from datetime import date
from pathlib import Path

import pytest

import modeleon as mo

DATA = json.loads((Path(__file__).parent / "data" / "excel_functions_measured.json").read_text())


def _d(s: str) -> date:
    return date.fromisoformat(s)


def _value(fn):
    try:
        v = fn().value
    except (ValueError, ZeroDivisionError, OverflowError):
        return "ERR"
    if isinstance(v, float) and not math.isfinite(v):
        return "ERR"
    return v


def _same(got, want, tol: float) -> bool:
    if want == "ERR" or got == "ERR":
        return got == want
    return abs(float(got) - float(want)) <= tol * max(1.0, abs(float(want)))


@pytest.mark.parametrize("case", DATA["yearfrac"], ids=lambda c: f"{c[0]}-{c[1]}-b{c[2]}")
def test_yearfrac_matches_excel(case) -> None:
    a, b, basis, want = case
    assert _same(_value(lambda: mo.YEARFRAC(_d(a), _d(b), basis)), want, 1e-12), case


@pytest.mark.parametrize("case", DATA["days"], ids=lambda c: f"{c[0]}-{c[1]}")
def test_days_matches_excel(case) -> None:
    a, b, want = case
    assert _value(lambda: mo.DAYS(_d(b), _d(a))) == want, case


@pytest.mark.parametrize("case", DATA["irr"], ids=lambda c: str(len(c[0])))
def test_irr_matches_excel(case) -> None:
    cf, want = case
    assert _same(_value(lambda: mo.IRR(cf)), want, 1e-7), case


@pytest.mark.parametrize("case", DATA["xnpv"], ids=lambda c: f"r{c[0]}")
def test_xnpv_matches_excel(case) -> None:
    rate, cf, ds, want = case
    assert _same(_value(lambda: mo.XNPV(rate, cf, [_d(s) for s in ds])), want, 1e-9), case


def test_days_is_written_with_its_file_prefix(tmp_path) -> None:
    """Excel 2013+ functions live in the file as ``_xlfn.DAYS``; written
    bare, Excel shows #NAME?."""
    from openpyxl import load_workbook

    m = mo.Model("m")
    m.start = mo.Variable(date(2024, 1, 31))
    m.end = mo.Variable(date(2024, 3, 1))
    m.n = mo.DAYS(m.end, m.start)
    m.f = mo.YEARFRAC(m.start, m.end, 1)
    path = tmp_path / "days.xlsx"
    m.to_excel(str(path))
    ws = load_workbook(str(path)).active
    formulas = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert any(f.startswith("=_xlfn.DAYS(") for f in formulas), formulas
    assert any(f.startswith("=YEARFRAC(") for f in formulas), formulas
