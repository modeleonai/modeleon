# SPDX-License-Identifier: Apache-2.0
"""Renderer fallback-to-value contract.

When a renderer encounters an AST node it cannot natively render — a
``FuncCall`` whose ``render_backends`` explicitly excludes the current
renderer — it must fall back to inlining ``_value`` as a literal in its
target format, not crash.

This locks in the graceful-degradation contract that makes
backend-specific functions (pandas-native ``EWMA``, SQL-only ``LAG``,
etc.) usable in other renderers. See ADR-010.
"""

from __future__ import annotations

import warnings

import modeleon as mo
from modeleon.core.expr import FuncCall, VarRef


def _make_backend_specific_var(name: str, value, render_backends):
    """Build a Variable whose ``_expr`` is a ``FuncCall`` declared with
    ``render_backends``, mimicking a future backend-specific function
    like pandas ``EWMA`` or SQL ``LAG``."""
    input_var = mo.Variable(value, display_name="input")
    result = mo.Variable(value, display_name="result")
    result._expr = FuncCall(
        func=name,
        args=[VarRef(input_var)],
        render_backends=frozenset(render_backends),
    )
    return input_var, result


class TestExcelFallback:
    """Excel renderer inlines value when FuncCall isn't Excel-native."""

    def test_backend_specific_funccall_inlines_value(self, tmp_path):
        input_var, smoothed = _make_backend_specific_var(
            "EWMA", 42.5, render_backends={"pandas"}
        )
        model = mo.MultiVariable("M")
        model.s = mo.MultiVariable("S", excel_props={'tab': True})
        with model.s as s:
                s.src = input_var
                s.smoothed = smoothed

        with warnings.catch_warnings(record=True):
            warnings.simplefilter("always")
            model.to_excel(tmp_path / "out.xlsx")
            path = tmp_path / "out.xlsx"
        from openpyxl import load_workbook
        wb = load_workbook(path)
        ws = wb["S"]
        formulas = [c.value for row in ws.iter_rows() for c in row
                    if c.value is not None]
        # The fallback should have inlined the literal 42.5 somewhere
        # instead of emitting "=EWMA(...)". The cell either holds the
        # raw value (no =) or a formula that contains "42.5".
        assert any(
            "EWMA" not in str(f) and ("42.5" in str(f) or f == 42.5)
            for f in formulas
        ), f"Expected inlined value, got formulas: {formulas}"


class TestJsonFallbackIsStructural:
    """JSON renderer always preserves structure — backend-specific
    functions serialize fully, including the ``render_backends`` hint."""

    def test_backend_specific_funccall_serializes_with_hint(self):
        _, smoothed = _make_backend_specific_var(
            "EWMA", 42.5, render_backends={"pandas", "sql"}
        )
        out = mo.to_json(smoothed._expr)
        assert out["kind"] == "funccall"
        assert out["func"] == "EWMA"
        assert out["render_backends"] == ["pandas", "sql"]

    def test_universal_funccall_omits_backend_fields(self):
        a = mo.Variable(10)
        b = mo.Variable(20)
        total = mo.SUM(a, b)
        out = mo.to_json(total._expr)
        # Universal functions don't surface either compat field.
        assert "render_backends" not in out
        assert "compute_backends" not in out
