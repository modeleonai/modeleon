# SPDX-License-Identifier: Apache-2.0
"""The text helpers write the spreadsheet's own calls — one per period
(``=LEFT(I!C2, 2)``), a scalar in one cell, ``TEXTJOIN`` over the row as
a range (``=_xlfn.TEXTJOIN(", ", TRUE, I!C2:E2)``) — and the cells, once
LibreOffice recalculates them, agree with the Python values.
"""

from __future__ import annotations

import logging
import shutil
import subprocess
import sys
from datetime import date, timedelta
from pathlib import Path

import pytest
from openpyxl import load_workbook

import modeleon as mo
from modeleon.core.errors import CrossScopeReferenceWarning
from modeleon.functions.text import (
    FIND,
    LEFT,
    MID,
    REPT,
    RIGHT,
    SEARCH,
    SUBSTITUTE,
    TEXT,
    TEXTJOIN,
    TRIM,
    VALUE,
)

PERIODS = 3

#: Every line the model writes on sheet X.
LINES = (
    "left", "left1", "leftnum", "leftdate", "leftnums", "leftbig", "leftbig2", "leftspan",
    "leftbool", "right", "mid", "midpast", "trim", "find", "find2", "findempty", "findempty2",
    "search", "search2", "searchempty", "searchempty2", "tilde", "sub", "sub2", "value", "value2",
    "valuebool", "valuebig", "valuetime", "valuetime2", "valuetime3", "text0", "text2", "textp",
    "textg", "textgbig", "textz", "textpad", "textgroup", "textd", "textd2", "textd3", "textm",
    "textb", "textlong", "textwide", "textempty", "rept", "reptbig", "lit", "join", "join2",
    "joinx", "quote", "chain", "err", "blank", "blanklist", "blankfind", "blanksub", "blankrept",
    "blankfmt", "blankjoin", "blanklen", "blankupper", "blankcat",
)

#: Where the engine and LibreOffice part — where Excel and LibreOffice
#: part, the engine follows Excel — the engine's answer, then
#: LibreOffice's. Both are pinned so a change on either side is seen.
DIVERGENT = {
    # a negative that rounds to zero keeps its sign
    "text0": (["1235", "-0", "3"], ["1235", "0", "3"]),
    # TRUE is the word, not 1
    "textb": (["TRUE", "FALSE", "TRUE"], ["1", "0", "1"]),
    "leftbool": (["TR", "FA", "TR"], ["1", "0", "1"]),
    # General text keeps 15 significant digits, writes a whole number out
    # in full below 1E20, and a fraction while it fits 20 characters;
    # LibreOffice keeps a 16th digit and writes three-digit exponents
    "leftbig": (["1234567890123460", "123456789012346000", "1.23456789012346E-05"],
                ["1234567890123456", "1.23456789012346E+017", "0.0000123456789012346"]),
    "textgbig": (["1234567890123460", "123456789012346000", "1.23456789012346E-05"],
                 ["1.23456789012346E+15", "1.23456789012346E+17", "0.0000123456789012346"]),
    # past those bounds E notation has two exponent digits at least;
    # LibreOffice writes three, and turns to E notation below about 1E-14
    "leftbig2": (["1.23456789012346E+20", "0.000000000000001", "-1E-19"],
                 ["1.23456789012346E+020", "1E-015", "-1E-019"]),
    # a number format shows 15 significant digits too; LibreOffice a 16th
    "textgroup": (["1,234,567,890,123,460", "123,456,789,012,346,000", "0"],
                  ["1,234,567,890,123,456", "123,456,789,012,346,000", "0"]),
    # empty find_text is found at start_num
    "findempty": ([1, 1, "#VALUE!"], ["#VALUE!", "#VALUE!", "#VALUE!"]),
    "searchempty2": ([2, 2, 2], ["#VALUE!", "#VALUE!", "#VALUE!"]),
    # VALUE reads no locale: a currency sign is not a number, nor is TRUE
    "value2": ([1234567, "#VALUE!", 46023], [1234567, 1234, 46023]),
    "valuebool": ("#VALUE!", 1),
    # nor a T between a date and its time, nor a negative time
    "valuetime2": (["#VALUE!", "#VALUE!", 0], [46053.5, -0.0625, 0]),
    # a time is hours, then minutes and seconds below 60; LibreOffice also
    # reads a bare "12:" and "0:90" (90 minutes) — Excel not measured
    "valuetime3": ([0.0625, "#VALUE!", "#VALUE!"], [0.0625, 0.5, 0.0625]),
    # a number too large for a cell; LibreOffice clamps to the largest double
    "valuebig": ("#NUM!", sys.float_info.max),
    # a text longer than a cell holds (32,767 characters); LibreOffice has no limit
    "reptbig": ("#VALUE!", "a" * 32768),
}


def _model():
    m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=PERIODS)
    with m:
        m.i = mo.MultiVariable("I")
        m.i.names = mo.Variable(["Alpha Fund", "beta co", None], display_name="names")
        m.i.codes = mo.Variable(["AB-123", "CD-45", "EF-6"], display_name="codes")
        m.i.amounts = mo.Variable([1234.5678, -0.4, 2.5], display_name="amounts")
        m.i.texts = mo.Variable(["1,234", "12%", "(1,234)"], display_name="texts")
        m.i.dates = mo.Variable([date(2026, 1, 31), date(2026, 2, 28), date(2026, 12, 25)],
                                display_name="dates")
        m.i.k = mo.Variable(3, display_name="k")
        m.i.sep = mo.Variable(" | ", display_name="sep")
        m.i.padded = mo.Variable("  a   b  ", display_name="padded")
        m.i.zero = mo.Variable(0.0, display_name="zero")
        m.i.bad = (m.i.amounts / m.i.zero).set_display_name("bad")
        m.i.big = (m.i.amounts > 1).set_display_name("big")
        m.i.nums = mo.Variable([1e15, 0.00001, 1 / 3], display_name="nums")
        m.i.texts2 = mo.Variable(["1234,567", "$1,234", "2026-01-01"], display_name="texts2")
        m.i.flag = mo.Variable(True, display_name="flag")
        m.i.bigs = mo.Variable([1234567890123456, 1.2345678901234568e17, 1.2345678901234568e-5],
                               display_name="bigs")
        m.i.times = mo.Variable(["1:30", "3:00 PM", "2026-01-31 18:00"], display_name="times")
        m.i.times2 = mo.Variable(["2026-01-31T12:00", "-1:30", "12:00 AM"], display_name="times2")
        m.i.prec = mo.Variable([1 / 3, 0.1 + 0.2, 2.675], display_name="prec")
        m.i.span = mo.Variable(timedelta(hours=36), display_name="span")
        m.i.bigs2 = mo.Variable([1.2345678901234568e20, 1e-15, -1e-19], display_name="bigs2")
        m.i.times3 = mo.Variable(["+1:30", "12:", "0:90"], display_name="times3")
        m.i.pads = mo.Variable([1234, 5, -1234.5], display_name="pads")
        m.x = mo.MultiVariable("X")
        m.x.left = LEFT(m.i.names, 2).set_display_name("left")
        m.x.left1 = LEFT(m.i.names).set_display_name("left1")
        m.x.leftnum = LEFT(m.i.amounts, 3).set_display_name("leftnum")
        m.x.leftdate = LEFT(m.i.dates, 4).set_display_name("leftdate")
        m.x.leftnums = LEFT(m.i.nums, 20).set_display_name("leftnums")
        m.x.leftbig = LEFT(m.i.bigs, 25).set_display_name("leftbig")
        m.x.leftbig2 = LEFT(m.i.bigs2, 25).set_display_name("leftbig2")
        m.x.leftspan = LEFT(m.i.span, 3).set_display_name("leftspan")
        m.x.leftbool = LEFT(m.i.big, 2).set_display_name("leftbool")
        m.x.right = RIGHT(m.i.codes, m.i.k).set_display_name("right")
        m.x.mid = MID(m.i.codes, 4, 2).set_display_name("mid")
        m.x.midpast = MID(m.i.codes, 9, 2).set_display_name("midpast")
        m.x.trim = TRIM(m.i.padded).set_display_name("trim")
        m.x.find = FIND("-", m.i.codes).set_display_name("find")
        m.x.find2 = FIND("a", m.i.names, 2).set_display_name("find2")
        m.x.findempty = FIND("", m.i.names).set_display_name("findempty")
        m.x.findempty2 = FIND("", "abc", 4).set_display_name("findempty2")
        m.x.search = SEARCH("a*f", m.i.names).set_display_name("search")
        m.x.search2 = SEARCH("A~?", "xa?b").set_display_name("search2")
        m.x.searchempty = SEARCH("", "abc", 4).set_display_name("searchempty")
        m.x.searchempty2 = SEARCH("", m.i.codes, 2).set_display_name("searchempty2")
        m.x.tilde = SEARCH("~a", "xab").set_display_name("tilde")
        m.x.sub = SUBSTITUTE(m.i.codes, "-", "").set_display_name("sub")
        m.x.sub2 = SUBSTITUTE(m.i.names, "a", "A", 1).set_display_name("sub2")
        m.x.value = VALUE(m.i.texts).set_display_name("value")
        m.x.value2 = VALUE(m.i.texts2).set_display_name("value2")
        m.x.valuebool = VALUE(m.i.flag).set_display_name("valuebool")
        m.x.valuebig = VALUE("1e999").set_display_name("valuebig")
        m.x.valuetime = VALUE(m.i.times).set_display_name("valuetime")
        m.x.valuetime2 = VALUE(m.i.times2).set_display_name("valuetime2")
        m.x.valuetime3 = VALUE(m.i.times3).set_display_name("valuetime3")
        m.x.text0 = TEXT(m.i.amounts, "0").set_display_name("text0")
        m.x.text2 = TEXT(m.i.amounts, "#,##0.00").set_display_name("text2")
        m.x.textp = TEXT(m.i.amounts, "0.0%").set_display_name("textp")
        m.x.textg = TEXT(m.i.amounts, "General").set_display_name("textg")
        m.x.textgbig = TEXT(m.i.bigs, "General").set_display_name("textgbig")
        m.x.textz = TEXT(m.i.k, "000").set_display_name("textz")
        m.x.textpad = TEXT(m.i.pads, "#,##000000").set_display_name("textpad")
        m.x.textgroup = TEXT(m.i.bigs, "#,##0").set_display_name("textgroup")
        m.x.textd = TEXT(m.i.dates, "mmm yyyy").set_display_name("textd")
        m.x.textd2 = TEXT(m.i.dates, "dd.mm.yyyy").set_display_name("textd2")
        m.x.textd3 = TEXT(m.i.dates, "yyyy-mm-dd").set_display_name("textd3")
        m.x.textm = TEXT(m.i.dates, "mmmmm").set_display_name("textm")
        m.x.textb = TEXT(m.i.big, "0").set_display_name("textb")
        m.x.textlong = TEXT(m.i.prec, "0.000000000000000000").set_display_name("textlong")
        m.x.textwide = TEXT(123456789012345678.0, "0").set_display_name("textwide")
        m.x.textempty = TEXT(m.i.amounts, "").set_display_name("textempty")
        m.x.rept = REPT("#", m.i.k).set_display_name("rept")
        m.x.reptbig = REPT("a", 32768).set_display_name("reptbig")
        m.x.lit = LEFT(["ab", "cd", "ef"], 1).set_display_name("lit")
        m.x.join = TEXTJOIN(", ", True, m.i.names).set_display_name("join")
        m.x.join2 = TEXTJOIN(m.i.sep, False, m.i.codes, "end").set_display_name("join2")
        m.x.joinx = TEXTJOIN(",", "x", m.i.names).set_display_name("joinx")
        m.x.quote = SUBSTITUTE(m.i.names, "a", '"').set_display_name("quote")
        m.x.chain = mo.UPPER(LEFT(m.i.names, 3)).set_display_name("chain")
        m.x.err = LEFT(m.i.bad, 1).set_display_name("err")
        # a literal blank where text is read goes in as "", not as 0
        m.x.blank = LEFT(None, 2).set_display_name("blank")
        m.x.blanklist = LEFT(["ab", None, "cd"], 1).set_display_name("blanklist")
        m.x.blankfind = FIND("0", None).set_display_name("blankfind")
        m.x.blanksub = SUBSTITUTE("a0b", "0", None).set_display_name("blanksub")
        m.x.blankrept = REPT(None, 3).set_display_name("blankrept")
        m.x.blankfmt = TEXT(m.i.amounts, None).set_display_name("blankfmt")
        m.x.blankjoin = TEXTJOIN("-", False, "a", None, "b").set_display_name("blankjoin")
        m.x.blanklen = mo.LEN(None).set_display_name("blanklen")
        m.x.blankupper = mo.UPPER(["ab", None, "cd"]).set_display_name("blankupper")
        m.x.blankcat = mo.CONCAT("a", None, "b").set_display_name("blankcat")
    return m


def _rows(ws):
    out = {}
    for row in ws.iter_rows():
        label = row[0].value
        if label:
            out[str(label)] = [c.value for c in row[1:]]
    return out


def _formulas(model, tmp_path):
    """Sheet X's formulas by line — the period cells of a series, the one
    cell of a scalar (the sheet keeps a column for scalars, blank on a
    series row)."""
    path = tmp_path / "text.xlsx"
    model.to_excel(path)
    return {label: [v for v in cells if v is not None]
            for label, cells in _rows(load_workbook(path)["X"]).items()}


class TestWhatTheWorkbookSays:
    def test_one_call_per_period(self, tmp_path):
        x = _formulas(_model(), tmp_path)
        assert x["left"] == ['=LEFT(I!C2, 2)', '=LEFT(I!D2, 2)', '=LEFT(I!E2, 2)']
        assert x["left1"] == ['=LEFT(I!C2)', '=LEFT(I!D2)', '=LEFT(I!E2)']
        assert x["leftdate"][0] == '=LEFT(I!C6, 4)'
        assert x["right"] == ['=RIGHT(I!C3, I!B7)', '=RIGHT(I!D3, I!B7)', '=RIGHT(I!E3, I!B7)']
        assert x["mid"][0] == '=MID(I!C3, 4, 2)'
        assert x["find"][0] == '=FIND("-", I!C3)'
        assert x["find2"] == ['=FIND("a", I!C2, 2)', '=FIND("a", I!D2, 2)', '=FIND("a", I!E2, 2)']
        assert x["findempty"][0] == '=FIND("", I!C2)'
        assert x["search"][0] == '=SEARCH("a*f", I!C2)'
        assert x["sub"][0] == '=SUBSTITUTE(I!C3, "-", "")'
        assert x["sub2"][0] == '=SUBSTITUTE(I!C2, "a", "A", 1)'
        assert x["value"] == ['=VALUE(I!C5)', '=VALUE(I!D5)', '=VALUE(I!E5)']
        assert x["text0"][0] == '=TEXT(I!C4, "0")'
        assert x["text2"][0] == '=TEXT(I!C4, "#,##0.00")'
        assert x["textp"][0] == '=TEXT(I!C4, "0.0%")'
        assert x["textd"] == ['=TEXT(I!C6, "mmm yyyy")', '=TEXT(I!D6, "mmm yyyy")',
                              '=TEXT(I!E6, "mmm yyyy")']
        assert x["textd2"][0] == '=TEXT(I!C6, "dd.mm.yyyy")'
        assert x["textm"][0] == '=TEXT(I!C6, "mmmmm")'
        assert x["textpad"] == ['=TEXT(I!C23, "#,##000000")', '=TEXT(I!D23, "#,##000000")',
                                '=TEXT(I!E23, "#,##000000")']
        assert x["valuetime3"][0] == '=VALUE(I!C22)'
        assert x["leftbig2"][0] == '=LEFT(I!C21, 25)'
        assert x["textb"][0] == '=TEXT(I!C12, "0")'
        assert x["leftnums"][0] == '=LEFT(I!C13, 20)'
        assert x["value2"][0] == '=VALUE(I!C14)'
        assert x["err"][0] == '=LEFT(I!C11, 1)'

    def test_a_scalar_call_is_one_cell(self, tmp_path):
        x = _formulas(_model(), tmp_path)
        assert x["trim"] == ['=TRIM(I!B9)']
        assert x["rept"] == ['=REPT("#", I!B7)']
        assert x["search2"] == ['=SEARCH("A~?", "xa?b")']
        assert x["tilde"] == ['=SEARCH("~a", "xab")']
        assert x["findempty2"] == ['=FIND("", "abc", 4)']
        assert x["textz"] == ['=TEXT(I!B7, "000")']
        assert x["valuebool"] == ['=VALUE(I!B15)']
        assert x["valuebig"] == ['=VALUE("1e999")']

    def test_textjoin_joins_the_row_as_a_range(self, tmp_path):
        x = _formulas(_model(), tmp_path)
        assert x["join"] == ['=_xlfn.TEXTJOIN(", ", TRUE, I!C2:E2)']
        assert x["join2"] == ['=_xlfn.TEXTJOIN(I!B8, FALSE, I!C3:E3, "end")']
        assert x["joinx"] == ['=_xlfn.TEXTJOIN(",", "x", I!C2:E2)']

    def test_textjoin_over_a_literal_list_is_written_as_its_value(self, tmp_path):
        # The range-function law: a list with no cells in the workbook
        # cannot be a range, so the call's computed text is written in
        # its place, with a warning. An array constant would keep it
        # live; the renderer does not write one.
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=PERIODS)
        with m:
            m.x = mo.MultiVariable("X")
            m.x.joinlist = TEXTJOIN(",", True, ["a", "b"], "c").set_display_name("joinlist")
        with pytest.warns(CrossScopeReferenceWarning, match="literal list"):
            m.to_excel(tmp_path / "text.xlsx")
        assert _rows(load_workbook(tmp_path / "text.xlsx")["X"])["joinlist"][0] == "a,b,c"

    def test_a_literal_list_is_this_period_s_item(self, tmp_path):
        x = _formulas(_model(), tmp_path)
        assert x["lit"] == ['=LEFT("ab", 1)', '=LEFT("cd", 1)', '=LEFT("ef", 1)']

    def test_a_literal_blank_where_text_is_read_is_empty_text(self, tmp_path):
        x = _formulas(_model(), tmp_path)
        assert x["blank"] == ['=LEFT("", 2)']
        assert x["blanklist"] == ['=LEFT("ab", 1)', '=LEFT("", 1)', '=LEFT("cd", 1)']
        assert x["blankfind"] == ['=FIND("0", "")']
        assert x["blanksub"] == ['=SUBSTITUTE("a0b", "0", "")']
        assert x["blankrept"] == ['=REPT("", 3)']
        assert x["blankfmt"][0] == '=TEXT(I!C4, "")'
        assert x["blankjoin"] == ['=_xlfn.TEXTJOIN("-", FALSE, "a", "", "b")']
        assert x["blanklen"] == ['=LEN("")']
        assert x["blankupper"] == ['=UPPER("ab")', '=UPPER("")', '=UPPER("cd")']
        assert x["blankcat"] == ['="a" & "" & "b"']

    def test_a_quote_in_a_text_is_doubled(self, tmp_path):
        x = _formulas(_model(), tmp_path)
        assert x["quote"][0] == '=SUBSTITUTE(I!C2, "a", """")'

    def test_a_call_inside_a_call(self, tmp_path):
        x = _formulas(_model(), tmp_path)
        assert x["chain"] == ['=UPPER(LEFT(I!C2, 3))', '=UPPER(LEFT(I!D2, 3))',
                              '=UPPER(LEFT(I!E2, 3))']

    def test_every_name_is_known_to_the_renderer(self, tmp_path, caplog):
        with caplog.at_level(logging.WARNING, logger="modeleon.compile.excel.renderer"):
            _model().to_excel(tmp_path / "text.xlsx")
        assert not [r for r in caplog.records if "Unknown function" in r.getMessage()]


def _soffice():
    for candidate in ("soffice", "libreoffice"):
        found = shutil.which(candidate)
        if found:
            return found
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


def _cell(v):
    return "" if v is None else v


@pytest.mark.slow
def test_the_workbook_recalculates_to_the_python_values(tmp_path):
    soffice = _soffice()
    if soffice is None:
        pytest.skip("LibreOffice not installed")
    m = _model()
    path = tmp_path / "text.xlsx"
    m.to_excel(path)
    subprocess.run(
        [soffice, f"-env:UserInstallation={(tmp_path / 'prof').as_uri()}",
         "--headless", "--convert-to", "xlsx", "--outdir",
         str(tmp_path / "calc"), str(path)],
        check=True, capture_output=True,
    )
    x = _rows(load_workbook(tmp_path / "calc" / "text.xlsx", data_only=True)["X"])
    assert set(DIVERGENT) <= set(LINES)
    for name in LINES:
        var = getattr(m.x, name)
        if var.var_type == 'list':
            got = [_cell(v) for v in x[name][-PERIODS:]]
        else:
            got = _cell(x[name][0])
        engine, libre = DIVERGENT.get(name, (var.value, var.value))
        assert var.value == engine, name
        assert got == libre, name
