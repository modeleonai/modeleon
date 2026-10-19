# SPDX-License-Identifier: Apache-2.0
"""``MIN`` / ``MAX`` over two or more operands is one call per period.

A hand-built workbook writes a cap or a floor in every column —
``=MIN(C2, C3)`` for a sweep no larger than the balance, ``=MAX(C4,
Inputs!B2)`` for a cash floor — and that is what ``mo.MIN(sweep,
balance)`` and ``mo.MAX(cash, inp.min_cash)`` mean: a series, period by
period, scalars broadcast as ``a + b`` broadcasts them, inside a loop
container too. ``mo.MIN(row)`` over ONE operand stays the reduction it
always was (``=MIN(B2:F2)``), and over scalars only ``mo.MIN(a, b)``
stays one number (``=MIN(B2, B3)``).
"""

import shutil
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook

import modeleon as mo


def _rows(ws):
    out = {}
    for row in ws.iter_rows():
        label = row[0].value
        if label:
            out[str(label)] = [c.value for c in row[1:]]
    return out


def _book(model, tmp_path, sheet):
    path = tmp_path / "book.xlsx"
    model.to_excel(path)
    return _rows(load_workbook(path)[sheet])


def _model():
    m = mo.Model("m", default_grain="year", default_start="2026", default_periods=3)
    with m:
        m.i = mo.MultiVariable("I")
        m.i.floor = mo.Variable(4.0, display_name="floor")
        m.i.p = mo.Variable(3.0, display_name="p")
        m.i.q = mo.Variable(9.0, display_name="q")
        m.x = mo.MultiVariable("X")
        m.x.a = mo.Variable([3.0, 8.0, 5.0], display_name="a")
        m.x.b = mo.Variable([4.0, 4.0, 4.0], display_name="b")
    return m


# ══════════════════════════════════════════════════════════════════════
# The values
# ══════════════════════════════════════════════════════════════════════


class TestTwoOrMoreOperandsArePerPeriod:
    def test_series_by_series(self):
        a = mo.Variable([3.0, 8.0, 5.0])
        b = mo.Variable([4.0, 4.0, 4.0])
        assert mo.MIN(a, b).value == [3.0, 4.0, 4.0]
        assert mo.MAX(a, b).value == [4.0, 8.0, 5.0]
        assert mo.MIN(a, b).var_type == "list"

    def test_a_scalar_broadcasts_on_either_side(self):
        a = mo.Variable([3.0, 8.0, 5.0])
        assert mo.MAX(a, 4.0).value == [4.0, 8.0, 5.0]
        assert mo.MAX(4.0, a).value == [4.0, 8.0, 5.0]
        assert mo.MIN(a, mo.Variable(4.0)).value == [3.0, 4.0, 4.0]
        assert mo.MAX(mo.Variable(0.0), a).value == [3.0, 8.0, 5.0]

    def test_a_raw_list_is_a_series(self):
        a = mo.Variable([3.0, 8.0, 5.0])
        assert mo.MIN(a, [4.0, 4.0, 4.0]).value == [3.0, 4.0, 4.0]
        assert mo.MAX([1.0, 2.0, 3.0], 2.0).value == [2.0, 2.0, 3.0]

    def test_three_operands(self):
        a = mo.Variable([3.0, 8.0, 5.0])
        b = mo.Variable([4.0, 4.0, 4.0])
        assert mo.MIN(a, b, 3.5).value == [3.0, 3.5, 3.5]
        assert mo.MAX(a, b, mo.Variable(6.0)).value == [6.0, 8.0, 6.0]

    def test_nested_in_arithmetic(self):
        a = mo.Variable([3.0, 8.0, 5.0])
        b = mo.Variable([4.0, 4.0, 4.0])
        assert (mo.MIN(a, b) * 2).value == [6.0, 8.0, 8.0]
        assert mo.MIN(a * 2, b).value == [4.0, 4.0, 4.0]
        assert (mo.MAX(a, 4.0) - mo.MIN(a, b)).value == [1.0, 4.0, 1.0]

    def test_a_one_value_series_broadcasts_as_the_operators_do(self):
        a = mo.Variable([3.0, 8.0, 5.0])
        assert mo.MAX(mo.Variable([4.0]), a).value == [4.0, 8.0, 5.0]

    def test_series_of_different_lengths_are_refused_as_the_operators_refuse_them(self):
        a = mo.Variable([3.0, 8.0, 5.0])
        with pytest.raises(ValueError, match="different lengths"):
            mo.MAX(mo.Variable([1.0, 2.0]), a)

    def test_the_formula_names_the_operands(self):
        m = _model()
        assert mo.MIN(m.x.a, m.x.b).formula == "MIN(a, b)"
        assert mo.MAX(m.x.a, 0).formula == "MAX(a, 0)"

    def test_value_type_follows_the_first_variable(self):
        a = mo.Variable([3, 8, 5])
        assert mo.MAX(a, 4).value_type == "int"


class TestOneOperandOrScalarsOnlyAreUnchanged:
    def test_one_operand_is_a_reduction(self):
        a = mo.Variable([3.0, 8.0, 5.0])
        assert mo.MIN(a).value == 3.0
        assert mo.MAX(a).value == 8.0
        assert mo.MIN(a).var_type == "scalar"
        assert mo.MAX([3.0, 8.0, 5.0]).value == 8.0

    def test_scalars_only_are_one_number(self):
        lo = mo.MIN(mo.Variable(3.0), mo.Variable(4.0))
        assert lo.value == 3.0 and lo.var_type == "scalar"
        hi = mo.MAX(5, 10, 3)
        assert hi.value == 10 and hi.formula == "MAX(5, 10, 3)"

    def test_a_reduction_as_an_operand_is_a_scalar(self):
        a = mo.Variable([3.0, 8.0, 5.0])
        worst = mo.MAX(mo.MAX(a), -mo.MIN(a))
        assert worst.value == 8.0 and worst.var_type == "scalar"


class TestUnitsAndErrors:
    def test_a_shared_unit_is_kept_and_a_mixed_one_dropped(self):
        a = mo.Variable([3.0, 8.0, 5.0], unit="$")
        b = mo.Variable([4.0, 4.0, 4.0], unit="$")
        c = mo.Variable([4.0, 4.0, 4.0], unit="hr")
        assert str(mo.MIN(a, b)._unit) == "$"
        assert str(mo.MAX(a, 0.0)._unit) == "$"
        assert mo.MIN(a, c)._unit is None

    def test_an_error_cell_is_the_period_s_result(self):
        e = mo.Variable([2.0, 2.0, 2.0]) / mo.Variable([1.0, 0.0, 2.0])
        assert e.value == [2.0, "#DIV/0!", 1.0]
        assert mo.MAX(e, 1.5).value == [2.0, "#DIV/0!", 1.5]
        assert mo.MIN(1.5, e).value == [1.5, "#DIV/0!", 1.0]

    def test_a_blank_is_skipped_as_in_the_spreadsheet(self):
        blank = mo.Variable([None, 5.0])
        assert mo.MIN(blank, 2.0).value == [2.0, 2.0]
        assert mo.MAX(blank, mo.Variable([None, None])).value == [0, 5.0]

    def test_tracked_operands_lift_per_track(self):
        m = mo.Model("m", tracks=mo.Tracks("plan", "actual"),
                     default_grain="month", default_start="2026-01", default_periods=3)
        with m:
            m.x = mo.Variable(plan=[1.0, 5.0, 3.0], actual=[2.0, 2.0, 2.0])
            m.cap = mo.MIN(m.x, 2.5)
            m.hi = mo.MAX(m.x, m.x * 0.0 + 1.5)
        assert dict(m.cap._value.items()) == {"plan": [1.0, 2.5, 2.5], "actual": [2.0, 2.0, 2.0]}
        assert dict(m.hi._value.items()) == {"plan": [1.5, 5.0, 3.0], "actual": [2.0, 2.0, 2.0]}
        assert m.cap.var_type == "list"


# ══════════════════════════════════════════════════════════════════════
# Inside a loop container
# ══════════════════════════════════════════════════════════════════════


def _loop(m):
    with m:
        m.d = mo.MultiVariable("D", loop=True)
        with m.d as d:
            d.sweep = mo.MIN(m.x.a, mo.lag(d.balance, fill=m.i.q)).set_display_name("sweep")
            d.balance = (mo.lag(d.balance, fill=m.i.q) - d.sweep).set_display_name("balance")
            d.floored = mo.MAX(d.balance - 2.0, 0.0).set_display_name("floored")
    return m


class TestInsideALoop:
    def test_a_cap_on_a_row_read_through_lag(self):
        m = _loop(_model())
        assert m.d.sweep.value == [3.0, 6.0, 0.0]
        assert m.d.balance.value == [6.0, 0.0, 0.0]
        assert m.d.floored.value == [4.0, 0.0, 0.0]

    def test_a_one_operand_reduction_stays_refused_and_points_at_the_pair(self):
        m = mo.Model("m", default_grain="year", default_start="2026", default_periods=4)
        with pytest.raises(ValueError, match="MAX reads a whole series") as err:
            with m:
                m.x = mo.MultiVariable("X", loop=True)
                with m.x as x:
                    x.a = mo.lag(x.b) + 1.0
                    x.peak = mo.MAX(x.a)
                    x.b = x.a - x.peak
        assert "mo.MAX(x, 0)" in str(err.value)
        assert "mo.IF" not in str(err.value)


# ══════════════════════════════════════════════════════════════════════
# The three renderers and check_compat agree
# ══════════════════════════════════════════════════════════════════════


class TestTheWorkbook:
    def test_one_call_per_period(self, tmp_path):
        m = _model()
        with m:
            m.x.lo = mo.MIN(m.x.a, m.x.b).set_display_name("lo")
            m.x.fl = mo.MAX(m.x.a, m.i.floor).set_display_name("fl")
            m.x.tri = mo.MIN(m.x.a, m.x.b, 4.5).set_display_name("tri")
            m.x.lit = mo.MAX(m.x.a, [1.0, 9.0, 1.0]).set_display_name("lit")
        x = _book(m, tmp_path, "X")
        assert x["lo"] == ["=MIN(B2, B3)", "=MIN(C2, C3)", "=MIN(D2, D3)"], x["lo"]
        assert x["fl"] == ["=MAX(B2, I!B1)", "=MAX(C2, I!B1)", "=MAX(D2, I!B1)"], x["fl"]
        assert x["tri"] == ["=MIN(B2, B3, 4.5)", "=MIN(C2, C3, 4.5)", "=MIN(D2, D3, 4.5)"], x["tri"]
        assert x["lit"] == ["=MAX(B2, 1.0)", "=MAX(C2, 9.0)", "=MAX(D2, 1.0)"], x["lit"]

    def test_inside_a_larger_formula_and_over_an_unlaid_operand(self, tmp_path):
        m = _model()
        with m:
            m.x.twice = (mo.MIN(m.x.a, m.x.b) * 2).set_display_name("twice")
            m.x.inner = mo.MIN(m.x.a * 2, m.x.b).set_display_name("inner")
        x = _book(m, tmp_path, "X")
        assert x["twice"] == ["=MIN(B2, B3) * 2", "=MIN(C2, C3) * 2", "=MIN(D2, D3) * 2"], x["twice"]
        assert x["inner"] == ["=MIN(B2 * 2, B3)", "=MIN(C2 * 2, C3)", "=MIN(D2 * 2, D3)"], x["inner"]

    def test_the_reduction_and_the_scalar_pair_are_unchanged(self, tmp_path):
        m = _model()
        with m:
            m.x.red = mo.MIN(m.x.a).set_display_name("red")
            m.x.pair = mo.MIN(m.i.p, m.i.q).set_display_name("pair")
        # A scalar line gives the sheet a constants column (B); the
        # periods move to C..E.
        x = _book(m, tmp_path, "X")
        assert x["red"][0] == "=MIN(C2:E2)", x["red"]
        assert x["pair"][0] == "=MIN(I!B2, I!B3)", x["pair"]

    def test_a_loop_s_cap_reads_the_previous_column(self, tmp_path):
        m = _loop(_model())
        d = _book(m, tmp_path, "D")
        assert d["sweep"] == ["=MIN(X!B2, I!B3)", "=MIN(X!C2, B3)", "=MIN(X!D2, C3)"], d["sweep"]
        assert d["balance"] == ["=I!B3 - B2", "=B3 - C2", "=C3 - D2"], d["balance"]
        assert d["floored"] == ["=MAX(B3 - 2.0, 0.0)", "=MAX(C3 - 2.0, 0.0)",
                                "=MAX(D3 - 2.0, 0.0)"], d["floored"]

    def test_check_compat_lists_nothing(self):
        m = _loop(_model())
        with m:
            m.x.lo = mo.MIN(m.x.a, m.x.b).set_display_name("lo")
            m.x.twice = (mo.MIN(m.x.a, m.x.b) * 2).set_display_name("twice")
        assert mo.check_compat(m, "excel") == []


class TestTheOtherRenderers:
    def test_json_is_the_call_over_its_operands(self):
        m = _model()
        with m:
            m.x.lo = mo.MIN(m.x.a, m.x.b).set_display_name("lo")
        assert mo.to_json(m.x.lo._expr) == {
            "kind": "funccall", "func": "MIN",
            "args": [{"kind": "varref", "path": "m.x.a"}, {"kind": "varref", "path": "m.x.b"}],
        }
        assert m.x.lo.var_type == "list"
        assert m.x.to_dict()["components_data"]["lo"]["values"] == [3.0, 4.0, 4.0]

    def test_the_notebook_shows_the_per_cell_formula(self):
        m = _model()
        with m:
            m.x.lo = mo.MIN(m.x.a, m.x.b).set_display_name("lo")
        html = m._repr_html_()
        assert "MIN(B2, B3)" in html
        assert "B2:D2" not in html


# ══════════════════════════════════════════════════════════════════════
# The spreadsheet computes the same numbers
# ══════════════════════════════════════════════════════════════════════


def _soffice():
    for candidate in ("soffice", "libreoffice"):
        found = shutil.which(candidate)
        if found:
            return found
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


@pytest.mark.slow
def test_the_workbook_recalculates_to_the_python_values(tmp_path):
    soffice = _soffice()
    if soffice is None:
        pytest.skip("LibreOffice not installed")
    m = _loop(_model())
    with m:
        m.x.lo = mo.MIN(m.x.a, m.x.b).set_display_name("lo")
        m.x.fl = mo.MAX(m.x.a, m.i.floor).set_display_name("fl")
        m.x.tri = mo.MIN(m.x.a, m.x.b, 4.5).set_display_name("tri")
        m.x.twice = (mo.MIN(m.x.a, m.x.b) * 2).set_display_name("twice")
        m.x.inner = mo.MIN(m.x.a * 2, m.x.b).set_display_name("inner")
        # On the scalars' sheet, so X keeps its period columns at B..D.
        m.i.red = mo.MIN(m.x.a).set_display_name("red")
    path = tmp_path / "minmax.xlsx"
    m.to_excel(path)
    subprocess.run(
        [soffice, f"-env:UserInstallation={(tmp_path / 'prof').as_uri()}",
         "--headless", "--convert-to", "xlsx", "--outdir",
         str(tmp_path / "calc"), str(path)],
        check=True, capture_output=True,
    )
    book = load_workbook(tmp_path / "calc" / "minmax.xlsx", data_only=True)
    x = _rows(book["X"])
    d = _rows(book["D"])
    assert m.d.sweep.value == [3.0, 6.0, 0.0]
    for name in ("lo", "fl", "tri", "twice", "inner"):
        assert [round(v, 6) for v in x[name][:3]] == getattr(m.x, name).value, name
    assert round(_rows(book["I"])["red"][0], 6) == m.i.red.value == 3.0
    for name in ("sweep", "balance", "floored"):
        assert [round(v, 6) for v in d[name][:3]] == getattr(m.d, name).value, name
