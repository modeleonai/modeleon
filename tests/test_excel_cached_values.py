# SPDX-License-Identifier: Apache-2.0
"""Cached values: a workbook's formula cells carry what the engine computed.

A reader that never recalculates (a file preview, a library reading
values) shows a formula cell's cached value, and an empty one reads as 0
or as a date in 1899. These tests pin that every cell the engine computed
carries its value beside its formula, that a header caption pointing at a
known cell carries that cell's value, that a row total carries
``row_sum_value`` of its row (what its ``=SUM`` gives), and, with
LibreOffice, that every cached value is what the spreadsheet itself
computes.
"""

import re
import shutil
import subprocess
import warnings
import zipfile
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from pathlib import Path

import numpy as np
import openpyxl
import pytest
from openpyxl.utils.datetime import WINDOWS_EPOCH

import modeleon as mo
from modeleon.compile.excel.cached_values import (
    _FORMULA_CELL,
    _VALUES,
    _cached,
    _Points,
    _with_value,
    remember,
    remember_reference,
    save_workbook,
)
from modeleon.compile.excel.writer import row_sum_value


def model(n=4):
    """Inputs, a time header mirrored atop the other sheets with an opening
    column, a line with a row total, and a line of text."""
    m = mo.Model('cv', default_grain='month', default_periods=n)
    with m:
        m.default_excel_view = mo.ExcelView('book', timeline={'opening': 'Pre-start'},
                                            meta={'fields': [('row_sum', 'Total')]})
        m.inputs = mo.MultiVariable('Inputs')
        with m.inputs as i:
            i.start = mo.Variable(date(2026, 1, 1))
            i.price = mo.Variable(12.5)
            i.units = mo.Variable([10.0 * (k + 1) for k in range(n)])
        m.default_start = m.inputs.start
        m.header = mo.MultiVariable('Time')
        with m.header as h:
            h.end = mo.time.end
            h.days = mo.time.days
            h.n = mo.time.index
        m.default_header = m.header
        m.sales = mo.MultiVariable('Sales')
        with m.sales as s:
            s.revenue = (m.inputs.units * m.inputs.price).set_style(row_sum=True)
            s.size = mo.IF(s.revenue > 300, 'big', 'small')
    return m


def write(m, tmp_path, name='book.xlsx'):
    path = tmp_path / name
    with warnings.catch_warnings():
        warnings.simplefilter('ignore')
        m.to_excel(str(path))
    return path


def formula_cells(path, sheet):
    """``{coordinate: (formula, cached value)}`` for every formula cell."""
    f = openpyxl.load_workbook(str(path))[sheet]
    v = openpyxl.load_workbook(str(path), data_only=True)[sheet]
    return {
        c.coordinate: (c.value, v[c.coordinate].value)
        for row in f.iter_rows() for c in row
        if isinstance(c.value, str) and c.value.startswith('=')
    }


def row_values(path, sheet, label):
    """The cached values of the formula cells on the row labelled ``label``."""
    f = openpyxl.load_workbook(str(path))[sheet]
    v = openpyxl.load_workbook(str(path), data_only=True)[sheet]
    for row in f.iter_rows():
        if any(c.value == label for c in row):
            return [v[c.coordinate].value for c in row
                    if isinstance(c.value, str) and c.value.startswith('=')
                    and not c.value.startswith('=SUM(')]
    raise AssertionError(f'no row labelled {label!r} on {sheet}')


def test_a_formula_cell_carries_the_value_the_engine_computed(tmp_path):
    m = model()
    path = write(m, tmp_path)
    assert row_values(path, 'Sales', 'Revenue') == pytest.approx(m.sales.revenue.value)
    assert row_values(path, 'Sales', 'Size') == m.sales.size.value
    # …and the cell is still the formula: the value rides beside it.
    cells = formula_cells(path, 'Sales')
    assert cells and all(f.startswith('=') for f, _ in cells.values())


def test_the_time_rows_and_their_captions_carry_their_dates(tmp_path):
    m = model()
    path = write(m, tmp_path)
    # The anchor's own rows, the mirrors of them atop another sheet, their
    # opening cells and the period captions pointing at a period's end all
    # show a value.
    for sheet in ('Time', 'Sales'):
        cells = formula_cells(path, sheet)
        header = {ref: v for ref, (f, v) in cells.items() if not f.startswith('=SUM(')}
        assert header and all(v is not None for v in header.values()), (sheet, header)
    ends = [d.date() if isinstance(d, datetime) else d for d in row_values(path, 'Time', 'End')]
    # The opening cell (the day before the first period), then each end.
    assert ends == [date(2025, 12, 31), *m.header.end.value]


def test_a_row_total_carries_the_sum_of_its_row(tmp_path):
    m = model()
    path = write(m, tmp_path)
    totals = [v for f, v in formula_cells(path, 'Sales').values() if f.startswith('=SUM(')]
    assert totals == pytest.approx([sum(m.sales.revenue.value)])


@pytest.mark.parametrize('series, total', [
    ([1.0, 2.5, None, 3], 6.5),
    ([None, None], None),
    (['a', True, 4.0], 4.0),
    ([np.int64(2), np.float64(0.5)], 2.5),
    ([], None),
    (7.0, None),
])
def test_a_row_total_adds_the_numbers_of_its_row(series, total):
    # As =SUM over the row's cells does: blanks, text and booleans skipped.
    assert row_sum_value(series) == total


def test_text_is_cached_as_text_and_numbers_as_numbers(tmp_path):
    path = write(model(), tmp_path)
    with zipfile.ZipFile(path) as z:
        names = {s.title: i + 1 for i, s in enumerate(openpyxl.load_workbook(str(path)).worksheets)}
        xml = z.read(f"xl/worksheets/sheet{names['Sales']}.xml").decode()
    assert '<v>small</v>' in xml and 't="str"' in xml
    assert '<v>125</v>' in xml


def test_the_book_still_recalculates_when_opened(tmp_path):
    path = write(model(), tmp_path)
    with zipfile.ZipFile(path) as z:
        assert 'fullCalcOnLoad="1"' in z.read('xl/workbook.xml').decode()


@pytest.mark.parametrize('value, cached', [
    (None, None),
    (True, ('b', '1')),
    (False, ('b', '0')),
    (np.bool_(True), ('b', '1')),
    (7, ('', '7')),
    (np.int64(7), ('', '7')),
    pytest.param(10 ** 400, None, id='int-past-a-double'),
    pytest.param(10 ** 5000, None, id='int-past-str-limit'),
    (2.5, ('', '2.5')),
    (np.float64(0.1), ('', '0.1')),
    (0.1 + 0.2, ('', '0.3')),
    (Decimal('1.25'), ('', '1.25')),
    (float('nan'), None),
    (float('inf'), None),
    ('a < b & c', ('str', 'a &lt; b &amp; c')),
    ('line\rbreak', ('str', 'line&#13;break')),
    ('bell\x07', None),
    pytest.param('x' * 40_000, None, id='text-past-a-cell'),
    ('#DIV/0!', ('e', '#DIV/0!')),
    (date(2045, 12, 31), ('', '53327')),
    (datetime(2045, 12, 31, 12, 0), ('', '53327.5')),
    (time(12, 0), ('', '0.5')),
    (timedelta(hours=6), ('', '0.25')),
    (np.timedelta64(1, 'D'), None),
    (np.datetime64('2026-01-01'), None),
    (object(), None),
])
def test_each_kind_of_value_has_its_spreadsheet_form(value, cached):
    assert _cached(value, WINDOWS_EPOCH) == cached


@pytest.mark.parametrize('reference, points', [
    ('H4', _Points(None, 'H4')),
    ('$H$4', _Points(None, 'H4')),
    ("'Time'!H12", _Points('Time', 'H12')),
    ("'It''s'!B2", _Points("It's", 'B2')),
    ('Timing!AC9', _Points('Timing', 'AC9')),
    ('Plan_Übersicht!D2', _Points('Plan_Übersicht', 'D2')),  # letters beyond ASCII
    ('H4+1', None),
    ('SUM(H4:H9)', None),
])
def test_a_caption_points_only_at_a_single_cell(reference, points):
    ws = openpyxl.Workbook().active
    remember_reference(ws, 'A1', reference)
    assert (_VALUES.get(ws) or {}).get('A1') == points


@pytest.mark.parametrize('empty', ['<v></v>', '<v />', '<v/>'])
def test_either_empty_value_takes_the_value_and_an_unknown_one_none(empty):
    # lxml writes the empty value as <v></v>, the standard library (used
    # when lxml is not installed) as <v />: both are filled the same way.
    xml = f'<c r="B1" s="2"><f>A1*2</f>{empty}</c><c r="C1"><f>B1+1</f>{empty}</c>'
    out = _FORMULA_CELL.sub(lambda m: _with_value(m, {'B1': 4.0}, WINDOWS_EPOCH), xml)
    assert out == '<c r="B1" s="2"><f>A1*2</f><v>4</v></c><c r="C1"><f>B1+1</f></c>'


def test_a_caption_reading_another_sheet_shows_that_sheet_s_date(tmp_path):
    m = mo.Model('cv', default_grain='month', default_periods=3)
    with m:
        m.inputs = mo.MultiVariable('Inputs')
        m.inputs.start = mo.Variable(date(2026, 1, 1))
        m.default_start = m.inputs.start
        m.header = mo.MultiVariable('Time')
        with m.header as h:
            h.end = mo.time.end
            h.days = mo.time.days
        # Only the days are repeated on Sales: its captions read Time's ends.
        m.default_header = [m.header.days]
        m.sales = mo.MultiVariable('Sales')
        m.sales.units = mo.Variable([1.0, 2.0, 3.0])
    path = write(m, tmp_path)
    captions = {ref: v for ref, (f, v) in formula_cells(path, 'Sales').items()
                if f.startswith('=Time!') and ref.endswith('1')}
    got = [d.date() for _, d in sorted(captions.items())]
    assert got == m.header.end.value


def _soffice():
    found = shutil.which('soffice') or shutil.which('libreoffice')
    if found:
        return found
    mac = Path('/Applications/LibreOffice.app/Contents/MacOS/soffice')
    return str(mac) if mac.exists() else None


# LibreOffice keeps a formula's cached value on load unless told to
# recalculate: without this profile the check would compare the cache
# with itself.
_RECALC_ALWAYS = (
    '<?xml version="1.0" encoding="UTF-8"?>'
    '<oor:items xmlns:oor="http://openoffice.org/2001/registry" '
    'xmlns:xs="http://www.w3.org/2001/XMLSchema" '
    'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
    '<item oor:path="/org.openoffice.Office.Calc/Formula/Load">'
    '<prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>'
    '</oor:items>'
)


def _without_cached_values(src: Path, dst: Path) -> None:
    # The same book with every formula's value removed: what LibreOffice
    # returns for it can only be its own computation.
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, 'w') as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename.startswith('xl/worksheets/'):
                data = re.sub(rb'(</f>)<v>[^<]*</v>', rb'\1', data)
            zout.writestr(info, data)


@pytest.mark.slow
def test_every_cached_value_is_what_the_spreadsheet_computes(tmp_path):
    soffice = _soffice()
    if soffice is None:
        pytest.skip('LibreOffice is not installed')
    path = write(model(6), tmp_path, 'src.xlsx')
    bare = tmp_path / 'bare.xlsx'
    _without_cached_values(path, bare)
    profile = tmp_path / 'lo-profile'
    (profile / 'user').mkdir(parents=True)
    (profile / 'user' / 'registrymodifications.xcu').write_text(_RECALC_ALWAYS)
    outdir = tmp_path / 'calc'
    subprocess.run(
        [soffice, f'-env:UserInstallation={profile.as_uri()}', '--headless',
         '--convert-to', 'xlsx', '--outdir', str(outdir), str(bare)],
        check=True, capture_output=True, timeout=120,
    )
    recalculated = openpyxl.load_workbook(str(outdir / 'bare.xlsx'), data_only=True)
    checked = 0
    for ws in openpyxl.load_workbook(str(path)).worksheets:
        for ref, (_formula, cached) in formula_cells(path, ws.title).items():
            if cached is None:
                continue
            got = recalculated[ws.title][ref].value
            if isinstance(cached, (int, float)) and not isinstance(cached, bool):
                assert got == pytest.approx(cached), (ws.title, ref)
            else:
                assert got == cached, (ws.title, ref, got, cached)
            checked += 1
    assert checked > 20


@pytest.mark.slow
def test_libreoffice_computes_a_formula_rather_than_keep_its_cache(tmp_path):
    # The ground every LibreOffice parity check in this suite stands on
    # (conftest seeds each profile to recalculate): a formula whose cached
    # value is wrong comes back computed, not as the cache.
    soffice = _soffice()
    if soffice is None:
        pytest.skip('LibreOffice is not installed')
    wb = openpyxl.Workbook()
    ws = wb.active
    ws['A1'] = 5
    ws['B1'] = '=A1*2'
    ws['C1'] = '="a"&"b"'
    remember(ws, 'B1', 99)
    remember(ws, 'C1', 'zz')
    src = tmp_path / 'canary.xlsx'
    save_workbook(wb, src)
    subprocess.run(
        [soffice, f'-env:UserInstallation={(tmp_path / "lo-profile").as_uri()}', '--headless',
         '--convert-to', 'xlsx', '--outdir', str(tmp_path / 'calc'), str(src)],
        check=True, capture_output=True, timeout=120,
    )
    got = openpyxl.load_workbook(str(tmp_path / 'calc' / 'canary.xlsx'), data_only=True).active
    assert (got['B1'].value, got['C1'].value) == (10, 'ab')
