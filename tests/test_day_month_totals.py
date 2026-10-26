# SPDX-License-Identifier: Apache-2.0
"""A day sheet totals by month — "30 31 | Total Jan | 01 02 03 | Total Feb".

A plan keys the sheet's OWN periods ``('period', i)`` and a bucket by
its ruler (``('month', bi)``); they used to share the spelling
``'month'``, so a month total over days took a day's address. Now the
month is a bucket like a quarter: its own columns, its own tier, its
own outline group, a live sum over its days.
"""
from openpyxl import load_workbook

import modeleon as mo
from modeleon.compile.excel.totals import NATIVE, build_totals_plan, header_rows
from modeleon.core.time import DefaultWindow, _period_labels


def _book(tmp_path, kinds):
    m = mo.Model('m', default_grain='day', default_start='2026-01-30', default_periods=5,
                 default_excel_view=mo.ExcelView(timeline={'totals': kinds}))
    with m:
        m.units = mo.Variable([1.0, 2.0, 3.0, 4.0, 5.0], regrain=mo.up('sum'), display_name='Units')
        m.price = mo.Variable(10.0, display_name='Price')
        m.rev = (m.units * m.price).set_regrain(mo.up('sum')).set_display_name('Revenue')
        m.stock = mo.Variable([7.0, 8.0, 9.0, 10.0, 11.0], regrain=mo.up('last'), display_name='Stock')
    path = tmp_path / 'day.xlsx'
    m.to_excel(path)
    return load_workbook(path)['M']


class TestThePlan:
    def test_native_columns_and_month_buckets_never_share_a_key(self):
        plan = build_totals_plan(('month',), 'day', '2026-01-30', 5, 2)
        assert plan.columns == [
            (NATIVE, 0), (NATIVE, 1), ('month', 0),
            (NATIVE, 2), (NATIVE, 3), (NATIVE, 4), ('month', 1),
        ]
        assert len(plan.col_of) == len(plan.columns) == 7
        assert plan.native_col(2) == 5 and plan.month_col(2) == 5   # the old name still answers

    def test_the_month_tier_stands_over_the_days(self):
        plan = build_totals_plan(('month', 'year'), 'day', '2026-01-30', 5, 2)
        rows = header_rows(plan, _period_labels('2026-01-30', None, 5, 'day', 'finance'))
        assert [r['cells'][0]['kind'] for r in rows] == ['year', 'month', NATIVE]
        month_tier = [(c['label'], c['col'], c['colspan']) for c in rows[1]['cells']]
        assert month_tier == [('Jan', 2, 2), ('Total Jan', 4, 1), ('Feb', 5, 3), ('Total Feb', 8, 1)]
        assert [c['label'] for c in rows[2]['cells']] == ['30', '31', '01', '02', '03']


class TestTheBook:
    def test_month_totals_are_live_sums_over_their_days(self, tmp_path):
        ws = _book(tmp_path, ['month'])
        # Column B holds the constants (Price); the periods start at C.
        assert [c.value for c in ws[1]][:9] == [None, None, 'Jan', None, 'Total Jan', 'Feb', None, None, 'Total Feb']
        assert [c.value for c in ws[2]][:9] == [None, None, '30', '31', None, '01', '02', '03', None]
        rows = {r[0].value: [c.value for c in r[2:9]] for r in ws.iter_rows(min_row=3) if r[0].value}
        assert rows['Units'] == [1, 2, '=SUM(C3,D3)', 3, 4, 5, '=SUM(F3,G3,H3)']
        assert rows['Revenue'][2] == '=SUM(C5,D5)'
        # A stock carries the month's last day.
        assert rows['Stock'][2] == '=D6' and rows['Stock'][6] == '=H6'

    def test_the_days_group_under_their_month(self, tmp_path):
        ws = _book(tmp_path, ['month', 'year'])
        levels = {k: v.outline_level for k, v in ws.column_dimensions.items() if v.outline_level}
        # Days sit inside their month and the year; a month's total inside the year.
        assert levels['C'] == 2 and levels['F'] == 2
        assert levels['E'] == 1 and levels['I'] == 1
        # The year total sums the days themselves.
        row = next(r for r in ws.iter_rows() if r[0].value == 'Units')
        assert row[9].value == '=SUM(C4,D4,F4,G4,H4)'


class TestGrowthByTheCalendar:
    def test_a_day_grows_on_the_same_day_a_year_back(self):
        from modeleon.compile.excel.writer import _growth_by_calendar
        child = mo.Variable([float(i + 1) for i in range(400)])
        window = DefaultWindow('2025-12-30', 'day', 400)
        out = _growth_by_calendar(child, window, 'day', 400, child.value)
        assert out is not None
        # 2026-12-30 is index 365; a year back is index 0.
        assert out.value[365] == 366.0 / 1.0
        assert all(v is None for v in out.value[:365])
        assert out._metric_no_bucket

    def test_a_window_shorter_than_a_year_has_no_growth(self):
        from modeleon.compile.excel.writer import _growth_by_calendar
        child = mo.Variable([1.0] * 30)
        assert _growth_by_calendar(child, DefaultWindow('2026-01-01', 'day', 30), 'day', 30, child.value) is None
