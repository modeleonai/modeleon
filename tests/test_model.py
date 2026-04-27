# SPDX-License-Identifier: Apache-2.0
"""Tests for the top-level ``MultiVariable`` container.

Style note: every test uses **explicit attribute assignment**
(``mv.child = mo.Variable(...)``). The ``with`` block is purely a
short-alias convenience — it never auto-attaches anything.
"""

from openpyxl import load_workbook

import modeleon as mo


class TestRole:
    """MultiVariables carry a ``role``. Plain ones are ``'group'``;
    ``excel_props={'tab': True}`` marks an Excel tab."""

    def test_sheet_role(self):
        s = mo.MultiVariable("P&L", excel_props={'tab': True})
        assert s._role == 'sheet'
        assert s._is_sheet is True

    def test_plain_mv_role(self):
        mv = mo.MultiVariable(display_name="section")
        assert mv._role == 'group'
        assert mv._is_sheet is False


class TestComponentAttachment:
    """Components attach via ``parent.child = ...`` only."""

    def test_attribute_assignment_attaches_sub_mv(self):
        model = mo.MultiVariable("Acme")
        model.pnl = mo.MultiVariable("P&L", excel_props={'tab': True})

        assert "pnl" in model._components
        assert model._components["pnl"] is model.pnl
        assert model.pnl._role == 'sheet'
        assert model.pnl._parent is model

    def test_attribute_assignment_attaches_variable(self):
        sheet = mo.MultiVariable("P&L", excel_props={'tab': True})
        sheet.revenue = mo.Variable(1_000_000, display_name='Revenue')

        assert "revenue" in sheet._components
        assert sheet.revenue._owner is sheet

    def test_with_block_is_just_an_alias(self):
        """``with mv as alias:`` binds ``alias`` to the same object as
        ``mv`` for the body. It does not auto-attach anything."""
        model = mo.MultiVariable("Acme")
        model.pnl = mo.MultiVariable("P&L", excel_props={'tab': True})
        with model.pnl as pnl:
            assert pnl is model.pnl
            pnl.revenue = mo.Variable(1_000_000)
        assert "revenue" in model.pnl._components

    def test_multiple_sheets(self):
        model = mo.MultiVariable("Acme")
        model.pnl = mo.MultiVariable("P&L", excel_props={'tab': True})
        model.pnl.revenue = mo.Variable(1_000_000)
        model.bs = mo.MultiVariable("Balance Sheet", excel_props={'tab': True})
        model.bs.cash = mo.Variable(100_000)

        assert set(model._components.keys()) == {"pnl", "bs"}

    def test_variables_belong_to_their_owning_sheet_only(self):
        """A Variable assigned to ``pnl`` is in ``pnl._components`` —
        it is NOT also a child of the model."""
        model = mo.MultiVariable("Acme")
        model.pnl = mo.MultiVariable("P&L", excel_props={'tab': True})
        model.pnl.revenue = mo.Variable(1_000_000, display_name='Revenue')
        model.pnl.cogs = model.pnl.revenue * 0.6

        assert "revenue" in model.pnl._components
        assert "cogs" in model.pnl._components
        assert "revenue" not in model._components
        assert "cogs" not in model._components


class TestToExcel:
    """``model.to_excel()`` emits only the sheets reachable from the root."""

    def test_single_model_emits_its_sheets(self, tmp_path):
        model = mo.MultiVariable("Acme")
        model.pnl = mo.MultiVariable("P&L", excel_props={'tab': True})
        model.pnl.revenue = mo.Variable(1_000_000, display_name='Revenue')
        model.pnl.cogs = model.pnl.revenue * 0.6

        model.to_excel(tmp_path / "acme.xlsx")

        path = tmp_path / "acme.xlsx"
        wb = load_workbook(path)
        assert wb.sheetnames == ["P&L"]

    def test_two_models_no_cross_pollution(self, tmp_path):
        """Emission scope is per-Model — a.xlsx doesn't see b's sheets."""
        a = mo.MultiVariable("A")
        a.sheet_a = mo.MultiVariable("Sheet A", excel_props={'tab': True})
        a.sheet_a.x = mo.Variable(100, display_name='X')

        b = mo.MultiVariable("B")
        b.sheet_b = mo.MultiVariable("Sheet B", excel_props={'tab': True})
        b.sheet_b.y = mo.Variable(200, display_name='Y')

        a.to_excel(tmp_path / "a.xlsx")

        a_path = tmp_path / "a.xlsx"
        b.to_excel(tmp_path / "b.xlsx")
        b_path = tmp_path / "b.xlsx"
        assert load_workbook(a_path).sheetnames == ["Sheet A"]
        assert load_workbook(b_path).sheetnames == ["Sheet B"]

    def test_formula_resolves_within_model(self, tmp_path):
        """Cross-sheet refs inside a Model produce valid Excel formulas."""
        model = mo.MultiVariable("Acme")
        model.pnl = mo.MultiVariable("P&L", excel_props={'tab': True})
        model.pnl.revenue = mo.Variable(1_000_000, display_name='Revenue')
        model.pnl.cogs = model.pnl.revenue * 0.6
        model.summary = mo.MultiVariable("Summary", excel_props={'tab': True})
        model.summary.pnl_cogs = model.pnl.cogs * 2

        model.to_excel(tmp_path / "acme.xlsx")

        path = tmp_path / "acme.xlsx"
        wb = load_workbook(path)
        summary = wb["Summary"]
        for row in summary.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str) and cell.value.startswith("="):
                    assert "{" not in cell.value


class TestPolicyAExternalRefs:
    """When a subtree's formulas reference Variables outside the subtree,
    the translator inlines the external value as an Excel literal
    (Policy A) instead of emitting a broken ``_var_N`` identifier."""

    def test_external_scalar_ref_inlines_value(self, tmp_path):
        rate = mo.Variable(0.05, display_name='Rate')  # lives outside any sheet
        model = mo.MultiVariable("Solo")
        model.derived = mo.MultiVariable("Derived", excel_props={'tab': True})
        model.derived.result = rate * 1000

        model.to_excel(tmp_path / "solo.xlsx")

        path = tmp_path / "solo.xlsx"
        wb = load_workbook(path)
        formulas = [
            cell.value for row in wb["Derived"].iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        ]
        assert any("0.05" in f and "_var_" not in f for f in formulas), formulas

    def test_external_list_ref_inlines_per_period_value(self, tmp_path):
        # Under new recurrence semantics period 0 = start (no formula),
        # periods 1..N-1 apply the formula. So growth[1] and growth[2]
        # land in formulas; growth[0] is unused.
        growth = mo.Variable([0.05, 0.06, 0.07], display_name='Growth')
        model = mo.MultiVariable("GrowthModel")
        model.revenue_sheet = mo.MultiVariable("Revenue", excel_props={'tab': True})
        model.revenue_sheet.revenue = mo.recurrence(
            start=1000,
            formula="{prev} * (1 + {g})",
            g=growth,
            periods=3,
        )

        model.to_excel(tmp_path / "growth.xlsx")

        path = tmp_path / "growth.xlsx"
        wb = load_workbook(path)
        formulas = [
            cell.value for row in wb["Revenue"].iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        ]
        assert any("0.06" in f for f in formulas)
        assert any("0.07" in f for f in formulas)
        assert all("_var_" not in f for f in formulas), formulas


class TestSheetAsRoot:
    """A Sheet alone is a valid emission root — no Model wrapper required."""

    def test_sheet_emits_standalone(self, tmp_path):
        sheet = mo.MultiVariable("Standalone", excel_props={'tab': True})
        sheet.v = mo.Variable(42, display_name='V')

        sheet.to_excel(tmp_path / "sheet.xlsx")

        path = tmp_path / "sheet.xlsx"
        assert load_workbook(path).sheetnames == ["Standalone"]


class TestAnyMVCanEmit:
    """``.to_excel()`` lives on MultiVariableBase — any MV can emit itself."""

    def test_sheet_directly_emits(self, tmp_path):
        sheet = mo.MultiVariable("P&L", excel_props={'tab': True})
        sheet.revenue = mo.Variable(1_000_000, display_name='Revenue')
        sheet.cogs = sheet.revenue * 0.6
        sheet.to_excel(tmp_path / "pnl.xlsx")
        path = tmp_path / "pnl.xlsx"
        assert load_workbook(path).sheetnames == ["P&L"]

    def test_plain_multivariable_with_sheets_emits(self, tmp_path):
        """A plain MultiVariable acting as a root emits its child Sheets."""
        root = mo.MultiVariable(display_name="root")
        root.sheet_a = mo.MultiVariable("A", excel_props={'tab': True})
        root.sheet_a.x = mo.Variable(1, display_name='X')
        root.sheet_b = mo.MultiVariable("B", excel_props={'tab': True})
        root.sheet_b.y = mo.Variable(2, display_name='Y')
        root.to_excel(tmp_path / "any_mv.xlsx")
        path = tmp_path / "any_mv.xlsx"
        assert sorted(load_workbook(path).sheetnames) == ["A", "B"]
