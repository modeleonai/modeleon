# SPDX-License-Identifier: Apache-2.0
"""A ``+`` chain too long for one Excel formula folds into ``SUM``.

A sum of a thousand rows renders as ``=A!B2 + A!B3 + …`` and stays
live up to Excel's 8,192 characters; past them the cell used to hold
its number. When the terms are exactly the cells of one run — side by
side along a row or down a column, in order — a hand-built book writes
``=SUM(A!B2:B1001)``, and so does the writer now. Any other chain past
the limit (a minus in it, a term that is not a cell) still ships its
number: a range never stands for something it is not.
"""

import functools
import logging
import operator
import shutil
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook

import modeleon as mo

#: Every reference repeats the quoted sheet name — ~40 characters a
#: term, so 300 rows run past Excel's 8,192.
LONG = "Inputs of a long portfolio"


def _rows(ws):
    out = {}
    for row in ws.iter_rows():
        label = row[0].value
        if label:
            out[str(label)] = [c.value for c in row[1:]]
    return out


def _model(n, build):
    m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=2)
    lines = {f"r{i}": mo.Variable([float(i), i + 0.5], display_name=f"R{i}") for i in range(n)}
    with m:
        m.inputs = mo.MultiVariable(LONG, **lines)
        m.t = mo.MultiVariable("T")
        rows = [getattr(m.inputs, f"r{i}") for i in range(n)]
        m.t.total = build(rows, m).set_display_name("Total")
    return m


def _total(model, tmp_path):
    path = tmp_path / "book.xlsx"
    model.to_excel(path)
    return _rows(load_workbook(path)["T"])["Total"]


class TestAChainPastTheLimit:
    def test_folds_into_one_range(self, tmp_path, caplog):
        m = _model(300, lambda rows, m: functools.reduce(operator.add, rows))
        with caplog.at_level(logging.WARNING, logger="modeleon"):
            total = _total(m, tmp_path)
        assert total[:2] == [f"=SUM('{LONG}'!B2:B301)", f"=SUM('{LONG}'!C2:C301)"], total
        assert "past Excel's limit" not in caplog.text

    def test_with_a_minus_in_it_stays_a_number(self, tmp_path, caplog):
        m = _model(300, lambda rows, m: functools.reduce(operator.add, rows[:-1]) - rows[-1])
        with caplog.at_level(logging.WARNING, logger="modeleon"):
            total = _total(m, tmp_path)
        assert total[:2] == pytest.approx(m.t.total.value)
        assert "past Excel's limit of 8,192" in caplog.text

    def test_with_a_term_that_is_not_a_cell_stays_a_number(self, tmp_path, caplog):
        m = _model(300, lambda rows, m: functools.reduce(operator.add, rows) + 1.0)
        with caplog.at_level(logging.WARNING, logger="modeleon"):
            total = _total(m, tmp_path)
        assert total[:2] == pytest.approx(m.t.total.value)
        assert "past Excel's limit of 8,192" in caplog.text


class TestAChainWithinTheLimit:
    def test_stays_a_plain_chain(self, tmp_path):
        m = _model(30, lambda rows, m: functools.reduce(operator.add, rows))
        total = _total(m, tmp_path)
        assert total[0].startswith(f"='{LONG}'!B2 + '{LONG}'!B3 + ") and "SUM(" not in total[0]
        assert total[0].count(" + ") == 29


def _soffice():
    for candidate in ("soffice", "libreoffice"):
        found = shutil.which(candidate)
        if found:
            return found
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


@pytest.mark.slow
def test_the_workbook_recalculates_to_the_python_values(tmp_path):
    soffice = _soffice()
    if soffice is None:
        pytest.skip("LibreOffice not installed")
    m = _model(300, lambda rows, m: functools.reduce(operator.add, rows))
    path = tmp_path / "long.xlsx"
    m.to_excel(path)
    subprocess.run(
        [soffice, f"-env:UserInstallation={(tmp_path / 'prof').as_uri()}",
         "--headless", "--convert-to", "xlsx", "--outdir",
         str(tmp_path / "calc"), str(path)],
        check=True, capture_output=True,
    )
    total = _rows(load_workbook(tmp_path / "calc" / "long.xlsx", data_only=True)["T"])["Total"]
    assert total[:2] == pytest.approx(m.t.total.value)
