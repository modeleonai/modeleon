# SPDX-License-Identifier: Apache-2.0
"""The workbook writes the loan schedule and depreciation functions as
live formulas — ``=IPMT(I!B1, 2, I!B3, I!B2)`` in every column of a
schedule row, an optional argument only when given — and LibreOffice
recalculates them to the engine's numbers.
"""

from __future__ import annotations

import shutil
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook

import modeleon as mo
from modeleon.functions.financial import (
    CUMIPMT,
    CUMPRINC,
    DB,
    DDB,
    IPMT,
    NPER,
    PPMT,
    PV,
    RATE,
    SLN,
    SYD,
)


def _rows(ws):
    out = {}
    for row in ws.iter_rows():
        label = row[0].value
        if label:
            out[str(label)] = [c.value for c in row[1:]]
    return out


def _book(model, tmp_path):
    path = tmp_path / "book.xlsx"
    model.to_excel(path)
    book = load_workbook(path)
    return {name: _rows(book[name]) for name in book.sheetnames}


def _model(periods=4):
    m = mo.Model("m", default_grain="year", default_start="2026", default_periods=periods)
    with m:
        m.i = mo.MultiVariable("I")
        m.i.rate = mo.Variable(0.1, display_name="rate")
        m.i.loan = mo.Variable(8000.0, display_name="loan")
        m.i.term = mo.Variable(periods, display_name="term")
        m.i.pmt = mo.Variable(-200.0, display_name="pmt")
        m.i.cost = mo.Variable(2400.0, display_name="cost")
        m.i.salvage = mo.Variable(300.0, display_name="salvage")
        m.s = mo.MultiVariable("S")
    return m


class TestTheFormulasWritten:
    def test_a_scalar_call_is_one_cell_with_the_optional_arguments_given(self, tmp_path):
        m = _model()
        with m:
            m.i.n = NPER(m.i.rate, m.i.pmt, m.i.loan).set_display_name("n")
            m.i.n1 = NPER(m.i.rate, m.i.pmt, m.i.loan, when=1).set_display_name("n1")
            m.i.r = RATE(m.i.term, m.i.pmt, m.i.loan).set_display_name("r")
            m.i.rg = RATE(m.i.term, m.i.pmt, m.i.loan, guess=0.05).set_display_name("rg")
            m.i.ci = CUMIPMT(m.i.rate, m.i.term, m.i.loan, 1, 2).set_display_name("ci")
            m.i.cp = CUMPRINC(m.i.rate, m.i.term, m.i.loan, 1, 2, 1).set_display_name("cp")
            m.i.sl = SLN(m.i.cost, m.i.salvage, m.i.term).set_display_name("sl")
            m.i.i2 = IPMT(m.i.rate, 2, m.i.term, m.i.loan, 0, 1).set_display_name("i2")
            m.i.p2 = PPMT(m.i.rate, 2, m.i.term, m.i.loan, 1000.0).set_display_name("p2")
        i = _book(m, tmp_path)["I"]
        assert i["n"] == ["=NPER(B1, B4, B2)"], i["n"]
        assert i["n1"] == ["=NPER(B1, B4, B2, 0, 1)"], i["n1"]
        assert i["r"] == ["=RATE(B3, B4, B2)"], i["r"]
        assert i["rg"] == ["=RATE(B3, B4, B2, 0, 0, 0.05)"], i["rg"]
        assert i["ci"] == ["=CUMIPMT(B1, B3, B2, 1, 2, 0)"], i["ci"]
        assert i["cp"] == ["=CUMPRINC(B1, B3, B2, 1, 2, 1)"], i["cp"]
        assert i["sl"] == ["=SLN(B5, B6, B3)"], i["sl"]
        assert i["i2"] == ["=IPMT(B1, 2, B3, B2, 0, 1)"], i["i2"]
        assert i["p2"] == ["=PPMT(B1, 2, B3, B2, 1000.0)"], i["p2"]

    def test_a_schedule_row_is_one_call_per_column_with_its_period(self, tmp_path):
        m = _model()
        with m:
            m.s.interest = IPMT(m.i.rate, mo.time.index, m.i.term, m.i.loan).set_display_name("i")
            m.s.principal = PPMT(m.i.rate, mo.time.index, m.i.term, m.i.loan).set_display_name("p")
            m.s.db = DB(m.i.cost, m.i.salvage, m.i.term, mo.time.index, 7).set_display_name("db")
            m.s.ddb = DDB(m.i.cost, m.i.salvage, m.i.term, mo.time.index).set_display_name("ddb")
            m.s.ddb15 = DDB(m.i.cost, m.i.salvage, m.i.term, mo.time.index,
                            1.5).set_display_name("ddb15")
            m.s.syd = SYD(m.i.cost, m.i.salvage, m.i.term, mo.time.index).set_display_name("syd")
            m.s.lit = IPMT(0.1, [1, 2, 3, 4], 4, 8000.0).set_display_name("lit")
        s = _book(m, tmp_path)["S"]
        assert s["i"] == [f"=IPMT(I!B1, {p}, I!B3, I!B2)" for p in (1, 2, 3, 4)], s["i"]
        assert s["p"] == [f"=PPMT(I!B1, {p}, I!B3, I!B2)" for p in (1, 2, 3, 4)], s["p"]
        assert s["db"] == [f"=DB(I!B5, I!B6, I!B3, {p}, 7)" for p in (1, 2, 3, 4)], s["db"]
        assert s["ddb"] == [f"=DDB(I!B5, I!B6, I!B3, {p})" for p in (1, 2, 3, 4)], s["ddb"]
        assert s["ddb15"] == [f"=DDB(I!B5, I!B6, I!B3, {p}, 1.5)" for p in (1, 2, 3, 4)], s["ddb15"]
        assert s["syd"] == [f"=SYD(I!B5, I!B6, I!B3, {p})" for p in (1, 2, 3, 4)], s["syd"]
        assert s["lit"] == [f"=IPMT(0.1, {p}, 4, 8000.0)" for p in (1, 2, 3, 4)], s["lit"]

    def test_the_cached_values_are_the_engine_s(self, tmp_path):
        m = _model()
        with m:
            m.s.interest = IPMT(m.i.rate, mo.time.index, m.i.term, m.i.loan).set_display_name("i")
            m.i.r = RATE(m.i.term, m.i.pmt, m.i.loan).set_display_name("r")
        path = tmp_path / "book.xlsx"
        m.to_excel(path)
        book = load_workbook(path, data_only=True)
        assert _rows(book["S"])["i"][:4] == [pytest.approx(v) for v in m.s.interest.value]
        assert _rows(book["I"])["r"][0] == pytest.approx(m.i.r.value)


# ══════════════════════════════════════════════════════════════════════
# The spreadsheet computes the same numbers
# ══════════════════════════════════════════════════════════════════════


def _soffice():
    for candidate in ("soffice", "libreoffice"):
        found = shutil.which(candidate)
        if found:
            return found
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


def _parity_model():
    """Every loan schedule and depreciation function, several argument sets each: the
    documented examples, a payment at the start of the period (type 1), a
    zero rate and a rate of -100%, a future value, a guess, fractional
    periods, the rates LibreOffice finds only from its further starts,
    lives under two periods, a rate and periods read as a cell shows them,
    numeric text, schedule rows over the model's periods, the cells that
    are errors on both sides, and the cells LibreOffice reads its own way."""
    m = _model(periods=5)
    with m:
        i, s = m.i, m.s
        i.zero = mo.Variable(0.0, display_name="zero")
        i.r01 = mo.Variable(0.01, display_name="r01")
        i.pmt100 = mo.Variable(-100.0, display_name="pmt100")
        i.pv1000 = mo.Variable(-1000.0, display_name="pv1000")
        i.fv10000 = mo.Variable(10000.0, display_name="fv10000")
        i.monthly = mo.Variable(0.09 / 12, display_name="monthly")
        i.n360 = mo.Variable(360, display_name="n360")
        i.big = mo.Variable(125000.0, display_name="big")
        i.months = mo.Variable(7, display_name="months")
        i.life4 = mo.Variable(4, display_name="life4")
        i.loan_text = mo.Variable("8000", display_name="loan_text")
        # NPER
        i.n_t1 = NPER(i.r01, i.pmt100, i.pv1000, i.fv10000, 1).set_display_name("n_t1")
        i.n_t0 = NPER(i.r01, i.pmt100, i.pv1000, i.fv10000).set_display_name("n_t0")
        i.n_nofv = NPER(i.r01, i.pmt100, i.pv1000).set_display_name("n_nofv")
        i.n_zero = NPER(i.zero, i.pmt100, -i.pv1000).set_display_name("n_zero")
        # IPMT / PPMT, one cell each
        i.i_t1 = IPMT(i.rate, 2, 3, i.loan, 0, 1).set_display_name("i_t1")
        i.i_fv = IPMT(i.rate, 2, 3, i.loan, 1000.0).set_display_name("i_fv")
        i.i_zero = IPMT(i.zero, 2, 10, i.loan).set_display_name("i_zero")
        i.i_frac = IPMT(i.rate, 2.5, 3, i.loan).set_display_name("i_frac")
        i.p_t1 = PPMT(i.rate, 2, 3, i.loan, 0, 1).set_display_name("p_t1")
        i.p_zero = PPMT(i.zero, 2, 10, i.loan).set_display_name("p_zero")
        # A rate of -100%, and a first period on a term too long for the
        # payment: computed as the spreadsheet computes them.
        i.i_m1 = IPMT(-1, 1, 3, i.loan).set_display_name("i_m1")
        i.p_m1 = PPMT(-1, 1, 3, i.loan).set_display_name("p_m1")
        i.p_m1b = PPMT(-1, 2, 3, i.loan).set_display_name("p_m1b")
        i.i_long = IPMT(i.rate, 1, 10000, i.loan).set_display_name("i_long")
        # A small rate: (1 + rate) ** nper - 1 loses no digits.
        i.p_small = PPMT(1e-12, 2, 10, 1000.0).set_display_name("p_small")
        i.cp_small = CUMPRINC(1e-9, 10, 1000.0, 1, 10, 0).set_display_name("cp_small")
        # RATE
        i.r_a = RATE(48, i.pmt, i.loan).set_display_name("r_a")
        i.r_t1 = RATE(48, i.pmt, i.loan, 0, 1).set_display_name("r_t1")
        i.r_fv = RATE(10, i.pmt100, 0, 1500.0).set_display_name("r_fv")
        i.r_g = RATE(48, i.pmt, i.loan, 0, 0, 0.05).set_display_name("r_g")
        i.r_frac = RATE(10.5, i.pmt100, 800.0).set_display_name("r_frac")
        i.r_neg = RATE(5, 50.0, i.pv1000).set_display_name("r_neg")
        # Newton from 10% runs below -100% on a loan repaid in a few
        # payments its own size, and the spreadsheet starts again at 20%,
        # 5%, 30%, ...; a long term settles inside the 150 steps.
        i.r_high = RATE(5, -1000.0, 1000.0).set_display_name("r_high")
        i.r_twelve = RATE(12, -1000.0, 1000.0).set_display_name("r_twelve")
        i.r_long0 = RATE(1000, -1.0, 1000.0).set_display_name("r_long0")
        i.r_long = RATE(1200, -1.0, 1000.0).set_display_name("r_long")
        i.r_due = RATE(3, -600.0, 1000.0, 0, 1).set_display_name("r_due")
        i.r_frac_high = RATE(5.5, -1000.0, 1000.0).set_display_name("r_frac_high")
        i.r_hardly = RATE(2, -1.0, 1000.0).set_display_name("r_hardly")
        # No flow at all: every rate balances, the guess is the answer.
        i.r_nothing = RATE(10, i.zero, i.zero).set_display_name("r_nothing")
        # Numeric text, written as a literal and held in a cell.
        i.r_text = RATE(48, i.pmt, "8000").set_display_name("r_text")
        i.r_text_cell = RATE(48, i.pmt, i.loan_text).set_display_name("r_text_cell")
        # CUMIPMT / CUMPRINC
        i.ci = CUMIPMT(i.monthly, i.n360, i.big, 13, 24, 0).set_display_name("ci")
        i.ci_t1 = CUMIPMT(i.monthly, i.n360, i.big, 1, 12, 1).set_display_name("ci_t1")
        i.ci_frac = CUMIPMT(i.monthly, i.n360, i.big, 13.7, 24.2, 0).set_display_name("ci_frac")
        i.cp = CUMPRINC(i.monthly, i.n360, i.big, 13, 24, 0).set_display_name("cp")
        i.cp_t1 = CUMPRINC(i.monthly, i.n360, i.big, 1, 12, 1).set_display_name("cp_t1")
        # 12.999999999999998 shows as 13: the count starts at period 13.
        i.ci_shown = CUMIPMT(i.monthly, i.n360, i.big, 12.999999999999998, 24,
                             0).set_display_name("ci_shown")
        i.cp_shown = CUMPRINC(i.monthly, i.n360, i.big, 12.999999999999998, 24,
                              0).set_display_name("cp_shown")
        # SLN
        i.sl = SLN(i.cost, i.salvage, i.term).set_display_name("sl")
        # DB: a period that reads as 1 is the first period's charge, even
        # past a life under two periods; the rate, a period and a month
        # are read as a cell shows them.
        i.db_short = DB(1000.0, 100.0, 0.5, 1).set_display_name("db_short")
        i.db_short15 = DB(1000.0, 100.0, 0.5, 1.5, 6).set_display_name("db_short15")
        i.db_life09 = DB(1000.0, 100.0, 0.9, 1, 6).set_display_name("db_life09")
        i.db_life1 = DB(1000.0, 100.0, 1, 1.5, 6).set_display_name("db_life1")
        i.db_life12 = DB(1000.0, 100.0, 1.2, 1.5, 6).set_display_name("db_life12")
        i.db_stub12 = DB(1000.0, 100.0, 1.2, 2, 6).set_display_name("db_stub12")
        i.db_round = DB(i.cost, 318.0, 1, 1).set_display_name("db_round")
        i.db_round7 = DB(i.cost, 318.0, 1, 2, i.months).set_display_name("db_round7")
        i.db_shown = DB(1000.0, 100.0, 5, 2.9999999999999996).set_display_name("db_shown")
        i.db_month_shown = DB(1000.0, 100.0, 5, 1,
                              2.9999999999999996).set_display_name("db_month_shown")
        # Schedule rows over the five periods: a five-year loan, a
        # four-year asset whose fifth period is DB's stub, and one of
        # four and a half years.
        s.interest = IPMT(i.rate, mo.time.index, i.term, i.loan).set_display_name("interest")
        s.principal = PPMT(i.rate, mo.time.index, i.term, i.loan).set_display_name("principal")
        s.interest_t1 = IPMT(i.rate, mo.time.index, i.term, i.loan, 0,
                             1).set_display_name("interest_t1")
        s.principal_t1 = PPMT(i.rate, mo.time.index, i.term, i.loan, 0,
                              1).set_display_name("principal_t1")
        s.interest_zero = IPMT(i.zero, mo.time.index, i.term,
                               i.loan).set_display_name("interest_zero")
        s.principal_zero = PPMT(i.zero, mo.time.index, i.term,
                                i.loan).set_display_name("principal_zero")
        s.nper_row = NPER(i.rate, i.pmt100, mo.time.index * 100.0).set_display_name("nper_row")
        s.rate_row = RATE(mo.time.index + 5, i.pmt100, 800.0).set_display_name("rate_row")
        s.db = DB(i.cost, i.salvage, i.life4, mo.time.index, i.months).set_display_name("db")
        s.db12 = DB(i.cost, i.salvage, i.life4, mo.time.index).set_display_name("db12")
        s.db_frac = DB(i.cost, i.salvage, 4.5, mo.time.index, i.months).set_display_name("db_frac")
        s.ddb = DDB(i.cost, i.salvage, i.term, mo.time.index).set_display_name("ddb")
        s.ddb15 = DDB(i.cost, i.salvage, i.term, mo.time.index, 1.5).set_display_name("ddb15")
        s.syd = SYD(i.cost, i.salvage, i.term, mo.time.index).set_display_name("syd")
        # Errors on both sides: a period past the term, a rate with no
        # root, an iteration that runs off from the guess given or from
        # every start.
        i.i_past = IPMT(i.rate, 6, i.term, i.loan).set_display_name("i_past")
        i.r_none = RATE(10, 100.0, 1000.0).set_display_name("r_none")
        i.r_poor = RATE(10, i.pmt100, 800.0, 0, 0, -0.5).set_display_name("r_poor")
        i.r_runaway = RATE(2, -5000.0, 1000.0).set_display_name("r_runaway")
        # The cells LibreOffice reads its own way (DIFFERS below).
        i.ci_t09 = CUMIPMT(i.monthly, i.n360, i.big, 1, 12, 0.9).set_display_name("ci_t09")
        i.syd_past = SYD(i.cost, i.salvage, i.life4, 5).set_display_name("syd_past")
        i.db_half = DB(i.cost, i.salvage, i.life4, 0.5).set_display_name("db_half")
        i.r_trap = RATE(120, i.pmt100, 1000.0, 0, 1).set_display_name("r_trap")
        i.r_flat = RATE(10, i.zero, 1000.0).set_display_name("r_flat")
    return m


SCALARS = ("n_t1", "n_t0", "n_nofv", "n_zero", "i_t1", "i_fv", "i_zero", "i_frac", "p_t1",
           "p_zero", "i_m1", "p_m1", "p_m1b", "i_long", "p_small", "cp_small", "r_a", "r_t1",
           "r_fv", "r_g", "r_frac", "r_neg", "r_high", "r_twelve", "r_long0", "r_long",
           "r_due", "r_frac_high", "r_hardly", "r_nothing", "r_text", "r_text_cell", "ci",
           "ci_t1", "ci_frac", "cp", "cp_t1", "ci_shown", "cp_shown", "sl", "db_short",
           "db_short15", "db_life09", "db_life1", "db_life12", "db_stub12", "db_round",
           "db_round7", "db_shown", "db_month_shown")
#: Errors in the engine and in LibreOffice, which spells them its own way:
#: Err:502 is written as #VALUE! for a period past the term, and a RATE
#: that does not converge as #N/A, where Excel shows #NUM!, the engine's
#: marker.
ERRORS = ("i_past", "r_none", "r_poor", "r_runaway")
#: Where LibreOffice departs from Excel's documentation, which the engine
#: follows, and what its cell holds instead. A fractional ``type`` in
#: CUMIPMT is truncated by Excel and refused by LibreOffice (#VALUE!).
#: SYD past the life and DB before period 1 are outside Excel's domain
#: (#NUM!), where LibreOffice computes the formula as written (0). RATE's
#: iteration ends, in LibreOffice, at -100% on an annuity due (the engine
#: goes on to the rate the payments carry) and at -98.2% on a loan with no
#: payment, which only -100% balances (the engine finds no rate). Each
#: pins the engine's reading and lets the spreadsheet refuse the cell,
#: agree, or hold its own value.
DIFFERS = {"ci_t09": None, "syd_past": 0, "db_half": 0, "r_trap": -1.0,
           "r_flat": -0.981934476404913}
ROWS = ("interest", "principal", "interest_t1", "principal_t1", "interest_zero",
        "principal_zero", "nper_row", "rate_row", "db", "db12", "db_frac", "ddb", "ddb15", "syd")


@pytest.mark.slow
def test_the_workbook_recalculates_to_the_python_values(tmp_path):
    soffice = _soffice()
    if soffice is None:
        pytest.skip("LibreOffice not installed")
    m = _parity_model()
    path = tmp_path / "financial.xlsx"
    m.to_excel(path)
    subprocess.run(
        [soffice, f"-env:UserInstallation={(tmp_path / 'prof').as_uri()}",
         "--headless", "--convert-to", "xlsx", "--outdir",
         str(tmp_path / "calc"), str(path)],
        check=True, capture_output=True,
    )
    book = load_workbook(tmp_path / "calc" / "financial.xlsx", data_only=True)
    i = _rows(book["I"])
    s = _rows(book["S"])
    # The engine's numbers are not the spreadsheet's defaults: a wrong
    # one would be caught here.
    assert m.i.n_t1.value == pytest.approx(59.6738657, abs=5e-8)
    assert m.s.db.value[4] == pytest.approx(DB(2400, 300, 4, 5, 7).value) and m.s.db.value[4] > 0
    assert m.i.db_short.value == 990.0 and m.i.db_round.value == pytest.approx(2083.2)
    for name in SCALARS:
        assert i[name][0] == pytest.approx(getattr(m.i, name).value, rel=1e-9, abs=1e-9), name
    for name in ROWS:
        got = s[name][:5]
        assert got == pytest.approx(getattr(m.s, name).value, rel=1e-9, abs=1e-9), name
    for name in ERRORS:
        assert getattr(m.i, name).value == '#NUM!' and str(i[name][0]).startswith("#"), name
    # The cells were recalculated, not handed back as the engine cached
    # them: LibreOffice spells some of these errors its own way.
    assert {str(i[name][0]) for name in ERRORS} != {'#NUM!'}
    # The engine's own readings, exactly.
    assert m.i.ci_t09.value == pytest.approx(CUMIPMT(0.09 / 12, 360, 125000.0, 1, 12, 0).value)
    assert m.i.syd_past.value == m.i.db_half.value == m.i.r_flat.value == '#NUM!'
    assert m.i.r_trap.value == pytest.approx(0.1111107124217, rel=1e-9)
    assert PV(m.i.r_trap.value, 120, -100.0, 0, 1).value == pytest.approx(1000.0)
    for name, own in DIFFERS.items():
        got, ours = i[name][0], getattr(m.i, name).value
        allowed = [v for v in (ours, own) if isinstance(v, int | float)]
        assert str(got).startswith("#") or any(got == pytest.approx(v) for v in allowed), \
            (name, got)
