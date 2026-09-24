"""Track-authored formulas read their operands at the SAME coordinate.

The formula of a ``plan=`` track must not resolve its tracked operands
to their blend/live rows — that would be a formula that computes a
different number than the cell holds. These pin the law: inside a
track coordinate a tracked operand means its same-track row (or its
value when that row is not laid out), a given track filled from a
shared list is a real splice (the head stays data), and the track
writer never publishes a coordinate's expression as the line's.
"""

import os
import tempfile

import openpyxl

import modeleon as mo
from modeleon.compile.excel.view import ExcelView


def _book(model, name):
    path = os.path.join(tempfile.mkdtemp(), "t.xlsx")
    model.to_excel(path)
    return openpyxl.load_workbook(path)[name]


def _rows(ws):
    return {ws.cell(r, 1).value: r for r in range(1, ws.max_row + 1)}


def _cells(ws, label):
    r = _rows(ws)[label]
    return [ws.cell(r, c).value for c in range(2, ws.max_column + 1)
            if ws.cell(r, c).value is not None]


def _model(mode):
    m = mo.Model("m", display_name="M",
                 tracks=mo.Tracks("plan", "actual",
                                  blend=mo.blend(given="actual", follow="plan",
                                                 until="2026-02")),
                 default_grain="month", default_start="2026-01",
                 default_periods=4)
    m.default_excel_view = ExcelView(tracks=mode)
    m.s = mo.MultiVariable("S", excel_props={"tab": True})
    with m.s as s:
        s.driver = mo.Variable(plan=[1.0, 2.0, 3.0, 4.0],
                               actual=[10.0, 20.0, None, None],
                               display_name="Driver")
        s.rev = mo.Variable(plan=s.driver * 2, actual=[5.0, None, None, None],
                            display_name="Rev")
        s.alias = mo.Variable(plan=s.driver, display_name="Alias")
        s.traj = mo.Variable(plan=mo.recurrence(100.0, "{prev} + {g}", g=s.driver),
                             display_name="Traj")
        s.cum = mo.Variable(plan=mo.cumsum(s.driver), display_name="Cum")
        s.only = mo.Variable(plan=s.driver * 3, display_name="Only")
        s.mix = mo.Variable([9.0, 8.0, None, None], plan=s.driver * 2,
                            display_name="Mix")
    return m


def test_rows_book_reads_same_track_rows():
    ws = _book(_model("rows"), "S")
    rows = _rows(ws)
    dp = rows["Driver · plan"]
    assert _cells(ws, "Rev · plan")[0] == f"=B{dp} * 2"
    assert _cells(ws, "Alias · plan")[0] == f"=B{dp}"
    traj = _cells(ws, "Traj · plan")
    assert traj[1] == f"=B{rows['Traj · plan']} + C{dp}", traj
    cum = _cells(ws, "Cum · plan")
    assert cum[0] == f"=B{dp}" and cum[1] == f"=B{rows['Cum · plan']} + C{dp}", cum


def test_blend_book_degrades_operand_without_a_coordinate_row():
    # No plan subrow is laid out in the blend book, and Driver's head
    # is its live splice — so the plan formula inlines Driver's PLAN
    # value rather than pointing at a row holding a different number.
    ws = _book(_model("blend"), "S")
    only = _cells(ws, "Only")
    assert only[:2] == ["=1.0 * 3", "=2.0 * 3"], only


def test_given_filled_from_shared_list_keeps_the_splice():
    m = _model("rows")
    ws = _book(m, "S")
    rows = _rows(ws)
    # the actual subrow is data — its numbers stand, no plan formula
    assert _cells(ws, "Mix · actual")[:2] == [9.0, 8.0]
    # the head is the splice over its subrows, not the plan expression
    head = _cells(ws, "Mix")[0]
    assert head.startswith("=IF(") and "* 2" not in head, head
    # the value-side law: a real splice has no shown expression, while
    # a follow-only line (given absent) shows the follow's expression
    assert m.s.mix.shown_track_expr() is None
    assert m.s.only.shown_track_expr() is not None


def test_blend_book_mix_head_stays_a_value_row():
    ws = _book(_model("blend"), "S")
    mix = _cells(ws, "Mix")
    assert mix[:2] == [9.0, 8.0], mix
