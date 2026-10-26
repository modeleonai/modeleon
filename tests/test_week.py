# SPDX-License-Identifier: Apache-2.0
"""ISO weeks, and the ISO year that totals them.

A week runs Monday to Sunday and is numbered in its ISO year (week 1
holds the year's first Thursday): ``2026-W01`` starts on 29 December
2025, and 2026 has 53 weeks. A week lies inside one ISO year and in no
month, quarter or calendar year — so a weekly line rolls up to its ISO
year and is refused anywhere else, by name, never split between two
months.
"""
import warnings
from datetime import date

import openpyxl
import pytest

import modeleon as mo
from modeleon.compile.excel.layout import LayoutEngine
from modeleon.compile.excel.totals import TotalsWarning, build_totals_plan, header_rows
from modeleon.core.time import _period_labels

from test_excel_time_header import _recalculated, _serial


def _weekly(n=6, start='2026-W51', totals=None):
    kw = {'default_excel_view': mo.ExcelView(timeline={'totals': totals})} if totals else {}
    m = mo.Model('m', default_grain='week', default_start=start, default_periods=n, **kw)
    with m:
        m.units = mo.Variable([float(i + 1) for i in range(n)], regrain=mo.up('sum'), display_name='Units')
        m.level = mo.Variable([float(10 + i) for i in range(n)], regrain=mo.up('last'), display_name='Level')
    return m


class TestTheWindow:
    def test_weeks_run_monday_to_sunday_across_week_53(self):
        assert _period_labels('2026-W51', None, 4, 'week', 'finance') == [
            'W51 2026', 'W52 2026', 'W53 2026', 'W01 2027']
        m = mo.Model('m', default_grain='week', default_start='2026-W01', default_periods=2)
        with m:
            m.start = mo.time.start
            m.end = mo.time.end
            m.days = mo.time.days
            m.index = mo.time.index
        assert m.start.value == [date(2025, 12, 29), date(2026, 1, 5)]
        assert m.end.value == [date(2026, 1, 4), date(2026, 1, 11)]
        assert m.days.value == [7, 7] and m.index.value == [1, 2]

    def test_a_week_window_starts_on_a_monday(self):
        m = mo.Model('m', default_grain='week', default_start=mo.DATE(2026, 1, 5), default_periods=2)
        assert m.default_start == '2026-W02'
        with pytest.raises(ValueError, match="A week window starts on a Monday; 2026-01-07 is a Wednesday. Start it on 2026-01-05"):
            mo.Model('m', default_grain='week', default_start=mo.DATE(2026, 1, 7), default_periods=2)

    def test_a_week_has_no_months(self):
        m = mo.Model('m', default_grain='week', default_start='2026-W01', default_periods=2)
        with m:
            with pytest.raises(ValueError, match="A week window has none"):
                mo.time.months


class TestRollUps:
    def test_weeks_roll_up_to_their_iso_year(self):
        m = _weekly()
        iy = m.units.at('iso_year')
        assert iy.value == [6.0, 15.0]           # W51–W53 2026, W01–W03 2027
        assert iy.time.start == '2026' and iy.time.grain == 'iso_year'
        assert m.level.at('iso_year').value == [12.0, 15.0]

    def test_weeks_never_roll_up_to_months(self):
        m = _weekly()
        with pytest.raises(ValueError, match="A week does not lie inside one month: weeks straddle months"):
            m.units.at('month')
        with pytest.raises(ValueError, match="weeks straddle years"):
            m.units.at('year')

    def test_days_roll_up_to_weeks(self):
        d = mo.Model('d', default_grain='day', default_start='2026-01-01', default_periods=10)
        with d:
            d.x = mo.Variable([1.0] * 10, regrain=mo.up('sum'))
        assert d.x.at('week').value == [4.0, 6.0]        # Thu–Sun of W01, then W02's Mon–Sat
        assert d.x.at('week').time.start == '2026-W01'

    def test_a_week_line_and_a_month_line_never_meet(self):
        w = _weekly()
        m = mo.Model('mm', default_grain='month', default_start='2026-12', default_periods=2)
        with m:
            m.a = mo.Variable([1.0, 2.0])
        with pytest.raises(ValueError, match="grains that do not nest: months straddle weeks"):
            m.a + w.units


class TestTheBook:
    def test_iso_year_totals_over_weeks(self, tmp_path):
        m = _weekly(totals=['iso_year'])
        path = tmp_path / 'w.xlsx'
        m.to_excel(str(path))
        ws = openpyxl.load_workbook(str(path))['M']
        assert [c.value for c in ws[1]][:9] == [None, '2026', None, None, 'Total 2026', '2027', None, None, 'Total 2027']
        assert [c.value for c in ws[2]][:9] == [None, 'W51', 'W52', 'W53', None, 'W01', 'W02', 'W03', None]
        units = next(r for r in ws.iter_rows() if r[0].value == 'Units')
        assert [c.value for c in units[1:9]] == [1, 2, 3, '=SUM(B3,C3,D3)', 4, 5, 6, '=SUM(F3,G3,H3)']
        level = next(r for r in ws.iter_rows() if r[0].value == 'Level')
        assert level[4].value == '=D4' and level[8].value == '=H4'

    def test_week_totals_over_days_under_the_iso_year(self):
        plan = build_totals_plan(('week', 'iso_year'), 'day', '2025-12-31', 6, 2)
        rows = header_rows(plan, _period_labels('2025-12-31', None, 6, 'day', 'finance'))
        assert [r['cells'][0]['kind'] for r in rows] == ['iso_year', 'week', 'period']
        # Six days from Wednesday 31 Dec: W01's five, then Monday of W02 — a
        # partial week the window closes.
        assert [c['label'] for c in rows[1]['cells']] == ['W01', 'Total W01', 'W02', 'Total W02']

    def test_a_weekly_header_is_a_live_chain(self, tmp_path):
        m = mo.Model('m', default_grain='week', default_start='2026-W01', default_periods=3)
        with m:
            m.header = mo.MultiVariable('Time')
            with m.header as h:
                h.start = mo.time.start
                h.end = mo.time.end
                h.days = mo.time.days
            m.default_header = m.header
            m.ops = mo.MultiVariable('Ops')
            with m.ops as o:
                o.fee = mo.Variable(10.0, display_name='Fee') * mo.time.days
        path = tmp_path / 'h.xlsx'
        with warnings.catch_warnings():
            warnings.simplefilter('ignore')
            m.to_excel(str(path))
        ws = openpyxl.load_workbook(str(path))['Time']
        rows = {r[0].value: [c.value for c in r[1:4]] for r in ws.iter_rows() if r[0].value}
        assert rows['Start'][1] == '=B3+1'                    # the previous end + 1
        assert rows['End'] == ['=B2+6', '=C2+6', '=D2+6']
        assert rows['Days'][0] == '=B3-B2+1'

    def test_a_totals_kind_a_week_cannot_draw_warns_and_is_not_drawn(self):
        m = _weekly(totals=['month'])
        with pytest.warns(TotalsWarning, match="'month' totals are not drawn: a week does not lie inside one month"):
            eng = LayoutEngine([m])
            eng.compute_addresses()
        assert eng.totals_notes and eng.totals_by_sheet == {}

    def test_weeks_and_months_never_tier_one_sheet(self):
        d = mo.Model('d', default_grain='day', default_start='2026-01-01', default_periods=10,
                     default_excel_view=mo.ExcelView(timeline={'totals': ['week', 'month']}))
        with d:
            d.x = mo.Variable([1.0] * 10, regrain=mo.up('sum'))
        with pytest.warns(TotalsWarning, match="'month' totals are not drawn beside 'week'"):
            eng = LayoutEngine([d])
            eng.compute_addresses()
        plan = next(iter(eng.totals_by_sheet.values()))
        assert plan.kinds == ('week',)

    @pytest.mark.slow
    def test_the_weekly_book_recalculates_to_the_engine(self, tmp_path):
        m = mo.Model('m', default_grain='week', default_start='2026-W51', default_periods=5,
                     default_excel_view=mo.ExcelView(timeline={'totals': ['iso_year']}))
        with m:
            m.header = mo.MultiVariable('Time')
            with m.header as h:
                h.start = mo.time.start
                h.end = mo.time.end
                h.days = mo.time.days
            m.default_header = m.header
            m.ops = mo.MultiVariable('Ops')
            with m.ops as o:
                o.rate = mo.Variable(2.0, display_name='Rate')
                o.fee = (o.rate * mo.time.days).set_display_name('Fee')
        wb = _recalculated(m, tmp_path)
        # W51 W52 W53 | Total 2026 | W01 W02 | Total 2027: the weeks are the
        # columns whose caption row says "Wnn".
        for sheet, mv, names in (('Time', m.header, ('start', 'end', 'days')), ('Ops', m.ops, ('fee',))):
            ws = wb[sheet]
            caption = next(r for r in ws.iter_rows(max_row=4)
                           if any(isinstance(c.value, str) and c.value.startswith('W') for c in r))
            weeks = [c.column for c in caption if isinstance(c.value, str) and c.value.startswith('W')]
            assert len(weeks) == 5
            rows = {r[0].value: r for r in ws.iter_rows() if r[0].value is not None}
            for name in names:
                var = getattr(mv, name)
                got = [_serial(rows[var.display_name][c - 1].value) for c in weeks]
                want = [_serial(v) for v in var.value]
                assert got == pytest.approx(want), (sheet, name, got, want)


def test_growth_by_the_week_a_year_back():
    from modeleon.compile.excel.writer import _growth_by_calendar
    from modeleon.core.time import DefaultWindow
    child = mo.Variable([float(i + 1) for i in range(60)])
    out = _growth_by_calendar(child, DefaultWindow('2025-W01', 'week', 60), 'week', 60, child.value)
    # 2026-W01 (index 52) grows on the week holding the same date a year
    # before — 2025-W01 (index 0).
    assert out.value[52] == 53.0 / 1.0
    assert all(v is None for v in out.value[:52])


def _iso_year_model(start):
    m = mo.Model('iy', default_grain='iso_year', default_periods=3)
    with m:
        m.inputs = mo.MultiVariable('Inputs')
        with m.inputs as i:
            i.go = mo.Variable(start, display_name='Go')
        m.default_start = m.inputs.go
        m.header = mo.MultiVariable('Time')
        with m.header as h:
            h.start = mo.time.start
            h.end = mo.time.end
            h.days = mo.time.days
            h.index = mo.time.index
        m.default_header = m.header
        m.ops = mo.MultiVariable('Ops')
        with m.ops as o:
            o.rate = mo.Variable(2.0, display_name='Rate')
            o.fee = (o.rate * mo.time.days).set_display_name('Fee')
    return m


class TestTheIsoYearBook:
    """An ISO year is 52 or 53 weeks from the Monday of the week holding
    4 January: its header is spelled off the start cell — the ISO year
    of a day is YEAR(day + 3), its first Monday DATE(y,1,4) -
    WEEKDAY(DATE(y,1,4),3)."""

    def test_an_iso_year_header_is_spelled_off_the_start(self):
        from modeleon.compile.excel.header import spell_anchor, spell_inline
        assert spell_inline('start', 1, 'iso_year', 'B2') == \
            '(DATE(YEAR(B2+3)+1,1,4)-WEEKDAY(DATE(YEAR(B2+3)+1,1,4),3))'
        cell = lambda f, j: None if j < 0 else f'{f[0].upper()}{j + 3}'  # noqa: E731
        assert spell_anchor('start', 1, 'iso_year', 'B2', cell) == 'E3+1'
        assert spell_anchor('end', 0, 'iso_year', 'B2', cell) == \
            'DATE(YEAR(S3+3)+1,1,4)-WEEKDAY(DATE(YEAR(S3+3)+1,1,4),3)-1'
        assert spell_anchor('days', 0, 'iso_year', 'B2', cell) == 'E3-S3+1'

    @pytest.mark.slow
    @pytest.mark.parametrize('start, days', [
        (date(2025, 12, 29), [371, 364, 364]),   # ISO 2026 has 53 weeks
        (date(2026, 5, 4), [245, 364, 364]),     # opened on a Monday inside it
    ])
    def test_the_iso_year_book_recalculates_to_the_engine(self, tmp_path, start, days):
        m = _iso_year_model(mo.DATE(start.year, start.month, start.day))
        assert m.header.days.value == days
        wb = _recalculated(m, tmp_path)
        rows = {r[0].value: r for r in wb['Time'].iter_rows() if r[0].value is not None}
        for name in ('start', 'end', 'days', 'index'):
            var = getattr(m.header, name)
            got = [_serial(rows[var.display_name][c].value) for c in range(1, 4)]
            assert got == pytest.approx([_serial(v) for v in var.value]), (name, got)
        fee = next(r for r in wb['Ops'].iter_rows() if r[0].value == 'Fee')
        assert [c.value for c in fee[2:5]] == pytest.approx(m.ops.fee.value)

    @pytest.mark.slow
    def test_a_weekly_book_projected_to_iso_years_reads_its_time(self, tmp_path):
        from modeleon.core.projection import project_model
        m = mo.Model('wk', default_grain='week', default_start='2026-W51', default_periods=5)
        with m:
            m.header = mo.MultiVariable('Time')
            with m.header as h:
                h.start = mo.time.start
                h.days = mo.time.days
            m.default_header = m.header
            m.ops = mo.MultiVariable('Ops')
            with m.ops as o:
                o.units = mo.Variable([1.0, 2.0, 3.0, 4.0, 5.0], regrain=mo.up('sum'), display_name='Units')
                o.fee = (mo.Variable(2.0, display_name='Rate') * mo.time.days).set_display_name('Fee')
        with warnings.catch_warnings(record=True) as seen:
            warnings.simplefilter('always')
            wb = _recalculated(project_model(m, 'iso_year'), tmp_path)
        assert not [w for w in seen if 'translate' in str(w.message)]
        ops = wb['Ops']
        assert [c.value for c in next(r for r in ops.iter_rows() if r[0].value == 'Units')[1:3]] == [6, 9]
        assert [c.value for c in next(r for r in ops.iter_rows() if r[0].value == 'Fee')[1:3]] == [42, 28]
