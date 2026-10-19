# SPDX-License-Identifier: Apache-2.0
"""A per-period line declared directly on an ``orient='columns'`` board
is laid out beneath the matrix, in the period columns — never dropped.

The board's rows are its instances' scalar fields; its own direct
scalars (a total across the instances) already followed beneath. A
direct SERIES — ``board.series = [...]`` — fitted neither and vanished
from the workbook without a word. It now follows beneath as an ordinary
row of the sheet: its cells in the period columns, which the masthead
names when the sheet has one, while a direct scalar keeps the sheet's
scalar column. A sheet that IS a board keeps its rule — its columns are
instances, no period header — and the series stands beneath all the
same.
"""

import shutil
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook

import modeleon as mo


def _grid(model, sheet, tmp_path, cols=5):
    path = tmp_path / "t.xlsx"
    model.to_excel(str(path))
    ws = load_workbook(str(path))[sheet]
    return [
        [ws.cell(r, c).value for c in range(1, cols + 1)]
        for r in range(1, ws.max_row + 1)
    ]


def _fill(board):
    board.p1 = mo.MultiVariable("P1")
    board.p1.cost = mo.Variable(50.0, display_name="Cost")
    board.p2 = mo.MultiVariable("P2")
    board.p2.cost = mo.Variable(80.0, display_name="Cost")
    board.series = mo.Variable([1.0, 2.0], display_name="Series")
    board.scaled = mo.Variable(board.series * board.p1.cost, display_name="Scaled")
    board.total = mo.Variable(board.p1.cost + board.p2.cost, display_name="Total")


def _model():
    m = mo.Model("m", display_name="M", default_grain="year", default_start="2026",
                 default_periods=2)
    with m:
        m.sheet = mo.MultiVariable("Sheet", excel_props={"tab": True})
        with m.sheet as s:
            s.board = mo.MultiVariable("Board", excel_props={"orient": "columns"})
            with s.board as b:
                _fill(b)
    return m


@pytest.fixture
def grid(tmp_path):
    return _grid(_model(), "Sheet", tmp_path)


def _labels(grid):
    return [row[0] for row in grid]


def _first_instance_column(grid):
    captions = next(row for row in grid if "P1" in row)
    first = captions.index("P1")
    assert captions[first:first + 2] == ["P1", "P2"]
    return first


def test_the_series_stands_beneath_the_matrix_in_the_period_columns(grid):
    labels = _labels(grid)
    assert "Series" in labels
    first = _first_instance_column(grid)
    assert grid[labels.index("Series")][first:first + 2] == [1, 2]


def test_the_masthead_names_the_series_columns(grid):
    # The board's series is the sheet's only per-period line: it alone
    # earns the sheet its period header, over the columns it stands in.
    first = _first_instance_column(grid)
    assert grid[0][first:first + 2] == ["2026", "2027"]


def test_a_formula_series_reads_the_board_cells(grid):
    labels = _labels(grid)
    first = _first_instance_column(grid)
    cost_row = labels.index("Cost") + 1
    series_row = labels.index("Series") + 1
    col = [chr(ord("A") + first), chr(ord("A") + first + 1)]
    assert grid[labels.index("Scaled")][first:first + 2] == [
        f"={col[0]}{series_row} * {col[0]}{cost_row}",
        f"={col[1]}{series_row} * {col[0]}{cost_row}",
    ]


def test_the_direct_scalar_keeps_the_scalar_column(grid):
    labels = _labels(grid)
    first = _first_instance_column(grid)
    cost_row = labels.index("Cost") + 1
    col = [chr(ord("A") + first), chr(ord("A") + first + 1)]
    total = grid[labels.index("Total")]
    assert f"={col[0]}{cost_row} + {col[1]}{cost_row}" in total
    # ...and not in the period columns: a total is no 2026 figure.
    assert total[first:first + 2] == [None, None]


def test_the_lines_follow_the_matrix_in_declaration_order(grid):
    labels = _labels(grid)
    i = labels.index("Cost")
    assert labels[i:i + 4] == ["Cost", "Series", "Scaled", "Total"]


def test_a_sheet_that_is_a_board_keeps_the_series_without_a_masthead(tmp_path):
    m = mo.Model("m", display_name="M", default_grain="year", default_start="2026",
                 default_periods=2)
    with m:
        m.board = mo.MultiVariable("Board", excel_props={"tab": True, "orient": "columns"})
        with m.board as b:
            _fill(b)
    g = _grid(m, "Board", tmp_path)
    # Its columns are instances: the captions stay in row 1, no period
    # header — and the series stands beneath in those same columns.
    assert g[0][:3] == [None, "P1", "P2"]
    assert _labels(g)[1:5] == ["Cost", "Series", "Scaled", "Total"]
    assert g[2][:3] == ["Series", 1, 2]
    assert g[3][:3] == ["Scaled", "=B3 * B2", "=C3 * B2"]
    assert g[4][:2] == ["Total", "=B2 + C2"]


def test_a_series_on_a_board_of_scalars_across_stands_beneath_them(tmp_path):
    m = mo.Model("m", display_name="M", default_grain="year", default_start="2026",
                 default_periods=2)
    with m:
        m.sheet = mo.MultiVariable("Sheet", excel_props={"tab": True})
        with m.sheet as s:
            s.kpis = mo.MultiVariable("KPIs", excel_props={"orient": "columns"})
            with s.kpis as k:
                k.irr = mo.Variable(0.2, display_name="IRR")
                k.moic = mo.Variable(2.5, display_name="MOIC")
                k.series = mo.Variable([1.0, 2.0], display_name="Series")
    g = _grid(m, "Sheet", tmp_path)
    captions = next(row for row in g if "IRR" in row)
    first = captions.index("IRR")
    assert captions[first:first + 2] == ["IRR", "MOIC"]
    assert g[g.index(captions) + 1][first:first + 2] == [0.2, 2.5]
    series = next(row for row in g if row[0] == "Series")
    assert series[first:first + 2] == [1, 2]
    assert g[0][first:first + 2] == ["2026", "2027"]


def test_a_header_row_line_rides_the_section_row_and_no_other(tmp_path):
    # A ``header_row`` line is hoisted onto the board's section row by
    # the parent's walk; the board must not lay it out a second time.
    m = mo.Model("m", display_name="M")
    with m:
        m.book = mo.MultiVariable("Book", excel_props={"tab": True})
        with m.book as bk:
            bk.board = mo.MultiVariable("Board", excel_props={"orient": "columns"})
            with bk.board as b:
                b.p1 = mo.MultiVariable("P1")
                b.p1.cost = mo.Variable(50.0, display_name="Cost")
                b.p2 = mo.MultiVariable("P2")
                b.p2.cost = mo.Variable(80.0, display_name="Cost")
                b.total = mo.Variable(
                    b.p1.cost + b.p2.cost, display_name="Total", excel_props={"header_row": True}
                )
    g = _grid(m, "Book", tmp_path)
    cells = [v for row in g for v in row]
    assert "Total" not in cells
    assert sum(1 for v in cells if isinstance(v, str) and v.startswith("=")) == 1


def _soffice():
    for candidate in ("soffice", "libreoffice"):
        found = shutil.which(candidate)
        if found:
            return found
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


@pytest.mark.slow
def test_the_board_series_recalculates_to_the_python_values(tmp_path):
    soffice = _soffice()
    if soffice is None:
        pytest.skip("LibreOffice not installed")
    m = _model()
    path = tmp_path / "board.xlsx"
    m.to_excel(path)
    subprocess.run(
        [soffice, f"-env:UserInstallation={(tmp_path / 'prof').as_uri()}",
         "--headless", "--convert-to", "xlsx", "--outdir",
         str(tmp_path / "calc"), str(path)],
        check=True, capture_output=True,
    )
    ws = load_workbook(tmp_path / "calc" / "board.xlsx", data_only=True)["Sheet"]
    rows = [[ws.cell(r, c).value for c in range(1, 5)] for r in range(1, ws.max_row + 1)]
    labels = [row[0] for row in rows]
    first = _first_instance_column(rows)
    assert rows[labels.index("Scaled")][first:first + 2] == [50.0, 100.0]
    assert m.sheet.board.scaled.value == [50.0, 100.0]
    assert rows[labels.index("Total")][first - 1] == 130.0 == m.sheet.board.total.value
