# SPDX-License-Identifier: Apache-2.0
"""A blank counts as zero, as in Excel.

A month nobody entered is a blank cell: a formula over it computes over 0,
a function reads it as Excel reads a blank cell, and only an input nobody
filled stays blank. Whether an actual was entered is the blend's question,
asked of the entries a formula is built from.
"""

from __future__ import annotations

from datetime import date

import pytest

import modeleon as mo


def _model(periods: int = 3, **kw) -> "mo.Model":
    return mo.Model('m', default_grain='month', default_start='2026-01',
                    default_periods=periods, **kw)


def _blended(periods: int = 3) -> "mo.Model":
    return _model(periods, tracks=mo.Tracks(
        'plan', 'forecast', 'actual',
        blend=mo.blend(given='actual', follow='forecast', name='live')))


class TestArithmetic:
    def test_a_blank_operand_reads_as_zero(self):
        a = mo.Variable([1.0, None, None])
        b = mo.Variable([2.0, 3.0, None])
        assert (a + b).value == [3.0, 3.0, 0]
        assert (a * b).value == [2.0, 0.0, 0]

    def test_dividing_by_a_blank_is_div0(self):
        assert (mo.Variable([4.0, 4.0]) / mo.Variable([2.0, None])).value == [2.0, '#DIV/0!']

    def test_minus_a_blank_is_zero(self):
        assert (-mo.Variable([None, 2.0, '#N/A'])).value == [0, -2.0, '#N/A']

    def test_minus_a_line_with_tracks_goes_track_by_track(self):
        m = _model(2, tracks=mo.Tracks('plan', 'actual'))
        with m:
            m.x = mo.Variable(plan=[1.0, 2.0], actual=[1.0, None])
            m.y = -m.x
        assert m.y.value['plan'] == [-1.0, -2.0]
        assert m.y.value['actual'] == [-1.0, 0]


class TestTime:
    def test_a_balance_carries_over_a_blank_month(self):
        bal = mo.recurrence(100.0, "{prev} + {x}", x=mo.Variable([None, 10.0, None]))
        assert bal.value == [100.0, 110.0, 110.0]

    def test_a_division_by_a_blank_month_is_div0_and_carries_on(self):
        m = _model(3)
        with m:
            m.x = mo.Variable([1.0, None, 2.0])
            m.d = mo.recurrence(100.0, "{prev} / {x}", x=m.x)
        assert m.d.value == [100.0, '#DIV/0!', '#DIV/0!']

    def test_a_running_total_with_tracks_runs_over_a_blank(self):
        m = _model(4, tracks=mo.Tracks('plan', 'actual'))
        with m:
            m.x = mo.Variable(plan=[1.0, 2.0, 3.0, 4.0], actual=[1.0, None, 3.0, None])
            m.c = mo.cumsum(m.x)
        assert m.c.value['actual'] == [1.0, 1.0, 4.0, 4.0]

    def test_a_lagged_blank_is_zero(self):
        assert mo.lag(mo.Variable([1.0, None, 3.0])).value == [0.0, 1.0, 0]

    def test_a_lagged_blank_date_stays_blank(self):
        # A line of dates has no zero.
        lagged = mo.lag(mo.Variable([date(2026, 1, 1), None, date(2026, 3, 1)]))
        assert lagged.value[2] is None


class TestFunctions:
    def test_round_abs_int_of_a_blank_are_zero(self):
        assert mo.ROUND(mo.Variable([None, 2.46]), 1).value == [0, 2.5]
        assert mo.ABS(mo.Variable([None, -2.0])).value == [0, 2.0]
        assert mo.INT(mo.Variable([None, 2.7])).value == [0, 2]

    def test_mod_by_a_blank_is_div0(self):
        assert mo.MOD(mo.Variable([5.0]), mo.Variable(None)).value == ['#DIV/0!']

    def test_if_landing_on_a_blank_shows_zero(self):
        assert mo.IF(mo.Variable([1, 0]), mo.Variable([None, 1.0]), 5).value == [0, 5]

    def test_choose_landing_on_a_blank_shows_zero(self):
        assert mo.CHOOSE(1, mo.Variable([None, 2.0])).value == [0, 2.0]

    def test_and_or_skip_a_blank(self):
        assert mo.AND(mo.Variable([True, None]), True).value == [True, True]
        assert mo.OR(mo.Variable([None]), mo.Variable([None])).value == ['#VALUE!']

    def test_text_of_a_blank_is_empty(self):
        assert mo.LEN(mo.Variable([None, 'ab'])).value == [0, 2]
        assert mo.UPPER(mo.Variable([None, 'ab'])).value == ['', 'AB']

    def test_npv_and_irr_skip_a_blank(self):
        # Excel's NPV and IRR skip a blank cell; the flows after it close up.
        npv = mo.NPV(0.1, mo.Variable([100.0, None, 100.0])).value
        assert npv == pytest.approx(100 / 1.1 + 100 / 1.21)
        assert mo.IRR(mo.Variable([-100.0, None, 60.0, 60.0])).value == pytest.approx(
            mo.IRR(mo.Variable([-100.0, 60.0, 60.0])).value)

    def test_pmt_reads_a_blank_argument_as_zero(self):
        assert mo.PMT(0.01, 12, mo.Variable(None)).value == 0


class TestTheBlendAsksTheEntries:
    def test_a_side_through_a_function_takes_the_forecast_where_nothing_was_entered(self):
        m = _blended()
        with m:
            m.a = mo.Variable([1.0] * 3, forecast=[2.0] * 3, actual=[1.5, None, None])
            m.r = mo.Variable(mo.ROUND(m.a * 10, 0), plan=[10.0] * 3)
            m.i = mo.Variable(mo.IF(m.a > 1, m.a, 0), plan=[1.0] * 3)
        assert m.r.value['live'] == [15.0, 20.0, 20.0]
        assert m.i.value['live'] == [1.5, 2.0, 2.0]

    def test_a_line_without_tracks_is_no_actuals_entry(self):
        # A seasonality filled in a month the actual is not: that month
        # is still un-entered.
        m = _blended(2)
        with m:
            m.a = mo.Variable([1.0] * 2, forecast=[2.0] * 2, actual=[1.0, None])
            m.season = mo.Variable([1.0, 1.0])
            m.x = mo.Variable(m.a * m.season, plan=[1.0] * 2)
        assert m.x.value['live'] == [1.0, 2.0]

    def test_neither_side_shows_zero_as_the_book_does(self):
        m = _blended(3)
        with m:
            m.a = mo.Variable(plan=[1.0] * 3, actual=[2.0, None, None])
        assert m.a.value['live'] == [2.0, 0, 0]

    def test_the_entered_signal_rides_the_line(self):
        m = _blended(3)
        with m:
            m.a = mo.Variable([1.0] * 3, forecast=[2.0] * 3, actual=[1.0, None, None])
            m.b = mo.Variable([1.0] * 3, forecast=[2.0] * 3, actual=[None, 3.0, None])
            m.t = mo.Variable(m.a + m.b, plan=[2.0] * 3)
        assert m.t._blend_entered == [True, True, False]
        assert m.t.value['live'] == [1.0, 3.0, 4.0]


class TestTheEntriesOfAnActualWrittenForItsTrack:
    """An actual written for its own track - through slices of other
    lines' actuals, or from loaded actuals in a line without tracks - is
    asked of what it reads, like a shared formula."""

    def test_an_actual_of_slices_takes_the_forecast_where_nothing_was_entered(self):
        m = _blended(3)
        with m:
            m.a = mo.Variable([10.0] * 3, forecast=[11.0] * 3, actual=[9.0, 8.0, None])
            m.b = mo.Variable([20.0] * 3, forecast=[21.0] * 3, actual=[19.0, None, None])
            m.t = mo.Variable([30.0] * 3, forecast=[32.0] * 3, actual=m.a.actual + m.b.actual)
        assert m.t.value['live'] == [28.0, 8.0, 32.0]

    def test_an_actual_from_loaded_actuals_takes_the_forecast_where_none_were_loaded(self):
        m = _blended(3)
        with m:
            m.raw = mo.Variable([5.0, None, None])
            m.x = mo.Variable([1.0] * 3, forecast=[2000.0] * 3, actual=m.raw * 1000)
            m.y = mo.Variable([1.0] * 3, forecast=[2.0] * 3, actual=m.raw)
        assert m.x.value['live'] == [5000.0, 2000.0, 2000.0]
        assert m.y.value['live'] == [5.0, 2.0, 2.0]
        # the bound actual is a reference: its blanks read as 0
        assert m.y.value['actual'] == [5.0, 0, 0]

    def test_a_reference_to_a_line_of_dates_keeps_its_blanks(self):
        m = _model(2)
        with m:
            m.d = mo.Variable([date(2026, 1, 1), None])
            m.w = mo.Variable(m.d)
        assert m.w.value == [date(2026, 1, 1), None]

    def test_a_slice_of_another_track_is_no_entry(self):
        m = _blended(2)
        with m:
            m.a = mo.Variable([1.0] * 2, forecast=[2.0] * 2, actual=[1.0, None])
            # the plan of another line is entered every month - no actual
            m.x = mo.Variable([1.0] * 2, forecast=[3.0] * 2,
                              actual=m.a.at(track='actual') + m.a.at(track='plan'))
        assert m.x.value['live'] == [2.0, 3.0]


class TestADerivedLineShowsWhereItWasEntered:
    def test_a_derived_actual_carries_the_entered_mask(self):
        m = _blended(3)
        with m:
            m.rev = mo.Variable([5.0] * 3, forecast=[6.0] * 3, actual=[4.0, None, None])
            m.cost = mo.Variable([1.0] * 3, forecast=[2.0] * 3, actual=[1.0, None, None])
            m.margin = m.rev - m.cost
        assert m.margin._blend_entered == [True, False, False]
        # its value is untouched: live from the operands' live
        assert m.margin.value['live'] == [3.0, 4.0, 4.0]

    def test_a_line_paid_a_month_late_shows_through_its_lag(self):
        m = _blended(3)
        with m:
            m.accrued = mo.Variable([5.0] * 3, forecast=[6.0] * 3, actual=[4.0, 4.0, None])
            m.paid = mo.lag(m.accrued)
        assert m.paid._blend_entered == [True, True, True]
        m2 = _blended(4)
        with m2:
            m2.accrued = mo.Variable([5.0] * 4, forecast=[6.0] * 4, actual=[4.0, None, None, None])
            m2.paid = mo.lag(m2.accrued)
        # the lag's first month reads its fill; then the accrual a month back
        assert m2.paid._blend_entered == [True, True, False, False]

    def test_this_months_entries_decide_over_last_months(self):
        m = _blended(3)
        with m:
            m.rev = mo.Variable([5.0] * 3, forecast=[6.0] * 3, actual=[4.0, 4.0, None])
            m.tax = mo.lag(m.rev) * 0.1
            m.net = m.rev - m.tax
        assert m.net._blend_entered == [True, True, False]


class TestMoreFunctions:
    def test_mod_over_a_series_divisor(self):
        assert mo.MOD(mo.Variable([5.0, 7.0, 9.0]), mo.Variable([2.0, 3.0, 4.0])).value == [1.0, 1.0, 1.0]
        assert mo.MOD(mo.Variable([5.0, None]), mo.Variable([0.0, 3.0])).value == ['#DIV/0!', 0.0]
        assert mo.MOD(-7, 3).value == 2

    def test_sum_of_single_dates_per_track_stands(self):
        m = _model(1, tracks=mo.Tracks('plan', 'actual'))
        with m:
            m.d = mo.Variable(plan=date(2026, 1, 1), actual=date(2026, 2, 1))
            m.s = mo.SUM(m.d)
        assert m.s.value['actual'] == date(2026, 2, 1)


class TestTheBookAsksTheSameCells:
    def test_an_actual_bound_to_a_plain_line_asks_that_lines_cell(self, tmp_path):
        # The actual row is ``=<raw>``, which Excel computes as 0 over a
        # blank - only the input cell itself can say nothing was entered.
        import openpyxl

        m = _blended(3)
        with m:
            m.s = mo.MultiVariable('S', excel_props={'tab': True})
            with m.s as s:
                s.raw = mo.Variable([5.0, None, None], display_name='Raw')
                s.x = mo.Variable([1.0] * 3, forecast=[2.0] * 3, actual=s.raw,
                                  display_name='X')
        path = tmp_path / 'b.xlsx'
        m.to_excel(str(path))
        ws = openpyxl.load_workbook(str(path))['S']
        rows = {ws.cell(r, 1).value: r for r in range(1, ws.max_row + 1)}
        raw, x = rows['Raw'], rows['X']
        assert ws.cell(x, 3).value == (
            f'=IF(C{raw}="",C{rows["X · forecast"]},C{rows["X · actual"]})')
