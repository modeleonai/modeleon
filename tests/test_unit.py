# SPDX-License-Identifier: Apache-2.0
"""Unit propagation through arithmetic and Excel emission."""

import warnings

import pytest
from openpyxl import load_workbook

import modeleon as mo


class TestBasicConstruction:
    def test_unit_stored_on_variable(self):
        v = mo.Variable(10, unit="$")
        assert str(v._unit) == "$"

    def test_compound_unit_parsed(self):
        v = mo.Variable(10, unit="$/hr")
        assert str(v._unit) == "$/hr"

    def test_no_unit_is_none(self):
        v = mo.Variable(10)
        assert v._unit is None


class TestMultiplication:
    def test_same_unit_composes(self):
        a = mo.Variable(5, unit="hr")
        b = mo.Variable(2, unit="hr")
        assert str((a * b)._unit) == "hr^2"

    def test_different_units_compose(self):
        price = mo.Variable(50, unit="$")
        hours = mo.Variable(10, unit="hr")
        assert str((price * hours)._unit) == "$ * hr"

    def test_scalar_preserves_unit(self):
        h = mo.Variable(10, unit="hr")
        assert str((h * 2)._unit) == "hr"
        assert str((2 * h)._unit) == "hr"


class TestDivision:
    def test_division_simplifies(self):
        revenue = mo.Variable(1000, unit="$")
        hours = mo.Variable(10, unit="hr")
        assert str((revenue / hours)._unit) == "$/hr"

    def test_cancellation_yields_dimensionless(self):
        a = mo.Variable(10, unit="$")
        b = mo.Variable(2, unit="$")
        assert (a / b)._unit is None

    def test_compound_division_cancels(self):
        cost = mo.Variable(100, unit="$")
        rate = mo.Variable(5, unit="$/hr")
        # $ / ($/hr) = hr
        assert str((cost / rate)._unit) == "hr"


class TestParsing:
    """A unit string reads the way it is written: everything after the
    first ``/`` is the denominator, and each further ``/`` divides again."""

    def test_a_chain_of_slashes_divides_again(self):
        assert mo.Unit("USD/kW/month") == mo.Unit("USD/(kW * month)")

    def test_a_chain_prints_as_written(self):
        assert str(mo.Unit("USD/kW/month")) == "USD/kW/month"
        assert str(mo.Unit(" USD/kW/month ")) == "USD/kW/month"

    def test_a_chain_cancels_through_arithmetic(self):
        capacity = mo.Variable(20_000.0, unit="kW")
        rent = mo.Variable(130.0, unit="USD/kW/month")
        assert str((capacity * rent).unit) == "USD/month"

    def test_a_chain_keeps_its_spelling_in_the_workbook(self, tmp_path):
        lease = mo.MultiVariable("Lease")
        lease.capacity = mo.Variable(20_000.0, unit="kW", display_name="Capacity")
        lease.rent = mo.Variable(130.0, unit="USD/kW/month", display_name="Rent")
        lease.monthly = (lease.capacity * lease.rent).set_display_name("Monthly rent")
        path = tmp_path / "lease.xlsx"
        lease.to_excel(path)
        labels = [row[0] for row in load_workbook(path).worksheets[0].iter_rows(values_only=True)]
        assert "Rent (USD/kW/month)" in labels
        assert "Monthly rent (USD/month)" in labels

    def test_a_parenthesised_quotient_in_the_denominator_inverts(self):
        assert mo.Unit("USD/(kW/month)") == mo.Unit("USD * month/kW")

    def test_a_parenthesised_quotient_in_the_numerator(self):
        assert mo.Unit("(USD/kW)/month") == mo.Unit("USD/kW/month")

    @pytest.mark.parametrize("text, rebuilt", [
        ("$/hr", "$/hr"),
        ("$/hr * Gb", "$/(hr * Gb)"),
        ("$/(hr * Gb)", "$/(hr * Gb)"),
        ("$ * hr / Gb", "$ * hr/Gb"),
        ("1/hr", "1/hr"),
        ("USD m", "USD m"),
        ("EUR/kWh", "EUR/kWh"),
    ])
    def test_one_slash_reads_as_before(self, text, rebuilt):
        assert str(mo.Unit(text)) == rebuilt

    @pytest.mark.parametrize("text", [
        "USD/kW/month", "USD/(kW/month)", "(USD/kW)/month", "$/(hr * Gb)", "1/hr",
    ])
    def test_the_printed_unit_reads_back_as_itself(self, text):
        unit = mo.Unit(text)
        assert mo.Unit(str(unit)) == unit

    def test_a_unit_built_from_a_chain_reads_back_as_itself(self):
        unit = mo.Unit("USD/kW/month") * mo.Unit("kW")
        assert str(unit) == "USD/month"
        assert mo.Unit(str(unit)) == unit

    def test_a_word_with_a_note_cancels_in_a_chain(self):
        capacity = mo.Variable(20_000.0, unit="kW (AC)")
        rent = mo.Variable(130.0, unit="USD/kW (AC)/month")
        assert str((capacity * rent).unit) == "USD/month"
        real = mo.Variable(20_000.0, unit="kW") * mo.Variable(130.0, unit="USD (real)/kW/month")
        total = real + mo.Variable(1.0, unit="USD (real)/month")
        assert str(total.unit) == "USD (real)/month"
        assert str(mo.Unit("kW (AC)")) == "kW (AC)"

    def test_parentheses_inside_a_group_change_nothing(self):
        assert mo.Unit("USD/(kW * (month))/yr") == mo.Unit("USD/(kW * month)/yr")
        assert mo.Unit("USD/(kW (AC) * month)/yr") == mo.Unit("USD/(month * kW (AC))/yr")

    def test_the_convention_after_a_later_slash(self):
        assert mo.Unit("USD/kW * month/year") == mo.Unit("USD/(kW * month * year)")
        assert mo.Unit("USD/(kW*(month/year))") == mo.Unit("USD * year/(kW * month)")

    def test_spaces_around_the_slashes(self):
        assert mo.Unit("USD / kW / month") == mo.Unit("USD/kW/month")
        assert str(mo.Unit("USD / kW / month")) == "USD / kW / month"

    def test_parentheses_around_the_whole_unit_do_not_print(self):
        assert str(mo.Unit("(USD/kWh)")) == "USD/kWh"
        assert str(mo.Unit("($/hr)")) == "$/hr"

    def test_unbalanced_parentheses_read_back_as_themselves(self):
        unit = mo.Unit("USD/(kW/month")
        assert mo.Unit(str(unit)) == unit

    def test_a_unit_restored_without_its_written_text_prints(self):
        restored = mo.Unit.__new__(mo.Unit)
        restored.__dict__.update({"_unit_components": {"USD m": 1}, "_display_symbol": None})
        assert str(restored) == "USD m"

    def test_a_chain_adds_to_its_parenthesised_spelling(self):
        a = mo.Variable(1.0, unit="USD/kW/month")
        b = mo.Variable(2.0, unit="USD/(kW * month)")
        assert (a + b).value == 3.0
        assert str((a + b).unit) == "USD/kW/month"

    def test_an_amount_per_month_still_names_the_month(self):
        monthly = mo.Variable(20_000.0, unit="kW") * mo.Variable(130.0, unit="USD/kW/month")
        with pytest.raises(ValueError, match="per 'month'"):
            monthly + mo.Variable(1.0, unit="USD")


class TestUnitReuse:
    """When ``*`` / ``/`` produce a unit one operand already carries, that
    operand's own Unit is the result (as ``+`` / ``-`` do) — equal to what
    the algebra would build, without building it."""

    def test_times_unitless_keeps_the_same_unit(self):
        a = mo.Variable(5, unit="$")
        r = a * mo.Variable(2)
        assert r._unit is a._unit and str(r._unit) == "$"
        assert (mo.Variable(2) * a)._unit is a._unit

    def test_times_dimensionless_keeps_the_same_unit(self):
        a = mo.Variable(5, unit="$")
        pct = mo.Variable(0.2, unit="%")
        assert (a * pct)._unit is a._unit
        assert (pct * a)._unit is a._unit

    def test_divided_by_unitless_keeps_the_same_unit(self):
        a = mo.Variable(12, unit="$/hr")
        assert (a / mo.Variable(4))._unit is a._unit
        assert (a / mo.Variable(4, unit="%"))._unit is a._unit

    def test_unitless_over_a_unit_still_inverts_it(self):
        r = mo.Variable(1) / mo.Variable(4, unit="hr")
        assert str(r._unit) == "1/hr"

    def test_dimensionless_results_stay_none(self):
        pct = mo.Variable(0.2, unit="%")
        assert (pct * pct)._unit is None
        assert (pct * mo.Variable(3))._unit is None
        assert (pct / mo.Variable(3))._unit is None

    def test_no_new_unit_is_built_for_a_reused_result(self, monkeypatch):
        from modeleon.core.unit import Unit

        a = mo.Variable(5, unit="$")
        built = []
        original = Unit.__init__

        def counting(self, *args, **kwargs):
            built.append(1)
            original(self, *args, **kwargs)

        monkeypatch.setattr(Unit, "__init__", counting)
        _ = a * mo.Variable(2) * mo.Variable(0.2, unit="%") / mo.Variable(12)
        # The only Unit built here is the one the "%" literal parses.
        assert len(built) == 1


class TestCountedLabels:
    """bps and pp label a pure number without rescaling it: 50 bps is
    written 0.005. A line holding the count (50) added to a % line holding
    fractions is refused - that sum is wrong under any reading. Anything
    the label cannot tell apart keeps the plain arithmetic."""

    def test_a_count_of_bps_added_to_a_rate_is_refused(self):
        base = mo.Variable(0.05, unit="%")
        spread = mo.Variable(50, unit="bps")
        with pytest.raises(ValueError, match=r"50 bps is 0\.005"):
            base + spread
        with pytest.raises(ValueError, match="never rescales"):
            spread - base

    def test_a_count_of_pp_added_to_a_margin_is_refused(self):
        margin = mo.Variable(0.30, unit="%")
        change = mo.Variable(2, unit="pp")
        with pytest.raises(ValueError, match=r"2 pp is 0\.02"):
            margin + change

    def test_a_count_in_one_period_is_enough(self):
        base = mo.Variable([0.05, 0.05], unit="%")
        spread = mo.Variable([0.005, 50], unit="bps")
        with pytest.raises(ValueError):
            base + spread

    def test_the_fraction_adds(self):
        base = mo.Variable(0.05, unit="%")
        spread = mo.Variable(0.005, unit="bps")
        assert (base + spread).value == pytest.approx(0.055)

    def test_an_explicit_conversion_adds(self):
        base = mo.Variable(0.05, unit="%")
        spread = mo.Variable(50, unit="bps")
        assert (base + spread / 10_000).value == pytest.approx(0.055)

    def test_the_excel_habit_still_computes(self):
        principal = mo.Variable(1_000_000.0, unit="USD")
        spread = mo.Variable(50, unit="bps")
        assert (principal * spread / 10_000).value == pytest.approx(5_000.0)

    def test_points_written_as_points_add(self):
        margin = mo.Variable(30, unit="%")
        change = mo.Variable(2, unit="pp")
        assert (margin + change).value == 32

    def test_counts_of_the_same_label_add(self):
        a = mo.Variable(50, unit="bps")
        b = mo.Variable(25, unit="bps")
        assert (a + b).value == 75
        assert str((a + b).unit) == "bps"


class TestCountedLabelsInALoop:
    """The check ``+`` makes of a count added to a share waits while a
    loop is open; it is made on the settled lines when the loop closes,
    under the row that adds."""

    @staticmethod
    def _model():
        m = mo.Model("M", default_grain="month", default_start="2026-01", default_periods=4)
        with m:
            m.inp = mo.MultiVariable("Inputs")
            with m.inp as i:
                i.base0 = mo.Variable(0.05, unit="%")         # an opening rate, a fraction
                i.base = mo.Variable([0.05] * 4, unit="%")
                i.drift = mo.Variable([0.0] * 4, unit="%")
                i.step = mo.Variable(50, unit="bps")          # a COUNT: 50 bps is 0.005
                i.step0 = mo.Variable(50, unit="bps")
                i.pp = mo.Variable(2, unit="pp")              # a COUNT of points
                i.frac_step = mo.Variable(0.005, unit="bps")  # 50 bps written right
                i.pts0 = mo.Variable(30.0, unit="%")          # a % written as points
                i.zero0 = mo.Variable(0.0, unit="%")
        return m

    def test_a_count_stepping_a_rate_inside_a_loop_is_refused_when_it_closes(self):
        m = self._model()
        with pytest.raises(ValueError, match=r"x\.rate: 'step' is in bps and holds a count"
                                             r".*50 bps is 0\.005"):
            with m:
                m.x = mo.MultiVariable("X", loop=True)
                with m.x as d:
                    d.rate = mo.lag(d.rate, fill=m.inp.base0) + m.inp.step

    def test_a_pp_count_on_a_fraction_inside_a_loop_is_refused(self):
        m = self._model()
        with pytest.raises(ValueError, match=r"x\.rate: 'pp' is in pp and holds a count"
                                             r".*2 pp is 0\.02"):
            with m:
                m.x = mo.MultiVariable("X", loop=True)
                with m.x as d:
                    d.rate = mo.lag(d.rate, fill=m.inp.base0) + m.inp.pp

    def test_a_row_read_before_its_line_is_checked_when_the_loop_closes(self):
        m = self._model()
        with pytest.raises(ValueError, match=r"x\.all_in: 'step' is in bps and holds a count"):
            with m:
                m.x = mo.MultiVariable("X", loop=True)
                with m.x as d:
                    d.all_in = d.rate + m.inp.step
                    d.rate = mo.lag(d.rate, fill=m.inp.base0) + m.inp.drift

    def test_a_count_row_of_the_loop_beside_a_share_is_refused(self):
        m = self._model()
        with pytest.raises(ValueError, match=r"x\.all_in: .*is in bps and holds a count"):
            with m:
                m.x = mo.MultiVariable("X", loop=True)
                with m.x as d:
                    d.all_in = m.inp.base + d.spread_row
                    d.spread_row = mo.lag(d.spread_row, fill=m.inp.step0) + m.inp.step * 0

    def test_a_sum_inside_a_loop_rows_formula_is_checked(self):
        m = self._model()
        with pytest.raises(ValueError, match=r"'step' is in bps and holds a count"):
            with m:
                m.x = mo.MultiVariable("X", loop=True)
                with m.x as d:
                    d.rate = (mo.lag(d.rate, fill=m.inp.base0) + m.inp.step) + m.inp.drift

    def test_the_fraction_steps_inside_a_loop(self):
        m = self._model()
        with m:
            m.x = mo.MultiVariable("X", loop=True)
            with m.x as d:
                d.rate = mo.lag(d.rate, fill=m.inp.base0) + m.inp.frac_step
        assert m.x.rate.value == pytest.approx([0.055, 0.06, 0.065, 0.07])
        assert str(m.x.rate.unit) == "%"

    def test_points_step_by_points_inside_a_loop(self):
        m = self._model()
        with m:
            m.x = mo.MultiVariable("X", loop=True)
            with m.x as d:
                d.rate = mo.lag(d.rate, fill=m.inp.pts0) + m.inp.pp
        assert m.x.rate.value == pytest.approx([32.0, 34.0, 36.0, 38.0])

    def test_a_zero_seed_inside_a_loop_reads_as_a_fraction(self):
        # As outside a loop: ``0 % + 2 pp`` is refused, a zero tells nothing
        # of points, and a count on a fraction is the error the check is for.
        m = self._model()
        with pytest.raises(ValueError, match=r"'pp' is in pp and holds a count"):
            with m:
                m.x = mo.MultiVariable("X", loop=True)
                with m.x as d:
                    d.rate = mo.lag(d.rate, fill=m.inp.zero0) + m.inp.pp

    def test_a_pick_inside_a_loop_is_checked_when_it_closes(self):
        # MAX(lag(rate, fill=0.05 %), 50 bps) settles the row to counts, so
        # the fraction and the counts meet in the lag: its source holds the
        # counts (judged by its first period), its fill the fraction.
        m = self._model()
        with pytest.raises(ValueError, match=r"'rate' is in bps and holds a count, beside "
                                             r"'base0' \(%\) in mo\.lag"):
            with m:
                m.x = mo.MultiVariable("X", loop=True)
                with m.x as d:
                    d.rate = mo.MAX(mo.lag(d.rate, fill=m.inp.base0), m.inp.step)

    def test_a_recurrence_template_inside_a_loop_is_checked_when_it_closes(self):
        # The row read before its line is settled when the loop closes;
        # the template's + is read then, {prev} as the start.
        m = self._model()
        with pytest.raises(ValueError, match=r"x\.cum: recurrence\(.*\): 'steps' is in bps and "
                                             r"holds a count, added to 'base0' \(%\)"):
            with m:
                m.x = mo.MultiVariable("X", loop=True)
                with m.x as d:
                    d.cum = mo.recurrence(m.inp.base0, "{prev} + {s}", s=d.steps)
                    d.steps = m.inp.step + mo.Variable([0.0, 0.0, 0.0, 0.0])


class TestCountedLabelsThroughFunctions:
    """The check ``+`` makes of a count added to a share is made wherever
    a function adds, compares or picks the lines - a total, a pick, a
    lag's fill, a conditional total - and over a recurrence template."""

    @staticmethod
    def _lines():
        return dict(
            share=mo.Variable([0.05, 0.06, 0.07], unit="%", display_name="Share"),
            share0=mo.Variable(0.05, unit="%", display_name="Opening share"),
            step=mo.Variable(50, unit="bps", display_name="Step"),
            steps=mo.Variable([50, 50, 50], unit="bps", display_name="Steps"),
            frac=mo.Variable(0.005, unit="bps", display_name="Frac"),
            pts=mo.Variable([30.0, 32.0, 34.0], unit="%", display_name="Points"),
            pp=mo.Variable(2, unit="pp", display_name="Pp"),
        )

    @pytest.mark.parametrize("where, build", [
        ("mo.SUM", lambda v: mo.SUM(v["share"], v["step"])),
        ("mo.MAX", lambda v: mo.MAX(v["step"], v["share"])),
        ("mo.MIN", lambda v: mo.MIN(v["share"], v["steps"])),
        ("mo.IF", lambda v: mo.IF(v["share"] > 0.06, v["share"], v["step"])),
        ("mo.CHOOSE", lambda v: mo.CHOOSE(1, v["share"], v["step"])),
        ("mo.CEILING", lambda v: mo.CEILING(v["step"], v["share"])),
        ("mo.IFS", lambda v: mo.IFS(v["share"] > 0.06, v["step"], True, v["share"])),
        ("mo.lag", lambda v: mo.lag(v["share"], fill=v["step"])),
        (".shift()", lambda v: v["share"].shift(1, fill_value=v["step"])),
        ("mo.SUMIF", lambda v: mo.SUMIF([mo.Variable(1.0), mo.Variable(2.0)], ">0",
                                         [v["share0"], v["step"]])),
    ])
    def test_a_count_beside_a_share_is_refused_where_a_function_combines_them(
            self, where, build):
        v = self._lines()
        with pytest.raises(ValueError, match=rf"'Steps?' is in bps and holds a count, beside "
                                             rf"'(Share|Opening share)' \(%\) in {where}"
                                             rf"[.(),]*, which holds a fraction"):
            build(v)

    def test_a_fraction_step_or_points_pass(self):
        v = self._lines()
        assert mo.SUM(v["share"], v["frac"]).unit is None       # not refused: two labels name neither
        assert mo.MAX(v["pts"], v["pp"]).value == [30.0, 32.0, 34.0]
        assert mo.lag(v["share"], fill=v["frac"]).value == pytest.approx([0.005, 0.05, 0.06])
        assert mo.IF(v["pts"] > 31, v["pts"], v["pp"]).value == [2.0, 32.0, 34.0]

    def test_a_function_over_many_lines_reads_each_line_once(self, monkeypatch):
        # A line is a count, a fraction or neither on its own: the check
        # reads each line once and names the first count and the first
        # fraction - never a pass over every two lines.
        from modeleon.core import variable_ops
        pairs = []
        check = variable_ops._refuse_a_count_added_to_a_share
        monkeypatch.setattr(variable_ops, "_refuse_a_count_added_to_a_share",
                            lambda *a, **k: pairs.append(a[:2]) or check(*a, **k))
        amounts = [mo.Variable(float(i), unit="USD") for i in range(300)]
        assert mo.SUM(*amounts).value == pytest.approx(sum(range(300)))
        assert pairs == []
        v = self._lines()
        more = mo.Variable([60, 60, 60], unit="bps", display_name="More steps")
        other = mo.Variable([0.01, 0.01, 0.01], unit="%", display_name="Other share")
        with pytest.raises(ValueError, match=r"'Steps' is in bps and holds a count, "
                                             r"beside 'Share' \(%\) in mo\.SUM"):
            mo.SUM(v["pts"], v["share"], v["steps"], more, other)
        assert len(pairs) == 1

    def test_a_recurrence_template_adding_a_count_to_a_share_is_refused(self):
        v = self._lines()
        with pytest.raises(ValueError, match=r'recurrence\("\{prev\} \+ \{x\}"\): \'Steps\' is in '
                                             r'bps and holds a count, added to \'Opening share\''):
            mo.recurrence(v["share0"], "{prev} + {x}", x=v["steps"])
        with pytest.raises(ValueError, match="holds a count, added to"):
            mo.recurrence_sum(v["share0"], v["steps"])
        pts0 = mo.Variable(30.0, unit="%")
        ramp = mo.recurrence(pts0, "{prev} + {x}", x=mo.Variable([2, 2, 2], unit="pp"))
        assert str(ramp.unit) == "%" and ramp.value[0] == pytest.approx(30.0)


class TestUnitsThroughFunctions:
    """A function that adds, compares or picks its values refuses two
    different units, as + does; a comparison of two lines does too. A
    number or an unlabelled line goes with any; a condition, an index or
    a compared case is not a value. A recurrence carries its start's unit."""

    def setup_method(self):
        self.usd = mo.Variable(1000.0, unit="USD", display_name="Revenue")
        self.usd2 = mo.Variable(400.0, unit="USD")
        self.hr = mo.Variable(160.0, unit="hr", display_name="Hours")

    @pytest.mark.parametrize("fn", [mo.SUM, mo.AVERAGE, mo.MIN, mo.MAX])
    def test_two_units_are_refused(self, fn):
        with pytest.raises(ValueError, match="'Revenue' and 'Hours' are in USD and hr"):
            fn(self.usd, self.hr)

    @pytest.mark.parametrize("fn", [mo.SUM, mo.AVERAGE, mo.MIN, mo.MAX])
    def test_one_unit_is_kept(self, fn):
        assert str(fn(self.usd, self.usd2).unit) == "USD"

    def test_a_number_goes_with_any_unit(self):
        assert str(mo.SUM(self.usd, 5).unit) == "USD"
        assert str(mo.MAX(self.usd, 0).unit) == "USD"

    def test_picked_values_in_two_units_are_refused(self):
        with pytest.raises(ValueError, match="mo.IF takes its values in one unit"):
            mo.IF(self.usd > 0, self.usd, self.hr)
        with pytest.raises(ValueError, match="mo.CHOOSE takes its values in one unit"):
            mo.CHOOSE(1, self.usd, self.hr)
        with pytest.raises(ValueError, match="mo.IFERROR takes its values in one unit"):
            mo.IFERROR(self.usd, self.hr)
        assert str(mo.IF(self.usd > 0, self.usd, 0).unit) == "USD"
        assert str(mo.IF(self.hr > 100, self.usd, self.usd2).unit) == "USD"

    def test_a_comparison_of_two_units_is_refused(self):
        with pytest.raises(ValueError, match="Cannot compare values with incompatible units: USD and hr"):
            self.usd > self.hr
        with pytest.raises(ValueError, match="Cannot compare"):
            mo.Variable(12.0, unit="month") == mo.Variable(12.0, unit="months")
        assert (self.usd > self.usd2).value is True
        assert (self.usd > 0).value is True

    def test_mixed_by_nature_functions_still_run(self):
        rate = mo.Variable(0.1, unit="%")
        flows = mo.Variable([100.0, 100.0], unit="USD")
        assert mo.NPV(rate, flows).value == pytest.approx(173.553719, rel=1e-6)

    def test_a_recurrence_carries_its_start_unit(self):
        opening = mo.Variable(100.0, unit="USD")
        balance = mo.recurrence(opening, "{prev} * 1.01", periods=3)
        assert str(balance.unit) == "USD"
        topped = mo.recurrence_sum(opening, mo.Variable([10.0, 10.0, 10.0], unit="USD"))
        assert str(topped.unit) == "USD"
        assert str(mo.recurrence(opening, "{prev} * 2", periods=2, unit="EUR").unit) == "EUR"
        assert mo.recurrence(100.0, "{prev} * 2", periods=2).unit is None

    def test_a_recurrence_template_refuses_two_units(self):
        # The template is read for units as any formula is: {prev} in the
        # start's unit beside a variable in another is refused, as + refuses
        # them, and a declared unit excuses no mix.
        eur = mo.Variable(100.0, unit="EUR", display_name="Opening")
        usd = mo.Variable([1.0, 2.0, 3.0], unit="USD", display_name="Fee")
        cap = mo.Variable(150.0, unit="USD", display_name="Cap")
        with pytest.raises(ValueError, match=r'recurrence\("\{prev\} \+ \{x\}"\): Cannot '
                                             r'add/subtract values with incompatible units: '
                                             r'EUR and USD'):
            mo.recurrence(eur, "{prev} + {x}", x=usd)
        with pytest.raises(ValueError, match="incompatible units: EUR and USD"):
            mo.recurrence_sum(eur, usd)
        with pytest.raises(ValueError, match="incompatible units: EUR and USD"):
            mo.recurrence(eur, "{prev} + {x}", x=usd, unit="EUR")
        with pytest.raises(ValueError, match="Cannot compare values with incompatible units: "
                                             "EUR and USD"):
            mo.recurrence(eur, "IF({prev} > {cap}, {cap}, {prev} + {x})", cap=cap, x=usd)
        a = mo.Variable([1.0] * 3, unit="EUR")
        b = mo.Variable([1.0] * 3, unit="USD")
        with pytest.raises(ValueError, match="incompatible units"):
            mo.recurrence(0.0, "{prev} + {a} - {b}", a=a, b=b)

    def test_a_recurrence_template_derives_the_lines_unit(self):
        usd = mo.Variable([1.0, 2.0, 3.0], unit="USD")
        eur = mo.Variable(100.0, unit="EUR")
        g = mo.Variable(0.1, unit="%")
        assert str(mo.recurrence(0.0, "{prev} + {x}", x=usd).unit) == "USD"
        assert str(mo.recurrence(eur, "{prev} * (1 + {g})", g=g, periods=3).unit) == "EUR"
        assert mo.recurrence(100.0, "{prev} * 2", periods=2).unit is None
        share = mo.Variable(0.5, unit="%")
        assert str(mo.recurrence(share, "{prev} * (1 + {g})", g=g, periods=3).unit) == "%"
        same = mo.Variable([1.0] * 3, unit="EUR")
        assert str(mo.recurrence(eur, "ABS({prev} - {x})", x=same).unit) == "EUR"

    def test_a_recurrence_template_that_does_not_settle_is_refused(self):
        eur = mo.Variable(100.0, unit="EUR")
        usd = mo.Variable([1.0, 2.0, 3.0], unit="USD")
        with pytest.raises(ValueError, match="grows a power of it, so the line's unit does "
                                             "not settle"):
            mo.recurrence(eur, "{prev} * {x}", x=usd)
        with pytest.raises(ValueError, match="the template gives USD, but the start is in EUR"):
            mo.recurrence(eur, "{x} * 2", x=usd)
        with pytest.raises(ValueError, match="does not settle"):
            mo.recurrence(0.0, "{prev} * {x}", x=usd)
        # A declared unit decides what the line prints: the product is its own.
        assert str(mo.recurrence(eur, "{prev} * {x}", x=usd, unit="EUR").unit) == "EUR"

    @pytest.mark.parametrize("build", ["sum", "max", "compare"])
    def test_a_loop_refuses_two_units_when_it_closes(self, build):
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=3)
        with pytest.raises(ValueError, match="takes its values in one unit|Cannot compare"):
            with m:
                m.inp = mo.MultiVariable("Inputs")
                with m.inp as i:
                    i.opening = mo.Variable(1000.0, unit="USD")
                    i.flow = mo.Variable([10.0] * 3, unit="USD")
                    i.hours = mo.Variable([160.0] * 3, unit="hr")
                m.debt = mo.MultiVariable("Debt", loop=True)
                with m.debt as d:
                    prev = mo.lag(d.balance, fill=m.inp.opening)
                    if build == "sum":
                        d.check = mo.SUM(prev, m.inp.hours)
                    elif build == "max":
                        d.check = mo.MAX(prev, m.inp.hours)
                    else:
                        d.check = mo.IF(prev > m.inp.hours, 1.0, 0.0)
                    d.balance = prev + m.inp.flow


class TestFunctionUnitLaws:
    """Each function's unit follows its own law — the one its result is
    built by and a loop container derives its rows' units again by: a
    power raises the unit, a root halves it, a product multiplies its
    operands' units, a remainder and a multiple share one unit; a
    conditional total is in the unit its totalled lines share, and
    compares a criterion in its range's unit."""

    def setup_method(self):
        self.usd = mo.Variable([3.0, 4.0], unit="USD", display_name="Amount")
        self.hr = mo.Variable(2.0, unit="hr", display_name="Hours")

    def test_a_power_raises_the_unit_and_a_root_halves_it(self):
        square = mo.POWER(self.usd, 2)
        assert str(square.unit) == "USD^2" and str(mo.POWER(self.usd, 2.0).unit) == "USD^2"
        assert str(mo.SQRT(square).unit) == "USD" and str((square ** 0.5).unit) == "USD"
        assert str(mo.POWER(self.usd, -1).unit) == "1/USD"
        assert str(mo.SQRT(mo.Variable(9.0, unit="USD^2")).unit) == "USD"
        # A fractional power of a dimension, or a power that is a line,
        # has no unit to write.
        assert mo.SQRT(self.usd).unit is None
        assert mo.POWER(self.usd, mo.Variable(2)).unit is None
        with pytest.raises(ValueError, match="USD\\^2 and USD"):
            _ = square + self.usd

    def test_a_product_multiplies_the_units(self):
        price = mo.Variable(5.0, unit="USD/hr")
        assert str(mo.PRODUCT(price, self.hr).unit) == "USD"
        assert str(mo.PRODUCT(price, self.hr, mo.Variable(0.9, unit="%")).unit) == "USD"
        with pytest.raises(ValueError, match="multiplies the cells of 'Amount', a row in USD"):
            mo.PRODUCT(self.usd)
        # A word on a row of growth factors is a dimension: the refusal says
        # how to make it the pure number it labels.
        with pytest.raises(ValueError, match=r"FACTOR = mo.Unit.dimensionless\('factor'\)"):
            mo.PRODUCT(mo.Variable([1.1, 1.2], unit="factor"))
        factor = mo.Unit.dimensionless("factor")
        assert mo.PRODUCT(mo.Variable([1.1, 1.2], unit=factor)).value == pytest.approx(1.32)

    def test_a_remainder_and_a_multiple_take_one_unit_a_rounding_keeps_it(self):
        for fn in (mo.MOD, mo.CEILING, mo.FLOOR):
            with pytest.raises(ValueError, match=f"mo.{fn.__name__} takes its values in one unit"):
                fn(self.usd, self.hr)
        assert str(mo.MOD(mo.Variable(14.0, unit="months"), 12).unit) == "months"
        assert str(mo.ROUND(self.usd, mo.Variable(1, unit="#")).unit) == "USD"

    def test_an_exponential_a_logarithm_and_a_sign_carry_none(self):
        for fn in (mo.EXP, mo.LN, mo.LOG10, mo.SIGN):
            assert fn(self.usd).unit is None

    def test_a_conditional_total_takes_its_lines_in_one_unit(self):
        a = mo.Variable(10.0, unit="USD", display_name="A")
        b = mo.Variable(5.0, unit="hr", display_name="B")
        calls = [
            lambda: mo.SUMIF([a, b], ">0"),
            lambda: mo.SUMIF([a, a], ">0", [a, b]),
            lambda: mo.AVERAGEIF([a, b], ">0"),
            lambda: mo.SUMIFS([a, b], [a, a], ">0"),
            lambda: mo.AVERAGEIFS([a, b], [a, a], ">0"),
            lambda: mo.MAXIFS([a, b], [a, a], ">0"),
            lambda: mo.MINIFS([a, b], [a, a], ">0"),
        ]
        for call in calls:
            with pytest.raises(ValueError, match="takes its values in one unit; 'A' and 'B' are"):
                call()
        assert str(mo.SUMIFS(self.usd, self.usd, ">0").unit) == "USD"
        # The criteria range's unit is never the result's; a count has none,
        # and lines in two units each meet a typed criteria.
        hours = mo.Variable([1.0, 2.0], unit="hr")
        assert str(mo.SUMIF(hours, ">1", self.usd).unit) == "USD"
        assert mo.COUNTIF([a, b], ">0").value == 2 and mo.COUNTIF([a, b], ">0").unit is None

    def test_a_criteria_compares_in_its_range_unit(self):
        cap = mo.Variable(3.5, unit="hr", display_name="Cap")
        calls = [
            lambda: mo.SUMIFS(self.usd, self.usd, (">=", cap)),
            lambda: mo.SUMIFS(self.usd, self.usd, mo.CONCAT(">=", cap)),
            lambda: mo.COUNTIF(self.usd, cap),
            lambda: mo.COUNTIFS(self.usd, (">=", cap)),
        ]
        for call in calls:
            with pytest.raises(ValueError, match="compares 'Amount' with the criteria 'Cap'"):
                call()
        floor = mo.Variable(3.5, unit="USD")
        assert mo.SUMIFS(self.usd, self.usd, (">=", floor)).value == 4.0

    @pytest.mark.parametrize("build", ["function", "operator", "pick"])
    def test_a_loop_derives_a_row_by_the_same_law(self, build):
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=3)
        with pytest.raises(ValueError, match="USD\\^2 and USD|takes its values in one unit"):
            with m:
                m.inp = mo.MultiVariable("Inputs")
                with m.inp as i:
                    i.opening = mo.Variable(1000.0, unit="USD")
                    i.flow = mo.Variable([10.0] * 3, unit="USD")
                    i.per_usd = mo.Variable(0.001, unit="1/USD")
                m.debt = mo.MultiVariable("Debt", loop=True)
                with m.debt as d:
                    prev = mo.lag(d.balance, fill=m.inp.opening)
                    if build == "function":
                        d.check = mo.POWER(prev, 2) + prev
                    elif build == "operator":
                        d.check = prev ** 2 + prev
                    else:
                        d.check = mo.IFS(prev > 0, prev, True, m.inp.per_usd)
                    d.balance = prev + m.inp.flow

    def test_a_loop_keeps_what_the_law_gives(self):
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=3)
        with m:
            m.inp = mo.MultiVariable("Inputs")
            with m.inp as i:
                i.opening = mo.Variable(1000.0, unit="USD")
                i.flow = mo.Variable([10.0] * 3, unit="USD")
                i.per_usd = mo.Variable(0.001, unit="1/USD")
            m.debt = mo.MultiVariable("Debt", loop=True)
            with m.debt as d:
                prev = mo.lag(d.balance, fill=m.inp.opening)
                d.square = mo.POWER(prev, 2)
                d.scaled = d.square * m.inp.per_usd
                d.root = mo.SQRT(d.square)
                d.balance = prev + m.inp.flow
        assert str(m.debt.square.unit) == "USD^2"
        assert str(m.debt.scaled.unit) == "USD" and str(m.debt.root.unit) == "USD"


def _loop_model(body, *, fill_unit="USD"):
    """A loop container built by ``body(d, inputs)`` over a three-month model."""
    m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=3)
    with m:
        m.inp = mo.MultiVariable("Inputs")
        with m.inp as i:
            i.start = mo.Variable(1000.0, unit=fill_unit)
            i.inflow = mo.Variable([10.0] * 3, unit="USD")
            i.cap = mo.Variable(500.0, unit="USD")
            i.per_usd = mo.Variable(0.001, unit="1/USD")
            i.growth = mo.Variable(0.01, unit="%")
        m.loop = mo.MultiVariable("Loop", loop=True)
        with m.loop as d:
            body(d, m.inp)
    return m.loop


class TestLoopUnitsSettle:
    """A loop container derives its rows' units again when it closes,
    from rows that are not derived yet: ``+`` takes the other side's
    unit, a product or a ratio waits for the row, a lag's fill anchors
    it — a row never takes a unit from a placeholder."""

    def test_a_shift_fill_gives_a_compounding_row_its_unit(self):
        def body(d, i):
            d.balance = d.balance.shift(1, fill_value=i.start) * (1 + i.growth)
        assert str(_loop_model(body).balance.unit) == "USD"

    def test_a_shift_fill_in_another_unit_is_refused_when_the_loop_closes(self):
        def body(d, i):
            d.opening = d.closing.shift(1, fill_value=i.start)
            d.closing = d.opening + i.inflow
        with pytest.raises(ValueError, match="takes its values in one unit|incompatible units"):
            _loop_model(body, fill_unit="EUR")

    def test_a_product_of_a_row_waits_for_its_unit(self):
        def body(d, i):
            d.opening = mo.lag(d.closing)
            d.square = d.opening * d.opening
            d.closing = d.opening + i.inflow + d.square * i.per_usd
        loop = _loop_model(body)
        assert str(loop.square.unit) == "USD^2" and str(loop.closing.unit) == "USD"

    def test_a_ratio_of_a_row_takes_no_unit_it_has_not_got(self):
        def body(d, i):
            d.opening = mo.lag(d.closing)
            d.used = d.opening / i.cap
            d.closing = d.opening + i.inflow * (1 + d.used)
        loop = _loop_model(body)
        # Nothing in the loop says what the opening is in: it is read as a
        # pure number, as the same formulas read an unlabelled line outside
        # a loop - the ratio is then per USD, the closing has no unit.
        assert str(loop.used.unit) == "1/USD" and loop.closing.unit is None

        def anchored(d, i):
            d.opening = mo.lag(d.closing, fill=i.start)
            d.used = d.opening / i.cap
            d.closing = d.opening + i.inflow * (1 + d.used)
        loop = _loop_model(anchored)
        assert loop.used.unit is None and str(loop.closing.unit) == "USD"

    def test_a_fill_gives_a_compounding_row_its_unit(self):
        def body(d, i):
            d.balance = mo.lag(d.balance, fill=i.start) * (1 + i.growth)
        assert str(_loop_model(body).balance.unit) == "USD"

    def test_a_fill_in_another_unit_is_refused(self):
        revenue = mo.Variable([1.0, 2.0], unit="USD", display_name="Revenue")
        with pytest.raises(ValueError, match="mo.lag takes its values in one unit"):
            mo.lag(revenue, fill=mo.Variable(5.0, unit="EUR", display_name="Opening"))
        assert str(mo.lag(revenue, fill=5.0).unit) == "USD"
        # A lag is its source's cells: an unlabelled line stays unlabelled.
        assert mo.lag(mo.Variable([1.0, 2.0]), fill=mo.Variable(5.0, unit="EUR")).unit is None

        def body(d, i):
            d.opening = mo.lag(d.closing, fill=i.start)
            d.closing = d.opening + i.inflow
        with pytest.raises(ValueError, match="takes its values in one unit|incompatible units"):
            _loop_model(body, fill_unit="EUR")

    def test_a_schedule_function_on_a_loop_row_keeps_the_unit(self):
        def body(d, i):
            d.book = mo.lag(d.closing, fill=i.start)
            d.dep = mo.SLN(d.book, i.cap, 10)
            d.closing = d.book - d.dep
        loop = _loop_model(body)
        assert str(loop.dep.unit) == "USD" and str(loop.closing.unit) == "USD"


class TestLoopUnitsJudgedWhenTheLoopCloses:
    """How a loop's units settle: provisional units
    are never judged before the loop closes, a pure-number row leaves its
    products their unit, rows settle the same way on every build, and a
    unit that never settles is refused."""

    @staticmethod
    def _stock(last_row):
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=3)
        with m:
            m.inp = mo.MultiVariable("Inputs")
            with m.inp as i:
                i.buy = mo.Variable([50.0, 50.0, 50.0], unit="t")
                i.sell = mo.Variable([40.0, 60.0, 30.0], unit="t")
                i.price = mo.Variable(250.0, unit="USD/t")
                i.step = mo.Variable(1000.0, unit="USD")
                i.salvage = mo.Variable(0.0, unit="USD")
                i.value0 = mo.Variable(25000.0, unit="USD")
                i.cap = mo.Variable(20000.0, unit="USD")
            m.s = mo.MultiVariable("Stock", loop=True)
            with m.s as s:
                s.inv_open = mo.lag(s.inv_close)
                s.value = s.inv_open * i.price           # provisional while the loop is open
                s.x = last_row(s, i)
                s.inv_close = s.inv_open + i.buy - i.sell
        return m.s

    @pytest.mark.parametrize("build", [
        lambda s, i: mo.CEILING(s.value, i.step),
        lambda s, i: mo.FLOOR(s.value, i.step),
        lambda s, i: mo.MOD(s.value, i.step),
        lambda s, i: mo.MAX(s.value, i.cap),
        lambda s, i: mo.IF(s.value > i.cap, s.value, i.cap),
        lambda s, i: mo.SLN(s.value, i.salvage, 12),
        lambda s, i: mo.IPMT(0.01, 1, 12, s.value, i.salvage),
        lambda s, i: mo.lag(s.value, fill=i.value0),
        lambda s, i: mo.COUNTIF(s.value, (">", i.cap)),
    ])
    def test_a_provisional_unit_is_judged_only_when_the_loop_closes(self, build):
        stock = self._stock(build)
        assert str(stock.value.unit) == "USD"

    @pytest.mark.parametrize("build", [
        lambda d, cap: mo.COUNTIF(d.balance, ("<", cap)),
        lambda d, cap: mo.SUMIF(d.balance, ("<", cap), d.balance),
        lambda d, cap: mo.AVERAGEIF(d.balance, (">=", cap)),
        lambda d, cap: mo.COUNTIFS(d.balance, ("<", cap)),
        lambda d, cap: mo.SUMIFS(d.balance, d.balance, ("<", cap)),
        lambda d, cap: mo.AVERAGEIFS(d.balance, d.balance, (">=", cap)),
        lambda d, cap: mo.MAXIFS(d.balance, d.balance, ("<", cap)),
        lambda d, cap: mo.MINIFS(d.balance, d.balance, (">=", cap)),
    ])
    def test_a_criteria_joined_to_a_line_in_another_unit_is_refused_when_the_loop_closes(
            self, build):
        # Re-run from its formula when the loop closes, the criteria is read
        # as given - an operator and a line - so its unit is compared with the
        # range's, as outside a loop; the join never becomes text.
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=4)
        with pytest.raises(ValueError, match="d.check: mo.[A-Z]+ compares 'balance' with the "
                                             "criteria 'Cap'; they are in USD and hr"):
            with m:
                m.cap = mo.Variable(500.0, unit="hr", display_name="Cap")
                m.repay = mo.Variable([200.0] * 4, unit="USD")
                m.d = mo.MultiVariable("D", loop=True)
                with m.d as d:
                    d.check = build(d, m.cap)
                    d.balance = mo.lag(d.balance, fill=1000.0) - m.repay

    def test_a_pure_number_row_leaves_its_products_their_unit(self):
        m = mo.Model("dcf", default_grain="year", default_start="2026", default_periods=4)
        with m:
            m.i = mo.MultiVariable("Inputs")
            m.i.rate = mo.Variable(0.08, unit="%")
            m.i.cf = mo.Variable([100.0, 110.0, 120.0, 130.0], unit="USD")
            m.i.hours = mo.Variable([1.0] * 4, unit="hr")
            m.d = mo.MultiVariable("Discounting", loop=True)
            with m.d as d:
                d.df_open = mo.lag(d.df_close, fill=1.0)
                d.pv = m.i.cf * d.df_open
                d.cum_open = mo.lag(d.cum_close)
                d.cum_close = d.cum_open + d.pv
                d.df_close = d.df_open / (1 + m.i.rate)
        assert str(m.d.pv.unit) == "USD" and str(m.d.cum_close.unit) == "USD"
        assert m.d.df_open.unit is None
        with pytest.raises(ValueError, match="incompatible units"):
            _ = m.d.pv + m.i.hours

    def test_a_row_whose_unit_never_settles_is_refused(self):
        def grows(d, i):
            d.x = mo.lag(d.x, fill=1.0) * i.cap
        with pytest.raises(ValueError, match="its unit does not settle"):
            _loop_model(grows)

    def test_an_unanchored_mix_is_still_refused(self):
        def body(d, i):
            d.open = mo.lag(d.close, fill=mo.Variable(1000.0))
            d.fee = d.open * i.cap
            d.close = d.open + d.fee
        with pytest.raises(ValueError, match="incompatible units"):
            _loop_model(body)

    def test_a_cycle_through_a_ratio_settles_its_units(self):
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=4)
        with m:
            m.inp = mo.MultiVariable("Inputs")
            with m.inp as i:
                i.demand = mo.Variable([100.0, 110.0, 120.0, 130.0], unit="t")
                i.capacity = mo.Variable(150.0, unit="t")
            m.plant = mo.MultiVariable("Plant", loop=True)
            with m.plant as d:
                d.output = i.demand / (1 + mo.lag(d.load))
                d.load = d.output / i.capacity
        assert str(m.plant.output.unit) == "t" and m.plant.load.unit is None

    def test_rows_settle_the_same_way_on_every_build(self):
        x = mo.Unit.dimensionless("x")
        outcomes = set()
        ballast = []
        for trial in range(8):
            ballast.append([object() for _ in range(trial * 37)])   # moves addresses only
            m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=3)
            with m:
                m.inp = mo.MultiVariable("Inputs")
                with m.inp as i:
                    i.multiple0 = mo.Variable(0.8, unit=x)
                    i.share0 = mo.Variable(0.05, unit="%")
                m.loop = mo.MultiVariable("Loop", loop=True)
                with m.loop as d:
                    d.a = mo.lag(d.b, fill=i.multiple0)
                    d.b = mo.lag(d.a, fill=i.share0)
            outcomes.add((repr(m.loop.a.unit), repr(m.loop.b.unit)))
        assert len(outcomes) == 1, outcomes


def _assign(container, name, value):
    setattr(container, name, value)


class TestLoopUnitsWhateverTheLineOrder:
    """Whatever the line order, a loop's units and values do not
    depend on the order its lines are written in; a row read through a lag
    with a fill in a unit is in that unit; a temporary whose recompute
    failed is never read."""

    LEASE_ORDERS = [["index", "rent", "billed", "fees"],
                    ["rent", "billed", "index", "fees"],
                    ["rent", "index", "billed", "fees"]]

    @staticmethod
    def _lease(order):
        m = mo.Model("lease", default_grain="month", default_start="2026-01", default_periods=4)
        with m:
            m.i = mo.MultiVariable("Inputs")
            with m.i as i:
                i.cpi = mo.Variable(0.004, unit="%")
                i.rent0 = mo.Variable(1000.0, unit="USD")
                i.area = mo.Variable(50.0, unit="m^2")
                i.cap_psm = mo.Variable(20.1, unit="USD/m^2")
                i.fee = mo.Variable([5.0] * 4, unit="USD")
            m.r = mo.MultiVariable("Rent", loop=True)
            with m.r as r:
                for name in order:
                    if name == "rent":
                        _assign(r, "rent", mo.lag(r.rent, fill=i.rent0)
                                * (r.index / mo.lag(r.index, fill=1.0)))
                    elif name == "index":
                        _assign(r, "index", mo.lag(r.index, fill=1.0) * (1 + i.cpi))
                    elif name == "billed":
                        _assign(r, "billed", mo.MIN(r.rent / i.area, i.cap_psm) * i.area + r.fees)
                    else:
                        _assign(r, "fees", i.fee + 0.0)
        return m.r

    @pytest.mark.parametrize("order", LEASE_ORDERS)
    def test_line_order_changes_neither_units_nor_values(self, order):
        r = self._lease(order)
        assert str(r.rent.unit) == "USD" and str(r.billed.unit) == "USD"
        assert r.billed.value == pytest.approx([1009.0, 1010.0, 1010.0, 1010.0])

    @pytest.mark.parametrize("order", [["value", "multiple", "qty"], ["multiple", "qty", "value"],
                                       ["value", "qty", "multiple"]])
    def test_a_fill_in_another_unit_is_refused_in_every_order(self, order):
        x = mo.Unit.dimensionless("x")
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=4)
        with pytest.raises(ValueError, match="mo.lag takes its values in one unit"):
            with m:
                m.i = mo.MultiVariable("Inputs")
                with m.i as i:
                    i.multiple = mo.Variable([1.5, 1.4, 1.3, 1.2], unit=x)
                    i.volume = mo.Variable([10.0] * 4, unit="t")
                    i.price0 = mo.Variable(2.0, unit="USD/t")
                m.d = mo.MultiVariable("D", loop=True)
                with m.d as d:
                    for name in order:
                        if name == "value":
                            _assign(d, "value", mo.lag(d.multiple, fill=i.price0) * d.qty)
                        elif name == "multiple":
                            _assign(d, "multiple", i.multiple + 0.0)
                        else:
                            _assign(d, "qty", i.volume + 0.0)

    @pytest.mark.parametrize("own_row", [False, True])
    def test_a_growth_curve_takes_its_fill_unit(self, own_row):
        m = mo.Model("saas", default_grain="month", default_start="2026-01", default_periods=6)
        with m:
            m.i = mo.MultiVariable("Inputs")
            with m.i as i:
                i.c0 = mo.Variable(100.0, unit="customers")
                i.market = mo.Variable(10_000.0, unit="customers")
                i.g = mo.Variable(0.30, unit="%")
            m.s = mo.MultiVariable("S-curve", loop=True)
            with m.s as s:
                if own_row:
                    s.penetration = mo.lag(s.customers) / i.market
                    s.customers = mo.lag(s.customers, fill=i.c0) * (1 + i.g * (1 - s.penetration))
                else:
                    s.customers = (mo.lag(s.customers, fill=i.c0)
                                   * (1 + i.g * (1 - mo.lag(s.customers) / i.market)))
        assert str(m.s.customers.unit) == "customers"
        # 130 * (1 + 0.3 * (1 - 130 / 10000))
        assert m.s.customers.value[1] == pytest.approx(168.493)

    @pytest.mark.parametrize("growth_first", [True, False])
    def test_a_compounding_line_is_in_its_fill_unit_whatever_comes_first(self, growth_first):
        m = mo.Model("saas", default_grain="month", default_start="2026-01", default_periods=6)
        with m:
            m.i = mo.MultiVariable("Inputs")
            with m.i as i:
                i.mrr0 = mo.Variable(10_000.0, unit="USD")
                i.churn = mo.Variable(0.02, unit="%")
                i.costs_eur = mo.Variable([9_000.0] * 6, unit="EUR")
            m.p = mo.MultiVariable("Plan", loop=True)
            with m.p as p:
                if growth_first:
                    p.growth = mo.lag(p.growth, fill=0.10) * 0.9
                p.mrr = mo.lag(p.mrr, fill=i.mrr0) * (1 + p.growth) - p.churned
                p.churned = mo.lag(p.mrr) * i.churn
                if not growth_first:
                    p.growth = mo.lag(p.growth, fill=0.10) * 0.9
        assert str(m.p.mrr.unit) == "USD" and str(m.p.churned.unit) == "USD"
        with pytest.raises(ValueError, match="incompatible units"):
            _ = m.p.mrr - m.i.costs_eur

    @pytest.mark.parametrize("service_first", [True, False])
    def test_a_row_typed_by_its_placeholder_keeps_its_unit_in_reports(self, service_first):
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=4)
        with m:
            m.i = mo.MultiVariable("Inputs")
            with m.i as i:
                i.cfads = mo.Variable([130.0, 90.0, 150.0, 160.0], unit="USD")
                i.debt0 = mo.Variable(1000.0, unit="USD")
                i.rate = mo.Variable(0.01, unit="%")
                i.fixed = mo.Variable(50.0, unit="USD")
            m.d = mo.MultiVariable("Debt", loop=True)
            with m.d as d:
                d.opening = mo.lag(d.closing, fill=i.debt0)
                if service_first:
                    d.service = d.opening * i.rate + i.fixed
                d.cover = i.cfads / d.service
                d.sweep = mo.IF(d.cover > 1.2, i.cfads - d.service, 0)
                if not service_first:
                    d.service = d.opening * i.rate + i.fixed
                d.closing = d.opening - d.sweep
        assert m.d.sweep.value_type == "float"
        for fn in (mo.SUM, mo.MAX):
            assert str(fn(m.d.sweep).unit) == "USD"
        assert str(mo.ROUND(m.d.sweep, 0).unit) == "USD"

    def test_a_temporary_that_failed_is_never_read(self):
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=3)
        with m:
            m.inp = mo.MultiVariable("Inputs")
            with m.inp as i:
                i.start = mo.Variable(1000.0, unit="USD")
                i.inflow = mo.Variable([10.0] * 3, unit="USD")
                i.hours = mo.Variable([8.0] * 3, unit="hr", display_name="Hours")
            m.loop = mo.MultiVariable("Loop", loop=True)
            with m.loop as d:
                d.open = mo.lag(d.close, fill=i.start)
                rem = mo.MOD(d.open, i.hours)
                twice = rem * 2
                d.close = d.open + i.inflow
                d.twice = twice
                d.copied = mo.Variable(rem)
        for line in (m.loop.twice, m.loop.copied):
            with pytest.raises(ValueError, match="mo.MOD takes its values in one unit"):
                _ = line.value


class TestALoopRefusedWhenItCloses:
    """A loop its unit check refuses when it closes hands out no numbers:
    every row keeps the refusal - its value, the lines built on it, its
    re-grain and the book say why - as a loop refused before it computes
    does (test_loops.py::test_rows_reading_a_failed_loop_say_why)."""

    @staticmethod
    def _refused(flat=False):
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=3)
        with m:
            m.inp = mo.MultiVariable("Inputs")
            with m.inp as i:
                i.opening = mo.Variable(1000.0, unit="EUR")
                i.fee = mo.Variable([1.0] * 3, unit="USD")
            m.debt = mo.MultiVariable("Debt", loop=True)
            with pytest.raises(ValueError, match="balance.*incompatible units: EUR and USD"):
                if flat:
                    m.debt.interest = mo.lag(m.debt.balance) * 0.05
                    m.debt.balance = (mo.lag(m.debt.balance, fill=i.opening)
                                      - m.debt.interest + i.fee)
                else:
                    with m.debt as d:
                        d.interest = mo.lag(d.balance) * 0.05
                        d.balance = mo.lag(d.balance, fill=i.opening) - d.interest + i.fee
        return m

    @pytest.mark.parametrize("flat", [False, True])
    def test_its_rows_hand_out_no_numbers(self, flat):
        m = self._refused(flat)
        assert m.debt._component_names == ["interest", "balance"]
        for row in (m.debt.interest, m.debt.balance):
            with pytest.raises(ValueError, match="loop could not be computed.*EUR and USD"):
                _ = row.value
            with pytest.raises(ValueError, match="loop could not be computed"):
                row.at("quarter")
        assert m.debt.to_dict()["components_data"]["balance"]["values"] is None

    def test_a_line_built_on_it_afterwards_says_why(self):
        m = self._refused()
        m.out = mo.MultiVariable("Out")
        m.out.total = m.debt.balance * 2
        with pytest.raises(ValueError, match="loop could not be computed.*EUR and USD"):
            _ = m.out.total.value

    def test_its_book_is_refused(self, tmp_path):
        m = self._refused()
        with pytest.raises(ValueError, match="loop could not be computed.*EUR and USD"):
            m.to_excel(tmp_path / "refused.xlsx")
        assert not (tmp_path / "refused.xlsx").exists()

    def test_a_unit_that_never_settles_leaves_no_numbers(self):
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=3)
        with m:
            m.cap = mo.Variable(500.0, unit="USD")
            m.loop = mo.MultiVariable("Loop", loop=True)
            with pytest.raises(ValueError, match="its unit does not settle"):
                with m.loop as d:
                    d.x = mo.lag(d.x, fill=1.0) * m.cap
        with pytest.raises(ValueError, match="loop could not be computed.*does not settle"):
            _ = m.loop.x.value

    def test_the_container_built_again_closes(self):
        m = self._refused()
        with m, warnings.catch_warnings():
            warnings.simplefilter("ignore")
            m.debt = mo.MultiVariable("Debt", loop=True)
            with m.debt as d:
                d.interest = mo.lag(d.balance) * 0.05
                d.balance = mo.lag(d.balance, fill=m.inp.opening) - d.interest
        assert m.debt.balance.value == pytest.approx([1000.0, 950.0, 902.5])
        assert str(m.debt.balance.unit) == "EUR"


class TestShiftUnits:
    """A lag keeps its source's unit; a fill in a unit is the source's."""

    def test_a_lag_of_an_unlabelled_line_stays_unlabelled(self):
        revenue = mo.Variable([100.0, 110.0, 121.0], display_name="Revenue")
        opening = mo.Variable(90.0, unit="USD")
        prior = mo.lag(revenue, fill=opening)
        assert prior.unit is None
        growth = revenue / prior - 1
        assert growth.unit is None
        assert (growth > mo.Variable(0.05, unit="%")).value is not None

    def test_shift_keeps_the_unit_and_checks_its_fill(self):
        revenue = mo.Variable([1.0, 2.0], unit="USD", display_name="Revenue")
        assert str(revenue.shift(1).unit) == "USD"
        assert str(revenue.shift(1, fill_value=mo.Variable(3.0, unit="USD")).unit) == "USD"
        with pytest.raises(ValueError, match="shift\\(\\) takes its values in one unit"):
            revenue.shift(1, fill_value=mo.Variable(3.0, unit="EUR", display_name="Opening"))

    def test_shift_and_lag_label_a_pure_number_alike(self):
        x = mo.Unit.dimensionless("x")
        share = mo.Variable([0.1, 0.2], unit="%", display_name="Share")
        for fill, label in ((mo.Variable(1.5, unit=x), None), (mo.Variable(0.5, unit="%"), "%"),
                            (0.5, "%")):
            shifted, lagged = share.shift(1, fill_value=fill), mo.lag(share, fill=fill)
            assert shifted.unit == lagged.unit
            assert (None if shifted.unit is None else str(shifted.unit)) == label

    def test_a_shift_fill_is_one_value_assigned_above(self):
        revenue = mo.Variable([1.0, 2.0], display_name="Revenue")
        with pytest.raises(TypeError, match="SCALAR"):
            revenue.shift(1, fill_value=mo.Variable([1.0, 2.0]))
        with pytest.raises(TypeError, match="assigned further down"):
            def body(d, i):
                d.a = d.b.shift(1, fill_value=d.c)
                d.b = d.a + i.inflow
                d.c = d.b * 0
            _loop_model(body)


class TestUnitLawsKeepTheirOwnFunctions:
    def test_a_python_function_has_no_unit_unless_declared(self):
        # The engine cannot see inside it: a ratio of two USD lines is a
        # pure number, a scaled USD series is still USD. Only the author
        # knows, so the result has none until declared - and a name like a
        # spreadsheet function's never lends that function's law.
        capacity = mo.Variable([10.0, 12.0], unit="MW")
        amounts = mo.Variable([5.0, 7.0], unit="USD")

        @mo.pyformula
        def power(cap, hours):
            return [c * hours for c in cap]

        @mo.pyformula
        def sumif(rows):
            return sum(rows)

        @mo.pyformula
        def ratio(a, b):
            return [x / y for x, y in zip(a, b)]

        @mo.pyformula(unit="USD")
        def product(rows, k):
            return [r * k for r in rows]

        assert power(capacity, 8760).unit is None
        assert sumif(amounts).unit is None
        assert ratio(amounts, amounts).unit is None
        assert str(product(amounts, 2).unit) == "USD"
        assert str((product(amounts, 2) + amounts).unit) == "USD"

    def test_a_loop_never_reruns_a_python_function_as_the_spreadsheet_one(self):
        @mo.pyformula
        def product(rows, k):
            return [r * k for r in rows]

        m = mo.Model("m", default_grain="year", default_start="2026", default_periods=3)
        with pytest.raises(ValueError, match="cannot be computed again when the loop closes"):
            with m:
                m.inp = mo.MultiVariable("Inputs")
                with m.inp as i:
                    i.start = mo.Variable(10.0, unit="MW")
                    i.add = mo.Variable([1.0] * 3, unit="MW")
                m.loop = mo.MultiVariable("Loop", loop=True)
                with m.loop as d:
                    d.cap_open = mo.lag(d.cap_close, fill=i.start)
                    d.x = product(d.cap_open, 2)
                    d.cap_close = d.cap_open + i.add

    def test_a_row_typed_by_an_error_marker_keeps_its_unit_through_min_max(self):
        actual = mo.Variable(["#N/A", 120.0, 130.0], unit="USD")
        assert str(mo.MAX(actual, 0).unit) == "USD"
        assert str(mo.MIN(actual, 100.0).unit) == "USD"

    def test_two_records_with_one_name_are_told_apart_by_path(self):
        m = mo.Model("fx")
        with m:
            m.ledger = mo.MultiVariable("Ledger")
            with m.ledger as led:
                led.a = mo.MultiVariable("A")
                led.a.amount = mo.Variable(10.0, unit="USD", display_name="Amount")
                led.b = mo.MultiVariable("B")
                led.b.amount = mo.Variable(8.0, unit="EUR", display_name="Amount")
        with pytest.raises(ValueError, match="ledger.a.amount' and 'fx.ledger.b.amount"):
            mo.SUM(m.ledger.a.amount, m.ledger.b.amount)


class TestFinancialUnits:
    """The amounts a financial function reads together are in one unit, the
    result's; a rate, a term, a period never are, and a rate or a count of
    periods carries none."""

    def setup_method(self):
        self.rate = mo.Variable(0.01, unit="%")
        self.term = mo.Variable(24, unit="months")
        self.loan = mo.Variable(10_000.0, unit="USD", display_name="Loan")
        self.balloon = mo.Variable(1_000.0, unit="EUR", display_name="Balloon")
        self.flows = mo.Variable([-100.0, 60.0, 60.0], unit="USD")

    def test_the_result_is_in_the_amounts_unit(self):
        assert str(mo.NPV(self.rate, self.flows).unit) == "USD"
        assert str(mo.PMT(self.rate, self.term, self.loan).unit) == "USD"
        assert str(mo.FV(self.rate, self.term, -100.0, self.loan).unit) == "USD"
        assert str(mo.PV(self.rate, self.term, mo.Variable(-100.0, unit="USD")).unit) == "USD"
        assert str(mo.IPMT(self.rate, 1, self.term, self.loan).unit) == "USD"
        assert str(mo.PPMT(self.rate, 1, self.term, self.loan).unit) == "USD"
        assert str(mo.CUMIPMT(self.rate, self.term, self.loan, 1, 12, 0).unit) == "USD"
        cost, salvage = mo.Variable(900.0, unit="USD"), mo.Variable(100.0, unit="USD")
        life = mo.Variable(8, unit="years")
        for fn in (mo.SLN, mo.SYD, mo.DB, mo.DDB):
            args = (cost, salvage, life) if fn is mo.SLN else (cost, salvage, life, 1)
            assert str(fn(*args).unit) == "USD", fn

    def test_a_rate_and_a_count_carry_none(self):
        assert mo.IRR(self.flows).unit is None
        assert mo.NPER(self.rate, mo.Variable(-500.0, unit="USD"), self.loan).unit is None
        assert mo.RATE(self.term, mo.Variable(-500.0, unit="USD"), self.loan).unit is None

    def test_amounts_in_two_units_are_refused(self):
        calls = [
            lambda: mo.PMT(self.rate, self.term, self.loan, self.balloon),
            lambda: mo.FV(self.rate, self.term, self.balloon, self.loan),
            lambda: mo.NPER(self.rate, self.balloon, self.loan),
            lambda: mo.IPMT(self.rate, 1, self.term, self.loan, self.balloon),
            lambda: mo.SLN(self.loan, self.balloon, 5),
        ]
        for call in calls:
            with pytest.raises(ValueError, match="takes its values in one unit; 'Loan' and"
                                                 " 'Balloon'|'Balloon' and 'Loan'"):
                call()


class TestFinancialUnitsOfAListOfLines:
    """A list of scalar lines is read line by line, as a row is read: one
    unit among the flows, the result's for NPV / XNPV, none for IRR / XIRR;
    two units refused; the book writes the lines' cells."""

    @staticmethod
    def _lines():
        return dict(
            rate=mo.Variable(0.01, unit="%"),
            loan=mo.Variable(10_000.0, unit="USD", display_name="Loan"),
            balloon=mo.Variable(1_000.0, unit="EUR", display_name="Balloon"),
            y1=mo.Variable(6_000.0, unit="USD", display_name="Y1"),
            y2=mo.Variable(6_000.0, unit="USD", display_name="Y2"),
            close=mo.Variable(mo.DATE(2025, 1, 1)),
            exit=mo.Variable(mo.DATE(2027, 1, 1)),
        )

    def test_the_result_is_in_the_flows_unit(self):
        v = self._lines()
        assert str(mo.NPV(v["rate"], [v["y1"], v["y2"]]).unit) == "USD"
        assert str(mo.NPV(0.1, [v["y1"], v["y2"]]).unit) == "USD"
        assert str(mo.NPV(v["rate"], [v["y1"], 50.0]).unit) == "USD"
        assert str(mo.XNPV(v["rate"], [-v["loan"], v["y1"]],
                           [v["close"], v["exit"]]).unit) == "USD"
        assert mo.IRR([-v["loan"], v["y1"], v["y2"]]).unit is None
        assert mo.XIRR([-v["loan"], v["y1"]], [v["close"], v["exit"]]).unit is None

    @pytest.mark.parametrize("call", ["NPV", "NPV typed rate", "IRR", "XNPV", "XIRR"])
    def test_flows_in_two_units_are_refused(self, call):
        v = self._lines()
        calls = {
            "NPV": lambda: mo.NPV(v["rate"], [v["y1"], v["balloon"]]),
            "NPV typed rate": lambda: mo.NPV(0.1, [v["y1"], v["balloon"]]),
            "IRR": lambda: mo.IRR([-v["loan"], v["y1"], v["balloon"]]),
            "XNPV": lambda: mo.XNPV(v["rate"], [v["y1"], v["balloon"]],
                                    [v["close"], v["exit"]]),
            "XIRR": lambda: mo.XIRR([-v["loan"], v["balloon"]], [v["close"], v["exit"]]),
        }
        name = call.split()[0]
        with pytest.raises(ValueError, match=rf"mo\.{name} takes its values in one unit; "
                                             r".* are in (USD and EUR|EUR and USD)"):
            calls[call]()

    def test_the_result_no_longer_adds_to_another_unit(self):
        v = self._lines()
        eur = mo.Variable(5.0, unit="EUR")
        with pytest.raises(ValueError, match="incompatible units: EUR and USD"):
            eur + mo.NPV(v["rate"], [v["y1"], v["y2"]])

    def test_the_book_writes_the_cells(self, tmp_path):
        # The list's cells are the call's range: live in the book, and the
        # lines are the call's dependencies.
        m = mo.Model("deal")
        with m:
            m.rate = mo.Variable(0.10, display_name="Rate")
            m.cost = mo.Variable(-100.0, display_name="Cost", unit="EUR")
            m.y1 = mo.Variable(50.0, display_name="Y1", unit="EUR")
            m.y2 = mo.Variable(60.0, display_name="Y2", unit="EUR")
            m.npv = mo.NPV(m.rate, [m.y1, m.y2]).set_display_name("NPV")
            m.irr = mo.IRR([m.cost, m.y1, m.y2]).set_display_name("IRR")
        assert [v.python_name for v in m.npv._dependency_refs] == ["rate", "y1", "y2"]
        m.to_excel(tmp_path / "npv.xlsx")
        ws = load_workbook(tmp_path / "npv.xlsx").active
        cells = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(1, ws.max_row + 1)}
        assert cells["NPV (EUR)"] == "=NPV(B1, B3:B4)"
        assert cells["IRR"] == "=IRR(B2:B4, 0.1)"


class TestAddition:
    def test_same_unit_keeps_unit(self):
        a = mo.Variable(10, unit="$")
        b = mo.Variable(20, unit="$")
        assert str((a + b)._unit) == "$"

    def test_mixed_units_raise(self):
        a = mo.Variable(10, unit="$")
        b = mo.Variable(20, unit="hr")
        with pytest.raises(ValueError, match="incompatible units"):
            _ = a + b

    def test_subtract_mixed_raise(self):
        a = mo.Variable(10, unit="$")
        b = mo.Variable(20, unit="hr")
        with pytest.raises(ValueError, match="incompatible units"):
            _ = a - b

    def test_scalar_plus_unitful_keeps_unit(self):
        a = mo.Variable(10, unit="$")
        # 5 is dimensionless/scalar — treat as compatible, keep unit
        assert str((a + 5)._unit) == "$"

    def test_a_rate_per_year_beside_an_amount_points_to_the_year_share(self):
        # 'USD/year' times days / 365 stays per year: the refusal says how
        # to make it an amount, and that way adds.
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=3)
        with m:
            m.rate = mo.Variable(1200.0, unit="USD/year")
            m.amount = mo.Variable([100.0] * 3, unit="USD")
            with pytest.raises(ValueError, match="per 'year'") as err:
                _ = m.amount + m.rate * mo.time.days / 365
            assert "year_share = (mo.time.months / 12).set_unit('year')" in str(err.value)
            year_share = (mo.time.months / 12).set_unit("year")
            total = m.amount + m.rate * year_share
        assert str(total._unit) == "USD"
        assert total.value == pytest.approx([200.0] * 3)

    def test_a_rate_per_another_word_names_it_without_a_year(self):
        a, b = mo.Variable(10.0, unit="USD/seat"), mo.Variable(1.0, unit="USD")
        with pytest.raises(ValueError, match="per 'seat'") as err:
            _ = b + a
        assert "year_share" not in str(err.value)

    def test_months_beside_years_write_out_the_step(self):
        # Dividing by 12 does not make months years: the label stays,
        # so the refusal writes the explicit step both ways.
        months = mo.Variable(18.0, unit="months")
        years = mo.Variable(2.0, unit="yrs")
        with pytest.raises(ValueError, match="never rescales") as err:
            _ = months / 12 + years
        text = str(err.value)
        assert "(x / 12).set_unit('yrs')" in text
        assert "(x * 12).set_unit('months')" in text
        assert "Otherwise convert" not in text
        total = (months / 12).set_unit("yrs") + years
        assert str(total._unit) == "yrs" and total.value == pytest.approx(3.5)

    def test_a_rate_per_month_beside_a_rate_per_year_writes_out_the_step(self):
        monthly = mo.Variable(100.0, unit="USD/month")
        yearly = mo.Variable(1200.0, unit="USD/yr")
        with pytest.raises(ValueError, match="never rescales") as err:
            _ = yearly + monthly
        assert "(x * 12).set_unit('USD/yr')" in str(err.value)
        assert "(x / 12).set_unit('USD/month')" in str(err.value)
        total = yearly + (monthly * 12).set_unit("USD/yr")
        assert str(total._unit) == "USD/yr" and total.value == pytest.approx(2400.0)

    def test_a_function_over_months_and_years_carries_the_step(self):
        months = mo.Variable(6.0, unit="months")
        years = mo.Variable(0.5, unit="yrs")
        with pytest.raises(ValueError, match="mo.IF takes its values in one unit") as err:
            mo.IF(months > 0, months / 12, years)
        assert "(x / 12).set_unit('yrs')" in str(err.value)

    def test_a_remainder_keeps_the_unit_both_sides_share(self):
        # ``%`` is MOD: a number goes with any, two units are refused.
        usd = mo.Variable([25.0, 31.0], unit="USD")
        assert str((usd % 12)._unit) == "USD" and (usd % 12).value == [1.0, 7.0]
        assert str((100 % usd)._unit) == "USD"
        assert str((usd % mo.Variable(10.0, unit="USD"))._unit) == "USD"
        with pytest.raises(ValueError, match="Cannot take the remainder of values with "
                                             "incompatible units: USD and hr"):
            _ = usd % mo.Variable(2.0, unit="hr")

    def test_a_reflected_floor_division_divides_the_units(self):
        usd = mo.Variable(25.0, unit="USD")
        assert str((100 // usd)._unit) == str((100 / usd)._unit)
        assert str((usd // 12)._unit) == "USD"

    def test_a_remainder_inside_a_loop_reads_the_same_law(self):
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=3)
        with m:
            m.i = mo.MultiVariable("I")
            m.i.add = mo.Variable([5.0, 5.0, 5.0], unit="USD")
            m.i.hours = mo.Variable(4.0, unit="hr")
            m.l = mo.MultiVariable("L", loop=True)
            with m.l as lp:
                lp.open = mo.lag(lp.close, fill=mo.Variable(0.0, unit="USD"))
                lp.rest = lp.open % 4
                lp.close = lp.open + m.i.add
        assert str(m.l.rest.unit) == "USD"
        m2 = mo.Model("m2", default_grain="month", default_start="2026-01", default_periods=3)
        with pytest.raises(ValueError, match="remainder of values with incompatible units"):
            with m2:
                m2.i = mo.MultiVariable("I")
                m2.i.add = mo.Variable([5.0, 5.0, 5.0], unit="USD")
                m2.i.hours = mo.Variable(4.0, unit="hr")
                m2.l = mo.MultiVariable("L", loop=True)
                with m2.l as lp:
                    lp.open = mo.lag(lp.close, fill=mo.Variable(0.0, unit="USD"))
                    lp.rest = lp.open % m2.i.hours
                    lp.close = lp.open + m2.i.add

    def test_without_a_hint_the_advice_does_not_open_with_otherwise(self):
        a, b = mo.Variable(1.0, unit="USD"), mo.Variable(1.0, unit="EUR")
        with pytest.raises(ValueError) as err:
            _ = a + b
        text = str(err.value)
        assert "USD and EUR. Convert one side to the other's unit" in text
        assert "Otherwise" not in text


class TestPower:
    def test_integer_exponent_raises_unit(self):
        t = mo.Variable(3, unit="hr")
        assert str((t ** 2)._unit) == "hr^2"
        assert str((t ** 3)._unit) == "hr^3"

    def test_zero_exponent_dimensionless(self):
        t = mo.Variable(3, unit="hr")
        # hr^0 = dimensionless
        assert (t ** 0)._unit is None

    def test_float_exponent_drops_unit(self):
        t = mo.Variable(3, unit="hr")
        # Fractional/non-int exponents have no clean unit rendering
        assert (t ** 0.5)._unit is None

    def test_a_whole_power_written_as_a_float_or_a_root_of_one_raises(self):
        t = mo.Variable(3, unit="hr")
        assert str((t ** 2.0)._unit) == "hr^2"
        assert str(((t * t) ** 0.5)._unit) == "hr"
        assert (t ** mo.Variable(2))._unit is None


class TestNegation:
    def test_unary_minus_preserves_unit(self):
        a = mo.Variable(10, unit="$")
        assert str((-a)._unit) == "$"


class TestExcelEmission:
    def test_a_format_keyed_by_a_chain_unit_follows_a_derived_line(self):
        # The derived line's unit prints canonically, not as the key is
        # written; the format is found by unit equality all the same.
        from modeleon.compile.excel.writer import resolve_number_format
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=2)
        m.default_excel_view = mo.ExcelView(formats={"USD/kW/month": "0.000", "%": "0.0%"})
        m.capacity = mo.Variable([5.0, 5.0], unit="kW")
        m.rent = mo.Variable([2.0, 2.0], unit="USD/kW/month")
        m.derived = (m.capacity * m.rent) / m.capacity
        m.share = mo.Variable(0.05, unit="%")
        m.step = mo.Variable(0.005, unit="bps")
        assert str(m.derived.unit) != "USD/kW/month"
        assert resolve_number_format(m.rent) == "0.000"
        assert resolve_number_format(m.derived) == "0.000"
        assert resolve_number_format(m.share) == "0.0%"
        assert resolve_number_format(m.step) is None      # '%' never dresses bps

    def test_unit_appears_in_label(self, tmp_path):
        m = mo.MultiVariable("UT")
        m.u = mo.MultiVariable("U", excel_props={'tab': True})
        m.u.revenue = mo.Variable(1000, unit="$")
        m.to_excel(tmp_path / "u.xlsx")
        path = tmp_path / "u.xlsx"
        ws = load_workbook(path)["U"]
        labels = [
            cell.value for row in ws.iter_rows() for cell in row
            if isinstance(cell.value, str)
        ]
        assert any("($)" in label for label in labels)

    def test_no_unit_no_suffix(self, tmp_path):
        m = mo.MultiVariable("UT")
        m.u = mo.MultiVariable("U", excel_props={'tab': True})
        m.u.plain = mo.Variable(1000)
        m.to_excel(tmp_path / "u.xlsx")
        path = tmp_path / "u.xlsx"
        ws = load_workbook(path)["U"]
        labels = [
            cell.value for row in ws.iter_rows() for cell in row
            if isinstance(cell.value, str)
        ]
        assert not any("(" in label for label in labels)


# ══════════════════════════════════════════════════════════════════════
# A word that labels a pure number
# ══════════════════════════════════════════════════════════════════════


FACTOR = mo.Unit.dimensionless("factor")
SHARE = mo.Unit.dimensionless("share")


class TestLabelledPureNumber:
    def test_a_label_prints_and_is_inert(self):
        cash = mo.Variable(100.0, unit="USD m")
        df = mo.Variable(0.9, unit=FACTOR)
        assert str(df.unit) == "factor" and FACTOR.is_dimensionless
        pv = cash * df
        assert pv.unit is cash.unit                     # the label adds nothing
        assert str((cash + pv).unit) == "USD m"         # and so the two add

    def test_a_string_is_always_a_dimension(self):
        cash = mo.Variable(100.0, unit="USD m")
        df = mo.Variable(0.9, unit="factor")
        assert mo.Unit("factor")._unit_components == {"factor": 1}
        with pytest.raises(ValueError, match="incompatible units"):
            cash + cash * df

    def test_pure_numbers_add_the_left_label_prints_a_product_drops_it(self):
        df = mo.Variable(0.9, unit=FACTOR)
        sh = mo.Variable(0.1, unit=SHARE)
        assert str((df + sh).unit) == "factor"
        assert str((sh + df).unit) == "share"
        assert (df * df).unit is None and (df * 2).unit is None
        pct, pp = mo.Variable(0.05, unit="%"), mo.Variable(0.01, unit="pp")
        assert str((pct + pp).unit) == "%"              # the conventional words, unchanged

    def test_the_conventional_words_are_the_same_kind_of_unit(self):
        assert mo.Unit("%") == mo.Unit.dimensionless("%") == mo.Unit.dimensionless()
        assert str(mo.Unit("%")) == "%"

    def test_repr_rebuilds_the_unit(self):
        for u in (FACTOR, mo.Unit("%"), mo.Unit("USD m"), mo.Unit("$/hr"),
                  mo.Unit.dimensionless()):
            back = eval(repr(u), {"Unit": mo.Unit})
            assert back == u and str(back) == str(u), repr(u)
        assert repr(FACTOR) == "Unit.dimensionless('factor')"

    @pytest.mark.parametrize("label", ["", "   ", "a/b", "a * b", "m^2"])
    def test_a_label_is_one_word(self, label):
        with pytest.raises(ValueError):
            mo.Unit.dimensionless(label)

    def test_what_is_not_a_unit_is_refused(self):
        with pytest.raises(TypeError):
            mo.Unit.dimensionless(3)
        with pytest.raises(TypeError):
            mo.Unit("factor", dimensionless=True)       # no stray keywords
        with pytest.raises(TypeError, match="a unit is a string"):
            mo.Variable(1.0, unit=5)
        with pytest.raises(TypeError, match="a unit is a string"):
            mo.Variable(1.0).set_unit(mo.Variable(2.0))

    def test_the_add_error_names_the_word_that_leaked_through_a_product(self):
        cash = mo.Variable(100.0, unit="USD m")
        with pytest.raises(ValueError) as err:
            cash + cash * mo.Variable(0.9, unit="factor")
        msg = str(err.value)
        assert "'factor' is a dimension here" in msg
        assert "mo.Unit.dimensionless('factor')" in msg

    @pytest.mark.parametrize("left, right", [
        ("EUR", "USD"), ("EUR/kWh", "USD/kWh"), ("USD m/yr", "EUR m/yr"),
        ("MWh/yr", "GWh/yr"), ("USD*hr", "USD*day"), ("USD/hr", "USD*hr"),
        ("USD*hr", "USD/hr"), ("USD", "USD*hr*hr"), ("kWh", "kWh*h*h"),
    ])
    def test_the_add_error_stays_plain_where_no_word_leaked(self, left, right):
        """Neither side is the other times words: a real mismatch, no hint."""
        with pytest.raises(ValueError) as err:
            mo.Variable(1.0, unit=left) + mo.Variable(1.0, unit=right)
        assert "dimensionless" not in str(err.value), (left, right)

    @pytest.mark.parametrize("left, right", [
        ("USD m", "USD m * factor"), ("USD m * factor", "USD m"),
        ("EUR/kWh", "EUR * flag/kWh"),
    ])
    def test_the_hint_fires_in_either_order(self, left, right):
        with pytest.raises(ValueError) as err:
            mo.Variable(1.0, unit=left) + mo.Variable(1.0, unit=right)
        assert "mo.Unit.dimensionless(" in str(err.value), (left, right)

    def test_the_add_error_tells_a_dimension_from_a_label_of_the_same_word(self):
        with pytest.raises(ValueError) as err:
            mo.Variable(1.0, unit="factor") + mo.Variable(1.0, unit=FACTOR)
        assert "factor (a dimension) and factor (a labelled pure number)" in str(err.value)

    def test_a_unit_is_a_value_never_a_row(self, tmp_path):
        m = mo.Model("m", default_grain="year", default_start="2026", default_periods=2)
        with m:
            m.F = FACTOR
            m.s = mo.MultiVariable("S")
            with m.s as s:
                s.df = mo.Variable([0.9, 0.8], unit=m.F, display_name="Discount factor")
        assert "F" not in m._component_names and m.F is FACTOR
        m.to_excel(tmp_path / "u.xlsx")
        assert load_workbook(tmp_path / "u.xlsx").sheetnames == ["S"]

    def test_the_label_survives_copy_and_pickle(self):
        import pickle
        df = mo.Variable(0.9, unit=FACTOR)
        assert str(df.copy().unit) == "factor"
        assert str(pickle.loads(pickle.dumps(df)).unit) == "factor"
        assert str(pickle.loads(pickle.dumps(FACTOR))) == "factor"

    def test_the_book_prints_the_label(self, tmp_path):
        m = mo.MultiVariable("UT")
        m.u = mo.MultiVariable("U", excel_props={'tab': True})
        m.u.df = mo.Variable(0.9, unit=FACTOR, display_name="Discount factor")
        m.to_excel(tmp_path / "u.xlsx")
        ws = load_workbook(tmp_path / "u.xlsx")["U"]
        labels = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
        assert "Discount factor (factor)" in labels

    def test_code_writes_the_constructor_for_a_label(self):
        df = mo.Variable(0.9, unit=FACTOR)
        assert "unit=mo.Unit.dimensionless('factor')" in df.code
        assert "unit='USD m'" in mo.Variable(1.0, unit="USD m").code
        assert "unit='%'" in mo.Variable(0.1, unit="%").code

    def test_an_aggregate_over_a_dimension_and_its_namesake_names_neither(self):
        dim = mo.Variable(1.0, unit="factor")
        lab = mo.Variable(1.0, unit=FACTOR)
        with pytest.raises(ValueError, match=r"factor \(a dimension\) and factor \(a labelled pure number\)"):
            mo.SUM(dim, lab)                           # refused, as dim + lab is
        assert mo.SUM(lab, mo.Variable(1.0, unit=SHARE)).unit is None
        assert str(mo.SUM(lab, lab).unit) == "factor"


class TestDeclaredUnitAcrossViews:
    def test_a_declared_unit_survives_a_grain_change(self):
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=6)
        with m:
            m.s = mo.MultiVariable("S")
            with m.s as s:
                s.a = mo.Variable([0.1] * 6, unit="%", regrain=mo.up("mean"))
                s.b = mo.Variable([0.2] * 6, unit="%", regrain=mo.up("mean"))
                s.f = (s.a + s.b).set_unit(FACTOR)
                s.c1 = mo.Variable([1.0] * 6, unit="USD", regrain=mo.up("sum"))
                s.c2 = (s.c1 + s.c1).set_unit("USD m")
        q = m.at("quarter")
        assert str(q.s.f.unit) == "factor"
        assert str(q.s.c2.unit) == "USD m"

    def test_a_container_in_the_units_place_carries_no_unit(self):
        stand_in = mo.MultiVariable("Elsewhere")
        assert mo.Variable(1.0, unit=stand_in).unit is None

    def test_a_power_written_with_a_caret_is_that_power(self):
        # A unit prints a power as hr^2, and the string reads back as it.
        a = mo.Variable([1.0, 2.0], unit="hr")
        assert str(((a * a) + mo.Variable([1.0, 1.0], unit="hr^2")).unit) == "hr^2"
        assert mo.Unit("USD/m^2") * mo.Unit("m^2") == mo.Unit("USD")
        assert mo.Unit("(kW * h)^2") == mo.Unit("kW") * mo.Unit("kW") * mo.Unit("h") * mo.Unit("h")
        assert eval(repr(mo.Unit("hr") ** -2), {"Unit": mo.Unit}) == mo.Unit("1/hr^2")
        assert str(mo.Unit("USD/m^2")) == "USD/m^2"
