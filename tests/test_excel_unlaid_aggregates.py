# SPDX-License-Identifier: Apache-2.0
"""An aggregate over an unnamed line: written natively where Excel has
the word for it, and listed where it is written as a number.

``mo.SUM(price * volume)`` reduces a line the book has no cells for.
Excel's word for it is ``SUMPRODUCT(price_range, volume_range)``, so
the cell stays live. ``MAX(ABS(row))`` has none: the cell carries its
number, and ``NPV(rate, cfads / min_dscr) * 2`` carries the NPV's
number inside a live formula. Both used to go into the book with one
log line; ``check_compat(model, "excel")`` now lists every such cell —
the line, the period, the sub-expression and the number — the cells
written as whole values apart from the ones with a number inlined.
"""

import shutil
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook

import modeleon as mo
from modeleon.compile import CompatIssue, check_compat


def _rows(ws):
    out = {}
    for row in ws.iter_rows():
        label = row[0].value
        if label:
            out[str(label)] = [c.value for c in row[1:]]
    return out


def _book(model, tmp_path):
    path = tmp_path / "book.xlsx"
    model.to_excel(path)
    return load_workbook(path)


def _model():
    m = mo.Model("m", default_grain="year", default_start="2026", default_periods=4)
    with m:
        m.a = mo.MultiVariable("A")
        m.a.price = mo.Variable([10.0, 11.0, 12.0, 13.0], display_name="price")
        m.a.volume = mo.Variable([100.0, 110.0, 120.0, 130.0], display_name="volume")
        m.a.rate = mo.Variable(0.1, display_name="rate")
        m.a.cfads = mo.Variable([50.0, 60.0, 70.0, 80.0], display_name="cfads")
        m.a.min_dscr = mo.Variable(1.3, display_name="min_dscr")
        m.a.row = mo.Variable([-5.0, 3.0, -8.0, 2.0], display_name="row")
        m.s = mo.MultiVariable("S")
    return m


class TestASumOfAProductIsSumproduct:
    def test_over_laid_out_rows(self, tmp_path):
        m = _model()
        with m:
            m.s.total = mo.SUM(m.a.price * m.a.volume).set_display_name("total")
            m.s.plus = (mo.SUM(m.a.price * m.a.volume) + m.a.rate).set_display_name("plus")
            m.s.three = mo.SUM(m.a.price * m.a.volume * m.a.cfads).set_display_name("three")
        s = _rows(_book(m, tmp_path)["S"])
        assert s["total"][0] == "=SUMPRODUCT(A!C2:F2, A!C3:F3)", s["total"]
        assert s["plus"][0] == "=SUMPRODUCT(A!C2:F2, A!C3:F3) + A!B4", s["plus"]
        assert s["three"][0] == "=SUMPRODUCT(A!C2:F2, A!C3:F3, A!C5:F5)", s["three"]
        # The value is the engine's SUM of the product, as before.
        assert m.s.total.value == 5340.0 and m.s.three.value == 358600.0

    def test_on_the_rows_own_sheet(self, tmp_path):
        m = _model()
        with m:
            m.a.total = mo.SUM(m.a.price * m.a.volume).set_display_name("total")
        a = _rows(_book(m, tmp_path)["A"])
        assert a["total"][0] == "=SUMPRODUCT(C2:F2, C3:F3)", a["total"]

    def test_any_other_shape_stays_a_number(self, tmp_path):
        m = _model()
        with m:
            m.s.scaled = mo.SUM(m.a.price * 2.0).set_display_name("scaled")
            m.s.avg = mo.AVERAGE(m.a.price * m.a.volume).set_display_name("avg")
            m.s.mx = mo.MAX(mo.ABS(m.a.row)).set_display_name("mx")
            m.s.named = mo.SUM(m.a.cfads).set_display_name("named")
        s = _rows(_book(m, tmp_path)["S"])
        assert s["scaled"][0] == 92.0 and s["avg"][0] == 1335.0 and s["mx"][0] == 8.0, s
        assert s["named"][0] == "=SUM(A!C5:F5)", s["named"]


class TestCheckCompatListsTheCells:
    def test_inlined_numbers_and_value_cells_apart(self):
        m = _model()
        with m:
            m.s.npv2 = (mo.NPV(m.a.rate, m.a.cfads / m.a.min_dscr) * 2).set_display_name("npv2")
            m.s.mx = mo.MAX(mo.ABS(m.a.row)).set_display_name("mx")
            m.s.live = mo.SUM(m.a.price * m.a.volume).set_display_name("live")
            m.s.named = mo.SUM(m.a.cfads).set_display_name("named")
        issues = check_compat(m, "excel")
        assert all(i.target == "excel" and i.declared == frozenset() for i in issues), issues
        by_kind = {k: [i for i in issues if i.kind == k] for k in ("inlined", "value", "render")}
        assert by_kind["render"] == []
        [inlined] = by_kind["inlined"]
        assert (inlined.qpath, inlined.period, inlined.func) == (m.s.npv2.id, 0, "NPV")
        assert inlined.inlined == "NPV(rate, (cfads / min_dscr))"
        assert float(inlined.literal) == pytest.approx(m.s.npv2.value / 2)
        assert str(inlined) == (
            f"{m.s.npv2.id}[0]: NPV(rate, (cfads / min_dscr)) is written as "
            f"{inlined.literal} inside a live formula")
        [value] = by_kind["value"]
        assert (value.qpath, value.period, value.func) == (m.s.mx.id, 0, "MAX")
        assert value.inlined == "MAX((ABS(row)))" and value.literal is None
        assert str(value) == f"{m.s.mx.id}[0]: written as its value — no formula stands for MAX((ABS(row)))"

    def test_every_period_of_a_series_is_listed(self):
        m = _model()
        with m:
            m.s.each = (m.a.price * mo.NPV(m.a.rate, m.a.cfads / m.a.min_dscr)).set_display_name("each")
        issues = check_compat(m, "excel")
        assert [(i.kind, i.qpath, i.period) for i in issues] == [
            ("inlined", m.s.each.id, p) for p in range(4)], issues

    def test_inputs_and_live_formulas_are_not_issues(self):
        m = _model()
        with m:
            m.s.rev = (m.a.price * m.a.volume).set_display_name("rev")
            m.s.total = mo.SUM(m.s.rev).set_display_name("total")
            m.s.mixed = mo.Variable([1.0, m.a.rate * 2, 3.0, 4.0], display_name="mixed")
        assert check_compat(m, "excel") == []

    def test_the_report_shape_of_a_backend_hint_is_unchanged(self):
        cf = mo.Variable([-100.0, 50.0, 60.0, 70.0])
        issues = check_compat(mo.IRR(cf), "pandas")
        assert issues == [CompatIssue("IRR", "render", frozenset({"excel"}), "pandas")]
        assert issues[0].qpath is None and issues[0].inlined is None
        assert check_compat(mo.IRR(cf), "excel") == []


def _soffice():
    for candidate in ("soffice", "libreoffice"):
        found = shutil.which(candidate)
        if found:
            return found
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


@pytest.mark.slow
def test_the_workbook_recalculates_to_the_python_values(tmp_path):
    soffice = _soffice()
    if soffice is None:
        pytest.skip("LibreOffice not installed")
    m = _model()
    with m:
        m.s.total = mo.SUM(m.a.price * m.a.volume).set_display_name("total")
        m.s.plus = (mo.SUM(m.a.price * m.a.volume) + m.a.rate).set_display_name("plus")
        m.s.three = mo.SUM(m.a.price * m.a.volume * m.a.cfads).set_display_name("three")
    path = tmp_path / "sumproduct.xlsx"
    m.to_excel(path)
    subprocess.run(
        [soffice, f"-env:UserInstallation={(tmp_path / 'prof').as_uri()}",
         "--headless", "--convert-to", "xlsx", "--outdir",
         str(tmp_path / "calc"), str(path)],
        check=True, capture_output=True,
    )
    s = _rows(load_workbook(tmp_path / "calc" / "sumproduct.xlsx", data_only=True)["S"])
    for name in ("total", "plus", "three"):
        assert round(s[name][0], 6) == round(getattr(m.s, name).value, 6), name
