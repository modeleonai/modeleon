# SPDX-License-Identifier: Apache-2.0
"""``orient='columns'`` on a sheet with lead columns — the board starts
where the periods would.

A board's columns are instances, run across the way periods are. On a
sheet whose view declares lead columns (a unit column, a projected
field, a constants column) the board used to start at column B
regardless, so the unit text of its first field row was written over
the first instance's value, and the board's formulas divided by that
text. The board now starts past the lead columns, in the first period
column, and the lead cells keep their own columns.
"""

import shutil
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook

import modeleon as mo

LEAD = {"unit": "Unit", "fields": [("description", "Note")], "constants": "Constant"}


def _model():
    m = mo.Model("m", display_name="M", default_grain="year", default_start="2026",
                 default_periods=3)
    with m:
        m.default_excel_view = mo.ExcelView(meta=LEAD)
        m.inputs = mo.MultiVariable("Inputs", excel_props={"tab": True})
        with m.inputs as i:
            i.ebitda = mo.Variable(100.0, display_name="LTM EBITDA", unit="USD m")
            i.debt = mo.Variable(400.0, display_name="Net debt", unit="USD m")
        m.sens = mo.MultiVariable("Sensitivity", excel_props={"tab": True})
        with m.sens as s:
            s.revenue = mo.Variable([10.0, 11.0, 12.0], display_name="Revenue", unit="USD m",
                                    description="Along the periods")
            s.exit_multiple = mo.Variable(8.0, display_name="Exit multiple", unit="x")
            s.entry = mo.MultiVariable("Sponsor equity by entry multiple",
                                       excel_props={"orient": "columns"})
            with s.entry as e:
                for mult in (7.0, 8.0, 9.0):
                    col = mo.MultiVariable(f"Entry {mult:.1f}x")
                    setattr(e, f"e{int(mult)}", col)
                    with col as c:
                        c.multiple = mo.Variable(mult, display_name="Entry multiple", unit="x",
                                                 description="EV / LTM EBITDA")
                        c.equity = mo.Variable(m.inputs.ebitda * c.multiple - m.inputs.debt,
                                               display_name="Sponsor equity")
                e.base = mo.Variable(5.0, display_name="Base case", unit="x")
    return m


def _sheet(m, tmp_path, name):
    path = tmp_path / "board.xlsx"
    m.to_excel(str(path))
    return load_workbook(str(path))[name]


def _row(ws, r, cols="ABCDEFG"):
    return [ws[f"{c}{r}"].value for c in cols]


class TestABoardUnderLeadColumns:
    def test_the_board_starts_in_the_first_period_column(self, tmp_path):
        ws = _sheet(_model(), tmp_path, "Sensitivity")
        # The sheet's own columns: label, Unit, Note, Constant, then the years.
        assert _row(ws, 1)[1:4] == ["Unit", "Note", "Constant"]
        assert ws["E1"].value is not None
        assert _row(ws, 2) == ["Revenue", "USD m", "Along the periods", None, 10, 11, 12]
        # The instance captions stand where the years do, not over the lead columns.
        assert _row(ws, 6) == [None, None, None, None, "Entry 7.0x", "Entry 8.0x", "Entry 9.0x"]

    def test_the_lead_cells_and_the_first_instance_keep_their_own_cells(self, tmp_path):
        ws = _sheet(_model(), tmp_path, "Sensitivity")
        # The defect: 'x' stood in B7 over the first instance's 7.0.
        assert _row(ws, 7) == ["Entry multiple", "x", "EV / LTM EBITDA", None, 7.0, 8.0, 9.0]
        assert _row(ws, 8)[:2] == ["Sponsor equity", "USD m"]

    def test_every_board_formula_reads_its_own_column(self, tmp_path):
        ws = _sheet(_model(), tmp_path, "Sensitivity")
        for col in "EFG":
            assert ws[f"{col}8"].value == f"=Inputs!D1 * {col}7 - Inputs!D2"
        assert _row(ws, 8)[2:4] == [None, None]

    def test_a_direct_scalar_under_the_board_stands_in_the_constants_column(self, tmp_path):
        ws = _sheet(_model(), tmp_path, "Sensitivity")
        # Like every other scalar on the sheet — and its value survives the
        # unit written beside its label.
        assert _row(ws, 9) == ["Base case", "x", None, 5.0, None, None, None]
        assert _row(ws, 3) == ["Exit multiple", "x", None, 8.0, None, None, None]

    def test_no_cell_is_written_twice(self, tmp_path):
        from modeleon.compile.excel.layout import LayoutEngine

        from openpyxl.utils.cell import coordinate_to_tuple

        m = _model()
        engine = LayoutEngine([m])
        addresses = engine.compute_addresses()
        taken: dict = {}
        for vid in engine.sheet_map["Sensitivity"]:
            a = addresses[vid]
            for cell in [a.name, *a.values]:
                if cell:
                    assert cell not in taken, (cell, vid, taken[cell])
                    taken[cell] = vid
            # The Unit and Note columns hold nobody's value cells.
            assert all(coordinate_to_tuple(cell)[1] not in (2, 3) for cell in a.values), vid


class TestABoardOnItsOwnSheet:
    def test_instances_start_past_the_unit_column(self, tmp_path):
        m = mo.Model("m", display_name="M")
        m.default_excel_view = mo.ExcelView(meta={"unit": "Unit"})
        m.board = mo.MultiVariable("Board", excel_props={"tab": True, "orient": "columns"})
        m.board.p1 = mo.MultiVariable("P1")
        m.board.p1.cost = mo.Variable(50.0, display_name="Cost", unit="USD")
        m.board.p2 = mo.MultiVariable("P2")
        m.board.p2.cost = mo.Variable(80.0, display_name="Cost", unit="USD")
        ws = _sheet(m, tmp_path, "Board")
        assert _row(ws, 1, "ABCD") == [None, None, "P1", "P2"]
        assert _row(ws, 2, "ABCD") == ["Cost", "USD", 50.0, 80.0]

    def test_scalar_children_across_start_past_the_lead_columns_too(self, tmp_path):
        m = mo.Model("m", display_name="M")
        m.default_excel_view = mo.ExcelView(meta={"unit": "Unit"})
        m.kpis = mo.MultiVariable("KPIs", excel_props={"tab": True, "orient": "columns"})
        m.kpis.irr = mo.Variable(0.2, display_name="IRR")
        m.kpis.moic = mo.Variable(2.5, display_name="MOIC")
        ws = _sheet(m, tmp_path, "KPIs")
        names = next(r for r in range(1, 5) if ws[f"C{r}"].value == "IRR")
        assert _row(ws, names, "BCD") == [None, "IRR", "MOIC"]
        assert _row(ws, names + 1, "BCD") == [None, 0.2, 2.5]

    def test_the_house_default_keeps_the_board_at_column_b(self, tmp_path):
        # No lead columns declared: byte-identical to before.
        m = mo.Model("m", display_name="M")
        m.board = mo.MultiVariable("Board", excel_props={"tab": True, "orient": "columns"})
        m.board.p1 = mo.MultiVariable("P1")
        m.board.p1.cost = mo.Variable(50.0, display_name="Cost")
        ws = _sheet(m, tmp_path, "Board")
        assert _row(ws, 1, "AB") == [None, "P1"]
        assert _row(ws, 2, "AB") == ["Cost", 50.0]


def _soffice():
    for candidate in ("soffice", "libreoffice"):
        found = shutil.which(candidate)
        if found:
            return found
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


@pytest.mark.slow
def test_the_board_recalculates_to_the_python_values(tmp_path):
    soffice = _soffice()
    if soffice is None:
        pytest.skip("LibreOffice not installed")
    m = _model()
    path = tmp_path / "board.xlsx"
    m.to_excel(path)
    subprocess.run(
        [soffice, f"-env:UserInstallation={(tmp_path / 'prof').as_uri()}",
         "--headless", "--convert-to", "xlsx", "--outdir",
         str(tmp_path / "calc"), str(path)],
        check=True, capture_output=True,
    )
    ws = load_workbook(tmp_path / "calc" / "board.xlsx", data_only=True)["Sensitivity"]
    assert _row(ws, 8)[4:] == [300.0, 400.0, 500.0]
    assert [m.sens.entry.e7.equity.value, m.sens.entry.e8.equity.value,
            m.sens.entry.e9.equity.value] == [300.0, 400.0, 500.0]
