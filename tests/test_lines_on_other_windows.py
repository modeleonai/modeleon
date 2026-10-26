# SPDX-License-Identifier: Apache-2.0
"""A line on another time window is read by DATE in every function.

``a + b`` aligns two lines on other windows by date: each is continued
by its own ``extend=`` rule — ``mo.zero()`` for a flow, ``mo.hold()``
for a rate — a line without one is refused, two grains are refused, and
the result runs on the combined window. Every function that combines two
or more lines one period at a time reads them the same way — ``IF``,
``AND`` / ``OR``, ``CHOOSE``, ``MOD``, the comparisons, ``CONCAT``, the
dates, the lookups, the statistics, a recurrence — and the workbook reads
such a line by date too: a formula on the model's window reading a line
on a sub-container's window inside it takes the cell of the same date,
and for a period the line has no column for, its rule (``0`` for a flow,
the first or last cell for a rate held). Where Python read a line by
position — a formula over ONE other window, which no second window
meets — the workbook reads it by position too, so the two always agree.
An unnamed formula (``rev * 1``, ``fee > 15``) runs on the window of the
lines it reads: it meets other windows by date and, holding no rule,
must cover the combined window — one that does not is refused, naming
the line to make of it. A one-value line on a window of more periods is
a constant, read in every period; a line along keys runs on no window;
a plain line beside a tracked one is read on the tracked line's window;
windows are ordered by date, never as text. A line of another length
with no window at all is refused, as ``a + b`` refuses it, while a list
typed in keeps its positional reading (its last value standing in, as
the cell writes it). A loop container reads such lines the same way:
each period of the loop reads the line's value of its date, and the
workbook reads the cell of that date. A formula whose operands met by
date on a window of its own, set on a sheet of another window, is
refused by the export as the bare sum is.

Lines along keys combine key by key, as ``a + b`` combines them: the
same keys in another order are read in the first line's order (the
workbook names each key's cell), other keys and a period series are
refused. A list of lines given to ``SUM`` (and its kin) as one argument
is one number over scalar lines and refused over series: per period is
``mo.SUM(a, b)``, the total ``mo.SUM(a) + mo.SUM(b)``.
"""

import shutil
import subprocess
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

import modeleon as mo

pytestmark = pytest.mark.filterwarnings("ignore::UserWarning")


def _rows(ws):
    out = {}
    for row in ws.iter_rows():
        label = row[0].value
        if label:
            out[str(label)] = [c.value for c in row[1:]]
    return out


def _lines():
    """A model from January (X, four months) with sub-containers on
    windows inside its own: T from March (two months: flows continued by
    zero, rates and dates held, one line without a rule), U from January
    ending early, O of one period (March: a flow, a rate held, one line
    without a rule), K on the model's window (a block to look up in), I
    of scalars."""
    m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=4)
    with m:
        m.x = mo.MultiVariable("X")
        m.x.rev = mo.Variable([1.0, 2.0, 3.0, 4.0], display_name="rev")
        m.x.flag = mo.Variable([0.0, 1.0, 0.0, 1.0], display_name="flag")
        m.x.idx = mo.Variable([1.0, 2.0, 1.0, 2.0], display_name="idx")
        m.x.big = mo.Variable([7.0, 9.0, 11.0, 13.0], display_name="big")
        m.x.tx = mo.Variable(["a", "b", "c", "d"], display_name="tx")
        m.x.dt = mo.Variable([date(2026, 1, 31), date(2026, 2, 28), date(2026, 3, 31),
                              date(2026, 4, 30)], display_name="dt")
        m.x.lk = mo.Variable([1.0, 10.0, 1.0, 10.0], display_name="lk")
        m.t = mo.MultiVariable("T", default_grain="month", default_start="2026-03",
                               default_periods=2)
        m.t.fee = mo.Variable([10.0, 20.0], display_name="fee", extend=mo.zero())
        m.t.rate = mo.Variable([0.5, 0.25], display_name="rate", extend=mo.hold())
        m.t.nofill = mo.Variable([7.0, 8.0], display_name="nofill")
        m.t.td = mo.Variable([date(2026, 6, 30), date(2026, 7, 31)], display_name="td",
                             extend=mo.hold())
        m.t.tw = mo.Variable(["p", "q"], display_name="tw", extend=mo.hold())
        m.t.kk = mo.Variable([1.0, 2.0], display_name="kk", extend=mo.hold())
        m.t.rw = mo.Variable([2.0, 1.0], display_name="rw", extend=mo.hold())
        m.t.cc = mo.Variable([3.0, 4.0], display_name="cc", extend=mo.hold())
        m.u = mo.MultiVariable("U", default_grain="month", default_start="2026-01",
                               default_periods=2)
        m.u.early = mo.Variable([100.0, 200.0], display_name="early", extend=mo.zero())
        m.u.cap = mo.Variable([9.0, 8.0], display_name="cap", extend=mo.hold())
        m.k = mo.MultiVariable("K")
        m.k.yrs = mo.Variable([1.0, 2.0, 3.0, 4.0], display_name="yrs")
        m.k.vals = mo.Variable([10.0, 20.0, 30.0, 40.0], display_name="vals")
        m.k.dec = mo.Variable([1.23456, 2.34567, 3.45678, 4.56789], display_name="dec")
        m.i = mo.MultiVariable("I")
        m.i.k1 = mo.Variable(1.0, display_name="k1")
        m.i.k2 = mo.Variable(3.0, display_name="k2")
        m.i.v1 = mo.Variable(10.0, display_name="v1")
        m.i.v2 = mo.Variable(30.0, display_name="v2")
        m.i.hol = mo.Variable(date(2026, 5, 1), display_name="hol")
        m.o = mo.MultiVariable("O", default_grain="month", default_start="2026-03",
                               default_periods=1)
        m.o.once = mo.Variable([50.0], display_name="once", extend=mo.zero())
        m.o.level = mo.Variable([2.0], display_name="level", extend=mo.hold())
        m.o.bare = mo.Variable([7.0], display_name="bare")
    return m


#: Each call on X over lines of T, U and the model's own window, with the
#: values read by date: fee (zero) [0, 0, 10, 20], rate (hold) [0.5, 0.5,
#: 0.5, 0.25], td (hold) [Jun 30, Jun 30, Jun 30, Jul 31], tw (hold) [p, p,
#: p, q], kk [1, 1, 1, 2], rw [2, 2, 2, 1], cc [3, 3, 3, 4], early (zero)
#: [100, 200, 0, 0], cap (hold) [9, 8, 8, 8].
_CALLS = {
    "if_fee": (lambda m: mo.IF(m.x.rev > 2, m.x.rev, m.t.fee), [0.0, 0.0, 3.0, 4.0]),
    "if_rate": (lambda m: mo.IF(m.x.rev > 2, m.t.rate, m.u.cap), [9.0, 8.0, 0.5, 0.25]),
    "if_flag": (lambda m: mo.IF(m.x.flag, m.t.fee, m.u.early), [100.0, 0.0, 0.0, 20.0]),
    "and_": (lambda m: mo.AND(m.x.rev, m.t.fee), [False, False, True, True]),
    "or_": (lambda m: mo.OR(m.x.flag, m.u.early), [True, True, False, True]),
    "ch_series": (lambda m: mo.CHOOSE(m.x.idx, m.t.fee, m.u.cap), [0.0, 8.0, 10.0, 8.0]),
    "ch_scalar": (lambda m: mo.CHOOSE(2, m.x.rev, m.t.fee), [0.0, 0.0, 10.0, 20.0]),
    "md": (lambda m: mo.MOD(m.x.big, m.u.cap), [7.0, 1.0, 3.0, 5.0]),
    "gt": (lambda m: m.x.rev > m.t.fee, [True, True, False, False]),
    "ge": (lambda m: m.u.cap >= m.x.big, [True, False, False, False]),
    "cc_": (lambda m: mo.CONCAT(m.x.tx, m.t.tw), ["ap", "bp", "cp", "dq"]),
    "dd": (lambda m: mo.DATE(2026, m.x.rev, m.t.fee),
           [date(2025, 12, 31), date(2026, 1, 31), date(2026, 3, 10), date(2026, 4, 20)]),
    "d360": (lambda m: mo.DAYS360(m.x.dt, m.t.td), [150, 120, 90, 90]),
    "yf": (lambda m: mo.YEARFRAC(m.x.dt, m.t.td, 3),
           [150 / 365, 122 / 365, 91 / 365, 92 / 365]),
    "dy": (lambda m: mo.DAYS(m.t.td, m.x.dt), [150, 122, 91, 92]),
    "nwd": (lambda m: mo.NETWORKDAYS(m.x.dt, m.t.td, m.i.hol), None),
    "xl": (lambda m: mo.XLOOKUP(m.x.rev, [m.i.k1, m.i.k2], [m.i.v1, m.i.v2], m.t.fee),
           [10.0, 0.0, 30.0, 20.0]),
    "hl": (lambda m: mo.HLOOKUP(m.x.rev, [m.k.yrs, m.k.vals], m.t.rw, False),
           [10.0, 20.0, 30.0, 4.0]),
    "ix": (lambda m: mo.INDEX([m.k.yrs, m.k.vals], m.x.idx, m.t.cc), [3.0, 30.0, 3.0, 40.0]),
    "vl": (lambda m: mo.VLOOKUP(m.x.lk, [m.k.yrs, m.k.vals], m.t.cc, False),
           [3.0, 30.0, 3.0, 40.0]),
    "med": (lambda m: mo.MEDIAN(m.x.rev, m.t.fee, m.u.cap), [1.0, 2.0, 8.0, 8.0]),
    "sdp": (lambda m: mo.STDEV_P(m.x.rev, m.u.early), [49.5, 99.0, 1.5, 2.0]),
    "mdl": (lambda m: mo.MEDIAN([m.x.rev, m.t.fee, m.u.cap]), [1.0, 2.0, 8.0, 8.0]),
    "pct": (lambda m: mo.PERCENTILE_INC([m.x.rev, m.t.fee, m.u.cap], 0.5),
            [1.0, 2.0, 8.0, 8.0]),
    "lg": (lambda m: mo.LARGE([m.x.rev, m.x.idx], m.t.kk), [1.0, 2.0, 3.0, 2.0]),
    "rk": (lambda m: mo.RANK_EQ(m.t.fee, [m.x.rev, m.t.fee]), [2, 2, 1, 1]),
    "rc": (lambda m: mo.recurrence(0.0, "{prev} + {a} + {b}", a=m.x.rev, b=m.t.fee),
           [0.0, 2.0, 15.0, 39.0]),
    "ie": (lambda m: mo.IFERROR(m.x.rev / m.t.fee, 0), [0.0, 0.0, 0.3, 0.2]),
    # O's one period, March: once (zero) [0, 0, 50, 0], level (hold) [2, 2, 2, 2].
    "o_if": (lambda m: mo.IF(m.x.rev > 3, m.x.rev, m.o.once), [0.0, 0.0, 50.0, 4.0]),
    "o_if2": (lambda m: mo.IF(m.x.flag, m.o.level, m.o.once), [0.0, 2.0, 50.0, 2.0]),
    "o_med": (lambda m: mo.MEDIAN(m.x.rev, m.o.once, m.o.level), [1.0, 2.0, 3.0, 2.0]),
    "o_cat": (lambda m: mo.CONCAT(m.x.tx, m.o.once), ["a0", "b0", "c50", "d0"]),
    "o_rnd": (lambda m: mo.ROUND(m.k.dec, m.o.level), [1.23, 2.35, 3.46, 4.57]),
    "o_ed": (lambda m: mo.EDATE(m.x.dt, m.o.level),
             [date(2026, 3, 31), date(2026, 4, 28), date(2026, 5, 31), date(2026, 6, 30)]),
    "o_pmt": (lambda m: mo.PMT(0.05, m.o.level, m.k.vals),
              [mo.PMT(0.05, 2, v).value for v in (10.0, 20.0, 30.0, 40.0)]),
    # A formula over X's lines runs on X's window, meeting others by date.
    "w_add": (lambda m: (m.x.rev * 1) + m.t.fee, [1.0, 2.0, 13.0, 24.0]),
    "w_if": (lambda m: mo.IF(m.x.rev * 1 > 2, m.x.rev * 1, m.t.fee), [0.0, 0.0, 3.0, 4.0]),
    "w_once": (lambda m: m.x.rev * 1 + m.o.once, [1.0, 2.0, 53.0, 4.0]),
    "w_max": (lambda m: mo.MAX(m.x.rev * 1, m.t.fee, m.u.cap), [9.0, 8.0, 10.0, 20.0]),
    "w_lag": (lambda m: mo.lag(m.x.rev, 1) + m.t.fee, [0.0, 1.0, 12.0, 23.0]),
    "w_cmp": (lambda m: mo.IF(m.x.rev > 2, m.t.fee, m.u.cap), [9.0, 8.0, 10.0, 20.0]),
}


def _model():
    """``_lines()`` with every call of ``_CALLS`` a row of X."""
    m = _lines()
    with m:
        for name, (make, _) in _CALLS.items():
            setattr(m.x, name, make(m).set_display_name(name))
    return m


def _networkdays():
    # The working days from each period's date to the held one, the
    # holiday on 1 May, each counted on its own.
    pairs = [(date(2026, 1, 31), date(2026, 6, 30)), (date(2026, 2, 28), date(2026, 6, 30)),
             (date(2026, 3, 31), date(2026, 6, 30)), (date(2026, 4, 30), date(2026, 7, 31))]
    return [mo.NETWORKDAYS(s, e, date(2026, 5, 1)).value for s, e in pairs]


# ══════════════════════════════════════════════════════════════════════
# Python reads by date
# ══════════════════════════════════════════════════════════════════════


class TestPythonReadsByDate:
    @pytest.mark.parametrize("name", list(_CALLS))
    def test_each_call_reads_each_line_on_the_date_s_cell(self, name):
        m = _lines()
        make, expected = _CALLS[name]
        got = make(m)
        if expected is None:
            expected = _networkdays()
        if isinstance(expected[0], float):
            assert got.value == pytest.approx(expected)
        else:
            assert got.value == expected
        if name == "ie":
            # The windows met in ``rev / fee``; IFERROR over that one
            # series and a number places nothing, as ``(rev / fee) * 2``.
            assert got.time is None
        else:
            # On the combined window, as an operator puts its result there.
            assert got.time is not None and got.time.start == "2026-01", name
        assert len(got.value) == 4

    def test_the_call_reads_what_the_operator_reads(self):
        m = _lines()
        rev, fee, rate = m.x.rev, m.t.fee, m.t.rate
        fee_on_x = (rev + fee - rev).value          # fee read on X's dates
        assert fee_on_x == [0.0, 0.0, 10.0, 20.0]
        assert (rev > fee).value == [a > b for a, b in zip(rev.value, fee_on_x)]
        assert (rev > fee).time == (rev + fee).time
        assert mo.MOD(rev, rate).value == (rev % rate).value
        assert mo.CONCAT(rev, fee).value == ["10", "20", "310", "420"]

    def test_three_windows_each_continued_by_its_own_rule(self):
        m = _lines()
        got = mo.IF(m.x.flag, m.t.fee, m.u.cap)
        assert got.value == [9.0, 0.0, 8.0, 20.0]
        assert mo.AND(m.x.rev, m.t.fee, m.u.early).value == [False, False, False, False]
        assert mo.MEDIAN(m.x.rev, m.t.fee, m.u.early).value == [1.0, 2.0, 3.0, 4.0]

    def test_equal_windows_are_read_by_position_as_before(self):
        m = _lines()
        x = m.x
        assert mo.IF(x.flag, x.rev, x.big).value == [7.0, 2.0, 11.0, 4.0]
        assert mo.AND(x.rev, x.flag).value == [False, True, False, True]
        assert mo.CHOOSE(x.idx, x.rev, x.big).value == [1.0, 9.0, 3.0, 13.0]
        assert (x.rev > x.idx).value == [False, False, True, True]
        assert mo.MOD(x.big, x.rev).value == [0.0, 1.0, 2.0, 1.0]
        # Plain lists keep their positional reading.
        assert mo.IF([1, 0, 1], [1.0, 2.0, 3.0], [9.0, 9.0, 9.0]).value == [1.0, 9.0, 3.0]
        assert mo.AND([1, 1, 0], [1]).value == [True, True, False]
        assert mo.AND([1, 1, 0], [1, 0]).value == [True, False, False]
        # Nothing aligned, nothing placed: the line takes its container's
        # window, as ``x.rev + x.big`` does.
        assert mo.IF(x.flag, x.rev, x.big).time is None
        assert (x.rev + x.big).time is None

    def test_a_formula_runs_on_the_window_of_the_lines_it_reads(self):
        # ``revz > 2`` holds X's four months: beside a line from November
        # the combined window runs six months, and a formula has no rule
        # to continue it — refused, naming it, where a positional pairing
        # would have read X's January in November.
        m = _lines()
        with m:
            m.t2 = mo.MultiVariable("T2", default_grain="month", default_start="2025-11",
                                    default_periods=2)
            m.t2.old = mo.Variable([5.0, 6.0], display_name="old", extend=mo.zero())
            m.x.revz = mo.Variable([1.0, 2.0, 3.0, 4.0], display_name="revz",
                                   extend=mo.zero())
        for formula, call in (
                ("revz > 2", lambda: mo.IF(m.x.revz > 2, m.x.revz, m.t2.old)),
                ("revz > 2", lambda: mo.AND(m.x.revz > 2, m.x.revz, m.t2.old)),
                ("idx \\* 1", lambda: mo.CHOOSE(m.x.idx * 1, m.x.revz, m.t2.old)),
                ("revz > 2", lambda: mo.CONCAT(m.x.revz > 2, m.x.revz, m.t2.old)),
                ("big \\* 1", lambda: mo.DATE(m.x.big * 1, m.x.revz, m.t2.old)),
                ("revz > 2", lambda: mo.MEDIAN(m.x.revz > 2, m.x.revz, m.t2.old)),
                ("revz \\* 1", lambda: (m.x.revz * 1) + m.t2.old)):
            with pytest.raises(ValueError, match=f"the line {formula} does not cover the "
                                                 f"combined window") as refused:
                call()
            assert "the periods of the lines it reads" in str(refused.value)
            assert "mo.Variable(" in str(refused.value)
        # The line itself, continued by its rule, is read by date.
        assert (m.x.revz + m.t2.old).value == [5.0, 6.0, 1.0, 2.0, 3.0, 4.0]


_NO_RULE = {
    "IF": lambda m: mo.IF(m.x.rev > 2, m.x.rev, m.t.nofill),
    "AND": lambda m: mo.AND(m.x.rev, m.t.nofill),
    "OR": lambda m: mo.OR(m.t.nofill, m.x.rev),
    "CHOOSE": lambda m: mo.CHOOSE(m.x.idx, m.t.nofill, m.t.fee),
    "CHOOSE_scalar": lambda m: mo.CHOOSE(1, m.t.nofill, m.x.rev),
    "MOD": lambda m: mo.MOD(m.x.rev, m.t.nofill),
    "compare": lambda m: m.x.rev > m.t.nofill,
    "CONCAT": lambda m: mo.CONCAT(m.x.tx, m.t.nofill),
    "DATE": lambda m: mo.DATE(2026, m.x.rev, m.t.nofill),
    "DAYS": lambda m: mo.DAYS(m.t.nofill, m.x.rev),
    "XLOOKUP": lambda m: mo.XLOOKUP(m.x.rev, [m.i.k1, m.i.k2], [m.i.v1, m.i.v2], m.t.nofill),
    "INDEX": lambda m: mo.INDEX([m.k.yrs, m.k.vals], m.x.idx, m.t.nofill),
    "MEDIAN": lambda m: mo.MEDIAN(m.x.rev, m.t.nofill),
    "MEDIAN_list": lambda m: mo.MEDIAN([m.x.rev, m.t.nofill]),
    "LARGE": lambda m: mo.LARGE([m.x.rev, m.x.idx], m.t.nofill),
    "recurrence": lambda m: mo.recurrence(0.0, "{prev} + {a}", a=m.t.nofill, b=m.x.rev),
}


class TestALineWithoutARuleIsRefused:
    @pytest.mark.parametrize("name", list(_NO_RULE))
    def test_as_a_plus_b_refuses_it(self, name):
        m = _lines()
        with pytest.raises(ValueError, match="'nofill' does not cover the combined window"):
            m.x.rev + m.t.nofill
        with pytest.raises(ValueError, match="'nofill' does not cover the combined window") \
                as refused:
            _NO_RULE[name](m)
        assert "extend=mo.zero()" in str(refused.value)
        assert "extend=mo.hold()" in str(refused.value)

    def test_a_formula_over_one_window_meets_another_by_date(self):
        # ``rev * 1`` runs on X's window, the one of the line it reads:
        # beside fee it is read by date, operators and functions alike,
        # where a positional reading paired January with March.
        m = _lines()
        wide = m.x.rev * 1
        assert (wide + m.t.fee).value == [1.0, 2.0, 13.0, 24.0]
        for call, expected in (
                (lambda: mo.IF(wide > 2, wide, m.t.fee), [0.0, 0.0, 3.0, 4.0]),
                (lambda: mo.AND(wide, m.t.fee), [False, False, True, True]),
                (lambda: mo.OR(m.t.fee, wide), [True, True, True, True]),
                (lambda: mo.CHOOSE(1, wide, m.t.fee), [1.0, 2.0, 3.0, 4.0]),
                (lambda: mo.CONCAT(wide, m.t.tw), ["1p", "2p", "3p", "4q"]),
                (lambda: mo.DATE(2026, wide, m.t.fee),
                 [date(2025, 12, 31), date(2026, 1, 31), date(2026, 3, 10),
                  date(2026, 4, 20)]),
                (lambda: mo.DAYS360(mo.EDATE(m.x.dt, 0), m.t.td), [150, 120, 90, 90]),
                (lambda: mo.DAYS(m.t.td, mo.EDATE(m.x.dt, 0)), [150, 122, 91, 92]),
                (lambda: mo.recurrence(0.0, "{prev} + {a} + {b}", a=wide, b=m.t.fee),
                 [0.0, 2.0, 15.0, 39.0]),
                (lambda: mo.recurrence(0.0, "{prev} + {b} + {a}", b=m.t.fee, a=wide),
                 [0.0, 2.0, 15.0, 39.0]),
                # A comparison is such a formula: the commonest condition.
                (lambda: mo.IF(m.x.rev > 2, m.t.fee, 0), [0.0, 0.0, 10.0, 20.0]),
                (lambda: mo.IF(m.x.rev > 2, m.t.fee, m.u.cap), [9.0, 8.0, 10.0, 20.0])):
            got = call()
            assert got.value == expected
            assert got.time.start == "2026-01"

    def test_a_formula_the_combined_window_outruns_is_refused(self):
        # Ops runs from July: ``cfads > 0`` holds July's periods, which no
        # rule continues back to January — refused, where a positional
        # pairing let July's cash decide January.
        m = _lines()
        with m:
            m.ops = mo.MultiVariable("Ops", default_grain="month", default_start="2026-07",
                                     default_periods=4)
            m.ops.cfads = mo.Variable([-5.0, -1.0, 2.0, 3.0], display_name="cfads")
            m.b = mo.MultiVariable("B", default_grain="month", default_start="2025-11",
                                   default_periods=2)
            m.b.bz = mo.Variable([1000.0, 2000.0], display_name="bz", extend=mo.zero())
            m.x.revz = mo.Variable([1.0, 2.0, 3.0, 4.0], display_name="revz",
                                   extend=mo.zero())
        # Every line here says what continues it; the formula cannot.
        cf, rev, fee = m.ops.cfads, m.x.revz, m.t.fee
        for formula, call in (
                ("cfads > 0", lambda: mo.IF(cf > 0, rev, fee)),
                ("cfads \\* 1", lambda: mo.MAX(cf * 1, rev, fee)),
                ("cfads \\* 1", lambda: mo.CHOOSE(1, cf * 1, rev)),
                ("cfads \\* 1", lambda: mo.SUM(cf * 1, rev, fee)),
                ("cfads \\* 1", lambda: mo.MEDIAN(cf * 1, rev, fee)),
                ("cfads \\* 1",
                 lambda: mo.recurrence(0.0, "{prev} + {a} + {b}", a=cf * 1, b=fee)),
                ("cfads / 3", lambda: mo.ROUND(cf / 3, rev)),
                ("cfads \\* 1", lambda: mo.PMT(0.05, rev, cf * 1)),
                ("cfads \\* 1", lambda: (cf * 1) + rev),
                ("cfads > 0", lambda: (cf > 0) + rev),
                # The combined window may start before the sheet's.
                ("revz > 3", lambda: mo.IF(rev > 3, m.b.bz, fee))):
            with pytest.raises(ValueError, match=f"the line {formula} does not cover"):
                call()
        with pytest.raises(ValueError, match="'cfads' does not cover"):
            cf + rev

    def test_the_formula_made_a_line_with_a_rule_is_read_by_date(self):
        m = _lines()
        with pytest.raises(ValueError, match="m.t.<name> = mo.Variable") as refused:
            mo.IF(m.t.fee > 15, m.x.rev, 0)
        assert "the container 'fee' runs on" in str(refused.value)
        with m:
            m.t.big = mo.Variable(m.t.fee > 15, extend=mo.zero())
        assert mo.IF(m.t.big, m.x.rev, 0).value == [0, 0, 0, 4.0]
        assert mo.OR(m.x.rev > 3, m.t.big).value == [False, False, False, True]

    def test_a_line_with_no_window_of_another_length_is_refused_by_name(self):
        m = _lines()
        pair = mo.Variable([1.0, 0.0], display_name="pair")
        for call in (lambda: mo.AND(m.x.rev, pair), lambda: mo.IF(m.x.rev > 2, pair, 0),
                     lambda: mo.CONCAT(m.x.tx, pair)):
            with pytest.raises(ValueError, match="'pair' cannot be paired with the call's "
                                                 "other series period by period "
                                                 r"\(different lengths: 2 vs 4\)"):
                call()

    def test_one_condition_chooses_a_branch_whole(self):
        # Nothing is paired: the branch comes back as it is.
        a, b = mo.Variable([1.0, 2.0, 3.0]), mo.Variable([5.0, 6.0])
        assert mo.IF(True, a, b).value == [1.0, 2.0, 3.0]
        assert mo.IF(False, a, b).value == [5.0, 6.0]

    def test_a_list_typed_in_keeps_its_positional_reading(self, tmp_path):
        # A list typed in has no cells: its last value stands in the
        # periods it lacks, as the cell writes it.
        m = _lines()
        assert mo.AND(m.x.rev, [1, 0]).value == [True, False, False, False]
        assert mo.CONCAT(m.x.tx, ["x", "y"]).value == ["ax", "by", "cy", "dy"]
        assert mo.DATE(2026, m.x.rev, [10, 20]).value[3] == date(2026, 4, 20)
        # IF too, where it raised IndexError past the list's end.
        with m:
            m.x.ty = mo.IF(m.x.rev > 9, m.x.rev, [7.0, 8.0]).set_display_name("ty")
        assert m.x.ty.value == [7.0, 8.0, 8.0, 8.0]
        m.to_excel(tmp_path / "typed.xlsx")
        assert _rows(load_workbook(tmp_path / "typed.xlsx")["X"])["ty"] == [
            "=IF(B2 > 9, B2, 7.0)", "=IF(C2 > 9, C2, 8.0)", "=IF(D2 > 9, D2, 8.0)",
            "=IF(E2 > 9, E2, 8.0)"]

    def test_a_one_period_line_without_a_rule_is_refused(self):
        m = _lines()
        for call in (lambda: m.x.rev + m.o.bare,
                     lambda: mo.IF(m.x.rev > 2, m.x.rev, m.o.bare),
                     lambda: mo.MEDIAN(m.x.rev, m.o.bare),
                     lambda: mo.ROUND(m.k.dec, m.o.bare)):
            with pytest.raises(ValueError, match="'bare' does not cover the combined window"):
                call()

    def test_two_grains_are_refused(self):
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=3)
        with m:
            m.x = mo.MultiVariable("X")
            m.x.a = mo.Variable([1.0, 2.0, 3.0], display_name="a")
            m.q = mo.MultiVariable("Q", default_grain="quarter", default_start="2026-Q1",
                                   default_periods=2)
            m.q.b = mo.Variable([1.0, 2.0], display_name="b")
        for call in (lambda: mo.IF(m.x.a, m.x.a, m.q.b), lambda: mo.AND(m.x.a, m.q.b),
                     lambda: mo.MOD(m.x.a, m.q.b), lambda: m.x.a > m.q.b,
                     lambda: mo.MEDIAN(m.x.a, m.q.b), lambda: mo.CONCAT(m.x.a, m.q.b)):
            with pytest.raises(ValueError, match="different grains"):
                call()


# ══════════════════════════════════════════════════════════════════════
# What runs on a window, and how windows are ordered
# ══════════════════════════════════════════════════════════════════════


def _windows():
    """``_lines()`` with a constant on X (one value on four months), lines
    along keys on X and on T, and a quarterly line."""
    m = _lines()
    with m:
        m.x.c = mo.Variable([5.0], display_name="c")
        m.x.k = mo.Variable({"Bear": 1.0, "Base": 2.0, "Bull": 3.0}, display_name="k")
        m.t.k2 = mo.Variable({"Bull": 30.0, "Base": 20.0, "Bear": 10.0}, display_name="k2")
        m.q = mo.MultiVariable("Q", default_grain="quarter", default_start="2026-Q1",
                               default_periods=2)
        m.q.h = mo.Variable([7.0, 8.0], display_name="h")
    return m


def _unpadded():
    # September written without its leading zero, a line from October.
    m = mo.Model("m", default_grain="month", default_start="2026-9", default_periods=6)
    with m:
        m.x = mo.MultiVariable("X")
        m.x.rev = mo.Variable([1.0, 2.0, 3.0, 4.0, 5.0, 6.0], display_name="rev")
        m.t = mo.MultiVariable("T", default_grain="month", default_start="2026-10",
                               default_periods=2)
        m.t.fee = mo.Variable([10.0, 20.0], display_name="fee", extend=mo.zero())
    return m


class TestWhatRunsOnAWindow:
    def test_a_one_value_line_is_a_constant(self):
        # One value on a window of four months is a constant (a bare list
        # holds one value or one per period), read in every period — by the operators too,
        # which refused it as a one-month line.
        m = _windows()
        rev, c = m.x.rev, m.x.c
        assert (rev + c).value == [6.0, 7.0, 8.0, 9.0]
        assert mo.MAX(rev, c).value == [5.0, 5.0, 5.0, 5.0]
        assert mo.SUM(rev, c).value == [6.0, 7.0, 8.0, 9.0]
        assert mo.IF(rev > 3, c, m.t.fee).value == [0.0, 0.0, 10.0, 5.0]
        assert (rev * 1 + c).value == [6.0, 7.0, 8.0, 9.0]

    def test_lines_along_keys_run_on_no_window(self):
        # k on X, k2 on T: keys, not periods — combined key by key, where
        # T's window refused them as a line of March.
        m = _windows()
        k, k2 = m.x.k, m.t.k2
        for got, expected in ((mo.MAX(k, k2), [10.0, 20.0, 30.0]),
                              (k + k2, [11.0, 22.0, 33.0]),
                              (mo.IF(k > 1, k2, 0), [0, 20.0, 30.0])):
            assert got.value == expected
            assert got._keys == ["Bear", "Base", "Bull"]

    def test_windows_are_ordered_by_date(self):
        # '2026-9' is before '2026-10', though not as text: September is
        # kept, not dropped.
        m = _unpadded()
        rev, fee = m.x.rev, m.t.fee
        for got, expected in ((rev + fee, [1.0, 12.0, 23.0, 4.0, 5.0, 6.0]),
                              (mo.SUM(rev, fee), [1.0, 12.0, 23.0, 4.0, 5.0, 6.0]),
                              (mo.MAX(rev, fee), [1.0, 10.0, 20.0, 4.0, 5.0, 6.0]),
                              (mo.IF(rev > 2, fee, rev), [1.0, 2.0, 20.0, 0.0, 0.0, 0.0])):
            assert got.value == expected
            assert got.time.start == "2026-9"
        assert mo.SUM(mo.SUM(rev, fee)).value == 51.0

    def test_a_formula_reads_lines_of_one_grain(self):
        # ``h * 1`` runs on quarters: beside months it is refused as the
        # line itself is, where its quarters were added to months.
        m = _windows()
        for call in (lambda: m.x.rev + m.q.h * 1, lambda: mo.MAX(m.x.rev * 1, m.q.h),
                     lambda: mo.IF(m.q.h > 7, m.x.rev, 0)):
            with pytest.raises(ValueError, match="different grains"):
                call()


# ══════════════════════════════════════════════════════════════════════
# The order a model is built in does not move a formula
# ══════════════════════════════════════════════════════════════════════


def _x_and_t():
    """X on the model's window (January, four months); T from March (two
    months, a flow continued by zero)."""
    m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=4)
    with m:
        m.x = mo.MultiVariable("X")
        m.t = mo.MultiVariable("T", default_grain="month", default_start="2026-03",
                               default_periods=2)
        m.t.fee = mo.Variable([10.0, 20.0], display_name="fee", extend=mo.zero())
    return m


def _attached_later():
    """Lines built floating — ``a``, ``b``, and over them the formulas
    ``d = a * 1`` and ``e = d + b`` — then attached to X, beside T and T4
    from March (four months)."""
    a = mo.Variable([1.0, 2.0, 3.0, 4.0], display_name="a")
    b = mo.Variable([5.0, 6.0, 7.0, 8.0], display_name="b")
    d = a * 1
    e = d + b
    m = _x_and_t()
    with m:
        m.x.a = a
        m.x.b = b
        m.t4 = mo.MultiVariable("T4", default_grain="month", default_start="2026-03",
                                default_periods=4)
        m.t4.fee = mo.Variable([10.0, 20.0, 30.0, 40.0], display_name="fee4",
                               extend=mo.zero())
    return m, d, e


def _six():
    return [mo.Variable([float(i + j) for j in range(4)], display_name=f"l{i}")
            for i in range(6)]


def _total(lines):
    total = lines[0]
    for v in lines[1:]:
        total = total + v
    return total


class TestTheBuildOrder:
    def test_a_floating_formula_attached_later_runs_on_its_lines_window(self):
        # ``d`` met ``b`` while ``a`` had no window: once ``a`` is on X,
        # ``d`` runs on January as a formula built after does — read by
        # date beside T, refused beside T4, which it does not cover.
        m, d, e = _attached_later()
        assert e.value == [6.0, 8.0, 10.0, 12.0]
        fresh = (m.x.a * 1) + m.t.fee
        assert (d + m.t.fee).value == fresh.value == [1.0, 2.0, 13.0, 24.0]
        assert mo.MAX(d, m.t.fee).value == [1.0, 2.0, 10.0, 20.0]
        assert (e + m.t.fee).value == [6.0, 8.0, 20.0, 32.0]
        for call in (lambda: d + m.t4.fee, lambda: mo.MAX(d, m.t4.fee),
                     lambda: e + m.t4.fee, lambda: (m.x.a * 1) + m.t4.fee):
            with pytest.raises(ValueError, match="does not cover the combined window"):
                call()

    def test_a_window_moved_after_a_formula_was_placed(self):
        # ``g`` was placed on January through ``rev``; the model moves to
        # March before ``g`` meets a line of March.
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=4)
        with m:
            m.x = mo.MultiVariable("X")
            m.x.rev = mo.Variable([1.0, 2.0, 3.0, 4.0], display_name="rev")
            m.x.one = mo.Variable([1.0, 1.0, 1.0, 1.0], display_name="one")
            g = m.x.rev * 1
            h = g + m.x.one
            m.default_start = "2026-03"
            m.t = mo.MultiVariable("T", default_grain="month", default_start="2026-03",
                                   default_periods=4)
            m.t.fee = mo.Variable([10.0, 20.0, 30.0, 40.0], display_name="fee",
                                  extend=mo.zero())
            expected = [11.0, 22.0, 33.0, 44.0]
            assert (g + m.t.fee).value == ((m.x.rev * 1) + m.t.fee).value == expected
            assert (m.x.rev + m.t.fee).value == expected
            assert (h + m.t.fee).value == [12.0, 23.0, 34.0, 45.0]

    def test_a_chain_answers_the_same_whatever_the_order_it_was_built_in(self):
        # A running total over six lines meets T the same whether it was
        # built before the lines were attached, after, or a step ahead of
        # each attachment.
        expected = [15.0, 21.0, 37.0, 53.0]
        lines = _six()
        before = _total(lines)
        m = _x_and_t()
        with m:
            for i, v in enumerate(lines):
                setattr(m.x, f"l{i}", v)
            after = _total(lines)
            assert (before + m.t.fee).value == (after + m.t.fee).value == expected
        m = _x_and_t()
        with m:
            total = None
            for i, v in enumerate(_six()):
                total = v if total is None else total + v
                setattr(m.x, f"l{i}", v)
            assert (total + m.t.fee).value == expected

    def test_the_workbook_reads_a_formula_attached_later_by_date(self, tmp_path):
        # ``s``, ``mx`` and ``es`` read ``a`` (B2) and ``b`` (B3) in their
        # own column and fee by date — the cells ``fs``, built after the
        # attachment, holds.
        m, d, e = _attached_later()
        with m:
            m.x.s = (d + m.t.fee).set_display_name("s")
            m.x.mx = mo.MAX(d, m.t.fee).set_display_name("mx")
            m.x.es = (e + m.t.fee).set_display_name("es")
            m.x.fs = ((m.x.a * 1) + m.t.fee).set_display_name("fs")
        path = tmp_path / "book.xlsx"
        m.to_excel(path)
        x = _rows(load_workbook(path)["X"])
        assert x["s"] == x["fs"] == ["=B2 * 1 + 0", "=C2 * 1 + 0", "=D2 * 1 + T!B2",
                                     "=E2 * 1 + T!C2"]
        assert x["mx"] == ["=MAX(B2 * 1, 0)", "=MAX(C2 * 1, 0)", "=MAX(D2 * 1, T!B2)",
                           "=MAX(E2 * 1, T!C2)"]
        assert x["es"] == ["=B2 * 1 + B3 + 0", "=C2 * 1 + C3 + 0", "=D2 * 1 + D3 + T!B2",
                           "=E2 * 1 + E3 + T!C2"]

    def test_a_copy_of_floating_lines_is_placed_as_a_fresh_formula(self):
        # A pickle or a copy of floating lines and a formula over them,
        # made after the model is built: once the copies are attached, the
        # copied formula answers what a formula built after the attachment
        # answers.
        import copy
        import pickle
        for clone in (copy.deepcopy, lambda o: pickle.loads(pickle.dumps(o))):
            m = _x_and_t()
            with m:
                m.t4 = mo.MultiVariable("T4", default_grain="month",
                                        default_start="2026-03", default_periods=4)
                m.t4.fee = mo.Variable([10.0, 20.0, 30.0, 40.0], display_name="fee4",
                                       extend=mo.zero())
            a = mo.Variable([1.0, 2.0, 3.0, 4.0], display_name="a")
            b = mo.Variable([5.0, 6.0, 7.0, 8.0], display_name="b")
            d = a * 1
            assert (d + b).value == [6.0, 8.0, 10.0, 12.0]
            a2, b2, d2 = clone((a, b, d))
            m.x.a = a2
            m.x.b = b2
            for call in (lambda: d2 + m.t4.fee, lambda: (m.x.a * 1) + m.t4.fee):
                with pytest.raises(ValueError, match="does not cover the combined window"):
                    call()

    def test_a_refusal_names_the_container_of_the_line_with_a_window(self):
        # ``z`` has no window yet; the formula runs on X's window through
        # ``a``, so the way out is a line of X.
        m, _d, _e = _attached_later()
        z = mo.Variable([0.0, 0.0, 0.0, 0.0], display_name="z")
        for build in (lambda: z + (m.x.a * 1), lambda: (m.x.a * 1) + z):
            with pytest.raises(ValueError, match=r"make it a line of m\.x, the container 'a'"):
                build() + m.t4.fee


# ══════════════════════════════════════════════════════════════════════
# A window changed elsewhere does not walk every formula again
# ══════════════════════════════════════════════════════════════════════


def _walks(monkeypatch, build, n):
    """How many times the placement walk runs while ``build(n)`` builds."""
    from modeleon.core import variable_ops
    count = [0]
    walk = variable_ops._scan

    def counted(*args, **kwargs):
        count[0] += 1
        return walk(*args, **kwargs)

    monkeypatch.setattr(variable_ops, "_scan", counted)
    build(n)
    monkeypatch.setattr(variable_ops, "_scan", walk)
    return count[0]


def _sections(n):
    # A running total over sections, each a container declaring the
    # model's window, built and filled one after another.
    m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=4)
    with m:
        m.assets = mo.MultiVariable("Assets")
    total = None
    for i in range(n):
        s = mo.MultiVariable(f"A{i}", default_grain="month", default_start="2026-01",
                             default_periods=4)
        setattr(m.assets, f"a{i}", s)
        s.cf = mo.Variable([float(i + j) for j in range(4)], extend=mo.zero())
        total = s.cf * 1 if total is None else total + s.cf
    return total


def _models(n):
    # A running total over the lines of one model while another small
    # model is built at every step (its constructor declares a window).
    m = _x_and_t()
    with m:
        for i in range(n):
            setattr(m.x, f"l{i}", mo.Variable([float(i + j) for j in range(4)]))
    total = None
    for i in range(n):
        mo.Model(f"side{i}", default_grain="month", default_start="2026-01",
                 default_periods=4)
        line = getattr(m.x, f"l{i}")
        total = line * 1 if total is None else total + line
    return total


class TestWindowChangesStayLinear:
    @pytest.mark.parametrize("build", [_sections, _models])
    def test_the_walks_grow_with_the_lines_not_their_square(self, monkeypatch, build):
        small = _walks(monkeypatch, build, 40)
        large = _walks(monkeypatch, build, 80)
        assert large <= 2.5 * max(small, 1), (small, large)


def _outcome(thunk):
    try:
        result = thunk()
    except ValueError as error:
        return ("refused", str(error).split(".")[0])
    return (result.value, result.time.start if result.time is not None else None)


def _kept_and_walked(thunk):
    """The answer of a formula's kept placement, then the answer of the
    same formula once every placement is walked again."""
    from modeleon.core import variable_ops
    kept = _outcome(thunk)
    variable_ops._windows_changed()
    return kept, _outcome(thunk)


class TestAKeptPlacementFollowsItsLines:
    """A formula's kept placement answers what a fresh walk answers after
    the lines it reads move — a container detached and set elsewhere, a
    window declared or deleted, a row read ahead becoming its line."""

    def test_a_variance_over_track_slices_after_its_container_moved(self):
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=4,
                     tracks=mo.Tracks("plan", "actual"))
        with m:
            m.x = mo.MultiVariable("X")
            m.x.trk = mo.Variable(plan=[10.0, 20.0, 30.0, 40.0],
                                  actual=[11.0, 22.0, 33.0, 44.0], display_name="trk")
            m.t = mo.MultiVariable("T", default_grain="month", default_start="2026-03",
                                   default_periods=2)
            m.t.fee = mo.Variable([1.0, 2.0], display_name="fee", extend=mo.zero())
        variance = m.x.trk.at(track="actual") - m.x.trk.at(track="plan")
        x = m.remove("x")
        with m:
            m.y = mo.MultiVariable("Y", default_grain="month", default_start="2026-03",
                                   default_periods=4)
            m.y.x = x
        kept, walked = _kept_and_walked(lambda: variance + m.t.fee)
        assert kept == walked == ([2.0, 4.0, 3.0, 4.0], "2026-03")

    def test_a_row_read_ahead_after_its_loop_container_moved(self):
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=4)
        with m:
            m.t = mo.MultiVariable("T", default_grain="month", default_start="2026-03",
                                   default_periods=2)
            m.t.fee = mo.Variable([10.0, 20.0], display_name="fee", extend=mo.zero())
            m.d = mo.MultiVariable("D", loop=True)
            with m.d as d:
                y = mo.lag(d.bal, fill=0.0) + mo.Variable(0.0)
                d.bal = mo.lag(d.bal, fill=1.0) + 1.0
        d = m.remove("d")
        with m:
            m.y = mo.MultiVariable("Y", default_grain="month", default_start="2026-03",
                                   default_periods=4)
            m.y.d = d
        kept, walked = _kept_and_walked(lambda: y + m.t.fee)
        fresh = _outcome(lambda: (mo.lag(m.y.d.bal, fill=0.0) + mo.Variable(0.0)) + m.t.fee)
        assert kept == walked == fresh

    def test_a_loop_container_declaring_its_window_after_a_row_was_read_ahead(self):
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=4)
        with m:
            m.t = mo.MultiVariable("T", default_grain="month", default_start="2026-03",
                                   default_periods=2)
            m.t.fee = mo.Variable([10.0, 20.0], display_name="fee", extend=mo.zero())
            m.d = mo.MultiVariable("D", loop=True)
            y = mo.lag(m.d.bal, fill=0.0) + mo.Variable(0.0)
            m.d.default_grain = "month"
            m.d.default_start = "2026-03"
            m.d.default_periods = 4
            m.d.bal = mo.lag(m.d.bal, fill=1.0) + 1.0
        kept, walked = _kept_and_walked(lambda: y + m.t.fee)
        fresh = _outcome(lambda: (mo.lag(m.d.bal, fill=0.0) + mo.Variable(0.0)) + m.t.fee)
        assert kept == walked == fresh

    def test_a_container_replaced_after_a_line_became_a_plain_number(self):
        # A notebook re-run: a line overwritten with a number, then the
        # cell that builds its container runs again.
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=4)
        with m:
            m.inputs = mo.MultiVariable("Inputs")
            m.inputs.rate = mo.Variable(0.05, display_name="rate")
        m.inputs.rate = 0.05
        old = m.inputs
        m.inputs = mo.MultiVariable("Inputs")
        assert old._parent is None and m.inputs is not old
        m.remove("inputs")
        assert "inputs" not in m._component_names

    def test_a_constant_moved_into_a_one_period_container(self):
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=4)
        with m:
            m.x = mo.MultiVariable("X")
            m.x.a = mo.Variable([1.0, 2.0, 3.0, 4.0], display_name="a")
            m.x.c = mo.Variable([5.0], display_name="c")
            m.t = mo.MultiVariable("T", default_grain="month", default_start="2026-03",
                                   default_periods=2)
            m.t.fee = mo.Variable([10.0, 20.0], display_name="fee", extend=mo.zero())
        d = m.x.c + m.x.a
        c = m.x.remove("c")
        with m:
            m.one = mo.MultiVariable("One", default_grain="month", default_start="2026-06",
                                     default_periods=1)
            m.one.c = c
        kept, walked = _kept_and_walked(lambda: d + m.t.fee)
        assert kept == walked

    @pytest.mark.parametrize("target_first", [True, False])
    def test_a_loop_class_instance_built_floating(self, target_first):
        class Debt(mo.MultiVariableClass, loop=True):
            def compute(self, rev=None):
                self._y = mo.lag(self.bal, fill=0.0) + rev
                self.bal = mo.lag(self.bal, fill=1.0) + 1.0

        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=4)
        with m:
            m.x = mo.MultiVariable("X")
            m.x.rev = mo.Variable([1.0, 2.0, 3.0, 4.0], display_name="rev")
            m.t = mo.MultiVariable("T", default_grain="month", default_start="2026-03",
                                   default_periods=2)
            m.t.fee = mo.Variable([10.0, 20.0], display_name="fee", extend=mo.zero())
            if target_first:
                m.y = mo.MultiVariable("Y", default_grain="month", default_start="2026-03",
                                       default_periods=4)
            inst = Debt(rev=m.x.rev)
            if not target_first:
                m.y = mo.MultiVariable("Y", default_grain="month", default_start="2026-03",
                                       default_periods=4)
            m.y.debt = inst
        kept, walked = _kept_and_walked(lambda: inst._y + m.t.fee)
        assert kept == walked

    def test_a_window_slot_deleted(self):
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=4)
        with m:
            m.x = mo.MultiVariable("X", default_grain="month", default_start="2026-03",
                                   default_periods=4)
            m.x.a = mo.Variable([1.0, 2.0, 3.0, 4.0], display_name="a")
            m.t = mo.MultiVariable("T", default_grain="month", default_start="2026-03",
                                   default_periods=2)
            m.t.fee = mo.Variable([10.0, 20.0], display_name="fee", extend=mo.zero())
        d = m.x.a + m.x.a
        del m.x.default_grain
        kept, walked = _kept_and_walked(lambda: d + m.t.fee)
        fresh = _outcome(lambda: (m.x.a + m.x.a) + m.t.fee)
        assert kept == walked == fresh

    @pytest.mark.parametrize("move", ["set", "del"])
    def test_track_slices_of_an_emptied_container_whose_window_changes(self, move):
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=4,
                     tracks=mo.Tracks("plan", "actual"))
        with m:
            if move == "del":
                m.x = mo.MultiVariable("X", default_grain="month", default_start="2026-03",
                                       default_periods=4)
            else:
                m.x = mo.MultiVariable("X")
            m.x.trk = mo.Variable(plan=[10.0, 20.0, 30.0, 40.0],
                                  actual=[11.0, 22.0, 33.0, 44.0], display_name="trk")
            m.t = mo.MultiVariable("T", default_grain="month", default_start="2026-03",
                                   default_periods=2)
            m.t.fee = mo.Variable([1.0, 2.0], display_name="fee", extend=mo.zero())
        actual = m.x.trk.at(track="actual")
        plan = m.x.trk.at(track="plan")
        m.x.remove("trk")
        variance = actual - plan
        if move == "del":
            del m.x.default_grain
        else:
            m.x.default_grain = "month"
            m.x.default_start = "2026-03"
            m.x.default_periods = 4
        kept, walked = _kept_and_walked(lambda: variance + m.t.fee)
        assert kept == walked

    def test_a_time_line_read_again_for_the_container_it_joins(self):
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=4)
        with m:
            m.x = mo.MultiVariable("X")
            m.t = mo.MultiVariable("T", default_grain="month", default_start="2026-03",
                                   default_periods=2)
            m.t.fee = mo.Variable([10.0, 20.0], display_name="fee", extend=mo.zero())
        side = mo.Model("side", default_grain="month", default_start="2026-03",
                        default_periods=4)
        with side:
            i = mo.time.index
            f = i * 1.0 + i
        m.x.i = i
        kept, walked = _kept_and_walked(lambda: f + m.t.fee)
        assert kept == walked == ([2.0, 4.0, 16.0, 28.0], "2026-01")

    def test_a_row_read_ahead_in_a_one_period_loop_container_is_no_constant(self):
        # Its window comes from the container it is read from: one period
        # of May beside four months is refused when the formula is built.
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=4)
        with m:
            m.x = mo.MultiVariable("X")
            m.x.rev = mo.Variable([1.0, 2.0, 3.0, 4.0], display_name="rev")
            m.d = mo.MultiVariable("D", loop=True, default_grain="month",
                                   default_start="2026-05", default_periods=1)
            with pytest.raises(ValueError, match="does not cover the combined window"):
                mo.lag(m.d.bal, fill=0.0) + m.x.rev

    def test_a_floating_loop_container_joining_with_a_row_still_read_ahead(self):
        from modeleon.core import variable_ops

        def build():
            m = mo.Model("m", default_grain="month", default_start="2026-01",
                         default_periods=4)
            with m:
                m.x = mo.MultiVariable("X")
                m.x.rev = mo.Variable([1.0, 2.0, 3.0, 4.0], display_name="rev")
                m.t = mo.MultiVariable("T", default_grain="month", default_start="2026-03",
                                       default_periods=2)
                m.t.fee = mo.Variable([10.0, 20.0], display_name="fee", extend=mo.zero())
                m.y = mo.MultiVariable("Y", default_grain="month", default_start="2026-03",
                                       default_periods=4)
                d = mo.MultiVariable("D", loop=True)
                y = mo.lag(d.bal, fill=0.0) + m.x.rev
            m.y.d = d
            return _outcome(lambda: y + m.t.fee)

        kept = build()
        saved = variable_ops._kept_placement
        variable_ops._kept_placement = lambda var: variable_ops._UNREAD
        try:
            fresh = build()
        finally:
            variable_ops._kept_placement = saved
        assert kept == fresh



def _tracked():
    m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=4,
                 tracks=mo.Tracks("plan", "actual"))
    with m:
        m.x = mo.MultiVariable("X")
        m.x.rev = mo.Variable([1.0, 2.0, 3.0, 4.0], display_name="rev")
        m.x.trk = mo.Variable(plan=[10.0, 20.0, 30.0, 40.0], actual=[11.0, 21.0, 31.0, 41.0],
                              display_name="trk")
        m.t = mo.MultiVariable("T", default_grain="month", default_start="2026-03",
                               default_periods=2)
        m.t.fee = mo.Variable([10.0, 20.0], display_name="fee", extend=mo.zero())
        m.o = mo.MultiVariable("O", default_grain="month", default_start="2026-04",
                               default_periods=1)
        m.o.once = mo.Variable([50.0], display_name="once", extend=mo.zero())
        m.u = mo.MultiVariable("U", default_grain="month", default_start="2025-11",
                               default_periods=2)
        m.u.early = mo.Variable([100.0, 200.0], display_name="early", extend=mo.zero())
    return m


#: Calls over a tracked line and lines on other windows, per track: once
#: (April, zero) [0, 0, 0, 50], fee (March, zero) [0, 0, 10, 20].
_TRACKED = {
    "t_sum": (lambda m: mo.SUM(m.x.rev, m.o.once, m.x.trk),
              [11.0, 22.0, 33.0, 94.0], [12.0, 23.0, 34.0, 95.0]),
    "t_med": (lambda m: mo.MEDIAN(m.x.rev, m.o.once, m.x.trk),
              [1.0, 2.0, 3.0, 40.0], [1.0, 2.0, 3.0, 41.0]),
    "t_add": (lambda m: m.x.trk + m.o.once, [10.0, 20.0, 30.0, 90.0], [11.0, 21.0, 31.0, 91.0]),
    "t_sub": (lambda m: m.x.trk - m.t.fee, [10.0, 20.0, 20.0, 20.0], [11.0, 21.0, 21.0, 21.0]),
    "t_if": (lambda m: mo.IF(m.x.rev > 2, m.x.trk, m.o.once),
             [0.0, 0.0, 30.0, 40.0], [0.0, 0.0, 31.0, 41.0]),
    "t_cmp": (lambda m: mo.IF(m.x.trk - m.t.fee > 15, m.o.once, 0),
              [0.0, 0.0, 0.0, 50.0], [0.0, 0.0, 0.0, 50.0]),
}


class TestTrackedLines:
    @pytest.mark.parametrize("name", list(_TRACKED))
    def test_a_plain_line_is_read_on_the_tracked_line_s_window(self, name):
        m = _tracked()
        make, plan, actual = _TRACKED[name]
        got = make(m)
        assert got.value["plan"] == pytest.approx(plan)
        assert got.value["actual"] == pytest.approx(actual)
        assert got.time.start == "2026-01"

    def test_a_call_reads_what_the_operators_read(self):
        m = _tracked()
        x = m.x
        avg = mo.AVERAGE(x.rev, m.t.fee, x.trk).value
        ops = ((x.rev + m.t.fee + x.trk) / 3).value
        for role in ("plan", "actual"):
            assert avg[role] == pytest.approx(ops[role])
        assert mo.SUM(x.rev, m.o.once, x.trk).value["plan"] == (
            (x.rev + m.o.once) + x.trk).value["plan"]

    def test_a_tracked_line_is_not_continued_past_its_window(self):
        m = _tracked()
        for call in (lambda: m.x.trk + m.u.early, lambda: mo.SUM(m.u.early, m.x.trk)):
            with pytest.raises(ValueError, match="'trk' holds tracks on 4 period"):
                call()

    def test_a_list_typed_in_beside_tracks_keeps_its_reading(self):
        m = _tracked()
        got = mo.IF(m.x.rev > 2, m.x.trk, [7.0, 8.0]).value
        assert got["plan"] == [7.0, 8.0, 30.0, 40.0]

    def test_the_workbook_reads_each_track_by_date(self, tmp_path):
        m = _tracked()
        with m:
            m.x.y = (m.x.trk + m.o.once).set_display_name("y")
            m.x.s = mo.SUM(m.x.rev, m.t.fee, m.x.trk).set_display_name("s")
        path = tmp_path / "tracked.xlsx"
        m.to_excel(path)
        x = _rows(load_workbook(path)["X"])
        assert x["y"] == ["=B3 + 0", "=C3 + 0", "=D3 + 0", "=E3 + O!B2"]
        assert x["y · actual"] == ["=B4 + 0", "=C4 + 0", "=D4 + 0", "=E4 + O!B2"]
        assert x["s · actual"] == ["=SUM(B2, 0, B4)", "=SUM(C2, 0, C4)",
                                   "=SUM(D2, T!B2, D4)", "=SUM(E2, T!C2, E4)"]


class TestRefusalsSayWhatToDo:
    def test_zero_continues_numbers_not_dates_or_text(self):
        m = _lines()
        with m:
            m.t.tdz = mo.Variable([date(2026, 6, 30), date(2026, 7, 31)], display_name="tdz",
                                  extend=mo.zero())
            m.t.twz = mo.Variable(["p", "q"], display_name="twz", extend=mo.zero())
        for call in (lambda: mo.DAYS(m.t.tdz, m.x.dt), lambda: mo.MAX(m.t.tdz, m.x.dt),
                     lambda: m.t.tdz > m.x.dt):
            with pytest.raises(ValueError, match="'tdz' holds dates, and extend=mo.zero"):
                call()
        with pytest.raises(ValueError, match="'twz' holds text") as refused:
            mo.CONCAT(m.x.tx, m.t.twz)
        assert "extend=mo.hold()" in str(refused.value)
        # Held, the date stands where the line has none.
        assert mo.DAYS(m.t.td, m.x.dt).value == [150, 122, 91, 92]

    def test_a_re_grained_line_says_its_rule_stays_behind(self):
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=6)
        with m:
            m.x = mo.MultiVariable("X")
            m.x.rev = mo.Variable([1.0, 2.0, 3.0, 4.0, 5.0, 6.0], display_name="rev",
                                  regrain=mo.up("sum"))
            m.u = mo.MultiVariable("U", default_grain="month", default_start="2026-01",
                                   default_periods=2)
            m.u.early = mo.Variable([100.0, 200.0], display_name="early", extend=mo.zero(),
                                    regrain=mo.up("sum"))
        with pytest.raises(ValueError, match="A re-grained line keeps no rule") as refused:
            m.u.early.at("quarter") + m.x.rev.at("quarter")
        assert "'early' declares extend=mo.zero()" in str(refused.value)
        both = (m.u.early + m.x.rev).set_regrain(mo.up("sum")).at("quarter")
        assert both.value == [306.0, 15.0]


# ══════════════════════════════════════════════════════════════════════
# A loop container reads a line on another window by date
# ══════════════════════════════════════════════════════════════════════


def _loop(m, make):
    """``make(m, d)`` run inside a loop container ``D`` of ``m``."""
    with m:
        m.d = mo.MultiVariable("D", loop=True)
        with m.d as d:
            make(m, d)
    return m.d


#: Rows of a loop over X's four months reading T, U and O by date: fee
#: [0, 0, 10, 20], rate [0.5, 0.5, 0.5, 0.25], cap [9, 8, 8, 8], once
#: [0, 0, 50, 0], level [2, 2, 2, 2].
_LOOPS = {
    "once": (lambda m, d: setattr(d, "bal", mo.lag(d.bal, fill=0.0) + m.o.once),
             {"bal": [0.0, 0.0, 50.0, 50.0]}),
    "fee": (lambda m, d: setattr(d, "bal", mo.lag(d.bal, fill=0.0) + m.t.fee),
            {"bal": [0.0, 0.0, 10.0, 30.0]}),
    "rate": (lambda m, d: setattr(d, "bal", mo.lag(d.bal, fill=100.0) * m.t.rate),
             {"bal": [50.0, 25.0, 12.5, 3.125]}),
    "cap": (lambda m, d: setattr(d, "bal", mo.lag(d.bal, fill=0.0) + m.u.cap),
            {"bal": [9.0, 17.0, 25.0, 33.0]}),
    "level": (lambda m, d: setattr(d, "bal", mo.lag(d.bal, fill=1.0) * m.o.level),
              {"bal": [2.0, 4.0, 8.0, 16.0]}),
    "mix": (lambda m, d: setattr(d, "bal", mo.lag(d.bal, fill=0.0) + m.x.rev + m.t.fee),
            {"bal": [1.0, 3.0, 16.0, 40.0]}),
    "max": (lambda m, d: setattr(d, "bal", mo.MAX(mo.lag(d.bal, fill=0.0), m.t.fee)),
            {"bal": [0.0, 0.0, 10.0, 20.0]}),
    "if": (lambda m, d: setattr(d, "bal",
                                mo.IF(m.x.rev > 2, mo.lag(d.bal, fill=0.0) + m.t.fee, 1.0)),
           {"bal": [1.0, 1.0, 11.0, 31.0]}),
    "two": (lambda m, d: (setattr(d, "a", mo.lag(d.b, fill=0.0) + m.o.once),
                          setattr(d, "b", d.a + m.t.fee)),
            {"a": [0.0, 0.0, 50.0, 60.0], "b": [0.0, 0.0, 60.0, 80.0]}),
}


class TestALoopContainerReadsByDate:
    @pytest.mark.parametrize("name", list(_LOOPS))
    def test_each_period_reads_the_line_s_value_of_its_date(self, name):
        make, expected = _LOOPS[name]
        d = _loop(_lines(), make)
        for row, want in expected.items():
            assert getattr(d, row).value == pytest.approx(want), row

    def test_the_row_reads_what_the_operators_read(self):
        # Outside a loop the same reads align by date: the loop agrees.
        m = _lines()
        fee_on_x = (m.x.rev + m.t.fee - m.x.rev).value
        assert fee_on_x == [0.0, 0.0, 10.0, 20.0]
        d = _loop(m, _LOOPS["fee"][0])
        bal, total = d.bal.value, 0.0
        for t, f in enumerate(fee_on_x):
            total += f
            assert bal[t] == total

    def test_a_line_without_a_rule_and_a_formula_are_refused(self):
        for make, message in (
                (lambda m, d: setattr(d, "bal", mo.lag(d.bal, fill=0.0) + m.t.nofill),
                 "'nofill' does not cover the combined window"),
                (lambda m, d: setattr(d, "bal", mo.lag(d.bal, fill=0.0) + m.o.bare),
                 "'bare' does not cover the combined window"),
                (lambda m, d: setattr(d, "bal", mo.lag(d.bal, fill=0.0) + m.t.fee * 2),
                 r"the line fee \* 2 does not cover the combined window"),
                (lambda m, d: setattr(d, "bal", mo.lag(d.bal, fill=0.0) * (1 + m.t.rate)),
                 r"the line rate \+ 1 does not cover the combined window"),
                (lambda m, d: setattr(d, "bal",
                                      mo.lag(d.bal, fill=0.0) + mo.lag(m.t.fee, fill=0.0)),
                 "does not cover the combined window")):
            with pytest.raises(ValueError, match=message):
                _loop(_lines(), make)
        with pytest.raises(ValueError, match="different grains"):
            _loop(_windows(), lambda m, d: setattr(d, "bal", mo.lag(d.bal, fill=0.0) + m.q.h))

    def test_the_workbook_reads_the_cell_of_each_date(self, tmp_path):
        m = _lines()
        with m:
            for name in ("once", "fee", "rate", "if", "two"):
                mv = mo.MultiVariable("D_" + name, loop=True)
                setattr(m, "d_" + name, mv)
                with mv as d:
                    _LOOPS[name][0](m, d)
        path = tmp_path / "loops.xlsx"
        m.to_excel(path)
        book = load_workbook(path)
        rows = {name: _rows(book["D_" + name]) for name in ("once", "fee", "rate", "if", "two")}
        # The row's own previous cell, then the line's cell of the same
        # date — O's one cell in March only, T's from March, a rate held.
        assert rows["once"]["Bal"] == ["=0 + 0", "=B2 + 0", "=C2 + O!B2", "=D2 + 0"]
        assert rows["fee"]["Bal"] == ["=0 + 0", "=B2 + 0", "=C2 + T!B2", "=D2 + T!C2"]
        assert rows["rate"]["Bal"] == ["=100.0 * T!B3", "=B2 * T!B3", "=C2 * T!B3",
                                       "=D2 * T!C3"]
        assert rows["if"]["Bal"] == ["=IF(X!B2 > 2, 0 + 0, 1.0)", "=IF(X!C2 > 2, B2 + 0, 1.0)",
                                     "=IF(X!D2 > 2, C2 + T!B2, 1.0)",
                                     "=IF(X!E2 > 2, D2 + T!C2, 1.0)"]
        assert rows["two"]["A"] == ["=0 + 0", "=B3 + 0", "=C3 + O!B2", "=D3 + 0"]
        assert rows["two"]["B"] == ["=B2 + 0", "=C2 + 0", "=D2 + T!B2", "=E2 + T!C2"]


# ══════════════════════════════════════════════════════════════════════
# The workbook reads by date where Python does, by position where it does
# ══════════════════════════════════════════════════════════════════════


class TestTheWorkbook:
    def test_each_line_is_read_on_its_date_s_cell(self, tmp_path):
        m = _model()
        path = tmp_path / "book.xlsx"
        m.to_excel(path)
        x = _rows(load_workbook(path)["X"])
        # rev B2, flag B3, idx B4, big B5, tx B6, dt B7, lk B8; on T fee
        # B2:C2, rate B3:C3, td B5:C5, tw B6:C6, rw B8:C8, cc B9:C9; on U
        # early B2:C2, cap B3:C3; K's block B2:E3; I's cells B1..B5.
        assert x["if_fee"] == ["=IF(B2 > 2, B2, 0)", "=IF(C2 > 2, C2, 0)",
                               "=IF(D2 > 2, D2, T!B2)", "=IF(E2 > 2, E2, T!C2)"]
        assert x["if_rate"] == ["=IF(B2 > 2, T!B3, U!B3)", "=IF(C2 > 2, T!B3, U!C3)",
                                "=IF(D2 > 2, T!B3, U!C3)", "=IF(E2 > 2, T!C3, U!C3)"]
        assert x["if_flag"] == ["=IF(B3, 0, U!B2)", "=IF(C3, 0, U!C2)", "=IF(D3, T!B2, 0)",
                                "=IF(E3, T!C2, 0)"]
        assert x["and_"] == ["=AND(B2, 0)", "=AND(C2, 0)", "=AND(D2, T!B2)", "=AND(E2, T!C2)"]
        assert x["or_"] == ["=OR(B3, U!B2)", "=OR(C3, U!C2)", "=OR(D3, 0)", "=OR(E3, 0)"]
        assert x["ch_series"] == ["=CHOOSE(B4, 0, U!B3)", "=CHOOSE(C4, 0, U!C3)",
                                  "=CHOOSE(D4, T!B2, U!C3)", "=CHOOSE(E4, T!C2, U!C3)"]
        assert x["ch_scalar"] == ["=CHOOSE(2, B2, 0)", "=CHOOSE(2, C2, 0)",
                                  "=CHOOSE(2, D2, T!B2)", "=CHOOSE(2, E2, T!C2)"]
        assert x["md"] == ["=MOD(B5, U!B3)", "=MOD(C5, U!C3)", "=MOD(D5, U!C3)",
                           "=MOD(E5, U!C3)"]
        assert x["gt"] == ["=B2 > 0", "=C2 > 0", "=D2 > T!B2", "=E2 > T!C2"]
        assert x["ge"] == ["=U!B3 >= B5", "=U!C3 >= C5", "=U!C3 >= D5", "=U!C3 >= E5"]
        assert x["cc_"] == ["=B6 & T!B6", "=C6 & T!B6", "=D6 & T!B6", "=E6 & T!C6"]
        assert x["dd"] == ["=DATE(2026, B2, 0)", "=DATE(2026, C2, 0)",
                           "=DATE(2026, D2, T!B2)", "=DATE(2026, E2, T!C2)"]
        assert x["d360"] == ["=DAYS360(B7, T!B5, FALSE)", "=DAYS360(C7, T!B5, FALSE)",
                             "=DAYS360(D7, T!B5, FALSE)", "=DAYS360(E7, T!C5, FALSE)"]
        assert x["yf"][1] == "=YEARFRAC(C7, T!B5, 3)" and x["yf"][3] == "=YEARFRAC(E7, T!C5, 3)"
        assert x["dy"] == ["=_xlfn.DAYS(T!B5, B7)", "=_xlfn.DAYS(T!B5, C7)",
                           "=_xlfn.DAYS(T!B5, D7)", "=_xlfn.DAYS(T!C5, E7)"]
        assert x["nwd"][0] == "=NETWORKDAYS(B7, T!B5, I!B5)"
        assert x["nwd"][3] == "=NETWORKDAYS(E7, T!C5, I!B5)"
        assert x["xl"] == ["=_xlfn.XLOOKUP(B2, I!B1:B2, I!B3:B4, 0)",
                           "=_xlfn.XLOOKUP(C2, I!B1:B2, I!B3:B4, 0)",
                           "=_xlfn.XLOOKUP(D2, I!B1:B2, I!B3:B4, T!B2)",
                           "=_xlfn.XLOOKUP(E2, I!B1:B2, I!B3:B4, T!C2)"]
        assert x["hl"] == ["=HLOOKUP(B2, K!B2:E3, T!B8, FALSE)",
                           "=HLOOKUP(C2, K!B2:E3, T!B8, FALSE)",
                           "=HLOOKUP(D2, K!B2:E3, T!B8, FALSE)",
                           "=HLOOKUP(E2, K!B2:E3, T!C8, FALSE)"]
        assert x["ix"] == ["=INDEX(K!B2:E3, B4, T!B9)", "=INDEX(K!B2:E3, C4, T!B9)",
                           "=INDEX(K!B2:E3, D4, T!B9)", "=INDEX(K!B2:E3, E4, T!C9)"]
        assert x["vl"][1] == "=VLOOKUP(C8, K!B2:E3, T!B9, FALSE)"
        assert x["med"] == ["=MEDIAN(B2, 0, U!B3)", "=MEDIAN(C2, 0, U!C3)",
                            "=MEDIAN(D2, T!B2, U!C3)", "=MEDIAN(E2, T!C2, U!C3)"]
        assert x["sdp"] == ["=_xlfn.STDEV.P(B2, U!B2)", "=_xlfn.STDEV.P(C2, U!C2)",
                            "=_xlfn.STDEV.P(D2, 0)", "=_xlfn.STDEV.P(E2, 0)"]
        # The recurrence's previous cell, then this period's cells by date.
        assert x["rc"][1:] == ["=B35 + C2 + 0", "=C35 + D2 + T!B2", "=D35 + E2 + T!C2"]
        assert x["ie"] == ["=IFERROR(B2 / 0, 0)", "=IFERROR(C2 / 0, 0)",
                           "=IFERROR(D2 / T!B2, 0)", "=IFERROR(E2 / T!C2, 0)"]
        # O's one cell: once (B2) in March only, 0 elsewhere; level (B3) held.
        assert x["o_if"] == ["=IF(B2 > 3, B2, 0)", "=IF(C2 > 3, C2, 0)",
                             "=IF(D2 > 3, D2, O!B2)", "=IF(E2 > 3, E2, 0)"]
        assert x["o_if2"] == ["=IF(B3, O!B3, 0)", "=IF(C3, O!B3, 0)", "=IF(D3, O!B3, O!B2)",
                              "=IF(E3, O!B3, 0)"]
        assert x["o_med"][2:] == ["=MEDIAN(D2, O!B2, O!B3)", "=MEDIAN(E2, 0, O!B3)"]
        assert x["o_cat"] == ["=B6 & 0", "=C6 & 0", "=D6 & O!B2", "=E6 & 0"]
        assert x["o_rnd"] == ["=ROUND(K!B4, O!B3)", "=ROUND(K!C4, O!B3)", "=ROUND(K!D4, O!B3)",
                              "=ROUND(K!E4, O!B3)"]
        assert x["o_ed"][0] == "=EDATE(B7, O!B3)"
        assert x["o_pmt"][3] == "=PMT(0.05, O!B3, K!E3, 0, 0)"
        # A formula over X's lines, written in X's column, beside the others' dates.
        assert x["w_add"] == ["=B2 * 1 + 0", "=C2 * 1 + 0", "=D2 * 1 + T!B2", "=E2 * 1 + T!C2"]
        assert x["w_if"][2:] == ["=IF(D2 * 1 > 2, D2 * 1, T!B2)", "=IF(E2 * 1 > 2, E2 * 1, T!C2)"]
        assert x["w_once"] == ["=B2 * 1 + 0", "=C2 * 1 + 0", "=D2 * 1 + O!B2", "=E2 * 1 + 0"]
        assert x["w_max"][0] == "=MAX(B2 * 1, 0, U!B3)"
        assert x["w_cmp"] == ["=IF(B2 > 2, 0, U!B3)", "=IF(C2 > 2, 0, U!C3)",
                              "=IF(D2 > 2, T!B2, U!C3)", "=IF(E2 > 2, T!C2, U!C3)"]
        # A list of lines on several sheets forms no run in one column: the
        # cell carries the value Python read by date, and check_compat says so.
        for name in ("mdl", "pct", "lg", "rk"):
            assert x[name] == getattr(m.x, name).value, name
        assert {i.qpath for i in mo.check_compat(m, "excel")} == {
            "m.x.mdl", "m.x.pct", "m.x.lg", "m.x.rk"}

    def test_a_formula_over_one_other_window_is_read_by_position_as_python_reads_it(
            self, tmp_path):
        # T2 spans four months from November: nothing on X's window meets
        # it in ``fee * 2``, so Python pairs January with November, and so
        # does the workbook — the cell of the same position, never of the
        # same date (which Python did not read).
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=4)
        with m:
            m.x = mo.MultiVariable("X")
            m.x.rev = mo.Variable([1.0, 2.0, 3.0, 4.0], display_name="rev")
            m.t2 = mo.MultiVariable("T2", default_grain="month", default_start="2025-11",
                                    default_periods=4)
            m.t2.fee = mo.Variable([10.0, 20.0, 30.0, 40.0], display_name="fee",
                                   extend=mo.zero())
            m.x.d = (m.t2.fee * 2).set_display_name("d")
            m.x.e = mo.ABS(m.t2.fee).set_display_name("e")
            m.x.c = (m.t2.fee > 15).set_display_name("c")
        assert m.x.d.value == [20.0, 40.0, 60.0, 80.0]
        path = tmp_path / "book.xlsx"
        m.to_excel(path)
        x = _rows(load_workbook(path)["X"])
        assert x["d"] == ["=T2!B2 * 2", "=T2!C2 * 2", "=T2!D2 * 2", "=T2!E2 * 2"]
        assert x["e"] == ["=ABS(T2!B2)", "=ABS(T2!C2)", "=ABS(T2!D2)", "=ABS(T2!E2)"]
        assert x["c"] == ["=T2!B2 > 15", "=T2!C2 > 15", "=T2!D2 > 15", "=T2!E2 > 15"]

    def test_constants_keys_and_lines_on_windows_of_their_own(self, tmp_path):
        # A constant's one cell in every column; lines along keys in another
        # order, on another sheet, key by key; a line on a window of its own
        # (start=) that reads another line by position, as Python does.
        m = _windows()
        with m:
            m.x.c_add = (m.x.rev + m.x.c).set_display_name("c_add")
            m.x.c_if = mo.IF(m.x.rev > 3, m.x.c, m.t.fee).set_display_name("c_if")
            m.x.k_max = mo.MAX(m.x.k, m.t.k2).set_display_name("k_max")
            m.t2 = mo.MultiVariable("T2", default_grain="month", default_start="2026-03",
                                    default_periods=4)
            m.t2.y = mo.Variable(m.x.rev, start="2026-03", grain="month").set_display_name("y")
        assert m.t2.y.value == [1.0, 2.0, 3.0, 4.0]
        path = tmp_path / "book.xlsx"
        m.to_excel(path)
        book = load_workbook(path)
        x = _rows(book["X"])
        assert x["c_add"] == ["=B2 + B9", "=C2 + B9", "=D2 + B9", "=E2 + B9"]
        assert x["c_if"] == ["=IF(B2 > 3, B9, 0)", "=IF(C2 > 3, B9, 0)",
                             "=IF(D2 > 3, B9, T!B2)", "=IF(E2 > 3, B9, T!C2)"]
        assert x["k_max"][:3] == ["=MAX(B11, T!D11)", "=MAX(C11, T!C11)", "=MAX(D11, T!B11)"]
        assert _rows(book["T2"])["y"] == ["=X!B2", "=X!C2", "=X!D2", "=X!E2"]

    def test_a_copy_a_track_slice_and_a_lag_keep_their_positions(self, tmp_path):
        # Placed on X by something other than alignment — a copy, a track
        # slice, a lag of a line on another window — a value carries no
        # mark and reads its line by position, as Python does, beside a
        # line read by date.
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=4)
        with m:
            m.x = mo.MultiVariable("X")
            m.x.rev = mo.Variable([1.0, 2.0, 3.0, 4.0], display_name="rev")
            m.t = mo.MultiVariable("T", default_grain="month", default_start="2026-03",
                                   default_periods=2)
            m.t.fee = mo.Variable([10.0, 20.0], display_name="fee", extend=mo.zero())
            m.t4 = mo.MultiVariable("T4", default_grain="month", default_start="2026-03",
                                    default_periods=4)
            m.t4.f4 = mo.Variable([10.0, 20.0, 30.0, 40.0], display_name="f4")
            m.t4.tr = mo.Variable(tracks={"plan": [10.0, 20.0, 30.0, 40.0],
                                          "actual": [11.0, 21.0, 31.0, 41.0]},
                                  display_name="tr")
            m.x.cp = m.t4.f4.copy().set_display_name("cp")
            m.x.sl = (m.t4.tr.at(track="plan") * 2).set_display_name("sl")
            m.x.lg = mo.lag(m.t4.f4, fill=0.0).set_display_name("lg")
            m.x.s = (m.x.rev + m.t.fee).set_display_name("s")
        assert m.x.cp.value == [10.0, 20.0, 30.0, 40.0]
        assert m.x.lg.value == [0.0, 10.0, 20.0, 30.0]
        assert not any(getattr(m.x, n)._read_by_date for n in ("cp", "sl", "lg"))
        path = tmp_path / "book.xlsx"
        m.to_excel(path)
        x = _rows(load_workbook(path)["X"])
        assert x["cp"] == ["=T4!B2", "=T4!C2", "=T4!D2", "=T4!E2"]
        assert x["sl"] == ["=T4!B3 * 2", "=T4!C3 * 2", "=T4!D3 * 2", "=T4!E3 * 2"]
        assert x["lg"] == ["=0", "=T4!B2", "=T4!C2", "=T4!D2"]
        assert x["s"] == ["=B2 + 0", "=C2 + 0", "=D2 + T!B2", "=E2 + T!C2"]

    def test_an_aligned_intermediate_on_a_window_of_its_own_is_refused_by_the_export(
            self, tmp_path):
        # ``bz + cz`` met two windows and runs from November; a function or
        # an operator over it alone keeps those dates, and X's columns
        # start in January — refused as the bare sum is, naming the sum.
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=6)
        with m:
            m.x = mo.MultiVariable("X")
            m.x.rev = mo.Variable([1.0, 2.0, 3.0, 4.0, 5.0, 6.0], display_name="rev")
            m.b = mo.MultiVariable("B", default_grain="month", default_start="2025-11",
                                   default_periods=4)
            m.b.bz = mo.Variable([1000.0, 2000.0, 3000.0, 4000.0], display_name="bz",
                                 extend=mo.zero())
            m.c = mo.MultiVariable("C", default_grain="month", default_start="2026-01",
                                   default_periods=4)
            m.c.cz = mo.Variable([1.0, 2.0, 3.0, 4.0], display_name="cz", extend=mo.zero())
            m.x.y1 = mo.IFERROR(m.b.bz + m.c.cz, 0).set_display_name("y1")
        assert m.x.y1.value == [1000.0, 2000.0, 3001.0, 4002.0, 3.0, 4.0]
        with pytest.raises(ValueError, match="window of its own") as refused:
            m.to_excel(tmp_path / "w1.xlsx")
        assert "'y1'" in str(refused.value) and "the line bz + cz" in str(refused.value)
        for name, make in (("y2", lambda: (m.b.bz + m.c.cz) * 1),
                           ("y3", lambda: mo.ROUND(m.b.bz + m.c.cz, 0))):
            with m:
                setattr(m.x, name, make().set_display_name(name))
            with pytest.raises(ValueError, match="window of its own"):
                m.to_excel(tmp_path / f"{name}.xlsx")
            delattr(m.x, name)
        delattr(m.x, "y1")
        # A total over it reads no dates; the same read inside the sheet's
        # window is written by date.
        with m:
            m.i = mo.MultiVariable("I")
            m.i.tot = mo.SUM(m.b.bz + m.c.cz).set_display_name("tot")
            m.x.ok = mo.IFERROR(m.x.rev + m.c.cz, 0).set_display_name("ok")
        m.to_excel(tmp_path / "ok.xlsx")
        book = load_workbook(tmp_path / "ok.xlsx")
        assert _rows(book["I"])["tot"][0] == 10010.0
        assert _rows(book["X"])["ok"] == ["=IFERROR(B2 + C!B2, 0)", "=IFERROR(C2 + C!C2, 0)",
                           "=IFERROR(D2 + C!D2, 0)", "=IFERROR(E2 + C!E2, 0)",
                           "=IFERROR(F2 + 0, 0)", "=IFERROR(G2 + 0, 0)"]

    def test_two_windows_reaching_beyond_the_sheet_s_are_refused_by_the_export(
            self, tmp_path):
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=4)
        with m:
            m.x = mo.MultiVariable("X")
            m.x.rev = mo.Variable([1.0, 2.0, 3.0, 4.0], display_name="rev", extend=mo.zero())
            m.x.on = mo.Variable([1.0, 0.0, 1.0, 0.0], display_name="on", extend=mo.zero())
            m.t2 = mo.MultiVariable("T2", default_grain="month", default_start="2025-11",
                                    default_periods=4)
            m.t2.fee = mo.Variable([10.0, 20.0, 30.0, 40.0], display_name="fee",
                                   extend=mo.zero())
            m.x.i = mo.IF(m.x.on, m.x.rev, m.t2.fee).set_display_name("i")
            m.x.g = (m.x.rev > m.t2.fee).set_display_name("g")
        # November to April: on [0, 0, 1, 0, 1, 0].
        assert m.x.i.value == [10.0, 20.0, 1.0, 40.0, 3.0, 0.0]
        assert m.x.i.time.start == "2025-11" and m.x.g.time.start == "2025-11"
        with pytest.raises(ValueError, match="window of its own") as refused:
            m.to_excel(tmp_path / "w.xlsx")
        assert "'i'" in str(refused.value) and "'g'" in str(refused.value)

    def test_an_intermediate_aligned_inside_a_formula_is_read_on_its_dates(self, tmp_path):
        # ``(rev + fee) * 2``: the sum is aligned by date, the product reads
        # it by position — the cells of the sum's dates.
        m = _lines()
        with m:
            m.x.tw2 = ((m.x.rev + m.t.fee) * 2).set_display_name("tw2")
            m.x.ng = (-(m.x.rev + m.t.fee)).set_display_name("ng")
        assert m.x.tw2.value == [2.0, 4.0, 26.0, 48.0]
        path = tmp_path / "book.xlsx"
        m.to_excel(path)
        x = _rows(load_workbook(path)["X"])
        assert x["tw2"] == ["=(B2 + 0) * 2", "=(C2 + 0) * 2", "=(D2 + T!B2) * 2",
                            "=(E2 + T!C2) * 2"]
        assert x["ng"][0] == "=-(B2 + 0)" and x["ng"][3] == "=-(E2 + T!C2)"


# ══════════════════════════════════════════════════════════════════════
# Lines along keys combine key by key, as a + b combines them
# ══════════════════════════════════════════════════════════════════════


def _keyed():
    m = mo.Model("m", default_grain="year", default_start="2026", default_periods=3)
    with m:
        m.s = mo.Variable([1.0, 2.0, 3.0], display_name="s")
        m.k = mo.Variable({"Bear": 0.9, "Base": 1.0, "Bull": 1.1}, display_name="k")
        m.k2 = mo.Variable({"Bear": 1.0, "Base": 2.0, "Bull": 3.0}, display_name="k2")
        m.k3 = mo.Variable({"Bull": 1.0, "Base": 2.0, "Bear": 3.0}, display_name="k3")
        m.k4 = mo.Variable({"Bear": 1.0, "Base": 2.0, "Mid": 3.0}, display_name="k4")
    return m


#: Two lines along the same keys in another order: ``k3`` read in ``k``'s
#: order (Bear 3, Base 2, Bull 1).
_KEYED = {
    "plus": (lambda a, b: a + b, [3.9, 3.0, 2.1]),
    "SUM": (lambda a, b: mo.SUM(a, b), [3.9, 3.0, 2.1]),
    "MAX": (lambda a, b: mo.MAX(a, b), [3.0, 2.0, 1.1]),
    "COUNT": (lambda a, b: mo.COUNT(a, b), [2, 2, 2]),
    "POWER": (lambda a, b: mo.POWER(a, b), [0.9 ** 3, 1.0, 1.1]),
    "IFERROR": (lambda a, b: mo.IFERROR(a / (b - 2), 0), [0.9, 0.0, -1.1]),
    "IF": (lambda a, b: mo.IF(a > 1, a, b), [3.0, 2.0, 1.1]),
    "AND": (lambda a, b: mo.AND(a > 0.95, b > 1.5), [False, True, False]),
    "MOD": (lambda a, b: mo.MOD(a, b), [0.9, 1.0, 0.1]),
    "CONCAT": (lambda a, b: mo.CONCAT(a, b), ["0.93", "12", "1.11"]),
    "MEDIAN": (lambda a, b: mo.MEDIAN(a, b), [1.95, 1.5, 1.05]),
    "compare": (lambda a, b: a > b, [False, False, True]),
}


class TestLinesAlongKeys:
    @pytest.mark.parametrize("name", list(_KEYED))
    def test_another_order_is_combined_key_by_key(self, name):
        m = _keyed()
        call, expected = _KEYED[name]
        got = call(m.k, m.k3)
        if isinstance(expected[0], float):
            assert got.value == pytest.approx(expected)
        else:
            assert got.value == expected
        assert got._keys == ["Bear", "Base", "Bull"]

    @pytest.mark.parametrize("name", list(_KEYED))
    def test_other_keys_are_refused(self, name):
        m = _keyed()
        with pytest.raises(ValueError, match="different key sets|other keys"):
            m.k + m.k4
        with pytest.raises(ValueError, match="different key sets|other keys"):
            _KEYED[name][0](m.k, m.k4)

    @pytest.mark.parametrize("name", list(_KEYED))
    def test_against_a_period_series_it_is_refused(self, name):
        m = _keyed()
        with pytest.raises(ValueError, match="runs along its keys"):
            _KEYED[name][0](m.k, m.s)

    def test_a_line_along_keys_among_many_series_is_refused(self):
        m = _keyed()
        many = [m.s * (i + 1) for i in range(50)]
        assert mo.SUM(*many).value == pytest.approx([1275.0, 2550.0, 3825.0])
        with pytest.raises(ValueError, match="runs along its keys"):
            mo.SUM(*many, m.k)
        with pytest.raises(ValueError, match="runs along its keys"):
            mo.IF(m.s > 1, m.s, m.k)

    def test_one_line_along_keys_keeps_them(self):
        m = _keyed()
        for got in (mo.ABS(m.k), mo.ROUND(m.k, 1), mo.POWER(m.k, 2), mo.IF(m.k > 1, m.k, 0),
                    mo.NOT(m.k > 1), mo.CONCAT(m.k, "x")):
            assert got._keys == ["Bear", "Base", "Bull"]
            with pytest.raises(ValueError, match="runs along its keys"):
                got + m.s

    def test_the_workbook_names_each_key_s_cell(self, tmp_path):
        m = _keyed()
        with m:
            m.p = (m.k + m.k3).set_display_name("p")
            m.pw = mo.POWER(m.k, m.k3).set_display_name("pw")
            m.same = (m.k + m.k2).set_display_name("same")
        path = tmp_path / "keys.xlsx"
        m.to_excel(path)
        rows = _rows(load_workbook(path)["M"])
        # k is B4:D4 (Bear, Base, Bull); k3 is B7:D7 (Bull, Base, Bear).
        assert rows["p"] == ["=B4 + D7", "=C4 + C7", "=D4 + B7"]
        assert rows["pw"] == ["=POWER(B4, D7)", "=POWER(C4, C7)", "=POWER(D4, B7)"]
        assert rows["same"] == ["=B4 + B5", "=C4 + C5", "=D4 + D5"]


# ══════════════════════════════════════════════════════════════════════
# A list of lines given as one argument
# ══════════════════════════════════════════════════════════════════════


def _listed():
    m = mo.Model("m", default_grain="year", default_start="2026", default_periods=3)
    with m:
        m.x = mo.MultiVariable("X")
        m.x.a = mo.Variable([1.0, 2.0, 3.0], display_name="a")
        m.x.b = mo.Variable([10.0, 20.0, 30.0], display_name="b")
        m.i = mo.MultiVariable("I")
        m.i.p = mo.Variable(4.0, display_name="p")
        m.i.q = mo.Variable(5.0, display_name="q")
    return m


_AGGREGATES = ("SUM", "AVERAGE", "MIN", "MAX", "PRODUCT", "COUNT", "COUNTA")


class TestAListOfLinesAsOneArgument:
    @pytest.mark.parametrize("name,expected", [
        ("SUM", 9.0), ("AVERAGE", 4.5), ("MIN", 4.0), ("MAX", 5.0), ("PRODUCT", 20.0),
        ("COUNT", 2), ("COUNTA", 2),
    ])
    def test_scalar_lines_are_one_number(self, name, expected):
        m = _listed()
        call = getattr(mo, name)
        got = call([m.i.p, m.i.q])
        assert got.value == expected and got.var_type == "scalar"
        assert call((m.i.p, m.i.q)).value == expected
        # Beside a series, the same values in every period.
        assert call([m.i.p, m.i.q], m.x.a).value == call(m.i.p, m.i.q, m.x.a).value

    @pytest.mark.parametrize("name", ("SUM", "AVERAGE", "MIN", "MAX", "PRODUCT"))
    def test_the_five_read_the_lines_one_by_one(self, name):
        m = _listed()
        got = getattr(mo, name)([m.i.p, m.i.q])
        assert got.formula == f"{name}(p, q)"
        with m:
            m.i.p.set_unit("USD")
            m.i.q.set_unit("EUR")
        if name != "PRODUCT":
            with pytest.raises(ValueError, match="one unit"):
                getattr(mo, name)([m.i.p, m.i.q])

    @pytest.mark.parametrize("name", _AGGREGATES)
    def test_a_series_in_the_list_is_refused_naming_both_spellings(self, name):
        m = _listed()
        for arg in ([m.x.a, m.x.b], (m.x.a, m.x.b), [m.x.a], [m.x.a, m.i.p]):
            with pytest.raises(ValueError) as refused:
                getattr(mo, name)(arg)
            text = str(refused.value)
            assert f"mo.{name}([" in text and "per period" in text, text
            assert "over all their periods" in text, text
        with pytest.raises(ValueError, match="per period, the lines one by one, "
                                             rf"mo\.{name}\(a, b\)"):
            getattr(mo, name)([m.x.a, m.x.b], m.i.p)

    def test_the_sum_s_two_readings_are_spelled_out(self):
        m = _listed()
        with pytest.raises(ValueError) as refused:
            mo.SUM([m.x.a, m.x.b])
        assert ("per period, the lines one by one, mo.SUM(a, b); the total over all their "
                "periods, mo.SUM(a) + mo.SUM(b)") in str(refused.value)

    def test_the_statistics_read_a_list_of_series_one_sample_per_period(self):
        m = _listed()
        assert mo.MEDIAN([m.x.a, m.x.b]).value == [5.5, 11.0, 16.5]
        assert mo.MEDIAN([m.i.p, m.i.q]).value == 4.5

    def test_the_workbook_writes_each_line_s_cell(self, tmp_path):
        m = _listed()
        with m:
            for name in _AGGREGATES:
                setattr(m.i, name.lower(), getattr(mo, name)([m.i.p, m.i.q])
                        .set_display_name(name.lower()))
        path = tmp_path / "list.xlsx"
        m.to_excel(path)
        rows = _rows(load_workbook(path)["I"])
        for name in ("SUM", "AVERAGE", "MIN", "MAX", "PRODUCT"):
            assert rows[name.lower()][0] == f"={name}(B1, B2)", name
        # A count reads the list as the range it is, as before.
        assert rows["count"][0] == "=COUNT(B1:B2)"
        assert rows["counta"][0] == "=COUNTA(B1:B2)"


# ══════════════════════════════════════════════════════════════════════
# The spreadsheet computes the same numbers
# ══════════════════════════════════════════════════════════════════════


def _soffice():
    for candidate in ("soffice", "libreoffice"):
        found = shutil.which(candidate)
        if found:
            return found
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


def _recalculated(m, tmp_path, name):
    soffice = _soffice()
    if soffice is None:
        pytest.skip("LibreOffice not installed")
    path = tmp_path / f"{name}.xlsx"
    m.to_excel(path)
    subprocess.run(
        [soffice, f"-env:UserInstallation={(tmp_path / 'prof').as_uri()}",
         "--headless", "--convert-to", "xlsx", "--outdir",
         str(tmp_path / "calc"), str(path)],
        check=True, capture_output=True,
    )
    return load_workbook(tmp_path / "calc" / f"{name}.xlsx", data_only=True)


def _same(got, want):
    if isinstance(want, bool):
        return got == want
    if isinstance(want, date):
        return isinstance(got, datetime) and got.date() == want
    if isinstance(want, str):
        return got == want
    return got == pytest.approx(want)


@pytest.mark.slow
def test_the_workbook_recalculates_every_call_by_date(tmp_path):
    m = _model()
    with m:
        m.x.tw2 = ((m.x.rev + m.t.fee) * 2).set_display_name("tw2")
    x = _rows(_recalculated(m, tmp_path, "dated")["X"])
    for name in [*_CALLS, "tw2"]:
        want = getattr(m.x, name).value
        got = x[name][:4]
        assert len(got) == len(want) == 4, (name, got, want)
        assert all(_same(g, w) for g, w in zip(got, want)), (name, got, want)


@pytest.mark.slow
def test_the_workbook_recalculates_constants_keys_and_windows_of_their_own(tmp_path):
    m = _windows()
    with m:
        m.x.c_add = (m.x.rev + m.x.c).set_display_name("c_add")
        m.x.c_max = mo.MAX(m.x.rev, m.x.c).set_display_name("c_max")
        m.x.c_if = mo.IF(m.x.rev > 3, m.x.c, m.t.fee).set_display_name("c_if")
        m.x.k_max = mo.MAX(m.x.k, m.t.k2).set_display_name("k_max")
        m.x.k_add = (m.x.k + m.t.k2).set_display_name("k_add")
        m.t2 = mo.MultiVariable("T2", default_grain="month", default_start="2026-03",
                                default_periods=4)
        m.t2.y = mo.Variable(m.x.rev, start="2026-03", grain="month").set_display_name("y")
        m.t2.z = mo.Variable(m.x.rev * 2, start="2026-03", grain="month").set_display_name("z")
    book = _recalculated(m, tmp_path, "windows")
    x, t2 = _rows(book["X"]), _rows(book["T2"])
    for name in ("c_add", "c_max", "c_if", "k_max", "k_add"):
        want = getattr(m.x, name).value
        assert x[name][:len(want)] == pytest.approx(want), name
    for name in ("y", "z"):
        assert t2[name][:4] == pytest.approx(getattr(m.t2, name).value), name


@pytest.mark.slow
def test_the_workbook_recalculates_each_track_by_date(tmp_path):
    m = _tracked()
    with m:
        for name, (make, _, _) in _TRACKED.items():
            setattr(m.x, name, make(m).set_display_name(name))
    x = _rows(_recalculated(m, tmp_path, "tracked")["X"])
    for name, (_, plan, actual) in _TRACKED.items():
        assert x[name][:4] == pytest.approx(plan), name
        assert x[f"{name} · actual"][:4] == pytest.approx(actual), name


@pytest.mark.slow
def test_the_workbook_recalculates_a_window_written_without_its_zero(tmp_path):
    m = _unpadded()
    with m:
        m.x.a = (m.x.rev + m.t.fee).set_display_name("a")
        m.x.s = mo.SUM(m.x.rev, m.t.fee).set_display_name("s")
        m.x.i = mo.IF(m.x.rev > 2, m.t.fee, m.x.rev).set_display_name("i")
    x = _rows(_recalculated(m, tmp_path, "unpadded")["X"])
    for name in ("a", "s", "i"):
        assert x[name][:6] == pytest.approx(getattr(m.x, name).value), name


@pytest.mark.slow
def test_the_workbook_recalculates_lines_along_keys_key_by_key(tmp_path):
    m = _keyed()
    with m:
        m.p = (m.k + m.k3).set_display_name("p")
        m.pw = mo.POWER(m.k, m.k3).set_display_name("pw")
        m.sm = mo.SUM(m.k, m.k3).set_display_name("sm")
        m.mx = mo.MAX(m.k, m.k3).set_display_name("mx")
        m.f = mo.IF(m.k > 1, m.k, m.k3).set_display_name("f")
    rows = _rows(_recalculated(m, tmp_path, "keys")["M"])
    for name in ("p", "pw", "sm", "mx", "f"):
        assert rows[name][:3] == pytest.approx(getattr(m, name).value), name


@pytest.mark.slow
def test_the_workbook_recalculates_a_list_of_scalar_lines(tmp_path):
    m = _listed()
    with m:
        for name in _AGGREGATES:
            setattr(m.i, name.lower(), getattr(mo, name)([m.i.p, m.i.q])
                    .set_display_name(name.lower()))
    rows = _rows(_recalculated(m, tmp_path, "listed")["I"])
    for name in _AGGREGATES:
        assert rows[name.lower()][0] == pytest.approx(getattr(m.i, name.lower()).value), name


@pytest.mark.slow
def test_the_workbook_recalculates_a_loop_container_by_date(tmp_path):
    m = _lines()
    with m:
        for name, (make, _) in _LOOPS.items():
            mv = mo.MultiVariable("D_" + name, loop=True)
            setattr(m, "d_" + name, mv)
            with mv as d:
                make(m, d)
    book = _recalculated(m, tmp_path, "loops")
    for name, (_, expected) in _LOOPS.items():
        rows = _rows(book["D_" + name])
        for row, want in expected.items():
            assert rows[row.capitalize()][:4] == pytest.approx(want), (name, row)
