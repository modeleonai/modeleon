# SPDX-License-Identifier: Apache-2.0
"""The function registry: one place that says how each Excel function is
shaped, read by the Excel renderer and the loop container in place of
the name sets they used to keep by hand.

Every name the engine knew before the registry keeps the spec those
sets gave it — the old literals are kept here as the expected values,
and the three names that gained a helper since are named — and the
whole registry equals those names and the table of the names added
since, so no name is registered, or changes its shape, silently.

:func:`elementwise`, the shared per-period helper ``MIN`` / ``MAX``
moved onto, must keep their law: scalars broadcast as ``a + b`` does, a
one-value series broadcasts, other length mismatches are refused,
tracked operands lift per coordinate, an error marker is the period's
result, a shared unit is kept.
"""

from __future__ import annotations

import dataclasses
import datetime as dt
import logging
import shutil
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook

import modeleon as mo
from modeleon import functions
from modeleon.core.expr import FuncCall, ListExpr, Literal, VarRef
from modeleon.functions import registry
from modeleon.functions._helpers import elementwise, is_error, is_series
from modeleon.functions.registry import (
    FunctionSpec,
    admits_in_loop,
    elementwise_names,
    excel_name,
    file_prefix,
    file_prefixes,
    is_elementwise,
    is_passthrough,
    is_range_taking,
    names,
    passthrough_names,
    per_period_arity,
    range_positions,
    range_taking_names,
    register,
    spec,
)

# The sets the Excel renderer kept by hand (compile/excel/renderer.py)
# before the registry.
OLD_RANGE_TAKING = {
    "SUM", "AVERAGE", "MIN", "MAX", "COUNT",
    "IRR", "NPV", "XIRR", "XNPV",
    "SUMPRODUCT",
}
OLD_PASSTHROUGH = OLD_RANGE_TAKING | {
    "IF", "AND", "OR", "NOT", "CHOOSE", "ISBLANK", "ABS", "ROUND", "INT",
    "MOD",
    "EDATE", "EOMONTH", "YEAR", "MONTH", "DAY", "DATE", "DAYS360", "TODAY",
    "YEARFRAC", "DAYS",
    "LEN", "UPPER", "LOWER", "TEXT",
    "PMT", "FV", "PV", "RATE",
}
OLD_FILE_PREFIX = {"DAYS": "_xlfn."}
# The set the loop container kept by hand (core/loops.py).
OLD_ELEMENTWISE = {
    'IF', 'AND', 'OR', 'NOT', 'ABS', 'ROUND', 'INT', 'MOD', 'CHOOSE',
    'ISBLANK', 'EDATE', 'EOMONTH', 'YEAR', 'MONTH', 'DAY', 'DATE',
    'DAYS360', 'TODAY', 'LEN', 'UPPER', 'LOWER', 'CONCAT',
}
# Every name the engine knew before the registry.
OLD_NAMES = OLD_PASSTHROUGH | OLD_ELEMENTWISE
# Of the range-taking ones, these read whole series and give one number;
# the others take fixed arguments.
OLD_REDUCTIONS = {"SUM", "AVERAGE", "MIN", "MAX", "COUNT"}
# Written as given, without a helper, until they gained one; COUNT stays
# a reduction.
GAINED_A_HELPER = {"COUNT", "TEXT", "RATE"}
# One call per period now, so a loop container runs them: TEXT and RATE,
# plain calls before, with the helpers they gained; PMT, FV and PV, which
# already read a series period by period, since a call over one period's
# numbers gives that period's cell (an error as the result, never a raise).
NOW_ELEMENTWISE = {"TEXT", "RATE", "PMT", "FV", "PV"}

# The names added since: shape, file prefix, the argument
# positions written as ranges (None: every list argument).
NEW_NAMES: dict[str, tuple[str, str, tuple[int, ...] | None]] = {
    # Logical and error (functions/logical.py)
    "IFERROR": ("elementwise", "", None),
    "ISERROR": ("elementwise", "", None),
    "ISNUMBER": ("elementwise", "", None),
    "XOR": ("elementwise", "_xlfn.", None),
    "IFS": ("elementwise", "_xlfn.", None),
    "SWITCH": ("elementwise", "_xlfn.", None),
    # Math and rounding (functions/mathfn.py)
    "ROUNDUP": ("elementwise", "", None),
    "ROUNDDOWN": ("elementwise", "", None),
    "TRUNC": ("elementwise", "", None),
    "CEILING": ("elementwise", "", None),
    "FLOOR": ("elementwise", "", None),
    "POWER": ("elementwise", "", None),
    "SQRT": ("elementwise", "", None),
    "EXP": ("elementwise", "", None),
    "LN": ("elementwise", "", None),
    "LOG": ("elementwise", "", None),
    "LOG10": ("elementwise", "", None),
    "SIGN": ("elementwise", "", None),
    "PRODUCT": ("reduction", "", None),
    # Financial (functions/financial.py)
    "NPER": ("elementwise", "", None),
    "IPMT": ("elementwise", "", None),
    "PPMT": ("elementwise", "", None),
    "CUMIPMT": ("elementwise", "", None),
    "CUMPRINC": ("elementwise", "", None),
    "SLN": ("elementwise", "", None),
    "DB": ("elementwise", "", None),
    "DDB": ("elementwise", "", None),
    "SYD": ("elementwise", "", None),
    # Dates (functions/dates.py): the holidays are a range every period's
    # call reads.
    "DATEDIF": ("elementwise", "", None),
    "WEEKDAY": ("elementwise", "", None),
    "WEEKNUM": ("elementwise", "", None),
    "NETWORKDAYS": ("range", "", (2,)),
    "WORKDAY": ("range", "", (2,)),
    # Text (functions/text.py)
    "LEFT": ("elementwise", "", None),
    "RIGHT": ("elementwise", "", None),
    "MID": ("elementwise", "", None),
    "TRIM": ("elementwise", "", None),
    "FIND": ("elementwise", "", None),
    "SEARCH": ("elementwise", "", None),
    "SUBSTITUTE": ("elementwise", "", None),
    "VALUE": ("elementwise", "", None),
    "REPT": ("elementwise", "", None),
    "TEXTJOIN": ("range", "_xlfn.", None),
    # Conditional aggregates and counts (functions/criteria.py)
    "SUMIF": ("range", "", None),
    "SUMIFS": ("range", "", None),
    "COUNTIF": ("range", "", None),
    "COUNTIFS": ("range", "", None),
    "AVERAGEIF": ("range", "", None),
    "AVERAGEIFS": ("range", "", None),
    "MAXIFS": ("range", "_xlfn.", None),
    "MINIFS": ("range", "_xlfn.", None),
    "COUNTA": ("reduction", "", None),
    # Lookup and reference (functions/lookup.py); NA() is the not-found
    # value an XLOOKUP is written with, without a helper of its own.
    "MATCH": ("range", "", (1,)),
    "INDEX": ("range", "", (0,)),
    "XLOOKUP": ("range", "_xlfn.", (1, 2)),
    "VLOOKUP": ("range", "", (1,)),
    "HLOOKUP": ("range", "", (1,)),
    "NA": ("scalar", "", None),
    # Statistics (functions/statistics.py): the reductions read every cell
    # of one argument, and two values or more with a series among them one
    # period at a time; the others one range (RANK's at its second
    # argument), a list of two or more lines with a series among them a
    # range per period. The names an older workbook carries are
    # their own entries, written as themselves.
    "MEDIAN": ("reduction", "", None),
    "STDEV_S": ("reduction", "_xlfn.", None),
    "STDEV_P": ("reduction", "_xlfn.", None),
    "VAR_S": ("reduction", "_xlfn.", None),
    "VAR_P": ("reduction", "_xlfn.", None),
    "STDEV": ("reduction", "", None),
    "STDEVP": ("reduction", "", None),
    "VAR": ("reduction", "", None),
    "VARP": ("reduction", "", None),
    "PERCENTILE_INC": ("range", "_xlfn.", (0,)),
    "PERCENTILE_EXC": ("range", "_xlfn.", (0,)),
    "QUARTILE_INC": ("range", "_xlfn.", (0,)),
    "QUARTILE_EXC": ("range", "_xlfn.", (0,)),
    "PERCENTILE": ("range", "", (0,)),
    "QUARTILE": ("range", "", (0,)),
    "LARGE": ("range", "", (0,)),
    "SMALL": ("range", "", (0,)),
    "RANK_EQ": ("range", "_xlfn.", (1,)),
    "RANK_AVG": ("range", "_xlfn.", (1,)),
    "RANK": ("range", "", (1,)),
}
# The names the workbook spells otherwise: Excel's dot is an underscore in
# Python.
EXCEL_SPELLING = {
    "STDEV_S": "STDEV.S", "STDEV_P": "STDEV.P", "VAR_S": "VAR.S", "VAR_P": "VAR.P",
    "PERCENTILE_INC": "PERCENTILE.INC", "PERCENTILE_EXC": "PERCENTILE.EXC",
    "QUARTILE_INC": "QUARTILE.INC", "QUARTILE_EXC": "QUARTILE.EXC",
    "RANK_EQ": "RANK.EQ", "RANK_AVG": "RANK.AVG",
}
# The names written as given, without a helper of their own.
BARE = {"SUMPRODUCT", "NA"}
# The new names that are one call per period from two values, a
# series among them: every statistic, PRODUCT and COUNTA — every aggregate
# follows the rule.
PER_PERIOD = {
    n: 2 for n in (
        "MEDIAN", "STDEV_S", "STDEV_P", "VAR_S", "VAR_P", "STDEV", "STDEVP", "VAR", "VARP",
        "PERCENTILE_INC", "PERCENTILE_EXC", "QUARTILE_INC", "QUARTILE_EXC", "PERCENTILE",
        "QUARTILE", "LARGE", "SMALL", "RANK_EQ", "RANK_AVG", "RANK", "PRODUCT", "COUNTA",
    )
}
# Of the old names, the reductions, one call per period from two operands,
# a series among them: MIN, MAX, SUM, AVERAGE and COUNT.
OLD_PER_PERIOD = {"MIN", "MAX", "SUM", "AVERAGE", "COUNT"}


@pytest.fixture
def scratch():
    """Register throwaway names in a test; forget them after it."""
    created: list[str] = []

    def _register(name: str, **kwargs):
        created.append(name)
        return register(name, **kwargs)

    yield _register
    for name in created:
        registry._SPECS.pop(name, None)


def _rows(ws):
    out = {}
    for row in ws.iter_rows():
        label = row[0].value
        if label:
            out[str(label)] = [c.value for c in row[1:]]
    return out


def _book(model, tmp_path, sheet):
    path = tmp_path / "book.xlsx"
    model.to_excel(path)
    return _rows(load_workbook(path)[sheet])


def _call_var(name, args, value):
    """A list-valued Variable whose formula is a hand-built call."""
    v = mo.Variable(formula=FuncCall(name, args))
    v._value = value
    v.var_type = 'list' if isinstance(value, list) else 'scalar'
    v.value_type = 'float'
    return v


# ══════════════════════════════════════════════════════════════════════
# The names the engine knew before the registry keep their specs
# ══════════════════════════════════════════════════════════════════════


class TestTheOldNamesKeepTheirSpecs:
    def test_range_taking(self):
        assert range_taking_names() & OLD_NAMES == OLD_RANGE_TAKING

    def test_passthrough(self):
        assert passthrough_names() & OLD_NAMES == OLD_PASSTHROUGH

    def test_elementwise(self):
        assert elementwise_names() & OLD_NAMES == OLD_ELEMENTWISE | NOW_ELEMENTWISE

    def test_file_prefix(self):
        assert {n: p for n, p in file_prefixes().items() if n in OLD_NAMES} == OLD_FILE_PREFIX

    def test_shapes(self):
        for name in OLD_NAMES:
            if name in OLD_REDUCTIONS:
                expected = "reduction"
            elif name in OLD_RANGE_TAKING:
                expected = "range"
            elif name in OLD_ELEMENTWISE | NOW_ELEMENTWISE:
                expected = "elementwise"
            else:
                expected = "scalar"
            assert spec(name).shape == expected, name

    def test_loop_admission(self):
        per_period = OLD_ELEMENTWISE | NOW_ELEMENTWISE
        for name in OLD_NAMES:
            for n_args in range(4):
                expected = name in per_period or (name in OLD_PER_PERIOD and n_args >= 2)
                assert admits_in_loop(name, n_args) == expected, (name, n_args)

    def test_per_period_arity_is_the_reductions_over_two_operands(self):
        assert {n for n in OLD_NAMES if per_period_arity(n) is not None} == OLD_PER_PERIOD
        assert {per_period_arity(n) for n in OLD_PER_PERIOD} == {2}
        assert OLD_PER_PERIOD == OLD_REDUCTIONS

    def test_the_three_that_gained_a_helper(self):
        assert not any(spec(n).bare for n in GAINED_A_HELPER)
        assert all(callable(getattr(functions, n, None)) for n in GAINED_A_HELPER)
        assert spec("COUNT").shape == "reduction"
        assert {spec(n).shape for n in NOW_ELEMENTWISE & GAINED_A_HELPER} == {"elementwise"}


# ══════════════════════════════════════════════════════════════════════
# The registry is exactly the old names and the new names' table
# ══════════════════════════════════════════════════════════════════════


class TestTheRegistryIsTheTable:
    def test_nothing_else_is_registered(self):
        # CONCAT is the one element-wise name the renderer does not pass
        # through: it writes the ``&`` operator.
        assert names() == OLD_NAMES | set(NEW_NAMES)
        assert not OLD_NAMES & set(NEW_NAMES)
        assert all(excel_name(n) == n for n in OLD_NAMES)
        assert not is_passthrough("CONCAT") and is_elementwise("CONCAT")

    @pytest.mark.parametrize("name", sorted(NEW_NAMES))
    def test_each_new_name_has_its_spec(self, name):
        shape, prefix, positions = NEW_NAMES[name]
        assert spec(name) == FunctionSpec(
            name=name, shape=shape, excel=EXCEL_SPELLING.get(name, name), file_prefix=prefix,
            range_args=positions, per_period_arity=PER_PERIOD.get(name), bare=name in BARE,
        )

    def test_the_names_one_call_per_period_from_two_values(self):
        assert {n for n in names() if per_period_arity(n) is not None} == (
            OLD_PER_PERIOD | set(PER_PERIOD))

    def test_the_views_over_the_whole_registry(self):
        new_ranges = {n for n, (shape, _, _) in NEW_NAMES.items()
                      if shape in ("reduction", "range")}
        new_elementwise = {n for n, (shape, _, _) in NEW_NAMES.items() if shape == "elementwise"}
        new_prefixes = {n: prefix for n, (_, prefix, _) in NEW_NAMES.items() if prefix}
        assert range_taking_names() == OLD_RANGE_TAKING | new_ranges
        assert passthrough_names() == OLD_PASSTHROUGH | set(NEW_NAMES)
        assert elementwise_names() == OLD_ELEMENTWISE | NOW_ELEMENTWISE | new_elementwise
        assert file_prefixes() == OLD_FILE_PREFIX | new_prefixes


# ══════════════════════════════════════════════════════════════════════
# The queries
# ══════════════════════════════════════════════════════════════════════


class TestQueries:
    def test_shapes(self):
        assert spec("SUM").shape == "reduction"
        assert spec("NPV").shape == "range"
        assert spec("IF").shape == "elementwise"
        assert spec("PMT").shape == "elementwise"
        assert spec("YEARFRAC").shape == "scalar"
        assert is_range_taking("SUM") and is_range_taking("NPV")
        assert not is_range_taking("IF") and not is_range_taking("PMT")

    def test_every_existing_range_function_ranges_every_argument(self):
        assert all(range_positions(n) is None for n in OLD_RANGE_TAKING)

    def test_an_unregistered_name_answers_the_safe_way(self):
        assert spec("NOPE") is None
        assert not is_passthrough("NOPE")
        assert not is_range_taking("NOPE")
        assert not is_elementwise("NOPE")
        assert file_prefix("NOPE") == ""
        assert excel_name("NOPE") == "NOPE"
        assert range_positions("NOPE") is None
        assert per_period_arity("NOPE") is None
        assert not admits_in_loop("NOPE", 1)

    def test_bare_names_have_no_helper_and_every_other_name_has_one(self):
        bare = {n for n in names() if spec(n).bare}
        assert bare == BARE
        for name in names() - bare:
            assert hasattr(functions, name), name
        for name in bare:
            assert not hasattr(functions, name), name

    def test_every_excel_cased_export_is_registered(self):
        exported = {n for n in functions.__all__ if n.isupper()}
        assert exported == {n for n in names() if not spec(n).bare}

    def test_loop_admission(self):
        assert admits_in_loop("IF", 3) and admits_in_loop("ABS", 1)
        assert admits_in_loop("TODAY", 0)
        assert not admits_in_loop("SUM", 1) and not admits_in_loop("AVERAGE", 1)
        assert admits_in_loop("SUM", 2) and admits_in_loop("AVERAGE", 3)
        assert admits_in_loop("MAX", 2) and admits_in_loop("MIN", 3)
        assert not admits_in_loop("MAX", 1)
        assert not admits_in_loop("NPV", 2) and not admits_in_loop("SUMPRODUCT", 2)
        assert admits_in_loop("PMT", 3) and admits_in_loop("PV", 5)
        assert not admits_in_loop("YEARFRAC", 2)

    def test_the_prefix_and_the_excel_name(self):
        assert file_prefix("DAYS") == "_xlfn." and excel_name("DAYS") == "DAYS"
        assert file_prefix("SUM") == "" and excel_name("SUM") == "SUM"
        assert file_prefix("PERCENTILE_INC") == "_xlfn."
        assert excel_name("PERCENTILE_INC") == "PERCENTILE.INC"
        assert {n for n in names() if excel_name(n) != n} == set(EXCEL_SPELLING)


# ══════════════════════════════════════════════════════════════════════
# Registering
# ══════════════════════════════════════════════════════════════════════


class TestRegister:
    def test_returns_a_frozen_spec_with_the_defaults_filled(self, scratch):
        s = scratch("ZZPLAIN", shape="elementwise")
        assert s == FunctionSpec(name="ZZPLAIN", shape="elementwise", excel="ZZPLAIN")
        assert s.file_prefix == "" and s.range_args is None
        assert s.per_period_arity is None and s.passthrough and not s.bare
        with pytest.raises(dataclasses.FrozenInstanceError):
            s.shape = "scalar"

    def test_the_same_spec_twice_is_one_entry(self, scratch):
        first = scratch("ZZSAME", shape="scalar", bare=True)
        again = scratch("ZZSAME", shape="scalar", bare=True)
        assert first == again
        assert "ZZSAME" in names()

    def test_a_different_spec_for_the_same_name_is_refused(self, scratch):
        scratch("ZZDIFF", shape="elementwise")
        with pytest.raises(ValueError, match="ZZDIFF.*already registered"):
            register("ZZDIFF", shape="scalar")
        assert spec("ZZDIFF").shape == "elementwise"

    def test_excel_name_and_prefix(self, scratch):
        scratch("ZZ_INC", shape="elementwise", excel="ZZ.INC", file_prefix="_xlfn.")
        assert excel_name("ZZ_INC") == "ZZ.INC"
        assert file_prefix("ZZ_INC") == "_xlfn."
        assert file_prefixes()["ZZ_INC"] == "_xlfn."
        assert is_passthrough("ZZ_INC")

    def test_range_positions(self, scratch):
        scratch("ZZPOS", shape="range", range_args=[0, 2])
        scratch("ZZALL", shape="range")
        assert range_positions("ZZPOS") == (0, 2)
        assert range_positions("ZZALL") is None
        assert is_range_taking("ZZPOS") and is_range_taking("ZZALL")
        assert "ZZPOS" in range_taking_names()

    def test_per_period_arity(self, scratch):
        scratch("ZZCAP", shape="reduction", per_period_arity=3)
        assert per_period_arity("ZZCAP") == 3
        assert admits_in_loop("ZZCAP", 3) and not admits_in_loop("ZZCAP", 2)
        # A range function counts the lines its ranges list.
        scratch("ZZPICK", shape="range", range_args=[0], per_period_arity=2)
        assert per_period_arity("ZZPICK") == 2
        assert admits_in_loop("ZZPICK", 2) and not admits_in_loop("ZZPICK", 1)

    def test_not_passthrough_and_bare(self, scratch):
        scratch("ZZOP", shape="elementwise", passthrough=False)
        scratch("ZZBARE", shape="scalar", bare=True)
        assert not is_passthrough("ZZOP") and is_elementwise("ZZOP")
        assert "ZZOP" not in passthrough_names() and "ZZOP" in elementwise_names()
        assert spec("ZZBARE").bare and is_passthrough("ZZBARE")

    @pytest.mark.parametrize("kwargs, message", [
        ({"shape": "table"}, "shape"),
        ({"shape": "elementwise", "range_args": [0]}, "range_args"),
        ({"shape": "range", "range_args": [-1]}, "range_args"),
        ({"shape": "reduction", "per_period_arity": 1}, "per_period_arity"),
        ({"shape": "range", "per_period_arity": 1}, "per_period_arity"),
        ({"shape": "elementwise", "per_period_arity": 2}, "per_period_arity"),
        ({"shape": "scalar", "per_period_arity": 2}, "per_period_arity"),
    ])
    def test_a_bad_spec_is_refused(self, kwargs, message):
        with pytest.raises(ValueError, match=message):
            register("ZZBAD", **kwargs)
        assert spec("ZZBAD") is None

    @pytest.mark.parametrize("name", ["", "sum", "Sum", "_xlfn.DAYS", "A B", "1X"])
    def test_a_name_must_be_an_upper_case_identifier(self, name):
        with pytest.raises(ValueError, match="upper-case"):
            register(name, shape="elementwise")


# ══════════════════════════════════════════════════════════════════════
# The Excel renderer reads the registry
# ══════════════════════════════════════════════════════════════════════


def _model():
    m = mo.Model("m", default_grain="year", default_start="2026", default_periods=3)
    with m:
        m.i = mo.MultiVariable("I")
        m.i.d0 = mo.Variable(dt.date(2026, 1, 1), display_name="d0")
        m.i.d1 = mo.Variable(dt.date(2026, 3, 1), display_name="d1")
        m.x = mo.MultiVariable("X")
        m.x.a = mo.Variable([3.0, 8.0, 5.0], display_name="a")
        m.x.k = mo.Variable([10.0, 20.0, 30.0], display_name="k")
    return m


class TestTheRendererReadsTheRegistry:
    def test_the_file_prefix_is_written(self, tmp_path):
        m = _model()
        with m:
            m.i.span = mo.DAYS(m.i.d1, m.i.d0).set_display_name("span")
        i = _book(m, tmp_path, "I")
        assert i["span"][0] == "=_xlfn.DAYS(B2, B1)", i["span"]

    def test_the_excel_name_and_its_prefix_are_written(self, scratch, tmp_path):
        scratch("ZZ_INC", shape="elementwise", excel="ZZ.INC", file_prefix="_xlfn.")
        m = _model()
        with m:
            m.x.z = _call_var("ZZ_INC", [VarRef(m.x.a)], [3.0, 8.0, 5.0])
            m.x.z.set_display_name("z")
        x = _book(m, tmp_path, "X")
        assert x["z"] == ["=_xlfn.ZZ.INC(B2)", "=_xlfn.ZZ.INC(C2)", "=_xlfn.ZZ.INC(D2)"], x["z"]

    def test_a_range_position_is_a_range_and_the_others_this_period_s_cell(
            self, scratch, tmp_path):
        scratch("ZZPOS", shape="range", range_args=[0])
        m = _model()
        with m:
            m.x.z = _call_var("ZZPOS", [VarRef(m.x.a), VarRef(m.x.k)], [26.0, 36.0, 46.0])
            m.x.z.set_display_name("z")
        x = _book(m, tmp_path, "X")
        assert x["z"] == ["=ZZPOS(B2:D2, B3)", "=ZZPOS(B2:D2, C3)", "=ZZPOS(B2:D2, D3)"], x["z"]

    def test_every_argument_is_a_range_when_no_positions_are_given(self, scratch, tmp_path):
        scratch("ZZALL", shape="range")
        m = _model()
        with m:
            m.i.z = _call_var("ZZALL", [VarRef(m.x.a), VarRef(m.x.k)], 76.0)
            m.i.z.set_display_name("z")
        i = _book(m, tmp_path, "I")
        assert i["z"][0] == "=ZZALL(X!B2:D2, X!B3:D3)", i["z"]

    def test_a_per_period_literal_at_a_cell_position_is_not_an_unlaid_list(
            self, scratch, tmp_path):
        # Only range positions can make the call fall back to its value:
        # a list literal at a cell position is this period's number.
        scratch("ZZPOS", shape="range", range_args=[0])
        m = _model()
        with m:
            m.x.z = _call_var("ZZPOS", [VarRef(m.x.a), Literal([1.0, 2.0, 3.0])],
                              [17.0, 18.0, 19.0])
            m.x.z.set_display_name("z")
        x = _book(m, tmp_path, "X")
        assert x["z"] == ["=ZZPOS(B2:D2, 1.0)", "=ZZPOS(B2:D2, 2.0)", "=ZZPOS(B2:D2, 3.0)"], x["z"]

    def test_a_list_of_lines_at_a_per_period_range_is_this_column_s_run(
            self, scratch, tmp_path):
        # A range function one call per period: the lines a range lists are
        # the run of their cells in this period's column, another argument
        # this period's cell.
        scratch("ZZPICK", shape="range", range_args=[0], per_period_arity=2)
        m = _model()
        with m:
            m.x.z = _call_var("ZZPICK", [ListExpr([VarRef(m.x.a), VarRef(m.x.k)]),
                                         VarRef(m.x.a)], [3.0, 8.0, 5.0])
            m.x.z.set_display_name("z")
            m.x.one = _call_var("ZZPICK", [VarRef(m.x.a), VarRef(m.x.k)], [3.0, 8.0, 5.0])
            m.x.one.set_display_name("one")
        x = _book(m, tmp_path, "X")
        assert x["z"] == ["=ZZPICK(B2:B3, B2)", "=ZZPICK(C2:C3, C2)", "=ZZPICK(D2:D3, D2)"], x["z"]
        # One line at the range position is the whole row, as before.
        assert x["one"] == ["=ZZPICK(B2:D2, B3)", "=ZZPICK(B2:D2, C3)", "=ZZPICK(B2:D2, D3)"]

    def test_an_unknown_name_is_written_as_given_and_points_at_the_registry(
            self, tmp_path, caplog):
        m = _model()
        with m:
            m.x.z = _call_var("ZZNOPE", [VarRef(m.x.a)], [3.0, 8.0, 5.0])
            m.x.z.set_display_name("z")
        with caplog.at_level(logging.WARNING, logger="modeleon"):
            x = _book(m, tmp_path, "X")
        assert x["z"] == ["=ZZNOPE(B2)", "=ZZNOPE(C2)", "=ZZNOPE(D2)"], x["z"]
        assert any("ZZNOPE" in r.getMessage() and "registry" in r.getMessage()
                   for r in caplog.records)

    def test_the_old_paths_still_hold(self, tmp_path):
        m = _model()
        with m:
            m.x.lo = mo.MIN(m.x.a, m.x.k).set_display_name("lo")
            m.x.r = mo.ROUND(m.x.a / 3, 1).set_display_name("r")
            m.i.total = mo.SUM(m.x.a).set_display_name("total")
        x = _book(m, tmp_path, "X")
        i = _book(m, tmp_path, "I")
        assert x["lo"] == ["=MIN(B2, B3)", "=MIN(C2, C3)", "=MIN(D2, D3)"], x["lo"]
        assert x["r"] == ["=ROUND(B2 / 3, 1)", "=ROUND(C2 / 3, 1)", "=ROUND(D2 / 3, 1)"], x["r"]
        assert i["total"][0] == "=SUM(X!B2:D2)", i["total"]


# ══════════════════════════════════════════════════════════════════════
# The loop container reads the registry
# ══════════════════════════════════════════════════════════════════════


class TestTheLoopReadsTheRegistry:
    def test_an_elementwise_function_runs_period_by_period(self):
        m = mo.Model("m", default_grain="year", default_start="2026", default_periods=3)
        with m:
            m.d = mo.MultiVariable("D", loop=True)
            with m.d as d:
                d.bal = mo.lag(d.r, fill=10.0) * 1.37
                d.cap = mo.MIN(d.bal, 15.0)
                d.r = mo.ROUND(d.cap, 1)
        assert d.bal.value == pytest.approx([13.7, 18.769, 20.55])
        assert d.cap.value == pytest.approx([13.7, 15.0, 15.0])
        assert d.r.value == pytest.approx([13.7, 15.0, 15.0])

    def test_a_reduction_is_refused(self):
        m = mo.Model("m", default_grain="year", default_start="2026", default_periods=3)
        with pytest.raises(ValueError, match="AVERAGE reads a whole series"):
            with m:
                m.d = mo.MultiVariable("D", loop=True)
                with m.d as d:
                    d.bal = mo.lag(d.x, fill=10.0) * 1.37
                    d.avg = mo.AVERAGE(d.bal)
                    d.x = d.bal - d.avg

    def test_a_name_the_engine_lacks_is_refused_by_name(self):
        m = mo.Model("m", default_grain="year", default_start="2026", default_periods=3)
        with pytest.raises(ValueError, match="ZZNOPE cannot be run one period at a time"):
            with m:
                m.d = mo.MultiVariable("D", loop=True)
                with m.d as d:
                    d.bal = mo.lag(d.z, fill=10.0) * 1.37
                    d.z = _call_var("ZZNOPE", [VarRef(d.bal)], [1.0, 1.0, 1.0])


# ══════════════════════════════════════════════════════════════════════
# The shared per-period helper
# ══════════════════════════════════════════════════════════════════════


def _hypot(a, b):
    return (a * a + b * b) ** 0.5


class TestElementwiseHelper:
    def test_scalars_only_is_one_call(self):
        m = _model()
        with m:
            m.i.p = mo.Variable(3.0)
        out = elementwise("ZZHYP", _hypot, [m.i.p, 4.0])
        assert out.value == 5.0 and out.var_type == "scalar"
        assert out.formula == "ZZHYP(p, 4.0)"
        assert isinstance(out._expr, FuncCall) and out._expr.func == "ZZHYP"
        assert isinstance(out._expr.args[0], VarRef)
        assert isinstance(out._expr.args[1], Literal) and out._expr.args[1].value == 4.0

    def test_a_series_gives_a_series_and_scalars_broadcast(self):
        a = mo.Variable([3.0, 6.0, 5.0])
        assert elementwise("ZZHYP", _hypot, [a, 4.0]).value == [5.0, _hypot(6.0, 4.0), _hypot(5.0, 4.0)]
        assert elementwise("ZZHYP", _hypot, [4.0, a]).var_type == "list"
        assert elementwise("ZZHYP", _hypot, [a, [4.0, 8.0, 12.0]]).value == [5.0, 10.0, 13.0]

    def test_a_one_value_series_broadcasts_and_other_lengths_are_refused(self):
        a = mo.Variable([3.0, 6.0, 5.0])
        assert elementwise("ZZHYP", _hypot, [a, mo.Variable([4.0])]).value[0] == 5.0
        with pytest.raises(ValueError, match="different lengths"):
            elementwise("ZZHYP", _hypot, [a, mo.Variable([4.0, 8.0])])

    def test_tracked_operands_lift_per_track(self):
        m = mo.Model("m", tracks=mo.Tracks("plan", "actual"),
                     default_grain="month", default_start="2026-01", default_periods=2)
        with m:
            m.x = mo.Variable(plan=[3.0, 6.0], actual=[5.0, 8.0])
            m.h = elementwise("ZZHYP", _hypot, [m.x, 4.0])
        assert dict(m.h._value.items()) == {"plan": [5.0, _hypot(6.0, 4.0)],
                                            "actual": [_hypot(5.0, 4.0), _hypot(8.0, 4.0)]}
        assert m.h.var_type == "list"

    def test_an_error_marker_is_the_period_s_result_first_in_reading_order(self):
        e = mo.Variable([2.0, 2.0, 2.0]) / mo.Variable([1.0, 0.0, 2.0])
        out = elementwise("ZZHYP", _hypot, [e, mo.Variable([1.0, "#VALUE!", 1.0])])
        assert out.value == [_hypot(2.0, 1.0), "#DIV/0!", _hypot(1.0, 1.0)]

    def test_errors_can_be_handed_to_the_function_instead(self):
        e = mo.Variable([2.0, "#DIV/0!"])
        caught = elementwise("ZZCATCH", lambda v, alt: alt if is_error(v) else v,
                             [e, 0.0], propagate_errors=False)
        assert caught.value == [2.0, 0.0]

    def test_a_blank_reaches_the_function_unless_a_stand_in_is_given(self):
        b = mo.Variable([None, 3.0])
        seen = elementwise("ZZSEEN", lambda v: v is None, [b], value_type='bool')
        assert seen.value == [True, False]
        assert elementwise("ZZZERO", lambda v: v + 1, [b], blank=0).value == [1, 4.0]

    def test_a_failed_call_is_the_spreadsheet_s_error(self):
        assert elementwise("ZZDIV", lambda a, b: a / b, [mo.Variable([1.0, 1.0]), [0.0, 2.0]]
                           ).value == ["#DIV/0!", 0.5]
        assert elementwise("ZZTYPE", lambda a: a + 1, [mo.Variable(["x"])]).value == ["#VALUE!"]
        assert elementwise("ZZBIG", lambda a: float(10 ** a), [mo.Variable([400])]
                           ).value == ["#NUM!"]

    def test_exceptions_raise_lets_a_failed_call_propagate(self):
        a = mo.Variable([1.0, 1.0])
        with pytest.raises(ZeroDivisionError):
            elementwise("ZZDIV", lambda x, y: x / y, [a, [0.0, 2.0]], exceptions='raise')
        with pytest.raises(TypeError):
            elementwise("ZZTYPE", lambda x: x + 1, [mo.Variable(["x"])], exceptions='raise')

    def test_a_function_that_cannot_take_the_operands_is_refused_at_entry(self):
        with pytest.raises(TypeError, match=r"ZZONE: the function given takes \(a\), "
                                             r"but the call has 2 operands"):
            elementwise("ZZONE", lambda a: a, [mo.Variable([1.0, 2.0]), 1.0])
        with pytest.raises(TypeError, match="has 1 operand"):
            elementwise("ZZTWO", _hypot, [mo.Variable([1.0, 2.0])])
        # A builtin without a signature is trusted.
        assert elementwise("ZZMAX", max, [mo.Variable([1.0, 5.0]), 3.0]).value == [3.0, 5.0]

    def test_value_type_follows_the_first_variable_unless_given(self):
        assert elementwise("ZZHYP", _hypot, [mo.Variable([3, 4]), 4]).value_type == "int"
        assert elementwise("ZZHYP", _hypot, [3.0, 4.0]).value_type == "float"
        assert elementwise("ZZLEN", lambda s: len(s), [mo.Variable(["ab"])],
                           value_type='int').value_type == "int"

    def test_a_shared_unit_is_kept_and_a_mixed_one_dropped(self):
        a = mo.Variable([3.0, 8.0], unit="$")
        b = mo.Variable([4.0, 4.0], unit="$")
        c = mo.Variable([4.0, 4.0], unit="hr")
        assert str(elementwise("ZZHYP", _hypot, [a, b])._unit) == "$"
        assert str(elementwise("ZZHYP", _hypot, [a, 1.0])._unit) == "$"
        assert elementwise("ZZHYP", _hypot, [a, c])._unit is None
        # A count or a span is not in the operands' unit.
        assert elementwise("COUNT", lambda *v: len(v), [a, b], value_type='int')._unit is None

    def test_source_code_and_render_backends_are_forwarded(self):
        a = mo.Variable([3.0, 8.0])
        out = elementwise("ZZHYP", _hypot, [a, 4.0], source_code="hyp(a, 4)",
                          render_backends=frozenset({'excel'}))
        assert out._source_code == "hyp(a, 4)"
        assert out._expr.render_backends == frozenset({'excel'})

    def test_an_axised_operand_is_refused(self):
        region = mo.Variable(["N", "S"], display_name="region")
        by_region = mo.Variable([1.0, 2.0], indexed_by=[region])
        with pytest.raises(ValueError, match="ZZHYP"):
            elementwise("ZZHYP", _hypot, [by_region, 4.0])

    def test_is_series_and_is_error(self):
        assert is_series([1.0]) and is_series(mo.Variable([1.0, 2.0]))
        assert not is_series(1.0) and not is_series(mo.Variable(1.0))
        assert is_error("#DIV/0!") and is_error("#N/A")
        assert not is_error("text") and not is_error(None) and not is_error(0)

    def test_min_and_max_still_use_it(self):
        a = mo.Variable([3.0, 8.0, 5.0])
        assert mo.MIN(a, 4.0).value == [3.0, 4.0, 4.0]
        assert mo.MAX(a, mo.Variable([None, None, None])).value == [3.0, 8.0, 5.0]

    def test_min_and_max_refuse_a_text_operand_at_build_time(self):
        # The spreadsheet skips text in a referenced cell, so =MAX(B2, B7)
        # over words is a number no marker could match; the call is
        # refused as it always was, and nothing is written.
        a = mo.Variable([3.0, -8.0, 0.0, 5.5])
        with pytest.raises(TypeError):
            mo.MAX(a, mo.Variable(["x", "Yy", "", "zZz"]))
        with pytest.raises(TypeError):
            mo.MIN(a, mo.Variable("abc"))
        with pytest.raises(TypeError):
            mo.MAX(a, ["x", "Yy", "", "zZz"])


# ══════════════════════════════════════════════════════════════════════
# The spreadsheet computes the same numbers
# ══════════════════════════════════════════════════════════════════════


def _soffice():
    for candidate in ("soffice", "libreoffice"):
        found = shutil.which(candidate)
        if found:
            return found
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


@pytest.mark.slow
def test_the_workbook_recalculates_to_the_python_values(tmp_path):
    """One function of each registry shape, written through the registry
    and recalculated by LibreOffice: a reduction over a range, a
    per-period cap, an element-wise call, a prefixed name, a range
    function with a cell argument, a scalar financial call."""
    soffice = _soffice()
    if soffice is None:
        pytest.skip("LibreOffice not installed")
    m = _model()
    with m:
        m.x.lo = mo.MIN(m.x.a, m.x.k).set_display_name("lo")
        m.x.r = mo.ROUND(m.x.a / 3, 1).set_display_name("r")
        m.x.pos = mo.ABS(m.x.a - m.x.k).set_display_name("pos")
        m.i.total = mo.SUM(m.x.a).set_display_name("total")
        m.i.span = mo.DAYS(m.i.d1, m.i.d0).set_display_name("span")
        m.i.rate = mo.Variable(0.1, display_name="rate")
        m.i.npv = mo.NPV(m.i.rate, m.x.k).set_display_name("npv")
        m.i.pmt = mo.PMT(m.i.rate, 3, 100.0).set_display_name("pmt")
    path = tmp_path / "registry.xlsx"
    m.to_excel(path)
    subprocess.run(
        [soffice, f"-env:UserInstallation={(tmp_path / 'prof').as_uri()}",
         "--headless", "--convert-to", "xlsx", "--outdir",
         str(tmp_path / "calc"), str(path)],
        check=True, capture_output=True,
    )
    book = load_workbook(tmp_path / "calc" / "registry.xlsx", data_only=True)
    x = _rows(book["X"])
    i = _rows(book["I"])
    for name in ("lo", "r", "pos"):
        assert [round(v, 6) for v in x[name][:3]] == getattr(m.x, name).value, name
    assert i["total"][0] == m.i.total.value == 16.0
    assert i["span"][0] == m.i.span.value == 59
    assert round(i["npv"][0], 6) == round(m.i.npv.value, 6)
    assert round(i["pmt"][0], 6) == round(m.i.pmt.value, 6)
