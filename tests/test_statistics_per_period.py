# SPDX-License-Identifier: Apache-2.0
"""Statistics over several series are one call per period.

The rule ``MIN`` / ``MAX`` follow: several series are combined period by
period, as ``a + b`` combines them; one series alone is reduced over its
periods; scalars only stay one number. ``mo.MEDIAN(a, b)`` over two rows
is the median of each period's two cells — ``=MEDIAN(C5, C9)`` in every
column — and a list of two or more rows (one field down the records of a
table, or sibling lines) is one sample per period: the median, the
percentile, the k-th largest and each line's rank of this period's cells,
written as the run those cells form in the column
(``=_xlfn.PERCENTILE.INC(Peers!B3:B5, 0.25)``). A scalar, a scalar line
and a one-value series stand in every period; series of other lengths
are refused; a line with tracks lifts per track; a loop container runs
the call one period at a time.
"""

import datetime as dt
import math

import pytest
from openpyxl import load_workbook

import modeleon as mo
from modeleon.compile import check_compat
from modeleon.core.expr import FuncCall, ListExpr, Literal, VarRef
from modeleon.functions.registry import admits_in_loop, per_period_arity

A = [3.0, 8.0, 5.0]
B = [4.0, 4.0, 6.0]
C = [1.0, 9.0, 5.0]


def _rows(ws):
    out = {}
    for row in ws.iter_rows():
        label = row[0].value
        if label:
            out[str(label).split(" (")[0]] = [c.value for c in row[1:]]
    return out


def _book(model, tmp_path):
    path = tmp_path / "book.xlsx"
    model.to_excel(path)
    return load_workbook(path)


class Peer(mo.MultiVariableClass):
    def compute(self, rev):
        self.rev = mo.Variable(rev, display_name="Revenue", unit="USD")


def _model():
    """Inputs on I; three sibling lines a, b, c on X (rows 2–4, periods in
    B–D); a records table of three peers with a revenue per year on Peers
    (rows 3–5); a Summary sheet."""
    m = mo.Model("m", display_name="M", default_grain="year", default_start="2026",
                 default_periods=3)
    with m:
        m.i = mo.MultiVariable("I")
        m.i.floor = mo.Variable(5.0, display_name="floor")
        m.i.k = mo.Variable(0.25, display_name="k")
        m.x = mo.MultiVariable("X")
        m.x.a = mo.Variable(list(A), display_name="a")
        m.x.b = mo.Variable(list(B), display_name="b")
        m.x.c = mo.Variable(list(C), display_name="c")
        m.peers = mo.MultiVariable("Peers", excel_props={"tab": True, "orient": "records"})
        with m.peers as p:
            p.a = Peer(rev=[10.0, 30.0, 25.0], display_name="A")
            p.b = Peer(rev=[20.0, 10.0, 45.0], display_name="B")
            p.c = Peer(rev=[35.0, 40.0, 10.0], display_name="C")
        m.s = mo.MultiVariable("Summary", excel_props={"tab": True})
    return m


def _revs(m):
    p = m.peers
    return [p.a.rev, p.b.rev, p.c.rev]


def _lines(m):
    return [m.x.a, m.x.b, m.x.c]


# ══════════════════════════════════════════════════════════════════════
# The reductions: MEDIAN and the spreads
# ══════════════════════════════════════════════════════════════════════


class TestTheReductionsPerPeriod:
    def test_two_series_are_combined_period_by_period(self):
        a, b = mo.Variable(list(A)), mo.Variable(list(B))
        out = mo.MEDIAN(a, b)
        assert out.value == [3.5, 6.0, 5.5]
        assert out.var_type == "list" and out.value_type == "float"
        assert mo.VAR_P(a, b).value == [0.25, 4.0, 0.25]
        assert mo.STDEV_S(a, b).value == pytest.approx(
            [math.sqrt(0.5), math.sqrt(8), math.sqrt(0.5)])
        assert mo.STDEV_P(a, b).value == [0.5, 2.0, 0.5]
        assert mo.VAR_S(a, b).value == [0.5, 8.0, 0.5]

    def test_three_series(self):
        a, b, c = (mo.Variable(list(v)) for v in (A, B, C))
        assert mo.MEDIAN(a, b, c).value == [3.0, 8.0, 5.0]
        assert mo.VARP(a, b, c).value == pytest.approx(
            [((3 - 8 / 3) ** 2 + (4 - 8 / 3) ** 2 + (1 - 8 / 3) ** 2) / 3,
             ((8 - 7) ** 2 + (4 - 7) ** 2 + (9 - 7) ** 2) / 3,
             ((5 - 16 / 3) ** 2 + (6 - 16 / 3) ** 2 + (5 - 16 / 3) ** 2) / 3])

    def test_the_older_names_follow_the_same_rule(self):
        a, b = mo.Variable(list(A)), mo.Variable(list(B))
        assert mo.STDEV(a, b).value == mo.STDEV_S(a, b).value
        assert mo.STDEVP(a, b).value == mo.STDEV_P(a, b).value
        assert mo.VAR(a, b).value == mo.VAR_S(a, b).value
        assert mo.VARP(a, b).value == mo.VAR_P(a, b).value

    def test_a_scalar_broadcasts_on_either_side_as_a_plus_b(self):
        a = mo.Variable(list(A))
        assert mo.MEDIAN(a, 5.0).value == [4.0, 6.5, 5.0]
        assert mo.MEDIAN(5.0, a).value == [4.0, 6.5, 5.0]
        assert mo.MEDIAN(a, mo.Variable(5.0)).value == [4.0, 6.5, 5.0]
        assert mo.STDEV_P(a, 5.0).value == [1.0, 1.5, 0.0]

    def test_a_one_value_series_broadcasts_and_other_lengths_are_refused(self):
        a = mo.Variable(list(A))
        assert mo.MEDIAN(mo.Variable([5.0]), a).value == [4.0, 6.5, 5.0]
        with pytest.raises(ValueError, match="different lengths"):
            mo.MEDIAN(mo.Variable([1.0, 2.0]), a)
        with pytest.raises(ValueError, match="different lengths"):
            mo.VAR_S(a, [1.0, 2.0])

    def test_a_list_typed_in_is_a_series_as_the_operators_read_it(self):
        a = mo.Variable(list(A))
        assert mo.MEDIAN(a, [1.0, 9.0, 1.0]).value == [2.0, 8.5, 3.0]
        assert mo.MEDIAN([1, 2, 3], [4, 5, 6]).value == [2.5, 3.5, 4.5]

    def test_one_series_alone_is_reduced_over_its_periods(self):
        a = mo.Variable(list(A))
        out = mo.MEDIAN(a)
        assert out.value == 5.0 and out.var_type == "scalar"
        assert mo.MEDIAN(list(A)).value == 5.0
        assert mo.VAR_P(a).value == pytest.approx(38 / 9)

    def test_scalars_only_stay_one_number(self):
        out = mo.MEDIAN(mo.Variable(3.0), mo.Variable(4.0))
        assert out.value == 3.5 and out.var_type == "scalar"
        assert mo.MEDIAN(1, 2, 3, 10).value == 2.5
        m = _model()
        field = [m.peers.a.rev[0], m.peers.b.rev[0], m.peers.c.rev[0]]
        assert mo.MEDIAN(field).value == 20.0              # one element of each row

    def test_a_list_of_series_lines_is_one_sample_per_period(self):
        m = _model()
        assert mo.MEDIAN(_lines(m)).value == mo.MEDIAN(m.x.a, m.x.b, m.x.c).value
        assert mo.MEDIAN(_revs(m)).value == [20.0, 30.0, 25.0]
        assert mo.STDEV_S(_revs(m)).value == pytest.approx(
            [12.583057392117917, 15.275252316519467, 17.559422921421234])
        assert mo.MEDIAN(_revs(m)).var_type == "list"

    def test_a_field_of_scalar_lines_with_a_series_is_read_whole_in_every_period(self):
        cells = [mo.Variable(1.0), mo.Variable(2.0), mo.Variable(10.0)]
        a = mo.Variable(list(A))
        assert mo.MEDIAN(cells, a).value == [2.5, 5.0, 3.5]
        assert mo.MEDIAN(a, (1.0, 2.0, 10.0)).value == [2.5, 5.0, 3.5]    # a typed array

    def test_a_list_holding_one_series_is_refused_with_both_readings(self):
        m = _model()
        with pytest.raises(ValueError, match=r"two or more") as err:
            mo.MEDIAN([m.x.a])
        assert "mo.MEDIAN(x.a) reduces the line over its periods" in str(err.value)
        with pytest.raises(ValueError, match=r"two or more") as err:
            mo.RANK_EQ(m.x.a, (m.x.a,))
        assert "mo.RANK_EQ(number, x.a, order)" not in str(err.value)
        assert "mo.RANK_EQ(number, x.a) reads the line over its periods" in str(err.value)

    def test_a_range_is_lines_and_values_never_a_list_inside_it(self):
        a = mo.Variable(list(A))
        with pytest.raises(TypeError, match="one value"):
            mo.MEDIAN([a, mo.Variable(list(B)), [1.0, 2.0]])

    def test_this_period_s_cells_as_the_spreadsheet_reads_them(self):
        # In a line a blank, text and a boolean are skipped; typed into the
        # call a number, TRUE and numeric text count.
        mixed = mo.Variable([None, "x", True])
        b = mo.Variable([1.0, 2.0, 3.0])
        assert mo.MEDIAN(mixed, b).value == [1.0, 2.0, 3.0]
        assert mo.MEDIAN(b, True).value == [1.0, 1.5, 2.0]
        assert mo.MEDIAN(b, "3").value == [2.0, 2.5, 3.0]
        assert mo.MEDIAN(b, None).value == [0.5, 1.0, 1.5]
        assert mo.MEDIAN(b, "abc").value == ["#VALUE!"] * 3

    def test_an_error_is_the_period_s_result_the_first_in_reading_order(self):
        e = mo.Variable([2.0, 2.0, 2.0]) / mo.Variable([1.0, 0.0, 2.0])
        na = mo.Variable(["#N/A", 1.0, 1.0])
        assert mo.MEDIAN(e, 4.0).value == [3.0, "#DIV/0!", 2.5]
        assert mo.MEDIAN(na, e).value == ["#N/A", "#DIV/0!", 1.0]
        assert mo.STDEV_S(e, na).value == ["#N/A", "#DIV/0!", 0.0]

    def test_too_few_numbers_in_a_period(self):
        blank = mo.Variable([None, 2.0, "x"])
        b = mo.Variable([None, None, 5.0])
        assert mo.MEDIAN(blank, b).value == ["#NUM!", 2.0, 5.0]
        assert mo.STDEV_S(blank, b).value == ["#DIV/0!", "#DIV/0!", "#DIV/0!"]
        assert mo.VAR_P(blank, b).value == ["#DIV/0!", 0.0, 0.0]

    def test_dates_give_a_date_each_period(self):
        d1 = mo.Variable([dt.date(2026, 1, 31), dt.date(2026, 2, 28)])
        d2 = mo.Variable([dt.date(2026, 3, 31), dt.date(2026, 3, 2)])
        out = mo.MEDIAN(d1, d2, dt.date(2026, 2, 28))
        assert out.value == [dt.date(2026, 2, 28), dt.date(2026, 2, 28)]
        assert out.value_type == "datetime"

    def test_the_values_are_in_one_unit_and_a_variance_in_its_square(self):
        usd = mo.Variable([1.0, 2.0, 4.0], unit="USD")
        more = mo.Variable([3.0, 2.0, 1.0], unit="USD")
        assert str(mo.MEDIAN(usd, more)._unit) == "USD"
        assert str(mo.STDEV_S(usd, more, 5.0)._unit) == "USD"
        assert str(mo.VAR_S(usd, more)._unit) == "USD^2"
        with pytest.raises(ValueError, match="mo.MEDIAN takes its values in one unit"):
            mo.MEDIAN(usd, mo.Variable([1.0, 1.0, 1.0], unit="hr"))
        with pytest.raises(ValueError, match="mo.VAR_P takes its values in one unit"):
            mo.VAR_P([usd, mo.Variable([1.0, 1.0, 1.0], unit="EUR")])

    def test_a_line_with_tracks_lifts_per_track(self):
        m = mo.Model("m", tracks=mo.Tracks("plan", "actual"), default_grain="month",
                     default_start="2026-01", default_periods=3)
        with m:
            m.x = mo.Variable(plan=[1.0, 5.0, 3.0], actual=[2.0, 2.0, 2.0])
            m.mid = mo.MEDIAN(m.x, 4.0)
            m.spread = mo.VAR_P([m.x, m.x * 0.0 + 3.0])
        assert dict(m.mid._value.items()) == {"plan": [2.5, 4.5, 3.5], "actual": [3.0, 3.0, 3.0]}
        assert dict(m.spread._value.items()) == {"plan": [1.0, 1.0, 0.0],
                                                 "actual": [0.25, 0.25, 0.25]}
        assert m.mid.var_type == "list"

    def test_one_line_with_tracks_alone_stays_refused(self):
        m = mo.Model("m", tracks=mo.Tracks("plan", "actual"), default_grain="month",
                     default_start="2026-01", default_periods=2)
        with m:
            m.x = mo.Variable(plan=[1.0, 2.0], actual=[3.0, 4.0])
        with pytest.raises(ValueError, match="slice a coordinate first"):
            mo.MEDIAN(m.x)

    def test_the_call_holds_its_operands(self):
        m = _model()
        out = mo.MEDIAN(m.x.a, m.x.b, 5.0)
        assert out._expr == FuncCall("MEDIAN", [VarRef(m.x.a), VarRef(m.x.b), Literal(5.0)])
        assert out.formula == "MEDIAN(a, b, 5.0)"
        typed = mo.MEDIAN(m.x.a, [1.0, 9.0, 1.0])
        assert typed._expr.args[1] == Literal([1.0, 9.0, 1.0])
        sample = mo.PERCENTILE_INC(_revs(m), 0.25)
        assert sample._expr.args[0] == ListExpr([VarRef(r) for r in _revs(m)])
        assert mo.to_json(sample._expr)["func"] == "PERCENTILE_INC"


# ══════════════════════════════════════════════════════════════════════
# One range: a list of two or more series lines is a sample per period
# ══════════════════════════════════════════════════════════════════════


class TestOneRangePerPeriod:
    def test_percentiles_quartiles_and_the_k_th_of_each_period(self):
        m = _model()
        lines = _lines(m)
        assert mo.PERCENTILE_INC(lines, 0.5).value == [3.0, 8.0, 5.0]
        assert mo.PERCENTILE_INC(lines, 0.25).value == [2.0, 6.0, 5.0]
        assert mo.PERCENTILE(lines, 0.25).value == [2.0, 6.0, 5.0]
        assert mo.PERCENTILE_EXC(lines, 0.5).value == [3.0, 8.0, 5.0]
        assert mo.QUARTILE_INC(lines, 1).value == [2.0, 6.0, 5.0]
        assert mo.QUARTILE(lines, 3).value == [3.5, 8.5, 5.5]
        assert mo.QUARTILE_EXC(lines, 2).value == [3.0, 8.0, 5.0]
        assert mo.LARGE(lines, 1).value == [4.0, 9.0, 6.0]
        assert mo.SMALL(lines, 1).value == [1.0, 4.0, 5.0]
        assert mo.LARGE(lines, 1).var_type == "list"

    def test_a_series_k_reads_this_period_s_k(self):
        m = _model()
        ks = mo.Variable([0.0, 0.5, 1.0])
        assert mo.PERCENTILE_INC(_lines(m), ks).value == [1.0, 8.0, 6.0]
        assert mo.SMALL(_lines(m), [1, 2, 3]).value == [1.0, 8.0, 6.0]
        with pytest.raises(ValueError, match="different lengths"):
            mo.PERCENTILE_INC(_lines(m), mo.Variable([0.0, 0.5]))

    def test_each_line_ranked_among_its_peers_each_period(self):
        m = _model()
        lines = _lines(m)
        assert mo.RANK_EQ(m.x.c, lines).value == [3, 1, 2]
        assert mo.RANK_AVG(m.x.c, lines).value == [3.0, 1.0, 2.5]
        assert mo.RANK_EQ(m.x.c, lines, 1).value == [1, 3, 1]
        assert mo.RANK(m.x.c, lines).value == [3, 1, 2]
        revs = _revs(m)
        assert [mo.RANK_EQ(r, revs).value for r in revs] == [[3, 2, 2], [2, 3, 1], [1, 1, 3]]
        assert mo.RANK_EQ(m.x.c, lines).value_type == "int"

    def test_a_number_typed_in_is_ranked_each_period(self):
        m = _model()
        assert mo.RANK_EQ(5.0, _lines(m)).value == ["#N/A", "#N/A", 2]

    def test_a_scalar_line_in_the_list_stands_in_every_period(self):
        m = _model()
        assert mo.PERCENTILE_INC([m.x.a, m.x.b, m.i.floor], 0.5).value == [4.0, 5.0, 5.0]
        assert mo.LARGE([m.x.a, 7.0], 1).value == [7.0, 8.0, 7.0]

    def test_a_single_series_or_a_list_of_scalars_is_read_as_before(self):
        m = _model()
        assert mo.PERCENTILE_INC(m.x.a, 0.5).value == 5.0
        assert mo.LARGE(m.x.b, mo.Variable([1, 2, 3])).value == [6.0, 4.0, 4.0]
        cells = [mo.Variable(8.0), mo.Variable(12.0), mo.Variable(10.0)]
        assert mo.PERCENTILE_INC(cells, 0.5).value == 10.0
        assert mo.RANK_EQ(cells[0], cells).value == 3

    def test_a_list_holding_one_series_is_refused(self):
        m = _model()
        with pytest.raises(ValueError, match="two or more"):
            mo.PERCENTILE_INC([m.x.a], 0.5)
        with pytest.raises(ValueError, match="two or more"):
            mo.RANK_EQ(m.x.a, [m.x.a])

    def test_an_error_no_number_and_dates_each_period(self):
        e = mo.Variable([2.0, 2.0, 2.0]) / mo.Variable([1.0, 0.0, 2.0])
        blank = mo.Variable([None, None, 3.0])
        assert mo.LARGE([e, blank], 1).value == [2.0, "#DIV/0!", 3.0]
        assert mo.PERCENTILE_INC([blank, mo.Variable(["x", None, 1.0])], 0.5).value == [
            "#NUM!", "#NUM!", 2.0]
        assert mo.RANK_EQ(1.0, [blank, mo.Variable(["x", None, 1.0])]).value == [
            "#N/A", "#N/A", 2]
        d1 = mo.Variable([dt.date(2026, 1, 31), dt.date(2026, 2, 28)])
        d2 = mo.Variable([dt.date(2026, 3, 31), dt.date(2026, 3, 2)])
        out = mo.SMALL([d1, d2], 1)
        assert out.value == [dt.date(2026, 1, 31), dt.date(2026, 2, 28)]
        assert out.value_type == "datetime"

    def test_units(self):
        usd = [mo.Variable([1.0, 2.0], unit="USD"), mo.Variable([3.0, 1.0], unit="USD")]
        assert str(mo.LARGE(usd, 1)._unit) == "USD"
        assert mo.RANK_EQ(usd[0], usd)._unit is None
        with pytest.raises(ValueError, match="mo.PERCENTILE_INC takes its values in one unit"):
            mo.PERCENTILE_INC([usd[0], mo.Variable([1.0, 1.0], unit="hr")], 0.5)
        with pytest.raises(ValueError, match="mo.RANK_EQ takes its values in one unit"):
            mo.RANK_EQ(mo.Variable([1.0, 1.0], unit="hr"), usd)

    def test_a_tracked_line_in_the_list_lifts_per_track(self):
        m = mo.Model("m", tracks=mo.Tracks("plan", "actual"), default_grain="month",
                     default_start="2026-01", default_periods=2)
        with m:
            m.x = mo.Variable(plan=[1.0, 5.0], actual=[2.0, 2.0])
            m.y = mo.Variable(plan=[3.0, 3.0], actual=[1.0, 4.0])
            m.top = mo.LARGE([m.x, m.y], 1)
        assert dict(m.top._value.items()) == {"plan": [3.0, 5.0], "actual": [2.0, 4.0]}


# ══════════════════════════════════════════════════════════════════════
# The registry and the loop container
# ══════════════════════════════════════════════════════════════════════

REDUCTIONS = ("MEDIAN", "STDEV_S", "STDEV_P", "VAR_S", "VAR_P", "STDEV", "STDEVP", "VAR", "VARP")
ONE_RANGE = ("PERCENTILE_INC", "PERCENTILE_EXC", "QUARTILE_INC", "QUARTILE_EXC", "PERCENTILE",
             "QUARTILE", "LARGE", "SMALL", "RANK_EQ", "RANK_AVG", "RANK")


class TestTheRegistry:
    @pytest.mark.parametrize("name", REDUCTIONS + ONE_RANGE)
    def test_two_values_make_a_call_one_per_period(self, name):
        assert per_period_arity(name) == 2
        assert not admits_in_loop(name, 1)
        assert admits_in_loop(name, 2) and admits_in_loop(name, 5)


def _loop_model():
    """A balance paid down by a cap the loop computes each period: the
    median of a row and last period's balance, the larger of two loop rows
    (a list of them), and their spread."""
    m = mo.Model("m", default_grain="year", default_start="2026", default_periods=3)
    with m:
        m.x = mo.MultiVariable("X")
        m.x.a = mo.Variable(list(A), display_name="a")
        m.x.c = mo.Variable(list(C), display_name="c")
        m.d = mo.MultiVariable("D", loop=True)
        with m.d as d:
            d.bal = (mo.lag(d.bal, fill=10.0) - d.paid).set_display_name("bal")
            d.mid = mo.MEDIAN(m.x.a, mo.lag(d.bal, fill=10.0)).set_display_name("mid")
            d.cap = (mo.lag(d.bal, fill=10.0) * 0.8).set_display_name("cap")
            d.top = mo.LARGE([d.mid, d.cap], 1).set_display_name("top")
            d.sd = mo.STDEV_P(d.mid, d.cap).set_display_name("sd")
            d.paid = mo.MIN(d.top - d.sd, mo.lag(d.bal, fill=10.0)).set_display_name("paid")
            d.mix = mo.LARGE([m.x.a, m.x.c, d.paid], 1).set_display_name("mix")
    return m


class TestInsideALoop:
    def test_a_statistic_over_two_or_more_values_runs_one_period_at_a_time(self):
        m = _loop_model()
        d = m.d
        assert d.mid.value == [6.5, 5.375, 2.5]
        assert d.cap.value == pytest.approx([8.0, 2.2, 0.0])
        assert d.top.value == [8.0, 5.375, 2.5]
        assert d.sd.value == pytest.approx([0.75, 1.5875, 1.25])
        assert d.paid.value == [7.25, 2.75, 0.0]
        assert d.bal.value == [2.75, 0.0, 0.0]
        assert d.mix.value == [7.25, 9.0, 5.0]

    def test_the_rows_match_the_same_calls_outside_a_loop(self):
        m = _loop_model()
        d = m.d
        assert mo.LARGE([d.mid, d.cap], 1).value == d.top.value
        assert mo.STDEV_P(d.mid, d.cap).value == d.sd.value

    def test_a_list_of_loop_rows_read_before_the_loop_closes(self):
        m = mo.Model("m", default_grain="year", default_start="2026", default_periods=3)
        with m:
            m.x = mo.MultiVariable("X")
            m.x.a = mo.Variable(list(A), display_name="a")
            m.d = mo.MultiVariable("D", loop=True)
            with m.d as d:
                d.early = mo.MEDIAN([d.bal, m.x.a]).set_display_name("early")
                d.bal = (mo.lag(d.bal, fill=10.0) * 0.5).set_display_name("bal")
        assert m.d.bal.value == [5.0, 2.5, 1.25]
        assert m.d.early.value == [4.0, 5.25, 3.125]

    def test_the_loop_reads_a_row_s_cell_as_a_cell(self):
        # A blank cell is skipped and an error a row holds is the result,
        # as the workbook reads its cells — never 0, nor the #VALUE! that
        # text typed into the call would give.
        m = mo.Model("m", default_grain="year", default_start="2026", default_periods=3)
        with m:
            m.x = mo.MultiVariable("X")
            m.x.den = mo.Variable([1.0, 0.0, 2.0], display_name="den")
            m.x.blank = mo.Variable([None, 4.0, None], display_name="blank")
            m.x.k = mo.Variable(["#N/A", 0.5, 1.0], display_name="k")
            m.d = mo.MultiVariable("D", loop=True)
            with m.d as d:
                d.r = (2.0 / m.x.den + mo.lag(d.n, fill=0.0) * 0.0).set_display_name("r")
                d.mid = mo.MEDIAN(d.r, m.x.blank).set_display_name("mid")
                d.big = mo.LARGE([d.r, m.x.blank], 1).set_display_name("big")
                d.pk = mo.PERCENTILE_INC([d.r, m.x.blank], m.x.k).set_display_name("pk")
                d.n = mo.IFERROR(d.mid + d.big + d.pk, 0.0).set_display_name("n")
        assert m.d.r.value == [2.0, "#DIV/0!", 1.0]
        assert m.d.mid.value == [2.0, "#DIV/0!", 1.0]
        assert m.d.big.value == [2.0, "#DIV/0!", 1.0]
        assert m.d.pk.value == ["#N/A", "#DIV/0!", 1.0]
        assert m.d.n.value == [0.0, 0.0, 3.0]
        # The same calls outside the loop read the same.
        assert mo.MEDIAN(m.d.r, m.x.blank).value == m.d.mid.value
        assert mo.PERCENTILE_INC([m.d.r, m.x.blank], m.x.k).value == m.d.pk.value

    @pytest.mark.parametrize("make", [
        lambda d: mo.MEDIAN(d.bal),
        lambda d: mo.STDEV_S(d.bal),
        lambda d: mo.PERCENTILE_INC(d.bal, 0.5),
        lambda d: mo.RANK_EQ(d.bal, d.bal),
    ])
    def test_one_row_alone_reads_a_whole_series_and_is_refused(self, make):
        m = mo.Model("m", default_grain="year", default_start="2026", default_periods=3)
        with pytest.raises(ValueError, match="reads a whole series") as err:
            with m:
                m.d = mo.MultiVariable("D", loop=True)
                with m.d as d:
                    d.bal = mo.lag(d.x, fill=10.0) * 1.37
                    d.stat = make(d)
                    d.x = d.bal - d.stat
        assert "one call per period" in str(err.value)

    def test_the_loop_derives_the_units_by_the_function_s_law(self):
        m = mo.Model("m", default_grain="year", default_start="2026", default_periods=2)
        with m:
            m.x = mo.MultiVariable("X")
            m.x.a = mo.Variable([1.0, 2.0], unit="USD", display_name="a")
            m.d = mo.MultiVariable("D", loop=True)
            with m.d as d:
                d.bal = (mo.lag(d.bal, fill=m.x.a[0]) + d.step).set_display_name("bal")
                d.mid = mo.MEDIAN(mo.lag(d.bal, fill=m.x.a[0]), m.x.a).set_display_name("mid")
                d.v = mo.VAR_P(mo.lag(d.bal, fill=m.x.a[0]), m.x.a).set_display_name("v")
                d.step = (d.mid * 0.5 + d.v / d.mid * 0.0).set_display_name("step")
        assert m.d.v.value == [0.0, 0.0625]
        assert m.d.bal.value == [1.5, 2.375]
        assert str(m.d.v._unit) == "USD^2"
        assert str(m.d.mid._unit) == "USD"
        assert str(m.d.step._unit) == "USD"

    def test_two_units_in_a_loop_are_refused(self):
        m = mo.Model("m", default_grain="year", default_start="2026", default_periods=2)
        with pytest.raises(ValueError, match="mo.MEDIAN takes its values in one unit"):
            with m:
                m.x = mo.MultiVariable("X")
                m.x.a = mo.Variable([1.0, 2.0], unit="USD", display_name="a")
                m.x.h = mo.Variable([1.0, 2.0], unit="hr", display_name="h")
                m.d = mo.MultiVariable("D", loop=True)
                with m.d as d:
                    d.bal = (mo.lag(d.bal, fill=m.x.a[0]) + (d.mid > 0.0) * 0.0
                             ).set_display_name("bal")
                    d.mid = mo.MEDIAN(mo.lag(d.bal, fill=m.x.a[0]), m.x.h).set_display_name("mid")


# ══════════════════════════════════════════════════════════════════════
# The workbook: one call per column over this column's cells
# ══════════════════════════════════════════════════════════════════════


class TestTheWorkbook:
    def test_operands_are_this_column_s_cells(self, tmp_path):
        m = _model()
        x, i = m.x, m.i
        with m:
            x.mid = mo.MEDIAN(x.a, x.b).set_display_name("mid")
            x.sd = mo.STDEV_S(x.a, x.b, x.c).set_display_name("sd")
            x.fl = mo.MEDIAN(x.a, i.floor).set_display_name("fl")
            x.lit = mo.VAR_P(x.a, [1.0, 9.0, 1.0]).set_display_name("lit")
            x.twice = (mo.MEDIAN(x.a, x.b) * 2).set_display_name("twice")
        rows = _rows(_book(m, tmp_path)["X"])
        assert rows["mid"] == ["=MEDIAN(B2, B3)", "=MEDIAN(C2, C3)", "=MEDIAN(D2, D3)"]
        assert rows["sd"] == [f"=_xlfn.STDEV.S({c}2, {c}3, {c}4)" for c in "BCD"]
        assert rows["fl"] == [f"=MEDIAN({c}2, I!B1)" for c in "BCD"]
        assert rows["lit"] == ["=_xlfn.VAR.P(B2, 1.0)", "=_xlfn.VAR.P(C2, 9.0)",
                               "=_xlfn.VAR.P(D2, 1.0)"]
        assert rows["twice"] == [f"=MEDIAN({c}2, {c}3) * 2" for c in "BCD"]
        assert check_compat(m, "excel") == []

    def test_a_list_of_sibling_lines_is_the_run_of_their_cells_in_each_column(self, tmp_path):
        m = _model()
        x, i = m.x, m.i
        with m:
            x.mid = mo.MEDIAN(_lines(m)).set_display_name("mid")
            x.q1 = mo.PERCENTILE_INC(_lines(m), i.k).set_display_name("q1")
            x.top = mo.LARGE(_lines(m), 1).set_display_name("top")
            x.rank = mo.RANK_EQ(x.c, _lines(m)).set_display_name("rank")
        rows = _rows(_book(m, tmp_path)["X"])
        assert rows["mid"] == [f"=MEDIAN({c}2:{c}4)" for c in "BCD"]
        assert rows["q1"] == [f"=_xlfn.PERCENTILE.INC({c}2:{c}4, I!B2)" for c in "BCD"]
        assert rows["top"] == [f"=LARGE({c}2:{c}4, 1)" for c in "BCD"]
        assert rows["rank"] == [f"=_xlfn.RANK.EQ({c}4, {c}2:{c}4)" for c in "BCD"]
        assert check_compat(m, "excel") == []

    def test_a_records_table_of_yearly_lines_gives_a_statistic_per_year(self, tmp_path):
        m = _model()
        revs = _revs(m)
        with m.s as s:
            s.median = mo.MEDIAN(revs).set_display_name("median")
            s.q1 = mo.PERCENTILE_INC(revs, m.i.k).set_display_name("q1")
            s.sd = mo.STDEV_S(revs).set_display_name("sd")
        with m.peers as p:
            for rec in (p.a, p.b, p.c):
                rec.rank = mo.RANK_EQ(rec.rev, revs).set_display_name("Rank")
        book = _book(m, tmp_path)
        rows = _rows(book["Summary"])
        assert rows["median"] == [f"=MEDIAN(Peers!{c}3:{c}5)" for c in "BCD"]
        assert rows["q1"] == [f"=_xlfn.PERCENTILE.INC(Peers!{c}3:{c}5, I!B2)" for c in "BCD"]
        assert rows["sd"] == [f"=_xlfn.STDEV.S(Peers!{c}3:{c}5)" for c in "BCD"]
        assert m.s.median.value == [20.0, 30.0, 25.0]
        assert m.s.q1.value == [15.0, 20.0, 17.5]
        ws = book["Peers"]
        ranks = {(cell.row, cell.column_letter): cell.value
                 for row in ws.iter_rows(min_row=8, max_row=10) for cell in row[1:4]}
        assert ranks == {
            (row, c): f"=_xlfn.RANK.EQ({c}{row - 5}, {c}3:{c}5)"
            for row in (8, 9, 10) for c in "BCD"
        }
        assert [r.rank.value for r in (m.peers.a, m.peers.b, m.peers.c)] == [
            [3, 2, 2], [2, 3, 1], [1, 1, 3]]
        assert check_compat(m, "excel") == []

    def test_cells_that_form_no_run_carry_the_value(self, tmp_path):
        m = _model()
        x = m.x
        with m:
            x.gap = mo.PERCENTILE_INC([x.a, x.c], 0.5).set_display_name("gap")
            x.back = mo.MEDIAN([x.c, x.b, x.a]).set_display_name("back")
            x.sheets = mo.LARGE([x.a, m.peers.a.rev], 1).set_display_name("sheets")
        rows = _rows(_book(m, tmp_path)["X"])
        assert rows["gap"] == [2.0, 8.5, 5.0]
        assert rows["back"] == [3.0, 8.0, 5.0]
        assert rows["sheets"] == [10.0, 30.0, 25.0]
        # One entry per cell: each period's call carries its value.
        assert [(i.kind, i.qpath) for i in check_compat(m, "excel")] == (
            [("value", x.gap.id)] * 3 + [("value", x.back.id)] * 3
            + [("value", x.sheets.id)] * 3)

    def test_a_field_of_scalar_lines_beside_a_series_is_its_range_in_every_column(
            self, tmp_path):
        m = _model()
        with m:
            m.i.lo = mo.Variable(1.0, display_name="lo")
            m.i.hi = mo.Variable(10.0, display_name="hi")
            m.x.mid = mo.MEDIAN([m.i.floor, m.i.k, m.i.lo, m.i.hi], m.x.a).set_display_name("mid")
        rows = _rows(_book(m, tmp_path)["X"])
        assert rows["mid"] == [f"=MEDIAN(I!B1:B4, {c}2)" for c in "BCD"]
        assert m.x.mid.value == [3.0, 5.0, 5.0]

    def test_the_reduction_of_one_line_and_the_scalars_are_unchanged(self, tmp_path):
        m = _model()
        with m:
            m.s.red = mo.MEDIAN(m.x.a).set_display_name("red")
            m.s.pct = mo.PERCENTILE_INC(m.x.a, m.i.k).set_display_name("pct")
            m.i.pair = mo.MEDIAN(m.i.floor, m.i.k).set_display_name("pair")
        book = _book(m, tmp_path)
        assert _rows(book["Summary"])["red"][0] == "=MEDIAN(X!B2:D2)"
        assert _rows(book["Summary"])["pct"][0] == "=_xlfn.PERCENTILE.INC(X!B2:D2, I!B2)"
        assert _rows(book["I"])["pair"][0] == "=MEDIAN(B1, B2)"

    def test_inside_a_loop_container(self, tmp_path):
        m = _loop_model()
        rows = _rows(_book(m, tmp_path)["D"])
        assert rows["mid"] == ["=MEDIAN(X!B2, 10.0)", "=MEDIAN(X!C2, B2)", "=MEDIAN(X!D2, C2)"]
        assert rows["top"] == [f"=LARGE({c}3:{c}4, 1)" for c in "BCD"]
        assert rows["sd"] == [f"=_xlfn.STDEV.P({c}3, {c}4)" for c in "BCD"]
        # Two sheets' cells are no run: the value, listed by check_compat.
        assert rows["mix"] == [7.25, 9, 5]
        assert [(i.kind, i.qpath) for i in check_compat(m, "excel")] == [
            ("value", m.d.mix.id)] * 3

    def test_the_notebook_shows_the_per_cell_formula(self):
        m = _model()
        with m:
            m.x.mid = mo.MEDIAN(m.x.a, m.x.b).set_display_name("mid")
        html = m._repr_html_()
        assert "MEDIAN(B2, B3)" in html
        assert "B2:D2" not in html
