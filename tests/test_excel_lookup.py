# SPDX-License-Identifier: Apache-2.0
"""Lookup and reference in the workbook.

A row is its cells' range (``Inputs!C2:E2``); a list of scalar lines
whose cells stand one under another is a vertical range (``Deals!B3:B5``);
a table of records is one block, the records' labels in its first
column and their scalar fields after (``Deals!A3:C5``); a list of rows
laid out one under another is one block too (``Inputs!C2:E3``, the
period block ``Deals!D8:F10``). A series key writes one call per
period over that period's key cell and the whole range. Cells that form
no rectangle — the keys typed into a lookup over a keyed line, records
laid out as sections, a unit column between a record's label and its
fields, rows on two sheets — cannot be one range, so the cell carries
its computed value and ``check_compat`` lists it. ``XLOOKUP`` is stored
as ``_xlfn.XLOOKUP``; a mode without ``if_not_found`` is written after
``NA()``. Every function recalculates in LibreOffice to the engine's
numbers, except where LibreOffice departs from Excel (a boolean read as
a number, text that reads as a number matching the number, a column or
a row number past the block), pinned as an expected failure.
"""

import datetime as dt
import shutil
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook

import modeleon as mo
from modeleon.compile import check_compat
from modeleon.functions.lookup import HLOOKUP, INDEX, MATCH, VLOOKUP, XLOOKUP


def _rows(ws):
    out = {}
    for row in ws.iter_rows():
        label = row[0].value
        if label:
            out[str(label).split(" (")[0]] = [c.value for c in row[1:]]
    return out


def _book(model, tmp_path):
    path = tmp_path / "book.xlsx"
    model.to_excel(path)
    return load_workbook(path)


class Deal(mo.MultiVariableClass):
    def compute(self, region, amount):
        self.region = mo.Variable(region, display_name="Region")
        self.amount = mo.Variable(amount, display_name="Amount", unit="USD")
        self.flow = mo.Variable([amount * 0.1, amount * 0.2, amount * 0.3], display_name="Flow")


def _model():
    m = mo.Model("m", display_name="M", default_grain="year", default_start="2026",
                 default_periods=3)
    with m:
        m.inputs = mo.MultiVariable("Inputs", excel_props={"tab": True})
        with m.inputs as i:
            i.years = mo.Variable([2026, 2027, 2028], display_name="Years")
            i.rates = mo.Variable([0.05, 0.04, 0.03], display_name="Rates")
            i.names = mo.Variable(["Alpha", "Beta", "Gamma"], display_name="Names")
            i.desc = mo.Variable([30, 20, 10], display_name="Desc")
            i.dup = mo.Variable([20, 10, 20], display_name="Dup")
            i.vals = mo.Variable(["a", "b", "c"], display_name="Vals")
            i.mixed = mo.Variable([1.0, None, 3.0], display_name="Mixed")
            i.dates = mo.Variable([dt.date(2026, 1, 1), dt.date(2027, 1, 1), dt.date(2028, 1, 1)],
                                  display_name="Dates")
            i.infl = mo.Variable({2026: 0.05, 2027: 0.04, 2028: 0.03}, display_name="Inflation")
            i.pick = mo.Variable(2027, display_name="Pick")
            i.word = mo.Variable("beta", display_name="Word")
            i.tpos = mo.Variable("2", display_name="TPos")
            i.blank = mo.Variable(None, display_name="Blank")
            i.zeros = mo.Variable([0, 1, 2], display_name="Zeros")
            i.err = (i.pick / 0).set_display_name("Err")
        m.deals = mo.MultiVariable("Deals", excel_props={"tab": True, "orient": "records"})
        with m.deals as d:
            d.a = Deal(region="North", amount=100.0, display_name="Alpha")
            d.b = Deal(region="South", amount=200.0, display_name="Beta")
            d.c = Deal(region="North", amount=300.0, display_name="Gamma")
        m.t = mo.MultiVariable("Time", excel_props={"tab": True})
        m.t.year = mo.YEAR(mo.time.start).set_display_name("Year")
        m.t.label = mo.Variable(["Alpha", "Gamma", "Delta"], display_name="Label")
        m.s = mo.MultiVariable("Summary", excel_props={"tab": True})
    return m


# ══════════════════════════════════════════════════════════════════════
# Rows are ranges, records and rows are blocks
# ══════════════════════════════════════════════════════════════════════


class TestRangesAndBlocks:
    def test_every_function_over_rows(self, tmp_path):
        m = _model()
        i = m.inputs
        with m.s as s:
            s.match = MATCH(i.pick, i.years, 0).set_display_name("match")
            s.match_default = MATCH(2027.5, i.years).set_display_name("match_default")
            s.index = INDEX(i.rates, 2).set_display_name("index")
            s.index_match = INDEX(i.rates, MATCH(i.pick, i.years, 0)).set_display_name("index_match")
            s.xl = XLOOKUP(i.pick, i.years, i.rates).set_display_name("xl")
            s.xl_text = XLOOKUP(i.word, i.names, i.rates).set_display_name("xl_text")
            s.xl_nf = XLOOKUP(2030, i.years, i.rates, 0).set_display_name("xl_nf")
            s.xl_modes = XLOOKUP(2026.5, i.years, i.rates, 0, 1).set_display_name("xl_modes")
            s.xl_na = XLOOKUP(2027.5, i.years, i.rates, match_mode=-1).set_display_name("xl_na")
            s.xl_last = XLOOKUP(20, i.dup, i.vals, search_mode=-1).set_display_name("xl_last")
            s.xl_date = XLOOKUP(dt.date(2027, 1, 1), i.dates, i.rates).set_display_name("xl_date")
            s.hl = HLOOKUP(i.pick, [i.years, i.rates], 2, False).set_display_name("hl")
            s.hl_default = HLOOKUP(2027.5, [i.years, i.rates], 2).set_display_name("hl_default")
            s.index_rows = INDEX([i.years, i.rates], 2, 3).set_display_name("index_rows")
        rows = _rows(_book(m, tmp_path)["Summary"])
        assert rows["match"][0] == "=MATCH(Inputs!B12, Inputs!C2:E2, 0)"
        assert rows["match_default"][0] == "=MATCH(2027.5, Inputs!C2:E2)"
        assert rows["index"][0] == "=INDEX(Inputs!C3:E3, 2)"
        assert rows["index_match"][0] == "=INDEX(Inputs!C3:E3, MATCH(Inputs!B12, Inputs!C2:E2, 0))"
        assert rows["xl"][0] == "=_xlfn.XLOOKUP(Inputs!B12, Inputs!C2:E2, Inputs!C3:E3)"
        assert rows["xl_text"][0] == "=_xlfn.XLOOKUP(Inputs!B13, Inputs!C4:E4, Inputs!C3:E3)"
        assert rows["xl_nf"][0] == "=_xlfn.XLOOKUP(2030, Inputs!C2:E2, Inputs!C3:E3, 0)"
        assert rows["xl_modes"][0] == "=_xlfn.XLOOKUP(2026.5, Inputs!C2:E2, Inputs!C3:E3, 0, 1)"
        assert rows["xl_na"][0] == "=_xlfn.XLOOKUP(2027.5, Inputs!C2:E2, Inputs!C3:E3, NA(), -1)"
        assert rows["xl_last"][0] == "=_xlfn.XLOOKUP(20, Inputs!C6:E6, Inputs!C7:E7, NA(), 0, -1)"
        assert rows["xl_date"][0] == "=_xlfn.XLOOKUP(DATE(2027, 1, 1), Inputs!C9:E9, Inputs!C3:E3)"
        assert rows["hl"][0] == "=HLOOKUP(Inputs!B12, Inputs!C2:E3, 2, FALSE)"
        assert rows["hl_default"][0] == "=HLOOKUP(2027.5, Inputs!C2:E3, 2)"
        assert rows["index_rows"][0] == "=INDEX(Inputs!C2:E3, 2, 3)"
        # The values the cells cache are the engine's.
        assert (m.s.match.value, m.s.index_match.value, m.s.xl.value) == (2, 0.04, 0.04)
        assert (m.s.xl_na.value, m.s.xl_last.value, m.s.hl.value) == (0.04, "c", 0.04)

    def test_a_table_of_records_is_one_block_the_label_column_first(self, tmp_path):
        m = _model()
        d = m.deals
        with m.s as s:
            s.vl = VLOOKUP("Beta", m.deals, 3, False).set_display_name("vl")
            s.vl_default = VLOOKUP("Bob", m.deals, 3).set_display_name("vl_default")
            s.vl_label = VLOOKUP("gamma", m.deals, 1, False).set_display_name("vl_label")
            s.index_block = INDEX(m.deals, 2, 3).set_display_name("index_block")
            s.index_label = INDEX(m.deals, 2, 1).set_display_name("index_label")
            s.hl_block = HLOOKUP("North", m.deals, 2, False).set_display_name("hl_block")
            s.match_vert = MATCH("South", [d.a.region, d.b.region, d.c.region], 0).set_display_name("match_vert")
            s.index_vert = INDEX([d.a.amount, d.b.amount, d.c.amount], 3).set_display_name("index_vert")
            s.xl_vert = XLOOKUP("South", [d.a.region, d.b.region, d.c.region],
                                [d.a.amount, d.b.amount, d.c.amount]).set_display_name("xl_vert")
        book = _book(m, tmp_path)
        rows = _rows(book["Summary"])
        assert rows["vl"][0] == '=VLOOKUP("Beta", Deals!A3:C5, 3, FALSE)'
        assert rows["vl_default"][0] == '=VLOOKUP("Bob", Deals!A3:C5, 3)'
        assert rows["vl_label"][0] == '=VLOOKUP("gamma", Deals!A3:C5, 1, FALSE)'
        assert rows["index_block"][0] == "=INDEX(Deals!A3:C5, 2, 3)"
        assert rows["index_label"][0] == "=INDEX(Deals!A3:C5, 2, 1)"
        assert rows["hl_block"][0] == '=HLOOKUP("North", Deals!A3:C5, 2, FALSE)'
        assert rows["match_vert"][0] == '=MATCH("South", Deals!B3:B5, 0)'
        assert rows["index_vert"][0] == "=INDEX(Deals!C3:C5, 3)"
        assert rows["xl_vert"][0] == '=_xlfn.XLOOKUP("South", Deals!B3:B5, Deals!C3:C5)'
        # The block the formula names holds what the engine looked through:
        # the labels in column A, the fields after them.
        ws = book["Deals"]
        grid = [[ws.cell(r, c).value for c in range(1, 4)] for r in range(3, 6)]
        assert grid == [["Alpha", "North", 100], ["Beta", "South", 200], ["Gamma", "North", 300]]
        assert (m.s.vl.value, m.s.vl_default.value, m.s.vl_label.value) == (200.0, 200.0, "Gamma")
        assert (m.s.index_block.value, m.s.index_label.value, m.s.hl_block.value) == (
            200.0, "Beta", "South")

    def test_the_rows_of_the_records_period_blocks_are_one_block(self, tmp_path):
        m = _model()
        d = m.deals
        with m.s as s:
            s.flows = INDEX([d.a.flow, d.b.flow, d.c.flow], 2, 3).set_display_name("flows")
            s.hl_flows = HLOOKUP(20, [d.a.flow, d.b.flow, d.c.flow], 3, False).set_display_name("hl_flows")
        rows = _rows(_book(m, tmp_path)["Summary"])
        assert rows["flows"][0] == "=INDEX(Deals!D8:F10, 2, 3)"
        assert rows["hl_flows"][0] == "=HLOOKUP(20, Deals!D8:F10, 3, FALSE)"
        assert m.s.flows.value == 60.0 and m.s.hl_flows.value == 60.0

    def test_a_series_key_is_one_call_per_period_over_the_whole_range(self, tmp_path):
        m = _model()
        i, t = m.inputs, m.t
        with m.s as s:
            s.series = XLOOKUP(t.year, i.years, i.rates, 0).set_display_name("series")
            s.series_match = MATCH(t.year, i.years, 0).set_display_name("series_match")
            s.series_index = INDEX(i.rates, MATCH(t.year, i.years, 0)).set_display_name("series_index")
            s.series_vl = VLOOKUP(t.label, m.deals, 3, False).set_display_name("series_vl")
            s.series_pos = INDEX(i.rates, t.year - 2025).set_display_name("series_pos")
        rows = _rows(_book(m, tmp_path)["Summary"])
        # A sheet of series only: the periods start at B.
        assert rows["series"] == [
            "=_xlfn.XLOOKUP(Time!B2, Inputs!C2:E2, Inputs!C3:E3, 0)",
            "=_xlfn.XLOOKUP(Time!C2, Inputs!C2:E2, Inputs!C3:E3, 0)",
            "=_xlfn.XLOOKUP(Time!D2, Inputs!C2:E2, Inputs!C3:E3, 0)",
        ]
        assert rows["series_match"][0] == "=MATCH(Time!B2, Inputs!C2:E2, 0)"
        assert rows["series_index"][2] == "=INDEX(Inputs!C3:E3, MATCH(Time!D2, Inputs!C2:E2, 0))"
        assert rows["series_vl"] == [
            '=VLOOKUP(Time!B3, Deals!A3:C5, 3, FALSE)',
            '=VLOOKUP(Time!C3, Deals!A3:C5, 3, FALSE)',
            '=VLOOKUP(Time!D3, Deals!A3:C5, 3, FALSE)',
        ]
        assert rows["series_pos"][1] == "=INDEX(Inputs!C3:E3, Time!C2 - 2025)"
        assert m.s.series.value == [0.05, 0.04, 0.03]
        assert m.s.series_vl.value == [100.0, 300.0, "#N/A"]
        assert m.s.series_pos.value == [0.05, 0.04, 0.03]

    def test_on_the_rows_own_sheet_and_inside_a_larger_formula(self, tmp_path):
        m = _model()
        with m.inputs as i:
            i.rate = XLOOKUP(i.pick, i.years, i.rates).set_display_name("rate")
            i.scaled = (INDEX(i.rates, MATCH(i.pick, i.years, 0)) * 100).set_display_name("scaled")
        rows = _rows(_book(m, tmp_path)["Inputs"])
        assert rows["rate"][0] == "=_xlfn.XLOOKUP(B12, C2:E2, C3:E3)"
        assert rows["scaled"][0] == "=INDEX(C3:E3, MATCH(B12, C2:E2, 0)) * 100"


# ══════════════════════════════════════════════════════════════════════
# No rectangle: the cell carries its value and check_compat lists it
# ══════════════════════════════════════════════════════════════════════


class TestNoRectangleShipsTheValue:
    def test_the_keys_of_a_keyed_line_have_no_cells(self, tmp_path):
        m = _model()
        i = m.inputs
        with m.s as s:
            s.keyed = XLOOKUP(i.pick, i.infl).set_display_name("keyed")
            s.keyed_series = XLOOKUP(m.t.year, i.infl, if_not_found=0).set_display_name("keyed_series")
        with pytest.warns(mo.CrossScopeReferenceWarning, match="XLOOKUP"):
            rows = _rows(_book(m, tmp_path)["Summary"])
        assert rows["keyed"][0] == 0.04
        assert rows["keyed_series"][1:] == [0.05, 0.04, 0.03]
        with pytest.warns(mo.CrossScopeReferenceWarning):
            issues = check_compat(m, "excel")
        assert [(x.kind, x.qpath, x.period) for x in issues] == [
            ("value", m.s.keyed.id, 0),
            ("value", m.s.keyed_series.id, 0),
            ("value", m.s.keyed_series.id, 1),
            ("value", m.s.keyed_series.id, 2),
        ]
        assert issues[0].inlined == "XLOOKUP(pick, [2026, 2027, 2028], infl)"

    def test_records_laid_out_as_sections(self, tmp_path):
        m = _model()
        with m:
            m.stack = mo.MultiVariable("Stack", excel_props={"tab": True})
            with m.stack as t:
                t.a = Deal(region="North", amount=100.0, display_name="Alpha")
                t.b = Deal(region="South", amount=200.0, display_name="Beta")
            with m.s as s:
                s.vl = VLOOKUP("Beta", m.stack, 3, False).set_display_name("vl")
        rows = _rows(_book(m, tmp_path)["Summary"])
        assert rows["vl"][0] == 200
        issues = check_compat(m, "excel")
        assert [(x.kind, x.qpath) for x in issues] == [("value", m.s.vl.id)]
        assert issues[0].inlined.startswith("[['Alpha', ")

    def test_a_unit_column_between_a_record_s_label_and_its_fields(self, tmp_path):
        m = _model()
        m.default_excel_view = mo.ExcelView(meta={"unit": "Unit"})
        with m.s as s:
            s.vl = VLOOKUP("Beta", m.deals, 3, False).set_display_name("vl")
            s.index = INDEX(m.deals, 2, 3).set_display_name("index")
        book = _book(m, tmp_path)
        ws = book["Deals"]
        # The fields start at C (the unit row under the captions pushes the
        # records to row 4): nothing names the label cell at A.
        assert [ws.cell(4, c).value for c in range(1, 5)] == ["Alpha", None, "North", 100]
        rows = _rows(book["Summary"])
        # (The Summary sheet has the unit column too: the values stand at C.)
        assert rows["vl"][1] == 200 and rows["index"][1] == 200
        assert [(x.kind, x.qpath) for x in check_compat(m, "excel")] == [
            ("value", m.s.vl.id), ("value", m.s.index.id)]

    def test_rows_on_two_sheets_a_typed_row_or_a_row_the_book_has_no_cells_for(self, tmp_path):
        m = _model()
        i = m.inputs
        with m.s as s:
            s.sheets = HLOOKUP(2027, [i.years, m.t.year], 2, False).set_display_name("sheets")
            s.typed = HLOOKUP(2027, [[2026, 2027, 2028], i.rates], 2, False).set_display_name("typed")
            s.unlaid = HLOOKUP(2027, [i.years, i.rates * 2], 2, False).set_display_name("unlaid")
        rows = _rows(_book(m, tmp_path)["Summary"])
        assert rows["sheets"][0] == 2027 and rows["typed"][0] == 0.04 and rows["unlaid"][0] == 0.08
        assert [(x.kind, x.qpath) for x in check_compat(m, "excel")] == [
            ("value", m.s.sheets.id), ("value", m.s.typed.id), ("value", m.s.unlaid.id)]

    def test_a_record_that_lacks_a_field(self, tmp_path):
        class Custom(mo.MultiVariableClass):
            def compute(self, salary, bonus=None):
                self.salary = mo.Variable(salary, display_name="Salary")
                if bonus is not None:
                    self.bonus = mo.Variable(bonus, display_name="Bonus")

        m = mo.Model("m", display_name="M")
        with m:
            m.staff = mo.MultiVariable("Staff", excel_props={"tab": True, "orient": "records"})
            with m.staff as t:
                t.a = Custom(salary=1, display_name="A")
                t.b = Custom(salary=2, bonus=9, display_name="B")
            m.s = mo.MultiVariable("Summary", excel_props={"tab": True})
            m.s.bonus = VLOOKUP("B", m.staff, 3, False).set_display_name("bonus")
        rows = _rows(_book(m, tmp_path)["Summary"])
        assert rows["bonus"][0] == 9
        assert [(x.kind, x.qpath) for x in check_compat(m, "excel")] == [("value", m.s.bonus.id)]

    def test_a_rectangle_is_listed_by_nobody(self, tmp_path):
        m = _model()
        i = m.inputs
        with m.s as s:
            s.vl = VLOOKUP("Beta", m.deals, 3, False).set_display_name("vl")
            s.hl = HLOOKUP(i.pick, [i.years, i.rates], 2, False).set_display_name("hl")
            s.xl = XLOOKUP(m.t.year, i.years, i.rates).set_display_name("xl")
        _book(m, tmp_path)
        assert check_compat(m, "excel") == []


# ══════════════════════════════════════════════════════════════════════
# LibreOffice recalculates these functions to the engine's numbers
# ══════════════════════════════════════════════════════════════════════


def _soffice():
    for candidate in ("soffice", "libreoffice"):
        found = shutil.which(candidate)
        if found:
            return found
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


def _recalculated(model, tmp_path):
    """The Summary sheet after LibreOffice recalculates the written
    workbook, by line label."""
    soffice = _soffice()
    if soffice is None:
        pytest.skip("LibreOffice not installed")
    path = tmp_path / "lookup.xlsx"
    model.to_excel(path)
    subprocess.run(
        [soffice, f"-env:UserInstallation={(tmp_path / 'prof').as_uri()}",
         "--headless", "--convert-to", "xlsx", "--outdir",
         str(tmp_path / "calc"), str(path)],
        check=True, capture_output=True,
    )
    return _rows(load_workbook(tmp_path / "calc" / "lookup.xlsx", data_only=True)["Summary"])


def _cells(calc_row, value):
    """What the line's cells hold after the recalculation: the one cell
    of a scalar line, the period cells of a series (past the sheet's
    constants column)."""
    if not isinstance(value, list):
        return calc_row[0]
    start = next(i for i, c in enumerate(calc_row) if c is not None)
    return calc_row[start:start + len(value)]


@pytest.mark.slow
def test_the_workbook_recalculates_to_the_python_values(tmp_path):
    m = _model()
    i, d, t = m.inputs, m.deals, m.t
    regions = [d.a.region, d.b.region, d.c.region]
    amounts = [d.a.amount, d.b.amount, d.c.amount]
    flows = [d.a.flow, d.b.flow, d.c.flow]
    with m.s as s:
        s.match_exact = MATCH(i.pick, i.years, 0).set_display_name("match_exact")
        s.match_approx = MATCH(2027.5, i.years).set_display_name("match_approx")
        s.match_desc = MATCH(25, i.desc, -1).set_display_name("match_desc")
        s.match_text = MATCH("b*", i.names, 0).set_display_name("match_text")
        s.match_na = MATCH(2030, i.years, 0).set_display_name("match_na")
        s.match_vert = MATCH("South", regions, 0).set_display_name("match_vert")
        s.index_row = INDEX(i.rates, 2).set_display_name("index_row")
        s.index_match = INDEX(i.rates, MATCH(i.pick, i.years, 0)).set_display_name("index_match")
        s.index_vert = INDEX(amounts, 3).set_display_name("index_vert")
        s.index_block = INDEX(m.deals, 2, 3).set_display_name("index_block")
        s.index_label = INDEX(m.deals, 2, 1).set_display_name("index_label")
        s.index_rows = INDEX([i.years, i.rates], 2, 3).set_display_name("index_rows")
        s.index_flows = INDEX(flows, 2, 3).set_display_name("index_flows")
        s.index_blank = INDEX(i.mixed, 2).set_display_name("index_blank")
        s.index_ref = INDEX(i.rates, 4).set_display_name("index_ref")
        s.xl = XLOOKUP(i.pick, i.years, i.rates).set_display_name("xl")
        s.xl_text = XLOOKUP(i.word, i.names, i.rates).set_display_name("xl_text")
        s.xl_nf = XLOOKUP(2030, i.years, i.rates, 0).set_display_name("xl_nf")
        s.xl_na = XLOOKUP(2030, i.years, i.rates).set_display_name("xl_na")
        s.xl_below = XLOOKUP(2027.5, i.years, i.rates, match_mode=-1).set_display_name("xl_below")
        s.xl_above = XLOOKUP(2026.5, i.years, i.rates, 0, 1).set_display_name("xl_above")
        s.xl_last = XLOOKUP(20, i.dup, i.vals, search_mode=-1).set_display_name("xl_last")
        s.xl_first = XLOOKUP(20, i.dup, i.vals).set_display_name("xl_first")
        s.xl_wild = XLOOKUP("g*", i.names, i.rates, 0, 2).set_display_name("xl_wild")
        s.xl_vert = XLOOKUP("South", regions, amounts).set_display_name("xl_vert")
        s.xl_date = XLOOKUP(dt.date(2027, 1, 1), i.dates, i.rates).set_display_name("xl_date")
        s.xl_blank = XLOOKUP(2027, i.years, i.mixed).set_display_name("xl_blank")
        s.xl_keyed = XLOOKUP(i.pick, i.infl).set_display_name("xl_keyed")
        s.vl = VLOOKUP("Beta", m.deals, 3, False).set_display_name("vl")
        s.vl_approx = VLOOKUP("Bob", m.deals, 3).set_display_name("vl_approx")
        s.vl_label = VLOOKUP("gamma", m.deals, 1, False).set_display_name("vl_label")
        s.vl_na = VLOOKUP("Delta", m.deals, 2, False).set_display_name("vl_na")
        s.vl_wild = VLOOKUP("G*", m.deals, 2, False).set_display_name("vl_wild")
        s.hl = HLOOKUP(i.pick, [i.years, i.rates], 2, False).set_display_name("hl")
        s.hl_approx = HLOOKUP(2027.5, [i.years, i.rates], 2).set_display_name("hl_approx")
        s.hl_block = HLOOKUP("North", m.deals, 2, False).set_display_name("hl_block")
        s.series = XLOOKUP(t.year, i.years, i.rates, 0).set_display_name("series")
        s.series_match = MATCH(t.year, i.years, 0).set_display_name("series_match")
        s.series_vl = VLOOKUP(t.label, m.deals, 3, False).set_display_name("series_vl")
        s.series_index = INDEX(i.rates, MATCH(t.year, i.years, 0)).set_display_name("series_index")
        # A position is read as a number: text that reads as one, a boolean.
        s.index_text = INDEX(i.rates, "2").set_display_name("index_text")
        s.index_true = INDEX(i.rates, True).set_display_name("index_true")
        s.index_text_cell = INDEX(i.rates, i.tpos).set_display_name("index_text_cell")
        s.index_block_text = INDEX(m.deals, "2", "3").set_display_name("index_block_text")
        s.vl_text_col = VLOOKUP("Beta", m.deals, "3", False).set_display_name("vl_text_col")
        s.vl_text_cell = VLOOKUP("Beta", m.deals, i.tpos, False).set_display_name("vl_text_cell")
        s.hl_text_row = HLOOKUP(i.pick, [i.years, i.rates], "2", False).set_display_name(
            "hl_text_row")
        s.index_date = INDEX(i.rates, dt.date(2026, 1, 1)).set_display_name("index_date")
        # A key read from a blank cell finds nothing, not the 0.
        s.match_blank = MATCH(i.blank, i.zeros, 0).set_display_name("match_blank")
        s.xl_blank_key = XLOOKUP(i.blank, i.zeros, i.rates, -1).set_display_name("xl_blank_key")
        # The column or row number is checked before the key: its error wins.
        s.vl_err_zero = VLOOKUP(i.err, m.deals, 0, False).set_display_name("vl_err_zero")
        s.hl_err_zero = HLOOKUP(i.err, [i.years, i.rates], 0, False).set_display_name(
            "hl_err_zero")
        s.vl_err_key = VLOOKUP(i.err, m.deals, 3, False).set_display_name("vl_err_key")
    with pytest.warns(mo.CrossScopeReferenceWarning, match="XLOOKUP"):
        calc = _recalculated(m, tmp_path)
    lines = [(name, line) for name, line in m.s.walk() if isinstance(line, mo.Variable)]
    assert len(lines) == 53
    assert (m.s.index_text.value, m.s.index_true.value, m.s.vl_text_cell.value) == (
        0.04, 0.05, "South")
    assert (m.s.match_blank.value, m.s.xl_blank_key.value) == ("#N/A", -1)
    assert (m.s.vl_err_zero.value, m.s.hl_err_zero.value, m.s.vl_err_key.value) == (
        "#VALUE!", "#VALUE!", "#DIV/0!")
    assert m.s.series_vl.value == [100.0, 300.0, "#N/A"]
    assert m.s.match_na.value == "#N/A" and m.s.index_ref.value == "#REF!"
    for name, line in lines:
        assert _cells(calc[name], line.value) == pytest.approx(line.value), (
            name, calc[name], line.value)


@pytest.mark.slow
@pytest.mark.xfail(strict=True, reason=(
    "LibreOffice reads a boolean as a number in a lookup, matches text that reads as a "
    "number against the number, and gives #VALUE! for a VLOOKUP column or HLOOKUP row "
    "number past the block; Excel does none of these (#REF! for the last), and the "
    "engine keeps Excel's rule"))
def test_where_libreoffice_departs_from_excel(tmp_path):
    m = _model()
    with m.inputs as i:
        i.bools = mo.Variable([1, True, "TRUE"], display_name="Bools")
        i.tyear = mo.Variable("2027", display_name="TYear")
    with m.s as s:
        # A boolean matches a boolean only: Excel 2, LibreOffice 1 (TRUE is 1 there).
        s.match_bool = MATCH(True, i.bools, 0).set_display_name("match_bool")
        s.xl_bool = XLOOKUP(True, i.bools, i.years).set_display_name("xl_bool")
        # Text is not a number: Excel #N/A, LibreOffice finds 2027.
        s.match_text = MATCH("2027", i.years, 0).set_display_name("match_text")
        s.match_text_cell = MATCH(i.tyear, i.years, 0).set_display_name("match_text_cell")
        s.xl_text_cell = XLOOKUP(i.tyear, i.years, i.vals).set_display_name("xl_text_cell")
        # Past the block: Excel #REF!, found or not; LibreOffice #VALUE!.
        s.vl_ref = VLOOKUP("Beta", m.deals, 4, False).set_display_name("vl_ref")
        s.hl_ref = HLOOKUP(i.pick, [i.years, i.rates], 3, False).set_display_name("hl_ref")
        s.vl_na_ref = VLOOKUP("Delta", m.deals, 4, False).set_display_name("vl_na_ref")
        s.hl_na_ref = HLOOKUP(2025, [i.years, i.rates], 3, False).set_display_name("hl_na_ref")
        s.vl_err_ref = VLOOKUP(i.err, m.deals, 4, False).set_display_name("vl_err_ref")
    assert [m.s.match_bool.value, m.s.xl_bool.value, m.s.match_text.value,
            m.s.match_text_cell.value, m.s.xl_text_cell.value] == [2, 2027, "#N/A", "#N/A", "#N/A"]
    assert {m.s.vl_ref.value, m.s.hl_ref.value, m.s.vl_na_ref.value, m.s.hl_na_ref.value,
            m.s.vl_err_ref.value} == {"#REF!"}
    calc = _recalculated(m, tmp_path)
    departing = [(name, calc[name][0], line.value) for name, line in m.s.walk()
                 if isinstance(line, mo.Variable) and calc[name][0] != line.value]
    assert departing == [], departing
