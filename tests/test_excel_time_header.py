# SPDX-License-Identifier: Apache-2.0
"""Time in the workbook: the anchor, the header mirrors, and where a
formula reading ``mo.time`` lands.

The model's formulas read ``mo.time.days`` and know nothing of sheets.
The workbook puts time on the ``default_header`` MultiVariable's sheet
(the anchor: real formulas off one start cell), mirrors those rows at
the top of every other sheet, and points each formula at its own
sheet's header row — else at the anchor, else spells the period off the
start. These tests pin the cells, and let LibreOffice recalculate the
book to prove it computes what the engine does.
"""

import shutil
import subprocess
import warnings
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pytest

import modeleon as mo
from modeleon.core.variable_ops import _to_serial


def target(n=4, start=date(2026, 1, 1)):
    """The shape of a deal model: inputs above the header in the code,
    a header anchoring time, a loop reading days, a sheet with its own
    header, and a sheet with none."""
    m = mo.Model('pf', default_grain='month', default_periods=n)
    with m:
        m.inputs = mo.MultiVariable('Inputs')
        with m.inputs as i:
            i.financial_close = mo.Variable(start)
            i.construction_months = mo.Variable(2)
            i.cod = mo.EDATE(i.financial_close, i.construction_months)
            i.rate = mo.Variable(0.12)
            i.loan = mo.Variable(1000.0)
            i.capex = mo.Variable([10.0 * (k + 1) for k in range(n)])
        m.default_start = m.inputs.financial_close
        m.header = mo.MultiVariable('Time')
        with m.header as h:
            h.start = mo.time.start
            h.end = mo.time.end
            h.days = mo.time.days
            h.construction = mo.IF(h.end <= m.inputs.cod, 1, 0)
        m.default_header = m.header
        m.debt = mo.MultiVariable('Debt', loop=True)
        with m.debt as d:
            d.interest = (mo.lag(d.balance, fill=m.inputs.loan) * m.inputs.rate
                          * mo.time.days / 365)
            d.balance = mo.lag(d.balance, fill=m.inputs.loan) + d.interest
        m.revenue = mo.MultiVariable('Revenue')
        m.revenue.default_header = [m.header.end]
        with m.revenue as r:
            r.sales = m.inputs.capex * mo.time.days / 30
        m.notes = mo.MultiVariable('Notes', default_header=None)
        with m.notes as nt:
            nt.accrual = m.inputs.capex * mo.time.days / 365
    return m


def book(model, tmp_path, name='book.xlsx'):
    path = tmp_path / name
    with warnings.catch_warnings():
        warnings.simplefilter('ignore')
        model.to_excel(str(path))
    return openpyxl.load_workbook(str(path))


def cells(ws):
    return {c.coordinate: c.value for row in ws.iter_rows() for c in row
            if c.value is not None}


# ══════════════════════════════════════════════════════════════════════
# The anchor
# ══════════════════════════════════════════════════════════════════════


def test_the_anchor_runs_off_the_start_cell(tmp_path):
    time = cells(book(target(), tmp_path)['Time'])
    assert time['B2'] == '=Inputs!B7'          # the financial close
    assert time['C2'] == '=B3+1'               # previous end + 1
    assert time['B3'] == '=EOMONTH(B2,0)'
    assert time['B4'] == '=B3-B2+1'
    assert time['B5'] == '=IF(B3 <= Inputs!B9, 1, 0)'


def test_a_string_start_is_a_typed_date(tmp_path):
    m = mo.Model('m', default_grain='quarter', default_start='2026-Q2', default_periods=2)
    with m:
        m.header = mo.MultiVariable('Time')
        with m.header as h:
            h.start = mo.time.start
            h.end = mo.time.end
        m.default_header = m.header
    time = cells(book(m, tmp_path)['Time'])
    assert time['B2'] == '=DATE(2026,4,1)'
    assert time['B3'] == '=EOMONTH(B2,2)'


# ══════════════════════════════════════════════════════════════════════
# Mirrors
# ══════════════════════════════════════════════════════════════════════


def test_every_sheet_mirrors_the_header_whatever_the_code_order(tmp_path):
    wb = book(target(), tmp_path)
    inputs = cells(wb['Inputs'])            # declared ABOVE the header
    assert [inputs[f'A{r}'] for r in (2, 3, 4, 5)] == ['Start', 'End', 'Days', 'Construction']
    # Constants hold column B; the periods (and so the mirrors) start at C.
    assert inputs['C2'] == '=Time!B2' and inputs['D4'] == '=Time!C4'
    debt = cells(wb['Debt'])
    assert debt['B4'] == '=Time!B4'


def test_a_sheet_shows_its_own_rows(tmp_path):
    revenue = cells(book(target(), tmp_path)['Revenue'])
    assert revenue['A2'] == 'End' and revenue['B2'] == '=Time!B3'
    assert 'A3' not in revenue


def test_none_means_no_header(tmp_path):
    notes = cells(book(target(), tmp_path)['Notes'])
    # Only the captions above the data — no header rows, no gap.
    assert notes['A2'] == 'Accrual'
    assert sorted(k for k in notes if k.startswith('A')) == ['A2']


def test_the_header_is_frozen_with_the_captions(tmp_path):
    wb = book(target(), tmp_path)
    assert wb['Debt'].freeze_panes == 'B6'
    assert wb['Time'].freeze_panes == 'B2'


# ══════════════════════════════════════════════════════════════════════
# Where a formula reads time
# ══════════════════════════════════════════════════════════════════════


def test_a_formula_reads_its_own_header_row(tmp_path):
    debt = cells(book(target(), tmp_path)['Debt'])
    assert debt['B7'] == '=(Inputs!B11 * Inputs!B10 * B4) / 365'


def test_without_the_row_it_reads_the_anchor(tmp_path):
    wb = book(target(), tmp_path)
    assert cells(wb['Revenue'])['B4'] == '=(Inputs!C12 * Time!B4) / 30'
    assert cells(wb['Notes'])['B2'] == '=(Inputs!C12 * Time!B4) / 365'


def test_captions_are_the_period_end(tmp_path):
    wb = book(target(), tmp_path)
    debt, revenue = wb['Debt'], wb['Revenue']
    assert debt['B1'].value == '=B3' and debt['B1'].number_format == 'mmm yyyy'
    assert revenue['C1'].value == '=C2'
    assert wb['Notes']['B1'].value == '=Time!B3'


def test_a_lag_of_time_reads_the_previous_period(tmp_path):
    m = target()
    with m:
        m.debt.default_header = None
    with m.revenue as r:
        r.prev_end = mo.lag(mo.time.end, fill=date(2025, 12, 31))
    revenue = cells(book(m, tmp_path)['Revenue'])
    assert revenue['C5'] == '=B2'


# ══════════════════════════════════════════════════════════════════════
# No header: time is spelled off the start, the book is otherwise as before
# ══════════════════════════════════════════════════════════════════════


def simple(start='2026-01'):
    m = mo.Model('s', default_grain='month', default_start=start, default_periods=3)
    with m:
        m.calc = mo.MultiVariable('Calc')
        with m.calc as c:
            c.rate = mo.Variable(0.12)
            c.bal = mo.Variable([1000.0, 900, 800])
            c.interest = c.bal * c.rate * mo.time.days / 365
    return m


def test_without_a_header_time_is_spelled_off_the_start(tmp_path):
    calc = cells(book(simple(), tmp_path)['Calc'])
    assert calc['C1'] == 'Jan 2026'                    # captions stay text
    assert calc['D4'] == ('=(D3 * B2 * (EOMONTH(DATE(2026,1,1),1)'
                          '-EOMONTH(DATE(2026,1,1),0))) / 365')


def test_without_a_header_a_start_cell_still_anchors(tmp_path):
    m = mo.Model('s', default_grain='month', default_periods=2)
    with m:
        m.inputs = mo.MultiVariable('Inputs')
        m.inputs.fc = mo.Variable(date(2026, 3, 1))
        m.default_start = m.inputs.fc
        m.calc = mo.MultiVariable('Calc')
        m.calc.end = mo.Variable(1.0) * 0 + mo.time.days
    calc = cells(book(m, tmp_path)['Calc'])
    assert calc['B2'] == '=1.0 * 0 + (EOMONTH(Inputs!B1,0)-EOMONTH(Inputs!B1,-1))'


def test_a_book_without_time_keeps_its_shape(tmp_path):
    m = mo.Model('plain', default_grain='month', default_start='2026-01', default_periods=2)
    with m:
        m.p = mo.MultiVariable('P')
        m.p.x = mo.Variable([1.0, 2.0])
    p = cells(book(m, tmp_path)['P'])
    assert p == {'B1': 'Jan 2026', 'C1': 'Feb 2026', 'A2': 'X', 'B2': 1, 'C2': 2}


# ══════════════════════════════════════════════════════════════════════
# The book computes what the engine does
# ══════════════════════════════════════════════════════════════════════


def _soffice():
    found = shutil.which('soffice') or shutil.which('libreoffice')
    if found:
        return found
    mac = Path('/Applications/LibreOffice.app/Contents/MacOS/soffice')
    return str(mac) if mac.exists() else None


def _recalculated(model, tmp_path):
    soffice = _soffice()
    if soffice is None:
        pytest.skip('LibreOffice is not installed')
    src = tmp_path / 'calc-src.xlsx'
    with warnings.catch_warnings():
        warnings.simplefilter('ignore')
        model.to_excel(str(src))
    outdir = tmp_path / 'calc'
    profile = (tmp_path / 'lo-profile').as_uri()
    subprocess.run(
        [soffice, f'-env:UserInstallation={profile}', '--headless',
         '--convert-to', 'xlsx', '--outdir', str(outdir), str(src)],
        check=True, capture_output=True, timeout=120,
    )
    return openpyxl.load_workbook(str(outdir / 'calc-src.xlsx'), data_only=True)


def _serial(v):
    if isinstance(v, datetime):
        v = v.date()
    return round(_to_serial(v), 6) if isinstance(v, (int, float, date)) else v


@pytest.mark.parametrize('grain_book', ['month', 'quarter'])
@pytest.mark.slow
def test_the_workbook_recalculates_to_the_engine(tmp_path, grain_book):
    m = target(6)
    if grain_book == 'quarter':
        # A quarterly book of the same model (the debt loop re-grains by
        # its balances; the rest is recomputed for quarters).
        m.debt.interest.set_regrain(mo.frozen(values='sum'))
        m.debt.balance.set_regrain(mo.frozen(values='last'))
        m.inputs.capex.set_regrain(mo.up('sum'))
        m = m.at('quarter')
    wb = _recalculated(m, tmp_path)
    rows = {
        'Inputs': ['rate', 'loan', 'cod'],
        'Time': ['start', 'end', 'days', 'construction'],
        'Debt': ['interest', 'balance'],
        'Revenue': ['sales'],
        'Notes': ['accrual'],
    }
    for sheet, names in rows.items():
        ws = wb[sheet]
        by_label = {r[0].value: [c.value for c in r[1:] if c.value is not None]
                    for r in ws.iter_rows() if r[0].value is not None}
        mv = getattr(m, sheet.lower() if sheet != 'Time' else 'header')
        for name in names:
            var = getattr(mv, name)
            label = var.display_name
            got = [_serial(v) for v in by_label[label]]
            value = var.value
            want = [_serial(v) for v in (value if isinstance(value, list) else [value])]
            assert got == pytest.approx(want), (sheet, name, got, want)


def test_a_lag_of_time_past_the_window_is_its_fill(tmp_path):
    m = mo.Model('m', default_grain='month', default_start='2026-01', default_periods=3)
    with m:
        m.c = mo.MultiVariable('C')
        m.c.nxt = mo.lag(mo.time.days, -1, fill=0)
    c = cells(book(m, tmp_path)['C'])
    assert c['D2'] == '=0'


@pytest.mark.slow
def test_a_lag_of_an_unnamed_time_formula_recalculates(tmp_path):
    m = mo.Model('m', default_grain='month', default_start='2026-01', default_periods=4)
    with m:
        m.c = mo.MultiVariable('C')
        m.c.g = mo.lag(mo.time.days * 2, fill=0)
    wb = _recalculated(m, tmp_path)
    assert _values(wb, 'C', 'G') == [[0, 62, 56, 62]] == [m.c.g.value]


@pytest.mark.slow
def test_a_header_sheet_with_an_apostrophe_recalculates(tmp_path):
    m = mo.Model('m', default_grain='month', default_start='2026-01', default_periods=3)
    with m:
        m.header = mo.MultiVariable("Bob's time")
        m.header.days = mo.time.days
        m.default_header = m.header
        m.c = mo.MultiVariable('C')
        m.c.x = mo.time.days * 2
    assert cells(book(m, tmp_path)['C'])['B2'] == "='Bob''s time'!B2"
    wb = _recalculated(m, tmp_path)
    assert _values(wb, 'C', 'X') == [[62, 56, 62]]


def _values(wb, sheet, label):
    ws = wb[sheet]
    return [[_serial(c.value) for c in r[1:] if c.value is not None]
            for r in ws.iter_rows() if r[0].value == label]


@pytest.mark.slow
def test_each_window_counts_its_own_periods(tmp_path):
    m = mo.Model('w', default_grain='month', default_start='2026-01', default_periods=6)
    with m:
        m.qtr = mo.MultiVariable('Qtr', default_grain='quarter', default_start='2026-Q2',
                                 default_periods=2)
        with m.qtr as q:
            q.d = mo.time.days * 1.0
        m.mon = mo.MultiVariable('Mon')
        with m.mon as a:
            a.d = mo.time.days * 1.0
    wb = _recalculated(m, tmp_path)
    assert _values(wb, 'Qtr', 'D') == [[91, 92]]
    assert _values(wb, 'Mon', 'D') == [[31, 28, 31, 30, 31, 30]]


@pytest.mark.slow
def test_a_quarterly_copy_in_the_same_book_reads_its_own_anchor(tmp_path):
    t = target(6)
    t.debt.interest.set_regrain(mo.frozen(values='sum'))
    t.debt.balance.set_regrain(mo.frozen(values='last'))
    t.inputs.capex.set_regrain(mo.up('sum'))
    w = mo.MultiVariable('W')
    w.mon, w.qtr = t, t.at('quarter')
    wb = _recalculated(w, tmp_path)
    assert _values(wb, 'Pf1', 'Days') == [[90, 91]]
    assert _values(wb, 'Pf1', 'Interest') == [pytest.approx(w.qtr.debt.interest.value)]


def test_a_tracked_sheet_keeps_its_header(tmp_path):
    m = mo.Model('tk', default_grain='month', default_periods=3,
                 tracks=mo.Tracks('plan', 'actual'))
    with m:
        m.inputs = mo.MultiVariable('Inputs')
        m.inputs.fc = mo.Variable(date(2026, 1, 1))
        m.default_start = m.inputs.fc
        m.header = mo.MultiVariable('Time')
        with m.header as h:
            h.end = mo.time.end
            h.days = mo.time.days
        m.default_header = m.header
        m.rev = mo.MultiVariable('Revenue', default_header=[m.header.days])
        with m.rev as r:
            r.vol = mo.Variable(plan=[10.0, 20, 30], actual=[11.0, 19, 29])
            r.x = r.vol * mo.time.days
    rev = cells(book(m, tmp_path)['Revenue'])
    assert rev['A2'] == 'Days' and rev['B2'] == '=Time!B3'
    assert rev['B6'] == '=B4 * B2' and rev['B7'] == '=B5 * B2'     # plan and actual


def test_months_is_a_header_row_and_a_formula_reads_it(tmp_path):
    """``mo.time.months`` writes as its grain's constant (a quarter holds 3)
    on the anchor sheet; a sheet's header mirrors it and a formula reading
    months points at the row, as one reading ``days`` does."""
    m = mo.Model('m', default_grain='quarter', default_start='2026-Q1', default_periods=2)
    with m:
        m.a = mo.MultiVariable('Inputs')
        with m.a as a:
            a.rent = mo.Variable(100.0)
        m.header = mo.MultiVariable('Time')
        with m.header as h:
            h.months = mo.time.months
        m.default_header = m.header
        m.opex = mo.MultiVariable('Opex')
        with m.opex as o:
            o.rent = m.a.rent * mo.time.months
        m.notes = mo.MultiVariable('Notes', default_header=None)
        with m.notes as nt:
            nt.rent = m.a.rent * mo.time.months
    wb = book(m, tmp_path)
    time_sheet = cells(wb['Time'])
    assert time_sheet['A2'] == 'Months' and time_sheet['B2'] == '=3' and time_sheet['C2'] == '=3'
    opex = cells(wb['Opex'])
    assert opex['B2'] == '=Time!B2'                 # the mirrored header row
    assert opex['B4'] == '=Inputs!B4 * B2'          # the formula reads the mirror
    notes = cells(wb['Notes'])
    assert notes['B2'] == '=Inputs!B4 * Time!B2'    # no header here: the anchor's row


def test_without_any_header_months_is_spelled_inline(tmp_path):
    m = mo.Model('m', default_grain='year', default_start='2026', default_periods=2)
    with m:
        m.a = mo.MultiVariable('Inputs')
        with m.a as a:
            a.rent = mo.Variable(100.0)
        m.opex = mo.MultiVariable('Opex')
        with m.opex as o:
            o.rent = m.a.rent * mo.time.months
    opex = cells(book(m, tmp_path)['Opex'])
    formulas = [v for v in opex.values() if isinstance(v, str) and v.startswith('=')]
    assert formulas and all(v.endswith('* 12') for v in formulas), formulas
