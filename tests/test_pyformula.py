# SPDX-License-Identifier: Apache-2.0
"""``@pyformula`` decorator — turn opaque Python into an AST-tracked node.

The decorator's job is to preserve dependency edges for functions that
can't be expressed as Variable arithmetic (scipy, trained models, API
calls). Result is a Variable with a :class:`FuncCall` AST whose
``render_backends`` is the empty set, so every renderer falls back to
inlining ``_value``. The call itself still surfaces in JSON, in
``check_compat``, and as a dep edge for graph walkers.
"""

from __future__ import annotations

import modeleon as mo


class TestBareDecorator:
    """``@mo.pyformula`` without arguments uses defaults."""

    def test_returns_variable_with_funccall_ast(self):
        @mo.pyformula
        def double(x):
            return x * 2

        v = mo.Variable(5)
        r = double(v)

        assert isinstance(r, mo.Variable)
        assert r._value == 10
        # AST is a FuncCall named after the function (uppercased).
        assert r._expr.func == "DOUBLE"
        # Argument is a VarRef back to the original Variable — this is
        # what preserves the dep edge.
        (arg,) = r._expr.args
        assert arg.var is v

    def test_dep_edge_present(self):
        @mo.pyformula
        def triple(x):
            return x * 3

        v = mo.Variable(4)
        r = triple(v)
        # dependencies is path-keyed; v's id should show up as a dep.
        assert v.id in r.dependencies


class TestRenderBackendsEmpty:
    """Empty ``render_backends`` — every renderer falls back."""

    def test_json_shows_empty_render_backends(self):
        @mo.pyformula
        def magic(x):
            return 42.0

        v = mo.Variable(1.0)
        r = magic(v)
        out = mo.to_json(r._expr)
        assert out["kind"] == "funccall"
        assert out["func"] == "MAGIC"
        assert out["render_backends"] == []

    def test_check_compat_flags_against_any_target(self):
        @mo.pyformula
        def opaque(x):
            return 1.0

        v = mo.Variable(2.0)
        r = opaque(v)
        # Empty set excludes every target, so any target is incompatible.
        for target in ("excel", "pandas", "sql"):
            issues = mo.check_compat(r, target)
            assert len(issues) == 1
            assert issues[0].func == "OPAQUE"


class TestDecoratorWithArgs:
    """``@mo.pyformula(name=..., value_type=...)`` overrides defaults."""

    def test_custom_name(self):
        @mo.pyformula(name="SCIPY_MIN")
        def minimize(x):
            return x * 0.5

        r = minimize(mo.Variable(10.0))
        assert r._expr.func == "SCIPY_MIN"

    def test_custom_value_type(self):
        @mo.pyformula(value_type='int')
        def floor_div(x):
            return int(x)

        r = floor_div(mo.Variable(7.9))
        assert r.value_type == 'int'


class TestMixedArgs:
    """Scalar and Variable args are each represented correctly."""

    def test_scalar_becomes_literal(self):
        @mo.pyformula
        def scale(x, factor):
            return x * factor

        v = mo.Variable(10.0)
        r = scale(v, 3.5)
        args = r._expr.args
        assert len(args) == 2
        # First arg: Variable → VarRef
        assert args[0].var is v
        # Second arg: scalar → Literal
        assert args[1].value == 3.5


class TestListResult:
    """List-returning function → ``var_type='list'``."""

    def test_list_var_type(self):
        @mo.pyformula
        def repeat(x):
            return [x] * 3

        r = repeat(mo.Variable(7))
        assert r.var_type == 'list'
        assert r._value == [7, 7, 7]


class TestRendersAsLiteral:
    """Excel renderer falls back to the inlined value, no broken formula."""

    def test_excel_inlines_value(self, tmp_path):
        @mo.pyformula
        def compute(x):
            return float(x) * 10

        model = mo.MultiVariable("M")
        model.s = mo.MultiVariable("S", excel_props={'tab': True})
        with model.s as s:
                s.src = mo.Variable(7.0)
                s.out = compute(s.src)

        model.to_excel(tmp_path / "out.xlsx")

        path = tmp_path / "out.xlsx"
        from openpyxl import load_workbook
        ws = load_workbook(path)["S"]
        formulas = [c.value for row in ws.iter_rows() for c in row
                    if c.value is not None]
        # Fallback should inline 70.0 (no =COMPUTE(...) formula).
        assert any(
            "COMPUTE" not in str(f) and ("70" in str(f) or f == 70.0)
            for f in formulas
        ), f"Expected inlined value, got formulas: {formulas}"
