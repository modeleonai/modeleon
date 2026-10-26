# SPDX-License-Identifier: Apache-2.0
"""``CHOOSE`` over a line that runs along keys is one choice per key.

A keyed line (``mo.Variable({"Bear": 0.9, "Base": 1.0, "Bull": 1.1})``)
or a line on one finite axis that is not time runs along KEYS, not
periods. ``mo.CHOOSE(pick, growth)`` used to read it as ONE series
choice: its keys were lost, a 3-key line in a 4-period window failed
the length check (``'g' has 3 value(s) in a 4-period window``), in a
3-period window it silently became a period series, and the workbook
held ``=CHOOSE(B1, B3)`` in every period — one option, ``#VALUE!`` in
Excel. Such a line is now spread into its elements, one choice per key
in key order, so ``mo.CHOOSE(pick, growth)`` is ``mo.CHOOSE(pick,
growth["Bear"], growth["Base"], growth["Bull"])``: value ``1.0`` for
``pick = 2``, written ``=CHOOSE(B1, B3, C3, D3)`` over the line's cells
(one such formula per period when the index is a series). A tracked
line, an index that runs along keys and a line on two axes are refused
with an error that names the fix. Bare series choices keep their branch:
``mo.CHOOSE(scenario, base, bull)`` still writes one CHOOSE per period
over that period's cells.
"""

import pytest
from openpyxl import load_workbook

import modeleon as mo


def _rows(model, sheet, tmp_path, n, cols="ABCDEFGH"):
    path = tmp_path / "t.xlsx"
    model.to_excel(str(path))
    ws = load_workbook(str(path))[sheet]
    return [[ws[f"{c}{r}"].value for c in cols] for r in range(1, n + 1)]


def _keyed():
    return mo.Variable({"Bear": 0.9, "Base": 1.0, "Bull": 1.1}, display_name="Growth")


def _scenario_axis():
    return mo.Variable(["Bear", "Base", "Bull"])


def _model(pick):
    m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=4)
    with m:
        m.inputs = mo.MultiVariable("Inputs", excel_props={"tab": True})
        with m.inputs as i:
            i.pick = mo.Variable(pick, display_name="Pick")
            i.growth = _keyed()
        m.out = mo.MultiVariable("Out", excel_props={"tab": True})
        m.out.g = mo.CHOOSE(i.pick, i.growth).set_display_name("G")
    return m


def _tracked_model():
    m = mo.Model("t", default_grain="month", default_start="2026-01", default_periods=3,
                 tracks=mo.Tracks("plan", "actual"))
    with m:
        m.x = mo.Variable(plan=[1, 2, 3], actual=[1, 2, 3], display_name="X")
        m.pick = mo.Variable(1)
    return m


def _grid():
    a = mo.Variable(["x", "y"])
    b = mo.Variable(["p", "q"])
    return mo.Variable([1, 2, 3, 4], indexed_by=[a, b], display_name="Grid")


class TestTheValue:
    def test_a_keyed_choice_is_one_choice_per_key(self):
        g = _keyed()
        r = mo.CHOOSE(mo.Variable(2), g)
        assert r._value == 1.0
        assert r.var_type == "scalar"
        assert mo.CHOOSE(3, g)._value == 1.1
        assert mo.CHOOSE(4, g)._value == "#VALUE!"

    def test_the_short_form_is_the_long_form(self):
        g = _keyed()
        pick = mo.Variable(2)
        short = mo.CHOOSE(pick, g)
        long = mo.CHOOSE(pick, g["Bear"], g["Base"], g["Bull"])
        assert short.formula == long.formula
        assert short._value == long._value
        # The spread is an AST of element references, never a value.
        assert short._expr is not None
        assert len(short._expr.args) == 4

    def test_a_line_on_a_finite_axis_spreads_by_position(self):
        scen = _scenario_axis()
        ga = mo.Variable([0.9, 1.0, 1.1], indexed_by=[scen], display_name="Growth")
        assert mo.CHOOSE(3, ga)._value == 1.1
        assert mo.CHOOSE(3, ga).formula == mo.CHOOSE(3, ga[0], ga[1], ga[2]).formula

    def test_a_per_period_index_picks_a_key_each_period(self):
        r = mo.CHOOSE(mo.Variable([1, 2, 3, 2]), _keyed())
        assert r._value == [0.9, 1.0, 1.1, 1.0]
        assert r.var_type == "list"

    def test_a_keyed_choice_beside_other_choices_is_refused(self):
        # Beside other choices the numbering would depend on how many keys
        # the line has, and two keyed lines would flatten into one row.
        with pytest.raises(ValueError, match="beside 2 other choice") as e:
            mo.CHOOSE(2, 0.5, _keyed(), 2.0)
        assert 'mo.CHOOSE(pick, Growth)' in str(e.value)
        assert "Growth['Bear'], Growth['Base'], Growth['Bull']" in str(e.value)
        base = mo.Variable({"A": 10.0, "B": 20.0}, display_name="Base")
        bull = mo.Variable({"A": 12.0, "B": 25.0}, display_name="Bull")
        with pytest.raises(ValueError, match="beside 1 other choice"):
            mo.CHOOSE(2, base, bull)

    def test_a_keyed_choice_beside_a_series_choice_is_refused_too(self):
        with pytest.raises(ValueError, match="beside 1 other choice"):
            mo.CHOOSE(1, _keyed(), mo.Variable([10, 20, 30]))

    def test_integer_keys_spread_by_position(self):
        # An int key inside the positional range is read as a POSITION by
        # ``[]`` and by the workbook alike, so the spread goes by position
        # and both agree — key order, whatever the ints say.
        g = mo.Variable({1: 0.9, 2: 1.0, 3: 1.1})
        assert [mo.CHOOSE(p, g)._value for p in (1, 2, 3)] == [0.9, 1.0, 1.1]
        g2 = mo.Variable({3: 0.9, 2: 1.0, 1: 1.1})
        assert [mo.CHOOSE(p, g2)._value for p in (1, 2, 3)] == [0.9, 1.0, 1.1]

    def test_a_computed_keyed_choice_is_refused_until_named(self, tmp_path):
        m = mo.Model("m", default_grain="month", default_start="2026-01", default_periods=3)
        with m:
            m.pick = mo.Variable(2, display_name="Pick")
            m.growth = _keyed()
            with pytest.raises(ValueError, match="computed line with no cells") as e:
                mo.CHOOSE(m.pick, m.growth * 2)
            assert "growth * 2" in str(e.value)
            m.growth_adj = mo.Variable(m.growth * 2, display_name="Growth adj")
            m.g = mo.CHOOSE(m.pick, m.growth_adj).set_display_name("G")
        assert m.g._value == 2.0
        rows = _rows(m, "M", tmp_path, 6, "ABCDE")
        g_row = next(r for r in rows if r[0] == "G")
        # live: one CHOOSE over the three cells of the NAMED computed line
        assert g_row[1] == "=CHOOSE(B1, B4, C4, D4)"

    def test_the_lines_unit_carries(self):
        from modeleon.core.unit import unit_key

        g = mo.Variable({"Bear": 0.9, "Base": 1.0, "Bull": 1.1}, unit="$")
        r = mo.CHOOSE(1, g)
        assert unit_key(r._unit) == unit_key(g._unit)


class TestTheRenderers:
    def test_the_workbook_writes_one_choose_over_the_lines_cells(self, tmp_path):
        m = _model(2)
        # No per-period line on Inputs, so no period header: Pick is row
        # 1, the caption row of keys is row 2, Growth's cells are row 3.
        assert _rows(m, "Inputs", tmp_path, 3, "ABCD") == [
            ["Pick", 2, None, None],
            [None, "Bear", "Base", "Bull"],
            ["Growth", 0.9, 1, 1.1],
        ]
        assert _rows(m, "Out", tmp_path, 1, "AB") == [
            ["G", "=CHOOSE(Inputs!B1, Inputs!B3, Inputs!C3, Inputs!D3)"],
        ]

    def test_a_series_index_writes_one_choose_per_period_over_the_same_cells(
        self, tmp_path,
    ):
        m = _model([1, 2, 3, 2])
        assert m.out.g._value == [0.9, 1.0, 1.1, 1.0]
        # Inputs now carries a period header (row 1): Pick is row 2, the
        # caption row 3, Growth row 4. Each period's CHOOSE reads its own
        # Pick cell and the same three Growth cells.
        rows = _rows(m, "Out", tmp_path, 2, "ABCDE")
        assert rows[1] == ["G"] + [
            f"=CHOOSE(Inputs!{c}2, Inputs!B4, Inputs!C4, Inputs!D4)" for c in "BCDE"
        ]

    def test_the_notebook_repr_shows_the_same_formula(self):
        assert "=CHOOSE(Inputs!B1, Inputs!B3, Inputs!C3, Inputs!D3)" in _model(2)._repr_html_()

    def test_the_json_carries_one_argument_per_choice(self):
        out = mo.to_json(_model(2).out.g._expr)
        assert out["kind"] == "funccall"
        assert out["func"] == "CHOOSE"
        assert [a["kind"] for a in out["args"]] == ["varref"] * 4
        assert out["args"][0]["path"] == "m.inputs.pick"

    def test_the_choice_depends_on_the_line(self):
        m = _model(2)
        deps = m.out.g._dependency_refs
        assert deps[0] is m.inputs.pick
        assert len(deps) == 4
        assert all(e._dependency_refs == [m.inputs.growth] for e in deps[1:])


class TestTheRefusals:
    def test_tracked_choices_still_switch_track_by_track(self, tmp_path):
        # The usual scenario switch over plan/actual rows: one index, tracked
        # choices — computed per track, one CHOOSE row per track in the book.
        m = mo.Model("t", default_grain="month", default_start="2026-01", default_periods=3,
                     tracks=mo.Tracks("plan", "actual"))
        with m:
            m.scenario = mo.Variable(2, display_name="Scenario")
            m.base = mo.Variable(plan=[100, 110, 120], actual=[101, 111, 121], display_name="Base")
            m.bull = mo.Variable(plan=[150, 170, 190], actual=[151, 171, 191], display_name="Bull")
            m.pick = mo.CHOOSE(m.scenario, m.base, m.bull).set_display_name("Pick")
        assert list(m.pick._value['plan']) == [150, 170, 190]
        assert list(m.pick._value['actual']) == [151, 171, 191]
        rows = _rows(m, "T", tmp_path, 8, "ABCD")
        formulas = [r for r in rows if r[0] and str(r[0]).startswith("Pick")]
        assert len(formulas) == 2
        assert all(str(r[2]).startswith("=CHOOSE(") for r in formulas)

    def test_a_series_index_over_tracked_choices_is_refused(self):
        m = mo.Model("t", default_grain="month", default_start="2026-01", default_periods=3,
                     tracks=mo.Tracks("plan", "actual"))
        with m:
            m.base = mo.Variable(plan=[100, 110, 120], actual=[101, 111, 121], display_name="Base")
            m.flag = mo.Variable([1, 2, 1], display_name="Flag")
            with pytest.raises(ValueError, match="each period would pick a whole tracked") as e:
                mo.CHOOSE(m.flag, m.base, 5)
            assert "base.plan" in str(e.value)

    def test_a_tracked_index_is_refused(self):
        m = _tracked_model()
        with pytest.raises(ValueError, match="tracked line"):
            mo.CHOOSE(m.x, 1, 2)

    def test_a_choice_on_two_axes_is_refused(self):
        grid = _grid()
        with pytest.raises(ValueError, match="2 axes"):
            mo.CHOOSE(1, grid)
        with pytest.raises(ValueError, match="2 axes"):
            mo.CHOOSE(grid, 1, 2)

    def test_an_index_along_keys_is_refused(self):
        with pytest.raises(ValueError, match="runs along keys") as e:
            mo.CHOOSE(mo.Variable({"a": 1, "b": 2}, display_name="Pick"), 10, 20)
        assert "Pick['a']" in str(e.value)
        scen = _scenario_axis()
        axis_pick = mo.Variable([1, 2, 3], indexed_by=[scen], display_name="Pick")
        with pytest.raises(ValueError, match="runs along keys") as e:
            mo.CHOOSE(axis_pick, 10, 20, 30)
        assert "Pick[0]" in str(e.value)


class TestWhatStays:
    def test_a_line_on_the_time_axis_is_still_a_series(self):
        t = mo.Time("2026-01", periods=3)
        tl = mo.Variable([1, 2, 3], indexed_by=[t], display_name="TL")
        first = mo.CHOOSE(1, tl, 5)
        assert first._value == [1, 2, 3]
        assert first.var_type == "list"
        second = mo.CHOOSE(2, tl, 5)
        assert second._value == [5, 5, 5]
        assert second.var_type == "list"

    def test_bare_series_choices_are_unchanged(self):
        r = mo.CHOOSE(mo.Variable(2), mo.Variable([100, 110, 120]), mo.Variable([150, 170, 190]))
        assert r._value == [150, 170, 190]
        # A series is one choice, never spread.
        assert len(r._expr.args) == 3
        assert mo.CHOOSE(1, 125, mo.Variable([150, 170, 190]))._value == [125, 125, 125]
        assert mo.CHOOSE(mo.Variable([0, 0, 1, 1]) + 1, "A", "F")._value == ["A", "A", "F", "F"]

    def test_line_keys_lives_in_core_and_the_layout_shares_it(self):
        from modeleon.compile.excel import layout
        from modeleon.core.shape import line_keys

        assert layout.line_keys is line_keys
        scen = _scenario_axis()
        assert line_keys(_keyed()) == ["Bear", "Base", "Bull"]
        axised = mo.Variable([0.9, 1.0, 1.1], indexed_by=[scen])
        assert line_keys(axised) == ["Bear", "Base", "Bull"]
        on_time = mo.Variable([1, 2, 3], indexed_by=[mo.Time("2026-01", periods=3)])
        assert line_keys(on_time) is None
        assert line_keys(mo.Variable([1, 2, 3])) is None
        assert line_keys(mo.Variable(5)) is None
        assert line_keys(_grid()) is None
        assert line_keys(_tracked_model().x) is None


def test_a_wrapper_keeps_the_keys_and_axes_of_the_line_it_wraps():
    g = mo.Variable({"Bear": 0.9, "Base": 1.0, "Bull": 1.1})
    assert mo.Variable(g * 2, display_name="Twice")._keys == ["Bear", "Base", "Bull"]
    scen = mo.Variable(["Bear", "Base", "Bull"])
    a = mo.Variable([1, 2, 3], indexed_by=[scen])
    assert mo.Variable(a * 2, display_name="Twice")._indexed_by == (scen,)
    assert mo.Variable(g, keys=["x", "y", "z"])._keys == ["x", "y", "z"], "a declared key set wins"
