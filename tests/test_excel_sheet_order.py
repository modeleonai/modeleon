# SPDX-License-Identifier: Apache-2.0
"""The book's sheets in the order the model declares: ``ExcelView(sheet={'order': [...]})``.

A model lays its sheets out in the order it declares them. A book that opens on
its summary says so once, on the model's own view, by the attributes the sheets
are assigned to; the sheets it leaves out follow in their declared order. Every
renderer that lists the sheets reads the same declaration.
"""

import re
import warnings

import openpyxl
import pytest

import modeleon as mo
from modeleon.compile.excel.view import SheetOrderWarning


def _model(order=None, tracks=False):
    kwargs = {"tracks": mo.Tracks("plan", "actual")} if tracks else {}
    m = mo.Model("m", default_grain="year", default_start="2026", default_periods=3, **kwargs)
    with m:
        m.default_excel_view = mo.ExcelView(sheet={"order": order} if order is not None else None)
        m.inputs = mo.MultiVariable("Inputs")
        m.inputs.price = mo.Variable(5.0, display_name="Price")
        m.ops = mo.MultiVariable("Operations")
        if tracks:
            m.ops.volume = mo.Variable(plan=[1.0, 2.0, 3.0], actual=[1.0, 2.5, 3.0],
                                       display_name="Volume")
        else:
            m.ops.volume = mo.Variable([1.0, 2.0, 3.0], display_name="Volume")
        m.summary = mo.MultiVariable("Summary")
        m.summary.revenue = (m.ops.volume * m.inputs.price).set_display_name("Revenue")
    return m


def _book(model, tmp_path):
    path = tmp_path / "book.xlsx"
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        model.to_excel(str(path))
    return openpyxl.load_workbook(str(path))


def test_without_an_order_the_sheets_stand_as_declared(tmp_path):
    assert _book(_model(), tmp_path).sheetnames == ["Inputs", "Operations", "Summary"]


def test_a_declared_order_puts_its_sheets_first(tmp_path):
    assert _book(_model(["summary"]), tmp_path).sheetnames == [
        "Summary", "Inputs", "Operations",
    ]
    assert _book(_model(["ops", "summary", "inputs"]), tmp_path).sheetnames == [
        "Operations", "Summary", "Inputs",
    ]


def test_references_follow_the_sheets_they_name(tmp_path):
    # Moving a sheet moves no cell: Summary's formula still reads the
    # Operations and Inputs cells it read before.
    plain = _book(_model(), tmp_path)["Summary"]
    moved = _book(_model(["summary"]), tmp_path)["Summary"]
    formulas = [c.value for row in plain.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=")]
    assert formulas
    assert formulas == [c.value for row in moved.iter_rows() for c in row
                        if isinstance(c.value, str) and c.value.startswith("=")]
    assert any("Operations!" in f and "Inputs!" in f for f in formulas)


def test_a_tracked_book_keeps_the_order(tmp_path):
    assert _book(_model(["summary"], tracks=True), tmp_path).sheetnames[0] == "Summary"


def test_a_name_that_is_no_sheet_is_skipped_and_the_rest_placed(tmp_path):
    # A typo never stops the book: the name is skipped, said once, with the
    # sheets there are, and the names that are sheets still lead.
    with pytest.warns(SheetOrderWarning) as caught:
        _model(["Summary", "ops"]).to_excel(str(tmp_path / "book.xlsx"))
    assert openpyxl.load_workbook(str(tmp_path / "book.xlsx")).sheetnames == [
        "Operations", "Inputs", "Summary",
    ]
    message = str(caught[0].message)
    assert "'Summary'" in message and "skips it" in message
    assert "'inputs', 'ops', 'summary'" in message


def test_one_name_or_a_repeat_is_an_order_too(tmp_path):
    assert _book(_model("summary"), tmp_path).sheetnames[0] == "Summary"
    assert _book(_model(["summary", "summary"]), tmp_path).sheetnames == [
        "Summary", "Inputs", "Operations",
    ]


@pytest.mark.parametrize("order", [[1], 5, ["summary", None]])
def test_the_order_is_made_of_names(order):
    with pytest.raises(ValueError):
        mo.ExcelView(sheet={"order": order})


def test_a_part_of_the_model_laid_out_alone_reads_no_order(tmp_path):
    # The order describes the whole book; the summary written on its own
    # is no such book, and the names it does not have are no error there.
    m = _model(["summary"])
    path = tmp_path / "part.xlsx"
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        m.summary.to_excel(str(path))
    assert openpyxl.load_workbook(str(path)).sheetnames == ["Summary"]


def test_the_notebook_tabs_follow_the_order():
    html = _model(["summary"])._repr_html_()
    assert html.index(">Summary<") < html.index(">Inputs<") < html.index(">Operations<")


def _with_own_lines(order=None):
    m = mo.Model("m", default_grain="year", default_start="2026", default_periods=3)
    with m:
        m.default_excel_view = mo.ExcelView(sheet={"order": order})
        m.note = mo.Variable(1.0, display_name="Note")
        m.inputs = mo.MultiVariable("Inputs")
        m.inputs.price = mo.Variable(5.0, display_name="Price")
        m.summary = mo.MultiVariable("Summary")
        m.summary.total = (m.inputs.price * 2).set_display_name("Total")
    return m


def test_the_tab_of_the_models_own_lines_goes_by_the_models_name(tmp_path):
    own, inputs, summary = _book(_with_own_lines(), tmp_path).sheetnames
    assert _book(_with_own_lines(["summary", "m"]), tmp_path).sheetnames == [
        summary, own, inputs,
    ]
    with pytest.warns(SheetOrderWarning, match=re.escape("'m', 'inputs', 'summary'")):
        _with_own_lines(["Summary"]).to_excel(str(tmp_path / "own.xlsx"))


def test_a_view_that_declares_no_order_is_not_read_for_one():
    # Whatever else the view holds, the notebook draws the tabs as it did.
    m = _model()
    m.default_excel_view = mo.ExcelView(sheet={"columns": 5})
    assert ">Inputs<" in m._repr_html_()
