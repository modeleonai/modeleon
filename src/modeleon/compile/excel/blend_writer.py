# SPDX-License-Identifier: Apache-2.0
"""The blend row, written as FORMULAS — the splice stays alive in Excel.

A blended line («live») is not data: it is ``given`` up to the close
boundary and ``follow`` after it, per cell, with a fallback to the
other side wherever one is missing. Emitted as numbers it would be
frozen in the written workbook — enter a fact in the file and the
blended line would sit unchanged, which is exactly the dead-values
habit Excel parity exists to kill.

So the blend row references its own track subrows:

* before the boundary — ``=IF(факт="", план, факт)``;
* after it            — ``=IF(план="", факт, план)``;
* one side only       — a bare ``=<that cell>``.

Enter a fact in the workbook and the blended line moves with it,
exactly as the engine would have re-blended.

Only DATA lines take these — the rows ``expand_tracked_tree`` marked
``_blend_is_splice``, using the engine's own predicate (a line with
authored role kwargs, or one with no expression at all). A purely
DERIVED line's blend row already carries its own formula over the
operands' blend rows (a derived line's live series is computed from
its operands' live series, never spliced), and rewriting it would
replay the splice on an output.

A line can be both: ``mo.Variable(<план formula>, факт=[...],
прогноз=[...])`` computes its план track from an expression and takes
the other two as authored series. It is a DATA line — its blend is a
splice of the authored rows — so the base expression is the wrong
formula for the blend row. Left there, the blend row would show план
where the engine computes the splice, and a fact entered in the file
would never reach it.
"""

from typing import Any, Dict, List, Optional

from .addresses import VariableAddresses
from ...core.multi_variable import MultiVariableBase
from ...core.variable import Variable

#: The subrow attribute infix ``expand_tracked_tree`` emits.
TRACK_INFIX = "__track_"


def _before_boundary(sheet_mv: MultiVariableBase, spec: Any,
                     n: int) -> List[bool]:
    """Which periods take ``given`` first — the same splice the engine
    computes, spelled over the sheet's own window."""
    from ...core.time import _parse, _period_labels, resolve_default_window

    if spec.until is None:
        # Boundary-less: given wins wherever it carries a value.
        return [True] * n
    window = resolve_default_window(sheet_mv)
    if window is None or window.start is None or window.grain is None:
        return [True] * n
    labels = _period_labels(window.start, None, n, window.grain)
    try:
        anchor = _parse(spec.until, window.grain)
    except (ValueError, TypeError):
        return [True] * n
    return [_parse(lb, window.grain) <= anchor for lb in labels]


def _cell(addr: Optional[VariableAddresses], i: int) -> Optional[str]:
    if addr is None or not addr.values:
        return None
    if len(addr.values) == 1:
        return addr.values[0]          # a scalar broadcasts
    return addr.values[i] if i < len(addr.values) else None


def write_blend_links(
    ws: Any,
    sheet_mv: MultiVariableBase,
    addresses: Dict[str, VariableAddresses],
    var_to_sheet: Dict[str, str],
    sheet_name: str,
) -> None:
    """Stamp the splice formulas onto this sheet's blend rows."""
    from ...core.tracks_decl import resolve_tracks_decl

    decl = resolve_tracks_decl(sheet_mv)
    spec = getattr(decl, "blend", None) if decl is not None else None
    if spec is None:
        return

    for _, comp in sheet_mv.walk():
        if not isinstance(comp, Variable) or not comp.id:
            continue
        if TRACK_INFIX in comp.id:
            continue                      # a subrow, not the blend row
        if not getattr(comp, "_blend_is_splice", False):
            continue                      # derived: its own formula stands
        if var_to_sheet.get(comp.id) != sheet_name:
            continue
        addr = addresses.get(comp.id)
        if addr is None or not addr.values:
            continue
        given = addresses.get(f"{comp.id}{TRACK_INFIX}{spec.given}")
        follow = addresses.get(f"{comp.id}{TRACK_INFIX}{spec.follow}")
        if given is None and follow is None:
            continue                      # nothing to reference
        splice = _before_boundary(sheet_mv, spec, len(addr.values))
        for i, ref in enumerate(addr.values):
            g = _cell(given, i)
            f = _cell(follow, i)
            first, second = (g, f) if splice[i] else (f, g)
            if first and second:
                ws[ref] = f'=IF({first}="",{second},{first})'
            elif first:
                ws[ref] = f"={first}"
            elif second:
                ws[ref] = f"={second}"
            if g:
                _mark_settled(ws, ref, g)


#: The settled-cell ink — the finance-Excel input blue.
SETTLED_FONT_COLOR = "FF2563EB"


def _mark_settled(ws: Any, ref: str, given_ref: str) -> None:
    """Ink the blended cell blue WHILE a fact stands behind it.

    A conditional rule, not a stamped font: the cell itself is a live
    splice formula, so a baked colour would freeze the provenance the
    moment the file is written — the same dead-values habit this
    module exists to kill. Typed into the file, a fact turns its own
    month blue on the spot.

    Font only, never a fill: a fill would win over the house style's
    section bands and cut holes in them.
    """
    try:
        from openpyxl.formatting.rule import FormulaRule
        from openpyxl.styles import Font
    except ImportError:      # pragma: no cover - openpyxl is a hard dep
        return
    ws.conditional_formatting.add(
        ref,
        FormulaRule(formula=[f'{given_ref}<>""'],
                    font=Font(color=SETTLED_FONT_COLOR)),
    )
