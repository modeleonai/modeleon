# SPDX-License-Identifier: Apache-2.0
"""Cached values: what a workbook shows before anything recalculates.

A formula cell in ``.xlsx`` holds two things, the formula and the value
it last produced (its ``<v>``). openpyxl writes the formula with an empty
value, so a reader that does not recalculate (a file preview, a library
reading values) shows the cell empty, as 0, or as a date in 1899. The
engine has already computed every value those formulas produce: the
writer records it beside each cell it fills (:func:`remember`), and
:func:`workbook_bytes` puts it into the file. Excel recalculates on open
(the book keeps ``fullCalcOnLoad``); LibreOffice, by default, shows the
cached values, as every preview and library does.

Only values the engine computed are cached: a row's own cells, the cells
of a header row that repeats another row, an opening cell, a caption that
only points at one of those (:func:`remember_reference`), and a row total,
which carries ``row_sum_value`` of the row (the number its ``=SUM``
gives). Any other formula cell,
and one whose value has no spreadsheet form, is written with no value at
all, so a preview shows it blank rather than as a made-up 0.

Caching never breaks the workbook: a value that cannot be written is left
out.
"""

from __future__ import annotations

import math
import numbers
import re
import xml.etree.ElementTree as ET
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from typing import Any, NamedTuple, Optional
from weakref import WeakKeyDictionary
from xml.sax.saxutils import escape
from zipfile import ZipFile

import numpy as np
from openpyxl import Workbook
from openpyxl.utils.cell import column_index_from_string, get_column_letter
from openpyxl.utils.datetime import time_to_days, to_excel
from openpyxl.worksheet.worksheet import Worksheet

# The computed value of each filled cell, per sheet, until the book is
# saved. Weak: a sheet that is never saved takes its values with it.
_VALUES: "WeakKeyDictionary[Worksheet, dict[str, Any]]" = WeakKeyDictionary()

_ERRORS = frozenset({
    "#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A",
})

# A formula cell as openpyxl writes it, with an empty value: ``<v></v>``
# when lxml writes the sheet, ``<v />`` when the standard library does.
# Only this shape is rewritten; any other cell passes through untouched.
_FORMULA_CELL = re.compile(
    r'<c r="([A-Z]{1,3}[0-9]+)"([^>]*)><f([^>]*)>([^<]*)</f>'
    r'(?:<v></v>|<v\s*/>)</c>'
)
_TYPE_ATTR = re.compile(r'\st="[^"]*"')
# A reference as the renderer names it: ``H4``, ``'Any name'!H4``, or a
# bare name (letters of any script, digits, underscore) and ``!H4``.
_REFERENCE = re.compile(
    r"^(?:'((?:[^']|'')+)'!|([A-Za-z_]\w*)!)?\$?([A-Z]{1,3})\$?([0-9]+)$"
)
# Characters XML 1.0 cannot carry: a value holding one is not cached.
_NOT_XML = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]")
# A cell holds at most this many characters.
_MAX_TEXT = 32_767
# Beyond this an integer is no longer a spreadsheet number.
_MAX_NUMBER = 1.7976931348623157e308


class _Points(NamedTuple):
    sheet: Optional[str]
    coordinate: str


def remember(ws: Worksheet, coordinate: str, value: Any) -> None:
    """Record what the engine computed for the cell at ``coordinate``.

    Used only if the cell still holds a formula when the book is saved:
    a cell written as a value already carries it."""
    _VALUES.setdefault(ws, {})[coordinate] = value


def remember_reference(ws: Worksheet, coordinate: str, reference: str) -> None:
    """Record that the cell at ``coordinate`` only shows another cell,
    named as its formula names it (``H4``, ``'Sheet'!H4``): once the book
    is saved it carries that cell's value, if that one is known."""
    m = _REFERENCE.match(reference)
    if m is None:
        return
    quoted, bare, col, row = m.groups()
    sheet = quoted.replace("''", "'") if quoted is not None else bare
    _VALUES.setdefault(ws, {})[coordinate] = _Points(sheet, f"{col}{row}")


def workbook_bytes(wb: Workbook) -> bytes:
    """The saved workbook, every remembered formula cell carrying its value."""
    buf = BytesIO()
    wb.save(buf)
    by_title = {ws.title: _VALUES.get(ws) or {} for ws in wb.worksheets}
    resolved = {
        title: {ref: _follow(by_title, title, value) for ref, value in values.items()}
        for title, values in by_title.items()
    }
    return _with_values(buf.getvalue(), resolved, wb.epoch)


def save_workbook(wb: Workbook, path: str | Path) -> None:
    """:func:`workbook_bytes`, written to ``path``."""
    Path(path).write_bytes(workbook_bytes(wb))


def _follow(by_title: dict[str, dict[str, Any]], title: str, value: Any) -> Any:
    # A pointing cell's value: the cell it shows, along a short chain of
    # such cells; None when the end of the chain is not known.
    for _ in range(8):
        if not isinstance(value, _Points):
            return value
        title = value.sheet or title
        value = by_title.get(title, {}).get(value.coordinate)
    return None


def _with_values(data: bytes, by_title: dict[str, dict[str, Any]], epoch: datetime) -> bytes:
    src = ZipFile(BytesIO(data))
    rewritten: dict[str, bytes] = {}
    for title, part in _sheet_parts(src).items():
        values = by_title.get(title, {})
        xml = src.read(part).decode("utf-8")
        xml = _FORMULA_CELL.sub(lambda m: _with_value(m, values, epoch), xml)
        rewritten[part] = _with_ignored_errors(xml).encode("utf-8")
    out = BytesIO()
    with ZipFile(out, "w") as dst:
        for info in src.infolist():
            dst.writestr(info, rewritten.get(info.filename) or src.read(info.filename))
    return out.getvalue()


def _sheet_parts(book: ZipFile) -> dict[str, str]:
    # {sheet title: its part in the package}, read from the workbook's own
    # index rather than assumed from the order sheets were written in.
    rels = ET.fromstring(book.read("xl/_rels/workbook.xml.rels"))
    target = {
        r.get("Id"): r.get("Target", "")
        for r in rels.iter("{http://schemas.openxmlformats.org/package/2006/relationships}"
                           "Relationship")
    }
    parts: dict[str, str] = {}
    workbook = ET.fromstring(book.read("xl/workbook.xml"))
    rel_id = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    for sheet in workbook.iter("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheet"):
        name, path = sheet.get("name"), target.get(sheet.get(rel_id), "")
        if name and path:
            parts[name] = path.lstrip("/") if path.startswith("/") else f"xl/{path}"
    return parts


# Excel checks a sheet as it opens it and marks what looks like a mistake
# with a green triangle in the cell's corner, which a reader takes for a
# note that is not there. Two kinds are the writer's own, by design: text
# that reads as a number (a line's number "1.1", a year label "2026"), and
# the formulas it wrote, which differ from their neighbours on purpose (a
# quarter's total between its months, a deviation between two tracks, a
# range that starts right after the opening column). The sheet records
# them as checked, the way Excel records a reader's own "Ignore error". A
# formula that computes an error keeps its mark, and so does whatever a
# reader later types into a cell that held no formula.
_TEXT_CELL = re.compile(
    r'<c r="([A-Z]{1,3})([0-9]+)"[^>]*?\st="inlineStr"[^>]*>'
    r'<is><t(?:\s[^>]*)?>([^<]*)</t></is></c>'
)
# Text Excel may read as a number or a date in some locale: digits with
# separators, a sign, a percent, brackets ("1", "1.1", "1.2.3", "2026").
_NUMBER_LIKE = re.compile(r"^\s*[-+(]?\s*\d[\d\s.,/'%()+-]*$")
_FORMULA_AT = re.compile(r'<c r="([A-Z]{1,3})([0-9]+)"[^>]*><f[\s>/]')
# What a worksheet puts after <ignoredErrors>; it goes before the first.
_AFTER_IGNORED = re.compile(
    r"<(?:smartTags|drawing|legacyDrawing|legacyDrawingHF|drawingHF|picture"
    r"|oleObjects|controls|webPublishItems|tableParts|extLst)[\s/>]"
)
# One entry names at most this many ranges.
_MAX_RANGES = 8000


def _with_ignored_errors(xml: str) -> str:
    end = xml.find("</sheetData>")
    if end < 0 or "<ignoredErrors" in xml:
        return xml
    texts = [
        (col, int(row)) for col, row, text in _TEXT_CELL.findall(xml, 0, end)
        if _NUMBER_LIKE.match(text)
    ]
    formulas = [(col, int(row)) for col, row in _FORMULA_AT.findall(xml, 0, end)]
    # Two sets that never meet, since a cell holds text or a formula:
    # Excel keeps one entry per cell, and overlapping entries lose rules.
    entries = ""
    for cells, rules in (
        (texts, 'numberStoredAsText="1" twoDigitTextYear="1"'),
        (formulas, 'formula="1" formulaRange="1"'),
    ):
        ranges = _ranges(cells)
        if 0 < len(ranges) <= _MAX_RANGES:
            entries += f'<ignoredError sqref="{" ".join(ranges)}" {rules}/>'
    if not entries:
        return xml
    after = _AFTER_IGNORED.search(xml, end)
    at = after.start() if after is not None else xml.rindex("</worksheet>")
    return f"{xml[:at]}<ignoredErrors>{entries}</ignoredErrors>{xml[at:]}"


def _ranges(cells: list[tuple[str, int]]) -> list[str]:
    """``["A3:A10", "C1:E1"]`` for a set of cells: the runs of rows down each
    column, and neighbouring columns with the same run joined into one."""
    rows_by_col: dict[int, set[int]] = {}
    for col, row in cells:
        rows_by_col.setdefault(column_index_from_string(col), set()).add(row)
    cols_by_run: dict[tuple[int, int], list[int]] = {}
    for col, rows in rows_by_col.items():
        for run in _runs(sorted(rows)):
            cols_by_run.setdefault(run, []).append(col)
    out: list[str] = []
    for (top, bottom), cols in sorted(cols_by_run.items()):
        for first, last in _runs(sorted(cols)):
            a, b = f"{get_column_letter(first)}{top}", f"{get_column_letter(last)}{bottom}"
            out.append(a if a == b else f"{a}:{b}")
    return out


def _runs(numbers: list[int]) -> list[tuple[int, int]]:
    """Sorted ``numbers`` as ``(first, last)`` runs of consecutive ones."""
    runs: list[tuple[int, int]] = []
    for n in numbers:
        if runs and n == runs[-1][1] + 1:
            runs[-1] = (runs[-1][0], n)
        else:
            runs.append((n, n))
    return runs


def _with_value(m: "re.Match[str]", values: dict[str, Any], epoch: datetime) -> str:
    ref, attrs, f_attrs, formula = m.group(1), m.group(2), m.group(3), m.group(4)
    if "dataTable" in f_attrs:
        return m.group(0)
    cached = _cached(values.get(ref), epoch)
    if cached is None:
        # No value to show: none at all, rather than an empty one a
        # preview reads as 0.
        return f'<c r="{ref}"{attrs}><f{f_attrs}>{formula}</f></c>'
    kind, text = cached
    attrs = _TYPE_ATTR.sub("", attrs) + (f' t="{kind}"' if kind else "")
    return f'<c r="{ref}"{attrs}><f{f_attrs}>{formula}</f><v>{text}</v></c>'


def _cached(value: Any, epoch: datetime) -> Optional[tuple[str, str]]:
    """``(cell type, value text)`` for one computed value, or ``None`` when
    it has no spreadsheet form (an empty period, NaN, an object, a value
    beyond what a cell holds). Never raises: a value that cannot be cached is left out."""
    try:
        return _spreadsheet_form(value, epoch)
    except (ArithmeticError, TypeError, ValueError):
        return None


def _spreadsheet_form(value: Any, epoch: datetime) -> Optional[tuple[str, str]]:
    if value is None or isinstance(value, (np.datetime64, np.timedelta64)):
        return None
    if isinstance(value, (bool, np.bool_)):
        return "b", "1" if value else "0"
    if isinstance(value, numbers.Integral):
        n = int(value)
        return ("", str(n)) if abs(n) <= _MAX_NUMBER else None
    if isinstance(value, (numbers.Real, Decimal)):
        return _number(float(value))
    if isinstance(value, str):
        if value in _ERRORS:
            return "e", value
        if len(value) > _MAX_TEXT or _NOT_XML.search(value):
            return None
        return "str", escape(value, {"\r": "&#13;"})
    if isinstance(value, (datetime, date)):
        if isinstance(value, datetime):
            value = value.replace(tzinfo=None)
        return _number(float(to_excel(value, epoch)))
    if isinstance(value, time):
        return _number(float(time_to_days(value)))
    if isinstance(value, timedelta):
        return _number(value / timedelta(days=1))
    return None


def _number(x: float) -> Optional[tuple[str, str]]:
    # Fifteen significant digits: all a spreadsheet keeps, and far shorter
    # in the file than a float's full repr.
    return ("", format(x, ".15g")) if math.isfinite(x) else None
