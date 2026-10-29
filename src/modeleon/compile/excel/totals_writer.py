# SPDX-License-Identifier: Apache-2.0
"""Writing the subtotal cells — the CONTENT half of ``totals``.

Geometry (which column each bucket owns) is the layout's
:class:`~.totals.TotalsPlan`; this module fills the cells so they are
ALIVE:

* rule lines — ``=SUM(jan,feb,mar)`` / ``=AVERAGE(...)`` / the last
  month's ref, straight from the line's regrain rule;
* formula lines — the line's own formula rendered AT the bucket, via a
  translator whose address book maps every periodic line to its bucket
  cells (:func:`project_model` supplies the bucket-grain expression
  tree, rebuilt under the sheet's own component names);
* lines the engine computes month by month and sums into the bucket (an
  IF, a ROUND, a product of two flows) — ``=SUM`` of the line's
  own month cells;
* ruleless literals — the projection's ``#VALUE!`` teaching token.

Edit January in the written file and the quarter and year move —
the parity promise, kept in the totals too.
"""

import logging
from numbers import Real
from typing import Any, Dict

from .addresses import VariableAddresses
from .cached_values import remember
from .layout import line_keys
from .translator import ExcelTranslator
from .totals import NATIVE, TotalsPlan, listed_or_ranged, regrain_rule_of, rule_formula
from .writer import resolve_number_format
from ...core.multi_variable import MultiVariableBase
from ...core.variable import Variable

logger = logging.getLogger(__name__)


def _walk_vars(mv: MultiVariableBase) -> Dict[str, Variable]:
    out: Dict[str, Variable] = {}
    for _, comp in mv.walk():
        if isinstance(comp, Variable) and comp.id:
            out[comp.id] = comp
    return out


def _vars_by_place(mv: MultiVariableBase, prefix: str = "") -> Dict[str, Variable]:
    """Every Variable under ``mv`` keyed by its PLACE — the dotted
    component names from ``mv`` down (``tax``, ``income.revenue``)."""
    out: Dict[str, Variable] = {}
    for name in mv._component_order:
        comp = mv._component(name)
        if isinstance(comp, Variable):
            out.setdefault(f"{prefix}{name}", comp)
        elif isinstance(comp, MultiVariableBase):
            out.update(_vars_by_place(comp, f"{prefix}{name}."))
    return out


def bucket_cells(
    sheet_mv: MultiVariableBase,
    plan: TotalsPlan,
    addresses: Dict[str, VariableAddresses],
    var_to_sheet: Dict[str, str],
    sheet_name: str,
    projected: Dict[str, Dict[str, Variable]] | None = None,
    mirrors: Any = None,
    plans: Any = None,
) -> Dict[str, Dict[str, Any]]:
    """Every bucket cell's CONTENT — ``{ref: {"formula": str|None,
    "value": Any, "id": vid}}`` — the single computation behind the
    .xlsx writer and any other renderer of the totals, so no two
    renderings of a bucket cell can disagree.

    ``mirrors`` — the header rows repeated at the top of this sheet, as
    ``(source, mirror)`` pairs (the source a Variable or its id): each
    bucket of a mirror reads the bucket cell of the row it repeats —
    ``=Timing!F7`` — found through ``plans``, the plans of every sheet
    by name (its own included); where that row has no such cell the
    bucket is the row's projected value. Their cells carry
    ``"mirror": True`` and the SOURCE's id.

    ``projected`` — pre-computed ``{kind: {place: projected Var}}``
    from a caller that already ran :func:`project_model` (e.g. one that
    memoizes one projection per kind across sheets), keyed by each
    line's dotted place below ``sheet_mv`` — which is also its projected
    id without the projected root's prefix; a kind it lacks (``None``:
    every kind) is projected here from the sheet alone, which is fine
    for the writer's one-shot export."""
    from ...core.projection import project_model

    native = _walk_vars(sheet_mv)
    # A line meets its projected self by its PLACE in the sheet — the
    # dotted names from the sheet down — because the projection rebuilds
    # the sheet under those very names. Not by id: a tab the writer
    # makes on its own (from a container not marked as a tab, from the
    # Variables placed straight on the model, or for a container written
    # with its own ``to_excel()``) is a stand-in whose id is unrelated to
    # the ids its lines keep, so no id prefix would ever line them up.
    place_of: Dict[str, str] = {}
    for place, var in _vars_by_place(sheet_mv).items():
        if var.id:
            place_of.setdefault(var.id, place)

    projected = dict(projected or {})
    for kind in plan.kinds:
        if kind in projected:
            continue
        try:
            projected[kind] = _vars_by_place(
                project_model(sheet_mv, kind, on_error="absorb")
            )
        except Exception as exc:  # projection must never kill the book
            logger.warning("totals: %s projection failed: %s", kind, exc)
            projected[kind] = {}

    # One translator per kind: every periodic line's "periods" are its
    # BUCKET cells, everything else (constants, cross-sheet operands)
    # keeps its native address — which is exactly how a hand-built
    # book's quarter formula reads. A line read as a range (``MAX(n)``,
    # ``COUNTIF(n, ">50")``) is its months, as the engine's total reads
    # it: a reduction over the line is one number at every grain.
    translators: Dict[str, ExcelTranslator] = {}
    bucket_refs: Dict[str, Dict[str, list]] = {}
    books: Dict[str, Any] = {}
    for kind in plan.kinds:
        merged = dict(addresses)
        sheets = dict(var_to_sheet)
        months: Dict[str, VariableAddresses] = {}
        books[kind] = (merged, sheets, months)
        refs_of: Dict[str, list] = {}
        for vid, addr in addresses.items():
            if var_to_sheet.get(vid) != sheet_name:
                continue
            if vid not in native:
                continue
            pvar = projected[kind].get(place_of.get(vid, ""))
            if getattr(native[vid], '_metric_no_bucket', False):
                # A growth-% metric has no prior-year bucket — the
                # file shows empty there, never a projected lag whose
                # step no longer fits the grain.
                continue
            if _is_periodic(native[vid]):
                row = _row_of(addr.values[0])
                # The line's SLOT inside each group — a track column
                # totals into its own bucket subcolumn, exactly under
                # its months.
                _slot = getattr(native[vid], '_col_slot', 0) or 0
                refs = [
                    f"{_col_letter(plan.cell((kind, bi), _slot))}{row}"
                    for bi in range(len(plan.buckets[kind]))
                ]
                refs_of[vid] = refs
                slot = VariableAddresses(
                    name="", formula=refs[0], values=refs,
                )
                months[vid] = addr
            else:
                # A constant reads the same in every bucket — its
                # native cell IS its quarter and its year.
                slot = addr
            merged[vid] = slot
            if pvar is not None and pvar.id:
                # The projected expr's refs point at PROJECTED objects
                # (``plan.revenue``), the address book at native ids
                # (``m.plan.revenue``) — alias, or every operand falls
                # back to an inlined literal and the bucket goes dead.
                merged[pvar.id] = slot
                sheets[pvar.id] = sheet_name
        translators[kind] = ExcelTranslator(merged, sheets, periods=months)
        bucket_refs[kind] = refs_of

    out: Dict[str, Dict[str, Any]] = {}
    position_memo: Dict[Any, Any] = {}
    for vid, refs in _stable(bucket_refs, plan):
        kind = refs["kind"]
        var = native[vid]
        row_refs = bucket_refs[kind][vid]
        month_row = _row_of(addresses[vid].values[0])
        rule = regrain_rule_of(var)
        pvar = projected.get(kind, {}).get(place_of.get(vid, ""))
        # A row of a tracked formula line folds by its line's formula in
        # its own track, over its operands' bucket cells on the same track.
        # A track authored with its own formula (``plan=<expr>``) on a line
        # with no rule: the engine refuses the line at a coarser grain.
        authored = rule is None and (getattr(var, "_track_role", None) is not None
                                     or getattr(var, "_expr_is_track_specific", False))
        tracked = (None if rule is not None or authored else
                   _tracked_fold(var, native, kind, *books[kind], bucket_refs[kind],
                                 position_memo))
        if (rule is None and not authored
                and _sums_its_months(var, pvar, tracked, plan.buckets[kind], sheet_name)):
            # Computed month by month and summed into the bucket (an IF, a
            # ROUND, a product of two flows): its bucket is the SUM of its
            # own month cells, live as a rule line's.
            rule, tracked = "sum", None
        elif rule is None and not authored and (picked := _picked_cell(var, pvar)):
            # A time row over hours: the day's first start or last stop is
            # this very row's hour cell, live as a rule line's.
            rule, tracked = picked, None
        # A compare column (Var, Var %, Share) of a line whose track rows
        # have no fold of their own: its formula over those rows' bucket
        # cells, all on this sheet.
        metric = (rule is None and "__metric_" in vid
                  and getattr(var, "_expr", None) is not None
                  and all(getattr(r, "id", None) in bucket_refs[kind]
                          for r in var._expr.iter_refs()))
        own_time = pvar is not None and _reads_its_own_time(var, pvar)
        for bi, (_label, lo, hi) in enumerate(plan.buckets[kind]):
            formula: Any = None
            value: Any = None
            inlined = False
            if rule is not None:
                _slot = getattr(var, '_col_slot', 0) or 0
                month_refs = [
                    f"{_col_letter(plan.cell((NATIVE, i), _slot))}{month_row}"
                    for i in range(lo, min(hi, plan.n))
                    if i < len(addresses[vid].values)
                ]
                end = lo + len(month_refs)
                if month_refs and isinstance(rule, str):
                    # Live over the bucket's months, an un-entered one
                    # blank and so zero - as the engine folds a hole -
                    # but in a line of dates or words, where a blank has
                    # no zero and the engine's bucket stays blank.
                    weights = _bucket_weights(var, lo, end) if rule == "mean" else None
                    formula = ('=""' if _stays_blank(var, rule, lo, end, weights)
                               else rule_formula(rule, month_refs, weights,
                                                 blanks=_has_blank(var, lo, end)))
                elif month_refs:
                    formula = _ratio_formula(var, rule, lo, end,
                                             addresses, var_to_sheet, sheet_name)
                    if formula is None and getattr(var, "_track_copy", False) \
                            and "__track_" in vid:
                        # A track's row the named lines do not carry: the
                        # engine refuses it, and a ratio of the display
                        # rows would be another track's number.
                        value = "#VALUE!"
                if formula is None and value is None and pvar is not None:
                    value = _bucket_value(pvar, bi, var)
            elif authored:
                value = "#VALUE!"
            elif tracked is not None:
                formula, value, inlined = _tracked_bucket(tracked, bi, sheet_name,
                                                          books[kind][0].get(vid), pvar, var)
            elif metric and getattr(pvar, "_expr", None) is None:
                try:
                    formula = translators[kind].translate(
                        var._expr, bi, current_sheet=sheet_name,
                        self_address=books[kind][0].get(vid), self_var=var)
                    inlined = bool(getattr(translators[kind], "last_render_inlined", False))
                except Exception:  # noqa: BLE001 — degrade to the number
                    value = _bucket_value(pvar, bi, var) if pvar is not None else None
                if formula is not None and translators[kind].last_render_lossy:
                    # No formula stands for it here (a reader of one range
                    # over months that are not one run): the number.
                    formula = None
                    value = _bucket_value(pvar, bi, var) if pvar is not None else None
            elif pvar is not None and getattr(pvar, "_expr", None) is not None:
                try:
                    formula = translators[kind].translate(
                        pvar._expr,
                        bi,
                        current_sheet=sheet_name,
                        # The row's bucket cells: a running total reads
                        # its own previous bucket.
                        self_address=books[kind][0].get(vid),
                        self_var=pvar,
                    )
                    inlined = bool(getattr(translators[kind], "last_render_inlined", False))
                except Exception:
                    value = _bucket_value(pvar, bi, var)
                if formula is not None and translators[kind].last_render_lossy:
                    formula, value = None, _bucket_value(pvar, bi, var)
                if isinstance(formula, str) and "TrackValues(" in formula:
                    # An operand the address book cannot place inlined a
                    # whole tracked value: the number, never that text.
                    formula, value = None, _bucket_value(pvar, bi, var)
            elif pvar is not None:
                # Ruleless literal → the projection already answered:
                # a value where it could, its teaching token where it
                # could not. Never an invented sum.
                value = _bucket_value(pvar, bi, var)
            out[row_refs[bi]] = {
                "formula": formula, "value": value, "id": vid,
                # True when a formula's text carries a computed value baked
                # in: that formula is correct for the current inputs only.
                "inlined": bool(inlined),
                # What the engine computed for a time row's bucket on a
                # clock sheet, kept beside its formula for a reader that does not recalculate.
                "cached": (_bucket_value(pvar, bi, var)
                           if own_time and pvar is not None and formula is not None else None),
            }
    _mirror_bucket_cells(mirrors, plan, plans, addresses, var_to_sheet, sheet_name, out)
    return out


def _qualified(ref: str, sheet: Any, here: str) -> str:
    """``ref`` as a formula on ``here`` names it: bare on its own sheet,
    else with the sheet, quoted when its name needs it."""
    from .renderer import _UNQUOTED_SHEET_NAME

    if not sheet or sheet == here:
        return ref
    if _UNQUOTED_SHEET_NAME.match(sheet):
        return f"{sheet}!{ref}"
    return "'" + sheet.replace("'", "''") + f"'!{ref}"


def _source_of(src: Any, mirror: Any) -> tuple:
    """``(id, Variable)`` of the row a header mirror repeats — the pair
    names it by the object or by its id; the mirror knows the object."""
    src_id = src if isinstance(src, str) else getattr(src, "id", None)
    source = getattr(mirror, "_mirror_of", None)
    if source is None and not isinstance(src, str):
        source = src
    return src_id, source


def _mirror_bucket_cells(
    mirrors: Any,
    plan: TotalsPlan,
    plans: Any,
    addresses: Dict[str, VariableAddresses],
    var_to_sheet: Dict[str, str],
    sheet_name: str,
    out: Dict[str, Dict[str, Any]],
) -> None:
    """Into ``out``: the bucket cells of the header rows repeated at the
    top of the sheet (``default_header``). A mirror is no row of the
    model — the layout makes it — so the lines' pass never meets it, and
    its buckets stood empty while its periods read the row they repeat.
    Each reads the bucket cell of its row (``=Timing!F7``, ``=F7`` on the
    anchor's own sheet); a row with no such cell (its sheet draws no such
    bucket) gives its projected value instead."""
    from ...core.projection import project_line

    for src, mirror in mirrors or ():
        src_id, source = _source_of(src, mirror)
        maddr = addresses.get(getattr(mirror, "id", "") or "")
        saddr = addresses.get(src_id or "")
        if (src_id is None or source is None or maddr is None or not maddr.values
                or saddr is None or not saddr.values or not _is_periodic(source)):
            continue
        src_sheet = var_to_sheet.get(src_id)
        src_plan = (plans or {}).get(src_sheet) if plans is not None else (
            plan if src_sheet == sheet_name else None)
        mrow, srow = _row_of(maddr.values[0]), _row_of(saddr.values[0])
        mslot = getattr(mirror, "_col_slot", 0) or 0
        sslot = getattr(source, "_col_slot", 0) or 0
        projected: Dict[str, Any] = {}
        for kind in plan.kinds:
            for bi in range(len(plan.buckets[kind])):
                ref = f"{_col_letter(plan.cell((kind, bi), mslot))}{mrow}"
                if ref in out:
                    continue
                formula = value = target = None
                if (src_plan is not None and (kind, bi) in src_plan.col_of
                        and bi < len(src_plan.buckets.get(kind, ()))):
                    target = f"{_col_letter(src_plan.cell((kind, bi), sslot))}{srow}"
                    formula = "=" + _qualified(target, src_sheet, sheet_name)
                else:
                    if kind not in projected:
                        try:
                            projected[kind] = project_line(source, kind, {}, on_error="absorb")
                        except Exception as exc:  # noqa: BLE001 — the book never dies of a header
                            logger.warning("totals: a header row's %s projection failed: %s",
                                           kind, exc)
                            projected[kind] = None
                    pv = projected[kind]
                    value = _bucket_value(pv, bi, source) if pv is not None else None
                out[ref] = {
                    "formula": formula, "value": value, "id": src_id, "inlined": False,
                    "cached": None, "mirror": True,
                    # The cell the formula shows: its value once known.
                    "reference": None if formula is None else formula[1:],
                }


def write_totals(
    ws: Any,
    sheet_mv: MultiVariableBase,
    plan: TotalsPlan,
    addresses: Dict[str, VariableAddresses],
    var_to_sheet: Dict[str, str],
    sheet_name: str,
    mirrors: Any = None,
    plans: Any = None,
    mirror_look: Dict[str, Any] | None = None,
) -> None:
    """Write the sheet's bucket cells — the lines' own and, with
    ``mirrors`` and ``plans`` (see :func:`bucket_cells`), the header
    rows' — bold, in each row's number format; a header row's in
    ``mirror_look`` too, as its period cells are."""
    from openpyxl.styles import Font

    from .cached_values import remember_reference
    from .writer import _apply_styling

    native = _walk_vars(sheet_mv)
    sources = dict(_source_of(src, mirror) for src, mirror in mirrors or ())
    cells = bucket_cells(sheet_mv, plan, addresses, var_to_sheet, sheet_name,
                         mirrors=mirrors, plans=plans)
    for ref, spec in cells.items():
        cell = ws[ref]
        if spec["formula"] is not None:
            cell.value = spec["formula"]
            if spec.get("cached") is not None:
                remember(ws, ref, spec["cached"])
            elif spec.get("reference"):
                remember_reference(ws, ref, spec["reference"])
        elif spec["value"] is not None:
            cell.value = spec["value"]
        cell.font = Font(bold=True)
        if spec.get("mirror") and mirror_look:
            _apply_styling(cell, {**mirror_look, "bold": True})
        var = native.get(spec["id"])
        if var is None and spec.get("mirror"):
            var = sources.get(spec["id"])
        fmt = resolve_number_format(var) if var is not None else None
        if fmt:
            cell.number_format = fmt


def _bucket_weights(var: Variable, lo: int, hi: int) -> Any:
    """The day counts of the line's native periods ``lo``..``hi`` — what
    its mean folds by, as the engine's re-grain does."""
    from ...core.time import period_weights

    loc = getattr(var, "time", None)
    if loc is None or loc.start is None or loc.grain is None:
        return None
    try:
        return period_weights(loc.grain, loc.start, hi, loc.first_day)[lo:hi]
    except (ValueError, TypeError):
        return None


def _shown_series(var: Variable) -> Any:
    """The series a line's row shows: a tracked line's display track (the
    blend when the model blends, else its first), else its value."""
    value = getattr(var, "_value", None)
    if type(value).__name__ != "TrackValues":
        return value
    from ...core.tracks_decl import resolve_tracks_decl

    blend = getattr(getattr(resolve_tracks_decl(var), "blend", None), "name", None)
    roles = list(value.roles)
    return value[blend if blend in roles else roles[0]]


def _reads_its_own_time(var: Variable, pvar: Any) -> bool:
    """A time row whose projection reads its own cells — a clock sheet's
    start, end or days (``core.projection._project_time_row``)."""
    from ...core.expr import Regrain
    from ...core.time import time_field_of

    expr = getattr(pvar, "_expr", None)
    return (isinstance(expr, Regrain) and time_field_of(var) is not None
            and getattr(expr.source, "id", None) == getattr(var, "id", None))


def _picked_cell(var: Variable, pvar: Any) -> str | None:
    """``'first'`` or ``'last'`` when the projection reads each bucket off
    this very row's own cells, as a clock sheet's time row does (a day's
    first start, its last stop: ``core.projection._project_time_row``)."""
    from ...core.expr import Regrain

    expr = getattr(pvar, "_expr", None)
    if (isinstance(expr, Regrain) and expr.recipe in ("first", "last")
            and getattr(expr.source, "id", None) == getattr(var, "id", None)):
        return str(expr.recipe)
    return None


def _sums_its_months(var: Variable, pvar: Any, tracked: Any, buckets: Any,
                     sheet_name: str) -> bool:
    """The engine's bucket is the sum of the months this row shows, held
    as numbers, not as a formula: the projection summed a series equal to
    the row's months (a tracked line: it holds the row's track as numbers),
    and in every bucket its number is that sum. Anything else - a bucket
    the row does not reach, buckets on another window - keeps the
    engine's number."""
    from ...core.expr import Regrain
    from ...core.regrain import apply_recipe, hole_value

    months = _shown_series(var)
    if not isinstance(months, list) or not all(
            v is None or (isinstance(v, Real) and not isinstance(v, bool)) for v in months):
        return False        # Excel's SUM skips a TRUE that the engine counts as 1
    if tracked is None:
        expr = getattr(pvar, "_expr", None)
        values = getattr(pvar, "_value", None)
        if not (isinstance(expr, Regrain) and expr.recipe == "sum"
                and getattr(expr.source, "_value", None) == months
                and isinstance(values, list)):
            return False

        def engine(bi: int) -> Any:
            return values[bi] if bi < len(values) else None
    else:
        if getattr(tracked[0], "_expr", None) is not None:
            return False                # its formula renders at the bucket

        def engine(bi: int) -> Any:
            formula, value, _inlined = _tracked_bucket(tracked, bi, sheet_name, None, pvar, var)
            return value if formula is None else _NOT_A_NUMBER
    blank = hole_value(months)
    for bi, (_label, lo, hi) in enumerate(buckets):
        bucket = months[lo:hi]
        if not bucket:
            return False
        try:
            if engine(bi) != apply_recipe(bucket, "sum", blank=blank):
                return False
        except (TypeError, ValueError):
            return False
    return True


#: Never equal to a number: a bucket the engine writes as a formula.
_NOT_A_NUMBER = object()


def _has_blank(var: Variable, lo: int, hi: int) -> bool:
    """A period of the bucket is not entered (or the line is no series)."""
    value = _shown_series(var)
    return not isinstance(value, list) or any(v is None for v in value[lo:hi])


def _stays_blank(var: Variable, rule: str, lo: int, hi: int, weights: Any) -> bool:
    """The engine's bucket is blank: a hole in a line of dates or words."""
    value = _shown_series(var)
    bucket = value[lo:hi] if isinstance(value, list) else None
    if not bucket or not any(v is None for v in bucket):
        return False
    from ...core.regrain import apply_recipe, hole_value

    blank = (getattr(var, "_hole_value", None) if hasattr(var, "_hole_value")
             else hole_value(getattr(var, "_value", None)))
    try:
        return apply_recipe(bucket, rule, weights, blank=blank) is None
    except (TypeError, ValueError):
        return False


def _operands(expr: Any) -> list:
    """Every line ``expr`` reads, through the unnamed steps between."""
    out: list = []
    seen: set = set()
    todo = [expr]
    while todo:
        for ref in todo.pop().iter_refs():
            if id(ref) in seen:
                continue
            seen.add(id(ref))
            out.append(ref)
            if getattr(ref, "_owner", None) is None and getattr(ref, "_expr", None) is not None:
                todo.append(ref._expr)
    return out


def _tracked_fold(var: Variable, native: Dict[str, Variable], kind: str,
                  merged: Dict[str, Any], sheets: Dict[str, Any],
                  months: Dict[str, VariableAddresses],
                  periodic: Dict[str, list], position_memo: Dict[Any, Any]) -> Any:
    """How a row of a tracked formula line folds into a bucket, or None.

    The line's display row and each track's row (``x__track_actual``) carry
    the line's formula - the track's row borrows its line's - and read
    their operands' rows on the same track. The formula is projected to
    the bucket with its own cache, so every projected operand is known by
    the line it came from and maps to that line's bucket cells: the
    display row's, and each track's row's; what the position law found is
    shared across the lines (``position_memo``). A line read as a range is
    its ``months``. Returns ``(projected, track, translator)``."""
    from .blend_writer import TRACK_INFIX
    from ...core.projection import position_cache, project_variable

    vid = getattr(var, "id", "") or ""
    track = getattr(var, "_track_role", None)
    source = var if getattr(var, "_expr", None) is not None else None
    if TRACK_INFIX in vid:
        head_id, row_track = vid.split(TRACK_INFIX, 1)
        track = track or row_track
        if source is None:
            head = native.get(head_id)
            if (head is None or getattr(head, "_expr", None) is None
                    or getattr(head, "_expr_is_track_specific", False)):
                return None
            source = head
    if source is None:
        return None
    lines = _operands(source._expr)
    if not any(type(getattr(v, "_value", None)).__name__ == "TrackValues" for v in lines):
        return None
    cache = position_cache(position_memo)
    try:
        projected = project_variable(source, kind, cache)
    except Exception:  # noqa: BLE001 — a projection must never kill the book
        return None
    if not isinstance(projected, Variable):
        return None
    book, book_sheets = dict(merged), dict(sheets)
    for line in lines:
        mirror = cache.get(id(line))
        if not isinstance(mirror, Variable) or not mirror.id or line.id not in merged:
            continue
        if line.id in periodic or not _is_periodic(line):
            # Its bucket cells on this sheet, or a constant's own cell.
            book[mirror.id] = merged[line.id]
            book_sheets[mirror.id] = sheets.get(line.id)
        prefix = f"{line.id}{TRACK_INFIX}"
        for key in periodic:
            if key.startswith(prefix):
                twin = f"{mirror.id}{TRACK_INFIX}{key[len(prefix):]}"
                book[twin] = merged[key]
                book_sheets[twin] = sheets.get(key)
    return projected, track, ExcelTranslator(book, book_sheets, periods=months)


def _tracked_bucket(tracked: Any, bi: int, sheet_name: str, slot: Any,
                    pvar: Any, line: Any = None) -> Any:
    """``(formula, value, inlined)`` of a tracked formula row's bucket
    ``bi``: the line's formula there, or the engine's number for the row's
    track when no live formula can stand; ``inlined`` says the formula's
    text carries a value baked in (it holds for the current inputs only)."""
    projected, track, translator = tracked
    if getattr(projected, "_expr", None) is not None:
        try:
            formula = translator.translate(projected._expr, bi, current_sheet=sheet_name,
                                           self_address=slot, self_var=projected,
                                           track_role=track)
        except Exception:  # noqa: BLE001 — degrade to the number
            formula = None
        if formula and "TrackValues(" not in formula and not translator.last_render_lossy:
            return formula, None, bool(getattr(translator, "last_render_inlined", False))
    value = getattr(projected, "_value", None)
    if type(value).__name__ == "TrackValues":
        if track is None or track not in value.roles:
            return None, (_bucket_value(pvar, bi, line) if pvar is not None else None), False
        value = value[track]
    if isinstance(value, list):
        return None, (value[bi] if bi < len(value) else None), False
    return None, value, False


def _track_suffix(var: Variable) -> str:
    """``'__track_actual'`` for a track's copy of a line in the book, else ''."""
    name = getattr(var, "_component_name", None) or getattr(var, "_python_name", None) or ""
    return name[name.index("__track_"):] if "__track_" in name else ""


def _ratio_formula(var: Variable, rule: Any, lo: int, hi: int,
                   addresses: Dict[str, VariableAddresses],
                   var_to_sheet: Dict[str, str], sheet_name: str) -> Any:
    """``=SUM(num)/SUM(den)`` over the bucket's native cells of the two
    lines a ``mo.ratio`` rule names — wherever they stand, on the same
    track as the line. ``None`` (the projected value is written) when
    they cannot be found."""
    from ...core.projection import _resolve_ratio_source
    from .renderer import _UNQUOTED_SHEET_NAME

    suffix = _track_suffix(var)
    parts = []
    for name in (rule.num, rule.den):
        try:
            source = _resolve_ratio_source(var, name + suffix)
            if source is None and suffix:
                # A line without tracks has one row, the same on every
                # track: the track's ratio divides by it.
                plain = _resolve_ratio_source(var, name)
                if plain is not None and not getattr(plain, "_track_copy", False) \
                        and type(getattr(plain, "_value", None)).__name__ != "TrackValues":
                    source = plain
        except (ValueError, AttributeError):
            return None
        if source is None:
            return None
        addr = addresses.get(getattr(source, "id", None) or "")
        if addr is None or not addr.values:
            return None
        # A constant (a track holding one value) has one cell, the same in
        # every period: its sum over the bucket is that cell times its
        # periods, as the engine spreads it.
        constant = len(addr.values) == 1 and len(addr.values) < hi
        if not constant and len(addr.values) < hi:
            return None
        sheet = var_to_sheet.get(source.id)

        def qualify(ref: str, sheet: Any = sheet) -> str:
            if not sheet or sheet == sheet_name:
                return ref
            if _UNQUOTED_SHEET_NAME.match(sheet):
                return f"{sheet}!{ref}"
            return "'" + sheet.replace("'", "''") + f"'!{ref}"

        parts.append(f"{qualify(str(addr.values[0]))}*{hi - lo}" if constant else
                     listed_or_ranged("sum", list(addr.values[lo:hi]), qualify=qualify))
    return f"={parts[0]}/({parts[1]})" if "*" in parts[1] else f"={parts[0]}/{parts[1]}"


def _stable(bucket_refs: Dict[str, Dict[str, list]], plan: TotalsPlan):
    for kind in plan.kinds:
        for vid in bucket_refs[kind]:
            yield vid, {"kind": kind}


def _bucket_value(pvar: Variable, bi: int, line: Any = None) -> Any:
    v = getattr(pvar, "_value", None)
    if type(v).__name__ == "TrackValues":
        # A tracked line's projected value is the whole role→series
        # container — written raw it lands in the cell as machine
        # repr. Resolve to the display default (the blend when the
        # model blends, else the first declared track) — the same
        # default the line's own row shows. A line projected on its own
        # (:func:`project_line`) is floating — no owner chain of its own
        # — so its default is the one of ``line``, the line it stands for.
        from ...core.tracks_decl import resolve_tracks_decl

        owner = getattr(pvar, "_owner", None)
        if owner is None and line is not None:
            owner = getattr(line, "_owner", None) or getattr(line, "_parent", None)
        decl = resolve_tracks_decl(owner)
        blend_name = getattr(getattr(decl, "blend", None), "name", None)
        roles = list(v.roles)
        v = v[blend_name if blend_name in roles else roles[0]]
    if isinstance(v, list):
        return v[bi] if bi < len(v) else None
    return v


def _is_periodic(var: Variable) -> bool:
    # A line that runs along the periods, and so folds into buckets. A
    # keyed line (``line_keys``) takes a series' place on the sheet but
    # runs along keys of its own: no months to fold, no bucket cell — the
    # layout keeps it off ``period_var_ids``, and its cells run across the
    # bucket columns, where a bucket written on its row landed on a value.
    value = getattr(var, "_value", None)
    periodic = (isinstance(value, list) or type(value).__name__ == "TrackValues"
                or bool(getattr(var, "_indexed_by", ()) or ()))
    return periodic and line_keys(var) is None


def _row_of(address: str) -> int:
    i = 0
    while i < len(address) and address[i].isalpha():
        i += 1
    return int(address[i:] or 0)


def _col_letter(col: int) -> str:
    result = ""
    while col > 0:
        col -= 1
        result = chr(65 + col % 26) + result
        col //= 26
    return result
